# -*- coding: utf-8 -*-
"""Sync README + Word column order: why is the 5th (last) column."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

CANVAS = Path(
    r"C:\Users\novikova-eyu\.cursor\projects\c-arm-gs\canvases\access-three-buttons-breakdown.canvas.tsx"
)
README = Path(r"c:\arm_gs\app\fuel\services\calculation\README.md")
DOCX = Path(r"c:\arm_gs\app\fuel\services\calculation\access_three_buttons_breakdown.docx")
DOCX_ROOT = Path(r"c:\arm_gs\access_three_buttons_breakdown.docx")


def parse_const_steps(src: str, const_name: str) -> list[dict[str, str]]:
    m = re.search(rf"const {const_name}: Step\[\] = \[(.*?)\n\];", src, re.DOTALL)
    if not m:
        raise SystemExit(f"const {const_name} not found")
    body = m.group(1)
    steps: list[dict[str, str]] = []
    for mobj in re.finditer(r'id:\s*"(?P<id>[^"]+)"', body):
        next_m = re.search(r'\nid:\s*"', body[mobj.end() :])
        end = mobj.end() + next_m.start() if next_m else len(body)
        chunk = body[mobj.start() : end]

        def bt(key: str) -> str:
            mm = re.search(rf"{key}:\s*`(.*?)`", chunk, re.DOTALL)
            return mm.group(1) if mm else ""

        def qs(key: str) -> str:
            mm = re.search(rf'{key}:\s*\n?\s*"((?:\\.|[^"\\])*)"', chunk, re.DOTALL)
            return mm.group(1) if mm else ""

        steps.append(
            {
                "id": mobj.group("id"),
                "access": bt("access"),
                "words": qs("words"),
                "why": qs("why"),
                "py": bt("py"),
                "ui": bt("ui"),
            }
        )
    return steps


def md_cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", "<br>")


def md_code_cell(text: str) -> str:
    return "`" + text.replace("|", "\\|").replace("\n", "<br>") + "`"


def md_access_cell(sid: str, access: str) -> str:
    acc = access.replace("|", "\\|").replace("\n", "<br>")
    return f"**{sid}**<br>`{acc}`"


def build_readme_table(steps: list[dict[str, str]]) -> str:
    lines = [
        "| № / Access (VBA) | Участок в АРМ ГС (Python) | Что сделать словами | Страница / что читается и пишется | Почему / зачем |",
        "|---|---|---|---|---|",
    ]
    for s in steps:
        lines.append(
            "| "
            + " | ".join(
                [
                    md_access_cell(s["id"], s["access"]),
                    md_code_cell(s["py"]),
                    md_cell(s["words"]),
                    md_cell(s["ui"]),
                    md_cell(s["why"]),
                ]
            )
            + " |"
        )
    return "\n".join(lines)


def update_readme(text: str, steps_map: dict[str, list[dict[str, str]]]) -> str:
    text = text.replace(
        "5 колонок: VBA · Python · словами · почему/зачем · страница/поля.",
        "5 колонок: VBA · Python · словами · страница/поля · почему/зачем.",
    )
    text = text.replace(
        "5 колонок: VBA · Python · словами · почему/зачем · страница/поля",
        "5 колонок: VBA · Python · словами · страница/поля · почему/зачем",
    )
    sections = [
        ("## 1. Коэфф", "KOEFF", "## 2. Распред"),
        ("## 2. Распред", "RASPRED", "## 2a."),
        ("## 3. Топливо", "TOPLIVO", "## Сводка порта"),
    ]
    for start_h, const_name, end_h in sections:
        start = text.index(start_h)
        end = start + text[start:].index(end_h)
        block = text[start:end]
        table_pos = block.index("| № / Access")
        intro = block[:table_pos].rstrip() + "\n\n"
        intro = intro.replace("почему/зачем · страница/поля", "страница/поля · почему/зачем")
        new_table = build_readme_table(steps_map[const_name])
        text = text[:start] + intro + new_table + "\n\n" + text[end:]
    return text


def set_cell_shading(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def write_cell(cell, text: str, *, bold: bool = False, mono: bool = False, size: int = 8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    font = "Consolas" if mono else "Calibri"
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:cs"), font)


def add_step_table(doc: Document, title: str, steps: list[dict[str, str]]) -> None:
    doc.add_heading(title, level=2)
    headers = [
        "№ / Access (VBA)",
        "Участок в АРМ ГС (Python)",
        "Что сделать словами",
        "Страница / что читается и пишется",
        "Почему / зачем",
    ]
    table = doc.add_table(rows=1 + len(steps), cols=5)
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        write_cell(hdr.cells[i], h, bold=True, size=9)
        set_cell_shading(hdr.cells[i], "D9E2F3")
    # highlight why header
    set_cell_shading(hdr.cells[4], "FFF2CC")
    for ri, s in enumerate(steps, start=1):
        row = table.rows[ri]
        write_cell(row.cells[0], f"{s['id']}\n{s['access']}", mono=True, size=7)
        write_cell(row.cells[1], s["py"], mono=True, size=7)
        write_cell(row.cells[2], s["words"], size=8)
        write_cell(row.cells[3], s["ui"], size=8)
        write_cell(row.cells[4], s["why"], size=8)
        if ri % 2 == 0:
            for c in row.cells:
                set_cell_shading(c, "F2F2F2")
        set_cell_shading(row.cells[4], "FFFBEA" if ri % 2 else "FFF2CC")


def build_docx(steps_map: dict[str, list[dict[str, str]]]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, Cm(1.2))
    doc.add_heading('Разбор кнопок «Коэфф / Распред / Топливо»', level=1)
    p = doc.add_paragraph(
        "5 столбцов: VBA · Python · словами · страница/поля · Почему/зачем (5-й)."
    )
    for run in p.runs:
        run.font.size = Pt(10)
    add_step_table(doc, "1. Коэфф — Кнопка5_Click", steps_map["KOEFF"])
    add_step_table(doc, "2. Распред — Кнопка27_Click", steps_map["RASPRED"])
    add_step_table(doc, "3. Топливо — Кнопка49_Click", steps_map["TOPLIVO"])
    doc.save(DOCX)
    try:
        doc.save(DOCX_ROOT)
    except Exception:
        pass


def main() -> None:
    src = CANVAS.read_text(encoding="utf-8")
    steps_map = {
        "KOEFF": parse_const_steps(src, "KOEFF"),
        "RASPRED": parse_const_steps(src, "RASPRED"),
        "TOPLIVO": parse_const_steps(src, "TOPLIVO"),
    }
    README.write_text(update_readme(README.read_text(encoding="utf-8"), steps_map), encoding="utf-8")
    build_docx(steps_map)
    print("OK", {k: len(v) for k, v in steps_map.items()})


if __name__ == "__main__":
    main()
