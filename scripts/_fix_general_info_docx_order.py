# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DST = r"z:\НИО-10\АРМ ГС\JSON\Перечень_показателей_json_станции_агрегаты.docx"
doc = Document(DST)

heading_p = body_p = gentypes = None
for p in doc.paragraphs:
    t = p.text.strip()
    if t == "general_info":
        heading_p = p
    elif t.startswith("{") and "general_info" in t and "Текущий год" in t:
        body_p = p
    elif t == "gentypes_info":
        gentypes = p

if not heading_p or not body_p or not gentypes:
    raise SystemExit("blocks missing")

heading_text = heading_p.text
body_text = body_p.text
for el in (heading_p._element, body_p._element):
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def insert_before(paragraph, text):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    np = Paragraph(new_p, paragraph._parent)
    np.text = text
    return np


body_para = insert_before(gentypes, body_text)
insert_before(body_para, heading_text)
doc.save(DST)

print("final:")
for i, p in enumerate(Document(DST).paragraphs):
    t = p.text.strip()
    if t:
        print(i, t[:100].replace("\n", " / "))
