# -*- coding: utf-8 -*-
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

PATH = ROOT / "app/СВОД 2024 год/02.06.26 Расчет электроемкости. Таблица 1.xlsx"
FD = {"СЗФО": (51, 144), "ЦФО": (144, 237)}


def inspect_fd(ws_f, ws_v, lo: int, hi: int, fd_code: str) -> None:
    print("=" * 72)
    print(fd_code)
    print("=" * 72)

    # 1) Summary VRP block (first accum + intensity at FO start)
    for r in range(lo + 1, min(lo + 12, hi)):
        b = str(ws_f.cell(r, 2).value or "")
        if b.startswith("Накопленные Инвестиции"):
            acc, inten = r, r + 1
            b_int = str(ws_f.cell(inten, 2).value or "")
            print("\n[Свод ФО / ВРП]")
            print(f"  R{acc}: X — накопленные инвестиции (сумма по ВЭД)")
            print(f"  R{inten}: Y — {b_int}")
            print(f"  2010: X={ws_v.cell(acc, 4).value}, Y={ws_v.cell(inten, 4).value}")
            print(f"  2024: X={ws_v.cell(acc, 18).value}, Y={ws_v.cell(inten, 18).value}")
            # check calc row after summary
            for r2 in range(inten + 1, inten + 5):
                b2 = str(ws_f.cell(r2, 2).value or "")
                if "РАСЧЕТ" in b2.upper():
                    print(f"  R{r2}: расчётная Y, formula 2010 = {ws_f.cell(r2, 4).value}")
                    break
                if b2.startswith("Накопленные") or b2 == "А=":
                    print("  Расчётной строки/коэффициентов у свода ВРП нет (только факт)")
                    break
            break

    # 2) Typical VED block
    print("\n[Типовой блок ВЭД — структура строк]")
    for r in range(lo, min(lo + 100, hi)):
        if str(ws_f.cell(r, 2).value or "") != "А=":
            continue
        # ved name above
        ved = ""
        for r2 in range(r - 1, max(lo, r - 6), -1):
            cand = str(ws_f.cell(r2, 2).value or ws_f.cell(r2, 1).value or "").strip()
            if cand and cand not in ("А=", "Х=", "ДЕЛЬТА") and not cand.startswith(
                ("Накоплен", "Электро", "Выпуск", "Потребление")
            ):
                ved = cand
                break
        print(f"\n  ВЭД: {ved}")
        print(f"  R{r}:   А=  значение в C: {ws_v.cell(r, 3).value}")
        print(f"  R{r+1}: Х=  формула: {ws_f.cell(r+1, 3).value}")
        print(f"         Х cached: {ws_v.cell(r+1, 3).value}")
        for r2 in range(r + 2, min(r + 12, hi)):
            b3 = str(ws_f.cell(r2, 2).value or "")
            if b3.startswith("Выпуск"):
                print(f"  R{r2}: выпуск продукции (исходные данные)")
            elif b3.startswith("Потребление ЭЭ"):
                print(f"  R{r2}: потребление ээ (исходные данные)")
            elif b3.startswith("Накопленные Инвестиции"):
                acc = r2
                inten = r2 + 1
                calc = r2 + 3
                print(f"  R{acc}: X графика — накопленные инвестиции")
                print(f"  R{inten}: Y факт — электроёмкость, formula 2011: {ws_f.cell(inten, 5).value}")
                # graph points row
                gp = None
                for r3 in range(inten + 1, calc):
                    if ws_f.cell(r3, 5).value and str(ws_f.cell(r3, 5).value).startswith("=LOG"):
                        gp = r3
                        break
                if gp:
                    print(f"  R{gp}: характерные точки, formula 2011: {ws_f.cell(gp, 5).value}")
                print(f"  R{calc}: Y расчёт, formula 2010: {ws_f.cell(calc, 4).value}")
                print(f"         Y расчёт 2010 cached: {ws_v.cell(calc, 4).value}")
                break
        break  # one VED example per FO


def main() -> None:
    wb_f = openpyxl.load_workbook(PATH, data_only=False)
    wb_v = openpyxl.load_workbook(PATH, data_only=True)
    ws_f = wb_f.active
    ws_v = wb_v.active
    for code, (lo, hi) in FD.items():
        inspect_fd(ws_f, ws_v, lo, hi, code)


if __name__ == "__main__":
    main()
