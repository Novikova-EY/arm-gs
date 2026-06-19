# -*- coding: utf-8 -*-
"""
Диагностическая разметка строки «Характерные точки графика» в Excel.

Ищет блоки по «Х=» в колонке B, подсвечивает красным текст и добавляет
комментарий к ячейкам этой строки (расчёт по электроёмкости и инвестициям).
Значения и формулы не меняются.

Использование:
    python scripts/mark_electrical_intensity_xlsx_diagnostics.py path/to/file.xlsx
    python scripts/mark_electrical_intensity_xlsx_diagnostics.py path/to/file.xlsx -o out.xlsx
"""

from __future__ import annotations

import argparse
import re
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.energy_consumption.electrical_intensity.services.electrical_intensity_diagnostic_services import (  # noqa: E402
    compute_intensity_cell_diagnostics,
)

RED_FONT = Font(color="FF0000")
_YEAR_RE = re.compile(r"(19|20)\d{2}")


def _normalize_label(value) -> str:
    return " ".join(str(value or "").split()).strip().lower()


def _parse_decimal(value) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    s = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    if not s or s in ("—", "-", "–"):
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _parse_year_header(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)) and float(value) == int(value):
        y = int(value)
        if 1900 <= y <= 2200:
            return y
    m = _YEAR_RE.search(str(value))
    return int(m.group(0)) if m else None


def _detect_year_columns(ws, header_row: int = 1) -> dict[int, int]:
    cols: dict[int, int] = {}
    for col in range(3, (ws.max_column or 1) + 1):
        y = _parse_year_header(ws.cell(row=header_row, column=col).value)
        if y is not None:
            cols[col] = y
    return cols


def _detect_current_year(year_cols: dict[int, int]) -> int | None:
    if not year_cols:
        return None
    return max(year_cols.values())


def _row_series(ws, row: int, year_cols: dict[int, int]) -> dict[int, Decimal | None]:
    return {
        year: _parse_decimal(ws.cell(row=row, column=col).value)
        for col, year in year_cols.items()
    }


def _is_x_marker(value) -> bool:
    if value is None:
        return False
    s = str(value).strip().upper().replace(" ", "")
    return s.startswith("Х=") or s.startswith("X=")


def mark_workbook(ws) -> int:
    year_cols = _detect_year_columns(ws)
    if not year_cols:
        raise ValueError("Не найдены столбцы с годами (ожидаются с 3-го столбца).")
    display_years = sorted(year_cols.values())
    current_year = _detect_current_year(year_cols)
    marked = 0

    for row_idx in range(1, (ws.max_row or 1) + 1):
        if not _is_x_marker(ws.cell(row=row_idx, column=2).value):
            continue
        investment_row = row_idx + 3
        energy_row = row_idx + 4
        graph_point_row = row_idx + 5
        inv_label = _normalize_label(ws.cell(row=investment_row, column=2).value)
        ei_label = _normalize_label(ws.cell(row=energy_row, column=2).value)
        gp_label = _normalize_label(ws.cell(row=graph_point_row, column=2).value)
        if "инвестиции" not in inv_label or "электроемкость" not in ei_label:
            continue
        if "характер" not in gp_label and "точк" not in gp_label:
            continue

        intensity_cells = _row_series(ws, energy_row, year_cols)
        investment_cells = _row_series(ws, investment_row, year_cols)
        notes = compute_intensity_cell_diagnostics(
            display_years=display_years,
            intensity_cells=intensity_cells,
            investment_cells=investment_cells,
            current_year=current_year,
        )
        col_by_year = {year: col for col, year in year_cols.items()}
        for year, text in notes.items():
            col = col_by_year.get(year)
            if col is None:
                continue
            cell = ws.cell(row=graph_point_row, column=col)
            cell.font = RED_FONT
            cell.comment = Comment(text, "Диагностика")
            marked += 1
    return marked


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Диагностическая разметка «Характерные точки графика» в Excel."
    )
    parser.add_argument("input", type=Path, help="Исходный .xlsx")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="Файл результата (по умолчанию — перезапись входного)",
    )
    args = parser.parse_args()
    out_path = args.output or args.input
    wb = load_workbook(args.input)
    ws = wb.active
    count = mark_workbook(ws)
    wb.save(out_path)
    print(f"Размечено ячеек: {count}. Сохранено: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
