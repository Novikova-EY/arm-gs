# -*- coding: utf-8 -*-
"""Детальное сравнение графиков СЗФО: Excel vs страница (версия БД 20)."""
from __future__ import annotations

import math
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

EXCEL = ROOT / "app/СВОД 2024 год/02.06.26 Расчет электроемкости. Таблица 1.xlsx"
YEARS = list(range(2010, 2026))


def yc(year: int) -> int:
    return 4 + (year - 2010)


def _f(v, n=4):
    if v is None:
        return None
    try:
        x = float(v)
        return None if math.isnan(x) else round(x, n)
    except (TypeError, ValueError):
        return None


def parse_excel_szfo(ws) -> dict[str, dict]:
    """Блоки ВЭД: заголовок в col B, затем A=/Х=, затем строки графика."""
    blocks: dict[str, dict] = {}
    r = 51
    pending_name: str | None = None
    while r < 144:
        b2 = str(ws.cell(r, 2).value or "").strip()
        if b2 == "А=":
            coef_a = ws.cell(r, 3).value
            coef_x = ws.cell(r + 1, 3).value if str(ws.cell(r + 1, 2).value or "") == "Х=" else None
            name = pending_name or f"block@{r}"
            acc_row = inten_row = calc_row = None
            for r2 in range(r + 2, min(r + 15, 144)):
                b3 = str(ws.cell(r2, 2).value or "").strip()
                if b3.startswith("Накопленные Инвестиции"):
                    acc_row = r2
                    inten_row = r2 + 1
                    calc_row = r2 + 3
                    break
            if acc_row and name:
                blocks[name] = {
                    "A": coef_a,
                    "X": coef_x,
                    "acc": {y: ws.cell(acc_row, yc(y)).value for y in YEARS},
                    "y": {y: ws.cell(inten_row, yc(y)).value for y in YEARS},
                    "calc": {y: ws.cell(calc_row, yc(y)).value for y in YEARS},
                }
            pending_name = None
        elif (
            b2
            and b2 not in ("СЗФО", "А=", "Х=", "ДЕЛЬТА")
            and not b2.startswith("Накопленные")
            and not b2.startswith("Электроемкость")
            and not b2.startswith("Выпуск")
            and not b2.startswith("Потребление")
            and ws.cell(r, 4).value is None
        ):
            pending_name = b2
        r += 1
    blocks["__summary__"] = {
        "acc": {y: ws.cell(57, yc(y)).value for y in YEARS},
        "y": {y: ws.cell(58, yc(y)).value for y in YEARS},
        "calc": {},
    }
    blocks["__population__"] = {
        "A": ws.cell(135, 3).value,
        "X": ws.cell(136, 3).value,
        "acc": {y: ws.cell(139, yc(y)).value for y in YEARS},
        "y": {},  # separate row for pop intensity
    }
    # population intensity row is after monetary income - find it
    for r2 in range(139, 144):
        b2 = str(ws.cell(r2, 2).value or "")
        if b2.startswith("Электроемкость") and "РАСЧЕТ" not in b2.upper():
            blocks["__population__"]["y"] = {y: ws.cell(r2, yc(y)).value for y in YEARS}
            blocks["__population__"]["calc"] = {
                y: ws.cell(r2 + 2, yc(y)).value for y in YEARS
            }
            break
    return blocks


def match_excel_block(name: str, blocks: dict[str, dict]) -> tuple[str, dict] | None:
    if name == "Население":
        return "__population__", blocks["__population__"]
    for key, val in blocks.items():
        if key.startswith("__"):
            continue
        if key in name or name in key:
            return key, val
        # first word match
        if key.split()[0][:6] == name.split()[0][:6]:
            return key, val
    return None


def main():
    from flask import g, session
    from run import app
    from app.electrical_intensity.services.electrical_intensity_page_services import (
        parse_electrical_intensity_page_kwargs,
    )
    from app.electrical_intensity.services.electrical_intensity_services import (
        build_electrical_intensity_page_context,
    )

    ws = openpyxl.load_workbook(EXCEL, data_only=True).active
    excel_blocks = parse_excel_szfo(ws)

    url = "/electrical_intensity_fo/?rounding_digits=1&start_year=2010&end_year=2042"
    with app.app_context():
        with app.test_request_context(url):
            session["current_db_version"] = 20
            g.current_db_version = 20
            kw = parse_electrical_intensity_page_kwargs(rounding_digits=1)
            for k in (
                "data_start_year",
                "data_end_year",
                "lt_ei_initial_visible_years",
                "lt_ei_year_seg_state",
            ):
                kw.pop(k, None)
            ctx = build_electrical_intensity_page_context(**kw)

    block = next(b for b in ctx["territory_blocks"] if b.get("abbr") == "СЗФО")
    print("СЗФО — сравнение фактических и расчётных точек графика\n")

    ok = mism = 0
    for section in block.get("ved_sections") or []:
        name = section.get("ved_name") or ""
        chart = section.get("scatter_chart")
        if not chart:
            print(f"[—] {name}: график не построен на странице")
            continue
        matched = match_excel_block(name, excel_blocks)
        if not matched:
            print(f"[!] {name}: нет блока в Excel")
            continue
        ex_key, ex = matched
        print(f"=== {name} (Excel: {ex_key}) ===")
        pa = section.get("coefficient_a_manual") or section.get("coefficient_a")
        px = section.get("coefficient_x")
        print(f"  A: Excel={_f(ex.get('A'), 2)} Page={_f(pa, 2)} Acalc={_f(section.get('coefficient_a_computed'), 2)}")
        print(f"  X: Excel={_f(ex.get('X'), 4)} Page={_f(px, 4)}")
        if ex.get("X") is not None and px is not None and abs(float(ex["X"]) - float(px)) > 0.001:
            print("  ** X отличается")
            mism += 1
        else:
            ok += 1
        for year in YEARS:
            fact = next((p for p in chart.get("fact", []) if p["year"] == year), None)
            calc = next((p for p in chart.get("calc", []) if p["year"] == year), None)
            issues = []
            if fact and ex.get("y"):
                if _f(fact["x"], 1) != _f(ex["acc"].get(year), 1):
                    issues.append(f"x: page={_f(fact['x'])} excel={_f(ex['acc'].get(year))}")
                if _f(fact["y"], 2) != _f(ex["y"].get(year), 2):
                    issues.append(f"y_fact: page={_f(fact['y'],2)} excel={_f(ex['y'].get(year),2)}")
            if calc and ex.get("calc"):
                if _f(calc["y"], 2) != _f(ex["calc"].get(year), 2):
                    issues.append(f"y_calc: page={_f(calc['y'],2)} excel={_f(ex['calc'].get(year),2)}")
            if issues:
                print(f"  {year}: " + "; ".join(issues))
                mism += len(issues)
            else:
                ok += 1

    # summary
    sc = block.get("fd_summary", {}).get("scatter_chart")
    ex = excel_blocks["__summary__"]
    print("=== Свод ФО (Электроемкость ВРП) ===")
    for year in [2010, 2015, 2020, 2024]:
        fact = next((p for p in sc.get("fact", []) if p["year"] == year), None)
        if fact:
            match = _f(fact["x"], 1) == _f(ex["acc"][year], 1) and _f(fact["y"], 2) == _f(
                ex["y"][year], 2
            )
            print(
                f"  {year}: page=({_f(fact['x'])}, {_f(fact['y'],2)}) excel=({_f(ex['acc'][year])}, {_f(ex['y'][year],2)}) {'OK' if match else 'DIFF'}"
            )
            ok += match
            mism += not match

    print(f"\nИтого проверок: ok≈{ok}, расхождений≈{mism}")


if __name__ == "__main__":
    main()
