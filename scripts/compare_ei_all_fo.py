# -*- coding: utf-8 -*-
"""Сравнение графиков электроёмкости по всем ФО: Excel vs страница (БД v20)."""
from __future__ import annotations

import math
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

EXCEL = ROOT / "app/СВОД 2024 год/02.06.26 Расчет электроемкости. Таблица 1.xlsx"
FD_RANGES = {
    "СЗФО": (51, 144),
    "ЦФО": (144, 237),
    "ЮФО": (237, 330),
    "ПФО": (330, 423),
    "СКФО": (423, 516),
    "УФО": (516, 609),
    "СФО": (609, 702),
    "ДФО": (702, 837),
}
YEARS = list(range(2010, 2026))


def yc(year: int) -> int:
    return 4 + (year - 2010)


def _f(v, n=4):
    if v is None:
        return None
    try:
        x = float(v)
        return None if math.isnan(x) else round(x, n)
    except (TypeError, ValueError):
        return None


def parse_fd_excel(ws, lo: int, hi: int) -> dict:
    blocks: dict[str, dict] = {}
    pending_name: str | None = None
    summary_acc = summary_y = None
    for r in range(lo, hi):
        b2 = str(ws.cell(r, 2).value or "").strip()
        if b2.startswith("Накопленные Инвестиции") and summary_acc is None and pending_name is None:
            # first accum in FO block = summary VRP
            summary_acc = {y: ws.cell(r, yc(y)).value for y in YEARS}
            summary_y = {y: ws.cell(r + 1, yc(y)).value for y in YEARS}
        if b2 == "А=":
            coef_a = ws.cell(r, 3).value
            coef_x = (
                ws.cell(r + 1, 3).value
                if str(ws.cell(r + 1, 2).value or "") == "Х="
                else None
            )
            name = pending_name or f"@r{r}"
            acc_row = None
            for r2 in range(r + 2, min(r + 15, hi)):
                b3 = str(ws.cell(r2, 2).value or "").strip()
                if b3.startswith("Накопленные Инвестиции"):
                    acc_row = r2
                    break
                if b3.startswith("Накопленные денежные"):
                    acc_row = r2
                    break
            if acc_row:
                inten_row = acc_row + 1
                calc_row = acc_row + 3
                blocks[name] = {
                    "A": coef_a,
                    "X": coef_x,
                    "acc": {y: ws.cell(acc_row, yc(y)).value for y in YEARS},
                    "y": {y: ws.cell(inten_row, yc(y)).value for y in YEARS},
                    "calc": {
                        y: ws.cell(calc_row, yc(y)).value
                        for y in YEARS
                        if ws.cell(calc_row, yc(y)).value is not None
                    },
                    "is_pop": b3.startswith("Накопленные денежные"),
                }
            pending_name = None
        elif (
            b2
            and b2 not in ("А=", "Х=", "ДЕЛЬТА")
            and not b2.startswith("Накопленные")
            and not b2.startswith("Электроемкость")
            and not b2.startswith("Выпуск")
            and not b2.startswith("Потребление")
            and ws.cell(r, 4).value is None
            and b2 != str(ws.cell(lo, 1).value or "").strip()
        ):
            pending_name = b2
    return {
        "blocks": blocks,
        "summary": {"acc": summary_acc or {}, "y": summary_y or {}},
    }


def match_name(page_name: str, blocks: dict[str, dict]) -> tuple[str, dict] | None:
    if page_name == "Население":
        for k, v in blocks.items():
            if v.get("is_pop"):
                return k, v
        return None
    for k, v in blocks.items():
        if v.get("is_pop"):
            continue
        if k in page_name or page_name in k:
            return k, v
        if k.split()[0][:5] == page_name.split()[0][:5]:
            return k, v
    return None


def compare_section(name: str, section: dict, ex: dict) -> list[str]:
    issues = []
    chart = section.get("scatter_chart") or {}
    pa = section.get("coefficient_a_manual") or section.get("coefficient_a")
    px = section.get("coefficient_x")
    if ex.get("A") is not None and pa is not None and _f(ex["A"], 2) != _f(pa, 2):
        issues.append(f"A: page={_f(pa,2)} excel={_f(ex['A'],2)}")
    if ex.get("X") is not None and px is not None and _f(ex["X"], 4) != _f(px, 4):
        issues.append(f"X: page={_f(px,4)} excel={_f(ex['X'],4)}")
    for year in YEARS:
        fact = next((p for p in chart.get("fact", []) if p["year"] == year), None)
        calc = next((p for p in chart.get("calc", []) if p["year"] == year), None)
        if fact:
            if _f(fact["x"], 1) != _f(ex["acc"].get(year), 1):
                issues.append(f"{year} x: {_f(fact['x'])} vs {_f(ex['acc'].get(year))}")
            if ex["y"].get(year) is not None and _f(fact["y"], 2) != _f(ex["y"].get(year), 2):
                issues.append(f"{year} y: {_f(fact['y'],2)} vs {_f(ex['y'].get(year),2)}")
        if calc and ex.get("calc") and year in ex["calc"]:
            if _f(calc["y"], 2) != _f(ex["calc"].get(year), 2):
                issues.append(f"{year} y_calc: {_f(calc['y'],2)} vs {_f(ex['calc'].get(year),2)}")
    return issues


def main():
    from flask import g, session
    from run import app
    from app.electrical_intensity.services.electrical_intensity_page_services import (
        parse_electrical_intensity_page_kwargs,
    )
    from app.electrical_intensity.services.electrical_intensity_services import (
        build_electrical_intensity_page_context,
    )

    ws = openpyxl.load_workbook(EXCEL, data_only=True).active
    excel = {fd: parse_fd_excel(ws, lo, hi) for fd, (lo, hi) in FD_RANGES.items()}

    url = "/electrical_intensity_fo/?rounding_digits=1&start_year=2010&end_year=2042"
    with app.app_context():
        with app.test_request_context(url):
            session["current_db_version"] = 20
            g.current_db_version = 20
            kw = parse_electrical_intensity_page_kwargs(rounding_digits=1)
            for k in (
                "data_start_year",
                "data_end_year",
                "lt_ei_initial_visible_years",
                "lt_ei_year_seg_state",
            ):
                kw.pop(k, None)
            ctx = build_electrical_intensity_page_context(**kw)

    total_issues = 0
    total_sections = 0
    for block in ctx["territory_blocks"]:
        if block.get("territory_kind") != "fd":
            continue
        abbr = block.get("abbr") or ""
        ex_fd = excel.get(abbr)
        if not ex_fd:
            print(f"[!] Нет Excel для {abbr}")
            continue
        print(f"\n{'='*60}\n{abbr} — {block.get('label')}\n{'='*60}")
        fo_issues = 0
        for section in block.get("ved_sections") or []:
            name = section.get("ved_name") or ""
            if not section.get("scatter_chart"):
                if name != "Промышленное производство":
                    print(f"  [—] {name}: нет графика")
                continue
            total_sections += 1
            matched = match_name(name, ex_fd["blocks"])
            if not matched:
                print(f"  [!] {name}: не найден в Excel")
                fo_issues += 1
                continue
            ex_key, ex_block = matched
            issues = compare_section(name, section, ex_block)
            if issues:
                print(f"  ** {name} (Excel: {ex_key[:40]})")
                for iss in issues[:8]:
                    print(f"     {iss}")
                if len(issues) > 8:
                    print(f"     ... ещё {len(issues)-8}")
                fo_issues += len(issues)
            else:
                print(f"  OK {name}")
        sc = (block.get("fd_summary") or {}).get("scatter_chart")
        if sc and ex_fd["summary"]["acc"]:
            sum_issues = []
            for year in [2010, 2015, 2020, 2024]:
                fact = next((p for p in sc["fact"] if p["year"] == year), None)
                if not fact:
                    continue
                if _f(fact["x"], 1) != _f(ex_fd["summary"]["acc"].get(year), 1):
                    sum_issues.append(f"{year} x")
                if _f(fact["y"], 2) != _f(ex_fd["summary"]["y"].get(year), 2):
                    sum_issues.append(f"{year} y")
            if sum_issues:
                print(f"  ** Свод ВРП: {', '.join(sum_issues)}")
                fo_issues += len(sum_issues)
            else:
                print("  OK Свод ФО (ВРП)")
        print(f"  Итог {abbr}: расхождений {fo_issues}")
        total_issues += fo_issues

    print(f"\n{'='*60}")
    print(f"Всего секций с графиками: {total_sections}")
    print(f"Всего расхождений: {total_issues}")
    if total_issues == 0:
        print("ВЫВОД: графики на странице полностью совпадают с Excel (версия БД 20).")


if __name__ == "__main__":
    main()
