# -*- coding: utf-8 -*-
"""Добавить general_info в Перечень_показателей_json_станции_агрегаты.docx."""
from __future__ import annotations

import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8")

DST = r"z:\НИО-10\АРМ ГС\JSON\Перечень_показателей_json_станции_агрегаты.docx"

GENERAL_HEADING = "general_info"
GENERAL_BODY = """{
    "dataset": "general_info",
    "database_version": 46,
    "version": 46,
    "comment": "",
    "rows": [
        {
            "Текущий год": 2024,
            "Начало периода": 2024,
            "Конец периода": 2042
        }
    ]
}
Источник: /refdata/ — «Текущий год», расчётный период версии по полям СиПР / Генеральная схема
(начало = year_sipr_start−2; конец = year_sipr_end для СиПР или 2042 для ГС).
СиПР и ГС различаются выбором версии БД в /api/database_versions (planning_scheme: sipr|gs),
а не отдельным API наборов."""


def insert_paragraph_before(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.text = text
    return new_para


def main() -> None:
    doc = Document(DST)

    # Update datasets list in API blurb
    for p in doc.paragraphs:
        if "список имён наборов:" in p.text and "general_info" not in p.text:
            p.text = p.text.replace(
                "gentypes_info, generation_objects",
                "general_info, gentypes_info, generation_objects",
            )

    # Insert general_info block before gentypes_info heading if missing
    has_general = any(p.text.strip() == GENERAL_HEADING for p in doc.paragraphs)
    if not has_general:
        target = None
        for p in doc.paragraphs:
            if p.text.strip() == "gentypes_info":
                target = p
                break
        if target is None:
            raise SystemExit("gentypes_info heading not found")
        insert_paragraph_before(target, GENERAL_BODY)
        insert_paragraph_before(target, GENERAL_HEADING)

    doc.save(DST)
    print("updated", DST)


if __name__ == "__main__":
    main()
