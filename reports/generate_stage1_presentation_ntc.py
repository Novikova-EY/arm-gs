# -*- coding: utf-8 -*-
"""
Презентация Этапа 1 АРМ «Генеральная схема» на корпоративном шаблоне НТЦ ЕЭС.

База: АРМ_результаты 2026.pptx (шаблон 16:9 НТЦ ЕЭС 2025, 25.4×14.29 см).
Дополнительно: свежие скриншоты и визуальные акценты под отчет Этапа 1.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt


ROOT = Path(r"c:\arm_gs\reports")
SHOTS = ROOT / "screenshots"
MATERIALS = Path(r"Z:\НИО-10\АРМ ГС\АРМ_результаты 2026.pptx")

OUT_PATHS = [
    ROOT / "Презентация_Этап1_АРМ_ГС-2026.pptx",
    Path(r"Z:\НИО-10\АРМ ГС\Отчеты НИР\2026 год")
    / "Презентация_Этап1_АРМ_ГС-2026.pptx",
    Path(r"Z:\НИО-10\АРМ ГС\ТЗ на 2026 год")
    / "Презентация_Этап1_АРМ_ГС-2026.pptx",
]

# Палитра шаблона НТЦ ЕЭС 2025
TEAL = RGBColor(0x00, 0x79, 0x76)
CYAN = RGBColor(0x00, 0xBC, 0xB8)
NAVY = RGBColor(0x00, 0x3B, 0x8A)
BLUE = RGBColor(0x00, 0x67, 0xBC)
GOLD2 = RGBColor(0xAF, 0x98, 0x5D)
DARK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xE8, 0xF4, 0xF3)

FONT = "Century Gothic"
SLIDE_W = Cm(25.4)
SLIDE_H = Cm(14.2875)


def _run(p, text, *, size=14, bold=False, color=DARK, align=None):
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def _textbox(slide, left, top, width, height, text, *, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, size=size, bold=bold, color=color)
    return box


def _round_rect(slide, left, top, width, height, *, fill=CARD, line=TEAL, line_w=Pt(1.0)):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = line_w
    try:
        sh.adjustments[0] = 0.1
    except Exception:
        pass
    return sh


def _accent_bar(slide, left, top, width, height=Cm(0.1), color=GOLD2):
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar


def _circle_num(slide, left, top, num, *, size=Cm(0.7), fill=TEAL):
    c = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, left, top, size, size)
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    tf = c.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, str(num), size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    return c


def _pic(slide, stem, left, top, width, height=None):
    path = SHOTS / f"{stem}.png"
    if not path.exists():
        return None
    if height is None:
        return slide.shapes.add_picture(str(path), left, top, width=width)
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def _shadow_frame(slide, left, top, width, height):
    shadow = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        left - Cm(0.06),
        top - Cm(0.06),
        width + Cm(0.12),
        height + Cm(0.12),
    )
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = RGBColor(0xD0, 0xE0, 0xDF)
    shadow.line.fill.background()
    try:
        shadow.adjustments[0] = 0.06
    except Exception:
        pass
    return shadow


def _set_placeholder_text(shape, lines: list[str], *, size=None, bold=None, color=None):
    tf = shape.text_frame
    while len(tf.paragraphs) > 1:
        p = tf.paragraphs[-1]._p
        p.getparent().remove(p)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if p.runs:
            p.runs[0].text = line
            for r in p.runs[1:]:
                r.text = ""
            run = p.runs[0]
        else:
            run = p.add_run()
            run.text = line
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color
        run.font.name = FONT


def _replace_all_text(slide, mapping: dict[str, str]) -> None:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            full = "".join(r.text for r in para.runs) if para.runs else para.text
            for old, new in mapping.items():
                if old in full:
                    if para.runs:
                        para.runs[0].text = full.replace(old, new)
                        for r in para.runs[1:]:
                            r.text = ""
                    else:
                        para.text = full.replace(old, new)


def _move_slide(prs: Presentation, old_index: int, new_index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst
    els = list(sld_id_lst)
    el = els[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


def _find_title_ph(slide):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        try:
            if shape.is_placeholder and shape.placeholder_format.idx == 0:
                return shape
        except Exception:
            continue
    for shape in slide.shapes:
        if shape.has_text_frame and "Заголовок" in getattr(shape, "name", ""):
            return shape
    return None


def _header(slide, title: str, *, fill=NAVY):
    header = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Cm(0), Cm(0), SLIDE_W, Cm(1.55)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = fill
    header.line.fill.background()
    _textbox(slide, Cm(0.7), Cm(0.35), Cm(23.5), Cm(0.95), title, size=20, bold=True, color=WHITE)
    _accent_bar(slide, Cm(0.7), Cm(1.55), Cm(7.5), Cm(0.09), CYAN if fill != TEAL else GOLD2)


def update_title(prs: Presentation) -> None:
    slide = prs.slides[0]
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        name = getattr(shape, "name", "")
        text = shape.text_frame.text
        if "Заголовок" in name or "Результаты" in text:
            _set_placeholder_text(
                shape,
                [
                    "Результаты Этапа 1 развития программного комплекса",
                    "АРМ «Генеральная схема» (2026)",
                ],
                size=26,
                bold=True,
                color=WHITE,
            )
        elif "Подзаголовок" in name or "Беляев" in text:
            _set_placeholder_text(
                shape,
                [
                    "ТЗ – Приложение 1 · график – Приложение 3",
                    "Период: 12.01.2026 – 30.09.2026  ·  АО «НТЦ ЕЭС» · НИО-10",
                ],
                size=14,
                bold=False,
                color=WHITE,
            )


def update_results_slide(prs: Presentation) -> None:
    for slide in prs.slides:
        joined = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
        if "Результат развития АРМ в 2026" not in joined:
            continue
        _replace_all_text(
            slide,
            {"Результат развития АРМ в 2026 году": "Итоги Этапа 1 (2026)"},
        )
        break


def update_thanks(prs: Presentation) -> None:
    for slide in prs.slides:
        joined = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
        if "Спасибо за внимание" not in joined:
            continue
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if "Беляев" in shape.text_frame.text or "Подзаголовок" in getattr(shape, "name", ""):
                _set_placeholder_text(
                    shape,
                    [
                        "АО «НТЦ ЕЭС» · НИО-10",
                        "АРМ «Генеральная схема» · Этап 1 · 2026",
                    ],
                    size=14,
                    color=WHITE,
                )
        break


def add_section_divider(prs: Presentation, title: str, subtitle: str) -> None:
    layout = prs.slide_layouts[2]
    slide = prs.slides.add_slide(layout)
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        try:
            idx = shape.placeholder_format.idx
        except Exception:
            continue
        if idx == 0:
            _set_placeholder_text(shape, [title], size=28, bold=True, color=WHITE)
        elif idx in (1, 2):
            _set_placeholder_text(shape, [subtitle], size=14, color=WHITE)
    _accent_bar(slide, Cm(1.0), Cm(8.8), Cm(6.5), Cm(0.12), GOLD2)


def add_stage1_agenda(prs: Presentation) -> None:
    layout = prs.slide_layouts[9]
    slide = prs.slides.add_slide(layout)
    _header(slide, "Этап 1: контур реализованных модулей")

    modules = [
        ("01", "Инфраструктура", "версии БД, роли, журналы"),
        ("02", "Справочники", "территории, топлива, НСИ"),
        ("03", "Генерация", "станции, агрегаты, площадки"),
        ("04", "Нагрузки", "потребление, максимумы, коэфф."),
        ("05", "Спрос", "объемы, приросты, ГАЭС"),
        ("06", "Балансы ЭЭ/ЭМ", "выработка, своды, передача"),
        ("07", "Топливо", "группы, расчет, итоги"),
        ("08", "Электроемкость", "ФО, ВЭД, экономика"),
        ("09", "Документы", "JSON / выходные формы"),
    ]
    for i, (num, title, desc) in enumerate(modules):
        col, row = i % 3, i // 3
        left = Cm(0.55) + Cm(8.2) * col
        top = Cm(2.0) + Cm(3.85) * row
        fill = WHITE if i % 2 == 0 else CARD
        _round_rect(slide, left, top, Cm(7.85), Cm(3.5), fill=fill, line=TEAL)
        _circle_num(slide, left + Cm(0.3), top + Cm(0.35), num, fill=TEAL if row < 2 else BLUE)
        _textbox(slide, left + Cm(1.25), top + Cm(0.4), Cm(6.2), Cm(0.65), title, size=13, bold=True, color=NAVY)
        _accent_bar(slide, left + Cm(0.35), top + Cm(1.3), Cm(7.1), Cm(0.07), GOLD2)
        _textbox(slide, left + Cm(0.35), top + Cm(1.6), Cm(7.1), Cm(1.5), desc, size=11, color=MUTED)


def add_infra_showcase(prs: Presentation) -> None:
    layout = prs.slide_layouts[9]
    slide = prs.slides.add_slide(layout)
    _header(slide, "Инфраструктура АРМ: версии, роли, журналы")

    cards = [
        ("02_db_versions", "Версии базы данных", "Выбор рабочей версии и изолированная работа"),
        ("02c_users_roles", "Пользователи и роли", "Админ, гость, админ/редактор"),
        ("02b_user_action_logs", "Журналы действий", "Вход, загрузка, правки – прослеживаемость"),
    ]
    for i, (stem, title, caption) in enumerate(cards):
        left = Cm(0.45) + Cm(8.25) * i
        top = Cm(1.95)
        _round_rect(slide, left, top, Cm(7.95), Cm(11.8), fill=WHITE, line=TEAL)
        _circle_num(slide, left + Cm(0.25), top + Cm(0.25), i + 1, fill=TEAL if i < 2 else BLUE)
        _textbox(slide, left + Cm(1.1), top + Cm(0.3), Cm(6.4), Cm(0.55), title, size=12, bold=True, color=NAVY)
        _accent_bar(slide, left + Cm(0.3), top + Cm(0.95), Cm(7.3), Cm(0.07), GOLD2)
        pic_top = top + Cm(1.15)
        pic_h = Cm(8.9)
        _shadow_frame(slide, left + Cm(0.3), pic_top, Cm(7.35), pic_h)
        _pic(slide, stem, left + Cm(0.3), pic_top, Cm(7.35), pic_h)
        _textbox(slide, left + Cm(0.3), top + Cm(10.3), Cm(7.35), Cm(1.2), caption, size=10, color=MUTED)


def add_generation_highlights(prs: Presentation) -> None:
    layout = prs.slide_layouts[9]
    slide = prs.slides.add_slide(layout)
    _header(slide, "Генерация: изменения УМ, перспективные площадки, итоги", fill=TEAL)

    tiles = [
        ("05f_station_power_changes", "Изменения УМ ЕГО"),
        ("05h_prospective_places_aes", "Перспективные АЭС"),
        ("05g_prospective_places_ges", "Перспективные ГЭС"),
        ("05i_generation_totals", "Итоговые суммы"),
    ]
    positions = [
        (Cm(0.4), Cm(1.9)),
        (Cm(12.85), Cm(1.9)),
        (Cm(0.4), Cm(8.0)),
        (Cm(12.85), Cm(8.0)),
    ]
    for (stem, title), (left, top) in zip(tiles, positions):
        w, h = Cm(12.15), Cm(5.8)
        _round_rect(slide, left, top, w, h, fill=WHITE, line=CYAN)
        _textbox(slide, left + Cm(0.3), top + Cm(0.15), Cm(11), Cm(0.45), title, size=11, bold=True, color=NAVY)
        _pic(slide, stem, left + Cm(0.25), top + Cm(0.65), Cm(11.65), Cm(4.9))


def add_fuel_pipeline(prs: Presentation) -> None:
    layout = prs.slide_layouts[9]
    slide = prs.slides.add_slide(layout)
    _header(slide, "Топливо: расчетная цепочка Коэфф → Распред → Топливо")

    steps = [
        ("1", "Коэфф", "Параметры и коэффициенты расчета"),
        ("2", "Распред", "Распределение выработки / топлива"),
        ("3", "Топливо", "Итоговые топливные показатели"),
    ]
    y = Cm(1.9)
    for i, (num, title, desc) in enumerate(steps):
        left = Cm(0.5) + Cm(8.25) * i
        _round_rect(slide, left, y, Cm(7.9), Cm(2.35), fill=CARD, line=TEAL)
        _circle_num(slide, left + Cm(0.3), y + Cm(0.65), num, size=Cm(0.95), fill=TEAL)
        _textbox(slide, left + Cm(1.5), y + Cm(0.35), Cm(6), Cm(0.6), title, size=16, bold=True, color=NAVY)
        _textbox(slide, left + Cm(1.5), y + Cm(1.1), Cm(6), Cm(0.9), desc, size=11, color=MUTED)
        if i < 2:
            arrow = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
                left + Cm(7.55),
                y + Cm(0.75),
                Cm(1.0),
                Cm(0.7),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GOLD2
            arrow.line.fill.background()

    shots = [
        ("14_fuel_calculation", "Интерфейс расчета"),
        ("14b_fuel_balance", "Топливный баланс"),
        ("14e_fuel_totals", "Итоги по топливу"),
    ]
    for i, (stem, cap) in enumerate(shots):
        left = Cm(0.45) + Cm(8.25) * i
        top = Cm(4.55)
        _shadow_frame(slide, left, top, Cm(7.95), Cm(8.0))
        _pic(slide, stem, left, top, Cm(7.95), Cm(8.0))
        _textbox(
            slide, left, top + Cm(8.05), Cm(7.95), Cm(0.4),
            cap, size=10, bold=True, color=TEAL, align=PP_ALIGN.CENTER,
        )


def add_ei_economics(prs: Presentation) -> None:
    layout = prs.slide_layouts[9]
    slide = prs.slides.add_slide(layout)
    _header(slide, "Электроемкость и экономика: своды, ВЭД, журналы", fill=BLUE)

    left_panel = [
        ("15_electrical_intensity", "Хаб электроемкости"),
        ("16_electrical_intensity_fo", "Таблица по ФО"),
        ("16b_ei_ved_consumption", "ВЭД / потребление"),
    ]
    right_panel = [
        ("16c_economics_hub", "Хаб экономики"),
        ("16d_economics_afci", "АФКИ / инвестиции"),
        ("16h_ved_change_log", "Журнал изменений"),
    ]

    _textbox(slide, Cm(0.5), Cm(1.75), Cm(11), Cm(0.4), "Электроемкость", size=12, bold=True, color=TEAL)
    _textbox(slide, Cm(13.0), Cm(1.75), Cm(11), Cm(0.4), "Экономика", size=12, bold=True, color=NAVY)

    for i, (stem, cap) in enumerate(left_panel):
        top = Cm(2.2) + Cm(3.85) * i
        _round_rect(slide, Cm(0.4), top, Cm(12.1), Cm(3.6), fill=WHITE, line=CYAN)
        _pic(slide, stem, Cm(0.55), top + Cm(0.2), Cm(11.8), Cm(2.85))
        _textbox(slide, Cm(0.55), top + Cm(3.1), Cm(11.8), Cm(0.35), cap, size=10, color=MUTED)

    for i, (stem, cap) in enumerate(right_panel):
        top = Cm(2.2) + Cm(3.85) * i
        _round_rect(slide, Cm(12.9), top, Cm(12.1), Cm(3.6), fill=WHITE, line=TEAL)
        _pic(slide, stem, Cm(13.05), top + Cm(0.2), Cm(11.8), Cm(2.85))
        _textbox(slide, Cm(13.05), top + Cm(3.1), Cm(11.8), Cm(0.35), cap, size=10, color=MUTED)


def add_visual_accents_to_existing(prs: Presentation) -> None:
    keys = (
        "Верхнеуровневая структура",
        "Модуль Генерация",
        "Модуль Балансы",
        "Модуль Топливо",
        "Модуль Нагрузки",
        "Модуль Спрос",
        "Долгосрочный прогноз",
        "Оценка наполненности",
    )
    for slide in prs.slides:
        title_shape = _find_title_ph(slide)
        if title_shape is None:
            continue
        title = title_shape.text_frame.text
        if not any(k in title for k in keys):
            continue
        top = title_shape.top + title_shape.height - Cm(0.04)
        _accent_bar(slide, title_shape.left, top, Cm(5.2), Cm(0.08), GOLD2)


def _clear_slide_number(shape, value: str) -> None:
    """Полностью перезаписывает номер слайда (без склейки со старым значением)."""
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED


def _renumber_slides(prs: Presentation) -> None:
    for i, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            name = getattr(shape, "name", "")
            if "Номер" not in name:
                continue
            if not shape.has_text_frame:
                continue
            try:
                _clear_slide_number(shape, str(i))
            except Exception:
                pass


def _find_idx(prs: Presentation, substr: str) -> int:
    for i, s in enumerate(prs.slides):
        text = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if substr in text:
            return i
    return -1


def _find_idx_exact_title(prs: Presentation, title: str) -> int:
    """Ищет слайд, у которого заголовок (первый существенный текст) совпадает."""
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip().split("\n")[0].strip()
            if t == title:
                return i
    return -1


def build() -> Path:
    if not MATERIALS.exists():
        raise FileNotFoundError(MATERIALS)

    work = ROOT / "_stage1_ppt_work.pptx"
    shutil.copy2(MATERIALS, work)
    prs = Presentation(str(work))

    update_title(prs)
    update_results_slide(prs)
    update_thanks(prs)
    add_visual_accents_to_existing(prs)

    add_stage1_agenda(prs)
    add_infra_showcase(prs)
    add_generation_highlights(prs)
    add_fuel_pipeline(prs)
    add_ei_economics(prs)
    add_section_divider(
        prs,
        "Итоги Этапа 1",
        "Единая среда данных · рабочие модули · прослеживаемость",
    )

    # agenda + infra after structure
    idx = _find_idx(prs, "Этап 1: контур реализованных модулей")
    if idx >= 0:
        _move_slide(prs, idx, 2)
    idx = _find_idx(prs, "Инфраструктура АРМ: версии")
    if idx >= 0:
        _move_slide(prs, idx, 3)

    idx = _find_idx(prs, "Генерация: изменения УМ")
    after = _find_idx(prs, "Модуль Генерация. Передача данных")
    if after >= 0 and idx >= 0:
        _move_slide(prs, idx, after + 1)

    idx = _find_idx(prs, "Топливо: расчетная цепочка")
    after = _find_idx(prs, "состав и характеристики групп оборудования")
    if after < 0:
        after = _find_idx(prs, "Модуль Топливо (интерфейс) – 2")
    if after >= 0 and idx >= 0:
        _move_slide(prs, idx, after + 1)

    idx = _find_idx(prs, "Электроемкость и экономика")
    after = _find_idx(prs, "выпуск продукции и накопленные инвестиции")
    if after < 0:
        after = _find_idx(prs, "М. Долгосрочный прогноз")
    if after >= 0 and idx >= 0:
        _move_slide(prs, idx, after + 1)

    idx = _find_idx(prs, "Единая среда данных · рабочие модули")
    if idx < 0:
        idx = _find_idx_exact_title(prs, "Итоги Этапа 1")
    before = _find_idx(prs, "Оценка наполненности базы данных")
    if before >= 0 and idx >= 0:
        _move_slide(prs, idx, before)

    # Восстановить порядок: наполненность → результаты → спасибо
    idx_full = _find_idx(prs, "Оценка наполненности базы данных")
    idx_res = _find_idx(prs, "Итоги Этапа 1 (2026)")
    idx_thanks = _find_idx(prs, "Спасибо за внимание")
    if idx_full >= 0 and idx_res >= 0 and idx_res < idx_full:
        _move_slide(prs, idx_res, idx_full + 1)
    idx_res = _find_idx(prs, "Итоги Этапа 1 (2026)")
    idx_thanks = _find_idx(prs, "Спасибо за внимание")
    if idx_res >= 0 and idx_thanks >= 0 and idx_thanks < idx_res:
        _move_slide(prs, idx_thanks, idx_res + 1)

    _renumber_slides(prs)

    primary = OUT_PATHS[0]
    prs.save(str(primary))
    for dest in OUT_PATHS[1:]:
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(primary, dest)

    work.unlink(missing_ok=True)
    return primary


if __name__ == "__main__":
    out = build()
    prs = Presentation(str(out))
    print(f"OK: {out}")
    print(f"slides: {len(prs.slides)}")
    for i, s in enumerate(prs.slides, 1):
        title = ""
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                title = sh.text_frame.text.strip().replace("\n", " / ")[:75]
                break
        print(f"{i:02d}. {title}")
