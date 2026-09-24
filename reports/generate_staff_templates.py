# -*- coding: utf-8 -*-
"""Заготовки Word/Excel для заполнения работниками (тестирование / предметка)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(r"Z:\НИО-10\АРМ ГС\04. Тестирование") / "00_Шаблоны_для_заполнения"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
EXAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"


def autosize(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_example_row(ws, row: int, values: list) -> None:
    for col, val in enumerate(values, 1):
        cell = ws.cell(row, col, val)
        cell.fill = EXAMPLE_FILL
        cell.alignment = WRAP
        cell.border = THIN


def blank_rows(ws, start: int, n: int, ncols: int) -> None:
    for r in range(start, start + n):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c, "")
            cell.border = THIN
            cell.alignment = WRAP


def add_list_validation(ws, col_letter: str, options: str, start: int = 3, end: int = 100) -> None:
    dv = DataValidation(type="list", formula1=options, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{start}:{col_letter}{end}")


def add_meta_sheet(wb, title: str, how_to: list[str], fields: list[tuple[str, str]]) -> None:
    ws = wb.create_sheet("Инструкция", 0)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A2"] = f"Дата шаблона: {date.today().isoformat()}"
    ws["A3"] = (
        "Жёлтая строка — пример. Заголовки не удаляйте. "
        "Заполняйте со следующей пустой строки."
    )
    ws["A4"] = "Обязательно: версия БД, годы, автор (E1/E2/S), дата."
    ws.merge_cells("A1:D1")
    ws["A6"] = "Как заполнять:"
    ws["A6"].font = Font(bold=True)
    r = 7
    for line in how_to:
        ws.cell(r, 1, line)
        r += 1
    r += 1
    ws.cell(r, 1, "Поля:").font = Font(bold=True)
    r += 1
    for name, desc in fields:
        ws.cell(r, 1, name).font = Font(bold=True)
        ws.cell(r, 2, desc)
        r += 1
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70


def make_guide_docx(name: str, title: str, sections: list[tuple[str, list[str]]]) -> Path:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2)

    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    meta = doc.add_paragraph()
    meta.add_run("Модуль: ").bold = True
    meta.add_run("…\t")
    meta.add_run("Автор: ").bold = True
    meta.add_run("…\t")
    meta.add_run("Дата: ").bold = True
    meta.add_run("…\t")
    meta.add_run("Версия БД: ").bold = True
    meta.add_run("…")

    for sec, hints in sections:
        h = doc.add_paragraph()
        r = h.add_run(sec)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        for hint in hints:
            doc.add_paragraph(hint, style="List Bullet")

    path = ROOT / name
    doc.save(path)
    return path


def make_readme_task() -> Path:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2)

    p = doc.add_paragraph()
    run = p.add_run("README задачи")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    def field(label: str, hint: str = "") -> None:
        para = doc.add_paragraph()
        r = para.add_run(f"{label}: ")
        r.bold = True
        para.add_run(hint)

    field("Код задачи", "например E1-W1-02")
    field("Автор", "E1 / E2 / S")
    field("Дата", "ГГГГ-ММ-ДД")
    field("Модуль", "Топливо / Нагрузки / …")
    field("Версия БД", "СиПР 2026–2031 (рабочая)")
    field("Годы / срез", "например 2028, ОЭС Центр")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Что сделано (5–10 строк)").bold = True
    doc.add_paragraph("1. …")
    doc.add_paragraph("2. …")
    doc.add_paragraph("3. …")

    p = doc.add_paragraph()
    p.add_run("Результат (имя файла)").bold = True
    doc.add_paragraph("…")

    p = doc.add_paragraph()
    p.add_run("Ограничения / допущения").bold = True
    doc.add_paragraph("…")

    p = doc.add_paragraph()
    p.add_run("Блокеры для L").bold = True
    doc.add_paragraph("- нет / …")

    p = doc.add_paragraph()
    p.add_run("Баги (ID из BUGS.xlsx)").bold = True
    doc.add_paragraph("- нет / BUG-2026-…")

    p = doc.add_paragraph()
    p.add_run("DoD").bold = True
    for t in [
        "файл в папке недели",
        "версия БД и годы указаны",
        "блокеры перечислены",
        "баги в BUGS",
        "README заполнен",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    path = ROOT / "00_README_задачи.docx"
    doc.save(path)
    return path


def make_bugs() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "BUGS"
    headers = [
        "ID",
        "Дата",
        "Автор",
        "Модуль",
        "Версия БД",
        "Роль",
        "Шаги",
        "Ожидалось",
        "Факт",
        "Скрин",
        "Критичность",
        "Статус",
        "Комментарий L",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            "BUG-2026-001",
            date.today().isoformat(),
            "E1",
            "Топливо",
            "СиПР 2026–2031 (рабочая)",
            "редактор топлива",
            "1. Открыть станцию X\n2. Нажать «Коэфф»\n3. Сравнить поле Y",
            "Значение = 0.85",
            "Значение = 0.00",
            r"screens\bug001.png",
            "высокий",
            "новый",
            "",
        ],
    )
    blank_rows(ws, 3, 40, len(headers))
    autosize(ws, [14, 12, 8, 14, 28, 16, 36, 22, 22, 24, 12, 14, 24])
    add_list_validation(ws, "C", '"E1,E2,S,L"')
    add_list_validation(ws, "K", '"блокер,высокий,средний,низкий"')
    add_list_validation(ws, "L", '"новый,подтверждён,исправлен,проверен,отклонён"')
    add_meta_sheet(
        wb,
        "Реестр багов BUGS.xlsx",
        [
            "1. Каждому багу — уникальный ID: BUG-YYYY-NNN.",
            "2. Критичность «блокер» — мешает сдаче этапа; сообщите L в день обнаружения.",
            "3. После исправления L ставит «исправлен»; автор обязан сделать ретест и «проверен».",
            "4. Скрин кладите в подпапку screens/ рядом с этим файлом.",
        ],
        [(h, "см. пример") for h in headers],
    )
    path = ROOT / "01_BUGS.xlsx"
    wb.save(path)
    return path


def make_control_objects() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Контрольные объекты"
    headers = [
        "№",
        "Станция / объект",
        "ОЭС",
        "Тип (ТЭС/АЭС/ГЭС/…)",
        "Группа оборудования",
        "Годы с данными",
        "Источник эталона (Access/СиПР/файл)",
        "Путь/ссылка",
        "Статус пригодности",
        "Причина непригодности",
        "Приоритет сверки",
        "Комментарий",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            1,
            "ТЭС Примерная",
            "Центр",
            "ТЭС",
            "К-300",
            "2024–2031",
            "Access топливо 2025",
            r"Z:\...\etalon.mdb",
            "пригоден",
            "",
            "высокий",
            "",
        ],
    )
    blank_rows(ws, 3, 20, len(headers))
    autosize(ws, [4, 28, 12, 14, 18, 14, 28, 28, 14, 24, 12, 20])
    add_list_validation(ws, "I", '"пригоден,частично,непригоден"', end=50)
    add_meta_sheet(
        wb,
        "Реестр контрольных объектов",
        [
            "Выберите 8–12 объектов разных типов и ОЭС.",
            "Статус «пригоден» — есть эталон и полнота данных для сверки.",
            "Непригодные оставьте в списке с причиной (это тоже результат).",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "02_Реестр_контрольных_объектов.xlsx"
    wb.save(path)
    return path


def make_fuel_compare() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сверка полей"
    headers = [
        "Объект",
        "ОЭС",
        "Год",
        "Шаг (Коэфф/Распред/Топливо)",
        "Поле / показатель",
        "Ед. изм.",
        "АРМ",
        "Access/эталон",
        "Δ",
        "Допуск",
        "Класс Δ (данные/методика/баг ПО)",
        "Комментарий",
        "BUG ID",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        ["ТЭС Примерная", "Центр", 2028, "Коэфф", "k топлива", "—", 0.85, 0.85, 0, 0.01, "ок", "", ""],
    )
    blank_rows(ws, 3, 50, len(headers))
    autosize(ws, [22, 10, 8, 18, 28, 10, 12, 14, 10, 10, 22, 22, 12])
    add_list_validation(ws, "K", '"данные,методика,баг ПО,ок,не проверялось"', end=200)

    ws2 = wb.create_sheet("Свод расхождений")
    h2 = ["Тип расхождения", "Кол-во", "Типовой пример", "Нужен код L?", "Приоритет"]
    ws2.append(h2)
    style_header(ws2, len(h2))
    blank_rows(ws2, 2, 15, len(h2))
    autosize(ws2, [22, 10, 40, 12, 12])

    add_meta_sheet(
        wb,
        "Сверка топлива АРМ ↔ эталон",
        [
            "Прогон: Коэфф → Распред → Топливо.",
            "Каждое поле — отдельная строка. Классифицируйте Δ.",
            "Блокеры для кода — только класс «баг ПО».",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "03_Сверка_топливо.xlsx"
    wb.save(path)
    return path


def make_loads_compare() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сверка нагрузок"
    headers = [
        "Версия БД",
        "Год",
        "Разрез (ОЭС/ФО/ЭЗ)",
        "Имя разреза",
        "Показатель",
        "Ед.",
        "АРМ",
        "Эталон",
        "Δ",
        "Допуск",
        "Статус",
        "Комментарий",
        "BUG ID",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            "СиПР 2026–2031",
            2028,
            "ОЭС",
            "Центр",
            "Совмещённый максимум",
            "МВт",
            12500,
            12510,
            -10,
            50,
            "ок",
            "",
            "",
        ],
    )
    blank_rows(ws, 3, 40, len(headers))
    autosize(ws, [20, 8, 14, 16, 28, 8, 12, 12, 10, 10, 12, 24, 12])

    add_meta_sheet(
        wb,
        "Сверка нагрузок",
        [
            "Проверьте максимумы, коэффициенты совмещения, периметры.",
            "Сценарий повтора сверки опишите на листе «Сценарий» (30–40 мин).",
        ],
        [(h, "") for h in headers],
    )

    ws3 = wb.create_sheet("Сценарий")
    ws3["A1"] = "Сценарий повтора сверки (30–40 мин)"
    ws3["A1"].font = Font(bold=True, size=12, color="1F4E79")
    for i, t in enumerate(
        ["Предусловия:", "Шаги:", "Ожидаемый результат:", "Фактический результат:", "Замечания:"],
        3,
    ):
        ws3.cell(i, 1, t).font = Font(bold=True)
    ws3.column_dimensions["A"].width = 80

    path = ROOT / "04_Сверка_нагрузки.xlsx"
    wb.save(path)
    return path


def make_chain() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Цепочка"
    headers = [
        "Версия БД",
        "Год",
        "Разрез",
        "Модуль-источник",
        "Показатель",
        "Значение",
        "Модуль-приёмник",
        "Ожидалось у приёмника",
        "Факт у приёмника",
        "Δ",
        "Источник Δ",
        "BUG ID",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            "СиПР 2026–2031",
            2028,
            "ОЭС Центр",
            "Нагрузки",
            "Максимум, МВт",
            100,
            "Спрос",
            100,
            98,
            2,
            "округление/данные",
            "",
        ],
    )
    blank_rows(ws, 3, 40, len(headers))
    autosize(ws, [20, 8, 14, 16, 28, 12, 16, 18, 16, 10, 20, 12])
    add_meta_sheet(
        wb,
        "Цепочка Нагрузки → Спрос → Балансы",
        [
            "На 1–2 срезах покажите, сходятся ли итоги между модулями.",
            "Для каждого Δ укажите источник (данные / методика / баг / округление).",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "05_Цепочка_нагрузки_спрос_баланс.xlsx"
    wb.save(path)
    return path


def make_db_fill() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Наполненность"
    headers = [
        "Модуль",
        "Таблица / экран",
        "Ключевое поле",
        "Всего записей",
        "Заполнено",
        "Пусто",
        "% заполнения",
        "Критичность дыры",
        "Комментарий",
        "Дата проверки",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            "Генерация",
            "Карточка станции",
            "external_code",
            1200,
            1100,
            100,
            91.7,
            "средний",
            "нет кода у старых",
            "2026-09-16",
        ],
    )
    blank_rows(ws, 3, 40, len(headers))
    autosize(ws, [14, 24, 22, 12, 12, 10, 12, 14, 28, 14])

    ws2 = wb.create_sheet("Top-10 дыр")
    ws2.append(["№", "Модуль", "Проблема", "Влияние", "Кому эскалировать", "Статус"])
    style_header(ws2, 6)
    blank_rows(ws2, 2, 10, 6)
    autosize(ws2, [4, 14, 40, 24, 16, 12])

    add_meta_sheet(
        wb,
        "Таблица наполненности БД",
        [
            "Считайте % по ключевым полям, не по всем колонкам подряд.",
            "Top-10 дыр перенесите на отдельный лист.",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "06_Наполненность_БД.xlsx"
    wb.save(path)
    return path


def make_import_log() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Лог импортов"
    headers = [
        "Дата",
        "Автор",
        "Модуль",
        "Имя файла",
        "Версия БД",
        "Результат (успех/ошибка)",
        "Текст ошибки / предупреждение",
        "Рекомендация",
        "BUG ID",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            date.today().isoformat(),
            "S",
            "Топливо",
            "fuel_import.xlsx",
            "СиПР 2026–2031",
            "ошибка",
            "строка 15: пустой код станции",
            "проверить справочник",
            "BUG-2026-002",
        ],
    )
    blank_rows(ws, 3, 40, len(headers))
    autosize(ws, [12, 8, 14, 28, 22, 16, 40, 28, 12])
    add_list_validation(ws, "F", '"успех,ошибка,частично"')
    add_meta_sheet(
        wb,
        "Лог прогона импортов Excel",
        [
            "На каждый файл — отдельная строка.",
            "Даже успешный импорт фиксируйте (для регресса).",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "07_Лог_импортов.xlsx"
    wb.save(path)
    return path


def make_acceptance() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Чек-лист"
    headers = [
        "№",
        "Модуль",
        "Роль",
        "Сценарий",
        "Шаги",
        "Ожидаемый результат",
        "Пройдено (да/нет/частично)",
        "BUG ID",
        "Блокер сдачи?",
        "Комментарий",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            1,
            "Топливо",
            "редактор",
            "Прогон трёх кнопок",
            "Коэфф→Распред→Топливо",
            "Расчёт без ошибки, итоги ненулевые",
            "частично",
            "BUG-2026-001",
            "да",
            "падение на Распред",
        ],
    )
    blank_rows(ws, 3, 40, len(headers))
    autosize(ws, [4, 14, 14, 28, 36, 28, 14, 12, 12, 24])
    add_list_validation(ws, "G", '"да,нет,частично"')
    add_list_validation(ws, "I", '"да,нет"')

    ws2 = wb.create_sheet("Сводка")
    ws2["A1"] = "Всего сценариев"
    ws2["A2"] = "Пройдено"
    ws2["A3"] = "% пройдено"
    ws2["A4"] = "Блокеры сдачи"
    for row in ws2["A1:A4"]:
        row[0].font = Font(bold=True)

    add_meta_sheet(
        wb,
        "Чек-лист приёмки модуля",
        [
            "Скопируйте файл и переименуйте под модуль: acceptance_fuel.xlsx и т.п.",
            "Сценарии чтение и редактор — отдельно.",
            "Блокер сдачи = да только если без исправления этап сдать нельзя.",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "08_Чек-лист_приёмки.xlsx"
    wb.save(path)
    return path


def make_output_catalog() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Каталог форм"
    headers = [
        "Код формы",
        "Название",
        "Блок (Генерация/Нагрузки/…)",
        "Источник данных (модули)",
        "Ключевые поля",
        "Ед. изм.",
        "Статус (есть/частично/нет)",
        "Приоритет (2026/2027)",
        "Макет (путь)",
        "Владелец предметки",
        "Комментарий",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            "OUT-GEN-01",
            "Установленная мощность по ОЭС",
            "Генерация",
            "generation",
            "ОЭС, год, МВт",
            "МВт",
            "частично",
            "2026",
            r"mockups\OUT-GEN-01.xlsx",
            "E1",
            "",
        ],
    )
    blank_rows(ws, 3, 50, len(headers))
    autosize(ws, [12, 32, 18, 24, 28, 10, 14, 12, 28, 12, 20])
    add_list_validation(ws, "G", '"есть,частично,нет"', end=200)
    add_list_validation(ws, "H", '"2026,2027"', end=200)
    add_meta_sheet(
        wb,
        "Каталог выходных документов",
        [
            "Одна строка = одна форма/таблица комплекта.",
            "Статус «частично» — экран/выгрузка есть, но набор полей неполный.",
            "Приоритет: обязательно для сдачи текущего этапа vs задел.",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "09_Каталог_выходных_документов.xlsx"
    wb.save(path)
    return path


def make_captions() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Подписи скринов"
    headers = [
        "№",
        "Имя файла",
        "Модуль",
        "Что на экране",
        "Версия БД",
        "Роль",
        "Подпись для отчёта",
        "Путь",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            "01",
            "01_home.png",
            "Общее",
            "Главная страница",
            "СиПР 2026–2031",
            "admin",
            "Главная страница АРМ «Генеральная схема»",
            r"screens\01_home.png",
        ],
    )
    blank_rows(ws, 3, 40, len(headers))
    autosize(ws, [6, 22, 14, 28, 22, 12, 40, 28])
    add_meta_sheet(
        wb,
        "Скрин-пакет и подписи",
        [
            "Имена файлов — с номером по порядку отчёта.",
            "Подпись — одна фраза, без лишних деталей.",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "10_Скрин_пакет_подписи.xlsx"
    wb.save(path)
    return path


def make_equip_mapping() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Маппинг"
    headers = [
        "Группа оборудования (Access/источник)",
        "Код/имя в источнике",
        "Station в АРМ",
        "ID станции",
        "ОЭС",
        "Статус (сопоставлено/не сопоставлено/спорно)",
        "Правило сопоставления",
        "Приоритет",
        "Комментарий",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            "К-300 блок 1",
            "ST-001",
            "ТЭС Примерная",
            29649,
            "Центр",
            "сопоставлено",
            "по external_code",
            "высокий",
            "",
        ],
    )
    blank_rows(ws, 3, 40, len(headers))
    autosize(ws, [28, 18, 24, 12, 10, 22, 24, 12, 20])
    add_list_validation(ws, "F", '"сопоставлено,не сопоставлено,спорно"', end=200)
    add_meta_sheet(
        wb,
        "Маппинг групп оборудования ↔ Station",
        [
            "Сначала правила, потом массовое сопоставление.",
            "Спорные случаи — отдельным приоритетом для L/предметки.",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "11_Маппинг_групп_оборудования.xlsx"
    wb.save(path)
    return path


def make_control_slices() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Срезы"
    headers = [
        "№",
        "Версия БД",
        "Год",
        "Разрез",
        "Имя",
        "Эталонный файл",
        "Согласован с L (да/нет)",
        "Назначение сверки",
        "Комментарий",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    add_example_row(
        ws,
        2,
        [
            1,
            "СиПР 2026–2031",
            2028,
            "ОЭС",
            "Центр",
            r"etalons\oes_center_2028.xlsx",
            "да",
            "Нагрузки→Спрос→Баланс",
            "",
        ],
    )
    blank_rows(ws, 3, 20, len(headers))
    autosize(ws, [4, 20, 8, 10, 16, 36, 14, 28, 20])
    add_meta_sheet(
        wb,
        "Контрольные срезы потребления и балансов",
        [
            "Цель: 2 версии × 2 года × 2 разреза (ОЭС + ФО) — или согласованный минимум с L.",
            "Пока срез не согласован — массовые сверки не начинать.",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "12_Контрольные_срезы.xlsx"
    wb.save(path)
    return path


def make_etalons_opis() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Опись эталонов"
    headers = [
        "№",
        "Тип (Access/СиПР/ГИС ТЭК/Excel)",
        "Название",
        "Период/годы",
        "ОЭС/территория",
        "Путь к файлу",
        "Хеш/дата файла",
        "Для какой сверки",
        "Ответственный",
        "Комментарий",
    ]
    ws.append(headers)
    style_header(ws, len(headers))
    blank_rows(ws, 2, 30, len(headers))
    autosize(ws, [4, 18, 28, 14, 14, 40, 16, 22, 12, 20])
    add_meta_sheet(
        wb,
        "Архив эталонов — опись",
        [
            "Файлы эталонов кладите в etalons/ рядом с описью.",
            "Без описи архив не считается сданным.",
        ],
        [(h, "") for h in headers],
    )
    path = ROOT / "18_Опись_эталонов.xlsx"
    wb.save(path)
    return path


def make_index_docx() -> Path:
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("Шаблоны для заполнения работниками")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    doc.add_paragraph(f"Папка: {ROOT}")
    doc.add_paragraph(f"Дата: {date.today().isoformat()}")
    doc.add_paragraph(
        "Правило: копируйте шаблон в папку недели работ, переименовывайте, заполняйте. "
        "Исходники в 00_Шаблоны_для_заполнения не затирайте."
    )
    doc.add_paragraph()

    h = doc.add_paragraph()
    h.add_run("Состав комплекта").bold = True

    mapping = [
        ("00_README_задачи.docx", "все", "К каждой задаче — короткий README"),
        ("01_BUGS.xlsx", "все", "Общий реестр багов"),
        ("02_Реестр_контрольных_объектов.xlsx", "E1", "Выбор станций для сверки"),
        ("03_Сверка_топливо.xlsx", "E1", "АРМ ↔ эталон по топливу"),
        ("04_Сверка_нагрузки.xlsx", "E2", "Сверка нагрузок + сценарий"),
        ("05_Цепочка_нагрузки_спрос_баланс.xlsx", "E2", "Сквозная цепочка"),
        ("06_Наполненность_БД.xlsx", "S", "% заполнения и top-дыры"),
        ("07_Лог_импортов.xlsx", "S", "Прогон импортов Excel"),
        ("08_Чек-лист_приёмки.xlsx", "E1/E2", "Приёмка модуля"),
        ("09_Каталог_выходных_документов.xlsx", "E2", "Реестр форм"),
        ("10_Скрин_пакет_подписи.xlsx", "S", "Подписи к скриншотам"),
        ("11_Маппинг_групп_оборудования.xlsx", "E1", "Группы ↔ Station"),
        ("12_Контрольные_срезы.xlsx", "E2", "Утверждение срезов"),
        ("13_Методика_экологии_v01.docx", "E1", "Методика экологии"),
        ("14_Входы_рациональной_структуры_v01.docx", "E1", "Входы РС"),
        ("15_Обмен_с_АРМ_Экономика_as_is_to_be.docx", "E2", "Обмен с Экономикой"),
        ("16_Акт_пилота_Этапа1.docx", "E2", "Акт пилота"),
        ("17_Инструкция_модуля_ШАБЛОН.docx", "S", "Пользовательская инструкция"),
        ("18_Опись_эталонов.xlsx", "S", "Архив эталонов"),
    ]

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Файл"
    hdr[1].text = "Для кого"
    hdr[2].text = "Назначение"
    for fname, who, purpose in mapping:
        row = table.add_row().cells
        row[0].text = fname
        row[1].text = who
        row[2].text = purpose

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Куда сдавать заполненное").bold = True
    doc.add_paragraph("Z:\\НИО-10\\АРМ ГС\\Работы сотрудников\\YYYY-WW\\")
    doc.add_paragraph(
        "Рядом: результат + README + при необходимости screens/. "
        "Баги — в общий 01_BUGS.xlsx недели/проекта."
    )

    # silence unused import warning locally if any
    _ = WD_ALIGN_PARAGRAPH

    path = ROOT / "README_комплект_шаблонов.docx"
    doc.save(path)
    return path


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    (ROOT / "screens").mkdir(exist_ok=True)
    (ROOT / "etalons").mkdir(exist_ok=True)
    (ROOT / "mockups").mkdir(exist_ok=True)

    creators = [
        make_readme_task,
        make_bugs,
        make_control_objects,
        make_fuel_compare,
        make_loads_compare,
        make_chain,
        make_db_fill,
        make_import_log,
        make_acceptance,
        make_output_catalog,
        make_captions,
        make_equip_mapping,
        make_control_slices,
        lambda: make_guide_docx(
            "13_Методика_экологии_v01.docx",
            "Методика экологии v0.1 (заготовка)",
            [
                ("1. Назначение", ["Расчёт показателей ПГ и ЗВ для …", "Граница модуля: что входит / не входит"]),
                ("2. Входные данные", ["Из топлива: …", "Вручную: …", "Из ГИС ТЭК: …", "Единицы измерения: …"]),
                ("3. Формулы ПГ", ["Формула: …", "Численный пример: …"]),
                ("4. Формулы ЗВ", ["Формула: …", "Численный пример: …"]),
                ("5. Выходные показатели", ["Список полей и единиц", "Связь с выходными формами"]),
                ("6. Открытые вопросы для L", ["…"]),
            ],
        ),
        lambda: make_guide_docx(
            "14_Входы_рациональной_структуры_v01.docx",
            "Входы для «Рациональной структуры» v0.1 (заготовка)",
            [
                ("1. Цель", ["Какие результаты расчёта должны получиться на выходе"]),
                ("2. Показатели из топлива", ["…", "единицы"]),
                ("3. Показатели из генерации", ["…"]),
                ("4. Показатели из электроёмкости / экономики", ["…"]),
                ("5. Ручной ввод", ["Что нельзя взять автоматически"]),
                ("6. Ограничения и допущения", ["…"]),
            ],
        ),
        lambda: make_guide_docx(
            "15_Обмен_с_АРМ_Экономика_as_is_to_be.docx",
            "Обмен с АРМ «Экономика»: as-is / to-be (заготовка)",
            [
                ("1. As-is (сейчас)", ["Прямой канал: какие наборы, направление, формат", "Что уже работает / не работает"]),
                ("2. To-be (нужно)", ["Обратный канал", "Квитанции / подтверждения", "Конфликты версий и разрешения"]),
                ("3. Приоритеты 2026 vs 2027", ["…"]),
                ("4. Риски сдачи", ["…"]),
            ],
        ),
        lambda: make_guide_docx(
            "16_Акт_пилота_Этапа1.docx",
            "Внутренний акт пилота Этапа 1 (заготовка)",
            [
                ("1. Состав пилота", ["Версия БД", "Срезы / объекты", "Участники"]),
                ("2. Сценарии", ["…"]),
                ("3. Результаты", ["Пройдено / не пройдено"]),
                ("4. Top-баги", ["BUG-… — кратко"]),
                ("5. Риск сдачи", ["низкий / средний / высокий — почему"]),
                ("6. Рекомендации", ["…"]),
            ],
        ),
        lambda: make_guide_docx(
            "17_Инструкция_модуля_ШАБЛОН.docx",
            "Инструкция пользователя (шаблон)",
            [
                ("1. Назначение модуля", ["…"]),
                ("2. Роли и права", ["чтение / редактор / admin"]),
                ("3. Как выбрать версию БД", ["…"]),
                ("4. Основные экраны", ["Список", "Карточка", "Итоги"]),
                ("5. Типовой сценарий работы", ["Шаг 1…"]),
                ("6. Импорт / экспорт", ["…"]),
                ("7. Типовые ошибки и что делать", ["…"]),
                ("8. Куда писать замечания", ["BUGS.xlsx"]),
            ],
        ),
        make_etalons_opis,
        make_index_docx,
    ]

    for fn in creators:
        path = fn()
        print("OK", path.name)

    print("DIR", ROOT)
    print("COUNT", len(list(ROOT.iterdir())))


if __name__ == "__main__":
    main()
