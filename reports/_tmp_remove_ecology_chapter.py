# -*- coding: utf-8 -*-
"""Убрать раздел 5.8 «Экология» из DOCX (2 этап) без полной пересборки."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document

DOCX = Path(
    r"Z:\НИО-10\АРМ ГС\02. Отчеты НИР\2026 год"
) / "Описание_программного_комплекса_АРМ_Генеральная_схема.docx"
BACKUP = DOCX.with_name(DOCX.stem + "_before_remove_ecology.docx")
LOCAL = Path(r"c:\arm_gs\reports") / DOCX.name


def _set_text(para, text: str) -> None:
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def delete_paragraph(para) -> None:
    el = para._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def renumber_section_refs(text: str) -> str:
    """5.10→5.9, 5.9→5.8; не трогать прочие номера."""
    text = re.sub(r"\b5\.10\b", "§59§", text)
    text = re.sub(r"\b5\.9\b", "5.8", text)
    text = text.replace("§59§", "5.9")
    return text


def main() -> None:
    if not DOCX.exists():
        raise FileNotFoundError(DOCX)
    if not BACKUP.exists():
        shutil.copy2(DOCX, BACKUP)
        print("backup", BACKUP)
    else:
        # если уже патчили — работаем от свежего бэкапа
        shutil.copy2(BACKUP, DOCX)
        print("restored", BACKUP)

    doc = Document(str(DOCX))

    # --- удалить главу 5.8 Экология ---
    start = end = None
    for i, para in enumerate(doc.paragraphs):
        t = (para.text or "").strip()
        style = para.style.name if para.style else ""
        is_h = "Заголовок" in style or style.startswith("Heading")
        if is_h and t.startswith('5.8 Модуль «Экология»'):
            start = i
        if start is not None and is_h and t.startswith('5.9 Модуль'):
            end = i
            break
    if start is None or end is None:
        raise RuntimeError(f"ecology bounds: start={start} end={end}")

    for i in range(end - 1, start - 1, -1):
        delete_paragraph(doc.paragraphs[i])
    print(f"removed body chapter paragraphs [{start}:{end})")

    # --- удалить строку TOC ---
    for para in list(doc.paragraphs):
        t = (para.text or "").strip()
        if t == '5.8 Модуль «Экология»' or t.startswith('5.8 Модуль «Экология»'):
            delete_paragraph(para)
            print("removed TOC ecology")

    # --- перенумеровать 5.9/5.10 ---
    for para in doc.paragraphs:
        t = para.text or ""
        if not t:
            continue
        if "5.9" in t or "5.10" in t:
            new = renumber_section_refs(t)
            if new != t:
                _set_text(para, new)

    # --- список модулей ---
    for para in list(doc.paragraphs):
        t = para.text or ""
        if 'модуль «Экология»' in t and "ПГ/ЗВ" in t:
            delete_paragraph(para)
            print("removed modules-list ecology bullet")

    # --- ключевые слова ---
    for para in doc.paragraphs:
        t = para.text or ""
        if "ЭКОЛОГИЯ" in t and "КЛЮЧЕВ" not in t.upper() and "ОПИСАНИЕ ПРОГРАММНОГО" in t:
            _set_text(para, t.replace(", ЭКОЛОГИЯ", "").replace("ЭКОЛОГИЯ, ", ""))
            print("keywords cleaned")
        elif t.startswith("АРМ ГЕНЕРАЛЬНАЯ СХЕМА") and "ЭКОЛОГИЯ" in t:
            _set_text(para, t.replace(", ЭКОЛОГИЯ", "").replace("ЭКОЛОГИЯ, ", ""))
            print("keywords cleaned")

    # --- цепочки взаимодействия / заключение ---
    for para in doc.paragraphs:
        t = para.text or ""
        new = t
        new = new.replace("Топливо/Экология", "Топливо")
        new = new.replace("«Топливо», «Экология», «Электроемкость»", "«Топливо», «Электроемкость»")
        new = new.replace(", «Экология»", "")
        new = new.replace("«Экология», ", "")
        if new != t:
            _set_text(para, new)
            print("prose:", t[:50], "->", new[:50])

    # --- смягчить упоминания экологических расчетов в генерации (не про модуль) ---
    # оставляем «экологических» как предметную связь, это ок

    # --- таблица ролей: удалить строку ecology ---
    for table in doc.tables:
        for ri, row in enumerate(list(table.rows)):
            cell0 = row.cells[0].text.strip() if row.cells else ""
            if "ecology" in cell0.lower() or (
                len(row.cells) > 1 and row.cells[1].text.strip() == "Экология"
            ):
                row._tr.getparent().remove(row._tr)
                print("removed roles row", cell0[:40])
                break

    # --- индекс / FR: убрать строки про экологию ---
    for para in list(doc.paragraphs):
        t = (para.text or "").strip()
        if t.startswith("экология —") or t.startswith("экология -"):
            delete_paragraph(para)
            print("removed index ecology")
        elif "FR-024" in t and "Экология" in t:
            delete_paragraph(para)
            print("removed FR-024/TC ecology")

    # в таблицах FR/TC
    for table in doc.tables:
        remove_idxs = []
        for ri, row in enumerate(table.rows):
            joined = " | ".join(c.text for c in row.cells)
            if "FR-024" in joined or (
                "Экология:" in joined and "ГИС ТЭК" in joined
            ):
                remove_idxs.append(ri)
            if "экология — раздел" in joined.lower():
                remove_idxs.append(ri)
        for ri in reversed(remove_idxs):
            row = table.rows[ri]
            row._tr.getparent().remove(row._tr)
            print("removed table row", ri)

    # --- scope: если есть фраза со списком модулей включая Экологию ---
    for para in doc.paragraphs:
        t = para.text or ""
        if "«Топливо», «Экология»" in t:
            _set_text(para, t.replace("«Топливо», «Экология», ", "«Топливо», "))
            print("scope cleaned")

    # добавить в ограничения, если есть маркированный список LIMIT
    inserted = False
    for i, para in enumerate(doc.paragraphs):
        t = (para.text or "").strip()
        if t.startswith("электроемкость по субъектам"):
            # вставим после этого пункта упоминание экологии как 2 этапа
            # проще заменить/дописать рядом через новый run в соседнем пустом — skip if complex
            pass
        if (
            not inserted
            and "не входит в описание базовой реализации" in t
            and "информационной безопасности" in t
        ):
            # prepend ecology note into this bullet area: add new paragraph before
            new_p = para.insert_paragraph_before(
                "модуль «Экология» (отчетность ГИС ТЭК, расчеты по "
                "действующим и новым ТЭС) относится ко второму этапу "
                "работ и не входит в состав описания текущего этапа;"
            )
            # try keep style
            try:
                new_p.style = para.style
            except Exception:
                pass
            inserted = True
            print("added stage-2 limit bullet")

    if not inserted:
        # fallback: append near conclusion modules sentence already cleaned
        print("WARN: stage-2 limit bullet not inserted")

    doc.save(str(DOCX))
    shutil.copy2(DOCX, LOCAL)
    print("saved", DOCX)
    print("saved", LOCAL)

    # verify
    doc2 = Document(str(DOCX))
    eco = [
        (i, (p.text or "")[:100])
        for i, p in enumerate(doc2.paragraphs)
        if "эколог" in (p.text or "").lower() or "Эколог" in (p.text or "")
    ]
    print("remaining ecology mentions:", len(eco))
    for i, t in eco[:20]:
        print(f"  {i}: {t}")
    heads = [
        (p.text or "").strip()
        for p in doc2.paragraphs
        if (p.text or "").strip().startswith("5.")
        and (
            "Заголовок" in (p.style.name if p.style else "")
            or (p.style.name if p.style else "").startswith("Heading")
            or (p.text or "").strip()
            in {
                '5.8 Модуль «Электроемкость» и контур «Экономика»',
                "5.9 Передача данных и подготовка выходных материалов",
            }
        )
    ]
    print("sample heads after:")
    for h in heads:
        if h.startswith("5.7") or h.startswith("5.8") or h.startswith("5.9"):
            print(" ", h[:80])


if __name__ == "__main__":
    main()
