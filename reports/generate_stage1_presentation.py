# -*- coding: utf-8 -*-
"""Презентация по отчету Этапа 1 АРМ «Генеральная схема» (~15 слайдов)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Pt, Emu


ROOT = Path(r"c:\arm_gs\reports")
SHOTS = ROOT / "screenshots"
OUT_PATHS = [
    ROOT / "Презентация_Этап1_АРМ_ГС-2026.pptx",
    Path(r"Z:\НИО-10\АРМ ГС\Отчеты НИР\2026 год")
    / "Презентация_Этап1_АРМ_ГС-2026.pptx",
    Path(r"Z:\НИО-10\АРМ ГС\ТЗ на 2026 год")
    / "Презентация_Этап1_АРМ_ГС-2026.pptx",
]

# Цвета в духе интерфейса АРМ (синий / белый), без «фиолетового ИИ»
NAVY = RGBColor(0x1A, 0x3A, 0x5C)
ACCENT = RGBColor(0x2F, 0x6F, 0xA8)
LIGHT = RGBColor(0xE8, 0xF0, 0xF7)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _set_run(run, *, size=18, bold=False, color=DARK, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_textbox(slide, left, top, width, height, text, *, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color)
    return box


def _add_bullets(slide, left, top, width, height, items, *, size=16):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = item
        _set_run(run, size=size, color=DARK)
    return box


def _header_bar(slide, title: str, subtitle: str | None = None):
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Cm(0), Cm(0), Cm(33.87), Cm(2.4)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    _add_textbox(
        slide, Cm(0.8), Cm(0.35), Cm(32), Cm(1.2), title,
        size=26, bold=True, color=WHITE,
    )
    if subtitle:
        _add_textbox(
            slide, Cm(0.8), Cm(1.4), Cm(32), Cm(0.8), subtitle,
            size=14, color=LIGHT,
        )


def _footer(slide, n: int, total: int = 15):
    _add_textbox(
        slide, Cm(0.8), Cm(18.2), Cm(25), Cm(0.5),
        "АО «НТЦ ЕЭС» · НИО-10 · Этап 1 · 2026",
        size=11, color=MUTED,
    )
    _add_textbox(
        slide, Cm(28), Cm(18.2), Cm(5), Cm(0.5),
        f"{n} / {total}",
        size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def _try_picture(slide, stem: str, left, top, width):
    path = SHOTS / f"{stem}.png"
    if not path.exists():
        return None
    return slide.shapes.add_picture(str(path), left, top, width=width)


def build() -> list[str]:
    prs = Presentation()
    prs.slide_width = Cm(33.867)  # 16:9
    prs.slide_height = Cm(19.05)
    blank = prs.slide_layouts[6]
    total = 15

    # 1. Title
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    _add_textbox(
        s, Cm(1.5), Cm(4.5), Cm(30), Cm(1.2),
        "АО «Научно-технический центр Единой энергетической системы»",
        size=16, color=LIGHT, align=PP_ALIGN.CENTER,
    )
    _add_textbox(
        s, Cm(1.5), Cm(5.8), Cm(30), Cm(1.0),
        "НИО-10 · Отдел генерации и прогнозирования спроса",
        size=15, color=LIGHT, align=PP_ALIGN.CENTER,
    )
    _add_textbox(
        s, Cm(1.5), Cm(7.5), Cm(30), Cm(2.2),
        "АРМ «Генеральная схема»\nЭтап 1 – выполненные работы",
        size=32, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
    )
    _add_textbox(
        s, Cm(2.5), Cm(10.5), Cm(28.5), Cm(2.5),
        "Инициализация БД, подразделы: Справочники, Генерация, Нагрузки, Спрос,\n"
        "Балансы ЭЭ и ЭМ, Топливо, Электроемкость.\n"
        "Выходные документы и связь с АРМ «Экономика».",
        size=16, color=LIGHT, align=PP_ALIGN.CENTER,
    )
    _add_textbox(
        s, Cm(1.5), Cm(15.5), Cm(30), Cm(1.0),
        "Приказ от 12.01.2026 № 680-ОД/2026 · срок этапа: 12.01.2026 – 30.09.2026",
        size=14, color=LIGHT, align=PP_ALIGN.CENTER,
    )
    _add_textbox(s, Cm(28), Cm(18.2), Cm(5), Cm(0.5), "1 / 15", size=11, color=LIGHT, align=PP_ALIGN.RIGHT)

    # 2. Agenda
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Содержание")
    _add_bullets(
        s, Cm(1.2), Cm(3.2), Cm(15), Cm(14),
        [
            "1. Цели и состав Этапа 1",
            "2. Общая архитектура АРМ",
            "3. Инфраструктура: версии, пользователи, журналы",
            "4. Справочники",
            "5. Генерация",
            "6. Нагрузки и Спрос",
            "7. Балансы ЭЭ и ЭМ",
            "8. Топливо",
            "9. Электроемкость и Экономика",
            "10. Выходные документы",
            "11. Итоги и Этап 2",
        ],
        size=18,
    )
    _try_picture(s, "01_home", Cm(17.5), Cm(3.5), Cm(15))
    _footer(s, 2, total)

    # 3. Goals
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Цели Этапа 1", "По Техническому заданию (Приложение 1) и графику (Приложение 3)")
    _add_bullets(
        s, Cm(1.2), Cm(3.2), Cm(31), Cm(14),
        [
            "Инициализация базы данных и общесистемный каркас АРМ",
            "Создание подразделов: Справочники, Генерация, Нагрузки, Спрос,\n   Балансы электроэнергии и мощности, Топливо, Электроемкость",
            "Подготовка комплекта выходных материалов по блокам этапа",
            "Передача данных генерации во внешний контур АРМ «Экономика»",
            "Стек: Python/Flask, PostgreSQL, веб-интерфейс; данные – по схемам и версиям БД",
        ],
        size=18,
    )
    _footer(s, 3, total)

    # 4. Architecture
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Общая архитектура", "Единый вход, модули по блокам, общая версия данных")
    _add_bullets(
        s, Cm(1.2), Cm(3.0), Cm(16), Cm(14),
        [
            "Энергоресурсы и балансы: Генерация, Балансы, Топливо",
            "Потребление: Нагрузки, Спрос",
            "Прогноз и территории: долгосрочный прогноз (электроемкость)",
            "НСИ: Справочники; обмен: передача данных для документов",
            "Администрирование: версии БД, пользователи и роли, журналы",
            "Каждый модуль – своя схема PostgreSQL + общие версии и права",
        ],
        size=17,
    )
    _try_picture(s, "01_home", Cm(17.8), Cm(3.2), Cm(14.8))
    _footer(s, 4, total)

    # 5. Infra
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Инфраструктура", "Версии БД · пользователи и роли · журналирование")
    _add_bullets(
        s, Cm(1.2), Cm(3.0), Cm(15.5), Cm(7),
        [
            "Управление версиями базы: создание, активация, смена в сессии",
            "Пользователи и роли: матрица прав по модулям\n   (чтение / редактор / админ)",
            "Общий журнал «Запись логов» и журналы на карточках/таблицах",
            "Миграции схемы, кэш справочников, авторизация маршрутов",
        ],
        size=16,
    )
    _try_picture(s, "02_db_versions", Cm(1.2), Cm(10.5), Cm(10))
    _try_picture(s, "02c_users_roles", Cm(11.8), Cm(10.5), Cm(10))
    _try_picture(s, "02b_user_action_logs", Cm(22.4), Cm(10.5), Cm(10))
    _footer(s, 5, total)

    # 6. Refdata
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Справочники", "Централизованная НСИ для всех модулей")
    _add_bullets(
        s, Cm(1.2), Cm(3.0), Cm(16), Cm(10),
        [
            "Территории, энергосистемы, топливо, типы станций и оборудования",
            "Годы и горизонты планирования (СиПР / Генеральная схема)",
            "CRUD, поиск, Excel-импорт/экспорт, история изменений",
            "Прикладные модули берут значения из НСИ, без локальных классификаторов",
        ],
        size=17,
    )
    _try_picture(s, "03_refdata", Cm(17.8), Cm(3.2), Cm(14.8))
    _footer(s, 6, total)

    # 7. Generation
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Генерация", "Реестр · изменения мощности · перспективные площадки · итоги")
    _add_bullets(
        s, Cm(1.2), Cm(2.9), Cm(31), Cm(6.5),
        [
            "Реестр и карточки станций/агрегатов – объектная модель генерации",
            "Изменения установленной мощности ЕГО – табличный учет и Excel (в т.ч. СиПР)",
            "Перспективные площадки АЭС/ГЭС/ГАЭС: связь со станцией реестра по external_code",
            "Итоговые суммы по ЕЭС, ТИТЭС, ФО, СЗ, Россия",
            "Обмен: dataset’ы generation_stations / generation_machines → АРМ «Экономика»",
        ],
        size=16,
    )
    _try_picture(s, "04_generation", Cm(1.0), Cm(9.8), Cm(7.8))
    _try_picture(s, "05f_station_power_changes", Cm(9.0), Cm(9.8), Cm(7.8))
    _try_picture(s, "05h_prospective_places_aes", Cm(17.0), Cm(9.8), Cm(7.8))
    _try_picture(s, "05i_generation_totals", Cm(25.0), Cm(9.8), Cm(7.8))
    _footer(s, 7, total)

    # 8. Loads
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Нагрузки", "Спрос на мощность: своды, коэффициенты, ОЗП")
    _add_bullets(
        s, Cm(1.2), Cm(3.0), Cm(16), Cm(8),
        [
            "Параметры спроса на мощность по иерархии уровней",
            "Своды максимумов: ОЭС / ФО / энергозоны",
            "Режим «Сводная таблица» и кнопки показателей",
            "Коэффициенты совмещения и расчетные k",
            "Максимумы ОЗП; Excel и обменный набор load_forecast",
        ],
        size=17,
    )
    _try_picture(s, "07b_pd_summary_oes_all_toggles", Cm(17.5), Cm(3.2), Cm(15))
    _footer(s, 8, total)

    # 9. Demand
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Спрос", "Потребление ЭЭ: своды, ГАЭС на заряд")
    _add_bullets(
        s, Cm(1.2), Cm(3.0), Cm(16), Cm(8),
        [
            "Потребление электрической энергии по ОЭС / ФО / энергозонам",
            "Те же приемы UI, что в нагрузках (сводная таблица, кнопки)",
            "Отдельный раздел «ГАЭС на заряд»",
            "Фоновый Excel-импорт и массовый экспорт",
            "Без отдельного блока коэффициентов (в отличие от нагрузок)",
        ],
        size=17,
    )
    _try_picture(s, "09b_ec_summary_oes_all_toggles", Cm(17.5), Cm(3.2), Cm(15))
    _footer(s, 9, total)

    # 10. Balances
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Балансы электроэнергии и мощности", "Выработка · перетоки · балансы · ВЭС/СЭС")
    _add_bullets(
        s, Cm(1.2), Cm(3.0), Cm(16), Cm(8),
        [
            "Выработка ЭЭ и контрольные экраны",
            "Перетоки ЭЭ",
            "Балансы мощности и энергии по уровням",
            "Специализированный срез ВЭС/СЭС",
            "Опора на генерацию, нагрузки и спрос в одной версии БД",
        ],
        size=17,
    )
    _try_picture(s, "10_energy_balance", Cm(17.5), Cm(3.0), Cm(15))
    _try_picture(s, "12_power_balance", Cm(17.5), Cm(10.5), Cm(15))
    _footer(s, 10, total)

    # 11. Fuel
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Топливо", "Отчетные таблицы · расчет Коэфф → Распред → Топливо · итоги")
    _add_bullets(
        s, Cm(1.2), Cm(2.9), Cm(16), Cm(9),
        [
            "Отчетные сведения: словарь станций/групп, работа ТЭС, затраты, топливо, тарифы",
            "Расчетный экран: сценарий (параметр распределения) и три кнопки",
            "Порядок обязателен: Коэфф → Распред → Топливо (см. Приложение Б отчета)",
            "Структура топливного баланса, ТЭП, цена топлива",
            "Итоговые суммы по крупным территориальным разрезам",
        ],
        size=16,
    )
    _try_picture(s, "13_fuel", Cm(17.5), Cm(3.0), Cm(15))
    _try_picture(s, "14_fuel_calculation", Cm(17.5), Cm(10.5), Cm(15))
    _footer(s, 11, total)

    # 12. EI + Economics
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Электроемкость и Экономика", "Долгосрочный прогноз потребления")
    _add_bullets(
        s, Cm(1.2), Cm(2.9), Cm(16), Cm(8),
        [
            "Хаб «Долгосрочный прогноз потребления»",
            "Электроемкость по ФО (коэффициенты A / Арасч / X, периоды)",
            "Потребление по ВЭД",
            "Модуль «Экономика»: инвестиции, выпуск, население, доходы",
            "Журналы изменений на таблицах; разрез по субъектам РФ – заготовка",
        ],
        size=16,
    )
    _try_picture(s, "15_electrical_intensity", Cm(17.5), Cm(3.0), Cm(15))
    _try_picture(s, "16c_economics_hub", Cm(17.5), Cm(10.5), Cm(15))
    _footer(s, 12, total)

    # 13. Output docs
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Выходные документы", "Excel из модулей · JSON API · связь с АРМ «Экономика»")
    _add_bullets(
        s, Cm(1.2), Cm(3.0), Cm(16), Cm(8),
        [
            "Страница «Передача данных для выходных документов»",
            "Каталог dataset’ов по версии БД (common, gs_sipr, …)",
            "API: /api/config и /api/datasets/<версия>/<набор>",
            "Из генерации наружу: станции и агрегаты для АРМ «Экономика»",
            "Оперативно – Excel-выгрузки прямо из таблиц модулей",
        ],
        size=17,
    )
    _try_picture(s, "17_output_documents", Cm(17.5), Cm(3.2), Cm(15))
    _footer(s, 13, total)

    # 14. Results
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Итоги Этапа 1", "Что предъявляется по результатам работ")
    _add_bullets(
        s, Cm(1.2), Cm(3.2), Cm(31), Cm(14),
        [
            "Поднята инфраструктура: версии БД, роли, журналы",
            "Работают подразделы Этапа 1 с экранными формами и Excel-обменом",
            "Генерация: реестр, изменения мощности, перспективные площадки, итоги",
            "Нагрузки / Спрос / Балансы / Топливо / Электроемкость – в единой версии данных",
            "Подготовлен канал передачи данных генерации в АРМ «Экономика»",
            "Содержание Этапа 1 по ТЗ и графику закрыто",
        ],
        size=18,
    )
    _footer(s, 14, total)

    # 15. Stage 2 / close
    s = prs.slides.add_slide(blank)
    _header_bar(s, "Далее – Этап 2", "По Техническому заданию")
    _add_bullets(
        s, Cm(1.2), Cm(3.5), Cm(31), Cm(8),
        [
            "Подразделы «Экология» и «Рациональная структура»",
            "Расширение комплекта выходных документов",
            "Развитие интеграционных связей",
            "По электроемкости – разрез по субъектам РФ (в UI уже заготовка)",
        ],
        size=20,
    )
    _add_textbox(
        s, Cm(1.2), Cm(12.5), Cm(31), Cm(2),
        "Спасибо за внимание",
        size=28, bold=True, color=NAVY, align=PP_ALIGN.CENTER,
    )
    _add_textbox(
        s, Cm(1.2), Cm(15.0), Cm(31), Cm(1.2),
        "Вопросы и обсуждение",
        size=18, color=MUTED, align=PP_ALIGN.CENTER,
    )
    _footer(s, 15, total)

    saved: list[str] = []
    for path in OUT_PATHS:
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            prs.save(str(path))
            saved.append(str(path))
        except Exception as exc:  # noqa: BLE001
            print(f"WARN: cannot save {path}: {exc}")
    return saved


if __name__ == "__main__":
    paths = build()
    print("SAVED:")
    for p in paths:
        print(p)
