# -*- coding: utf-8 -*-
"""Обновить рисунки «Спрос» и «Экология» в существующем DOCX без полной пересборки."""

from __future__ import annotations

import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm

SHOTS = Path(r"c:\arm_gs\reports\screenshots")
DOCX = Path(
    r"Z:\НИО-10\АРМ ГС\02. Отчеты НИР\2026 год"
) / "Описание_программного_комплекса_АРМ_Генеральная_схема.docx"
BACKUP = DOCX.with_name(DOCX.stem + "_before_demand_ecology_figs.docx")
LOCAL = Path(r"c:\arm_gs\reports") / DOCX.name

DEMAND_MAP = {
    "rId39": "08_energy_consumption.png",
    "rId40": "09_energy_consumption_summary.png",
    "rId41": "09b_ec_summary_oes_all_toggles.png",
    "rId42": "09c_ec_summary_fo_compact.png",
    "rId43": "09d_ec_summary_ez_compact.png",
    "rId44": "09e_ec_gaes_charge.png",
    "rId45": "09f_ec_summary_table_hub.png",
    "rId46": "09g_ec_summary_table_oes.png",
}

ECOLOGY_FIGURES = [
    ("18_ecology", "Стартовая страница модуля «Экология»"),
    ("18b_ecology_gis_tek_pg", "Отчетные данные ГИС ТЭК: парниковые газы (ПГ)"),
    (
        "18c_ecology_gis_tek_zv",
        "Отчетные данные ГИС ТЭК: загрязняющие вещества (ЗВ)",
    ),
    (
        "18d_ecology_analysis_pg",
        "Анализ отчетных данных по парниковым газам",
    ),
    (
        "18e_ecology_operating_tes",
        "Расчет по действующим ТЭС (форма «Объемы»)",
    ),
    ("18f_ecology_new_tes", "Расчет по новым ТЭС (форма «Объемы»)"),
    (
        "18g_ecology_totals",
        "Агрегированные суммы модуля «Экология» (ПГ)",
    ),
]

SHIFT = len(ECOLOGY_FIGURES)


def replace_demand_images(doc: Document) -> None:
    for rid, filename in DEMAND_MAP.items():
        path = SHOTS / filename
        if not path.exists():
            raise FileNotFoundError(path)
        rel = doc.part.rels[rid]
        rel.target_part._blob = path.read_bytes()
        print("replaced", rid, filename, path.stat().st_size)


def _para_text(p_el) -> str:
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def find_ecology_note(doc: Document):
    for para in doc.paragraphs:
        t = para.text or ""
        if "экологических форм не включены" in t:
            return para
        if "основные экранные формы модуля" in t and "ГИС ТЭК" in t:
            return para
    raise RuntimeError("ecology screen-forms paragraph not found")


def figure_number(text: str) -> int | None:
    m = re.match(r"^Рисунок\s+(\d+)\.", (text or "").strip())
    return int(m.group(1)) if m else None


def last_figure_before(doc: Document, before_el) -> int:
    last = 0
    for p_el in doc.element.body.iterchildren(qn("w:p")):
        if p_el is before_el:
            break
        n = figure_number(_para_text(p_el))
        if n is not None:
            last = n
    return last


def first_figure_after(doc: Document, after_el) -> int | None:
    seen = False
    for p_el in doc.element.body.iterchildren(qn("w:p")):
        if p_el is after_el:
            seen = True
            continue
        if not seen:
            continue
        n = figure_number(_para_text(p_el))
        if n is not None:
            return n
    return None


def find_template_drawing_paragraph(doc: Document):
    body = doc.element.body
    for p_el in body.iterchildren(qn("w:p")):
        if p_el.find(".//" + qn("w:drawing")) is not None:
            return p_el
    raise RuntimeError("drawing template paragraph not found")


def find_caption_template(doc: Document):
    for p_el in doc.element.body.iterchildren(qn("w:p")):
        if _para_text(p_el).startswith("Рисунок 35."):
            return p_el
    for p_el in doc.element.body.iterchildren(qn("w:p")):
        if re.match(r"^Рисунок\s+\d+\.", _para_text(p_el).strip()):
            return p_el
    raise RuntimeError("caption template not found")


def insert_ecology_figures(doc: Document, first_new_number: int) -> None:
    note = find_ecology_note(doc)
    note_el = note._element
    for child in list(note_el):
        if child.tag == qn("w:r"):
            note_el.remove(child)
    run = note.add_run(
        "Модуль доступен из блока «Энергоресурсы и балансы» главной "
        "страницы. Ниже приведены основные экранные формы модуля: "
        "стартовая страница, отчетные данные ГИС ТЭК, аналитическая "
        "форма, расчеты по действующим и новым ТЭС, а также "
        "агрегированные итоги."
    )
    run.font.name = "Times New Roman"

    template_draw = find_template_drawing_paragraph(doc)
    caption_template = find_caption_template(doc)

    insert_after = note_el
    for offset, (stem, caption) in enumerate(ECOLOGY_FIGURES):
        path = SHOTS / f"{stem}.png"
        if not path.exists():
            raise FileNotFoundError(path)

        blank = deepcopy(note_el)
        for child in list(blank):
            if child.tag != qn("w:pPr"):
                blank.remove(child)
        insert_after.addnext(blank)
        insert_after = blank

        draw = deepcopy(template_draw)
        for child in list(draw):
            if child.tag != qn("w:pPr"):
                draw.remove(child)
        insert_after.addnext(draw)
        insert_after = draw

        pic_para = None
        for para in doc.paragraphs:
            if para._element is draw:
                pic_para = para
                break
        if pic_para is None:
            raise RuntimeError("picture paragraph wrapper not found")
        pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = pic_para.add_run()
        run.add_picture(str(path), width=Cm(16.0))

        cap = deepcopy(caption_template)
        texts = list(cap.iter(qn("w:t")))
        new_text = f"Рисунок {first_new_number + offset}. {caption}"
        if not texts:
            raise RuntimeError("caption has no text nodes")
        texts[0].text = new_text
        for t in texts[1:]:
            t.text = ""
        insert_after.addnext(cap)
        insert_after = cap
        print("inserted", new_text)


def renumber_figures(doc: Document, from_number: int, shift: int) -> None:
    pattern_cap = re.compile(r"^(Рисунок )(\d+)(\..*)$")
    pattern_ref = re.compile(r"(рисунк(?:е|а|ом|у)?\s+)(\d+)\b", re.I)

    for para in doc.paragraphs:
        text = para.text
        if not text:
            continue
        m = pattern_cap.match(text.strip())
        if m:
            num = int(m.group(2))
            if num >= from_number:
                new = f"{m.group(1)}{num + shift}{m.group(3)}"
                _set_paragraph_text_keep_style(para, new)
            continue

        if pattern_ref.search(text):

            def repl(mo: re.Match) -> str:
                n = int(mo.group(2))
                if n >= from_number:
                    return f"{mo.group(1)}{n + shift}"
                return mo.group(0)

            new_text = pattern_ref.sub(repl, text)
            if new_text != text:
                _set_paragraph_text_keep_style(para, new_text)

        if text.startswith("Документ ") and "рис." in text:

            def repl_ref(mo: re.Match) -> str:
                return f"{int(mo.group(1)) + shift} рис."

            new_text = re.sub(r"(\d+) рис\.", repl_ref, text, count=1)
            if new_text != text:
                _set_paragraph_text_keep_style(para, new_text)
                print("referat", new_text)


def _set_paragraph_text_keep_style(para, text: str) -> None:
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def update_sources_for_future() -> None:
    nir_path = Path(r"c:\arm_gs\reports\generate_stage1_report_nir.py")
    text = nir_path.read_text(encoding="utf-8")
    ecology_block = '''    "ecology": [
        ("18_ecology", "Стартовая страница модуля «Экология»"),
        (
            "18b_ecology_gis_tek_pg",
            "Отчетные данные ГИС ТЭК: парниковые газы (ПГ)",
        ),
        (
            "18c_ecology_gis_tek_zv",
            "Отчетные данные ГИС ТЭК: загрязняющие вещества (ЗВ)",
        ),
        (
            "18d_ecology_analysis_pg",
            "Анализ отчетных данных по парниковым газам",
        ),
        (
            "18e_ecology_operating_tes",
            "Расчет по действующим ТЭС (форма «Объемы»)",
        ),
        ("18f_ecology_new_tes", "Расчет по новым ТЭС (форма «Объемы»)"),
        (
            "18g_ecology_totals",
            "Агрегированные суммы модуля «Экология» (ПГ)",
        ),
    ],
'''
    if '"ecology":' not in text.split("CHAPTER_FIGURES", 1)[1][:5000]:
        text = text.replace('    "balance": [', ecology_block + '    "balance": [', 1)
        nir_path.write_text(text, encoding="utf-8")
        print("updated CHAPTER_FIGURES ecology")

    content_path = Path(r"c:\arm_gs\reports\arm_product_description_content.py")
    content = content_path.read_text(encoding="utf-8")
    old = (
        '            "Экранные формы": [\n'
        '                "Модуль доступен из блока «Энергоресурсы "\n'
        '                "и балансы» главной страницы. В настоящем "\n'
        '                "комплекте иллюстраций отдельные скриншоты "\n'
        '                "экологических форм не включены; общий контекст "\n'
        '                "навигации виден на рисунке главной страницы.",\n'
        "            ],"
    )
    new = (
        '            "Экранные формы": [\n'
        '                "Модуль доступен из блока «Энергоресурсы "\n'
        '                "и балансы» главной страницы. Ниже приведены "\n'
        '                "основные экранные формы: стартовая страница, "\n'
        '                "отчетные данные ГИС ТЭК, аналитическая форма, "\n'
        '                "расчеты по действующим и новым ТЭС и "\n'
        '                "агрегированные итоги.",\n'
        "            ],"
    )
    if old in content:
        content = content.replace(old, new)
        content_path.write_text(content, encoding="utf-8")
        print("updated ecology screen-forms prose")

    content2 = content_path.read_text(encoding="utf-8")
    if (
        '"ecology",\n        "Модуль «Экология»"' in content2
        and '},\n        "",\n    ),\n    (\n        "ei",' in content2
    ):
        content2 = content2.replace(
            '            ],\n        },\n        "",\n    ),\n    (\n        "ei",',
            '            ],\n        },\n        "ecology",\n    ),\n    (\n        "ei",',
            1,
        )
        content_path.write_text(content2, encoding="utf-8")
        print("updated ecology fig_key")


def main() -> None:
    if not BACKUP.exists():
        raise FileNotFoundError(f"нужен бэкап до правки: {BACKUP}")

    # восстановить исходник с правками пользователя (до прошлой вставки)
    shutil.copy2(BACKUP, DOCX)
    print("restored from", BACKUP)

    doc = Document(str(DOCX))
    replace_demand_images(doc)

    note = find_ecology_note(doc)
    last_before = last_figure_before(doc, note._element)
    first_after = first_figure_after(doc, note._element)
    first_new = last_before + 1
    print(
        "ecology slot:",
        f"after fig {last_before},",
        f"insert {first_new}–{first_new + SHIFT - 1},",
        f"shift from {first_after}" if first_after else "no later figures",
    )

    if first_after is not None:
        renumber_figures(doc, from_number=first_after, shift=SHIFT)
    insert_ecology_figures(doc, first_new_number=first_new)

    doc.save(str(DOCX))
    shutil.copy2(DOCX, LOCAL)
    print("saved", DOCX)
    print("saved", LOCAL)
    update_sources_for_future()


if __name__ == "__main__":
    main()
