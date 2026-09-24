# -*- coding: utf-8 -*-
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(r"Z:\НИО-10\АРМ ГС\04. Тестирование") / "Для заполнения"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
YELLOW = PatternFill("solid", fgColor="FFF2CC")
NOTE = PatternFill("solid", fgColor="D6EAF8")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def header_row(ws, headers):
    for col, name in enumerate(headers, 1):
        cell = ws.cell(1, col, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def widths(ws, sizes):
    for i, w in enumerate(sizes, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def paint_row(ws, row, values, fill):
    for col, val in enumerate(values, 1):
        cell = ws.cell(row, col, val)
        cell.fill = fill
        cell.alignment = WRAP
        cell.border = THIN
        cell.font = Font(name="Calibri", size=11)


def empty_rows(ws, start, n, ncols):
    for r in range(start, start + n):
        ws.row_dimensions[r].height = 45
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c, None)
            cell.border = THIN
            cell.alignment = WRAP


def dropdown(ws, col, options, start=3, end=80):
    dv = DataValidation(type="list", formula1=options, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}{start}:{col}{end}")


def make_errors():
    wb = Workbook()
    ws = wb.active
    ws.title = "Ошибки"

    headers = [
        "Номер",
        "Дата",
        "Кто нашёл",
        "Версия БД",
        "Где (модуль, страница)",
        "Что нажимали / что смотрели",
        "Как должно быть",
        "Как получилось",
        "Цифра АРМ",
        "Цифра эталона",
        "Насколько мешает",
        "Статус",
        "Комментарий программиста",
    ]
    header_row(ws, headers)
    widths(ws, [10, 12, 14, 24, 28, 36, 28, 28, 14, 16, 18, 16, 28])

    paint_row(
        ws,
        2,
        [
            "пример, строку можно удалить",
            date.today().strftime("%d.%m.%Y"),
            "Иванов",
            "СиПР 2026–2031 (рабочая)",
            "Топливо, расчёт станции «Примерная»",
            "Открыл станцию, нажал Коэфф, потом Распред, потом Топливо",
            "Расход газа как в Access, около 120",
            "В АРМ 0, расчёт прошёл без ошибки",
            0,
            120,
            "неверно считает",
            "новое",
            "",
        ],
        YELLOW,
    )
    empty_rows(ws, 3, 40, len(headers))

    dropdown(ws, "K", '"не открывается,неверно считает,пропали данные,неудобно,непонятно"')
    dropdown(ws, "L", '"новое,в работе,исправлено,проверить снова,не ошибка"')

    ws.auto_filter.ref = "A1:M42"
    ws.sheet_properties.tabColor = "1F4E79"

    note = wb.create_sheet("Как заполнять", 0)
    note.sheet_properties.tabColor = "2E86C1"
    note["A1"] = "Один файл на всех. Одна ошибка = одна строка."
    note["A1"].font = Font(bold=True, size=14, color="1F4E79", name="Calibri")
    note.merge_cells("A1:B1")
    lines = [
        "",
        "1. Не пишите длинные отчёты. Пишите как коллеге: где были, что нажали, что увидели.",
        "2. В «Версия БД» напишите, какую версию смотрели. Без этого ошибку не повторить.",
        "3. Если ошибка в цифре, заполните «Цифра АРМ» и «Цифра эталона». Без двух цифр расчёт не проверить.",
        "4. Скрин положите в папку «скрины» и напишите имя файла в «Как получилось».",
        "5. «Насколько мешает» выбирайте из списка.",
        "6. «Статус» после сдачи строки не меняйте, пока не попросят проверить снова.",
        "7. Жёлтая строка — пример. Её можно удалить.",
        "",
        "Колонка «Статус» для программиста:",
        "новое = ещё не смотрели",
        "в работе = чинят",
        "исправлено = можно проверять",
        "проверить снова = вы перепроверили и снова плохо",
        "не ошибка = так и задумано или данные кривые",
    ]
    for i, line in enumerate(lines, 3):
        note.cell(i, 1, line).font = Font(name="Calibri", size=12)
        note.cell(i, 1).fill = NOTE
    note.column_dimensions["A"].width = 110
    note.row_dimensions[1].height = 24

    path = ROOT / "1_Ошибки.xlsx"
    wb.save(path)
    return path


def make_numbers():
    wb = Workbook()
    ws = wb.active
    ws.title = "Цифры"

    headers = [
        "Дата",
        "Кто сверял",
        "Версия БД",
        "Модуль",
        "Объект / разрез",
        "Год",
        "Показатель",
        "Цифра в АРМ",
        "Цифра в эталоне",
        "Разница",
        "Эталон откуда (Access, СиПР, файл)",
        "Это ошибка программы?",
        "Номер ошибки из файла 1_Ошибки",
        "Комментарий",
    ]
    header_row(ws, headers)
    widths(ws, [12, 14, 24, 16, 28, 10, 28, 14, 16, 12, 28, 22, 24, 28])

    paint_row(
        ws,
        2,
        [
            date.today().strftime("%d.%m.%Y"),
            "Петров",
            "СиПР 2026–2031 (рабочая)",
            "Нагрузки",
            "ОЭС Центр",
            2028,
            "Совмещённый максимум, МВт",
            12500,
            12510,
            "=H2-I2",
            "файл СиПР, лист максимумы",
            "нет",
            "",
            "мелочь, меньше допуска",
        ],
        YELLOW,
    )
    empty_rows(ws, 3, 40, len(headers))
    for r in range(3, 43):
        ws.cell(r, 10, f'=IF(OR(H{r}="",I{r}=""),"",H{r}-I{r})')

    dropdown(ws, "L", '"да,нет,не знаю"')
    ws.auto_filter.ref = "A1:N42"

    note = wb.create_sheet("Как заполнять", 0)
    note["A1"] = "Сюда пишут только цифры. Не описания экранов."
    note["A1"].font = Font(bold=True, size=14, color="1F4E79", name="Calibri")
    note.merge_cells("A1:B1")
    lines = [
        "",
        "Нужно, чтобы было видно: АРМ считает так же, как эталон, или нет. Версию БД тоже укажите.",
        "Одна строка = один показатель на одном объекте/разрезе за один год.",
        "Если цифры не сходятся и вы считаете, что виновата программа, поставьте «да» и заведите строку в файле 1_Ошибки.",
        "Если виноваты исходные данные, поставьте «нет» и коротко напишите почему.",
        "Жёлтая строка — пример.",
    ]
    for i, line in enumerate(lines, 3):
        note.cell(i, 1, line).font = Font(name="Calibri", size=12)
        note.cell(i, 1).fill = NOTE
    note.column_dimensions["A"].width = 110

    path = ROOT / "2_Сверка_цифр.xlsx"
    wb.save(path)
    return path


def make_howto():
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2)

    p = doc.add_paragraph()
    r = p.add_run("Что заполнять")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    doc.add_paragraph("Два файла. Больше ничего не нужно.")
    doc.add_paragraph("1_Ошибки.xlsx — если что-то сломалось, не открылось или посчиталось не так.")
    doc.add_paragraph("2_Сверка_цифр.xlsx — если сравнивали АРМ с Access / СиПР / своим файлом.")
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.add_run("Как сдать").bold = True
    doc.add_paragraph("Пишите в эти же файлы. Не создавайте свои таблицы «как удобнее».")
    doc.add_paragraph("Если есть скрин, положите его в папку «скрины» рядом и напишите имя файла в строке ошибки.")
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.add_run("Чего не делать").bold = True
    doc.add_paragraph("Не описывать весь модуль целиком. Только конкретный случай.")
    doc.add_paragraph("Не писать «ничего не работает» без шагов и без цифр.")

    path = ROOT / "0_Что_заполнять.docx"
    doc.save(path)
    return path


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    (ROOT / "скрины").mkdir(exist_ok=True)
    for fn in (make_howto, make_errors, make_numbers):
        print("OK", fn().name)


if __name__ == "__main__":
    main()
