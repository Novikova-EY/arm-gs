from io import BytesIO
from decimal import Decimal
from types import SimpleNamespace
from urllib.parse import parse_qs, urlsplit

import pytest
from flask import Flask
from flask_login import LoginManager
from jinja2 import ChoiceLoader, DictLoader, FileSystemLoader
from openpyxl import load_workbook

from app.fuel.routes import station_power_comparison_routes as routes
from app.fuel.services.fuel_exports.export_station_power_comparison_services import (
    export_station_power_comparison_to_excel,
)


def _comparison(**overrides):
    data = {
        "selected_year": 2023,
        "show_p_rasp": False,
        "show_ved0": False,
        "show_boilers": False,
        "rounding_digits": 0,
        "display_rows": [],
        "pagination": None,
        "year_choices": [{"value": 2023, "label": "2023"}],
    }
    data.update(overrides)
    return data


@pytest.mark.parametrize("show_p_rasp", [False, True])
@pytest.mark.parametrize("digits,number_format", [(0, "#,##0.##########"), (2, "#,##0.00"), (-1, "#,##0")])
def test_export_preserves_numeric_values_hierarchy_and_mismatch_colors(show_p_rasp, digits, number_format):
    totals = {
        "p_ust": Decimal("48.1234"), "p_rasp": Decimal("35.5"),
        "nust": Decimal("48"), "nr": Decimal("35.5"),
        "p_ust_mismatch": True, "p_rasp_mismatch": False,
    }
    comparison = _comparison(show_p_rasp=show_p_rasp, rounding_digits=digits, display_rows=[
        {"kind": "res", "label": "РЭС Архангельской области", **totals},
        {
            "kind": "station_group", "label": "ТЭС-2 / ТЭС-3", "numb": 11,
            "group_name": "Общая группа", "group_type": "ТЭЦ", "machine_number": "1, 2, 3",
            "machine_type": "Турбина", "title": "Шесть агрегатов двух станций", **totals,
        },
        {"kind": "fuel", "group_id": 20, "group_name": "Оболочка", "ved": 0, "nust": Decimal("100"), "nr": None},
    ])

    output = export_station_power_comparison_to_excel(comparison)
    assert output.tell() == 0
    sheet = load_workbook(output).active

    assert sheet.max_row == 6
    assert sheet.max_column == (10 if show_p_rasp else 8)
    assert sheet["A4"].value == "РЭС Архангельской области"
    assert "A4:F4" in {str(region) for region in sheet.merged_cells.ranges}
    assert sheet["A5"].value == "ТЭС-2 / ТЭС-3"
    assert sheet["B5"].value == 11
    assert sheet["G5"].value == pytest.approx(48.1234)
    assert sheet["G5"].data_type == "n"
    assert sheet["G5"].number_format == number_format
    assert sheet["G5"].fill.fgColor.rgb.endswith("F8D7DA")
    nust_col = "I" if show_p_rasp else "H"
    assert sheet[f"{nust_col}5"].value == 48
    assert sheet[f"{nust_col}5"].fill.fgColor.rgb.endswith("F8D7DA")
    assert sheet["A5"].comment.text == "Шесть агрегатов двух станций"
    assert sheet["G6"].value == "—"
    assert sheet[f"{nust_col}6"].value == 100
    if show_p_rasp:
        assert sheet["H5"].value == 35.5
        assert sheet["J6"].value == "—"
    assert sheet.freeze_panes == "G4"


def test_export_writes_database_names_as_text_not_formulas():
    sheet = load_workbook(export_station_power_comparison_to_excel(_comparison(display_rows=[
        {"kind": "station", "label": "=1+1"},
        {"kind": "fuel", "group_name": '=HYPERLINK("https://example.com")', "nust": 0},
    ]))).active

    assert sheet["A4"].value == "=1+1"
    assert sheet["A4"].data_type == "s"
    assert sheet["C5"].value.startswith("=HYPERLINK")
    assert sheet["C5"].data_type == "s"
    assert sheet["H5"].value == 0
    assert sheet["H5"].data_type == "n"


def test_empty_yearless_export_is_a_valid_workbook():
    sheet = load_workbook(export_station_power_comparison_to_excel(_comparison(selected_year=""))).active
    assert sheet.max_row == 3
    assert "Год: Без года" in sheet["A2"].value
    assert sheet["G3"].value == "Руст, МВт"


@pytest.mark.parametrize("raw_year,expected_year", [("2023", 2023), ("", ""), ("bad", None)])
def test_export_route_keeps_options_but_ignores_current_page(monkeypatch, raw_year, expected_year):
    app = Flask(__name__)
    app.config["LOGIN_DISABLED"] = True
    app.add_url_rule("/export", view_func=routes.export_station_power_comparison)
    captured = {}

    def build(**options):
        captured.update(options)
        return _comparison(**{key: value for key, value in options.items() if key != "paginate"})

    monkeypatch.setattr(routes, "build_station_power_comparison", build)
    response = app.test_client().get("/export", query_string={
        "year": raw_year, "page": 99, "p_rasp": "1", "ved0": "1", "boilers": "1", "rounding_digits": "2",
    })

    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in response.headers["Content-Disposition"]
    assert ".xlsx" in response.headers["Content-Disposition"]
    assert captured == {
        "selected_year": expected_year, "show_p_rasp": True, "show_ved0": True,
        "show_boilers": True, "rounding_digits": 2, "paginate": False,
    }
    assert load_workbook(BytesIO(response.data)).active["H3"].value == "Ррасп, МВт"


def test_export_requires_authentication():
    app = Flask(__name__)
    app.secret_key = "test"
    login = LoginManager(app)
    login.user_loader(lambda user_id: None)
    app.add_url_rule("/export", view_func=routes.export_station_power_comparison)

    assert app.test_client().get("/export").status_code == 401


def test_excel_button_passes_current_filters_without_pagination():
    app = Flask(__name__)
    app.jinja_loader = ChoiceLoader([
        DictLoader({"base.html": "{% block content %}{% endblock %}"}),
        FileSystemLoader("app/templates"),
    ])
    for endpoint in ("stations_equipment_groups", "station_power_comparison", "export_station_power_comparison"):
        app.add_url_rule(f"/{endpoint}", endpoint=f"fuel_bp.{endpoint}", view_func=lambda: "")
    with app.test_request_context("/?page=7"):
        html = app.jinja_env.get_template(
            "fuel/equipment_groups/stations_equipment_groups_power_comparison.html",
        ).render(
            current_user=SimpleNamespace(is_authenticated=True),
            comparison=_comparison(show_p_rasp=True, show_ved0=True, show_boilers=True, rounding_digits=-1),
        )

    from html.parser import HTMLParser

    class Links(HTMLParser):
        def __init__(self):
            super().__init__()
            self.urls = []

        def handle_starttag(self, tag, attrs):
            if tag == "a":
                self.urls.append(dict(attrs).get("href", ""))

    links = Links()
    links.feed(html)
    export_urls = [url for url in links.urls if urlsplit(url).path == "/export_station_power_comparison"]
    assert len(export_urls) == 1
    assert parse_qs(urlsplit(export_urls[0]).query) == {
        "year": ["2023"], "p_rasp": ["1"], "ved0": ["1"], "boilers": ["1"], "rounding_digits": ["-1"],
    }
