# -*- coding: utf-8 -*-
from types import SimpleNamespace

from openpyxl import load_workbook

from app.fuel.services.equipment_groups.equipment_group_specific_fuel_consumption_services import (
    SPECIFIC_FUEL_CONSUMPTION_COLUMNS,
)
from app.fuel.services.fuel_exports import (
    export_stations_equipment_group_specific_fuel_consumption_services as svc,
)


def _param(**overrides):
    fields = {
        "name": None,
        "year_number": 2024,
        "k": None,
        "y": None,
        "y_calc": None,
        "btp": None,
        "btp_calc": None,
        "sntp": None,
        "sntp_calc": None,
        "bk": None,
        "bk_calc": None,
        "snk": None,
        "snk_calc": None,
        "numb1120": 100,
        "snbas": None,
        "ksn": None,
        "bbas": None,
        "kh": None,
    }
    fields.update(overrides)
    return SimpleNamespace(**fields)


def _hierarchy_from_rows(rows):
    group_blocks = [
        {"equipment_group": eg, "rows": [(eg, param)]}
        for eg, param in rows
    ]
    return [
        {
            "est_name": "ЕЭС",
            "ues_list": [
                {
                    "ues_name": "ОЭС",
                    "res_list": [
                        {
                            "res_name": "РЭС",
                            "station_blocks": [{"group_blocks": group_blocks}],
                        }
                    ],
                }
            ],
        }
    ]


def _patch_export_source(monkeypatch, rows):
    captured = {}

    def fake_data(**kwargs):
        return {"rows": rows}

    def fake_hierarchy(passed_rows):
        captured["rows"] = list(passed_rows)
        return _hierarchy_from_rows(passed_rows)

    monkeypatch.setattr(svc, "get_equipment_groups_with_specific_fuel_consumption_data", fake_data)
    monkeypatch.setattr(svc, "build_equipment_group_specific_fuel_consumption_hierarchy", fake_hierarchy)
    return captured


def _group_names_from_xlsx(excel_file):
    excel_file.seek(0)
    wb = load_workbook(excel_file)
    ws = wb.active
    names = []
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
        value = row[0]
        if value in (None, "ЕЭС", "ОЭС", "РЭС"):
            continue
        names.append(value)
    return names


def test_export_without_mismatch_filter_keeps_all_rows(monkeypatch):
    mismatch_eg = SimpleNamespace(id=1, name="Mismatch", name_ext=None)
    match_eg = SimpleNamespace(id=2, name="Match", name_ext=None)
    rows = [
        (mismatch_eg, _param(y=10, y_calc=11, numb1120=101)),
        (match_eg, _param(y=10, y_calc=10, numb1120=102)),
    ]
    captured = _patch_export_source(monkeypatch, rows)

    excel_file = svc.export_stations_equipment_group_specific_fuel_consumption_to_excel(
        {}, 2024, 2024, show_calc_mismatch=False
    )

    assert [eg.name for eg, row_param in captured["rows"]] == ["Mismatch", "Match"]
    assert _group_names_from_xlsx(excel_file) == ["Mismatch", "Match"]


def test_export_with_mismatch_filter_keeps_only_displayed_rows(monkeypatch):
    mismatch_eg = SimpleNamespace(id=1, name="Mismatch", name_ext=None)
    match_eg = SimpleNamespace(id=2, name="Match", name_ext=None)
    rows = [
        (mismatch_eg, _param(y=10, y_calc=11, numb1120=101)),
        (match_eg, _param(y=10, y_calc=10, numb1120=102)),
    ]
    captured = _patch_export_source(monkeypatch, rows)

    excel_file = svc.export_stations_equipment_group_specific_fuel_consumption_to_excel(
        {}, 2024, 2024, show_calc_mismatch=True
    )

    assert [eg.name for eg, row_param in captured["rows"]] == ["Mismatch"]
    assert _group_names_from_xlsx(excel_file) == ["Mismatch"]


def _cell_fill_rgb(cell):
    rgb = getattr(cell.fill.fgColor, "rgb", None) or ""
    return str(rgb).upper()


def _column_for_attr(attr):
    return 3 + [name for name, _label, _num in SPECIFIC_FUEL_CONSUMPTION_COLUMNS].index(attr)


def test_export_mismatch_cells_have_red_fill(monkeypatch):
    mismatch_eg = SimpleNamespace(id=1, name="Mismatch", name_ext=None)
    _patch_export_source(
        monkeypatch,
        [(mismatch_eg, _param(y=10, y_calc=11, btp=1, btp_calc=1, numb1120=101))],
    )

    excel_file = svc.export_stations_equipment_group_specific_fuel_consumption_to_excel(
        {}, 2024, 2024, show_calc_mismatch=True
    )
    excel_file.seek(0)
    ws = load_workbook(excel_file).active
    data_row = None
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row=row_idx, column=2).value == "Mismatch":
            data_row = row_idx
            break
    assert data_row is not None

    y_rgb = _cell_fill_rgb(ws.cell(row=data_row, column=_column_for_attr("y")))
    y_calc_rgb = _cell_fill_rgb(ws.cell(row=data_row, column=_column_for_attr("y_calc")))
    btp_rgb = _cell_fill_rgb(ws.cell(row=data_row, column=_column_for_attr("btp")))
    assert y_rgb.endswith("FFC7CE")
    assert y_calc_rgb.endswith("FFC7CE")
    assert not btp_rgb.endswith("FFC7CE")


def test_export_with_mismatch_filter_returns_none_when_nothing_displayed(monkeypatch):
    match_eg = SimpleNamespace(id=2, name="Match", name_ext=None)
    _patch_export_source(monkeypatch, [(match_eg, _param(y=10, y_calc=10))])

    excel_file = svc.export_stations_equipment_group_specific_fuel_consumption_to_excel(
        {}, 2024, 2024, show_calc_mismatch=True
    )

    assert excel_file is None
