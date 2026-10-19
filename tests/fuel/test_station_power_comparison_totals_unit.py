from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import Mock

import pytest

from app.fuel.services.stations import station_power_comparison_services as service


@pytest.fixture
def comparison_data(monkeypatch):
    from app.common.services.get_services.energy_systems import (
        energy_system_type_get_services as est_services,
        regional_energy_system_get_services as res_services,
        union_energy_system_get_services as ues_services,
    )
    from app.fuel.services.stations import stations_equipment_groups_services

    est = SimpleNamespace(id=10, name="ЕЭС")
    ues = SimpleNamespace(id=20, name="ОЭС", energy_system_type=est, id_energy_system_type=est.id)
    res = SimpleNamespace(id=30, name="ЭС Архангельской области", union_energy_system=ues)
    stations = [SimpleNamespace(id=1, name="ТЭС-2"), SimpleNamespace(id=2, name="ТЭС-3")]
    group_types = [SimpleNamespace(id=1, name="ТЭЦ"), SimpleNamespace(id=2, name="котельная")]
    groups = [
        SimpleNamespace(id=group_id, numb=group_id, main=None, name=name, regional_energy_system=region)
        for group_id, name, region in [
            (11, "ТЭС-2 и ТЭС-3", res),
            (12, "Котельная станции", res),
            (13, "Общая котельная", res),
            (14, "Котельная без РЭС", None),
            (15, "Котельная без станции", res),
            (16, "Оболочка", res),
        ]
    ]
    links = []
    set_stations = []
    for group_id, station_ids, type_id in [
        (11, (1, 2), 1), (12, (1,), 2), (13, (1, 2), 2),
        (14, (1, 2), 2), (16, (1,), None), (15, (None,), 2),
    ]:
        for station_id in station_ids:
            link_id = len(links) + 1
            set_stations.append(SimpleNamespace(
                id=link_id, station_id=station_id, equipment_group_type_id=type_id,
            ))
            links.append(SimpleNamespace(
                equipment_group_id=group_id, equipment_group_set_station_id=link_id,
            ))
    machines = [
        SimpleNamespace(
            id=machine_id, id_station=station_id, id_equipment_group=type_id,
            equipment_group=group_types[type_id - 1], machine_number=str(machine_id),
            machine_name="Агрегат", is_archived=False,
        )
        for machine_id, station_id, type_id in [(101, 1, 1), (102, 2, 1), (103, 1, 2)]
    ]
    powers = [
        SimpleNamespace(id=index, id_machine=machine_id, year_number=2023, p_ust=Decimal(p), p_rasp=Decimal(pr))
        for index, (machine_id, p, pr) in enumerate([
            (101, "30", "20"), (102, "18", "15.5"), (103, "5", "4"),
        ])
    ]
    params = [
        SimpleNamespace(
            id=group_id, equipment_group_id=group_id, year_number=2023,
            nust=Decimal(nust), nr=Decimal(nr), ved=ved,
        )
        for group_id, nust, nr, ved in [
            (11, "48", "35.5", 3), (12, "7", "6", 2), (13, "9", "8", 2),
            (14, "11", "10", 2), (15, "13", "12", 2), (16, "100", "90", 0),
        ]
    ]
    data = {
        service.Station: stations,
        service.Machine: machines,
        service.MachinePower: powers,
        service.EquipmentGroup: groups,
        service.EquipmentGroupType: group_types,
        service.EquipmentGroupFuelParam: params,
        service.EquipmentGroupSetStation: set_stations,
        service.MachineFuelParam: [
            SimpleNamespace(machine_id=101, equipment_group_id=11),
            SimpleNamespace(machine_id=102, equipment_group_id=11),
            SimpleNamespace(machine_id=103, equipment_group_id=12),
        ],
    }
    monkeypatch.setattr(service, "get_current_db_version_id", lambda: 37)
    monkeypatch.setattr(service, "get_current_year", lambda: 2023)
    monkeypatch.setattr(service, "_tpp_station_ids", lambda version: {station.id for station in stations})
    monkeypatch.setattr(service, "_versioned_rows", lambda model, version, *options: data[model])
    monkeypatch.setattr(service, "joinedload", lambda *args: Mock())
    monkeypatch.setattr(service, "EquipmentGroupSet", SimpleNamespace(
        query=SimpleNamespace(all=lambda: links),
    ))
    monkeypatch.setattr(stations_equipment_groups_services, "_hierarchy_key_from_station", lambda station: (
        10, 20, 30, None, None,
    ))
    monkeypatch.setattr(est_services, "get_energy_system_type_map", lambda: {10: "ЕЭС"})
    monkeypatch.setattr(ues_services, "get_union_energy_systems_map", lambda: {20: "ОЭС"})
    monkeypatch.setattr(ues_services, "get_union_energy_system_display_order_map", lambda: {})
    monkeypatch.setattr(res_services, "get_regional_energy_systems_map", lambda: {
        30: "Электроэнергетическая система Архангельской области",
    })
    return data


@pytest.mark.parametrize("show_ved0", [False, True])
@pytest.mark.parametrize("show_boilers", [False, True])
def test_fuel_totals_include_boilers_and_unlinked_groups_but_not_shells(
    comparison_data, show_ved0, show_boilers,
):
    result = service.build_station_power_comparison(
        2023, show_p_rasp=True, show_ved0=show_ved0, show_boilers=show_boilers,
    )
    rows = result["display_rows"]
    for kind in ("system", "ues", "res"):
        summaries = [row for row in rows if row["kind"] == kind]
        assert len(summaries) == 1
        summary = summaries[0]
        assert summary["p_ust"] == Decimal("48")
        assert summary["p_rasp"] == Decimal("35.5")
        assert summary["nust"] == Decimal("77")
        assert summary["nr"] == Decimal("61.5")
        assert summary["p_ust_mismatch"]
        assert summary["p_rasp_mismatch"]

    shared = [row for row in rows if row["kind"] == "station_group" and row["numb"] == 11]
    assert len(shared) == 1
    assert {station["id"] for station in shared[0]["stations"]} == {1, 2}
    assert shared[0]["p_ust"] == shared[0]["nust"] == Decimal("48")
    assert shared[0]["p_rasp"] == shared[0]["nr"] == Decimal("35.5")

    assert any(row.get("ved") == 0 for row in rows) == show_ved0
    assert any(row.get("group_id") == 12 for row in rows) == show_boilers
    assert any(row.get("machine_id") == 103 for row in rows) == show_boilers
    assert any(row.get("group_id") == 13 for row in rows) == show_boilers
    assert any(row.get("group_id") == 14 for row in rows) == show_boilers
    assert any(row.get("group_id") == 15 for row in rows) == show_boilers
    for row in rows:
        if row["kind"] in {"unassigned", "unassigned_station", "unclassified"}:
            assert row["p_ust"] == row["p_rasp"] == 0
    assert result["unclassified_fuel_totals"] == service._new_totals()
    if show_boilers:
        unmapped = next(row for row in rows if row["kind"] == "unassigned")
        assert unmapped["nust"] == Decimal("11")
        assert unmapped["nr"] == Decimal("10")
    if show_boilers:
        boiler = next(row for row in rows if row.get("group_id") == 12)
        assert boiler["nust"] == Decimal("7")
        assert boiler["nr"] == Decimal("6")


def test_distinct_res_ids_with_the_same_name_are_not_merged(comparison_data, monkeypatch):
    from app.common.services.get_services.energy_systems import regional_energy_system_get_services
    from app.fuel.services.stations import stations_equipment_groups_services

    comparison_data[service.Station].append(SimpleNamespace(id=3, name="Другая ТЭС"))
    comparison_data[service.Machine].append(SimpleNamespace(
        id=104, id_station=3, id_equipment_group=None, equipment_group=None,
        machine_number="1", machine_name="Агрегат", is_archived=False,
    ))
    comparison_data[service.MachinePower].append(SimpleNamespace(
        id=104, id_machine=104, year_number=2023, p_ust=Decimal("10"), p_rasp=Decimal("8"),
    ))
    monkeypatch.setattr(stations_equipment_groups_services, "_hierarchy_key_from_station", lambda station: (
        10, 20, 31 if station.id == 3 else 30, None, None,
    ))
    monkeypatch.setattr(regional_energy_system_get_services, "get_regional_energy_systems_map", lambda: {
        30: "Одинаковое название", 31: "Одинаковое название",
    })

    result = service.build_station_power_comparison(2023)
    regional_totals = [row for row in result["display_rows"] if row["kind"] == "res"]

    assert len(regional_totals) == 2
    assert sorted(row["p_ust"] for row in regional_totals) == [Decimal("10"), Decimal("48")]
    assert sorted(row["nust"] for row in regional_totals) == [Decimal("0"), Decimal("77")]


@pytest.mark.parametrize("group_type", ["котельная", " КОТЕЛЬНАЯ ", "ТЭЦ; котельная"])
def test_boiler_type_does_not_discard_electrical_fuel_power(group_type):
    totals = service._new_totals()
    service._add_fuel_totals(totals, {
        "group_type": group_type, "ved": 3, "nust": Decimal("450"), "nr": Decimal("400"),
    })
    assert totals["nust"] == Decimal("450")
    assert totals["nr"] == Decimal("400")
    assert totals["p_ust"] == totals["p_rasp"] == 0


def test_parent_parts_check_includes_active_boiler_fuel():
    result = service._ved0_parts_check(
        {"nust": Decimal("48"), "nr": Decimal("35.5")},
        [
            {"has_param": True, "ved": 3, "nust": Decimal("48"), "nr": Decimal("35.5")},
            {"has_param": True, "ved": 3, "group_type": "котельная", "nust": Decimal("7"), "nr": Decimal("6")},
        ],
        show_p_rasp=True,
    )
    assert result["parts_nust"] == Decimal("55")
    assert result["parts_nr"] == Decimal("41.5")
    assert result["nust_mismatch"]
    assert result["nr_mismatch"]


def test_unpaginated_comparison_exports_rows_beyond_the_first_page(comparison_data):
    from openpyxl import load_workbook
    from app.fuel.services.fuel_exports.export_station_power_comparison_services import (
        export_station_power_comparison_to_excel,
    )

    for station_id in range(100, 160):
        comparison_data[service.Station].append(SimpleNamespace(id=station_id, name=f"Станция {station_id}"))
        comparison_data[service.Machine].append(SimpleNamespace(
            id=station_id + 1000, id_station=station_id, id_equipment_group=None, equipment_group=None,
            machine_number="1", machine_name="Агрегат", is_archived=False,
        ))
        comparison_data[service.MachinePower].append(SimpleNamespace(
            id=station_id + 1000, id_machine=station_id + 1000, year_number=2023, p_ust=Decimal("1"), p_rasp=None,
        ))
    page = service.build_station_power_comparison(2023)
    complete = service.build_station_power_comparison(2023, page=99, paginate=False)
    sheet = load_workbook(export_station_power_comparison_to_excel(complete)).active

    assert page["pagination"]["total_pages"] > 1
    assert complete["pagination"] is None
    assert len(complete["display_rows"]) > len(page["display_rows"])
    exported_names = {sheet.cell(row_number, 1).value for row_number in range(4, sheet.max_row + 1)}
    assert {f"Станция {station_id}" for station_id in range(100, 160)} <= exported_names
    regional_totals = [row for row in complete["display_rows"] if row["kind"] == "res"]
    assert len(regional_totals) == 1
    assert regional_totals[0]["nust"] == Decimal("77")


@pytest.mark.parametrize("show_boilers", [False, True])
def test_comparison_matches_fuel_summary_without_requiring_station_links(
    comparison_data, monkeypatch, show_boilers,
):
    from app.common.services import database_version_filter
    from app.common.services.get_services.energy_systems import energy_system_type_get_services
    from app.common.services.get_services.territories import federal_district_get_services
    from app.common.services.get_services.years import year_feature_services
    from app.fuel.services.totals_summary import fuel_totals_summary_services as summary

    groups = comparison_data[service.EquipmentGroup]
    region = groups[0].regional_energy_system
    # Группа с электрической мощностью, но без связи с Generation Station.
    groups.append(SimpleNamespace(id=17, numb=17, main=None, name="Без связи", regional_energy_system=region))
    comparison_data[service.EquipmentGroupFuelParam].append(SimpleNamespace(
        id=17, equipment_group_id=17, year_number=2023,
        nust=Decimal("38.297"), nr=Decimal("30.3348"), ved=4,
    ))
    group_by_id = {group.id: group for group in groups}
    raw_rows = [
        (group_by_id[param.equipment_group_id], param)
        for param in comparison_data[service.EquipmentGroupFuelParam]
    ]
    monkeypatch.setattr(summary, "_load_fuel_param_rows", lambda start, end: raw_rows)
    monkeypatch.setattr(summary, "build_fuel_param_columns_hierarchy", lambda columns: {"fuel_param_columns": columns})
    monkeypatch.setattr(database_version_filter, "get_current_db_version_id", lambda: 37)
    monkeypatch.setattr(energy_system_type_get_services, "get_energy_system_type_list_full", lambda: [
        SimpleNamespace(id=10, name="ЕЭС"),
    ])
    monkeypatch.setattr(federal_district_get_services, "get_federal_district_list_full", lambda: [])
    monkeypatch.setattr(federal_district_get_services, "get_federal_districts_map", lambda: {})
    monkeypatch.setattr(year_feature_services, "get_year_feature_dict", lambda: {})

    totals = summary.build_fuel_totals_summary_context(
        start_year=2023, end_year=2023, aggregation_types=["ees"],
        rounding_digits=0, filter_year_list=[2023],
    )
    comparison = service.build_station_power_comparison(
        2023, show_boilers=show_boilers, paginate=False,
    )
    system = next(row for row in comparison["display_rows"] if row["kind"] == "system")
    assert system["nust"] == Decimal("115.297")
    assert system["nr"] == Decimal("91.8348")
    assert [summary._format_cell(system[attr], 0) for attr in ("nust", "nr")] == totals["blocks"][0]["rows"][0]["cells"][:2]
    unlinked_rows = [row for row in comparison["display_rows"] if row.get("group_id") == 17]
    assert len(unlinked_rows) == 1
    assert unlinked_rows[0]["nust"] == Decimal("38.297")


def test_fuel_uses_group_territory_instead_of_linked_station_territory(comparison_data):
    group = comparison_data[service.EquipmentGroup][1]
    group.regional_energy_system = SimpleNamespace(
        id=31, union_energy_system=comparison_data[service.EquipmentGroup][0].regional_energy_system.union_energy_system,
    )
    comparison = service.build_station_power_comparison(2023, show_boilers=True)
    regional_totals = {row["label"]: row for row in comparison["display_rows"] if row["kind"] == "res"}
    assert regional_totals["Не указано"]["nust"] == Decimal("7")
    assert regional_totals["Электроэнергетическая система Архангельской области"]["nust"] == Decimal("70")


def test_child_group_inherits_parent_territory():
    parent = SimpleNamespace(regional_energy_system=SimpleNamespace(
        id=30, union_energy_system=SimpleNamespace(id=20, id_energy_system_type=10),
    ))
    child = SimpleNamespace(main=11, regional_energy_system=None)
    assert service._equipment_group_path(child, {11: parent}) == (10, 20, 30)


def test_hidden_boiler_only_region_keeps_its_fuel_summary(comparison_data):
    group = comparison_data[service.EquipmentGroup][4]
    group.regional_energy_system = SimpleNamespace(
        id=31, union_energy_system=comparison_data[service.EquipmentGroup][0].regional_energy_system.union_energy_system,
    )
    result = service.build_station_power_comparison(2023, show_boilers=False)
    regional = next(row for row in result["display_rows"] if row["kind"] == "res" and row["label"] == "Не указано")
    assert regional["nust"] == Decimal("13")
    assert regional["nr"] == Decimal("12")
    assert not any(row.get("group_id") == 15 for row in result["display_rows"])
