# -*- coding: utf-8 -*-
"""Checks territory totals on the fuel-parameters hierarchy."""

from types import SimpleNamespace

from app.fuel.services.equipment_groups.equipment_group_fuel_params_services import (
    build_fuel_params_aggregate_rows,
)
from app.fuel.services.equipment_groups.fuel_station_hierarchy_pagination import (
    paginate_fuel_eg_station_hierarchy,
)


def test_territory_total_uses_participating_rows_and_recomputes_ratios():
    active = SimpleNamespace(ved=1, e=8, nust=4, eust=10, eotp=2)
    parent = SimpleNamespace(ved=0, e=100, nust=100, eust=100, eotp=1)
    legacy = SimpleNamespace(ved=None, e=3, nust=1, eust=5, eotp=1)

    def station_block(group_ids, params):
        return {
            "group_blocks": [
                {
                    "rows": [(SimpleNamespace(id=group_id), param)],
                }
                for group_id, param in zip(group_ids, params)
            ]
        }

    aggregate_rows = build_fuel_params_aggregate_rows(
        [
            station_block([1, 2], [active, parent]),
            station_block([3], [legacy]),
        ]
    )

    summary = aggregate_rows[0]["summary"]
    assert summary["e"] == 11
    assert summary["nust"] == 5
    assert summary["eurt"] == 5000
    assert summary["h"] == 2200


def test_extra_fuel_values_are_added_to_full_scope_aggregates():
    from app.fuel.services.calculation.fuel_calculation_edit_data_services import (
        enrich_fuel_hierarchy_summaries_with_extra,
    )

    station_blocks = []
    for group_id in (1, 2):
        equipment_group = SimpleNamespace(id=group_id)
        param = SimpleNamespace(ved=1, year_number=2024)
        station_blocks.append(
            {
                "group_blocks": [
                    {
                        "rows": [(equipment_group, param)],
                    }
                ]
            }
        )

    ues_aggregate = {"summary": {}}
    est_aggregate = {"summary": {}}
    russia_aggregate = {"summary": {}}
    hierarchy = [
        {
            "aggregate_rows": [est_aggregate],
            "ues_list": [
                {
                    "aggregate_rows": [ues_aggregate],
                    "res_list": [
                        {
                            "res_summary": {},
                            "group_blocks": [],
                            "station_blocks": station_blocks,
                        }
                    ],
                }
            ],
        },
        {"is_russia_row": True, "aggregate_rows": [russia_aggregate]},
    ]
    extra_by_group_year = {
        "1:2024": SimpleNamespace(extra=2),
        "2:2024": SimpleNamespace(extra=5),
    }

    enrich_fuel_hierarchy_summaries_with_extra(
        hierarchy, extra_by_group_year, {"extra"}
    )

    assert ues_aggregate["summary"]["extra"] == 7
    assert est_aggregate["summary"]["extra"] == 7
    assert russia_aggregate["summary"]["extra"] == 7


def test_excel_export_includes_territory_aggregate_rows(monkeypatch):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.fuel.services.fuel_exports import (
        export_stations_equipment_group_fuel_params_services as export_service,
    )

    rows = [(SimpleNamespace(id=1), None)]
    hierarchy = [
        {
            "est_name": "ЕЭС России",
            "aggregate_rows": [{"summary": {"e": 10}}],
            "ues_list": [
                {
                    "ues_name": "ОЭС-1",
                    "aggregate_rows": [{"summary": {"e": 10}}],
                    "res_list": [],
                }
            ],
        },
        {"is_russia_row": True, "aggregate_rows": [{"summary": {"e": 10}}]},
    ]
    monkeypatch.setattr(
        export_service,
        "get_equipment_groups_with_fuel_params_data",
        lambda **kwargs: {"rows": rows},
    )
    monkeypatch.setattr(export_service, "build_name_maps_from_rows", lambda _: {})
    monkeypatch.setattr(
        export_service,
        "build_equipment_group_fuel_params_hierarchy",
        lambda *args, **kwargs: hierarchy,
    )
    monkeypatch.setattr(
        export_service, "should_suppress_aggregate_rows_for_filters", lambda _: False
    )

    output = export_service.export_stations_equipment_group_fuel_params_to_excel(
        {}, 2024, 2024
    )
    workbook = load_workbook(BytesIO(output.getvalue()), data_only=True)
    values = {
        str(cell.value).replace("\xa0", " ")
        for row in workbook.active.iter_rows()
        for cell in row
        if cell.value is not None
    }

    assert "ОЭС-1, всего" in values
    assert "ЕЭС России, всего" in values
    assert "Россия, всего" in values


def test_pagination_shows_each_full_scope_total_once_on_its_last_page(monkeypatch):
    from app.fuel.services.equipment_groups import fuel_station_hierarchy_pagination as pagination

    monkeypatch.setattr(pagination, "get_current_db_version_id", lambda: None)
    monkeypatch.setattr(pagination, "get_standalone_equipment_group_ids", lambda **kwargs: set())

    def res_block(res_id, group_id):
        equipment_group = SimpleNamespace(id=group_id, name=f"group-{group_id}")
        return {
            "res_id": res_id,
            "res_name": f"RES-{res_id}",
            "res_summary": {},
            "station_blocks": [
                {
                    "station_id": group_id,
                    "station_name": f"station-{group_id}",
                    "cluster_key": f"station:{group_id}",
                    "suppress_station_summary": True,
                    "group_blocks": [
                        {"equipment_group": equipment_group, "rows": [(equipment_group, None)]}
                    ],
                }
            ],
        }

    hierarchy = [
        {
            "est_id": 1,
            "est_name": "ЕЭС России",
            "aggregate_rows": [{"summary": {"e": 12}}],
            "ues_list": [
                {
                    "ues_id": 10,
                    "ues_name": "ОЭС-1",
                    "aggregate_rows": [{"summary": {"e": 5}}],
                    "res_list": [res_block(100, 1)],
                },
                {
                    "ues_id": 20,
                    "ues_name": "ОЭС-2",
                    "aggregate_rows": [{"summary": {"e": 7}}],
                    "res_list": [res_block(200, 2)],
                },
            ],
        },
        {"is_russia_row": True, "aggregate_rows": [{"summary": {"e": 12}}]},
    ]

    page1, _, _, _ = paginate_fuel_eg_station_hierarchy(hierarchy, 1, 1, ["e"])
    assert page1[0]["ues_list"][0]["aggregate_rows"] == [{"summary": {"e": 5}}]
    assert page1[0]["aggregate_rows"] == []
    assert not any(block.get("is_russia_row") for block in page1)

    page2, _, _, _ = paginate_fuel_eg_station_hierarchy(hierarchy, 2, 1, ["e"])
    assert page2[0]["ues_list"][0]["aggregate_rows"] == [{"summary": {"e": 7}}]
    assert page2[0]["aggregate_rows"] == [{"summary": {"e": 12}}]
    assert page2[-1]["is_russia_row"]
