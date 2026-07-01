# -*- coding: utf-8 -*-
"""Юнит-тесты экспорта выработки ЭЭ и перетоков."""

from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from app.energy_balance.services.ee_generation_export_services import export_ee_generation_to_excel
from app.energy_balance.services.ee_transfers_export_services import export_ee_transfers_to_excel


def _read_sheet(stream: BytesIO) -> list[list]:
    wb = load_workbook(stream)
    ws = wb.active
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def test_export_ee_generation_empty_hierarchy():
    page_data = {
        "hierarchy": [],
        "period_columns": [(2024, "2024")],
        "generation_aggregates": {},
        "should_show_totals": {},
        "res_show_rd_level_map": {},
        "res_verification_by_res": {},
    }
    with patch(
        "app.energy_balance.services.ee_generation_export_services.get_station_ee_generation_page_data",
        return_value=page_data,
    ):
        stream = export_ee_generation_to_excel(
            filters={},
            period_mode="years",
            selected_year=None,
            start_year=2024,
            end_year=2024,
            rounding_digits=1,
            show_totals=False,
            export_verification=False,
        )
    rows = _read_sheet(stream)
    assert rows[0][0] == "КТО"
    assert rows[0][1] == "Электростанция"
    assert rows[1][0] == "Электростанции не найдены"


def test_export_ee_generation_with_station_row():
    page_data = {
        "hierarchy": [
            {
                "est_name": "ЕЭС России",
                "ues_list": [
                    {
                        "ues_id": 1,
                        "ues_name": "ОЭС Центра",
                        "res_list": [
                            {
                                "res_id": 10,
                                "res_name": "РЭС Центра",
                                "rd_list": [
                                    {
                                        "rd_id": 100,
                                        "rd_name": "Московская область",
                                        "eu_list": [
                                            {
                                                "eu_id": 0,
                                                "eu_name": "—",
                                                "sign_groups": [
                                                    {
                                                        "stations": [
                                                            {
                                                                "kto_display": "310340",
                                                                "station_name": "ТЭЦ-1",
                                                                "station_type_name": "ТЭЦ",
                                                                "tes_types": "ТЭС",
                                                                "tes_machine_type": "—",
                                                                "primary_fuel": "Газ",
                                                                "fuel_so": "—",
                                                                "periods": {2024: 12.3},
                                                            }
                                                        ],
                                                    }
                                                ],
                                            }
                                        ],
                                    }
                                ],
                            }
                        ],
                    }
                ],
            }
        ],
        "period_columns": [(2024, "2024")],
        "generation_aggregates": {},
        "should_show_totals": {},
        "res_show_rd_level_map": {10: False},
        "res_verification_by_res": {},
    }
    with patch(
        "app.energy_balance.services.ee_generation_export_services.get_station_ee_generation_page_data",
        return_value=page_data,
    ):
        stream = export_ee_generation_to_excel(
            filters={},
            period_mode="years",
            selected_year=None,
            start_year=2024,
            end_year=2024,
            rounding_digits=1,
            show_totals=False,
            export_verification=False,
        )
    rows = _read_sheet(stream)
    flat = [cell for row in rows for cell in row if cell]
    assert "ТЭЦ-1" in flat
    assert "12,3" in flat or "12.3" in flat


def test_export_ee_transfers_with_row():
    page_data = {
        "hierarchy": [
            {
                "ues_id": 1,
                "ues_name": "ОЭС Центра",
                "res_list": [
                    {
                        "group_key": "res:10",
                        "res_id": 10,
                        "res_name": "РЭС Центра",
                        "transfer_rows": [
                            {
                                "to_name": "РЭС Северо-Запада",
                                "periods": {2024: -5.5},
                            }
                        ],
                    }
                ],
            }
        ],
        "period_columns": [(2024, "2024")],
        "transfer_aggregates": {},
        "should_show_totals": {},
    }
    with patch(
        "app.energy_balance.services.ee_transfers_export_services.get_ee_transfers_page_data",
        return_value=page_data,
    ):
        stream = export_ee_transfers_to_excel(
            filters={},
            period_mode="years",
            selected_year=None,
            start_year=2024,
            end_year=2024,
            rounding_digits=1,
            show_totals=False,
        )
    rows = _read_sheet(stream)
    flat = [cell for row in rows for cell in row if cell]
    assert "РЭС Северо-Запада" in flat
