# -*- coding: utf-8 -*-
"""Юнит-тесты витрины итоговых значений выработки ЭЭ."""

from __future__ import annotations

from decimal import Decimal

from openpyxl import load_workbook

from app.energy_balance.services.ee_generation_totals_summary_services import (
    _sorted_synchronous_area_ids,
    build_ee_generation_totals_rows,
    parse_ee_generation_aggregation_types,
    parse_show_res,
    parse_show_verification,
)
from app.energy_balance.services.export_ee_generation_totals_summary_services import (
    export_ee_generation_totals_summary_to_excel,
)


class _Args(dict):
    def getlist(self, key):
        value = self.get(key, [])
        if isinstance(value, list):
            return value
        return [value]


def _aggregates():
    return {
        "energy_system_types": {
            1: {2023: Decimal("100"), 2024: Decimal("110")},
            2: {2023: Decimal("20"), 2024: Decimal("22")},
        },
        "union_energy_systems": {
            10: {2023: Decimal("70"), 2024: Decimal("75")},
            20: {2023: Decimal("30"), 2024: Decimal("35")},
        },
        "regional_energy_systems": {
            101: {2023: Decimal("40"), 2024: Decimal("42")},
            102: {2023: Decimal("30"), 2024: Decimal("33")},
            201: {2023: Decimal("30"), 2024: Decimal("35")},
        },
        "total": {2023: Decimal("120"), 2024: Decimal("132")},
        "synchronous_areas": {
            7: {2023: Decimal("90"), 2024: Decimal("95")},
            8: {2023: Decimal("30"), 2024: Decimal("37")},
        },
    }


def _row_kwargs(**overrides):
    data = {
        "aggregates": _aggregates(),
        "period_columns": [(2023, "2023"), (2024, "2024")],
        "aggregation_types": ["ees", "oes"],
        "rounding_digits": 1,
        "energy_system_type_names": {1: "ЕЭС России", 2: "ТИТЭС"},
        "sorted_energy_system_type_ids": [1, 2],
        "union_energy_system_names": {10: "ОЭС Центра", 20: "ОЭС Северо-Запада"},
        "sorted_union_energy_system_ids": [10, 20],
        "regional_energy_system_names": {
            101: "Московская",
            102: "Тверская",
            201: "Ленинградская",
        },
        "ues_to_res_ids": {10: [101, 102], 20: [201]},
        "synchronous_area_names": {
            7: "Первая синхронная зона",
            8: "Синхронная зона Востока",
        },
        "sorted_synchronous_area_ids": [7, 8],
    }
    data.update(overrides)
    return data


def _visible_labels(rows):
    return [row["label"] for row in rows if not row.get("is_verification")]
    assert parse_ee_generation_aggregation_types(_Args()) == [
        "russia",
        "ees",
        "sz",
        "oes",
    ]
    assert parse_ee_generation_aggregation_types(_Args(start_year=2023)) == [
        "russia",
        "ees",
        "sz",
        "oes",
    ]


def test_parse_show_res_defaults_on_and_hides_when_zero():
    assert parse_show_res(_Args()) is True
    assert parse_show_res(_Args(show_res=["1"])) is True
    assert parse_show_res(_Args(show_res=["0"])) is False
    assert parse_show_res(_Args(show_res=["0", "1"])) is True


def test_parse_show_verification_defaults_off():
    assert parse_show_verification(_Args()) is False
    assert parse_show_verification(_Args(show_verification=["1"])) is True
    assert parse_show_verification(_Args(show_verification=["0"])) is False


def test_totals_rows_hide_res_when_show_res_false():
    rows = build_ee_generation_totals_rows(
        **_row_kwargs(aggregation_types=["oes"], show_res=False)
    )
    labels = _visible_labels(rows)
    assert labels == ["ОЭС Центра, всего", "ОЭС Северо-Запада, всего"]
    assert all(row["level"] != "res" for row in rows)
    args = _Args(aggregation_type=["ees", "fo", "oes"])
    assert parse_ee_generation_aggregation_types(args) == ["ees", "oes"]


def test_totals_rows_nest_res_inside_ues():
    rows = build_ee_generation_totals_rows(**_row_kwargs())
    labels = _visible_labels(rows)
    assert labels == [
        "ЕЭС России, всего",
        "ОЭС Центра, всего",
        "Московская, всего",
        "Тверская, всего",
        "ОЭС Северо-Запада, всего",
        "Ленинградская, всего",
    ]
    visible = [row for row in rows if not row.get("is_verification")]
    assert visible[0]["level"] == "est"
    assert visible[1]["level"] == "ues"
    assert visible[2]["level"] == "res"
    assert visible[2]["cells"][0] != "–"


def test_totals_rows_insert_sync_areas_between_ees_and_oes():
    rows = build_ee_generation_totals_rows(
        **_row_kwargs(aggregation_types=["ees", "sz", "oes"])
    )
    labels = _visible_labels(rows)
    assert labels[:4] == [
        "ЕЭС России, всего",
        "Первая синхронная зона, всего",
        "Синхронная зона Востока, всего",
        "ОЭС Центра, всего",
    ]
    northwest_idx = labels.index("ОЭС Северо-Запада, всего")
    assert labels.index("Синхронная зона Востока, всего") < northwest_idx
    assert [row["level"] for row in rows if not row.get("is_verification")][1] == "sa"


def test_totals_rows_russia_comes_before_ees():
    rows = build_ee_generation_totals_rows(
        **_row_kwargs(aggregation_types=["russia", "ees", "oes"])
    )
    labels = [row["label"] for row in rows]
    assert labels[0] == "Россия, всего"
    assert labels[1] == "ЕЭС России, всего"
    assert rows[0]["level"] == "russia"
    assert rows[0]["row_class"] == "table-danger"


def test_totals_rows_skip_unspecified_and_empty():
    rows = build_ee_generation_totals_rows(
        **_row_kwargs(
            aggregation_types=["oes"],
            union_energy_system_names={10: "ОЭС Центра", 30: "Не указано"},
            sorted_union_energy_system_ids=[10, 30],
            regional_energy_system_names={101: "Московская", 102: "Не указано"},
            ues_to_res_ids={10: [101, 102], 30: [999]},
        )
    )
    labels = _visible_labels(rows)
    assert "Не указано, всего" not in labels
    assert labels == ["ОЭС Центра, всего", "Московская, всего"]


def test_totals_rows_tites_uses_warning_class():
    rows = build_ee_generation_totals_rows(
        **_row_kwargs(aggregation_types=["tites"])
    )
    assert rows[0]["label"] == "ТИТЭС, всего"
    assert rows[0]["row_class"] == "table-warning"


def test_export_totals_summary_writes_ues_and_indented_res():
    rows = build_ee_generation_totals_rows(**_row_kwargs(aggregation_types=["oes"]))
    stream = export_ee_generation_totals_summary_to_excel(
        {
            "period_columns": [(2023, "2023"), (2024, "2024")],
            "rows": rows,
            "rounding_digits": 1,
            "year_features": {},
        }
    )
    wb = load_workbook(stream)
    ws = wb.active
    labels = [ws.cell(row=idx, column=1).value for idx in range(2, ws.max_row + 1)]
    assert labels[0] == "ОЭС Центра, всего"
    assert labels[1].startswith("    ")
    assert "Московская" in labels[1]
    assert ws.cell(row=2, column=2).value == 70


def test_totals_rows_res_are_editable_with_entity_id():
    rows = build_ee_generation_totals_rows(**_row_kwargs(aggregation_types=["oes"]))
    moscow = next(row for row in rows if row["label"] == "Московская, всего")
    assert moscow["level"] == "res"
    assert moscow["editable"] is True
    assert moscow["entity_id"] == 101
    assert moscow["input_values"][2023] != ""


def test_totals_rows_show_empty_res_for_editing():
    rows = build_ee_generation_totals_rows(
        **_row_kwargs(
            aggregation_types=["oes"],
            aggregates={
                "union_energy_systems": {},
                "regional_energy_systems": {},
            },
        )
    )
    labels = [row["label"] for row in rows]
    assert "Московская, всего" in labels
    moscow = next(row for row in rows if row["entity_id"] == 101)
    assert moscow["editable"] is True
    assert moscow["input_values"][2023] == ""


def test_overlay_hierarchy_control_overrides_parent_totals():
    from app.energy_balance.services.ee_generation_hierarchy_control_services import (
        overlay_hierarchy_control,
    )

    aggregates = _aggregates()
    overlay_hierarchy_control(
        aggregates,
        {
            "russia": {2023: Decimal("999")},
            "ues": {10: {2024: Decimal("1")}},
            "est": {},
            "sa": {},
        },
    )
    assert aggregates["total"][2023] == Decimal("999")
    assert aggregates["total"][2024] == Decimal("132")
    assert aggregates["union_energy_systems"][10][2024] == Decimal("1")
    assert aggregates["union_energy_systems"][10][2023] == Decimal("70")


def test_totals_rows_verification_oes_equals_ues_minus_res():
    rows = build_ee_generation_totals_rows(**_row_kwargs(aggregation_types=["oes"]))
    verify = next(row for row in rows if row["label"] == "Проверка для ОЭС Центра")
    assert verify["is_verification"] is True
    assert verify["editable"] is False
    assert "всем РЭС данной ОЭС" in verify["info_tooltip"]
    # 70 - (40+30) = 0
    assert verify["period_values"][2023] == Decimal("0")
    assert verify["cells"][0] == "0"


def test_totals_rows_verification_ees_minus_oes_and_tites():
    rows = build_ee_generation_totals_rows(
        **_row_kwargs(aggregation_types=["ees", "tites", "oes"])
    )
    verify = next(row for row in rows if row["label"] == "Проверка для ЕЭС России")
    assert verify["is_verification"] is True
    assert "всем ОЭС" in verify["info_tooltip"]
    assert "ТИТЭС" in verify["info_tooltip"]
    # ЕЭС 100 − ОЭС (70+30) − ТИТЭС 20 = −20
    assert verify["period_values"][2023] == Decimal("-20")
    assert verify["cell_nonzero"][0] is True


def test_totals_rows_verification_sa_minus_ues_in_zone():
    rows = build_ee_generation_totals_rows(
        **_row_kwargs(
            aggregation_types=["sz", "oes"],
            res_to_sa_ids={101: frozenset({7}), 102: frozenset({7}), 201: frozenset({8})},
        )
    )
    verify = next(
        row for row in rows if row["label"] == "Проверка для Первая синхронная зона"
    )
    assert verify["is_verification"] is True
    assert "ОЭС и РЭС" in verify["info_tooltip"]
    # СЗ 90 − ОЭС Центра 70 (целиком в СЗ 7) = 20
    assert verify["period_values"][2023] == Decimal("20")


def test_export_skips_verification_rows_by_default():
    rows = build_ee_generation_totals_rows(**_row_kwargs(aggregation_types=["oes"]))
    stream = export_ee_generation_totals_summary_to_excel(
        {
            "period_columns": [(2023, "2023"), (2024, "2024")],
            "rows": rows,
            "rounding_digits": 1,
            "year_features": {},
        }
    )
    wb = load_workbook(stream)
    ws = wb.active
    labels = [ws.cell(row=idx, column=1).value for idx in range(2, ws.max_row + 1)]
    assert all(label is None or not str(label).startswith("Проверка") for label in labels)


def test_sorted_synchronous_area_ids_follow_display_order():
    from types import SimpleNamespace

    areas = [
        SimpleNamespace(id=3, name="Калининградская", display_order=30),
        SimpleNamespace(id=1, name="Первая синхронная зона", display_order=10),
        SimpleNamespace(id=2, name="Синхронная зона Востока", display_order=20),
        SimpleNamespace(id=4, name="Без порядка", display_order=None),
        SimpleNamespace(id=0, name="Не указано", display_order=1),
    ]
    ids, names = _sorted_synchronous_area_ids(areas)
    assert ids == [1, 2, 3, 4]
    assert names[1] == "Первая синхронная зона"
    assert "Не указано" not in names.values()

