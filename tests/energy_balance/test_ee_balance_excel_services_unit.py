# -*- coding: utf-8 -*-
"""Юнит-тесты Excel-выгрузки расчета балансов электрической энергии."""

from __future__ import annotations

from unittest.mock import patch

from openpyxl import load_workbook

from app.energy_balance.services import ee_balance_excel_services
from app.energy_balance.services.ee_balance_excel_services import (
    ee_balance_excel_table_title,
    export_ee_balance_to_excel,
)
from app.energy_balance.services.ee_balance_page_services import (
    EUROPEAN_EES_SLUG,
    EUROPEAN_EES_WITHOUT_SOUTH_URAL_SLUG,
    SEVERO_ZAPAD_CENTER_SLUG,
    URAL_WITHOUT_TYUMEN_SLUG,
)
from app.energy_balance.services.power_balance_page_services import KIND_CHILD, KIND_TOTAL


def test_excel_table_title_uses_energy_unit():
    assert (
        ee_balance_excel_table_title(
            {
                "table_number": 4,
                "title": "Баланс электрической энергии ОЭС Юга",
                "sheet_name": "ОЭС Юга",
            }
        )
        == "Таблица 4 - Баланс электрической энергии ОЭС Юга, млн.кВт·ч"
    )


def _sample_sheets():
    return [
        {
            "slug": "centr",
            "sheet_name": "Центр",
            "title": "Баланс электрической энергии Центр",
            "table_number": 2,
            "unit": "млн.кВт·ч",
        },
        {
            "slug": "sibir",
            "sheet_name": "Сибирь",
            "title": "Баланс электрической энергии Сибирь",
            "table_number": 3,
            "unit": "млн.кВт·ч",
        },
    ]


def _sample_tables():
    return {
        "centr": {
            "sheet": _sample_sheets()[0],
            "rows": [
                {
                    "key": "consumption",
                    "label": "Потребление электрической энергии",
                    "kind": KIND_TOTAL,
                    "indent": 0,
                    "italic": False,
                    "year_values_raw": {2026: 10},
                    "year_values": {2026: "10"},
                    "hide_when_empty": False,
                    "is_flow_block": False,
                },
                {
                    "key": "generation_aes",
                    "label": "АЭС",
                    "kind": KIND_CHILD,
                    "indent": 1,
                    "italic": False,
                    "year_values_raw": {2026: 4},
                    "year_values": {2026: "4"},
                    "hide_when_empty": False,
                    "is_flow_block": False,
                },
            ],
        },
        "sibir": {
            "sheet": _sample_sheets()[1],
            "rows": [
                {
                    "key": "consumption",
                    "label": "Потребление электрической энергии",
                    "kind": KIND_TOTAL,
                    "indent": 0,
                    "italic": False,
                    "year_values_raw": {2026: 20},
                    "year_values": {2026: "20"},
                    "hide_when_empty": False,
                    "is_flow_block": False,
                },
            ],
        },
    }


def test_default_export_includes_virtual_and_subtab_balances():
    base_sheets = [
        {"slug": "ees-rossii", "sheet_name": "ЕЭС России", "layout": "ees_rossii", "group": "ees"},
        {"slug": "1-sz-ees", "sheet_name": "1-я СЗ ЕЭС", "layout": "sz1", "group": "sz"},
        {"slug": "2-sz-ees-vostok", "sheet_name": "2-я СЗ ЕЭС (ОЭС Востока)", "layout": "oes_vostok", "group": "sz"},
        {"slug": "ural", "sheet_name": "Урал", "layout": "oes_standard", "group": "oes"},
        {"slug": "severo-zapad", "sheet_name": "Северо-Запад", "layout": "oes_standard", "group": "oes"},
        {"slug": "tites-hub", "sheet_name": "ТИТЭС Востока", "layout": "oes_vostok", "group": "oes", "skip_table": True, "is_tites_east_hub": True},
        {"slug": "eu-101", "sheet_name": "Центральный энергорайон Камчатского края", "layout": "oes_standard", "group": "eu", "parent_slug": "tites-hub"},
    ]
    tyumen = {"slug": "res-tyumen", "sheet_name": "Тюменская ЭС", "layout": "oes_standard"}
    regions = [
        {"slug": f"res-{name}", "sheet_name": name, "layout": "oes_standard"}
        for name in ("archangelsk", "kaliningrad", "komi")
    ]

    def build_tables(years, *, rounding_digits, sheets):
        return {
            item["slug"]: {"sheet": item, "rows": []}
            for item in sheets
            if not item.get("skip_table")
        }

    def append_european(tables, years, digits, *, slug=EUROPEAN_EES_SLUG,
                        sheet_name="Европейская часть ЕЭС", source_slugs=()):
        sheet = {"slug": slug, "sheet_name": sheet_name, "title": sheet_name}
        tables[slug] = {"sheet": sheet, "rows": []}

    def append_ural(tables, tyumen_slug, years, digits):
        sheet = {"slug": URAL_WITHOUT_TYUMEN_SLUG,
                 "sheet_name": "ОЭС Урала (без Тюменской ЭС)"}
        tables[URAL_WITHOUT_TYUMEN_SLUG] = {"sheet": sheet, "rows": []}

    def append_center(tables, regional_slugs, years, digits):
        sheet = {"slug": SEVERO_ZAPAD_CENTER_SLUG,
                 "sheet_name": "ОЭС Северо-Запада (центр)"}
        tables[SEVERO_ZAPAD_CENTER_SLUG] = {"sheet": sheet, "rows": []}

    with patch.object(ee_balance_excel_services, "get_ee_balance_sheets", return_value=base_sheets), \
         patch.object(ee_balance_excel_services, "_ural_tyumen_energy_system", return_value=tyumen), \
         patch.object(ee_balance_excel_services, "_severo_zapad_region_sheets", return_value=regions), \
         patch.object(ee_balance_excel_services, "_ees_rossii_with_tites_source_slugs",
                      return_value=("ees-rossii", "eu-101")), \
         patch.object(ee_balance_excel_services, "_european_ees_navigation_sheet",
                      return_value={"slug": EUROPEAN_EES_SLUG,
                                    "sheet_name": "Европейская часть ЕЭС",
                                    "layout": "sz1"}), \
         patch.object(ee_balance_excel_services, "build_ee_balance_tables", side_effect=build_tables), \
         patch.object(ee_balance_excel_services, "_append_european_ees_table", side_effect=append_european), \
         patch.object(ee_balance_excel_services, "_append_ural_without_tyumen_table", side_effect=append_ural), \
         patch.object(ee_balance_excel_services, "_overlay_ural_without_tyumen_manual_generation"), \
         patch.object(ee_balance_excel_services, "_append_severo_zapad_center_table", side_effect=append_center):
        stream = export_ee_balance_to_excel(
            years=[2026], rounding_digits=1, year_features={}
        )

    sheetnames = load_workbook(stream).sheetnames
    assert "ЕЭС России" in sheetnames
    assert "Европейская часть ЕЭС" in sheetnames
    assert any(name.startswith("Европейская часть ЕЭС (без") for name in sheetnames)
    assert "Тюменская ЭС" in sheetnames
    assert "ОЭС Урала (без Тюменской ЭС)" in sheetnames
    assert "ОЭС Северо-Запада (центр)" in sheetnames
    assert {"archangelsk", "kaliningrad", "komi"}.issubset(sheetnames)
    assert sheetnames.index("ЭЭС России") < sheetnames.index("ЕЭС России")
    assert sheetnames.index("Европейская часть ЕЭС") < next(
        index for index, name in enumerate(sheetnames)
        if name.startswith("Европейская часть ЕЭС (без")
    ) < sheetnames.index("1-я СЗ ЕЭС")
    assert sheetnames.index("ОЭС Урала (с Тюменской ЭС)") < sheetnames.index(
        "ОЭС Урала (без Тюменской ЭС)"
    ) < sheetnames.index("Тюменская ЭС")
    assert sheetnames.index("ЭС ОЭС Северо-Запада") < sheetnames.index(
        "ОЭС Северо-Запада (центр)"
    ) < sheetnames.index("archangelsk") < next(
        index for index, name in enumerate(sheetnames)
        if name.startswith("Центральный энергорайон")
    )


def test_export_writes_sheet_per_tab_and_energy_title():
    with patch(
        "app.energy_balance.services.ee_balance_excel_services.build_ee_balance_tables",
        return_value=_sample_tables(),
    ), patch(
        "app.energy_balance.services.ee_balance_excel_services.get_ee_balance_sheets",
        return_value=_sample_sheets(),
    ), patch(
        "app.energy_balance.services.ee_balance_excel_services.get_ee_balance_year_columns",
        return_value=[2026],
    ), patch(
        "app.energy_balance.services.ee_balance_excel_services.get_ee_balance_year_features",
        return_value={2026: "План"},
    ):
        stream = export_ee_balance_to_excel(
            tables=_sample_tables(),
            sheets=_sample_sheets(),
            years=[2026],
            year_features={2026: "План"},
            include_type_breakdown=True,
        )
    wb = load_workbook(stream)
    assert wb.sheetnames == ["Центр", "Сибирь"]
    ws = wb["Центр"]
    assert "Баланс электрической энергии" in str(ws["A1"].value)
    assert "млн.кВт·ч" in str(ws["A1"].value)
    assert ws["A3"].value == "Потребление электрической энергии"
    assert ws["B3"].value == 10
    assert ws["A4"].value == "АЭС"
    assert ws.cell(2, 3).value == "Примечание"
