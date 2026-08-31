# -*- coding: utf-8 -*-
"""Юнит-тесты Excel-выгрузки расчета балансов мощности."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from app.energy_balance.services.power_balance_excel_services import (
    export_power_balance_to_excel,
    power_balance_excel_table_title,
)
from app.energy_balance.services.power_balance_page_services import KIND_CHILD, KIND_TOTAL


def test_excel_table_title_uses_system_name_and_number():
    assert (
        power_balance_excel_table_title(
            {
                "table_number": 4,
                "title": "Баланс мощности ОЭС Юга",
                "sheet_name": "ОЭС Юга",
            }
        )
        == "Таблица 4 - Баланс мощности ОЭС Юга, МВт"
    )


def _sample_sheets():
    return [
        {
            "slug": "centr",
            "sheet_name": "Центр",
            "title": "Баланс мощности Центр",
            "table_number": 2,
        },
        {
            "slug": "yug",
            "sheet_name": "Юг",
            "title": "Баланс мощности Юг",
            "table_number": 4,
        },
    ]


def _sample_tables():
    return {
        "centr": {
            "sheet": _sample_sheets()[0],
            "rows": [
                {
                    "key": "demand_max",
                    "label": "Максимум потребления мощности",
                    "kind": "value",
                    "indent": 0,
                    "year_values_raw": {2026: Decimal("10.5")},
                    "year_values": {2026: "10,5"},
                },
                {
                    "key": "installed_total",
                    "label": "Установленная мощность",
                    "kind": KIND_TOTAL,
                    "indent": 0,
                    "year_values_raw": {2026: Decimal("20")},
                    "year_values": {2026: "20"},
                },
                {
                    "key": "installed_aes",
                    "label": "АЭС",
                    "kind": KIND_CHILD,
                    "indent": 1,
                    "year_values_raw": {2026: Decimal("20")},
                    "year_values": {2026: "20"},
                },
                {
                    "key": "installed_machine_9",
                    "label": "Нулевой агрегат",
                    "kind": KIND_CHILD,
                    "indent": 2,
                    "hide_zero_capacity": True,
                    "year_values_raw": {2026: Decimal("0")},
                    "year_values": {2026: "0"},
                },
                {
                    "key": "export",
                    "label": "Экспорт мощности",
                    "kind": "value",
                    "indent": 0,
                    "year_values_raw": {},
                    "year_values": {},
                    "hide_when_empty": True,
                },
                {
                    "key": "surplus_deficit",
                    "label": "Дефицит (-)/избыток (+)",
                    "kind": KIND_TOTAL,
                    "indent": 0,
                    "year_values_raw": {2026: Decimal("-5")},
                    "year_values": {2026: "-5"},
                },
                {
                    "key": "flow_total",
                    "label": "Переток мощности в смежные энергосистемы выдача (-)/прием (+)",
                    "kind": KIND_TOTAL,
                    "indent": 0,
                    "is_flow_block": True,
                    "year_values_raw": {2026: Decimal("4")},
                    "year_values": {2026: "4"},
                },
                {
                    "key": "flow_in",
                    "label": "Получение мощности (+)",
                    "kind": "transfer",
                    "indent": 1,
                    "is_flow_block": True,
                    "year_values_raw": {2026: Decimal("7")},
                    "year_values": {2026: "7"},
                },
                {
                    "key": "surplus_deficit_with_flow",
                    "label": "Дефицит (-)/избыток (+) с учетом перетока мощности в смежные энергосистемы",
                    "kind": KIND_TOTAL,
                    "indent": 0,
                    "is_flow_block": True,
                    "year_values_raw": {2026: Decimal("-1")},
                    "year_values": {2026: "-1"},
                },
            ],
        },
        "yug": {
            "sheet": _sample_sheets()[1],
            "rows": [
                {
                    "key": "demand_max",
                    "label": "Максимум потребления мощности",
                    "kind": "value",
                    "indent": 0,
                    "year_values_raw": {2026: Decimal("1")},
                    "year_values": {2026: "1"},
                }
            ],
        },
    }


def test_export_writes_tab_per_sheet_and_title_row():
    stream = export_power_balance_to_excel(
        years=[2026],
        rounding_digits=1,
        include_type_breakdown=False,
        tables=_sample_tables(),
        sheets=_sample_sheets(),
        year_features={2026: "Отчет"},
    )
    wb = load_workbook(stream)
    assert wb.sheetnames == ["Центр", "Юг"]
    ws = wb["Центр"]
    assert ws.cell(1, 1).value == "Таблица 2 - Баланс мощности Центр, МВт"
    assert ws.cell(2, 1).value == "Наименование"
    assert "2026" in str(ws.cell(2, 2).value)
    labels = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)]
    assert "Максимум потребления мощности" in labels
    assert "Установленная мощность" in labels
    assert "АЭС" not in labels
    assert "Экспорт мощности" in labels
    assert ws.cell(2, 3).value == "Примечание"
    assert ws.cell(3, 2).value == 10.5


def test_export_includes_station_type_rows_when_checkbox_on():
    stream = export_power_balance_to_excel(
        years=[2026],
        rounding_digits=1,
        include_type_breakdown=True,
        tables=_sample_tables(),
        sheets=_sample_sheets(),
        year_features={},
    )
    ws = load_workbook(stream)["Центр"]
    labels = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)]
    assert "АЭС" in labels
    assert "Нулевой агрегат" not in labels


def test_export_writes_row_note_column():
    tables = _sample_tables()
    tables["centr"]["rows"][0]["note"] = "Комментарий к строке"
    stream = export_power_balance_to_excel(
        years=[2026],
        rounding_digits=1,
        include_type_breakdown=False,
        tables=tables,
        sheets=_sample_sheets(),
        year_features={},
    )
    ws = load_workbook(stream)["Центр"]
    assert ws.cell(2, 3).value == "Примечание"
    assert ws.cell(3, 3).value == "Комментарий к строке"
    labels = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)]
    assert "Примечание" not in labels


def test_export_omits_empty_rows_when_toggle_off():
    stream = export_power_balance_to_excel(
        years=[2026],
        rounding_digits=1,
        include_type_breakdown=False,
        show_empty_rows=False,
        tables=_sample_tables(),
        sheets=_sample_sheets(),
        year_features={},
    )
    ws = load_workbook(stream)["Центр"]
    labels = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)]
    assert "Максимум потребления мощности" in labels
    assert "Экспорт мощности" not in labels


def test_export_omits_flow_block_when_toggle_off():
    stream = export_power_balance_to_excel(
        years=[2026],
        rounding_digits=1,
        include_type_breakdown=False,
        show_flow_rows=False,
        tables=_sample_tables(),
        sheets=_sample_sheets(),
        year_features={},
    )
    ws = load_workbook(stream)["Центр"]
    labels = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)]
    assert "Дефицит (-)/избыток (+)" in labels
    assert "Переток мощности в смежные энергосистемы выдача (-)/прием (+)" not in labels
    assert "Получение мощности (+)" not in labels
    assert "Дефицит (-)/избыток (+) с учетом перетока мощности в смежные энергосистемы" not in labels


def test_export_groups_flow_children_like_excel_outline():
    stream = export_power_balance_to_excel(
        years=[2026],
        rounding_digits=1,
        include_type_breakdown=False,
        show_flow_rows=True,
        tables=_sample_tables(),
        sheets=_sample_sheets(),
        year_features={},
    )
    ws = load_workbook(stream)["Центр"]
    labels = {ws.cell(r, 1).value: r for r in range(3, ws.max_row + 1)}
    parent_row = labels["Переток мощности в смежные энергосистемы выдача (-)/прием (+)"]
    child_row = labels["Получение мощности (+)"]
    with_flow_row = labels[
        "Дефицит (-)/избыток (+) с учетом перетока мощности в смежные энергосистемы"
    ]
    assert (ws.row_dimensions[parent_row].outline_level or 0) == 0
    assert ws.row_dimensions[child_row].outline_level == 1
    assert (ws.row_dimensions[with_flow_row].outline_level or 0) == 0
    assert ws.sheet_properties.outlinePr.summaryBelow is False


def test_export_builds_tables_when_not_passed():
    with patch(
        "app.energy_balance.services.power_balance_excel_services.build_power_balance_tables",
        return_value=_sample_tables(),
    ) as build, patch(
        "app.energy_balance.services.power_balance_excel_services.get_power_balance_sheets",
        return_value=_sample_sheets(),
    ), patch(
        "app.energy_balance.services.power_balance_excel_services.get_power_balance_year_columns",
        return_value=[2026],
    ), patch(
        "app.energy_balance.services.power_balance_excel_services.get_power_balance_year_features",
        return_value={},
    ):
        stream = export_power_balance_to_excel(include_type_breakdown=False)
    build.assert_called_once()
    assert isinstance(stream, BytesIO)
    assert load_workbook(stream).sheetnames == ["Центр", "Юг"]
