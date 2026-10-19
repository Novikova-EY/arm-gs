# -*- coding: utf-8 -*-
"""Юнит-тесты Excel-выгрузки выходных документов."""

from decimal import Decimal

from openpyxl import load_workbook

from app.api.services.output_documents_excel_export_services import (
    GENERATION_STATIONS_GENERATION_COLUMN,
    build_dataset_excel_bytes,
)


def test_generation_stations_capacity_is_displayed_with_six_decimal_places():
    stream = build_dataset_excel_bytes(
        {
            "dataset": "generation_stations",
            "rows": [{"Установленная мощность, МВт": 16}],
        }
    )

    worksheet = load_workbook(stream).active
    capacity = worksheet.cell(row=2, column=1)

    assert capacity.value == 16
    assert capacity.number_format == "#,##0.000000"


def test_generation_stations_export_preserves_generation_precision():
    stream = build_dataset_excel_bytes(
        {
            "dataset": "generation_stations",
            "rows": [{GENERATION_STATIONS_GENERATION_COLUMN: Decimal("123.456789")}],
        }
    )

    worksheet = load_workbook(stream).active
    generation = worksheet.cell(row=2, column=1)

    assert generation.value == 123.456789
    assert generation.number_format == "General"


    stream = build_dataset_excel_bytes(
        {
            "dataset": "another_dataset",
            "rows": [{"Установленная мощность, МВт": 16}],
        }
    )

    worksheet = load_workbook(stream).active
    capacity = worksheet.cell(row=2, column=1)

    assert capacity.value == 16
    assert capacity.number_format == "General"
