# -*- coding: utf-8 -*-
"""Юнит-тесты числовых ячеек Excel (формат как на выгрузке «Выработка ЭЭ»)."""

from openpyxl import load_workbook

from app.common.services.excel_numeric_cell import excel_number_format, parse_excel_numeric
from app.energy_consumption.services.energy_consumption_summary_export_services import (
    build_demand_summary_excel_stream as build_ec_excel,
)
from app.power_demand.services.demand_summary_export_services import (
    build_demand_summary_excel_stream as build_pd_excel,
)


def test_excel_number_format_matches_ee_pattern():
    assert excel_number_format(-1) == "#,##0"
    assert excel_number_format(0) == "#,##0.##########"
    assert excel_number_format(1) == "#,##0.0"
    assert excel_number_format(2) == "#,##0.00"
    assert excel_number_format(3) == "#,##0.000"


def test_parse_excel_numeric_accepts_display_strings():
    assert parse_excel_numeric("1 796,65") == 1796.65
    assert parse_excel_numeric("—") is None
    assert parse_excel_numeric(12.3456) == 12.3456


def test_ec_summary_excel_writes_full_precision_float_with_number_format():
    stream = build_ec_excel(
        summary_rows=[
            {
                "show_entity_cell": True,
                "entity_depth": 0,
                "entity_label": "ОЭС Центра",
                "entity_kind": "default",
                "parameter_key": "energy_consumption_mln_kvt_ch",
                "parameter_label": "Потребление электрической энергии, млн кВт·ч",
                "year_values": ["12,3"],
                "year_numeric_tooltips": ["12,3456"],
            }
        ],
        years=[2024],
        sheet_title="Сводка ОЭС",
        rounding_digits=1,
    )
    wb = load_workbook(stream)
    ws = wb.active
    cell = ws.cell(row=2, column=3)
    assert isinstance(cell.value, float)
    assert cell.value == 12.3456
    assert cell.number_format == "#,##0.0"


def test_ec_summary_excel_yoy_uses_two_decimal_format():
    stream = build_ec_excel(
        summary_rows=[
            {
                "show_entity_cell": True,
                "entity_depth": 0,
                "entity_label": "ОЭС Центра",
                "entity_kind": "default",
                "parameter_key": "energy_consumption_yoy_growth_pct",
                "parameter_label": "Годовой темп прироста, %",
                "year_values": ["2,30"],
                "year_numeric_tooltips": ["2,30123"],
            }
        ],
        years=[2024],
        sheet_title="Сводка ОЭС",
        rounding_digits=1,
    )
    wb = load_workbook(stream)
    ws = wb.active
    cell = ws.cell(row=2, column=3)
    assert isinstance(cell.value, float)
    assert abs(cell.value - 2.30123) < 1e-9
    assert cell.number_format == "#,##0.00"


def test_pd_summary_excel_writes_full_precision_float_with_number_format():
    stream = build_pd_excel(
        summary_rows=[
            {
                "show_entity_cell": True,
                "entity_depth": 0,
                "entity_label": "ОЭС Центра",
                "parameter_key": "max_power",
                "parameter_label": "Максимум потребления, МВт",
                "hist_value": "—",
                "hist_numeric_tooltip": "",
                "year_values": ["1 234,6"],
                "year_numeric_tooltips": ["1 234,56789"],
            }
        ],
        years=[2024],
        sheet_title="Сводка нагрузок",
        rounding_digits=1,
        show_hist_col=True,
    )
    wb = load_workbook(stream)
    ws = wb.active
    # col1 entity, col2 param, col3 hist, col4 year
    cell = ws.cell(row=2, column=4)
    assert isinstance(cell.value, float)
    assert abs(cell.value - 1234.56789) < 1e-9
    assert cell.number_format == "#,##0.0"


def test_pd_summary_excel_integer_rounding_format():
    stream = build_pd_excel(
        summary_rows=[
            {
                "show_entity_cell": True,
                "entity_depth": 0,
                "entity_label": "ОЭС Центра",
                "parameter_key": "max_power",
                "parameter_label": "Максимум потребления, МВт",
                "hist_value": "—",
                "year_values": ["100"],
                "year_numeric_tooltips": ["100,4"],
            }
        ],
        years=[2024],
        sheet_title="Сводка нагрузок",
        rounding_digits=-1,
        show_hist_col=False,
    )
    wb = load_workbook(stream)
    ws = wb.active
    cell = ws.cell(row=2, column=3)
    assert isinstance(cell.value, float)
    assert cell.value == 100.4
    assert cell.number_format == "#,##0"
