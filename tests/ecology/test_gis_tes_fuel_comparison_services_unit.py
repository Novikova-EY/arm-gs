# -*- coding: utf-8 -*-
"""Юнит-тесты сопоставления ГИС ТЭС и Топливо."""

from __future__ import annotations

from openpyxl import Workbook

from decimal import Decimal

from app.ecology.services.gis_tes_fuel_comparison_constants import (
    COMPARISON_SHEETS,
    DEFAULT_SHEET_SLUG,
    PG_ECOLOGY_EDIT_RELATED_KEYS,
    PG_ECOLOGY_EDITABLE_KEYS,
    PG_ECOLOGY_OVERLAY_BY_KEY,
    SKIP_SHEET_NAME,
    ZV_ECOLOGY_EDIT_RELATED_KEYS,
    ZV_ECOLOGY_EDITABLE_KEYS,
    ZV_ECOLOGY_OVERLAY_BY_KEY,
    get_comparison_sheet,
    list_comparison_sheets_for_group,
    pg_ecology_header_labels,
)
from app.ecology.services.gis_tes_fuel_comparison_services import (
    _attach_column_formula_tooltips,
    _compute_differences,
    _display_is_nonzero,
    _drop_boiler_groups,
    _pg_ecology_values_by_key,
    _safe_subtract,
    _zv_ecology_values_by_key,
    apply_fuel_indicators,
    apply_fuel_tep_headers,
    apply_fuel_tce_headers,
    apply_pg_ecology_headers,
    apply_pg_ecology_overlay,
    apply_station_affiliation_cells,
    apply_station_generation_overlay,
    apply_zv_ecology_overlay,
    apply_pg_fuel_group_headers,
    apply_zv_gaseous_group_headers,
    apply_zv_emissions_banner_headers,
    apply_zv_analysis_banner_headers,
    apply_zv_leaf_group_headers,
    apply_zv_display_column_order,
    build_comparison_table_context,
    build_header_cells,
    clear_comparison_cache,
    drop_key_header_row,
    load_station_affiliations_by_numb_year,
    load_station_generation_overlay,
    parse_sheet_rows,
    prepare_visible_table,
    save_analysis_ecology_cell,
    save_analysis_pg_ecology_cell,
    save_analysis_zv_ecology_cell,
    station_generation_check_pairs,
    station_generation_mismatch,
)
from types import SimpleNamespace
from typing import Any


def _patch_pg_groups_from_fuel(
    monkeypatch,
    year_rows: list[tuple[int, int, str]],
    fuel_overlay: dict[tuple[int, int], dict[str, Any]] | None = None,
) -> None:
    """Мок состава групп из Топливо: [(numb, year, name), ...].

    filters.(...)_filter из страницы передаются в листинг Топлива; мок учитывает
    имя группы и ID-фильтры по тем же ключам, что и row_matches_territory_filters.
    """
    import app.common.services.database_version_filter as dvf
    import app.ecology.services.analysis_pg_param_services as aps
    import app.ecology.services.analysis_zv_param_services as azps
    import app.fuel.services.calculation.fuel_calculation_edit_data_services as fced
    import app.fuel.services.equipment_groups.equipment_group_fuel_params_services as egfps
    import app.fuel.services.stations.stations_equipment_groups_services as segs

    by_numb: dict[int, list[tuple[Any, Any]]] = {}
    for numb, year, name in year_rows:
        eg = SimpleNamespace(id=numb, numb=numb, name=name, name_ext=None)
        param = SimpleNamespace(numb1120=numb, year_number=year, ved=1)
        by_numb.setdefault(numb, []).append((eg, param))
    flat = [item for items in by_numb.values() for item in items]
    choices = sorted(by_numb)

    def _unique_egs(rows):
        seen = set()
        out = []
        for eg, _param in rows:
            if eg.id in seen:
                continue
            seen.add(eg.id)
            out.append(eg)
        return out

    _ID_FILTER_FIELDS = (
        ("energy_system_type_filter", "energy_system_type_id"),
        ("union_energy_system_filter", "ues_id"),
        ("regional_energy_system_filter", "res_id"),
        ("federal_district_filter", "federal_district_id"),
        ("regional_district_filter", "regional_district_id"),
    )

    def _fuel_listing(**kwargs):
        filters = kwargs.get("filters") or {}
        name_q = str(filters.get("equipment_group_name_filter") or "").strip().casefold()
        rows = flat
        if name_q:
            rows = [
                item
                for item in flat
                if name_q in str(getattr(item[0], "name", "") or "").casefold()
            ]
        for filter_key, field in _ID_FILTER_FIELDS:
            selected = filters.get(filter_key) or []
            if not selected:
                continue
            wanted = {int(value) for value in selected}
            rows = [
                item
                for item in rows
                if int(
                    ((fuel_overlay or {}).get((item[0].id, item[1].year_number)) or {}).get(
                        field
                    )
                    or 0
                )
                in wanted
            ]
        return {"rows": rows, "numb1120_filter_choices": choices}

    monkeypatch.setattr(dvf, "get_current_db_version_id", lambda: 1)
    monkeypatch.setattr(
        egfps,
        "get_equipment_groups_with_fuel_params_data",
        _fuel_listing,
    )
    monkeypatch.setattr(fced, "_expand_fuel_param_rows_for_year_interval", lambda rows, start, end: rows)
    monkeypatch.setattr(
        egfps,
        "build_equipment_group_fuel_params_hierarchy",
        lambda rows, **kwargs: [
            {
                "est_id": 1,
                "est_name": "ЕЭС",
                "ues_list": [
                    {
                        "ues_id": 2,
                        "ues_name": "ОЭС",
                        "res_list": [
                            {
                                "res_id": 3,
                                "res_name": "РЭС",
                                "station_blocks": [
                                    {
                                        "group_blocks": [
                                            {
                                                "equipment_group": eg,
                                                "rows": [
                                                    item
                                                    for item in rows
                                                    if item[0].id == eg.id
                                                ],
                                                "is_composite_parent": False,
                                                "is_composite_child": False,
                                                "show_composite_cluster": False,
                                            }
                                            for eg in _unique_egs(rows)
                                        ]
                                    }
                                ],
                            }
                        ],
                    }
                ],
            }
        ],
    )
    monkeypatch.setattr(egfps, "should_suppress_aggregate_rows_for_filters", lambda filters: False)
    monkeypatch.setattr(segs, "get_standalone_equipment_group_ids", lambda **kwargs: set())
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_pg_from_fuel_services.hidden_boiler_numbs_for_sheet",
        lambda sheet_slug, version_id=None: set(),
    )
    monkeypatch.setattr(aps, "seed_missing_analysis_pg_params_from_reports", lambda *a, **k: 0)
    monkeypatch.setattr(azps, "seed_missing_analysis_zv_params_from_reports", lambda *a, **k: 0)
    # Группы состава привязаны к станциям «Генерации»: иначе на «С котельными»
    # такие группы не показываются.
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_generation_station_links_by_numb",
        lambda numbs: {
            int(numb): [{"id": int(numb), "name": f"Станция {numb}"}]
            for numb in numbs
        },
    )


def test_drop_boiler_groups_hides_boilers_without_parent(monkeypatch):
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services._load_hidden_boiler_numbs",
        lambda sheet_slug: {38},
    )
    groups = [{"numb": 708}, {"numb": 38}, {"numb": None}]

    assert _drop_boiler_groups(groups, "zv-vse") == [{"numb": 708}, {"numb": None}]


def test_four_sheets_except_list1():
    slugs = [item["slug"] for item in COMPARISON_SHEETS]
    names = [item["excel_name"] for item in COMPARISON_SHEETS]
    assert slugs == ["zv-vse", "pg-vse", "zv-bez-kot", "pg-bez-kot"]
    assert names == ["ЗВ_всё", "ПГ_всё", "ЗВ_без кот", "ПГ_без кот"]
    assert SKIP_SHEET_NAME == "Лист1"
    assert DEFAULT_SHEET_SLUG == "zv-vse"
    assert get_comparison_sheet("pg-vse")["page_title"].startswith("Парниковые газы")
    assert get_comparison_sheet("missing") is None
    zv_sheets = list_comparison_sheets_for_group("zv")
    pg_sheets = list_comparison_sheets_for_group("pg")
    assert [item["slug"] for item in zv_sheets] == ["zv-vse", "zv-bez-kot"]
    assert [item["slug"] for item in pg_sheets] == ["pg-vse", "pg-bez-kot"]
    assert [item["nav_label"] for item in zv_sheets] == ["С котельными", "Без котельных"]
    assert [item["nav_label"] for item in pg_sheets] == ["С котельными", "Без котельных"]
    assert list_comparison_sheets_for_group("missing") == []
    from app.ecology.services.gis_tes_fuel_comparison_constants import (
        ANALYSIS_SHEET_SLUGS,
        BEZ_KOT_SHEET_SLUGS,
    )

    assert ANALYSIS_SHEET_SLUGS == {"zv-vse", "pg-vse", "zv-bez-kot", "pg-bez-kot"}
    assert BEZ_KOT_SHEET_SLUGS == {"pg-bez-kot", "zv-bez-kot"}


def test_header_cells_merge_empty_neighbors():
    grid = [
        ["Электростанция", "в том числе", None, "Принадлежность"],
        [None, "Твердые", "Сера", None],
    ]
    rows = build_header_cells(grid)
    assert rows[0][0]["text"] == "Электростанция"
    assert rows[0][0]["rowspan"] == 2
    assert rows[0][1]["text"] == "в том числе"
    assert rows[0][1]["colspan"] == 2
    assert rows[0][1]["rowspan"] == 1
    assert rows[1][0]["text"] == "Твердые"
    assert rows[1][1]["text"] == "Сера"


def test_parse_sheet_keeps_excel_captions_and_appends_key_row():
    meta = get_comparison_sheet("zv-bez-kot")
    rows = [
        ["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"],
        ["Федеральный округ", "ОЭС", "Субъект Федерации", "Электростанция", None, "Год", "Выработка"],
        [None, None, None, None, None, None, "млн. кВт.ч"],
        ["ФО", "ОЭС СЗ", "АО", "ТЭЦ-1", 1, 2017, 10.25],
        ["ФО", "ОЭС СЗ", "АО", "ТЭЦ-1", 1, 2018, 0],
        ["ФО", "ОЭС СЗ", "АО", "ТЭЦ-2", 2, 2017, 3],
    ]
    parsed = parse_sheet_rows(rows, meta)
    assert parsed["years"] == [2017, 2018]
    assert parsed["columns"][0]["sticky"] is True
    assert parsed["columns"][3]["key"] == "NAME"
    assert parsed["columns"][6]["numeric"] is True
    assert len(parsed["header_rows"]) == 3
    caption_texts = [cell["text"] for row in parsed["header_rows"][:-1] for cell in row]
    assert "Федеральный округ" in caption_texts
    assert "Код группы оборудования (numb1120)" in caption_texts
    assert "Выработка" in caption_texts
    assert "млн. кВт.ч" in caption_texts
    key_row = parsed["header_rows"][-1]
    assert all(cell.get("is_key") for cell in key_row)
    assert [cell["text"] for cell in key_row] == [
        "Fed_okr",
        "nameoes",
        "oblname",
        "NAME",
        "numb1120",
        "Year",
        "W_en",
    ]


def test_apply_fuel_indicators_by_numb_and_year():
    meta = get_comparison_sheet("zv-bez-kot")
    parsed = parse_sheet_rows(
        [
            ["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "nust", "W_en", "Teplo"],
            ["ФО", "ОЭС", "Субъект", "Станция", None, "Год", "Nust", "Выработка", "Тепло"],
            ["excel-fo", "excel-oes", "excel-obl", "excel-st", 1, 2017, 10, 20, 30],
        ],
        meta,
    )
    overlay = {
        (1, 2017): {
            "federal_district": "Северо-Западный",
            "ues": "ОЭС Северо-Запада",
            "regional_district": "Архангельская область",
            "station_name": "Архангельская ТЭЦ",
            "nust": 450,
            "e": 1810.029,
            "q": 2697.017,
        }
    }
    filled = apply_fuel_indicators(parsed["rows"][0], parsed["columns"], overlay)
    assert filled[0] == "Северо-Западный"
    assert filled[1] == "ОЭС Северо-Запада"
    assert filled[2] == "Архангельская область"
    assert filled[3] == "Архангельская ТЭЦ"
    assert filled[4] == 1
    assert filled[5] == 2017
    assert filled[6] == 450
    assert filled[7] == 1810.029
    assert filled[8] == 2697.017
    assert apply_fuel_indicators(parsed["rows"][0], parsed["columns"], {})[0] == "excel-fo"


def test_apply_fuel_indicators_scales_tut_to_tce():
    meta = get_comparison_sheet("pg-bez-kot")
    parsed = parse_sheet_rows(
        [
            [
                "Fed_okr",
                "nameoes",
                "oblname",
                "NAME",
                "numb1120",
                "Year",
                "GAZ_All",
                "GAZ_ESTEST",
                "MAZUT_6tp",
                "TvTopl_6tp",
                "PROCH_6tp",
            ],
            ["ФО", "ОЭС", "Субъект", "Станция", None, "Год", "газ", "ест", "мазут", "тв", "проч"],
            ["fo", "oes", "obl", "st", 708, 2017, 1, 2, 3, 117.531, 4],
        ],
        meta,
    )
    from decimal import Decimal

    overlay = {
        (708, 2017): {
            "gaz_all_tce": Decimal("1006581"),
            "gaz_estest_tce": Decimal("1006581"),
            "mazut_tce": Decimal("7338"),
            "ugol_tce": Decimal("117531"),
            "proch_tce": Decimal("12"),
        }
    }
    filled = apply_fuel_indicators(parsed["rows"][0], parsed["columns"], overlay)
    by_key = {
        col["key"]: filled[i] for i, col in enumerate(parsed["columns"])
    }
    assert by_key["GAZ_All"] == Decimal("1006581")
    assert by_key["GAZ_ESTEST"] == Decimal("1006581")
    assert by_key["MAZUT_6tp"] == Decimal("7338")
    assert by_key["TvTopl_6tp"] == Decimal("117531")
    assert by_key["PROCH_6tp"] == Decimal("12")


def test_apply_fuel_indicators_maps_zv_fuel_columns():
    meta = get_comparison_sheet("zv-vse")
    parsed = parse_sheet_rows(
        [
            [
                "Fed_okr",
                "nameoes",
                "oblname",
                "NAME",
                "numb1120",
                "Year",
                "GAZ_6tp",
                "GAZ_estestv",
                "MAZUT_6tp",
                "TVERD_TOPL_6tp",
            ],
            ["ФО", "ОЭС", "Субъект", "Станция", None, "Год", "газ", "ест", "мазут", "тв"],
            ["fo", "oes", "obl", "st", 1, 2017, 1, 2, 3, 4],
        ],
        meta,
    )
    from decimal import Decimal

    overlay = {
        (1, 2017): {
            "gaz_all_tce": Decimal("2000"),
            "gaz_estest_tce": Decimal("1500"),
            "mazut_tce": Decimal("800"),
            "solid_tce": Decimal("4000"),
        }
    }
    filled = apply_fuel_indicators(parsed["rows"][0], parsed["columns"], overlay)
    by_key = {col["key"]: filled[i] for i, col in enumerate(parsed["columns"])}
    assert by_key["GAZ_6tp"] == Decimal("2000")
    assert by_key["GAZ_estestv"] == Decimal("1500")
    assert by_key["MAZUT_6tp"] == Decimal("800")
    assert by_key["TVERD_TOPL_6tp"] == Decimal("4000")


def test_fuel_indicator_payload_scales_and_sums_tce():
    from decimal import Decimal
    from types import SimpleNamespace

    from app.ecology.services.gis_tes_fuel_comparison_services import (
        _fuel_indicator_payload,
    )

    param = SimpleNamespace(
        ved=1,
        name="ТЭЦ",
        nust=10,
        e=20,
        q=30,
        gaz=Decimal("1.5"),
        isk_gaz=Decimal("0.5"),
        mazut=Decimal("2"),
        ugol=Decimal("3"),
        torf=Decimal("1"),
        proch=Decimal("0.1"),
    )
    group = SimpleNamespace(
        name="ТЭЦ",
        regional_district=None,
        regional_energy_system=None,
        gen_company_external_mapping=None,
    )
    payload = _fuel_indicator_payload(param, group)
    assert payload["gaz_estest_tce"] == Decimal("1500.0")
    assert payload["gaz_all_tce"] == Decimal("2000.0")
    assert payload["mazut_tce"] == Decimal("2000.0")
    assert payload["ugol_tce"] == Decimal("3000.0")
    assert payload["solid_tce"] == Decimal("4000.0")
    assert payload["proch_tce"] == Decimal("100.0")


def test_apply_fuel_tce_headers_renames_consumption_columns():
    columns = [
        {"key": "GAZ_All"},
        {"key": "GAZ_ESTEST"},
        {"key": "MAZUT_6tp"},
        {"key": "TvTopl_6tp"},
        {"key": "PROCH_6tp"},
    ]
    headers = [
        [
            {"text": "старое", "column_key": "GAZ_All", "is_key": False, "colspan": 1},
            {"text": "старое", "column_key": "GAZ_ESTEST", "is_key": False, "colspan": 1},
            {"text": "старое", "column_key": "MAZUT_6tp", "is_key": False, "colspan": 1},
            {"text": "старое", "column_key": "TvTopl_6tp", "is_key": False, "colspan": 1},
            {"text": "старое", "column_key": "PROCH_6tp", "is_key": False, "colspan": 1},
        ]
    ]
    apply_fuel_tce_headers(columns, headers)
    assert headers[0][0]["text"] == "Расход газа, всего, т у.т."
    assert headers[0][0]["access_code"] == "GAZ + ISK_GAZ"
    assert headers[0][1]["text"] == "в т.ч. расход газа естественного, т у.т."
    assert headers[0][1]["access_code"] == "GAZ"
    assert headers[0][2]["text"] == "мазут"
    assert headers[0][2]["access_code"] == "MAZUT"
    assert headers[0][3]["text"] == "Расход твердого топлива, т у.т."
    assert headers[0][3]["access_code"] == "UGOL"
    assert headers[0][4]["text"] == "Расход прочих видов топлива, т у.т."
    assert headers[0][4]["access_code"] == "PROCH"
    assert all(col.get("fuel_tep_highlight") is True for col in columns)
    assert all(cell.get("fuel_tep_highlight") is True for cell in headers[0])


def test_thousand_tut_to_tce_helpers():
    from decimal import Decimal

    from app.ecology.services.gis_tes_fuel_comparison_services import (
        _sum_optional,
        _thousand_tut_to_tce,
    )

    assert _thousand_tut_to_tce(Decimal("117.531")) == Decimal("117531.000")
    assert _thousand_tut_to_tce(None) is None
    assert _sum_optional(Decimal("117.531"), Decimal("0")) == Decimal("117.531")
    assert _sum_optional(None, None) is None
    assert _thousand_tut_to_tce(_sum_optional(Decimal("1006.581"), Decimal("1"))) == Decimal(
        "1007581.000"
    )


def test_fuel_tep_headers_use_fuel_params_labels_and_codes():
    meta = get_comparison_sheet("pg-bez-kot")
    parsed = parse_sheet_rows(
        [
            ["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "NUST", "W_en", "Teplo"],
            ["ФО", "ОЭС", "Субъект", "Станция", None, "Год", "Nust excel", "W excel", "Q excel"],
            ["excel-fo", "excel-oes", "excel-obl", "excel-st", 1, 2017, 10, 20, 30],
        ],
        meta,
    )
    visible, headers, _sticky = prepare_visible_table(parsed["columns"], parsed["header_rows"])
    apply_fuel_tep_headers(visible, headers)
    by_key = {col["key"]: col for col in visible}
    assert by_key["NUST"]["fuel_tep_highlight"] is True
    assert by_key["W_en"]["fuel_tep_highlight"] is True
    assert by_key["Teplo"]["fuel_tep_highlight"] is True
    captions = {
        cell["column_key"]: cell
        for row in headers
        for cell in row
        if cell.get("column_key") in {"NUST", "W_en", "Teplo"} and not cell.get("is_key")
    }
    assert captions["NUST"]["text"] == "Суммарная установленная мощность, МВт"
    assert captions["NUST"]["access_code"] == "NUST"
    assert captions["W_en"]["text"] == "Выработка\u00a0ЭЭ, млн кВтч"
    assert captions["W_en"]["access_code"] == "E"
    assert captions["Teplo"]["text"] == "Отпуск\u00a0ТЭ, тыс.Гкал"
    assert captions["Teplo"]["access_code"] == "Q"
    key_row = next(row for row in headers if row and row[0].get("is_key"))
    key_texts = {
        cell["column_key"]: cell["text"]
        for cell in key_row
        if cell.get("column_key") in {"NUST", "W_en", "Teplo"}
    }
    assert key_texts == {"NUST": "NUST", "W_en": "E", "Teplo": "Q"}


def test_context_keeps_selected_year(tmp_path, monkeypatch):
    path = tmp_path / "cmp.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ЗВ_без кот"
    ws.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
    ws.append(["ФО", "ОЭС", "Субъект", "Станция", "Код", "Год", "Выработка"])
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 12.44])
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2018, 8])
    for title in ("ПГ_без кот", "ЗВ_всё", "ПГ_всё", "Лист1"):
        wb.create_sheet(title)
        extra = wb[title]
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 1])
    wb.save(path)
    clear_comparison_cache()
    _patch_pg_groups_from_fuel(
        monkeypatch, [(1, 2017, "ТЭЦ"), (1, 2018, "ТЭЦ")]
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {
            (1, 2017): {"e": 12.44, "ved": 1},
            (1, 2018): {"e": 8, "ved": 1},
        },
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {},
    )
    context = build_comparison_table_context(
        "zv-bez-kot",
        start_year=2017,
        end_year=2018,
        rounding_digits=1,
        source_path=path,
    )
    assert context is not None
    assert context["start_year"] == 2017
    assert context["end_year"] == 2018
    assert context["total_count"] == 1
    assert [col["key"] for col in context["columns"]] == ["NAME", "numb1120", "Year", "W_en"]
    assert context["sticky_col_count"] == 3
    group = context["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]
    assert len(group["rows"]) == 2
    assert group["rows"][0]["cells"][0] == "ТЭЦ"
    assert group["rows"][0]["cells"][-1] == "12,4"
    header_texts = [cell["text"] for row in context["header_rows"] for cell in row]
    assert "Федеральный округ" not in header_texts
    assert "ОЭС" not in header_texts
    assert "Субъект" not in header_texts
    assert not any(row and row[0].get("is_key") for row in context["header_rows"])
    assert "NAME" not in header_texts
    assert "Fed_okr" not in header_texts
    assert context["comparison_sheets"][0]["slug"] == "zv-vse"
    assert [item["slug"] for item in context["comparison_sheets"]] == ["zv-vse", "zv-bez-kot"]
    clear_comparison_cache()


def test_zv_vse_gaseous_total_is_sum_with_pollutants_formula(tmp_path, monkeypatch):
    path = tmp_path / "zv-gaseous.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ЗВ_всё"
    ws.append(
        [
            "Fed_okr",
            "nameoes",
            "oblname",
            "NAME",
            "numb1120",
            "Year",
            "GAZ and ZHID",
            "SO2_gaz",
            "CO_gaz",
            "NOx_gaz",
            "UGLEVOD",
            "LOS",
            "PROCH_ekol",
        ]
    )
    ws.append(
        [
            "ФО",
            "ОЭС",
            "Субъект",
            "Станция",
            "Код",
            "Год",
            "старое всего",
            "диоксид серы",
            "оксид углерода",
            "оксиды азота",
            "углеводороды",
            "ЛОС",
            "прочие",
        ]
    )
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 999, 1, 2, 3, 4, 5, 6])
    for title in ("ПГ_без кот", "ЗВ_без кот", "ПГ_всё", "Лист1"):
        extra = wb.create_sheet(title)
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 1])
    wb.save(path)
    clear_comparison_cache()
    _patch_pg_groups_from_fuel(monkeypatch, [(1, 2017, "ТЭЦ")])
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {2017: "отч"},
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {},
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_zv_ecology_overlay_by_numb_year",
        lambda sheet_slug, numbs, years: {
            (1, 2017): {
                "sulfur_dioxide": 1,
                "carbon_monoxide": 2,
                "nitrogen_oxides": 3,
                "hydrocarbons": 4,
                "voc": 5,
                "other_gaseous_and_liquid": 6,
            }
        },
    )
    context = build_comparison_table_context(
        "zv-vse",
        start_year=2017,
        end_year=2017,
        rounding_digits=0,
        source_path=path,
    )
    assert context is not None
    keys = [col["key"] for col in context["columns"]]
    total_idx = keys.index("GAZ and ZHID")
    group = context["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]
    assert group["rows"][0]["cells"][total_idx] == "21"
    top_texts = [cell["text"] for cell in context["header_rows"][0]]
    assert "Всего выброшено в атмосферу ЗВ за год, т" in top_texts
    banner = next(
        cell
        for cell in context["header_rows"][0]
        if cell["text"] == "Всего выброшено в атмосферу ЗВ за год, т"
    )
    assert banner["colspan"] == 7
    group_texts = [cell["text"] for cell in context["header_rows"][1]]
    assert "Газообразные и жидкие вещества" in group_texts
    group_cell = next(
        cell
        for cell in context["header_rows"][1]
        if cell["text"] == "Газообразные и жидкие вещества"
    )
    assert group_cell["colspan"] == 7
    leaf_by_key = {
        cell["column_key"]: cell
        for cell in context["header_rows"][2]
        if cell.get("column_key")
    }
    assert leaf_by_key["GAZ and ZHID"]["text"] == "всего"
    assert leaf_by_key["CO_gaz"]["text"] == "оксид углерода [CO]"
    assert leaf_by_key["LOS"]["text"] == "летучие органические соединения [ЛОС]"
    tooltip = leaf_by_key["GAZ and ZHID"].get("formula_tooltip") or ""
    assert "диоксид серы" in tooltip
    assert "оксид углерода [CO]" in tooltip
    assert "прочие газообразные и жидкие" in tooltip
    assert not any(row and row[0].get("is_key") for row in context["header_rows"])
    clear_comparison_cache()


def test_bundled_workbook_has_four_data_sheets():
    from app.ecology.services.gis_tes_fuel_comparison_constants import bundled_xlsx_path
    from app.ecology.services.gis_tes_fuel_comparison_services import load_comparison_sheets

    path = bundled_xlsx_path()
    if not path.is_file():
        return
    clear_comparison_cache()
    parsed = load_comparison_sheets(path)
    assert set(parsed) == {"zv-bez-kot", "pg-bez-kot", "zv-vse", "pg-vse"}
    assert parsed["zv-bez-kot"]["years"] == [2017, 2018, 2019, 2020, 2021, 2022, 2023]
    assert parsed["zv-bez-kot"]["header_rows"]
    key_row = parsed["zv-bez-kot"]["header_rows"][-1]
    assert [cell["text"] for cell in key_row[:8]] == [
        "Fed_okr",
        "nameoes",
        "oblname",
        "NAME",
        "numb1120",
        "Year",
        "nust",
        "W_en",
    ]
    captions = [cell["text"] for row in parsed["zv-bez-kot"]["header_rows"][:-1] for cell in row]
    assert "Выработка электроэнергии, млн. кВт.ч" in captions
    assert "в том числе" in captions
    assert parsed["zv-vse"]["columns"][0]["key"] == "Fed_okr"
    clear_comparison_cache()


def test_territory_filters_and_hidden_geo_columns(tmp_path, monkeypatch):
    path = tmp_path / "cmp.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ЗВ_без кот"
    ws.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
    ws.append(["ФО", "ОЭС", "Субъект", "Станция", "Код", "Год", "Выработка"])
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1", 1, 2017, 12.44])
    ws.append(["ЦФО", "ОЭС Ц", "МО", "ТЭЦ-2", 2, 2017, 8])
    for title in ("ПГ_без кот", "ЗВ_всё", "ПГ_всё", "Лист1"):
        wb.create_sheet(title)
        extra = wb[title]
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 1])
    wb.save(path)
    clear_comparison_cache()
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {},
    )
    overlay = {
        (1, 2017): {
            "energy_system_type_id": 10,
            "ues_id": 20,
            "res_id": 30,
            "federal_district_id": 40,
            "regional_district_id": 50,
            "station_name": "Архангельская ТЭЦ",
            "gen_company_name": "ТГК-2",
        },
        (2, 2017): {
            "energy_system_type_id": 11,
            "ues_id": 21,
            "res_id": 31,
            "federal_district_id": 41,
            "regional_district_id": 51,
            "station_name": "ТЭЦ-2",
            "gen_company_name": "ОГК-1",
        },
    }
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: overlay,
    )
    _patch_pg_groups_from_fuel(
        monkeypatch,
        [(1, 2017, "Архангельская ТЭЦ"), (2, 2017, "ТЭЦ-2")],
        fuel_overlay=overlay,
    )
    by_name = build_comparison_table_context(
        "zv-bez-kot",
        start_year=2017,
        end_year=2017,
        source_path=path,
        request_args={"station_name_filter": "Архангельская"},
    )
    assert by_name is not None
    assert by_name["total_count"] == 1
    group = by_name["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]
    assert group["rows"][0]["cells"][0] == "Архангельская ТЭЦ"

    by_est = build_comparison_table_context(
        "zv-bez-kot",
        start_year=2017,
        end_year=2017,
        source_path=path,
        request_args={"energy_system_type_filter": [11]},
    )
    assert by_est is not None
    assert by_est["total_count"] == 1
    assert by_est["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]["rows"][0]["cells"][0] == "ТЭЦ-2"

    missing = build_comparison_table_context(
        "zv-bez-kot",
        start_year=2017,
        end_year=2017,
        source_path=path,
        request_args={"energy_system_type_filter": [99]},
    )
    assert missing is not None
    assert missing["total_count"] == 0
    assert missing["hierarchy"] == []
    clear_comparison_cache()


def test_hierarchy_headers_and_pagination(tmp_path, monkeypatch):
    path = tmp_path / "cmp.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ЗВ_без кот"
    ws.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
    ws.append(["ФО", "ОЭС", "Субъект", "Станция", "Код", "Год", "Выработка"])
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1", 1, 2017, 1])
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-2", 2, 2017, 2])
    ws.append(["ЦФО", "ОЭС Ц", "МО", "ТЭЦ-3", 3, 2017, 3])
    for title in ("ПГ_без кот", "ЗВ_всё", "ПГ_всё", "Лист1"):
        wb.create_sheet(title)
        extra = wb[title]
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 1])
    wb.save(path)
    clear_comparison_cache()
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {},
    )
    _patch_pg_groups_from_fuel(
        monkeypatch, [(1, 2017, "ТЭЦ-1"), (2, 2017, "ТЭЦ-2"), (3, 2017, "ТЭЦ-3")]
    )
    overlay = {
        (1, 2017): {
            "energy_system_type_id": 10,
            "est_name": "ЕЭС",
            "ues_id": 20,
            "ues": "ОЭС Северо-Запада",
            "res_id": 30,
            "res_name": "Архангельская",
            "station_name": "ТЭЦ-1",
        },
        (2, 2017): {
            "energy_system_type_id": 10,
            "est_name": "ЕЭС",
            "ues_id": 20,
            "ues": "ОЭС Северо-Запада",
            "res_id": 31,
            "res_name": "Карельская",
            "station_name": "ТЭЦ-2",
        },
        (3, 2017): {
            "energy_system_type_id": 10,
            "est_name": "ЕЭС",
            "ues_id": 21,
            "ues": "ОЭС Центра",
            "res_id": 40,
            "res_name": "Московская",
            "station_name": "ТЭЦ-3",
        },
    }
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: overlay,
    )
    page1 = build_comparison_table_context(
        "zv-bez-kot",
        start_year=2017,
        end_year=2017,
        source_path=path,
        request_args={"per_page": 1, "page": 1},
    )
    assert page1 is not None
    assert page1["total_count"] == 3
    assert page1["total_pages"] == 3
    assert page1["current_page"] == 1
    assert page1["per_page"] == 1
    assert page1["hierarchy"][0]["est_name"] == "ЕЭС"
    assert page1["hierarchy"][0]["ues_list"][0]["ues_name"] == "ОЭС"
    assert page1["hierarchy"][0]["ues_list"][0]["res_list"][0]["res_name"] == "РЭС"
    assert len(page1["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"]) == 1

    page2 = build_comparison_table_context(
        "zv-bez-kot",
        start_year=2017,
        end_year=2017,
        source_path=path,
        request_args={"per_page": 1, "page": 2},
    )
    assert page2 is not None
    assert page2["current_page"] == 2
    page2_groups = page2["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"]
    assert [group["numb"] for group in page2_groups] == [2]

    by_code = build_comparison_table_context(
        "zv-bez-kot",
        start_year=2017,
        end_year=2017,
        source_path=path,
        request_args={"numb1120_filter": [2]},
    )
    assert by_code is not None
    assert by_code["total_count"] == 1
    assert by_code["numb1120_filter"] == [2]
    assert by_code["numb1120_filter_choices"] == [2]
    assert by_code["has_active_filters"] is True
    assert by_code["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]["numb"] == 2
    clear_comparison_cache()


def test_pg_ecology_overlay_mapping_keys():
    assert PG_ECOLOGY_OVERLAY_BY_KEY == {
        "GAZ_ek": "gas_consumption_tce",
        "CO2_gaz": "co2_from_gas_t",
        "MAZUT_ek": "oil_fuel_consumption_tce",
        "CO2_mazut": "co2_from_oil_fuel_t",
        "TvTopl_ek": "solid_fuel_consumption_tce",
        "CO2_ugol": "co2_from_solid_fuel_t",
    }
    assert PG_ECOLOGY_EDITABLE_KEYS == frozenset(PG_ECOLOGY_OVERLAY_BY_KEY)
    derived = (
        "d_mazut",
        "d_gaz",
        "d_tv",
        "K_gaz_ek",
        "K_mazut_ek",
        "K_ugol_ek",
        "CO2",
        "UD_CO2_ek",
        "UD_CO2_ener",
    )
    assert PG_ECOLOGY_EDIT_RELATED_KEYS["GAZ_ek"] == ("GAZ_ek",) + derived
    assert PG_ECOLOGY_EDIT_RELATED_KEYS["CO2_gaz"] == ("CO2_gaz",) + derived
    assert PG_ECOLOGY_EDIT_RELATED_KEYS["CO2_ugol"] == ("CO2_ugol",) + derived
    assert set(PG_ECOLOGY_EDIT_RELATED_KEYS) == set(PG_ECOLOGY_OVERLAY_BY_KEY)


def test_pg_ecology_values_by_key_recomputes_diff_and_k():
    by_key = _pg_ecology_values_by_key(
        {
            "gas_consumption_tce": 100,
            "co2_from_gas_t": 250,
            "oil_fuel_consumption_tce": 40,
            "co2_from_oil_fuel_t": 80,
            "solid_fuel_consumption_tce": 20,
            "co2_from_solid_fuel_t": 60,
        },
        {
            "gaz_estest_tce": 110,
            "mazut_tce": 50,
            "ugol_tce": 30,
            "e": 100,
            "q": 200,
        },
    )
    assert by_key["GAZ_ek"] == 100
    assert by_key["d_gaz"] == 10
    assert by_key["K_gaz_ek"] == 2.5
    assert by_key["d_mazut"] == 10
    assert by_key["K_mazut_ek"] == 2.0
    assert by_key["d_tv"] == 10
    assert by_key["K_ugol_ek"] == 3.0
    assert by_key["CO2"] == 390
    assert abs(by_key["UD_CO2_ek"] - 390 * 1000 / 160) < 1e-9
    assert abs(by_key["UD_CO2_ener"] - 390 / (100 + 200 / 859.8 * 1000)) < 1e-9


def test_save_analysis_pg_ecology_cell_rejects_unknown_column():
    try:
        save_analysis_pg_ecology_cell("pg-vse", 1, 2020, "W_en", 1)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "недоступен" in str(exc).lower() or "Недоступен" in str(exc)


def test_pg_ecology_header_labels_are_fuel_agnostic():
    labels = pg_ecology_header_labels()
    assert labels["GAZ_All"] == "Расход топлива, всего, т у.т."
    assert labels["GAZ_ESTEST"] == "в т.ч. расход газа естественного, т у.т."
    assert labels["GAZ_ek"] == "Расход топлива, т у.т."
    assert labels["d_gaz"] == "Разность по расходу топлива, т у.т."
    assert labels["CO2_gaz"] == "Выбросы СО2 от сжигания топлива, т"
    assert labels["K_gaz_ek"] == "Коэффициент эмиссии"
    assert labels["MAZUT_6tp"] == labels["GAZ_ek"]
    assert labels["MAZUT_ek"] == labels["GAZ_ek"]
    assert labels["d_mazut"] == labels["d_gaz"]
    assert labels["CO2_mazut"] == labels["CO2_gaz"]
    assert labels["K_mazut_ek"] == labels["K_gaz_ek"]
    assert labels["TvTopl_6tp"] == labels["GAZ_ek"]
    assert labels["TvTopl_ek"] == labels["GAZ_ek"]
    assert labels["d_tv"] == labels["d_gaz"]
    assert labels["CO2_ugol"] == labels["CO2_gaz"]
    assert labels["K_ugol_ek"] == labels["K_gaz_ek"]


def test_apply_pg_ecology_headers_rewrites_caption_cells():
    columns = [
        {"key": "GAZ_All"},
        {"key": "GAZ_ek"},
        {"key": "CO2_gaz"},
        {"key": "MAZUT_ek"},
        {"key": "d_gaz"},
        {"key": "K_gaz_ek"},
    ]
    headers = [
        [
            {"text": "Газ", "column_key": "GAZ_All", "is_key": False, "colspan": 4},
        ],
        [
            {"text": "старое", "column_key": "GAZ_All", "is_key": False, "colspan": 1, "access_code": "GAZ + ISK_GAZ"},
            {"text": "старое", "column_key": "GAZ_ek", "is_key": False, "colspan": 1},
            {"text": "старое", "column_key": "CO2_gaz", "is_key": False, "colspan": 1},
            {"text": "старое", "column_key": "MAZUT_ek", "is_key": False, "colspan": 1},
            {"text": "Разность", "column_key": "d_gaz", "is_key": False, "colspan": 1},
            {"text": "K газ", "column_key": "K_gaz_ek", "is_key": False, "colspan": 1},
        ],
        [
            {"text": "GAZ_ek", "column_key": "GAZ_ek", "is_key": True, "colspan": 1},
        ],
    ]
    apply_pg_ecology_headers(columns, headers)
    assert headers[0][0]["text"] == "Газ"
    assert headers[1][0]["text"] == "Расход топлива, всего, т у.т."
    assert headers[1][0]["access_code"] == "GAZ + ISK_GAZ"
    assert headers[1][1]["text"] == "Расход топлива, т у.т."
    assert headers[1][2]["text"] == "Выбросы СО2 от сжигания топлива, т"
    assert headers[1][3]["text"] == "Расход топлива, т у.т."
    assert headers[1][4]["text"] == "Разность по расходу топлива, т у.т."
    assert headers[1][5]["text"] == "Коэффициент эмиссии"
    assert headers[2][0]["text"] == "GAZ_ek"
    assert columns[1]["label"] == "Расход топлива, т у.т."


def test_drop_key_header_row_keeps_captions():
    headers = [
        [{"text": "Газ", "column_key": "GAZ_All", "is_key": False}],
        [
            {
                "text": "Расход топлива, т у.т.",
                "column_key": "GAZ_ek",
                "is_key": False,
            }
        ],
        [
            {
                "text": "GAZ_ek",
                "column_key": "GAZ_ek",
                "is_key": True,
                "formula_tooltip": "GAZ_ek = …",
            }
        ],
    ]
    drop_key_header_row(headers)
    assert len(headers) == 2
    assert all(not (row and row[0].get("is_key")) for row in headers)
    assert headers[1][0]["formula_tooltip"] == "GAZ_ek = …"
    drop_key_header_row(headers)
    assert len(headers) == 2


def test_apply_pg_fuel_group_headers_replaces_report_ecology_row():
    columns = [
        {"key": "Stancia", "source_index": 3, "sticky": True, "sticky_index": 0},
        {"key": "NUST", "source_index": 6, "sticky": False, "sticky_index": None},
        {"key": "W_en", "source_index": 7, "sticky": False, "sticky_index": None},
        {"key": "Teplo", "source_index": 8, "sticky": False, "sticky_index": None},
        {"key": "GAZ_All", "source_index": 9, "sticky": False, "sticky_index": None},
        {"key": "GAZ_ESTEST", "source_index": 10, "sticky": False, "sticky_index": None},
        {"key": "GAZ_ek", "source_index": 11, "sticky": False, "sticky_index": None},
        {"key": "d_gaz", "source_index": 12, "sticky": False, "sticky_index": None},
        {"key": "CO2_gaz", "source_index": 13, "sticky": False, "sticky_index": None},
        {"key": "K_gaz_ek", "source_index": 14, "sticky": False, "sticky_index": None},
        {"key": "MAZUT_6tp", "source_index": 15, "sticky": False, "sticky_index": None},
        {"key": "MAZUT_ek", "source_index": 16, "sticky": False, "sticky_index": None},
        {"key": "d_mazut", "source_index": 17, "sticky": False, "sticky_index": None},
        {"key": "CO2_mazut", "source_index": 18, "sticky": False, "sticky_index": None},
        {"key": "K_mazut_ek", "source_index": 19, "sticky": False, "sticky_index": None},
        {"key": "TvTopl_6tp", "source_index": 20, "sticky": False, "sticky_index": None},
        {"key": "TvTopl_ek", "source_index": 21, "sticky": False, "sticky_index": None},
        {"key": "d_tv", "source_index": 22, "sticky": False, "sticky_index": None},
        {"key": "CO2_ugol", "source_index": 23, "sticky": False, "sticky_index": None},
        {"key": "K_ugol_ek", "source_index": 24, "sticky": False, "sticky_index": None},
        {"key": "PROCH_6tp", "source_index": 25, "sticky": False, "sticky_index": None},
    ]
    headers = [
        [
            {"text": "Электростанция", "column_key": "Stancia", "rowspan": 3, "colspan": 1},
            {"text": "Отчет", "column_key": None, "col_index": 6, "rowspan": 1, "colspan": 5},
            {"text": "Макет Экология", "column_key": "GAZ_ek", "rowspan": 1, "colspan": 1},
            {
                "text": "Разность по израсходованному газу, т у.т.",
                "column_key": "d_gaz",
                "rowspan": 2,
                "colspan": 1,
            },
            {"text": "Макет Экология", "column_key": None, "col_index": 13, "rowspan": 1, "colspan": 2},
            {"text": "Отчет", "column_key": "MAZUT_6tp", "rowspan": 1, "colspan": 1},
            {"text": "Макет Экология", "column_key": "MAZUT_ek", "rowspan": 1, "colspan": 1},
            {
                "text": "Разность по израсходованному мазуту, т у.т.",
                "column_key": "d_mazut",
                "rowspan": 2,
                "colspan": 1,
            },
            {"text": "Макет Экология", "column_key": None, "col_index": 18, "rowspan": 1, "colspan": 2},
            {"text": "Отчет", "column_key": "TvTopl_6tp", "rowspan": 1, "colspan": 1},
            {"text": "Макет Экология", "column_key": "TvTopl_ek", "rowspan": 1, "colspan": 1},
            {
                "text": "Разность по израсходованному твердому топливу, т у.т.",
                "column_key": "d_tv",
                "rowspan": 2,
                "colspan": 1,
            },
            {"text": "Макет Экология", "column_key": None, "col_index": 23, "rowspan": 1, "colspan": 2},
            {"text": "Прочие", "column_key": "PROCH_6tp", "rowspan": 3, "colspan": 1},
        ],
        [
            {"text": "NUST leaf", "column_key": "NUST", "rowspan": 2, "colspan": 1},
            {"text": "W_en leaf", "column_key": "W_en", "rowspan": 2, "colspan": 1},
            {"text": "Teplo leaf", "column_key": "Teplo", "rowspan": 2, "colspan": 1},
            {"text": "GAZ_All leaf", "column_key": "GAZ_All", "rowspan": 2, "colspan": 1},
            {"text": "GAZ_ESTEST leaf", "column_key": "GAZ_ESTEST", "rowspan": 2, "colspan": 1},
            {"text": "GAZ_ek leaf", "column_key": "GAZ_ek", "rowspan": 2, "colspan": 1},
            {"text": "CO2_gaz leaf", "column_key": "CO2_gaz", "rowspan": 2, "colspan": 1},
            {"text": "K_gaz leaf", "column_key": "K_gaz_ek", "rowspan": 2, "colspan": 1},
            {"text": "MAZUT_6tp leaf", "column_key": "MAZUT_6tp", "rowspan": 2, "colspan": 1},
            {"text": "MAZUT_ek leaf", "column_key": "MAZUT_ek", "rowspan": 2, "colspan": 1},
            {"text": "CO2_mazut leaf", "column_key": "CO2_mazut", "rowspan": 2, "colspan": 1},
            {"text": "K_mazut leaf", "column_key": "K_mazut_ek", "rowspan": 2, "colspan": 1},
            {"text": "TvTopl_6tp leaf", "column_key": "TvTopl_6tp", "rowspan": 2, "colspan": 1},
            {"text": "TvTopl_ek leaf", "column_key": "TvTopl_ek", "rowspan": 2, "colspan": 1},
            {"text": "CO2_ugol leaf", "column_key": "CO2_ugol", "rowspan": 2, "colspan": 1},
            {"text": "K_ugol leaf", "column_key": "K_ugol_ek", "rowspan": 2, "colspan": 1},
        ],
        [
            {"text": "6тп-Экол", "column_key": "d_gaz", "rowspan": 1, "colspan": 1},
            {"text": "6тп-Экол", "column_key": "d_mazut", "rowspan": 1, "colspan": 1},
            {"text": "6тп-Экол", "column_key": "d_tv", "rowspan": 1, "colspan": 1},
        ],
        [{"text": "Stancia", "column_key": "Stancia", "is_key": True}],
    ]
    apply_pg_fuel_group_headers(columns, headers)
    top_texts = [(c["text"], c.get("colspan")) for c in headers[0]]
    assert ("Отчет", 5) not in top_texts
    assert "Макет Экология" not in [c["text"] for c in headers[0]]
    assert "ТЭП" not in [c["text"] for c in headers[0]]
    assert ("Газ", 6) in top_texts
    assert ("Мазут", 5) in top_texts
    assert ("Твердое топливо", 5) in top_texts
    top_by_key = {c["column_key"]: c for c in headers[0] if c.get("column_key")}
    assert top_by_key["NUST"]["rowspan"] == 2
    assert top_by_key["W_en"]["rowspan"] == 2
    assert top_by_key["Teplo"]["rowspan"] == 2
    mid_by_key = {c["column_key"]: c["text"] for c in headers[1]}
    assert "NUST" not in mid_by_key
    assert "W_en" not in mid_by_key
    assert "Teplo" not in mid_by_key
    assert "Разность по израсходованному газу, т у.т." == mid_by_key["d_gaz"]
    assert "Разность по израсходованному мазуту, т у.т." == mid_by_key["d_mazut"]
    assert "Разность по израсходованному твердому топливу, т у.т." == mid_by_key["d_tv"]
    assert all(
        "6тп-экол" not in str(cell.get("text") or "").casefold()
        for row in headers
        for cell in row
    )
    assert headers[0][0]["rowspan"] == 2
    assert mid_by_key  # mid row still present
    assert headers[1][0]["rowspan"] == 1
    assert len([row for row in headers if not (row and row[0].get("is_key"))]) == 2


def test_apply_zv_gaseous_group_headers_matches_pollutants_page():
    from app.ecology.services.gis_tes_fuel_comparison_constants import (
        zv_gaseous_header_group,
    )

    title, group_keys, labels = zv_gaseous_header_group()
    keys = ["NAME", "TVERD_V", *group_keys, "energy"]
    columns = [
        {
            "key": key,
            "source_index": index,
            "sticky": key == "NAME",
            "sticky_index": 0 if key == "NAME" else None,
        }
        for index, key in enumerate(keys)
    ]
    headers = [
        [
            {
                "text": f"old {key}",
                "column_key": key,
                "rowspan": 1,
                "colspan": 1,
                "col_index": index,
            }
            for index, key in enumerate(keys)
        ],
        [
            {
                "text": key,
                "column_key": key,
                "is_key": True,
                "rowspan": 1,
                "colspan": 1,
            }
            for key in keys
        ],
    ]
    apply_zv_gaseous_group_headers(columns, headers)
    _attach_column_formula_tooltips(
        headers,
        {"GAZ and ZHID": "всего = диоксид серы + оксид углерода [CO]"},
    )
    drop_key_header_row(headers)
    assert title == "Газообразные и жидкие вещества"
    assert not any(row and row[0].get("is_key") for row in headers)
    group = next(cell for cell in headers[0] if cell["text"] == title)
    assert group["colspan"] == 7
    assert group["rowspan"] == 1
    assert headers[0][0]["rowspan"] == 2
    leaf_by_key = {cell["column_key"]: cell for cell in headers[1]}
    assert [leaf_by_key[key]["text"] for key in group_keys] == [labels[key] for key in group_keys]
    assert leaf_by_key["GAZ and ZHID"]["text"] == "всего"
    assert "диоксид серы" in leaf_by_key["GAZ and ZHID"]["formula_tooltip"]
    assert group.get("formula_tooltip") in (None, "")


def test_apply_zv_emissions_banner_headers_spans_total_solid_and_gaseous():
    from app.ecology.services.gis_tek_constants import ZV_EMISSIONS_BANNER_TITLE
    from app.ecology.services.gis_tes_fuel_comparison_constants import (
        zv_gaseous_header_group,
    )

    title, group_keys, labels = zv_gaseous_header_group()
    keys = ["NAME", "VIBROS_ekol", "TVERD_V", *group_keys, "energy"]
    columns = [
        {
            "key": key,
            "source_index": index,
            "sticky": key == "NAME",
            "sticky_index": 0 if key == "NAME" else None,
        }
        for index, key in enumerate(keys)
    ]
    headers = [
        [
            {
                "text": f"old {key}",
                "column_key": key,
                "rowspan": 1,
                "colspan": 1,
                "col_index": index,
            }
            for index, key in enumerate(keys)
        ],
        [
            {
                "text": key,
                "column_key": key,
                "is_key": True,
                "rowspan": 1,
                "colspan": 1,
            }
            for key in keys
        ],
    ]
    apply_zv_gaseous_group_headers(columns, headers)
    apply_zv_emissions_banner_headers(columns, headers)
    _attach_column_formula_tooltips(
        headers,
        {
            "GAZ and ZHID": "всего = диоксид серы",
            "VIBROS_ekol": "всего = Твердые вещества + Газообразные и жидкие вещества (всего)",
        },
    )
    drop_key_header_row(headers)
    assert len(headers) == 3
    banner = next(cell for cell in headers[0] if cell["text"] == ZV_EMISSIONS_BANNER_TITLE)
    assert banner["colspan"] == 2 + len(group_keys)
    name_cell = next(cell for cell in headers[0] if cell.get("column_key") == "NAME")
    assert name_cell["rowspan"] == 3
    energy_cell = next(cell for cell in headers[0] if cell.get("column_key") == "energy")
    assert energy_cell["rowspan"] == 3
    total_cell = next(cell for cell in headers[1] if cell.get("column_key") == "VIBROS_ekol")
    assert total_cell["text"] == "всего"
    assert "Твердые вещества" in (total_cell.get("formula_tooltip") or "")
    group = next(cell for cell in headers[1] if cell["text"] == title)
    assert group["colspan"] == 7
    leaf_by_key = {cell["column_key"]: cell for cell in headers[2] if cell.get("column_key")}
    assert leaf_by_key["GAZ and ZHID"]["text"] == labels["GAZ and ZHID"]


def test_apply_zv_leaf_group_headers_spans_specific_and_normir_ingredients():
    from app.ecology.services.gis_tes_fuel_comparison_constants import (
        zv_normir_ingredient_header_group,
        zv_specific_emission_header_group,
    )

    specific_title, specific_keys, specific_labels = zv_specific_emission_header_group()
    normir_title, normir_keys, normir_labels = zv_normir_ingredient_header_group()
    keys = [
        "NAME",
        *specific_keys,
        "udel_na_energ",
        "normir_ZV_na_energ",
        *normir_keys,
    ]
    columns = [
        {
            "key": key,
            "source_index": index,
            "sticky": key == "NAME",
            "sticky_index": 0 if key == "NAME" else None,
        }
        for index, key in enumerate(keys)
    ]
    headers = [
        [
            {
                "text": f"old {key}",
                "column_key": key,
                "rowspan": 1,
                "colspan": 1,
                "col_index": index,
            }
            for index, key in enumerate(keys)
        ],
        [
            {
                "text": key,
                "column_key": key,
                "is_key": True,
                "rowspan": 1,
                "colspan": 1,
            }
            for key in keys
        ],
    ]
    apply_zv_leaf_group_headers(columns, headers)
    drop_key_header_row(headers)
    assert specific_title == "Удельные выбросы нормируемых ЗВ по ингредиентам, кг/т у.т."
    assert normir_title == "в том числе по ингредиентам"
    specific_group = next(cell for cell in headers[0] if cell["text"] == specific_title)
    normir_group = next(cell for cell in headers[0] if cell["text"] == normir_title)
    assert specific_group["colspan"] == 3
    assert specific_group["rowspan"] == 1
    assert normir_group["colspan"] == 3
    assert normir_group["rowspan"] == 1
    leaf_by_key = {cell["column_key"]: cell["text"] for cell in headers[1] if cell.get("column_key")}
    assert [leaf_by_key[key] for key in specific_keys] == [specific_labels[key] for key in specific_keys]
    assert [leaf_by_key[key] for key in normir_keys] == [normir_labels[key] for key in normir_keys]
    assert columns[keys.index("udel_tv")]["label"] == "твердое топливо"
    assert columns[keys.index("udel_SO2")]["label"] == "сера"
    assert columns[keys.index("udel_NOx")]["label"] == "азот"
    assert columns[keys.index("normir_TVERD_na_energ")]["label"] == "твердое топливо"
    assert columns[keys.index("normir_SO2_na_energ")]["label"] == "сера"
    assert columns[keys.index("normir_NOx_na_energ")]["label"] == "азот"
    name_cell = next(cell for cell in headers[0] if cell.get("column_key") == "NAME")
    assert name_cell["rowspan"] == 2
    mid_keys = {cell.get("column_key") for cell in headers[1]}
    assert "udel_na_energ" not in mid_keys
    assert "normir_ZV_na_energ" not in mid_keys


def test_apply_zv_analysis_banner_headers_spans_energy_specific_columns():
    from app.ecology.services.gis_tes_fuel_comparison_constants import (
        ZV_EE_SPECIFIC_KEY,
        ZV_ENERGY_SPECIFIC_BANNER_TITLE,
        ZV_ENERGY_SPECIFIC_LEAF_LABELS,
        zv_normir_ingredient_header_group,
        zv_specific_emission_header_group,
    )

    specific_title, specific_keys, _specific_labels = zv_specific_emission_header_group()
    normir_title, normir_keys, _normir_labels = zv_normir_ingredient_header_group()
    keys = [
        "NAME",
        *specific_keys,
        "udel_na_energ",
        ZV_EE_SPECIFIC_KEY,
        "normir_ZV_na_energ",
        *normir_keys,
    ]
    columns = [
        {
            "key": key,
            "source_index": index,
            "sticky": key == "NAME",
            "sticky_index": 0 if key == "NAME" else None,
        }
        for index, key in enumerate(keys)
    ]
    headers = [
        [
            {
                "text": f"old {key}",
                "column_key": key,
                "rowspan": 1,
                "colspan": 1,
                "col_index": index,
            }
            for index, key in enumerate(keys)
        ],
        [
            {
                "text": key,
                "column_key": key,
                "is_key": True,
                "rowspan": 1,
                "colspan": 1,
            }
            for key in keys
        ],
    ]
    apply_zv_leaf_group_headers(columns, headers)
    apply_zv_analysis_banner_headers(columns, headers)
    drop_key_header_row(headers)
    assert len(headers) == 3
    energy_banner = next(
        cell for cell in headers[0] if cell["text"] == ZV_ENERGY_SPECIFIC_BANNER_TITLE
    )
    assert energy_banner["colspan"] == 6
    assert energy_banner["rowspan"] == 1
    specific_group = next(cell for cell in headers[0] if cell["text"] == specific_title)
    assert specific_group["colspan"] == 3
    assert specific_group["rowspan"] == 2
    name_cell = next(cell for cell in headers[0] if cell.get("column_key") == "NAME")
    assert name_cell["rowspan"] == 3
    udel_cell = next(cell for cell in headers[1] if cell.get("column_key") == "udel_na_energ")
    ee_cell = next(cell for cell in headers[1] if cell.get("column_key") == ZV_EE_SPECIFIC_KEY)
    normir_cell = next(
        cell for cell in headers[1] if cell.get("column_key") == "normir_ZV_na_energ"
    )
    ingredient_group = next(cell for cell in headers[1] if cell["text"] == normir_title)
    assert udel_cell["text"] == ZV_ENERGY_SPECIFIC_LEAF_LABELS["udel_na_energ"]
    assert ee_cell["text"] == ZV_ENERGY_SPECIFIC_LEAF_LABELS[ZV_EE_SPECIFIC_KEY]
    assert normir_cell["text"] == ZV_ENERGY_SPECIFIC_LEAF_LABELS["normir_ZV_na_energ"]
    assert udel_cell["rowspan"] == 2
    assert ee_cell["rowspan"] == 2
    assert normir_cell["rowspan"] == 2
    assert ingredient_group["colspan"] == 3
    assert ingredient_group["rowspan"] == 1
    leaf_by_key = {
        cell["column_key"]: cell["text"] for cell in headers[2] if cell.get("column_key")
    }
    assert leaf_by_key["normir_TVERD_na_energ"] == "твердое топливо"
    assert leaf_by_key["normir_SO2_na_energ"] == "сера"
    assert leaf_by_key["normir_NOx_na_energ"] == "азот"
    assert "udel_na_energ" not in leaf_by_key
    assert columns[keys.index("udel_na_energ")]["label"] == udel_cell["text"]
    assert columns[keys.index("normir_ZV_na_energ")]["label"] == normir_cell["text"]


def test_apply_zv_fuel_consumption_group_headers():
    from app.ecology.services.gis_tes_fuel_comparison_constants import (
        FUEL_TCE_HEADER_META,
        ZV_FUEL_CONSUMPTION_BANNER_TITLE,
        zv_fuel_consumption_header_group,
    )

    title, fuel_keys, leaf_labels = zv_fuel_consumption_header_group()
    assert title == ZV_FUEL_CONSUMPTION_BANNER_TITLE
    assert title == "Расход топлива, т у.т."
    keys = ["NAME", *fuel_keys, "VIBROS_ekol"]
    columns = [
        {
            "key": key,
            "source_index": index,
            "sticky": key == "NAME",
            "sticky_index": 0 if key == "NAME" else None,
        }
        for index, key in enumerate(keys)
    ]
    headers = [
        [
            {
                "text": f"old {key}",
                "column_key": key,
                "rowspan": 1,
                "colspan": 1,
                "col_index": index,
            }
            for index, key in enumerate(keys)
        ],
        [
            {
                "text": key,
                "column_key": key,
                "is_key": True,
                "rowspan": 1,
                "colspan": 1,
            }
            for key in keys
        ],
    ]
    apply_zv_leaf_group_headers(columns, headers)
    apply_zv_analysis_banner_headers(columns, headers)
    drop_key_header_row(headers)
    fuel_banner = next(
        cell for cell in headers[0] if cell["text"] == ZV_FUEL_CONSUMPTION_BANNER_TITLE
    )
    assert fuel_banner["colspan"] == 4
    assert fuel_banner["rowspan"] == 2
    assert fuel_banner["fuel_tep_highlight"] is True
    leaf_by_key = {
        cell["column_key"]: cell for cell in headers[-1] if cell.get("column_key")
    }
    for key, label in leaf_labels.items():
        assert leaf_by_key[key]["text"] == label
        assert leaf_by_key[key]["access_code"] == FUEL_TCE_HEADER_META[key]["access_code"]
        assert leaf_by_key[key]["fuel_tep_highlight"] is True
        assert columns[keys.index(key)]["label"] == label
        assert columns[keys.index(key)]["access_code"] == FUEL_TCE_HEADER_META[key]["access_code"]
        assert columns[keys.index(key)]["fuel_tep_highlight"] is True
    assert leaf_by_key["GAZ_6tp"]["access_code"] == "GAZ + ISK_GAZ"
    assert leaf_by_key["GAZ_estestv"]["access_code"] == "GAZ"
    assert leaf_by_key["MAZUT_6tp"]["access_code"] == "MAZUT"
    assert leaf_by_key["TVERD_TOPL_6tp"]["access_code"] == "UGOL + TORF"


def test_apply_zv_display_column_order_swaps_vibros_and_count():
    columns = [
        {"key": "Teplo"},
        {"key": "VIBROS_ekol"},
        {"key": "VIBR_cnt"},
        {"key": "TVERD_V"},
    ]
    headers = [
        [
            {"text": "тепло", "column_key": "Teplo"},
            {"text": "Всего выброшено в атмосферу ЗВ за год, тонн", "column_key": "VIBROS_ekol"},
            {"text": "Количество отчитавшихся", "column_key": "VIBR_cnt"},
            {"text": "Твердые вещества", "column_key": "TVERD_V"},
        ]
    ]
    apply_zv_display_column_order(columns, headers)
    assert [col["key"] for col in columns] == ["Teplo", "VIBR_cnt", "VIBROS_ekol", "TVERD_V"]
    assert [cell["column_key"] for cell in headers[0]] == [
        "Teplo",
        "VIBR_cnt",
        "VIBROS_ekol",
        "TVERD_V",
    ]
    apply_zv_display_column_order(columns, headers)
    assert [col["key"] for col in columns] == ["Teplo", "VIBR_cnt", "VIBROS_ekol", "TVERD_V"]


def test_prepare_visible_table_hides_reported_columns_by_extra_keys():
    from app.ecology.services.gis_tes_fuel_comparison_constants import (
        REPORTED_TABLE_COLUMN_KEYS,
    )

    columns = [
        {"key": "NAME"},
        {"key": "nust"},
        {"key": "nust_vibr"},
        {"key": "W_en"},
        {"key": "W_en_vibr"},
        {"key": "Teplo"},
        {"key": "Teplo_vibr"},
        {"key": "VIBROS_ekol"},
        {"key": "VIBR_cnt"},
    ]
    headers = [
        [
            {"text": "NAME", "column_key": "NAME", "col_index": 0, "colspan": 1},
            {"text": "nust", "column_key": "nust", "col_index": 1, "colspan": 1},
            {"text": "nust_vibr", "column_key": "nust_vibr", "col_index": 2, "colspan": 1},
            {"text": "W_en", "column_key": "W_en", "col_index": 3, "colspan": 1},
            {"text": "W_en_vibr", "column_key": "W_en_vibr", "col_index": 4, "colspan": 1},
            {"text": "Teplo", "column_key": "Teplo", "col_index": 5, "colspan": 1},
            {"text": "Teplo_vibr", "column_key": "Teplo_vibr", "col_index": 6, "colspan": 1},
            {"text": "VIBROS_ekol", "column_key": "VIBROS_ekol", "col_index": 7, "colspan": 1},
            {"text": "VIBR_cnt", "column_key": "VIBR_cnt", "col_index": 8, "colspan": 1},
        ]
    ]
    hidden, hidden_headers, _sticky = prepare_visible_table(
        columns, headers, extra_hidden_keys=REPORTED_TABLE_COLUMN_KEYS
    )
    assert [col["key"] for col in hidden] == ["NAME", "nust", "W_en", "Teplo", "VIBROS_ekol"]
    assert [cell["column_key"] for cell in hidden_headers[0]] == [
        "NAME",
        "nust",
        "W_en",
        "Teplo",
        "VIBROS_ekol",
    ]
    shown, shown_headers, _sticky = prepare_visible_table(columns, headers)
    assert [col["key"] for col in shown] == [col["key"] for col in columns]
    assert [cell["column_key"] for cell in shown_headers[0]] == [col["key"] for col in columns]


def test_zv_vse_reported_columns_toggle(tmp_path, monkeypatch):
    path = tmp_path / "zv-reported.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ЗВ_всё"
    keys = [
        "Fed_okr",
        "nameoes",
        "oblname",
        "NAME",
        "numb1120",
        "Year",
        "nust",
        "nust_vibr",
        "W_en",
        "W_en_vibr",
        "Teplo",
        "Teplo_vibr",
        "VIBROS_ekol",
        "VIBR_cnt",
    ]
    labels = [
        "ФО",
        "ОЭС",
        "Субъект",
        "Станция",
        "Код",
        "Год",
        "Установленная мощность, МВт",
        "Уст. мощность отчитавшихся, МВт",
        "Выработка электроэнергии, млн. кВт.ч",
        "Выработка отчитавшихся, млн. кВт.ч",
        "Отпущенно тепла, тыс. Гкалл",
        "Тепло отчитавшихся, тыс. Гкалл",
        "всего",
        "Количество отчитавшихся",
    ]
    ws.append(keys)
    ws.append(labels)
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 10, 8, 20, 15, 30, 12, 40, 2])
    for title in ("ЗВ_без кот", "ПГ_без кот", "ПГ_всё", "Лист1"):
        wb.create_sheet(title)
        extra = wb[title]
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 1])
    wb.save(path)
    clear_comparison_cache()
    _patch_pg_groups_from_fuel(monkeypatch, [(1, 2017, "ТЭЦ")])
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {},
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {},
    )
    hidden = build_comparison_table_context(
        "zv-vse",
        start_year=2017,
        end_year=2017,
        rounding_digits=0,
        source_path=path,
    )
    assert hidden is not None
    assert hidden["show_reported"] is False
    assert hidden["has_reported_columns"] is True
    assert hidden["keep_query"]["show_reported"] == "0"
    assert [col["key"] for col in hidden["columns"]] == [
        "NAME",
        "numb1120",
        "Year",
        "nust",
        "nust_generation",
        "W_en",
        "W_en_generation",
        "Teplo",
        "station_affiliation",
        "VIBROS_ekol",
    ]
    shown = build_comparison_table_context(
        "zv-vse",
        start_year=2017,
        end_year=2017,
        rounding_digits=0,
        source_path=path,
        request_args={"show_reported": ["1"]},
    )
    assert shown is not None
    assert shown["show_reported"] is True
    assert shown["keep_query"]["show_reported"] == "1"
    assert [col["key"] for col in shown["columns"]] == [
        "NAME",
        "numb1120",
        "Year",
        "nust",
        "nust_generation",
        "nust_vibr",
        "W_en",
        "W_en_generation",
        "W_en_vibr",
        "Teplo",
        "station_affiliation",
        "Teplo_vibr",
        "VIBR_cnt",
        "VIBROS_ekol",
    ]
    header_by_key = {
        cell["column_key"]: cell["text"]
        for row in shown["header_rows"]
        for cell in row
        if cell.get("column_key")
    }
    assert header_by_key["nust_vibr"] == "Уст. мощность отчитавшихся, МВт"
    assert header_by_key["W_en_vibr"] == "Выработка отчитавшихся, млн. кВт.ч"
    assert header_by_key["Teplo_vibr"] == "Тепло отчитавшихся, тыс. Гкалл"
    assert header_by_key["VIBR_cnt"] == "Количество отчитавшихся"
    clear_comparison_cache()


def _station_generation_workbook(path):
    """«ЗВ_всё» с ТЭП-столбцами NUST/E — для столбцов сверки с «Генерацией»."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ЗВ_всё"
    ws.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "nust", "W_en"])
    ws.append(
        [
            "ФО",
            "ОЭС",
            "Субъект",
            "Станция",
            "Код",
            "Год",
            "Установленная мощность, МВт",
            "Выработка электроэнергии, млн. кВт.ч",
        ]
    )
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1 (котельная)", 708, 2017, 100, 500])
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1 (котельная)", 708, 2018, 100, 500])
    for title in ("ПГ_без кот", "ЗВ_без кот", "ПГ_всё", "Лист1"):
        wb.create_sheet(title)
        extra = wb[title]
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1 (котельная)", 708, 2017, 1])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1 (котельная)", 708, 2018, 1])
    wb.save(path)
    clear_comparison_cache()


def _station_generation_context(slug, path, overlay, monkeypatch, **kwargs):
    _patch_pg_groups_from_fuel(
        monkeypatch,
        [(708, 2017, "ТЭЦ-1 (котельная)"), (708, 2018, "ТЭЦ-1 (котельная)")],
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {},
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {
            (708, 2017): {"nust": 100, "e": 500, "q": 30, "ved": 1},
            (708, 2018): {"nust": 100, "e": 500, "q": 30, "ved": 1},
        },
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_station_generation_overlay",
        lambda numbs, years: overlay,
    )
    context = build_comparison_table_context(
        slug,
        start_year=kwargs.pop("start_year", 2017),
        end_year=kwargs.pop("end_year", 2018),
        rounding_digits=kwargs.pop("rounding_digits", 1),
        source_path=path,
        **kwargs,
    )
    clear_comparison_cache()
    return context


def test_zv_vse_station_generation_columns_follow_tep_columns(tmp_path, monkeypatch):
    path = tmp_path / "zv-generation.xlsx"
    _station_generation_workbook(path)
    context = _station_generation_context(
        "zv-vse",
        path,
        {
            "names_by_numb": {708: "ТЭЦ-1"},
            "station_links_by_numb": {708: [{"id": 501, "name": "ТЭЦ-1"}]},
            "values_by_numb_year": {
                (708, 2017): {
                    "nust_generation": Decimal("100.0"),
                    "W_en_generation": Decimal("500.5"),
                }
            },
        },
        monkeypatch,
    )
    assert context is not None
    keys = [column["key"] for column in context["columns"]]
    assert keys == [
        "NAME",
        "numb1120",
        "Year",
        "nust",
        "nust_generation",
        "W_en",
        "W_en_generation",
    ]
    assert context["station_check_pairs"] == {
        "nust_generation": "nust",
        "W_en_generation": "W_en",
    }
    header_by_key = {
        cell.get("column_key"): cell
        for row in context["header_rows"]
        for cell in row
        if cell.get("column_key")
    }
    assert header_by_key["nust_generation"]["text"] == "Руст, МВт"
    assert header_by_key["W_en_generation"]["text"] == "Выработка ЭЭ, млн кВтч (Генерация)"
    assert header_by_key["nust_generation"]["station_generation_highlight"] is True
    assert header_by_key["nust_generation"]["rowspan"] == len(context["header_rows"])
    assert header_by_key["W_en_generation"]["rowspan"] == len(context["header_rows"])

    rows = context["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]["rows"]
    group = context["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]
    assert group["station_links"] == [{"id": 501, "name": "ТЭЦ-1"}]
    assert [row["year"] for row in rows] == [2017, 2018]
    by_key = {key: rows[0]["cells"][index] for index, key in enumerate(keys)}
    assert by_key["NAME"] == "ТЭЦ-1"
    assert by_key["nust"] == "100"
    assert by_key["nust_generation"] == "100"
    assert by_key["W_en"] == "500"
    assert by_key["W_en_generation"] == "500,5"
    assert rows[0]["cell_mismatch"][keys.index("nust_generation")] is False
    assert rows[0]["cell_mismatch"][keys.index("W_en_generation")] is True
    # Год без станционных данных: ячейки пустые и не краснеют; название станции есть.
    assert rows[1]["cells"][keys.index("nust_generation")] == ""
    assert rows[1]["cells"][keys.index("W_en_generation")] == ""
    assert rows[1]["cells"][keys.index("NAME")] == "ТЭЦ-1"
    assert rows[1]["cell_mismatch"][keys.index("nust_generation")] is False


def test_station_generation_columns_absent_on_bez_kot_sheets(tmp_path, monkeypatch):
    path = tmp_path / "zv-generation-none.xlsx"
    _station_generation_workbook(path)
    context = _station_generation_context(
        "zv-bez-kot",
        path,
        {
            "names_by_numb": {708: "ТЭЦ-1"},
            "values_by_numb_year": {(708, 2017): {"nust_generation": 100.0}},
        },
        monkeypatch,
    )
    assert context is not None
    keys = [column["key"] for column in context["columns"]]
    assert "nust_generation" not in keys
    assert "W_en_generation" not in keys
    assert context["station_check_pairs"] == {}
    rows = context["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]["rows"]
    assert rows[0]["cells"][keys.index("NAME")] == "ТЭЦ-1 (котельная)"


def test_apply_station_generation_overlay_sets_name_and_values():
    columns = [
        {"key": "NAME", "numeric": False, "source_index": 0},
        {"key": "numb1120", "numeric": True, "source_index": 1},
        {"key": "Year", "numeric": True, "source_index": 2},
        {"key": "nust_generation", "numeric": True, "source_index": 3},
    ]
    filled = ["ТЭЦ-1 (котельная)", 708, 2017, None]
    overlay = {
        "names_by_numb": {708: "ТЭЦ-1"},
        "values_by_numb_year": {(708, 2017): {"nust_generation": Decimal("95.5")}},
    }
    assert apply_station_generation_overlay(filled, columns, overlay) == [
        "ТЭЦ-1",
        708,
        2017,
        Decimal("95.5"),
    ]
    # Группы с другой нумерацией и строки без года не меняются.
    assert apply_station_generation_overlay(
        ["Котельные Коми", 38, 2017, None], columns, overlay
    ) == ["Котельные Коми", 38, 2017, None]


def test_station_generation_mismatch_flags_only_filled_cells():
    columns = [
        {"key": "nust", "numeric": True},
        {"key": "nust_generation", "numeric": True},
        {"key": "W_en", "numeric": True},
        {"key": "W_en_generation", "numeric": True},
    ]
    pairs = station_generation_check_pairs(columns)
    assert pairs == {"nust_generation": "nust", "W_en_generation": "W_en"}
    assert station_generation_mismatch(
        ["100", "90", "500", "500"], columns, pairs
    ) == [False, True, False, False]
    assert station_generation_mismatch(
        ["100", "", "500", ""], columns, pairs
    ) == [False, False, False, False]
    # Столбца макета для сверки нет — красного шрифта тоже.
    assert station_generation_mismatch(
        ["90"], [{"key": "nust_generation", "numeric": True}], pairs
    ) == [False]


def test_load_station_generation_overlay_without_app_context():
    assert load_station_generation_overlay({708}, [2017]) == {}
    assert load_station_generation_overlay(set(), [2017]) == {}
    assert load_station_generation_overlay({708}, []) == {}


def _station_affiliation_workbook(path):
    """«ЗВ_всё» с ТЭП-столбцом «Отпуск ТЭ» — для столбца «Принадлежность»."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ЗВ_всё"
    ws.append(
        ["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "nust", "W_en", "Teplo"]
    )
    ws.append(
        [
            "ФО",
            "ОЭС",
            "Субъект",
            "Станция",
            "Код",
            "Год",
            "Установленная мощность, МВт",
            "Выработка электроэнергии, млн. кВт.ч",
            "Отпуск ТЭ, тыс.Гкал",
        ]
    )
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1 (котельная)", 708, 2017, 100, 500, 30])
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1 (котельная)", 708, 2018, 100, 500, 30])
    for title in ("ПГ_без кот", "ЗВ_без кот", "ПГ_всё", "Лист1"):
        wb.create_sheet(title)
        extra = wb[title]
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en", "Teplo"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1 (котельная)", 708, 2017, 1, 30])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1 (котельная)", 708, 2018, 1, 30])
    wb.save(path)
    clear_comparison_cache()


def _station_affiliation_context(slug, path, affiliations, monkeypatch):
    _patch_pg_groups_from_fuel(
        monkeypatch,
        [(708, 2017, "ТЭЦ-1 (котельная)"), (708, 2018, "ТЭЦ-1 (котельная)")],
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {},
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {
            (708, 2017): {"nust": 100, "e": 500, "q": 30, "ved": 1},
            (708, 2018): {"nust": 100, "e": 500, "q": 30, "ved": 1},
        },
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_station_affiliations_by_numb_year",
        lambda numbs, years: affiliations,
    )
    context = build_comparison_table_context(
        slug,
        start_year=2017,
        end_year=2018,
        rounding_digits=1,
        source_path=path,
    )
    clear_comparison_cache()
    return context


def test_zv_vse_station_affiliation_column_follows_teplo_column(tmp_path, monkeypatch):
    path = tmp_path / "zv-affiliation.xlsx"
    _station_affiliation_workbook(path)
    context = _station_affiliation_context(
        "zv-vse", path, {(708, 2017): "отрасль"}, monkeypatch
    )
    assert context is not None
    keys = [column["key"] for column in context["columns"]]
    assert keys == [
        "NAME",
        "numb1120",
        "Year",
        "nust",
        "nust_generation",
        "W_en",
        "W_en_generation",
        "Teplo",
        "station_affiliation",
    ]
    assert keys.index("Teplo") + 1 == keys.index("station_affiliation")
    assert context["columns"][keys.index("station_affiliation")]["numeric"] is False
    header = {
        cell.get("column_key"): cell
        for row in context["header_rows"]
        for cell in row
        if cell.get("column_key")
    }
    assert header["station_affiliation"]["text"] == "Принадлежность"
    assert header["station_affiliation"]["rowspan"] == len(context["header_rows"])
    assert "station_generation_highlight" not in header["station_affiliation"]

    rows = context["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]["rows"]
    assert [row["year"] for row in rows] == [2017, 2018]
    assert rows[0]["cells"][keys.index("station_affiliation")] == "отрасль"
    # Год без данных ГИС ТЭК — «—», как в столбце на страницах ГИС ТЭК.
    assert rows[1]["cells"][keys.index("station_affiliation")] == "—"


def test_station_affiliation_column_absent_on_bez_kot_sheets(tmp_path, monkeypatch):
    path = tmp_path / "zv-affiliation-none.xlsx"
    _station_affiliation_workbook(path)
    context = _station_affiliation_context(
        "zv-bez-kot", path, {(708, 2017): "отрасль"}, monkeypatch
    )
    assert context is not None
    keys = [column["key"] for column in context["columns"]]
    assert "station_affiliation" not in keys
    assert all(
        cell.get("column_key") != "station_affiliation"
        for row in context["header_rows"]
        for cell in row
    )


def test_apply_station_affiliation_cells_fills_unknown_with_dash():
    columns = [
        {"key": "NAME", "numeric": False},
        {"key": "numb1120", "numeric": True},
        {"key": "Year", "numeric": True},
        {"key": "station_affiliation", "numeric": False},
    ]
    affiliations = {(708, 2017): "промпредприятие"}
    assert apply_station_affiliation_cells(
        ["ТЭЦ-1", 708, 2017, None], columns, affiliations
    ) == ["ТЭЦ-1", 708, 2017, "промпредприятие"]
    # Нет данных по группе/году, нет года, нет столбца — «—» или строка без изменений.
    assert apply_station_affiliation_cells(
        ["ТЭЦ-1", 999, 2017, None], columns, affiliations
    ) == ["ТЭЦ-1", 999, 2017, "—"]
    assert apply_station_affiliation_cells(
        ["ТЭЦ-1", 708, None, None], columns, affiliations
    ) == ["ТЭЦ-1", 708, None, "—"]
    assert apply_station_affiliation_cells(
        ["ТЭЦ-1", 708, 2017, None],
        [column for column in columns if column["key"] != "station_affiliation"],
        affiliations,
    ) == ["ТЭЦ-1", 708, 2017, None]


def test_load_station_affiliations_without_app_context():
    assert load_station_affiliations_by_numb_year({708}, [2017]) == {}
    assert load_station_affiliations_by_numb_year(set(), [2017]) == {}
    assert load_station_affiliations_by_numb_year({708}, []) == {}


def test_apply_pg_ecology_overlay():
    # Year читается по фиксированному индексу YEAR_COL_INDEX (=5).
    columns = [
        {"key": "Fed_okr"},
        {"key": "nameoes"},
        {"key": "oblast"},
        {"key": "Stancia"},
        {"key": "numb1120"},
        {"key": "Year"},
        {"key": "GAZ_ek"},
        {"key": "CO2_gaz"},
        {"key": "MAZUT_ek"},
        {"key": "CO2_mazut"},
        {"key": "TvTopl_ek"},
        {"key": "CO2_ugol"},
        {"key": "d_gaz"},
    ]
    raw = ["СЗ", "ОЭС", "АО", "ТЭЦ", 708, 2017, 1, 2, 3, 4, 5, 6, 999]
    overlay = {
        (708, 2017): {
            "gas_consumption_tce": 100.5,
            "co2_from_gas_t": 200.5,
            "oil_fuel_consumption_tce": 10,
            "co2_from_oil_fuel_t": 20,
            "solid_fuel_consumption_tce": 30,
            "co2_from_solid_fuel_t": 40,
        }
    }
    out = apply_pg_ecology_overlay(raw, columns, overlay)
    assert out[6:12] == [100.5, 200.5, 10, 20, 30, 40]
    assert out[12] == 999
    assert apply_pg_ecology_overlay(raw, columns, {}) == raw


def test_apply_zv_ecology_overlay():
    columns = [
        {"key": "Fed_okr"},
        {"key": "nameoes"},
        {"key": "oblast"},
        {"key": "Stancia"},
        {"key": "numb1120"},
        {"key": "Year"},
        {"key": "TVERD_V"},
        {"key": "SO2_gaz"},
        {"key": "CO_gaz"},
        {"key": "NOx_gaz"},
        {"key": "UGLEVOD"},
        {"key": "LOS"},
        {"key": "PROCH_ekol"},
        {"key": "GAZ and ZHID"},
    ]
    raw = ["СЗ", "ОЭС", "АО", "ТЭЦ", 708, 2017, 1, 2, 3, 4, 5, 6, 7, 999]
    overlay = {
        (708, 2017): {
            "solid_substances": 10.5,
            "sulfur_dioxide": 20.5,
            "carbon_monoxide": 1,
            "nitrogen_oxides": 2,
            "hydrocarbons": 3,
            "voc": 4,
            "other_gaseous_and_liquid": 5,
        }
    }
    out = apply_zv_ecology_overlay(raw, columns, overlay)
    assert out[6:13] == [10.5, 20.5, 1, 2, 3, 4, 5]
    assert out[13] == 999
    assert apply_zv_ecology_overlay(raw, columns, {}) == raw


def test_pg_pages_use_ecology_overlay_before_differences(tmp_path, monkeypatch):
    path = tmp_path / "cmp.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ПГ_без кот"
    ws.append(
        [
            "Fed_okr",
            "nameoes",
            "oblast",
            "Stancia",
            "numb1120",
            "Year",
            "GAZ_ESTEST",
            "GAZ_ek",
            "d_gaz",
            "CO2_gaz",
            "MAZUT_6tp",
            "MAZUT_ek",
            "d_mazut",
            "CO2_mazut",
            "TvTopl_6tp",
            "TvTopl_ek",
            "d_tv",
            "CO2_ugol",
        ]
    )
    ws.append(
        [
            "ФО",
            "ОЭС",
            "Субъект",
            "Станция",
            "Код",
            "Год",
            "Газ естественный",
            "Газ макет",
            "Разность газ",
            "СО2 газ",
            "Мазут 6ТП",
            "Мазут макет",
            "Разность мазут",
            "СО2 мазут",
            "Твердое 6ТП",
            "Твердое макет",
            "Разность твердое",
            "СО2 уголь",
        ]
    )
    ws.append(
        [
            "СЗ",
            "ОЭС СЗ",
            "АО",
            "ТЭЦ-1",
            708,
            2017,
            50.0,
            1.0,
            999.0,
            2.0,
            100.0,
            3.0,
            999.0,
            4.0,
            200.0,
            5.0,
            999.0,
            6.0,
        ]
    )
    for title in ("ЗВ_без кот", "ЗВ_всё", "ПГ_всё", "Лист1"):
        extra = wb.create_sheet(title)
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 1])
    wb.save(path)
    clear_comparison_cache()
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {},
    )
    _patch_pg_groups_from_fuel(monkeypatch, [(708, 2017, "ТЭЦ-1")])
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {
            (708, 2017): {
                "energy_system_type_id": 1,
                "est_name": "ЕЭС",
                "ues_id": 2,
                "ues": "ОЭС",
                "res_id": 3,
                "res_name": "РЭС",
                "station_name": "ТЭЦ-1",
                "ved": 1,
                "gaz_estest_tce": 50.0,
                "mazut_tce": 100.0,
                "ugol_tce": 200.0,
            }
        },
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_pg_ecology_overlay_by_numb_year",
        lambda sheet_slug, numbs, years: {
            (708, 2017): {
                "gas_consumption_tce": 20.0,
                "co2_from_gas_t": 80.0,
                "oil_fuel_consumption_tce": 30.0,
                "co2_from_oil_fuel_t": 90.0,
                "solid_fuel_consumption_tce": 50.0,
                "co2_from_solid_fuel_t": 110.0,
            }
        },
    )
    ctx = build_comparison_table_context(
        "pg-bez-kot",
        start_year=2017,
        end_year=2017,
        rounding_digits=1,
        source_path=path,
    )
    assert ctx is not None
    group = ctx["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]
    by_key = {col["key"]: group["rows"][0]["cells"][i] for i, col in enumerate(ctx["columns"])}
    assert by_key["GAZ_ek"] == "20"
    assert by_key["CO2_gaz"] == "80"
    assert by_key["MAZUT_ek"] == "30"
    assert by_key["CO2_mazut"] == "90"
    assert by_key["TvTopl_ek"] == "50"
    assert by_key["CO2_ugol"] == "110"
    assert by_key["d_gaz"] == "30"
    assert by_key["d_mazut"] == "70"
    assert by_key["d_tv"] == "150"
    editable = {
        col["key"] for col in ctx["columns"] if col.get("editable")
    }
    assert editable == {
        "GAZ_ek",
        "CO2_gaz",
        "MAZUT_ek",
        "CO2_mazut",
        "TvTopl_ek",
        "CO2_ugol",
    }
    clear_comparison_cache()


def test_ved_filter_parsing_and_matching():
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        parse_territory_filters,
        row_matches_territory_filters,
    )

    parsed = parse_territory_filters({"ved_filter": ["industry", "junk"]})
    assert parsed["ved_filter"] == "industry"
    assert parse_territory_filters({"ved_filter": [""]})["ved_filter"] == ""
    columns = [{"key": "Stancia"}]
    assert row_matches_territory_filters(
        ["ТЭЦ"], columns, {"ved": 1}, {"ved_filter": "industry"}
    )
    assert not row_matches_territory_filters(
        ["ТЭЦ"], columns, {"ved": 3}, {"ved_filter": "industry"}
    )
    assert row_matches_territory_filters(
        ["ТЭЦ"], columns, {"ved": 3}, {"ved_filter": "industrial"}
    )
    assert not row_matches_territory_filters(
        ["ТЭЦ"], columns, {"ved": None}, {"ved_filter": "industry"}
    )
    assert row_matches_territory_filters(["ТЭЦ"], columns, {"ved": 99}, {"ved_filter": ""})


def test_header_station_and_affiliation_filters(monkeypatch):
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        _keep_query_args,
        apply_header_filters,
        parse_territory_filters,
    )

    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_station_affiliations_by_numb_year",
        lambda numbs, years: {
            (1, 2023): "отрасль",
            (2, 2023): "промпредприятие",
            (3, 2023): "—",
        },
    )
    columns = [{"key": "Stancia"}, {"key": "station_affiliation"}]
    groups = [
        {"numb": 1, "meta": {"station_name": "ТЭЦ-1"}},
        {"numb": 2, "meta": {"station_name": "ТЭЦ-2"}},
        {"numb": None, "meta": {"station_name": ""}},
    ]

    kept, station_choices, affiliation_choices = apply_header_filters(
        groups, columns, {}, [2023]
    )
    assert [group["numb"] for group in kept] == [1, 2, None]
    assert station_choices == ["ТЭЦ-1", "ТЭЦ-2"]
    # «—» (нет данных ГИС ТЭК) в списке принадлежности не предлагается.
    assert affiliation_choices == ["отрасль", "промпредприятие"]

    by_station, _, _ = apply_header_filters(
        groups, columns, {"station_header_filter": "ТЭЦ-2"}, [2023]
    )
    assert [group["numb"] for group in by_station] == [2]

    by_affiliation, _, _ = apply_header_filters(
        groups, columns, {"affiliation_filter": "промпредприятие"}, [2023]
    )
    assert [group["numb"] for group in by_affiliation] == [2]

    unknown, _, _ = apply_header_filters(
        groups, columns, {"affiliation_filter": "—"}, [2023]
    )
    assert [group["numb"] for group in unknown] == [None]

    # Группы без данных ГИС ТЭК дают пустой список принадлежностей.
    only_unknown, _, no_choices = apply_header_filters(
        [
            {"numb": 3, "meta": {"station_name": "ТЭЦ-3"}},
            {"numb": None, "meta": {"station_name": ""}},
        ],
        columns,
        {},
        [2023],
    )
    assert [group["numb"] for group in only_unknown] == [3, None]
    assert no_choices == []

    # Без столбца «Принадлежность» фильтр не применяется и вариантов нет.
    without_column, _, choices = apply_header_filters(
        groups, [{"key": "Stancia"}], {"affiliation_filter": "отрасль"}, [2023]
    )
    assert [group["numb"] for group in without_column] == [1, 2, None]
    assert choices == []

    parsed = parse_territory_filters(
        {
            "station_header_filter": [" Архангельская ТЭЦ "],
            "affiliation_filter": ["отрасль"],
        }
    )
    assert parsed["station_header_filter"] == "Архангельская ТЭЦ"
    assert parsed["affiliation_filter"] == "отрасль"
    empty = parse_territory_filters(None)
    assert empty["station_header_filter"] == ""
    assert empty["affiliation_filter"] == ""

    keep = _keep_query_args(
        2023,
        2024,
        1,
        None,
        {
            "station_header_filter": "Архангельская ТЭЦ",
            "affiliation_filter": "отрасль",
            "show_totals": False,
            "show_empty_rows": False,
            "show_reported": False,
        },
    )
    assert keep["station_header_filter"] == "Архангельская ТЭЦ"
    assert keep["affiliation_filter"] == "отрасль"


def test_safe_subtract():
    assert _safe_subtract(100, 30) == 70.0
    assert _safe_subtract(10.5, 10.5) == 0.0
    assert _safe_subtract(None, 10) is None
    assert _safe_subtract(10, None) is None
    assert _safe_subtract("", 10) is None
    assert _safe_subtract(10, "") is None
    assert _safe_subtract("abc", 10) is None
    assert _safe_subtract(True, 10) is None
    assert _safe_subtract(10, True) is None


def test_safe_math_ignores_source_float_noise():
    """Хвост double из БД не должен попадать в расчётные столбцы страницы."""
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        _format_value,
        _safe_divide,
        _safe_sum,
    )

    # d_mazut: 3121 − 3120,2000000000003 (так значение лежит в Numeric(36,16)).
    assert _safe_subtract(3121, Decimal("3120.2000000000003")) == 0.8
    assert _safe_sum(Decimal("1604.1000000000001"), Decimal("1582.1")) == 3186.2

    # K_gaz_ek = CO2 / расход газа: в «Не округлять» — только значащие цифры.
    ratio = _safe_divide(Decimal("1275805.8000000000000"), Decimal("802393.6000000001"))
    assert _format_value(ratio, True, 0) == "1,58999997008949"


def test_format_value_keeps_significant_digits_when_not_rounding():
    from app.ecology.services.gis_tes_fuel_comparison_services import _format_value

    assert _format_value(Decimal("802393.6000000001"), True, 0) == "802 393,6"
    assert _format_value(Decimal("3120.2000000000003"), True, 0) == "3 120,2"
    assert _format_value(Decimal("1275805.8000000000"), True, 0) == "1 275 805,8"
    assert _format_value(Decimal("0.00008782455212957862"), True, 0) == "0,0000878245521295786"
    # Округление до заданных знаков не изменилось.
    assert _format_value(Decimal("1275805.8000000000000"), True, 2) == "1 275 805,8"
    assert _format_value(Decimal("1275805.8000000000000"), True, 1) == "1 275 805,8"


def test_cell_titles_show_value_whole_regardless_of_rounding():
    """Подсказка к ячейке — значение целиком, как в БД, без учёта «Округл»."""
    from app.ecology.services.gis_tes_fuel_comparison_services import cell_titles

    columns = [
        {"key": "numb1120", "numeric": True, "source_index": 0},
        {"key": "Year", "numeric": True, "source_index": 1},
        {"key": "NAME", "numeric": False, "source_index": 2},
        {"key": "GAZ_ek", "numeric": True, "source_index": 3},
        {"key": "K_gaz_ek", "numeric": True, "source_index": 4},
        {"key": "Teplo", "numeric": True, "source_index": 5},
    ]
    filled = [
        1,
        2024,
        "Архангельская ТЭЦ",
        Decimal("802393.6000000001"),
        Decimal("1.589999970089492"),
        None,
    ]
    assert cell_titles(filled, columns) == [
        "",
        "",
        "",
        "802 393,6000000001",
        "1,589999970089492",
        "",
    ]


def test_safe_divide():
    from app.ecology.services.gis_tes_fuel_comparison_services import _safe_divide

    assert _safe_divide(100, 50) == 2.0
    assert abs(_safe_divide(1373100.2, 863585) - 1.59000005789818) < 1e-9
    assert _safe_divide(10, 0) is None
    assert _safe_divide(None, 10) is None
    assert _safe_divide(10, None) is None
    assert _safe_divide(True, 10) is None


def test_compute_emission_coefficients():
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        _compute_emission_coefficients,
    )

    columns = [
        {"key": "GAZ_ek"},
        {"key": "CO2_gaz"},
        {"key": "K_gaz_ek"},
        {"key": "MAZUT_ek"},
        {"key": "CO2_mazut"},
        {"key": "K_mazut_ek"},
        {"key": "TvTopl_ek"},
        {"key": "CO2_ugol"},
        {"key": "K_ugol_ek"},
    ]
    filled = [100.0, 159.0, 999, 50.0, 113.0, 999, 200.0, 540.0, 999]
    result = _compute_emission_coefficients(filled, columns)
    assert abs(result[2] - 1.59) < 1e-9
    assert abs(result[5] - 2.26) < 1e-9
    assert abs(result[8] - 2.7) < 1e-9
    filled2 = [0.0, 100.0, 999, None, 50.0, 999]
    result2 = _compute_emission_coefficients(filled2, columns[:6])
    assert result2[2] is None
    assert result2[5] is None


def test_compute_total_co2():
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        _compute_total_co2,
        _safe_sum,
    )

    assert _safe_sum(10, 20, 30) == 60.0
    assert _safe_sum(10, None, 30) == 40.0
    assert _safe_sum(None, None) is None
    assert _safe_sum(10, "x") is None

    columns = [
        {"key": "CO2_gaz"},
        {"key": "CO2_mazut"},
        {"key": "CO2_ugol"},
        {"key": "CO2"},
    ]
    result = _compute_total_co2([100.0, 50.0, 25.0, 999], columns)
    assert result[3] == 175.0
    result_partial = _compute_total_co2([100.0, None, 25.0, 999], columns)
    assert result_partial[3] == 125.0
    result_empty = _compute_total_co2([None, None, None, 999], columns)
    assert result_empty[3] is None


def test_compute_gaseous_and_liquid():
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        _compute_gaseous_and_liquid,
    )

    columns = [
        {"key": "GAZ and ZHID"},
        {"key": "SO2_gaz"},
        {"key": "CO_gaz"},
        {"key": "NOx_gaz"},
        {"key": "UGLEVOD"},
        {"key": "LOS"},
        {"key": "PROCH_ekol"},
    ]
    result = _compute_gaseous_and_liquid([999, 1, 2, 3, 4, 5, 6], columns)
    assert result[0] == 21.0
    result_partial = _compute_gaseous_and_liquid([999, 10, None, 5, None, None, None], columns)
    assert result_partial[0] == 15.0
    result_empty = _compute_gaseous_and_liquid([999, None, None, None, None, None, None], columns)
    assert result_empty[0] is None


def test_compute_zv_total():
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        _compute_zv_total,
    )

    columns = [
        {"key": "VIBROS_ekol"},
        {"key": "TVERD_V"},
        {"key": "GAZ and ZHID"},
    ]
    result = _compute_zv_total([999, 10, 21], columns)
    assert result[0] == 31.0
    result_partial = _compute_zv_total([999, 10, None], columns)
    assert result_partial[0] == 10.0
    result_empty = _compute_zv_total([999, None, None], columns)
    assert result_empty[0] is None


def test_compute_specific_co2():
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        _brought_energy_mln_kwh,
        _compute_specific_co2,
    )

    assert abs(_brought_energy_mln_kwh(1810.029, 2697.017) - (1810.029 + 2697.017 / 859.8 * 1000)) < 1e-9
    assert _brought_energy_mln_kwh(None, None) is None

    columns = [
        {"key": "GAZ_ek"},
        {"key": "MAZUT_ek"},
        {"key": "TvTopl_ek"},
        {"key": "CO2"},
        {"key": "W_en"},
        {"key": "Teplo"},
        {"key": "UD_CO2_ek"},
        {"key": "UD_CO2_ener"},
    ]
    filled = [863585, 18021, 0, 1413926.1, 1810.029, 2697.017, 999, 999]
    result = _compute_specific_co2(filled, columns)
    assert abs(result[6] - 1603.807256302702) < 1e-6
    assert abs(result[7] - 1413926.1 / (1810.029 + 2697.017 / 859.8 * 1000)) < 1e-9


def test_compute_zv_specific_metrics_matches_layout_formulas():
    from app.ecology.services.gis_tes_fuel_comparison_constants import ZV_EE_SPECIFIC_KEY
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        _brought_energy_mln_kwh,
        _compute_derived_columns,
    )

    columns = [
        {"key": "TVERD_V"},
        {"key": "SO2_gaz"},
        {"key": "NOx_gaz"},
        {"key": "CO_gaz"},
        {"key": "UGLEVOD"},
        {"key": "LOS"},
        {"key": "PROCH_ekol"},
        {"key": "GAZ and ZHID"},
        {"key": "VIBROS_ekol"},
        {"key": "W_en"},
        {"key": "Teplo"},
        {"key": "GAZ_6tp"},
        {"key": "MAZUT_6tp"},
        {"key": "TVERD_TOPL_6tp"},
        {"key": "energy"},
        {"key": "ZV_normir"},
        {"key": "udel_na_energ"},
        {"key": ZV_EE_SPECIFIC_KEY},
        {"key": "normir_ZV_na_energ"},
        {"key": "normir_TVERD_na_energ"},
        {"key": "normir_SO2_na_energ"},
        {"key": "normir_NOx_na_energ"},
        {"key": "udel_tv"},
        {"key": "udel_SO2"},
        {"key": "udel_NOx"},
    ]
    filled = [
        3.13,
        778.623,
        2883.29,
        366.529,
        0,
        0,
        19.863,
        999,
        999,
        1810.029,
        2697.017,
        863585,
        18021,
        0,
        999,
        999,
        999,
        999,
        999,
        999,
        999,
        999,
        999,
        999,
        999,
    ]
    result = _compute_derived_columns(filled, columns)
    energy = _brought_energy_mln_kwh(1810.029, 2697.017)
    gaseous = 778.623 + 366.529 + 2883.29 + 0 + 0 + 19.863
    total_zv = 3.13 + gaseous
    zv_normir = 3.13 + 778.623 + 2883.29
    by_key = {col["key"]: result[index] for index, col in enumerate(columns)}
    assert abs(by_key["GAZ and ZHID"] - gaseous) < 1e-9
    assert abs(by_key["VIBROS_ekol"] - total_zv) < 1e-9
    assert abs(by_key["energy"] - energy) < 1e-9
    assert abs(by_key["ZV_normir"] - zv_normir) < 1e-9
    assert abs(by_key["udel_na_energ"] - total_zv / energy) < 1e-9
    assert abs(by_key[ZV_EE_SPECIFIC_KEY] - zv_normir / 1810.029) < 1e-9
    assert abs(by_key["normir_ZV_na_energ"] - zv_normir / energy) < 1e-9
    assert abs(by_key["normir_TVERD_na_energ"] - 3.13 / energy) < 1e-9
    assert abs(by_key["normir_SO2_na_energ"] - 778.623 / energy) < 1e-9
    assert abs(by_key["normir_NOx_na_energ"] - 2883.29 / energy) < 1e-9
    assert abs(by_key["udel_tv"] - 3.13 * 1000 / 18021) < 1e-9
    assert abs(by_key["udel_SO2"] - 778.623 * 1000 / 18021) < 1e-9
    assert abs(by_key["udel_NOx"] - 2883.29 * 1000 / (863585 + 18021)) < 1e-9


def test_ensure_zv_electricity_specific_column_inserts_after_udel_na_energ():
    from app.ecology.services.gis_tes_fuel_comparison_constants import (
        ZV_EE_SPECIFIC_KEY,
        ZV_ENERGY_SPECIFIC_LEAF_LABELS,
    )
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        ensure_zv_electricity_specific_column,
    )

    columns = [
        {"key": "udel_na_energ", "numeric": True, "sticky": False},
        {"key": "normir_ZV_na_energ", "numeric": True, "sticky": False},
    ]
    headers = [
        [
            {"text": "old", "column_key": "udel_na_energ", "col_index": 0, "colspan": 1},
            {"text": "old2", "column_key": "normir_ZV_na_energ", "col_index": 1, "colspan": 1},
        ]
    ]
    rows = [[1.0, 2.0]]
    ensure_zv_electricity_specific_column(columns, headers, rows)
    assert [col["key"] for col in columns] == [
        "udel_na_energ",
        ZV_EE_SPECIFIC_KEY,
        "normir_ZV_na_energ",
    ]
    assert rows == [[1.0, None, 2.0]]
    assert headers[0][1]["column_key"] == ZV_EE_SPECIFIC_KEY
    assert headers[0][1]["text"] == ZV_ENERGY_SPECIFIC_LEAF_LABELS[ZV_EE_SPECIFIC_KEY]
    assert headers[0][2]["col_index"] == 2
    ensure_zv_electricity_specific_column(columns, headers, rows)
    assert [col["key"] for col in columns] == [
        "udel_na_energ",
        ZV_EE_SPECIFIC_KEY,
        "normir_ZV_na_energ",
    ]


def test_region_summary_sums_and_attaches_when_region_complete():
    from app.ecology.services.gis_tes_fuel_comparison_services import (
        attach_region_summaries,
        build_region_summary_block,
        parse_show_empty_rows_arg,
        parse_show_reported_arg,
        parse_show_totals_arg,
    )

    assert parse_show_totals_arg({"show_totals": ["1"]}) is True
    assert parse_show_totals_arg({"show_totals": ["0"]}) is False
    assert parse_show_totals_arg({}) is False
    assert parse_show_empty_rows_arg({}) is False
    assert parse_show_empty_rows_arg({"show_empty_rows": ["1"]}) is True
    assert parse_show_empty_rows_arg({"show_empty_rows": ["0"]}) is False
    assert parse_show_reported_arg({}) is False
    assert parse_show_reported_arg({"show_reported": ["1"]}) is True
    assert parse_show_reported_arg({"show_reported": ["0"]}) is False

    source_columns = [
        {"key": "NAME", "numeric": False},
        {"key": "Year", "numeric": True},
        {"key": "GAZ_ek", "numeric": True},
        {"key": "CO2_gaz", "numeric": True},
        {"key": "MAZUT_ek", "numeric": True},
        {"key": "CO2_mazut", "numeric": True},
        {"key": "TvTopl_ek", "numeric": True},
        {"key": "CO2_ugol", "numeric": True},
        {"key": "CO2", "numeric": True},
    ]
    visible = [
        {**col, "source_index": i, "diff_highlight": False}
        for i, col in enumerate(source_columns)
    ]
    g1 = {
        "id": 1,
        "meta": {"est_id": 10, "ues_id": 20, "res_id": 30},
        "filled_by_year": {
            2017: ["A", 2017, 10.0, 20.0, 1.0, 2.0, 0.0, 0.0, 999],
        },
    }
    g2 = {
        "id": 2,
        "meta": {"est_id": 10, "ues_id": 20, "res_id": 30},
        "filled_by_year": {
            2017: ["B", 2017, 5.0, 7.0, 3.0, 4.0, 0.0, 0.0, 999],
        },
    }
    block = build_region_summary_block(
        [g1, g2],
        "РЭС-1",
        source_columns=source_columns,
        visible_columns=visible,
        digits=1,
        key_prefix="sum-res",
    )
    assert block is not None
    assert block["label"] == "РЭС-1, всего"
    cells = block["rows"][0]["cells"]
    assert cells[0] == "РЭС-1, всего"
    assert cells[2] == "15"
    assert cells[3] == "27"
    assert cells[8] == "33"

    hierarchy = [
        {
            "est_id": 10,
            "est_name": "ЕЭС",
            "ues_list": [
                {
                    "ues_id": 20,
                    "ues_name": "ОЭС",
                    "res_list": [
                        {
                            "res_id": 30,
                            "res_name": "РЭС-1",
                            "groups": [g1, g2],
                        }
                    ],
                }
            ],
        }
    ]
    attach_region_summaries(
        hierarchy,
        [g1, g2],
        source_columns=source_columns,
        visible_columns=visible,
        digits=1,
    )
    assert hierarchy[0]["ues_list"][0]["res_list"][0]["summary"]["label"] == "РЭС-1, всего"
    assert hierarchy[0]["ues_list"][0]["summary"]["label"] == "ОЭС, всего"
    assert hierarchy[0]["summary"]["label"] == "ЕЭС, всего"


def test_compute_differences():
    columns = [
        {"key": "MAZUT_6tp"},
        {"key": "MAZUT_ek"},
        {"key": "d_mazut"},
        {"key": "GAZ_ESTEST"},
        {"key": "GAZ_ek"},
        {"key": "d_gaz"},
        {"key": "TvTopl_6tp"},
        {"key": "TvTopl_ek"},
        {"key": "d_tv"},
    ]
    # Все значения есть — разности вычисляются
    filled = [100.0, 30.0, 999, 50.0, 20.0, 999, 200.0, 50.0, 999]
    result = _compute_differences(filled, columns)
    assert result[2] == 70.0
    assert result[5] == 30.0
    assert result[8] == 150.0

    # Пустые значения — разность None
    filled2 = [100.0, None, 999, 50.0, 20.0, 999]
    columns2 = columns[:6]
    result2 = _compute_differences(filled2, columns2)
    assert result2[2] is None
    assert result2[5] == 30.0

    # Нет столбца d_mazut — не падает
    columns3 = [
        {"key": "MAZUT_6tp"},
        {"key": "MAZUT_ek"},
    ]
    result3 = _compute_differences([100.0, 30.0], columns3)
    assert result3 == [100.0, 30.0]


def test_display_is_nonzero():
    assert _display_is_nonzero("70") is True
    assert _display_is_nonzero("0") is False
    assert _display_is_nonzero("0.0") is False
    assert _display_is_nonzero("0,0") is False
    assert _display_is_nonzero("") is False
    assert _display_is_nonzero(None) is False
    assert _display_is_nonzero("-1.5") is True


def test_attach_d_mazut_formula_tooltip_on_caption_not_key_row():
    headers = [
        [
            {
                "text": "Разность по израсходованному мазуту, т у.т.",
                "column_key": "d_mazut",
                "is_key": False,
            },
            {"text": "ГАЗ", "column_key": "d_gaz", "is_key": False},
        ],
        [
            {"text": "d_mazut", "column_key": "d_mazut", "is_key": True},
            {"text": "d_gaz", "column_key": "d_gaz", "is_key": True},
        ],
    ]
    _attach_column_formula_tooltips(
        headers,
        {"d_mazut": "d_mazut = MAZUT_6tp − MAZUT_ek"},
    )
    assert headers[0][0]["formula_tooltip"] == "d_mazut = MAZUT_6tp − MAZUT_ek"
    assert "formula_tooltip" not in headers[1][0]
    assert "formula_tooltip" not in headers[0][1]


def test_pg_pages_recompute_difference_columns_and_attach_formulas(tmp_path, monkeypatch):
    path = tmp_path / "cmp.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ПГ_без кот"
    ws.append(
        [
            "Fed_okr",
            "nameoes",
            "oblast",
            "Stancia",
            "numb1120",
            "Year",
            "MAZUT_6tp",
            "MAZUT_ek",
            "d_mazut",
            "GAZ_ESTEST",
            "GAZ_ek",
            "d_gaz",
            "TvTopl_6tp",
            "TvTopl_ek",
            "d_tv",
        ]
    )
    ws.append(
        [
            "ФО",
            "ОЭС",
            "Субъект",
            "Станция",
            "Код",
            "Год",
            "Мазут",
            "Мазут",
            "Мазут",
            "Газ",
            "Газ",
            "Газ",
            "Твердое топливо",
            "Твердое топливо",
            "Твердое топливо",
        ]
    )
    ws.append(
        [
            "ФО",
            "ОЭС",
            "Субъект",
            "Станция",
            "Код",
            "Год",
            "Мазут 6ТП",
            "Мазут макет",
            "Разность по израсходованному мазуту, т у.т.",
            "Газ естественный",
            "Газ макет",
            "Разность по израсходованному газу, т у.т.",
            "Твердое 6ТП",
            "Твердое макет",
            "Разность по израсходованному твердому топливу, т у.т.",
        ]
    )
    ws.append(
        [
            "СЗ",
            "ОЭС СЗ",
            "АО",
            "ТЭЦ-1",
            1,
            2017,
            100.0,
            30.0,
            999.0,
            50.0,
            20.0,
            999.0,
            200.0,
            50.0,
            999.0,
        ]
    )
    ws.append(
        [
            "СЗ",
            "ОЭС СЗ",
            "АО",
            "ТЭЦ-1",
            1,
            2018,
            30.0,
            30.0,
            999.0,
            20.0,
            20.0,
            999.0,
            50.0,
            50.0,
            999.0,
        ]
    )
    for title in ("ЗВ_без кот", "ЗВ_всё", "ПГ_всё", "Лист1"):
        extra = wb.create_sheet(title)
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 1])
    wb.save(path)
    clear_comparison_cache()
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {},
    )
    _patch_pg_groups_from_fuel(
        monkeypatch,
        [(1, 2017, "ТЭЦ-1"), (1, 2018, "ТЭЦ-1")],
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {
            (1, 2017): {
                "energy_system_type_id": 1,
                "est_name": "ЕЭС",
                "ues_id": 2,
                "ues": "ОЭС",
                "res_id": 3,
                "res_name": "РЭС",
                "station_name": "ТЭЦ-1",
                "ved": 1,
                "gaz_estest_tce": 50.0,
                "mazut_tce": 100.0,
                "ugol_tce": 200.0,
            },
            (1, 2018): {
                "energy_system_type_id": 1,
                "est_name": "ЕЭС",
                "ues_id": 2,
                "ues": "ОЭС",
                "res_id": 3,
                "res_name": "РЭС",
                "station_name": "ТЭЦ-1",
                "ved": 1,
                "gaz_estest_tce": 20.0,
                "mazut_tce": 30.0,
                "ugol_tce": 50.0,
            },
        },
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_pg_ecology_overlay_by_numb_year",
        lambda sheet_slug, numbs, years: {
            (1, 2017): {
                "gas_consumption_tce": 20.0,
                "oil_fuel_consumption_tce": 30.0,
                "solid_fuel_consumption_tce": 50.0,
            },
            (1, 2018): {
                "gas_consumption_tce": 20.0,
                "oil_fuel_consumption_tce": 30.0,
                "solid_fuel_consumption_tce": 50.0,
            },
        },
    )
    ctx = build_comparison_table_context(
        "pg-bez-kot",
        start_year=2017,
        end_year=2018,
        rounding_digits=1,
        source_path=path,
    )
    assert ctx is not None
    group = ctx["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]
    expected = {
        "d_mazut": ("70", "MAZUT_6tp", "MAZUT_ek"),
        "d_gaz": ("30", "GAZ_ESTEST", "GAZ_ek"),
        "d_tv": ("150", "TvTopl_6tp", "TvTopl_ek"),
    }
    for key, (value, left, right) in expected.items():
        idx = next(i for i, col in enumerate(ctx["columns"]) if col["key"] == key)
        assert group["rows"][0]["cells"][idx] == value
        captions = [
            cell.get("formula_tooltip")
            for row in ctx["header_rows"]
            for cell in row
            if cell.get("column_key") == key and cell.get("formula_tooltip")
        ]
        assert captions
        assert left in captions[0]
        assert right in captions[0]
        assert ctx["columns"][idx]["diff_highlight"] is True
        header_flags = [
            cell.get("diff_highlight")
            for row in ctx["header_rows"]
            for cell in row
            if cell.get("column_key") == key
        ]
        assert header_flags
        assert all(header_flags)
        assert group["rows"][0]["cell_nonzero"][idx] is True
        assert group["rows"][1]["cells"][idx] == "0"
        assert group["rows"][1]["cell_nonzero"][idx] is False
    other = next(col for col in ctx["columns"] if col["key"] == "MAZUT_6tp")
    assert other["diff_highlight"] is False
    other_idx = next(i for i, col in enumerate(ctx["columns"]) if col["key"] == "MAZUT_6tp")
    assert group["rows"][0]["cell_nonzero"][other_idx] is False
    assert not any(row and row[0].get("is_key") for row in ctx["header_rows"])
    assert not any(
        str(cell.get("text") or "") in {"GAZ_ek", "d_gaz", "GAZ + ISK_GAZ"}
        and cell.get("is_key")
        for row in ctx["header_rows"]
        for cell in row
    )
    clear_comparison_cache()


def test_format_related_pg_cells_rounds_emission_coefficients_separately(monkeypatch):
    """«Округл (k эмис)» округляет только K_*, остальные ячейки — как «Округл»."""
    import app.ecology.services.gis_tes_fuel_comparison_services as svc

    monkeypatch.setattr(
        svc,
        "load_pg_ecology_overlay_by_numb_year",
        lambda sheet_slug, numbs, years: {
            (1, 2017): {
                "gas_consumption_tce": 802393.6,
                "co2_from_gas_t": 1275805.8,
            }
        },
    )
    monkeypatch.setattr(
        svc,
        "load_fuel_indicators_by_numb_year",
        lambda numbs, years: {(1, 2017): {}},
    )

    rounded = svc.format_related_pg_ecology_cells(
        "pg-bez-kot", 1, 2017, "GAZ_ek", rounding_digits=1, emission_rounding_digits=3
    )
    assert rounded["GAZ_ek"]["text"] == "802 393,6"
    assert rounded["K_gaz_ek"]["text"] == "1,59"

    whole = svc.format_related_pg_ecology_cells(
        "pg-bez-kot", 1, 2017, "GAZ_ek", rounding_digits=1, emission_rounding_digits=0
    )
    assert whole["K_gaz_ek"]["text"] == "1,58999997008949"


def test_emission_rounding_digits_apply_only_to_emission_columns(tmp_path, monkeypatch):
    path = tmp_path / "cmp-emission.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ПГ_без кот"
    ws.append(
        [
            "Fed_okr",
            "nameoes",
            "oblast",
            "Stancia",
            "numb1120",
            "Year",
            "GAZ_ek",
            "CO2_gaz",
            "K_gaz_ek",
            "W_en",
        ]
    )
    ws.append(
        [
            "ФО",
            "ОЭС",
            "Субъект",
            "Станция",
            "Код",
            "Год",
            "Газ макет",
            "Выбросы СО2 газа",
            "Коэффициент эмиссии (газ)",
            "Выработка",
        ]
    )
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ-1", 1, 2017, 802393.6, 1275805.8, 999.0, 500.0])
    for title in ("ЗВ_без кот", "ЗВ_всё", "ПГ_всё", "Лист1"):
        extra = wb.create_sheet(title)
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 1])
    wb.save(path)
    clear_comparison_cache()
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {},
    )
    _patch_pg_groups_from_fuel(monkeypatch, [(1, 2017, "ТЭЦ-1")])
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {(1, 2017): {"e": 500.0, "ved": 1}},
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_pg_ecology_overlay_by_numb_year",
        lambda sheet_slug, numbs, years: {
            (1, 2017): {
                "gas_consumption_tce": 802393.6,
                "co2_from_gas_t": 1275805.8,
            }
        },
    )

    def _cells(emission_rounding_digits):
        ctx = build_comparison_table_context(
            "pg-bez-kot",
            start_year=2017,
            end_year=2017,
            rounding_digits=1,
            emission_rounding_digits=emission_rounding_digits,
            source_path=path,
        )
        assert ctx is not None
        assert ctx["emission_rounding_digits"] == emission_rounding_digits
        keys = [column["key"] for column in ctx["columns"]]
        cells = ctx["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]["rows"][0][
            "cells"
        ]
        return {key: cells[index] for index, key in enumerate(keys)}

    three_digits = _cells(3)
    assert three_digits["K_gaz_ek"] == "1,59"
    assert three_digits["GAZ_ek"] == "802 393,6"

    not_rounded = _cells(0)
    assert not_rounded["K_gaz_ek"] == "1,58999997008949"
    assert not_rounded["GAZ_ek"] == "802 393,6"
    clear_comparison_cache()


def test_format_related_zv_cells_rounds_specific_columns_separately(monkeypatch):
    """«Округл (удельные)» округляет только удельные столбцы ЗВ."""
    import app.ecology.services.gis_tes_fuel_comparison_services as svc

    monkeypatch.setattr(
        svc,
        "load_zv_ecology_overlay_by_numb_year",
        lambda sheet_slug, numbs, years: {
            (1, 2017): {
                "solid_substances": 3.33333,
                "sulfur_dioxide": 10.123456,
                "nitrogen_oxides": 4.98765,
            }
        },
    )
    monkeypatch.setattr(
        svc,
        "load_fuel_indicators_by_numb_year",
        lambda numbs, years: {
            (1, 2017): {
                "e": 3.0,
                "q": 1.0,
                "gaz_all_tce": 100.0,
                "mazut_tce": 20.0,
                "solid_tce": 50.0,
            }
        },
    )

    rounded = svc.format_related_zv_ecology_cells(
        "zv-vse", 1, 2017, "SO2_gaz", rounding_digits=1, specific_rounding_digits=3
    )
    assert rounded["SO2_gaz"]["text"] == "10,1"
    assert rounded["udel_SO2"]["text"] == "144,621"
    assert rounded["normir_ZV_na_ee"]["text"] == "6,148"

    whole = svc.format_related_zv_ecology_cells(
        "zv-vse", 1, 2017, "SO2_gaz", rounding_digits=1, specific_rounding_digits=0
    )
    assert whole["udel_SO2"]["text"] == "144,6208"


def test_zv_specific_rounding_digits_apply_only_to_specific_columns(tmp_path, monkeypatch):
    path = tmp_path / "cmp-zv-specific.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ЗВ_без кот"
    keys = [
        "Fed_okr",
        "nameoes",
        "oblast",
        "Stancia",
        "numb1120",
        "Year",
        "SO2_gaz",
        "NOx_gaz",
        "TVERD_V",
        "GAZ_6tp",
        "MAZUT_6tp",
        "TVERD_TOPL_6tp",
        "W_en",
        "Teplo",
        "udel_na_energ",
        "normir_ZV_na_ee",
        "udel_tv",
        "udel_SO2",
        "udel_NOx",
        "normir_TVERD_na_energ",
        "normir_SO2_na_energ",
        "normir_NOx_na_energ",
    ]
    ws.append(keys)
    ws.append(keys)
    ws.append(
        [
            "СЗ",
            "ОЭС СЗ",
            "АО",
            "ТЭЦ-1",
            1,
            2017,
            10.123456,
            4.98765,
            3.33333,
            100.0,
            20.0,
            50.0,
            3.0,
            1.0,
            999.0,
            999.0,
            999.0,
            999.0,
            999.0,
            999.0,
            999.0,
            999.0,
        ]
    )
    for title in ("ПГ_без кот", "ЗВ_всё", "ПГ_всё", "Лист1"):
        extra = wb.create_sheet(title)
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 1, 2017, 1])
    wb.save(path)
    clear_comparison_cache()
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {},
    )
    _patch_pg_groups_from_fuel(monkeypatch, [(1, 2017, "ТЭЦ-1")])
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {
            (1, 2017): {
                "e": 3.0,
                "q": 1.0,
                "gaz_all_tce": 100.0,
                "mazut_tce": 20.0,
                "solid_tce": 50.0,
                "ved": 1,
            }
        },
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_zv_ecology_overlay_by_numb_year",
        lambda sheet_slug, numbs, years: {
            (1, 2017): {
                "solid_substances": 3.33333,
                "sulfur_dioxide": 10.123456,
                "nitrogen_oxides": 4.98765,
            }
        },
    )

    def _cells(specific_rounding_digits):
        ctx = build_comparison_table_context(
            "zv-bez-kot",
            start_year=2017,
            end_year=2017,
            rounding_digits=1,
            specific_rounding_digits=specific_rounding_digits,
            source_path=path,
        )
        assert ctx is not None
        assert ctx["specific_rounding_digits"] == specific_rounding_digits
        keys = [column["key"] for column in ctx["columns"]]
        cells = ctx["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]["rows"][0][
            "cells"
        ]
        return {key: cells[index] for index, key in enumerate(keys)}

    three_digits = _cells(3)
    assert three_digits["SO2_gaz"] == "10,1"
    assert three_digits["udel_SO2"] == "144,621"
    assert three_digits["normir_ZV_na_ee"] == "6,148"

    not_rounded = _cells(0)
    assert not_rounded["SO2_gaz"] == "10,1"
    assert not_rounded["udel_SO2"] == "144,6208"
    clear_comparison_cache()


def test_zv_ecology_edit_mapping_keys():
    assert ZV_ECOLOGY_OVERLAY_BY_KEY == {
        "TVERD_V": "solid_substances",
        "SO2_gaz": "sulfur_dioxide",
        "CO_gaz": "carbon_monoxide",
        "NOx_gaz": "nitrogen_oxides",
        "UGLEVOD": "hydrocarbons",
        "LOS": "voc",
        "PROCH_ekol": "other_gaseous_and_liquid",
    }
    assert ZV_ECOLOGY_EDITABLE_KEYS == frozenset(ZV_ECOLOGY_OVERLAY_BY_KEY)
    derived = (
        "VIBROS_ekol",
        "GAZ and ZHID",
        "energy",
        "ZV_normir",
        "udel_na_energ",
        "normir_ZV_na_ee",
        "normir_ZV_na_energ",
        "normir_TVERD_na_energ",
        "normir_SO2_na_energ",
        "normir_NOx_na_energ",
        "udel_tv",
        "udel_SO2",
        "udel_NOx",
    )
    assert ZV_ECOLOGY_EDIT_RELATED_KEYS["SO2_gaz"] == ("SO2_gaz",) + derived
    assert ZV_ECOLOGY_EDIT_RELATED_KEYS["TVERD_V"] == ("TVERD_V",) + derived
    assert set(ZV_ECOLOGY_EDIT_RELATED_KEYS) == set(ZV_ECOLOGY_OVERLAY_BY_KEY)


def test_zv_ecology_values_by_key_recomputes_totals_and_specific():
    by_key = _zv_ecology_values_by_key(
        {
            "solid_substances": 10,
            "sulfur_dioxide": 5,
            "carbon_monoxide": 4,
            "nitrogen_oxides": 3,
            "hydrocarbons": 2,
            "voc": 1,
            "other_gaseous_and_liquid": 0.5,
        },
        {
            "gaz_all_tce": 1000,
            "mazut_tce": 200,
            "solid_tce": 300,
            "e": 500,
            "q": 859.8,
        },
    )
    assert by_key["GAZ and ZHID"] == 15.5
    assert by_key["VIBROS_ekol"] == 25.5
    assert by_key["ZV_normir"] == 18
    assert by_key["energy"] == 1500
    assert abs(by_key["udel_na_energ"] - 25.5 / 1500) < 1e-9
    assert abs(by_key["normir_ZV_na_ee"] - 18 / 500) < 1e-9
    assert abs(by_key["normir_ZV_na_energ"] - 18 / 1500) < 1e-9
    assert abs(by_key["udel_tv"] - 10 * 1000 / 500) < 1e-9
    assert abs(by_key["udel_SO2"] - 5 * 1000 / 500) < 1e-9
    assert abs(by_key["udel_NOx"] - 3 * 1000 / 1500) < 1e-9


def test_save_analysis_zv_ecology_cell_saves_mapped_field(monkeypatch):
    calls = []

    def _upsert(sheet_slug, station_numb, year_number, values):
        calls.append((sheet_slug, station_numb, year_number, values))

    monkeypatch.setattr(
        "app.ecology.services.analysis_zv_param_services.upsert_analysis_zv_param_values",
        _upsert,
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_zv_ecology_overlay_by_numb_year",
        lambda sheet_slug, numbs, years: {
            (708, 2017): {
                "solid_substances": 10,
                "sulfur_dioxide": 15,
                "carbon_monoxide": 4,
                "nitrogen_oxides": 3,
                "hydrocarbons": 2,
                "voc": 1,
                "other_gaseous_and_liquid": 0.5,
            }
        },
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {},
    )

    values = save_analysis_zv_ecology_cell(
        "zv-vse", 708, 2017, "SO2_gaz", Decimal("15"), rounding_digits=1
    )

    assert calls == [("zv-vse", 708, 2017, {"sulfur_dioxide": Decimal("15")})]
    assert values["SO2_gaz"]["text"] == "15"
    assert values["GAZ and ZHID"]["text"] == "25,5"
    assert values["VIBROS_ekol"]["text"] == "35,5"
    assert values["SO2_gaz"]["nonzero"] is False


def test_related_zv_cells_return_full_value_as_title(monkeypatch):
    """В ответе на правку рядом с округлённым текстом — полное значение для подсказки."""
    import app.ecology.services.gis_tes_fuel_comparison_services as svc

    monkeypatch.setattr(
        svc,
        "load_zv_ecology_overlay_by_numb_year",
        lambda slug, numbs, years: {
            (708, 2017): {"sulfur_dioxide": Decimal("15.123456789")}
        },
    )
    monkeypatch.setattr(svc, "load_fuel_indicators_by_numb_year", lambda numbs, years: {})

    values = svc.format_related_zv_ecology_cells(
        "zv-vse", 708, 2017, "SO2_gaz", rounding_digits=1
    )

    assert values["SO2_gaz"]["text"] == "15,1"
    assert values["SO2_gaz"]["title"] == "15,123456789"


def test_save_analysis_zv_ecology_cell_rejects_unknown_column():
    try:
        save_analysis_zv_ecology_cell("zv-vse", 1, 2017, "W_en", 1)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "недоступен" in str(exc).lower()


def test_save_analysis_ecology_cell_dispatches_by_sheet(monkeypatch):
    import app.ecology.services.gis_tes_fuel_comparison_services as svc

    pg_calls = []
    zv_calls = []
    monkeypatch.setattr(
        svc, "save_analysis_pg_ecology_cell", lambda *args: pg_calls.append(args) or {}
    )
    monkeypatch.setattr(
        svc, "save_analysis_zv_ecology_cell", lambda *args: zv_calls.append(args) or {}
    )

    svc.save_analysis_ecology_cell("pg-vse", 1, 2017, "GAZ_ek", 1, 2, 3)
    svc.save_analysis_ecology_cell("zv-bez-kot", 2, 2017, "SO2_gaz", 2, 2, 2, 3)

    assert pg_calls == [("pg-vse", 1, 2017, "GAZ_ek", 1, 2, 3)]
    assert zv_calls == [("zv-bez-kot", 2, 2017, "SO2_gaz", 2, 2, 3)]
    try:
        svc.save_analysis_ecology_cell("zv-other", 1, 2017, "SO2_gaz", 1)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "недоступен" in str(exc).lower()


def test_zv_analysis_columns_are_editable_and_use_params_overlay(tmp_path, monkeypatch):
    path = tmp_path / "zv-edit.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ЗВ_всё"
    ws.append(
        [
            "Fed_okr",
            "nameoes",
            "oblname",
            "NAME",
            "numb1120",
            "Year",
            "VIBROS_ekol",
            "TVERD_V",
            "GAZ and ZHID",
            "SO2_gaz",
            "CO_gaz",
            "NOx_gaz",
            "UGLEVOD",
            "LOS",
            "PROCH_ekol",
        ]
    )
    ws.append(
        [
            "ФО",
            "ОЭС",
            "Субъект",
            "Станция",
            "Код",
            "Год",
            "Всего ЗВ",
            "Твердые",
            "всего",
            "диоксид серы",
            "оксид углерода",
            "оксиды азота",
            "углеводороды",
            "ЛОС",
            "прочие",
        ]
    )
    ws.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 708, 2017, 999, 1, 999, 1, 1, 1, 1, 1, 1])
    for title in ("ПГ_без кот", "ЗВ_без кот", "ПГ_всё", "Лист1"):
        extra = wb.create_sheet(title)
        extra.append(["Fed_okr", "nameoes", "oblname", "NAME", "numb1120", "Year", "W_en"])
        extra.append(["СЗ", "ОЭС СЗ", "АО", "ТЭЦ", 708, 2017, 1])
    wb.save(path)
    clear_comparison_cache()
    _patch_pg_groups_from_fuel(monkeypatch, [(708, 2017, "ТЭЦ-1")])
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.get_year_feature_dict",
        lambda: {2017: "отч"},
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_fuel_indicators_by_numb_year",
        lambda numbs, years: {
            (708, 2017): {
                "gaz_all_tce": 1000,
                "mazut_tce": 200,
                "solid_tce": 300,
                "e": 500,
                "q": 0,
                "ved": 1,
            }
        },
    )
    monkeypatch.setattr(
        "app.ecology.services.gis_tes_fuel_comparison_services.load_zv_ecology_overlay_by_numb_year",
        lambda sheet_slug, numbs, years: {
            (708, 2017): {
                "solid_substances": 10,
                "sulfur_dioxide": 15,
                "carbon_monoxide": 4,
                "nitrogen_oxides": 3,
                "hydrocarbons": 2,
                "voc": 1,
                "other_gaseous_and_liquid": 0.5,
            }
        },
    )
    context = build_comparison_table_context(
        "zv-vse",
        start_year=2017,
        end_year=2017,
        rounding_digits=1,
        source_path=path,
    )
    assert context is not None
    group = context["hierarchy"][0]["ues_list"][0]["res_list"][0]["groups"][0]
    by_key = {
        col["key"]: group["rows"][0]["cells"][i]
        for i, col in enumerate(context["columns"])
    }
    assert by_key["TVERD_V"] == "10"
    assert by_key["SO2_gaz"] == "15"
    assert by_key["GAZ and ZHID"] == "25,5"
    assert by_key["VIBROS_ekol"] == "35,5"
    editable = {col["key"] for col in context["columns"] if col.get("editable")}
    assert editable == set(ZV_ECOLOGY_OVERLAY_BY_KEY)
    clear_comparison_cache()
