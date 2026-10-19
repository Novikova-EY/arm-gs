# -*- coding: utf-8 -*-
"""Юнит-тесты выгрузки отчетных данных ГИС ТЭК в Excel."""

from __future__ import annotations

from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import app.ecology.services.gis_tek_export_services as export_services
from app.ecology.services.gis_tek_constants import REPORT_PG, REPORT_ZV, ZV_NORM_ATTR

IDENTITY_COLUMNS = [
    ("station_numb", "№_станции", True),
    ("gis_tek_station_name", "Электростанция (ГИС ТЭК)", False),
]
# В выгрузке перед станционными столбцами добавляются «ОЭС» и «РЭС».
EXPORT_IDENTITY_COUNT = len(IDENTITY_COLUMNS) + 2


def _context(value_columns, groups, metrics, norm_selected=()):
    """Контекст страницы в объеме, который читает выгрузка Excel."""
    return {
        "identity_columns": list(IDENTITY_COLUMNS),
        "value_columns": list(value_columns),
        "value_column_groups": list(groups),
        "value_column_banner": "Выбросы загрязняющих веществ",
        "norm_zv_selected": list(norm_selected),
        "show_empty_rows": False,
        "hierarchy": [
            {
                "summary": None,
                "res_list": [
                    {
                        "rows": [
                            {
                                "station_numb": 1,
                                "gis_tek_station_name": "Станция",
                                "metrics": dict(metrics),
                                "is_empty": False,
                            }
                        ],
                        "summary": None,
                    }
                ],
            }
        ],
    }


def _export(monkeypatch, kind, context, args=None):
    monkeypatch.setattr(
        export_services,
        "build_gis_tek_page_context",
        lambda *_args, **_kwargs: context,
    )
    workbook = export_services.build_gis_tek_excel_stream(kind, args or {})
    return load_workbook(workbook).active


def test_zv_export_headers_norm_zv_with_selected_parts(monkeypatch):
    value_columns = [("total", "всего"), (ZV_NORM_ATTR, "Нормируемые ЗВ")]
    groups = [("всего", [("total", "")]), ("Нормируемые ЗВ", [(ZV_NORM_ATTR, "")])]
    metrics = {"total": Decimal("10.5"), ZV_NORM_ATTR: Decimal("3.5")}
    context = _context(
        value_columns,
        groups,
        metrics,
        norm_selected=["solid_substances", "carbon_monoxide"],
    )

    ws = _export(monkeypatch, REPORT_ZV, context)

    # Шапка: 3-я строка — показатели, 4-я — группы, 5-я — подписи столбцов.
    norm_col = EXPORT_IDENTITY_COUNT + len(value_columns)
    caption = ws.cell(4, norm_col).value
    assert "Нормируемые ЗВ" in caption
    assert "Твердые вещества" in caption
    assert "оксид углерода [CO]" in caption
    assert "диоксид серы" not in caption
    # Значения строки остались в своих столбцах, сумма выбрана в последний.
    assert ws.cell(6, norm_col - 1).value == Decimal("10.5")
    assert ws.cell(6, norm_col).value == Decimal("3.5")
    # Подпись слагаемых шире обычных столбцов.
    assert ws.column_dimensions[get_column_letter(norm_col)].width > ws.column_dimensions["A"].width


def test_zv_export_marks_cleared_norm_zv_selection(monkeypatch):
    value_columns = [("total", "всего"), (ZV_NORM_ATTR, "Нормируемые ЗВ")]
    groups = [("всего", [("total", "")]), ("Нормируемые ЗВ", [(ZV_NORM_ATTR, "")])]
    metrics = {"total": Decimal("10.5"), ZV_NORM_ATTR: None}
    context = _context(value_columns, groups, metrics, norm_selected=[])

    ws = _export(monkeypatch, REPORT_ZV, context, args={"norm_zv": ""})

    assert ws.cell(4, EXPORT_IDENTITY_COUNT + 2).value == "Нормируемые ЗВ (слагаемые не выбраны)"
    assert ws.cell(6, EXPORT_IDENTITY_COUNT + 2).value is None


def test_pg_export_has_no_norm_zv_caption(monkeypatch):
    value_columns = [("gas_consumption_tce", "Расход газа, т у.т.")]
    groups = [("Газ", [("gas_consumption_tce", "")])]
    context = _context(value_columns, groups, {"gas_consumption_tce": Decimal("1.5")})

    ws = _export(monkeypatch, REPORT_PG, context)

    assert ws.cell(4, EXPORT_IDENTITY_COUNT + 1).value == "Газ"
    assert ws.cell(6, EXPORT_IDENTITY_COUNT + 1).value == Decimal("1.5")
    assert ws.row_dimensions[4].height is None
    values = " ".join(str(cell.value) for row in ws.iter_rows() for cell in row)
    assert "Нормируемые ЗВ" not in values
