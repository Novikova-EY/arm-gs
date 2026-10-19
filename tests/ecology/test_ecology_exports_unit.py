# -*- coding: utf-8 -*-
"""Выгрузка экологии в Excel: к каждой станции добавлены ОЭС и РЭС."""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import load_workbook

import app.ecology.services.gis_tek_export_services as gis_tek_export
from app.ecology.services.export_gis_tes_fuel_comparison_services import (
    export_comparison_table_to_excel,
)
from app.ecology.services.gis_tek_constants import REPORT_PG


def _gis_tek_context() -> dict:
    return {
        "identity_columns": [
            ("station_numb", "Код группы оборудования (numb)", True),
            ("gis_tek_station_name", "Электростанция (ГИС ТЭК)", False),
        ],
        "value_columns": [("total", "Всего")],
        "value_column_groups": [],
        "value_column_banner": "",
        "show_empty_rows": False,
        "norm_zv_selected": [],
        "hierarchy": [
            {
                "ues_name": "ОЭС Северо-Запада",
                "summary": None,
                "res_list": [
                    {
                        "res_name": "ЭС Архангельской области",
                        "summary": None,
                        "rows": [
                            {
                                "station_numb": 1,
                                "gis_tek_station_name": "Архангельская ТЭЦ",
                                "ues_name": "ОЭС Северо-Запада",
                                "res_name": "ЭС Архангельской области",
                                "metrics": {"total": Decimal("12.5")},
                                "is_empty": False,
                            }
                        ],
                    }
                ],
            }
        ],
    }


def _analysis_context() -> dict:
    columns = [
        {"key": "NAME", "numeric": False, "source_index": 0},
        {"key": "numb1120", "numeric": True, "source_index": 1},
        {"key": "Year", "numeric": True, "source_index": 2},
        {"key": "nust", "numeric": True, "source_index": 3},
    ]
    header_rows = [
        [
            {"column_key": "NAME", "text": "Электростанция"},
            {"column_key": "numb1120", "text": "Код группы оборудования (numb1120)"},
            {"column_key": "Year", "text": "Год"},
            {"column_key": "nust", "text": "Установленная мощность, МВт"},
        ]
    ]
    return {
        "sheet": {"nav_label": "С котельными", "page_title": "Парниковые газы"},
        "page_title": "Парниковые газы (с котельными)",
        "columns": columns,
        "header_rows": header_rows,
        "show_empty_rows": False,
        "hierarchy": [
            {
                "est_name": "ЕЭС России",
                "summary": None,
                "ues_list": [
                    {
                        "ues_name": "ОЭС Северо-Запада",
                        "summary": None,
                        "res_list": [
                            {
                                "res_name": "ЭС Архангельской области",
                                "summary": None,
                                "groups": [
                                    {
                                        "rows": [
                                            {
                                                "year": 2023,
                                                "is_group_first": True,
                                                "is_empty": False,
                                                "cells": ["Архангельская ТЭЦ", "1", "2023", "450"],
                                            },
                                            {
                                                "year": 2024,
                                                "is_group_first": False,
                                                "is_empty": False,
                                                "cells": ["", "", "2024", "460"],
                                            },
                                            {
                                                "year": 2025,
                                                "is_group_first": False,
                                                "is_empty": True,
                                                "cells": ["", "", "2025", ""],
                                            },
                                        ],
                                        "filled_by_year": {
                                            2023: ["Архангельская ТЭЦ", 1, 2023, Decimal("450")],
                                            2024: ["Архангельская ТЭЦ", 1, 2024, Decimal("460")],
                                            2025: ["Архангельская ТЭЦ", 1, 2025, None],
                                        },
                                    }
                                ],
                            }
                        ],
                    }
                ],
            }
        ],
    }


def _sheet(stream) -> "object":
    workbook = load_workbook(io.BytesIO(stream.getvalue()))
    return workbook.active


def test_gis_tek_export_adds_ues_and_res_columns(monkeypatch):
    monkeypatch.setattr(
        gis_tek_export,
        "build_gis_tek_page_context",
        lambda kind, args, user=None: _gis_tek_context(),
    )
    ws = _sheet(gis_tek_export.build_gis_tek_excel_stream(REPORT_PG, {}))
    headers = [ws.cell(3, col).value for col in range(1, 6)]
    assert headers == [
        "Код группы оборудования (numb)",
        "ОЭС",
        "РЭС",
        "Электростанция (ГИС ТЭК)",
        "Всего",
    ]
    assert [ws.cell(4, col).value for col in range(1, 6)] == [
        1,
        "ОЭС Северо-Запада",
        "ЭС Архангельской области",
        "Архангельская ТЭЦ",
        12.5,
    ]


def test_analysis_export_has_hidden_key_row():
    ws = _sheet(export_comparison_table_to_excel(_analysis_context()))
    # Служебная строка с кодами столбцов — по ней импорт находит столбцы.
    assert ws.cell(2, 1).value == "#keys"
    assert [ws.cell(2, col).value for col in range(3, 7)] == [
        "NAME",
        "numb1120",
        "Year",
        "nust",
    ]
    assert ws.row_dimensions[2].hidden is True
    # Фильтр Excel начинается со строки заголовков, а не со служебной строки.
    assert ws.auto_filter.ref == f"A3:F{ws.max_row}"


def test_analysis_export_adds_ues_and_res_columns():
    ws = _sheet(export_comparison_table_to_excel(_analysis_context()))
    headers = [ws.cell(3, col).value for col in range(1, 7)]
    assert headers == [
        "ОЭС",
        "РЭС",
        "Электростанция",
        "Код группы оборудования (numb1120)",
        "Год",
        "Установленная мощность, МВт",
    ]
    # Пустая строка 2025 скрыта, как на странице (show_empty_rows=0).
    assert [ws.cell(row, col).value for row in (4, 5) for col in (1, 2)] == [
        "ОЭС Северо-Запада",
        "ЭС Архангельской области",
    ] * 2
    assert [ws.cell(4, col).value for col in range(1, 7)] == [
        "ОЭС Северо-Запада",
        "ЭС Архангельской области",
        "Архангельская ТЭЦ",
        1,
        2023,
        450,
    ]
    # У повторной строки года столбцы станции заполняются из первой строки группы.
    assert [ws.cell(5, col).value for col in range(1, 7)] == [
        "ОЭС Северо-Запада",
        "ЭС Архангельской области",
        "Архангельская ТЭЦ",
        1,
        2024,
        460,
    ]
    assert ws.max_row == 5
