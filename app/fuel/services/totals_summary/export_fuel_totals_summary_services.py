# -*- coding: utf-8 -*-
"""Экспорт итоговых сумм модуля «Топливо» в Excel (как на экране)."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
_ROW_FILLS = {
    "table-primary": PatternFill(start_color="CFE2FF", end_color="CFE2FF", fill_type="solid"),
    "table-success": PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid"),
    "table-secondary": PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid"),
    "table-info": PatternFill(start_color="CFF4FC", end_color="CFF4FC", fill_type="solid"),
    "table-warning": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
    "table-danger": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
}


def export_fuel_totals_summary_to_excel(context: dict[str, Any]) -> BytesIO:
    """Выгрузка таблицы итоговых сумм топлива: блоки/годы/показатели как на экране."""
    columns = list(context.get("columns") or [])
    blocks = list(context.get("blocks") or [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Итоговые суммы"

    headers = ["", "Год"] + [label for _attr, label, _num in columns]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

    for block in blocks:
        fill = _ROW_FILLS.get(block.get("row_class") or "")
        label = block.get("label") or ""
        for row in block.get("rows") or []:
            values = [label, row.get("year")] + list(row.get("cells") or [])
            ws.append(values)
            excel_row = ws.max_row
            for col_idx in range(1, len(values) + 1):
                cell = ws.cell(row=excel_row, column=col_idx)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if fill:
                    cell.fill = fill
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    cell.font = Font(bold=True)

    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        width = 18 if col_idx == 1 else (10 if col_idx == 2 else 14)
        ws.column_dimensions[letter].width = width

    ws.freeze_panes = "C2"
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
