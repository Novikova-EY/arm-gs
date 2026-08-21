# -*- coding: utf-8 -*-
"""Импорт «Ограничения» из Excel во все версии БД (шапка как в Access)."""

from __future__ import annotations

import io
from decimal import Decimal
from typing import BinaryIO

from openpyxl import load_workbook

from app.extensions import db
from app.fuel.services.restrictions.fuel_restrictions_all_versions_services import (
    upsert_fuel_restriction_in_all_versions,
)
from app.fuel.services.distribution_parameters.import_distribution_parameters_services import _parse_decimal

# Канонические ключи (Access / gs_fue_restrictions).
_EXCEL_CANONICAL_HEADERS = frozenset(
    {
        "year",
        "oes",
        "name",
        "obl",
        "emin",
        "emax",
        "ecur",
        "h",
        "ecurdis",
        "hdis",
        "etp",
        "kobl",
        "kcur",
    }
)

# Обязательные для строки импорта.
_REQUIRED_HEADERS = frozenset({"year", "oes", "obl"})

# Синонимы шапки (Access, подписи АРМ, латиница).
_HEADER_ALIASES: dict[str, str] = {
    "year": "year",
    "год": "year",
    "oes": "oes",
    "оэс": "oes",
    "name": "name",
    "restriction_name": "name",
    "наименование": "name",
    "obl": "obl",
    "субъект (obl)": "obl",
    "субъект": "obl",
    "emin": "emin",
    "emin (ввод)": "emin",
    "emax": "emax",
    "emax (ввод)": "emax",
    "ecur": "ecur",
    "h": "h",
    "ecurdis": "ecurdis",
    "hdis": "hdis",
    "etp": "etp",
    "kobl": "kobl",
    "kcur": "kcur",
}


def _normalize_header_cell(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    key = _HEADER_ALIASES.get(s.lower())
    if key in _EXCEL_CANONICAL_HEADERS:
        return key
    return None


def _cell_str(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        t = value.strip()
        return t or None
    return str(value).strip() or None


def _row_key(year: Decimal, oes: Decimal, obl: Decimal) -> tuple[Decimal, Decimal, Decimal]:
    return (year, oes, obl)


def import_fuel_restrictions_from_excel(
    file_obj: BinaryIO,
    *,
    database_version_id: int | None = None,
) -> tuple[int, int, list[str]]:
    """
    Читает первый лист Excel. Первая строка — шапка (порядок столбцов любой).

    Сопоставление существующих строк: (year, oes, obl) **во всех версиях БД**.
    При совпадении — обновление; иначе — вставка в каждую версию.

    ``database_version_id`` сохранён для совместимости вызовов и не ограничивает
    синхронизацию (данные пишутся во все версии).

    Returns:
        (добавлено_копий, обновлено_копий, список_ошибок_или_предупреждений)
    """
    _ = database_version_id

    wb = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValueError("Файл пустой или нет строки заголовков.")

        col_index: dict[str, int] = {}
        for i, cell in enumerate(header_row):
            key = _normalize_header_cell(cell)
            if key and key not in col_index:
                col_index[key] = i

        missing = sorted(_REQUIRED_HEADERS - set(col_index.keys()))
        if missing:
            raise ValueError(
                "В первой строке не хватает столбцов: "
                + ", ".join(missing)
                + ". Ожидаются как минимум year, oes, obl "
                "(как в Access-таблице «Ограничения»)."
            )

        errors: list[str] = []
        prepared: list[dict] = []
        seen_in_file: set[tuple[Decimal, Decimal, Decimal]] = set()
        row_no = 1

        for data_row in rows_iter:
            row_no += 1
            if not data_row or all(v is None or str(v).strip() == "" for v in data_row):
                continue

            def col(key: str):
                j = col_index.get(key)
                if j is None:
                    return None
                return data_row[j] if j < len(data_row) else None

            year_v = _parse_decimal(col("year"))
            oes_v = _parse_decimal(col("oes"))
            obl_v = _parse_decimal(col("obl"))
            if year_v is None or oes_v is None or obl_v is None:
                errors.append(
                    f"Строка {row_no}: нужны year, oes и obl "
                    f"(year={col('year')!r}, oes={col('oes')!r}, obl={col('obl')!r})."
                )
                continue

            key = _row_key(year_v, oes_v, obl_v)
            if key in seen_in_file:
                errors.append(
                    f"Строка {row_no}: дубликат ключа year={year_v}, oes={oes_v}, obl={obl_v} в файле."
                )
                continue
            seen_in_file.add(key)

            name = _cell_str(col("name"))
            if name is not None and len(name) > 50:
                errors.append(f"Строка {row_no}: наименование длиннее 50 символов.")
                continue
            kcur = _cell_str(col("kcur"))
            if kcur is not None and len(kcur) > 50:
                errors.append(f"Строка {row_no}: kcur длиннее 50 символов.")
                continue

            optional_numeric = ("emin", "emax", "ecur", "h", "ecurdis", "hdis", "etp", "kobl")
            # Только колонки из файла — отсутствующие не затирают значения в БД.
            data: dict = {}
            if "name" in col_index:
                data["restriction_name"] = name
            if "kcur" in col_index:
                data["kcur"] = kcur
            for attr in optional_numeric:
                if attr in col_index:
                    data[attr] = _parse_decimal(col(attr))

            prepared.append(
                {
                    "year": year_v,
                    "oes": oes_v,
                    "obl": obl_v,
                    "data": data,
                }
            )

        if errors:
            return 0, 0, errors

        if not prepared:
            return 0, 0, []

        n_added = 0
        n_updated = 0
        warnings: list[str] = []
        try:
            for item in prepared:
                added, updated, warns = upsert_fuel_restriction_in_all_versions(
                    year=item["year"],
                    oes=item["oes"],
                    obl=item["obl"],
                    data=item["data"],
                )
                n_added += added
                n_updated += updated
                warnings.extend(warns)
            db.session.commit()
            return n_added, n_updated, warnings
        except Exception:
            db.session.rollback()
            raise
    finally:
        wb.close()


def import_fuel_restrictions_from_upload(file_storage) -> tuple[int, int, list[str]]:
    raw = file_storage.read()
    if not raw:
        raise ValueError("Пустой файл.")
    return import_fuel_restrictions_from_excel(io.BytesIO(raw))
