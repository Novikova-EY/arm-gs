# -*- coding: utf-8 -*-
"""Сервис экспорта топливных параметров групп оборудования в Excel."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.common.services.help_services import apply_nbsp_to_row, format_decimal_for_display
from app.fuel.models.fue_equipment_group_fuel_param_model import EquipmentGroupFuelParam
from app.fuel.models.fue_equipment_group_specific_fuel_consumption_model import (
    EquipmentGroupSpecificFuelConsumption,
)
from app.fuel.services.equipment_groups.equipment_group_fuel_params_services import (
    get_equipment_groups_with_fuel_params_data,
    build_name_maps_from_rows,
    build_equipment_group_fuel_params_hierarchy,
)
from app.fuel.services.fuel_exports.hierarchy_excel_layout import (
    EQUIPMENT_GROUP_ID_HEADER,
    FILL_EST,
    FILL_RES_HEADER,
    FILL_RES_SUMMARY,
    FILL_STATION_SUMMARY,
    FILL_UES,
    equipment_group_id_cell,
    next_data_row_index,
    style_data_row,
    write_merged_section_row,
)

FUEL_PARAM_COLUMNS = [
    ("numb1120", "Код группы оборудования", False),
    ("nust", EquipmentGroupFuelParam.NUST_COLUMN_LABEL, True),
    ("nr", EquipmentGroupFuelParam.NR_COLUMN_LABEL, True),
    ("e", "Выработка ЭЭ, тыс.кВтч", True),
    ("ewtp", "Теплофикационная выработка ЭЭ, тыс.кВтч", True),
    ("eotp", EquipmentGroupFuelParam.EOTP_COLUMN_LABEL, True),
    ("eust", "Расх топ ээ", True),
    ("eurt", EquipmentGroupFuelParam.EURT_COLUMN_LABEL, True),
    ("snk", EquipmentGroupSpecificFuelConsumption.SNK_COLUMN_LABEL, True),
    ("q", EquipmentGroupFuelParam.Q_COLUMN_LABEL, True),
    ("qotr", "Тепловое потребление (отборов турбин), тыс.Гкал", True),
    ("turt", EquipmentGroupFuelParam.TURT_COLUMN_LABEL, True),
    ("tust", EquipmentGroupFuelParam.TUST_COLUMN_LABEL, True),
    ("sn_t", "СН, кВтч/⁠Гкал", True),
    ("b", "Расход топлива, всего", True),
    ("gaz", "Газ", True),
    ("isk_gaz", "Иск. газ", True),
    ("mazut", "Мазут", True),
    ("torf", "Торф", True),
    ("slan", "Сланцы", True),
    ("proch", "Прочее", True),
    ("ugol", "Уголь", True),
    ("don", "Дон", True),
    ("podm", "Подм", True),
    ("pech", "Печ", True),
    ("arkt", "Арктикуголь", True),
    ("kuzn", "Кузбасс", True),
    ("ural", "Урал", True),
    ("bashk", "Башкортостан", True),
    ("kazah", "Казахстан", True),
    ("kan", "Канск", True),
    ("tung", "Тунгусск", True),
    ("irkut", "Иркутск", True),
    ("hak", "Хакасия", True),
    ("tuv", "Тува", True),
    ("bur", "Бурятия", True),
    ("chit", "Чита", True),
    ("yakut", "Якутия", True),
    ("amur", "Амур", True),
    ("urg", "Юрга", True),
    ("ushum", "Ушумун", True),
    ("prim", "Приморье", True),
    ("mag", "Магадан", True),
    ("chukot", "Чукотка", True),
    ("kamch", "Камчатка", True),
    ("sah", "Сахалин", True),
    ("nt", EquipmentGroupFuelParam.NT_COLUMN_LABEL, True),
    ("nt_sum", EquipmentGroupFuelParam.NT_SUM_COLUMN_LABEL, True),
]


def _format_cell_value(param, attr, is_numeric, name_maps, rounding_digits):
    """Форматирует значение ячейки по аналогии с шаблоном."""
    if param is None:
        return "—"
    val = getattr(param, attr, None)
    if val is None:
        return "—"
    if attr == "obor" and name_maps.get("obor_name_map"):
        return name_maps["obor_name_map"].get(val, val) or "—"
    if attr == "obl" and name_maps.get("obl_name_map"):
        return name_maps["obl_name_map"].get(val, val) or "—"
    if attr == "dep" and name_maps.get("dep_name_map"):
        return name_maps["dep_name_map"].get(val, val) or "—"
    if attr == "oes" and name_maps.get("oes_name_map"):
        return name_maps["oes_name_map"].get(val, val) or "—"
    if attr == "er" and name_maps.get("er_name_map"):
        return name_maps["er_name_map"].get(val, val) or "—"
    if attr == "gk" and name_maps.get("gk_name_map"):
        return name_maps["gk_name_map"].get(val, val) or "—"
    if attr == "be" and name_maps.get("be_name_map"):
        return name_maps["be_name_map"].get(val, val) or "—"
    if attr == "ved":
        s = str(val)
        if s == "1":
            return "КЭС"
        if s == "2":
            return "ТЭЦ"
        return val
    if attr == "ved_cyrillic":
        s = str(val)
        if s == "1":
            return "станция отрасли"
        if s == "2":
            return "пром.предприятие"
        return val
    if is_numeric:
        return format_decimal_for_display(val, digits=rounding_digits)
    return str(val) if val is not None else "—"


def _format_summary_value(summary, attr, is_numeric, rounding_digits):
    if summary is None or attr not in summary:
        return "—"
    val = summary.get(attr)
    if val is None:
        return "—"
    if is_numeric:
        return format_decimal_for_display(val, digits=rounding_digits)
    return str(val)


def export_stations_equipment_group_fuel_params_to_excel(
    filters,
    start_year,
    end_year,
    *,
    per_page="all",
    page=1,
    show_all=True,
    rounding_digits=1,
):
    """
    Экспортирует топливные параметры групп оборудования в Excel.
    Учитывает фильтры: субъекты, ОЭС, региональные энергосистемы и т.п.
    Всегда полный набор строк (per_page=all); иерархия и агрегации по полным данным.
    """
    fuel_data = get_equipment_groups_with_fuel_params_data(
        filters=filters,
        per_page=per_page,
        page=page,
        start_year=start_year,
        end_year=end_year,
        show_all=show_all,
    )
    rows = fuel_data.get("rows") or []
    if not rows:
        return None

    name_maps = build_name_maps_from_rows(rows)
    hierarchy = build_equipment_group_fuel_params_hierarchy(
        rows, use_equipment_group_hierarchy_only=False
    )
    if not hierarchy:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "Топливные параметры"

    columns = [EQUIPMENT_GROUP_ID_HEADER, "Группа оборудования"] + [
        label for _, label, _ in FUEL_PARAM_COLUMNS
    ]
    num_cols = len(columns)

    header_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num, column_title in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = apply_nbsp_to_row([column_title])[0]
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    max_lengths = [len(str(c)) for c in columns]

    def _bump_lengths(row_vals):
        for i, v in enumerate(row_vals):
            if v is not None and len(str(v)) > max_lengths[i]:
                max_lengths[i] = len(str(v))

    for est_block in hierarchy:
        title = est_block.get("est_name") or "—"
        _bump_lengths([title] + [""] * (num_cols - 1))
        write_merged_section_row(
            ws, next_data_row_index(ws), num_cols, title, FILL_EST
        )

        for ues_block in est_block.get("ues_list") or []:
            ues_title = ues_block.get("ues_name") or "—"
            _bump_lengths([ues_title] + [""] * (num_cols - 1))
            write_merged_section_row(
                ws, next_data_row_index(ws), num_cols, ues_title, FILL_UES
            )

            for res_block in ues_block.get("res_list") or []:
                res_title = res_block.get("res_name") or "—"
                _bump_lengths([res_title] + [""] * (num_cols - 1))
                write_merged_section_row(
                    ws,
                    next_data_row_index(ws),
                    num_cols,
                    res_title,
                    FILL_RES_HEADER,
                )

                for station_block in res_block.get("station_blocks") or []:
                    for group_block in station_block.get("group_blocks") or []:
                        group_entity = group_block.get("equipment_group")
                        group_name = (
                            (group_entity.name if group_entity and group_entity.name else None)
                            or (
                                group_entity.name_ext
                                if group_entity and group_entity.name_ext
                                else None
                            )
                            or "—"
                        )
                        for _eg, param in group_block.get("rows") or []:
                            row_data = [
                                equipment_group_id_cell(group_entity),
                                group_name,
                            ]
                            for attr, _label, is_numeric in FUEL_PARAM_COLUMNS:
                                row_data.append(
                                    _format_cell_value(
                                        param, attr, is_numeric, name_maps, rounding_digits
                                    )
                                )
                            _bump_lengths(row_data)
                            ws.append(apply_nbsp_to_row(row_data))

                    gb_count = len(station_block.get("group_blocks") or [])
                    if (
                        gb_count > 1
                        and not station_block.get("is_virtual")
                        and not station_block.get("suppress_station_summary")
                    ):
                        st_name = station_block.get("station_name") or "—"
                        label = f"{st_name}, всего"
                        summary = station_block.get("station_summary") or {}
                        row_data = ["—", label]
                        for attr, _label, is_numeric in FUEL_PARAM_COLUMNS:
                            row_data.append(
                                _format_summary_value(
                                    summary, attr, is_numeric, rounding_digits
                                )
                            )
                        _bump_lengths(row_data)
                        ws.append(apply_nbsp_to_row(row_data))
                        style_data_row(
                            ws, ws.max_row, bold=True, fill=FILL_STATION_SUMMARY
                        )

                if res_block.get("group_blocks"):
                    res_nm = res_block.get("res_name") or "—"
                    label = f"{res_nm}, всего"
                    summary = res_block.get("res_summary") or {}
                    row_data = ["—", label]
                    for attr, _label, is_numeric in FUEL_PARAM_COLUMNS:
                        row_data.append(
                            _format_summary_value(summary, attr, is_numeric, rounding_digits)
                        )
                    _bump_lengths(row_data)
                    ws.append(apply_nbsp_to_row(row_data))
                    style_data_row(ws, ws.max_row, bold=True, fill=FILL_RES_SUMMARY)

    for col_idx, max_length in enumerate(max_lengths, start=1):
        adjusted_width = min(max_length + 2, 60)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
