# -*- coding: utf-8 -*-
"""Общее оформление строк иерархии (ЕЭС / ОЭС / РЭС / итоги) для выгрузки в Excel."""

from openpyxl.styles import Alignment, Font, PatternFill


# Первый столбец данных в выгрузках с иерархией (перед «Группа оборудования»)
EQUIPMENT_GROUP_ID_HEADER = "EquipmentGroup.id"


def equipment_group_id_cell(equipment_group) -> str:
    """Значение ячейки идентификатора группы оборудования; для строк без группы — «—»."""
    if equipment_group is None:
        return "—"
    eid = getattr(equipment_group, "id", None)
    return str(eid) if eid is not None else "—"


def next_data_row_index(ws) -> int:
    """
    Следующая свободная строка после уже записанных (в т.ч. после ws.append).
    Нужна, чтобы строки merge-заголовков не пересекались с данными: нельзя вести
    отдельный счётчик current_row только по merge, игнорируя append.
    """
    return ws.max_row + 1

# Подобрано по смыслу Bootstrap table-light / primary / secondary / warning / info
FILL_EST = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
FILL_UES = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
FILL_RES_HEADER = PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")
FILL_STATION_SUMMARY = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FILL_RES_SUMMARY = PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid")


def write_merged_section_row(ws, row: int, num_cols: int, title: str, fill: PatternFill) -> int:
    """Одна строка на всю ширину таблицы (как colspan в HTML). Возвращает номер следующей строки."""
    if num_cols < 1:
        return row
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
    cell = ws.cell(row=row, column=1)
    cell.value = title
    cell.font = Font(bold=True, size=11)
    cell.fill = fill
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    return row + 1


def style_data_row(ws, row: int, *, bold: bool = False, fill=None) -> None:
    for col in range(1, ws.max_column + 1):
        c = ws.cell(row=row, column=col)
        if bold:
            c.font = Font(bold=True, size=11)
        if fill is not None:
            c.fill = fill
