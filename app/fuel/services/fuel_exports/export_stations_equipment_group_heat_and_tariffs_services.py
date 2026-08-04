# -*- coding: utf-8 -*-
"""Сервис экспорта «Тепло и тарифы из СТ» в Excel."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.common.services.help_services import (
    apply_nbsp_to_row,
    format_decimal_for_display,
)
from app.fuel.services.equipment_groups.equipment_group_heat_and_tariffs_services import (
    get_equipment_groups_with_heat_and_tariffs_data,
    build_equipment_group_heat_and_tariffs_hierarchy,
    HEAT_AND_TARIFFS_IDENTITY_COLUMNS,
    HEAT_AND_TARIFFS_METRIC_COLUMNS,
)
from app.fuel.services.fuel_exports.hierarchy_excel_layout import (
    EQUIPMENT_GROUP_ID_HEADER,
    FILL_EST,
    FILL_RES_HEADER,
    FILL_RES_SUMMARY,
    FILL_STATION_SUMMARY,
    FILL_UES,
    equipment_group_id_cell,
    next_data_row_index,
    style_data_row,
    write_merged_section_row,
)


def _format_numeric(val, rounding_digits):
    if val is None:
        return "—"
    return format_decimal_for_display(val, digits=rounding_digits)


def _format_identity(val):
    if val is None or val == "":
        return "—"
    return str(val)


def export_stations_equipment_group_heat_and_tariffs_to_excel(
    filters,
    start_year,
    end_year,
    *,
    per_page="all",
    page=1,
    show_all=True,
    rounding_digits=1,
):
    data = get_equipment_groups_with_heat_and_tariffs_data(
        filters=filters,
        per_page=per_page,
        page=page,
        start_year=start_year,
        end_year=end_year,
        show_all=show_all,
    )
    rows = data.get("rows") or []
    if not rows:
        return None

    years = list(range(int(start_year), int(end_year) + 1))
    hierarchy = build_equipment_group_heat_and_tariffs_hierarchy(rows, years=years)
    if not hierarchy:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "Тепло и тарифы из СТ"

    columns = (
        [EQUIPMENT_GROUP_ID_HEADER, "Группа оборудования"]
        + [label for _, label, _ in HEAT_AND_TARIFFS_IDENTITY_COLUMNS]
        + ["Показатель"]
        + [str(y) for y in years]
    )
    num_cols = len(columns)

    header_fill = PatternFill(
        start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
    )
    header_font = Font(bold=True, size=10)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num, column_title in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = apply_nbsp_to_row([column_title])[0]
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    max_lengths = [len(str(c)) for c in columns]

    def _bump_lengths(row_vals):
        for i, v in enumerate(row_vals):
            if v is not None and len(str(v)) > max_lengths[i]:
                max_lengths[i] = len(str(v))

    def _append_metric_rows(group_entity, group_name, identity, metric_rows):
        eg_id_cell = (
            equipment_group_id_cell(group_entity)
            if getattr(group_entity, "id", None) and group_entity.id > 0
            else "—"
        )
        identity_vals = [
            _format_identity(identity.get(attr))
            for attr, _label, _num in HEAT_AND_TARIFFS_IDENTITY_COLUMNS
        ]
        for metric in metric_rows or []:
            values_by_year = metric.get("values_by_year") or {}
            row_data = (
                [eg_id_cell, group_name]
                + identity_vals
                + [metric.get("label") or metric.get("attr") or "—"]
                + [
                    _format_numeric(values_by_year.get(y), rounding_digits)
                    for y in years
                ]
            )
            _bump_lengths(row_data)
            ws.append(apply_nbsp_to_row(row_data))

    def _append_summary_rows(label, summary, fill):
        summary = summary or {}
        q_by_year = summary.get("q") or {}
        for attr, metric_label, _is_num in HEAT_AND_TARIFFS_METRIC_COLUMNS:
            if attr == "q":
                year_vals = [
                    _format_numeric(q_by_year.get(y), rounding_digits) for y in years
                ]
            else:
                year_vals = ["—"] * len(years)
            row_data = (
                ["—", label]
                + ["—"] * len(HEAT_AND_TARIFFS_IDENTITY_COLUMNS)
                + [metric_label]
                + year_vals
            )
            _bump_lengths(row_data)
            ws.append(apply_nbsp_to_row(row_data))
            style_data_row(ws, ws.max_row, bold=True, fill=fill)

    for est_block in hierarchy:
        title = est_block.get("est_name") or "—"
        _bump_lengths([title] + [""] * (num_cols - 1))
        write_merged_section_row(ws, next_data_row_index(ws), num_cols, title, FILL_EST)

        for ues_block in est_block.get("ues_list") or []:
            ues_title = ues_block.get("ues_name") or "—"
            _bump_lengths([ues_title] + [""] * (num_cols - 1))
            write_merged_section_row(
                ws, next_data_row_index(ws), num_cols, ues_title, FILL_UES
            )

            for res_block in ues_block.get("res_list") or []:
                res_title = res_block.get("res_name") or "—"
                _bump_lengths([res_title] + [""] * (num_cols - 1))
                write_merged_section_row(
                    ws,
                    next_data_row_index(ws),
                    num_cols,
                    res_title,
                    FILL_RES_HEADER,
                )

                for station_block in res_block.get("station_blocks") or []:
                    for group_block in station_block.get("group_blocks") or []:
                        group_entity = group_block.get("equipment_group")
                        group_name = (
                            (
                                group_entity.name
                                if group_entity and group_entity.name
                                else None
                            )
                            or (
                                group_entity.name_ext
                                if group_entity and group_entity.name_ext
                                else None
                            )
                            or "—"
                        )
                        _append_metric_rows(
                            group_entity,
                            group_name,
                            group_block.get("identity") or {},
                            group_block.get("metric_rows") or [],
                        )

                    gb_count = len(station_block.get("group_blocks") or [])
                    if (
                        gb_count > 1
                        and not station_block.get("is_virtual")
                        and not station_block.get("suppress_station_summary")
                    ):
                        st_name = station_block.get("station_name") or "—"
                        _append_summary_rows(
                            f"{st_name}, всего",
                            station_block.get("station_summary"),
                            FILL_STATION_SUMMARY,
                        )

                if res_block.get("group_blocks"):
                    res_nm = res_block.get("res_name") or "—"
                    _append_summary_rows(
                        f"{res_nm}, всего",
                        res_block.get("res_summary"),
                        FILL_RES_SUMMARY,
                    )

    for i, length in enumerate(max_lengths, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(length + 2, 10), 40)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
