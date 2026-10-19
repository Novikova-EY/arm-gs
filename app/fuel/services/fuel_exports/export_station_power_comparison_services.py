"""Excel-выгрузка сравнения мощностей «Генерация» и «Топливо»."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.common.services.excel_numeric_cell import (
    excel_number_format,
    excel_numeric_cell_value,
)
from app.fuel.services.fuel_exports.hierarchy_excel_layout import (
    FILL_EST,
    FILL_RES_HEADER,
    FILL_STATION_SUMMARY,
    FILL_UES,
    style_data_row,
)


_SUMMARY_FILLS = {
    "system": FILL_EST,
    "ues": FILL_UES,
    "res": FILL_RES_HEADER,
    "unassigned": FILL_STATION_SUMMARY,
    "unassigned_station": FILL_STATION_SUMMARY,
    "unclassified": FILL_STATION_SUMMARY,
}
_SUMMARY_KINDS = {*_SUMMARY_FILLS, "station", "numb"}


def export_station_power_comparison_to_excel(comparison):
    """Использует готовые строки и итоги сравнения без повторного суммирования."""
    columns = [
        "Станция", "Код группы оборудования (numb)", "Группа оборудования",
        "Тип группы оборудования", "ст. №", "Тип агрегата",
    ]
    power_columns = [("p_ust", "Руст, МВт", "p_ust_mismatch")]
    if comparison["show_p_rasp"]:
        power_columns.append(("p_rasp", "Ррасп, МВт", "p_rasp_mismatch"))
    power_columns.append(("nust", "NUST, МВт", "p_ust_mismatch"))
    if comparison["show_p_rasp"]:
        power_columns.append(("nr", "NR, МВт", "p_rasp_mismatch"))
    columns.extend(title for _key, title, _mismatch in power_columns)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сравнение мощностей"
    sheet.append(['Сравнение мощностей в модулях «Топливо» и «Генерация»'])
    year = comparison["selected_year"]
    year_label = "Без года" if year in (None, "") else str(year)
    sheet.append([
        f"Год: {year_label}; VED=0: {'да' if comparison['show_ved0'] else 'нет'}; "
        f"Котельные: {'да' if comparison['show_boilers'] else 'нет'}. "
        "VED=0 и котельные не входят в расчётные суммы."
    ])
    for row_number in (1, 2):
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=len(columns))
        sheet.cell(row_number, 1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 30
    sheet.append(columns)
    style_data_row(sheet, 3, bold=True, fill=FILL_UES)
    for cell in sheet[3]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[3].height = 36

    number_format = excel_number_format(comparison["rounding_digits"])
    mismatch_fill = PatternFill(fill_type="solid", fgColor="F8D7DA")
    station_name = ""
    for row in comparison["display_rows"]:
        kind = row["kind"]
        if kind == "station":
            station_name = row["label"]
        elif kind in _SUMMARY_FILLS:
            station_name = ""
        if kind in _SUMMARY_KINDS:
            label = f"Итого по коду numb {row['numb']}" if kind == "numb" else row["label"]
            values = [label, None, None, None, None, None]
        else:
            values = [
                row.get("label") or row.get("station_name") or station_name,
                row.get("numb"), row.get("group_name"), row.get("group_type"),
                row.get("machine_number"), row.get("machine_type"),
            ]
        sheet.append(values)
        row_number = sheet.max_row
        for column, (key, _title, mismatch_key) in enumerate(power_columns, 7):
            value, is_numeric = excel_numeric_cell_value(row.get(key))
            cell = sheet.cell(row_number, column, value)
            if is_numeric:
                cell.number_format = number_format
            cell.alignment = Alignment(horizontal="right", vertical="center")
        if kind in _SUMMARY_KINDS:
            sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=6)
        style_data_row(
            sheet, row_number, bold=kind in _SUMMARY_KINDS or kind == "station_group",
            fill=_SUMMARY_FILLS.get(kind),
        )
        for cell in sheet[row_number][:6]:
            # Имена из БД остаются текстом, даже если начинаются со знака «=».
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        if row.get("title"):
            sheet.cell(row_number, 1).comment = Comment(row["title"], "АРМ ГС")
        for column, (_key, _title, mismatch_key) in enumerate(power_columns, 7):
            if row.get(mismatch_key):
                sheet.cell(row_number, column).fill = mismatch_fill

    widths = [45, 18, 55, 25, 18, 45] + [18] * len(power_columns)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "G4"
    sheet.print_title_rows = "1:3"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
