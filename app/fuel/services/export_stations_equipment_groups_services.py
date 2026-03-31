# -*- coding: utf-8 -*-
"""Сервис экспорта таблицы stations_equipment_groups в Excel."""

from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.common.services.help_services import apply_nbsp_to_row
from app.fuel.services.stations_equipment_groups_services import (
    build_equipment_group_hierarchy_eg_first,
)
from app.common.services.get_services.energy_systems.energy_system_type_get_services import (
    get_energy_system_type_map,
)
from app.common.services.get_services.energy_systems.union_energy_system_get_services import (
    get_union_energy_systems_map,
)
from app.common.services.get_services.energy_systems.regional_energy_system_get_services import (
    get_regional_energy_systems_map,
)


def export_stations_equipment_groups_to_excel(filters, start_year, end_year):
    """
    Экспортирует данные stations_equipment_groups в Excel.
    EquipmentGroup-first: иерархия строится из групп оборудования.
    """
    data_filters = {k: v for k, v in (filters or {}).items()}
    if (data_filters.get("equipment_group_name_filter") or "").strip():
        data_filters.pop("station_name_filter", None)
        data_filters.pop("gen_company_filter", None)

    est_names = get_energy_system_type_map()
    ues_names = get_union_energy_systems_map()
    res_names = get_regional_energy_systems_map()

    equipment_group_blocks_hierarchy = build_equipment_group_hierarchy_eg_first(
        filters=data_filters,
        start_year=start_year,
        end_year=end_year,
        energy_system_type_names=est_names,
        union_energy_system_names=ues_names,
        regional_energy_system_names=res_names,
    )
    
    # Создаем Excel файл
    wb = Workbook()
    ws = wb.active
    ws.title = "Группы оборудования"
    
    # Определяем столбцы ровно как на экране (см. stations_equipment_groups.html)
    columns = [
        "Группа оборудования",
        "Станция",
        "Тип",
        "ст. №",
        "Тип агрегата",
        "Субъект РФ",
        "Региональная энергосистема",
        "Генерирующая компания",
        "Признак группы оборудования (niv)",
        "Признак станции, разбитой на группы (comp)",
        "Код станции (main)",
        "Признак действующей станции (d)",
        "Признак расширяемой станции (r)",
        "Признак ФОРЭМ (forem)",
        "Ведомство (vedomstvo)",
        "Субъект РФ (obl)",
        "Департамент (dep)",
        "ОЭС (oes)",
        "Экономический район (er)",
        "Федеральный округ (fo)",
        "Код станции (numb)",
        "Типы турбин (tm)",
        "Мощность блока (вар. 1) (n1)",
        "Мощность блока (вар. 2) (n2)",
        "Давление пара (вар. 1) (p1)",
        "Давление пара (вар. 2) (p2)",
        "Порядковый номер станции (ordnumb)",
        "Адрес (addr)",
        "Примечание (note)",
        "Код города (codegor)",
        "Тип генерирующей компании (be)",
        "Генерирующая компания (gk)",
        "Филиал ГК (gkf)",
    ]
    
    # Заголовки
    header_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Предварительно считаем максимальные длины по заголовкам
    max_lengths = [len(str(col)) if col is not None else 0 for col in columns]
    
    for col_num, column_title in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = apply_nbsp_to_row([column_title])[0]
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    current_row = 2

    def _track_row_lengths(row):
        for idx, value in enumerate(row):
            if value is None:
                continue
            length = len(str(value))
            if length > max_lengths[idx]:
                max_lengths[idx] = length

    def _group_name(group_obj):
        if group_obj and group_obj.name:
            return group_obj.name
        if group_obj and group_obj.name_ext:
            return group_obj.name_ext
        return "—"

    def _write_group_fields(row, group_obj, station_for_company):
        if not group_obj:
            for col_idx in range(5, len(columns)):
                row[col_idx] = "—"
            return

        rd = getattr(group_obj, "regional_district", None)
        row[5] = rd.name if rd and rd.name else "—"
        res = getattr(group_obj, "regional_energy_system", None)
        row[6] = res.name if res and res.name else "—"
        row[7] = (
            getattr(station_for_company, "gen_companies", None) or "—"
            if station_for_company and hasattr(station_for_company, "gen_companies")
            else "—"
        )
        row[8] = group_obj.niv if group_obj.niv else "—"
        row[9] = group_obj.comp if group_obj.comp else "—"
        row[10] = group_obj.main if group_obj.main else "—"
        d_val = group_obj.d
        r_val = group_obj.r
        forem_val = group_obj.forem
        row[11] = "да" if d_val and str(d_val) == "1" else (d_val or "—")
        row[12] = "да" if r_val and str(r_val) == "1" else (r_val or "—")
        row[13] = "да" if forem_val and str(forem_val) == "1" else (forem_val or "—")
        _vedomstvo = group_obj.vedomstvo
        if _vedomstvo and str(_vedomstvo) == "1":
            row[14] = "станция\nотрасли"
        elif _vedomstvo and str(_vedomstvo) == "2":
            row[14] = "пром.\nпредприятие"
        else:
            row[14] = _vedomstvo or "—"
        terr_mapping = getattr(group_obj, "territories_energy_external_mapping", None)
        row[15] = (
            terr_mapping.external_name
            if terr_mapping and terr_mapping.external_name
            else (group_obj.obl or "—")
        )
        dep_mapping = getattr(group_obj, "department_external_mapping", None)
        row[16] = (
            dep_mapping.external_name
            if dep_mapping and dep_mapping.external_name
            else (group_obj.dep or "—")
        )
        ues_mapping = getattr(group_obj, "union_energy_system_external_mapping", None)
        row[17] = (
            ues_mapping.external_nameoes
            if ues_mapping and ues_mapping.external_nameoes
            else (group_obj.oes or "—")
        )
        er_mapping = getattr(group_obj, "economic_region_external_mapping", None)
        row[18] = (
            er_mapping.external_name
            if er_mapping and er_mapping.external_name
            else (group_obj.er or "—")
        )
        fd_mapping = getattr(group_obj, "federal_district_external_mapping", None)
        row[19] = (
            fd_mapping.external_name
            if fd_mapping and fd_mapping.external_name
            else (group_obj.fo or "—")
        )
        row[20] = group_obj.numb if group_obj.numb else "—"
        row[21] = group_obj.tm if group_obj.tm else "—"
        row[22] = group_obj.n1 if group_obj.n1 else "—"
        row[23] = group_obj.n2 if group_obj.n2 else "—"
        row[24] = group_obj.p1 if group_obj.p1 else "—"
        row[25] = group_obj.p2 if group_obj.p2 else "—"
        row[26] = group_obj.ordnumb if group_obj.ordnumb else "—"
        row[27] = group_obj.addr if group_obj.addr else "—"
        row[28] = group_obj.note if group_obj.note else "—"
        row[29] = (
            group_obj.cities_external_mapping.name
            if group_obj.cities_external_mapping
            else (group_obj.codegor or "—")
        )
        bu_mapping = getattr(group_obj, "business_unit_external_mapping", None)
        row[30] = (
            bu_mapping.external_name
            if bu_mapping and bu_mapping.external_name
            else (group_obj.be or "—")
        )
        gk_mapping = getattr(group_obj, "gen_company_external_mapping", None)
        row[31] = (
            gk_mapping.external_name
            if gk_mapping and gk_mapping.external_name
            else (group_obj.gk or "—")
        )
        gkf_mapping = getattr(group_obj, "gen_company_branch_external_mapping", None)
        row[32] = (
            gkf_mapping.external_name
            if gkf_mapping and gkf_mapping.external_name
            else (group_obj.gkf or "—")
        )

    def _render_station_blocks(station_blocks):
        nonlocal current_row

        for station_block in station_blocks:
            station = station_block.get("station")
            if not station:
                continue
            station_name = (
                station.name if hasattr(station, "name") and station.name else "—"
            )

            for group_block in station_block.get("group_blocks") or []:
                group_obj = group_block.get("equipment_group")
                group_first_row = None

                for link in group_block.get("links") or []:
                    link_first_row = None
                    group_type = link.get("equipment_group_type")
                    machines = link.get("machines") or [None]

                    for machine in machines:
                        row = [None] * len(columns)
                        row[1] = station_name
                        row[3] = (
                            getattr(machine, "machine_number", None) or "—"
                            if machine else "—"
                        )
                        row[4] = (
                            getattr(machine, "machine_name", None) or "—"
                            if machine else "—"
                        )

                        if group_first_row is None:
                            group_first_row = current_row
                            row[0] = _group_name(group_obj)

                        _write_group_fields(row, group_obj, station)

                        if link_first_row is None:
                            link_first_row = current_row
                            row[2] = (
                                group_type.name
                                if group_type and getattr(group_type, "name", None)
                                else "—"
                            )

                        _track_row_lengths(row)
                        ws.append(apply_nbsp_to_row(row))
                        current_row += 1

                    if link_first_row is not None and current_row - 1 > link_first_row:
                        ws.merge_cells(
                            start_row=link_first_row,
                            start_column=3,
                            end_row=current_row - 1,
                            end_column=3,
                        )

                if group_first_row is not None and current_row - 1 > group_first_row:
                    ws.merge_cells(
                        start_row=group_first_row,
                        start_column=1,
                        end_row=current_row - 1,
                        end_column=1,
                    )

    def _render_multi_station_group_blocks(group_blocks):
        nonlocal current_row

        for group_block in group_blocks:
            group_obj = group_block.get("equipment_group")
            group_first_row = None
            first_station = (
                group_block["station_entries"][0].get("station")
                if group_block.get("station_entries")
                else None
            )

            for station_entry in group_block.get("station_entries") or []:
                station = station_entry.get("station")
                if not station:
                    continue
                station_name = (
                    station.name if hasattr(station, "name") and station.name else "—"
                )

                for link in station_entry.get("links") or []:
                    link_first_row = None
                    group_type = link.get("equipment_group_type")
                    machines = link.get("machines") or [None]

                    for machine in machines:
                        row = [None] * len(columns)
                        row[1] = station_name
                        row[3] = (
                            getattr(machine, "machine_number", None) or "—"
                            if machine else "—"
                        )
                        row[4] = (
                            getattr(machine, "machine_name", None) or "—"
                            if machine else "—"
                        )

                        if group_first_row is None:
                            group_first_row = current_row
                            row[0] = _group_name(group_obj)
                            _write_group_fields(row, group_obj, first_station)

                        if link_first_row is None:
                            link_first_row = current_row
                            row[2] = (
                                group_type.name
                                if group_type and getattr(group_type, "name", None)
                                else "—"
                            )

                        _track_row_lengths(row)
                        ws.append(apply_nbsp_to_row(row))
                        current_row += 1

                    if link_first_row is not None and current_row - 1 > link_first_row:
                        ws.merge_cells(
                            start_row=link_first_row,
                            start_column=3,
                            end_row=current_row - 1,
                            end_column=3,
                        )

            if group_first_row is not None and current_row - 1 > group_first_row:
                ws.merge_cells(
                    start_row=group_first_row,
                    start_column=1,
                    end_row=current_row - 1,
                    end_column=1,
                )
                for col in range(6, len(columns) + 1):
                    ws.merge_cells(
                        start_row=group_first_row,
                        start_column=col,
                        end_row=current_row - 1,
                        end_column=col,
                    )

    for est_block in equipment_group_blocks_hierarchy:
        es_type_name = est_block.get("est_name", "—")
        row = [None] * len(columns)
        row[2] = es_type_name
        _track_row_lengths(row)
        ws.append(apply_nbsp_to_row(row))
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(columns))
        cell = ws.cell(row=current_row, column=1)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        current_row += 1

        for ues_block in est_block.get("ues_list") or []:
            ues_name = ues_block.get("ues_name", "—")
            row = [None] * len(columns)
            row[2] = ues_name
            _track_row_lengths(row)
            ws.append(apply_nbsp_to_row(row))
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(columns))
            cell = ws.cell(row=current_row, column=1)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            current_row += 1

            for res_block in ues_block.get("res_list") or []:
                res_name = res_block.get("res_name", "—")
                row = [None] * len(columns)
                row[2] = res_name
                _track_row_lengths(row)
                ws.append(apply_nbsp_to_row(row))
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(columns))
                cell = ws.cell(row=current_row, column=1)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCFFFF", end_color="CCFFFF", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                current_row += 1
                _render_station_blocks(res_block.get("station_blocks") or [])
                _render_multi_station_group_blocks(
                    res_block.get("multi_station_group_blocks") or []
                )
                _render_station_blocks(res_block.get("boiler_station_blocks") or [])
    
    # Автоподбор ширины столбцов (по предвычисленным длинам)
    for col_idx, max_length in enumerate(max_lengths, start=1):
        adjusted_width = min(max_length + 2, 60)
        column_letter = get_column_letter(col_idx)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Сохраняем в BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
