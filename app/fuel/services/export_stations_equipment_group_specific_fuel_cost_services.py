# -*- coding: utf-8 -*-
"""Сервис экспорта стоимости для групп оборудования в Excel."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.common.services.help_services import apply_nbsp_to_row, format_decimal_for_display
from app.fuel.services.equipment_group_specific_fuel_cost_services import (
    get_equipment_groups_with_specific_fuel_cost_data,
    SPECIFIC_FUEL_COST_COLUMNS,
)


def _format_cell_value(param, attr, is_numeric, rounding_digits):
    """Форматирует значение ячейки."""
    if param is None:
        return "—"
    val = getattr(param, attr, None)
    if val is None:
        return "—"
    if attr == "ved":
        s = str(val)
        if s == "1":
            return "КЭС"
        if s == "2":
            return "ТЭЦ"
        return val
    if is_numeric:
        return format_decimal_for_display(val, digits=rounding_digits)
    return str(val) if val is not None else "—"


def export_stations_equipment_group_specific_fuel_cost_to_excel(
    filters,
    start_year,
    end_year,
    *,
    per_page="all",
    page=1,
    show_all=True,
    rounding_digits=1,
):
    """
    Экспортирует стоимость для групп оборудования в Excel.
    Учитывает фильтры: субъекты, ОЭС, региональные энергосистемы и т.п.
    """
    data = get_equipment_groups_with_specific_fuel_cost_data(
        filters=filters,
        per_page=per_page,
        page=page,
        start_year=start_year,
        end_year=end_year,
        show_all=show_all,
    )
    rows = data.get("rows") or []
    if not rows:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "Стоимость"

    columns = ["Группа оборудования"] + [
        label for _, label, _ in SPECIFIC_FUEL_COST_COLUMNS
    ]

    header_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    header_font = Font(bold=True, size=10)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num, column_title in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = apply_nbsp_to_row([column_title])[0]
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    max_lengths = [len(str(c)) for c in columns]

    for row_idx, (group_entity, param) in enumerate(rows, start=2):
        group_name = (
            (group_entity.name if group_entity and group_entity.name else None)
            or (group_entity.name_ext if group_entity and group_entity.name_ext else None)
            or "—"
        )
        row_data = [group_name]
        for attr, _label, is_numeric in SPECIFIC_FUEL_COST_COLUMNS:
            cell_val = _format_cell_value(param, attr, is_numeric, rounding_digits)
            row_data.append(cell_val)
            for i, v in enumerate(row_data):
                if v is not None and len(str(v)) > max_lengths[i]:
                    max_lengths[i] = len(str(v))
        ws.append(apply_nbsp_to_row(row_data))

    for col_idx, max_length in enumerate(max_lengths, start=1):
        adjusted_width = min(max_length + 2, 20)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
