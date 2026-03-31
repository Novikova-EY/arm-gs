# -*- coding: utf-8 -*-
"""Сервис экспорта топливных параметров групп оборудования в Excel."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.common.services.help_services import apply_nbsp_to_row, format_decimal_for_display
from app.fuel.services.equipment_group_fuel_params_services import (
    get_equipment_groups_with_fuel_params_data,
    build_name_maps_from_rows,
)

FUEL_PARAM_COLUMNS = [
    ("numb1120", "Код станции", False),
    ("nust", "Руст", True),
    ("nr", "Ррасп", True),
    ("e", "Выр", True),
    ("ewtp", "Этц", True),
    ("eotp", "Отпуск ээ", True),
    ("eurt", "Уд.расх ээ", True),
    ("eust", "Расх топ ээ", True),
    ("snk", "СН, %", True),
    ("q", "Отпуск, Гкал", True),
    ("qotr", "Отраб", True),
    ("turt", "Уд.расх тэ", True),
    ("tust", "Расх топ тэ", True),
    ("sn_t", "СН, кВтч/Гкал", True),
    ("b", "Расх топл.", True),
    ("gaz", "Газ", True),
    ("isk_gaz", "Иск. газ", True),
    ("mazut", "Мазут", True),
    ("torf", "Торф", True),
    ("slan", "Сланцы", True),
    ("proch", "Прочее", True),
    ("ugol", "Уголь", True),
    ("don", "Дон", True),
    ("podm", "Подм", True),
    ("pech", "Печ", True),
    ("arkt", "Арктикуголь", True),
    ("kuzn", "Кузбасс", True),
    ("ural", "Урал", True),
    ("bashk", "Башкортостан", True),
    ("kazah", "Казахстан", True),
    ("kan", "Канск", True),
    ("tung", "Тунгусск", True),
    ("irkut", "Иркутск", True),
    ("hak", "Хакасия", True),
    ("tuv", "Тува", True),
    ("bur", "Бурятия", True),
    ("chit", "Чита", True),
    ("yakut", "Якутия", True),
    ("amur", "Амур", True),
    ("urg", "Юрга", True),
    ("ushum", "Ушумун", True),
    ("prim", "Приморье", True),
    ("mag", "Магадан", True),
    ("chukot", "Чукотка", True),
    ("kamch", "Камчатка", True),
    ("sah", "Сахалин", True),
    ("nt", "Тепл. мощн. отборов", True),
    ("nt_sum", "Сумма NT", True),
]


def _format_cell_value(param, attr, is_numeric, name_maps, rounding_digits):
    """Форматирует значение ячейки по аналогии с шаблоном."""
    if param is None:
        return "—"
    val = getattr(param, attr, None)
    if val is None:
        return "—"
    if attr == "obor" and name_maps.get("obor_name_map"):
        return name_maps["obor_name_map"].get(val, val) or "—"
    if attr == "obl" and name_maps.get("obl_name_map"):
        return name_maps["obl_name_map"].get(val, val) or "—"
    if attr == "dep" and name_maps.get("dep_name_map"):
        return name_maps["dep_name_map"].get(val, val) or "—"
    if attr == "oes" and name_maps.get("oes_name_map"):
        return name_maps["oes_name_map"].get(val, val) or "—"
    if attr == "er" and name_maps.get("er_name_map"):
        return name_maps["er_name_map"].get(val, val) or "—"
    if attr == "gk" and name_maps.get("gk_name_map"):
        return name_maps["gk_name_map"].get(val, val) or "—"
    if attr == "be" and name_maps.get("be_name_map"):
        return name_maps["be_name_map"].get(val, val) or "—"
    if attr == "ved":
        s = str(val)
        if s == "1":
            return "КЭС"
        if s == "2":
            return "ТЭЦ"
        return val
    if attr == "ved_cyrillic":
        s = str(val)
        if s == "1":
            return "станция отрасли"
        if s == "2":
            return "пром.предприятие"
        return val
    if is_numeric:
        return format_decimal_for_display(val, digits=rounding_digits)
    return str(val) if val is not None else "—"


def export_stations_equipment_group_fuel_params_to_excel(
    filters,
    start_year,
    end_year,
    *,
    per_page="all",
    page=1,
    show_all=True,
    rounding_digits=1,
):
    """
    Экспортирует топливные параметры групп оборудования в Excel.
    Учитывает фильтры: субъекты, ОЭС, региональные энергосистемы и т.п.
    """
    fuel_data = get_equipment_groups_with_fuel_params_data(
        filters=filters,
        per_page=per_page,
        page=page,
        start_year=start_year,
        end_year=end_year,
        show_all=show_all,
    )
    rows = fuel_data.get("rows") or []
    if not rows:
        return None

    name_maps = build_name_maps_from_rows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Топливные параметры"

    columns = ["Группа оборудования"] + [label for _, label, _ in FUEL_PARAM_COLUMNS]

    header_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num, column_title in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = apply_nbsp_to_row([column_title])[0]
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    max_lengths = [len(str(c)) for c in columns]

    for row_idx, (group_entity, param) in enumerate(rows, start=2):
        group_name = (
            (group_entity.name if group_entity and group_entity.name else None)
            or (group_entity.name_ext if group_entity and group_entity.name_ext else None)
            or "—"
        )
        row_data = [group_name]
        for attr, _label, is_numeric in FUEL_PARAM_COLUMNS:
            cell_val = _format_cell_value(
                param, attr, is_numeric, name_maps, rounding_digits
            )
            row_data.append(cell_val)
            for i, v in enumerate(row_data):
                if v is not None and len(str(v)) > max_lengths[i]:
                    max_lengths[i] = len(str(v))
        ws.append(apply_nbsp_to_row(row_data))

    for col_idx, max_length in enumerate(max_lengths, start=1):
        adjusted_width = min(max_length + 2, 60)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
