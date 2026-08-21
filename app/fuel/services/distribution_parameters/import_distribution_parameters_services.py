# -*- coding: utf-8 -*-
"""Импорт параметров распределения из Excel (формат Access «Параметры-распределения»)."""

from __future__ import annotations

import io
import re
from decimal import Decimal, InvalidOperation
from typing import BinaryIO

from openpyxl import load_workbook
from sqlalchemy.orm import joinedload

from app.extensions import db
from app.common.services.database_version_filter import get_current_db_version_id
from app.fuel.models.external_mapping.fue_em_union_energy_system_model import (
    UnionEnergySystemExternalMapping,
)
from app.fuel.models.fue_distribution_parameter_model import DistributionParameter
from app.refdata.models.energy_systems.union_energy_system_model import UnionEnergySystem
from app.refdata.models.years.year_model import Year

# Поля Access / шаблона экспорта АРМ (порядок столбцов любой).
_EXCEL_CANONICAL_HEADERS = frozenset(
    {
        "name",
        "year",
        "filter",
        "e",
        "kplus",
        "kmin",
        "k",
        "bkl",
        "wname",
        "uname",
        "toplname",
        "dopname",
        "byear",
        "kn",
        "knps",
        "kngt",
        "knpg",
        "hnps",
        "hngt",
        "hnpg",
        "numb",
        "doptim",
        "lim",
    }
)

# Минимально нужные для строки (остальные колонки — по возможности).
_REQUIRED_HEADERS = frozenset({"year"})

# Синонимы шапки: Access, экспорт АРМ, русские подписи.
_HEADER_ALIASES: dict[str, str] = {
    "name": "name",
    "оэс": "name",
    "объединенная энергосистема": "name",
    "year": "year",
    "год": "year",
    "расчетный год": "year",
    "расчитываемый год": "year",
    "filter": "filter",
    "filter_text": "filter",
    "фильтр": "filter",
    "e": "e",
    "выработка тэс": "e",
    "kplus": "kplus",
    "kmin": "kmin",
    "k": "k",
    "bkl": "bkl",
    "wname": "wname",
    "uname": "uname",
    "toplname": "toplname",
    "dopname": "dopname",
    "byear": "byear",
    "базовый год": "byear",
    "kn": "kn",
    "knps": "knps",
    "kngt": "kngt",
    "knpg": "knpg",
    "hnps": "hnps",
    "hngt": "hngt",
    "hnpg": "hnpg",
    "numb": "numb",
    "doptim": "doptim",
    "lim": "lim",
    # Служебное поле Access — игнорируется при разборе шапки.
    "wt_recid": "",
}

_OES_IN_FILTER_RE = re.compile(r"oes\s*=\s*(\d+)", re.IGNORECASE)

# Access «Параметры-распределения».name / nameoes → UnionEnergySystem.name (casefold).
# В БД Топливо oes=10 — «Норильск» / «Норильск.эн.р-н», в АРМ — «ТИТЭС Сибири».
_ACCESS_UES_LABEL_ALIASES: dict[str, str] = {
    "норильск": "титэс сибири",
    "норильск.эн.р-н": "титэс сибири",
}


def _normalize_header_cell(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    low = s.lower()
    if low in _HEADER_ALIASES:
        alias = _HEADER_ALIASES[low]
        return alias or None  # "" → игнор (WT_RECID)
    if low in _EXCEL_CANONICAL_HEADERS:
        return low
    return None


def _cell_str(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        t = value.strip()
        return t or None
    return str(value).strip() or None


def _parse_decimal(value):
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    s = str(value).strip().replace(",", ".")
    if not s:
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _parse_int(value):
    if value is None or value == "":
        return None
    if isinstance(value, int) and not isinstance(value, bool):
        return value
    if isinstance(value, float):
        if value != value:  # NaN
            return None
        return int(round(value))
    s = str(value).strip()
    if not s:
        return None
    try:
        return int(round(float(s.replace(",", "."))))
    except (ValueError, TypeError):
        return None


def _parse_year_num(value):
    return _parse_int(value)


def _parse_oes_from_filter(filter_text: str | None) -> str | None:
    if not filter_text:
        return None
    m = _OES_IN_FILTER_RE.search(str(filter_text))
    return m.group(1) if m else None


def _year_id_by_number(year_num: int | None, version_id: int | None) -> int | None:
    if year_num is None:
        return None
    q = Year.query.filter(Year.number == year_num)
    if version_id is not None:
        q = q.filter(Year.database_version_id == version_id)
    else:
        q = q.filter(Year.database_version_id.is_(None))
    row = q.first()
    return row.id if row else None


def _apply_access_ues_aliases(
    *,
    label_to_ues_id: dict[str, int],
    oes_to_ues_id: dict[str, int],
    name_to_ues_id: dict[str, int],
    unmatched_mappings: list,
) -> None:
    """Access-имена, которые не совпадают 1:1 с external_name / external_nameoes."""
    for access_label, ues_name in _ACCESS_UES_LABEL_ALIASES.items():
        uid = name_to_ues_id.get(ues_name)
        if uid is None:
            continue
        label_to_ues_id.setdefault(access_label, uid)

    for m in unmatched_mappings:
        labels: list[str] = []
        for raw in (getattr(m, "external_nameoes", None), getattr(m, "external_name", None)):
            if raw and str(raw).strip():
                labels.append(str(raw).strip().casefold())
        uid = None
        for lab in labels:
            alias_ues = _ACCESS_UES_LABEL_ALIASES.get(lab)
            if alias_ues:
                uid = name_to_ues_id.get(alias_ues)
            if uid is None:
                uid = label_to_ues_id.get(lab)
            if uid is not None:
                break
        if uid is None:
            continue
        ext_id = getattr(m, "external_id", None)
        if ext_id is not None and str(ext_id).strip():
            oes_to_ues_id.setdefault(str(ext_id).strip(), uid)
        for lab in labels:
            label_to_ues_id.setdefault(lab, uid)


def _build_ues_resolvers(version_id: int | None):
    """
    Сопоставление Access → UnionEnergySystem.id только через
    ``UnionEnergySystemExternalMapping`` (/fuel/refdata/union_energy_system):
    external_id (oes), external_name, external_nameoes → ref_uuid → UES.
    """
    ues_q = UnionEnergySystem.query
    if version_id is not None:
        ues_q = ues_q.filter(UnionEnergySystem.database_version_id == version_id)
    else:
        ues_q = ues_q.filter(UnionEnergySystem.database_version_id.is_(None))

    by_uuid: dict[str, int] = {}
    name_to_ues_id: dict[str, int] = {}
    for u in ues_q.all():
        if getattr(u, "ref_uuid", None):
            by_uuid[u.ref_uuid] = u.id
        if getattr(u, "name", None) and str(u.name).strip():
            name_to_ues_id[str(u.name).strip().casefold()] = u.id

    oes_to_ues_id: dict[str, int] = {}
    label_to_ues_id: dict[str, int] = {}
    unmatched_mappings: list = []

    mappings = UnionEnergySystemExternalMapping.query.all()
    for m in mappings:
        ues_id = by_uuid.get(m.union_energy_system_ref_uuid or "")
        if ues_id is None:
            unmatched_mappings.append(m)
            continue
        if m.external_id is not None and str(m.external_id).strip():
            oes_to_ues_id[str(m.external_id).strip()] = ues_id
        for label in (m.external_nameoes, m.external_name):
            if label and str(label).strip():
                label_to_ues_id[str(label).strip().casefold()] = ues_id
        if m.external_name and str(m.external_name).strip():
            # Access: «ТЭС Волгаэнерго» при external_name=«Волгаэнерго»
            tes = f"ТЭС {str(m.external_name).strip()}"
            label_to_ues_id[tes.casefold()] = ues_id

    _apply_access_ues_aliases(
        label_to_ues_id=label_to_ues_id,
        oes_to_ues_id=oes_to_ues_id,
        name_to_ues_id=name_to_ues_id,
        unmatched_mappings=unmatched_mappings,
    )
    return oes_to_ues_id, label_to_ues_id


def _resolve_union_energy_system_id(
    *,
    name_raw: str | None,
    filter_text: str | None,
    oes_to_ues_id: dict[str, int],
    label_to_ues_id: dict[str, int],
) -> int | None:
    """Только через маппинг /fuel/refdata/union_energy_system (не напрямую по UES.name)."""
    oes = _parse_oes_from_filter(filter_text)
    if oes is not None and oes in oes_to_ues_id:
        return oes_to_ues_id[oes]

    if not name_raw:
        return None

    key = name_raw.strip().casefold()
    if key in label_to_ues_id:
        return label_to_ues_id[key]
    if key.startswith("тэс "):
        short = key[4:].strip()
        if short in label_to_ues_id:
            return label_to_ues_id[short]
    return None


def _apply_fields_to_dp(dp: DistributionParameter, fields: dict) -> None:
    for attr, value in fields.items():
        setattr(dp, attr, value)


def _normalize_e_key(value) -> str | None:
    """Канонический ключ выработки для сравнения уникальности."""
    if value is None or value == "":
        return None
    d = value if isinstance(value, Decimal) else _parse_decimal(value)
    if d is None:
        return None
    # Без экспоненты и хвоста нулей: 61030 == 61030.0
    return format(d.normalize(), "f")


def _ues_ref_uuid_by_id(ues_id: int | None, version_id: int | None) -> str | None:
    if ues_id is None:
        return None
    q = UnionEnergySystem.query.filter(UnionEnergySystem.id == ues_id)
    if version_id is not None:
        q = q.filter(UnionEnergySystem.database_version_id == version_id)
    row = q.first()
    if row is None:
        row = db.session.get(UnionEnergySystem, ues_id)
    return row.ref_uuid if row else None


def _prefer_dp_row(
    a: DistributionParameter | None, b: DistributionParameter
) -> DistributionParameter:
    """При выборе донора данных: с заполненной выработкой, иначе с большим id."""
    if a is None:
        return b
    a_e = a.e is not None
    b_e = b.e is not None
    if a_e != b_e:
        return b if b_e else a
    return b if (b.id or 0) >= (a.id or 0) else a


def _prefer_dp_keep(
    a: DistributionParameter | None, b: DistributionParameter
) -> DistributionParameter:
    """Какую строку оставить: меньший id (стабильные FK на сводки «Коэфф»)."""
    if a is None:
        return b
    return a if (a.id or 0) <= (b.id or 0) else b


def _delete_dp_with_dependents(dp: DistributionParameter) -> None:
    """Удаляет параметр и зависимые результаты расчёта (ORM иначе ставит FK в NULL)."""
    from app.fuel.models.coefficient.distribution_coefficient_summary_model import (
        DistributionCoefficientSummary,
    )
    from app.fuel.models.coefficient.equipment_group_coefficient_result_model import (
        EquipmentGroupCoefficientResult,
    )

    if not getattr(dp, "id", None):
        return
    DistributionCoefficientSummary.query.filter_by(
        distribution_parameter_id=dp.id
    ).delete(synchronize_session=False)
    EquipmentGroupCoefficientResult.query.filter_by(
        distribution_parameter_id=dp.id
    ).delete(synchronize_session=False)
    db.session.delete(dp)


def _logical_row_key(
    ues_ref_uuid: str | None,
    year_num: int | None,
    byear_num: int | None,
    e_key: str | None,
) -> tuple[str | None, int | None, int | None, str | None]:
    """Уникальность: ОЭС (ref_uuid) + базовый год + расчётный год + выработка."""
    return (ues_ref_uuid, year_num, byear_num, e_key)


def _load_existing_dp_indexes() -> tuple[
    dict[tuple[str | None, int | None, int | None, str | None], DistributionParameter],
    dict[tuple[str | None, int | None, int | None], list[DistributionParameter]],
]:
    """
    Индексы существующих строк по всем версиям БД (id ОЭС/Year разные, смысл один).
    """
    rows = (
        DistributionParameter.query.options(
            joinedload(DistributionParameter.union_energy_system),
            joinedload(DistributionParameter.year),
            joinedload(DistributionParameter.base_year),
        ).all()
    )
    by_full: dict[
        tuple[str | None, int | None, int | None, str | None], DistributionParameter
    ] = {}
    by_oes_years: dict[
        tuple[str | None, int | None, int | None], list[DistributionParameter]
    ] = {}
    for r in rows:
        ues = r.union_energy_system
        uuid = (ues.ref_uuid if ues and ues.ref_uuid else None) or (
            ues.name if ues else None
        )
        yn = r.year.number if r.year else None
        byn = r.base_year.number if r.base_year else None
        ek = _normalize_e_key(r.e)
        full = _logical_row_key(uuid, yn, byn, ek)
        by_full[full] = _prefer_dp_row(by_full.get(full), r)
        oes_years = (uuid, yn, byn)
        by_oes_years.setdefault(oes_years, []).append(r)
    return by_full, by_oes_years


def _find_existing_dp_for_import(
    *,
    ues_ref_uuid: str | None,
    year_num: int | None,
    byear_num: int | None,
    e_key: str | None,
    by_full: dict,
    by_oes_years: dict,
) -> DistributionParameter | None:
    """
    Ищем существующую строку:
    1) точное совпадение ОЭС + годы + выработка;
    2) та же ОЭС + годы с пустой выработкой (слот под импорт).
    Среди кандидатов предпочитаем меньший id (сохраняем FK сводок).
    """
    candidates: list[DistributionParameter] = []
    exact = by_full.get(_logical_row_key(ues_ref_uuid, year_num, byear_num, e_key))
    if exact is not None:
        candidates.append(exact)
    if e_key is not None:
        vacant = by_full.get(_logical_row_key(ues_ref_uuid, year_num, byear_num, None))
        if vacant is not None and vacant not in candidates:
            candidates.append(vacant)
    group = by_oes_years.get((ues_ref_uuid, year_num, byear_num)) or []
    for r in group:
        if _normalize_e_key(r.e) is None and r not in candidates:
            candidates.append(r)
    if not candidates:
        return None
    keep: DistributionParameter | None = None
    for r in candidates:
        keep = _prefer_dp_keep(keep, r)
    return keep


def import_distribution_parameters_from_excel(
    file_obj: BinaryIO,
    *,
    database_version_id: int | None = None,
) -> tuple[int, int, list[str]]:
    """
    Читает первый лист Excel (Access «Параметры-распределения» / экспорт АРМ)
    и пишет строки **во все версии БД** (общий ``ref_uuid``).

    ОЭС: через ``UnionEnergySystemExternalMapping`` (/fuel/refdata/union_energy_system).
    Уникальность в файле: ОЭС + базовый год + расчётный год + выработка.
    Сопоставление в БД: ОЭС (ref_uuid) + базовый год + расчётный год.

    Returns:
        (добавлено копий по версиям, обновлено копий, список_ошибок/предупреждений)
    """
    from app.fuel.services.distribution_parameters.distribution_parameters_all_versions_services import (
        upsert_distribution_parameter_in_all_versions,
    )
    from app.refdata.services.refdata_all_versions_common import (
        all_database_version_ids_for_refdata,
    )

    if database_version_id is None:
        database_version_id = get_current_db_version_id()

    if not all_database_version_ids_for_refdata():
        raise ValueError(
            "В системе нет зарегистрированных версий БД — импорт во все версии невозможен."
        )

    wb = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValueError("Файл пустой или нет строки заголовков.")

        col_index: dict[str, int] = {}
        for i, cell in enumerate(header_row):
            key = _normalize_header_cell(cell)
            if key and key not in col_index:
                col_index[key] = i

        missing = sorted(_REQUIRED_HEADERS - set(col_index.keys()))
        if missing:
            raise ValueError(
                "В первой строке не хватает столбцов: "
                + ", ".join(missing)
                + ". Ожидается выгрузка Access «Параметры-распределения» "
                "(name, year, filter, e, …) или экспорт с этой страницы."
            )

        # Резолв Access-имён ОЭС через маппинг + UES текущей версии → ref_uuid
        oes_to_ues_id, label_to_ues_id = _build_ues_resolvers(database_version_id)
        ues_uuid_by_id: dict[int, str | None] = {}
        for uid in set(oes_to_ues_id.values()) | set(label_to_ues_id.values()):
            ues_uuid_by_id[uid] = _ues_ref_uuid_by_id(uid, database_version_id)

        prepared: list[dict] = []
        errors: list[str] = []
        row_no = 1
        seen_keys: set[tuple] = set()

        for data_row in rows_iter:
            row_no += 1
            if not data_row or all(v is None or str(v).strip() == "" for v in data_row):
                continue

            def col(key: str):
                j = col_index.get(key)
                if j is None or j >= len(data_row):
                    return None
                return data_row[j]

            year_num = _parse_year_num(col("year"))
            if year_num is None:
                errors.append(f"Строка {row_no}: не задан или не распознан год (year).")
                continue

            # Год должен существовать хотя бы в текущей версии (якорь для проверки файла)
            if _year_id_by_number(year_num, database_version_id) is None:
                errors.append(
                    f"Строка {row_no}: для года {year_num} не найдена запись Year "
                    f"в справочнике (текущая версия БД)."
                )
                continue

            name_raw = _cell_str(col("name"))
            filter_text = _cell_str(col("filter"))
            id_ues = _resolve_union_energy_system_id(
                name_raw=name_raw,
                filter_text=filter_text,
                oes_to_ues_id=oes_to_ues_id,
                label_to_ues_id=label_to_ues_id,
            )
            if name_raw and id_ues is None:
                errors.append(
                    f"Строка {row_no}: ОЭС «{name_raw}» не сопоставлена в "
                    f"/fuel/refdata/union_energy_system "
                    f"(oes из filter / external_name / external_nameoes → UUID ОЭС)."
                )
                continue
            if not name_raw and id_ues is None and filter_text:
                errors.append(
                    f"Строка {row_no}: не удалось определить ОЭС по filter «{filter_text}» "
                    f"(нет связи oes→UUID на /fuel/refdata/union_energy_system)."
                )
                continue

            byear_num = _parse_year_num(col("byear"))
            if byear_num is not None and _year_id_by_number(byear_num, database_version_id) is None:
                errors.append(
                    f"Строка {row_no}: базовый год {byear_num} (byear) не найден "
                    f"в справочнике Year."
                )
                continue

            e_val = _parse_decimal(col("e"))
            e_key = _normalize_e_key(e_val)
            ues_uuid = None
            if id_ues is not None:
                ues_uuid = ues_uuid_by_id.get(id_ues)
                if id_ues not in ues_uuid_by_id:
                    ues_uuid = _ues_ref_uuid_by_id(id_ues, database_version_id)
                    ues_uuid_by_id[id_ues] = ues_uuid

            logical_key = _logical_row_key(ues_uuid, year_num, byear_num, e_key)
            byear_label = str(byear_num) if byear_num is not None else "не задан"
            oes_label = name_raw if name_raw else "не задана"
            e_label = e_key if e_key is not None else "не задана"
            if logical_key in seen_keys:
                errors.append(
                    f"Строка {row_no}: в файле уже есть строка с той же комбинацией "
                    f"ОЭС «{oes_label}», базовый год {byear_label}, "
                    f"расчётный год {year_num}, выработка {e_label}."
                )
                continue
            seen_keys.add(logical_key)

            prepared.append(
                {
                    "ues_ref_uuid": ues_uuid,
                    "year_number": year_num,
                    "base_year_number": byear_num,
                    "data": {
                        "filter_text": filter_text,
                        "e": e_val,
                        "kplus": _parse_decimal(col("kplus")),
                        "kmin": _parse_decimal(col("kmin")),
                        "k": _parse_decimal(col("k")),
                        "bkl": _parse_decimal(col("bkl")),
                        "wname": _cell_str(col("wname")),
                        "uname": _cell_str(col("uname")),
                        "toplname": _cell_str(col("toplname")),
                        "dopname": _cell_str(col("dopname")),
                        "kn": _parse_decimal(col("kn")),
                        "knps": _parse_decimal(col("knps")),
                        "kngt": _parse_decimal(col("kngt")),
                        "knpg": _parse_decimal(col("knpg")),
                        "hnps": _parse_decimal(col("hnps")),
                        "hngt": _parse_decimal(col("hngt")),
                        "hnpg": _parse_decimal(col("hnpg")),
                        "numb": _parse_decimal(col("numb")),
                        "doptim": _parse_decimal(col("doptim")),
                        "lim": _parse_int(col("lim")),
                    },
                }
            )

        if errors:
            return 0, 0, errors

        if not prepared:
            return 0, 0, []

        n_added = 0
        n_updated = 0
        warnings: list[str] = []
        try:
            for item in prepared:
                _ref, added, updated, warns = upsert_distribution_parameter_in_all_versions(
                    ref_uuid=None,
                    ues_ref_uuid=item["ues_ref_uuid"],
                    year_number=item["year_number"],
                    base_year_number=item["base_year_number"],
                    data=item["data"],
                )
                n_added += added
                n_updated += updated
                warnings.extend(warns)
            db.session.commit()
            return n_added, n_updated, warnings
        except Exception:
            db.session.rollback()
            raise
    finally:
        wb.close()


def import_distribution_parameters_from_upload(file_storage) -> tuple[int, int, list[str]]:
    raw = file_storage.read()
    if not raw:
        raise ValueError("Пустой файл.")
    return import_distribution_parameters_from_excel(io.BytesIO(raw))
