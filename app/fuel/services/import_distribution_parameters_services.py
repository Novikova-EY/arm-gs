# -*- coding: utf-8 -*-
"""Импорт параметров распределения из Excel (шапка как в Access)."""

from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation
from typing import BinaryIO

from openpyxl import load_workbook
from sqlalchemy import func

from app.extensions import db
from app.common.services.database_version_filter import get_current_db_version_id
from app.fuel.models.fue_distribution_parameter_model import DistributionParameter
from app.refdata.models.energy_systems.union_energy_system_model import UnionEnergySystem
from app.refdata.models.years.year_model import Year

# Ключи после нормализации шапки (первая строка файла)
_EXCEL_CANONICAL_HEADERS = frozenset(
    {
        "name",
        "year",
        "filter",
        "e",
        "kplus",
        "kmin",
        "k",
        "bkl",
        "wname",
        "uname",
        "toplname",
        "dopname",
        "byear",
        "kn",
        "knps",
        "kngt",
        "knpg",
        "hnps",
        "hngt",
        "hnpg",
        "numb",
        "doptim",
        "lim",
    }
)


def _normalize_header_cell(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    low = s.lower()
    if low == "filter_text":
        return "filter"
    if low in _EXCEL_CANONICAL_HEADERS:
        return low
    return None


def _cell_str(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        t = value.strip()
        return t or None
    return str(value).strip() or None


def _parse_decimal(value):
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    s = str(value).strip().replace(",", ".")
    if not s:
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _parse_int(value):
    if value is None or value == "":
        return None
    if isinstance(value, int) and not isinstance(value, bool):
        return value
    if isinstance(value, float):
        if value != value:  # NaN
            return None
        return int(round(value))
    s = str(value).strip()
    if not s:
        return None
    try:
        return int(round(float(s.replace(",", "."))))
    except (ValueError, TypeError):
        return None


def _parse_year_num(value):
    n = _parse_int(value)
    if n is None:
        return None
    return n


def _union_energy_system_id_by_name(name: str | None, version_id: int | None) -> int | None:
    if not name or not str(name).strip():
        return None
    nm = str(name).strip()
    q = UnionEnergySystem.query.filter(func.trim(UnionEnergySystem.name) == nm)
    if version_id is not None:
        q = q.filter(UnionEnergySystem.database_version_id == version_id)
    else:
        q = q.filter(UnionEnergySystem.database_version_id.is_(None))
    row = q.first()
    return row.id if row else None


def _year_id_by_number(year_num: int | None, version_id: int | None) -> int | None:
    if year_num is None:
        return None
    q = Year.query.filter(Year.number == year_num)
    if version_id is not None:
        q = q.filter(Year.database_version_id == version_id)
    else:
        q = q.filter(Year.database_version_id.is_(None))
    row = q.first()
    return row.id if row else None


def import_distribution_parameters_from_excel(
    file_obj: BinaryIO,
    *,
    database_version_id: int | None = None,
) -> tuple[int, list[str]]:
    """
    Читает первый лист Excel. Первая строка — шапка с обязательным набором столбцов (порядок любой).

    Returns:
        (число добавленных строк, список предупреждений; пустой список при полном успехе)
    """
    if database_version_id is None:
        database_version_id = get_current_db_version_id()

    wb = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValueError("Файл пустой или нет строки заголовков.")

        col_index: dict[str, int] = {}
        for i, cell in enumerate(header_row):
            key = _normalize_header_cell(cell)
            if key:
                col_index[key] = i

        missing = sorted(_EXCEL_CANONICAL_HEADERS - set(col_index.keys()))
        if missing:
            raise ValueError(
                "В первой строке не хватает столбцов: "
                + ", ".join(missing)
                + ". Скачайте шаблон и заполните все колонки."
            )

        pending: list[DistributionParameter] = []
        errors: list[str] = []
        row_no = 1

        seen_keys: set[tuple[int, int | None, int | None]] = set()
        existing_keys = {
            (r[0], r[1], r[2])
            for r in db.session.query(
                DistributionParameter.id_year,
                DistributionParameter.id_base_year,
                DistributionParameter.id_union_energy_system,
            )
            .filter(DistributionParameter.database_version_id == database_version_id)
            .all()
        }

        for data_row in rows_iter:
            row_no += 1
            if not data_row or all(v is None or str(v).strip() == "" for v in data_row):
                continue

            def col(key: str):
                j = col_index[key]
                return data_row[j] if j < len(data_row) else None

            year_num = _parse_year_num(col("year"))
            if year_num is None:
                errors.append(f"Строка {row_no}: не задан или не распознан год (year).")
                continue

            id_year = _year_id_by_number(year_num, database_version_id)
            if id_year is None:
                errors.append(
                    f"Строка {row_no}: для года {year_num} не найдена запись Year "
                    f"в справочнике (текущая версия БД)."
                )
                continue

            name_raw = _cell_str(col("name"))
            id_ues = _union_energy_system_id_by_name(name_raw, database_version_id)
            if name_raw and id_ues is None:
                errors.append(
                    f"Строка {row_no}: ОЭС «{name_raw}» не найдена в UnionEnergySystem для текущей версии БД."
                )
                continue

            byear_num = _parse_year_num(col("byear"))
            id_base_year = None
            if byear_num is not None:
                id_base_year = _year_id_by_number(byear_num, database_version_id)
                if id_base_year is None:
                    errors.append(
                        f"Строка {row_no}: базовый год {byear_num} (byear) не найден в справочнике Year."
                    )
                    continue

            lim_v = _parse_int(col("lim"))

            key = (id_year, id_base_year, id_ues)
            byear_label = str(byear_num) if byear_num is not None else "не задан"
            oes_label = name_raw if name_raw else "не задана"
            if key in seen_keys:
                errors.append(
                    f"Строка {row_no}: в файле уже есть строка с той же комбинацией "
                    f"«Расчитываемый год» {year_num}, «Базовый год» {byear_label}, ОЭС «{oes_label}»."
                )
                continue
            if key in existing_keys:
                errors.append(
                    f"Строка {row_no}: в базе уже есть запись с комбинацией "
                    f"«Расчитываемый год» {year_num}, «Базовый год» {byear_label}, ОЭС «{oes_label}»."
                )
                continue
            seen_keys.add(key)

            dp = DistributionParameter(
                id_union_energy_system=id_ues,
                id_year=id_year,
                id_base_year=id_base_year,
                filter_text=_cell_str(col("filter")),
                e=_parse_decimal(col("e")),
                kplus=_parse_decimal(col("kplus")),
                kmin=_parse_decimal(col("kmin")),
                k=_parse_decimal(col("k")),
                bkl=_parse_decimal(col("bkl")),
                wname=_cell_str(col("wname")),
                uname=_cell_str(col("uname")),
                toplname=_cell_str(col("toplname")),
                dopname=_cell_str(col("dopname")),
                kn=_parse_decimal(col("kn")),
                knps=_parse_decimal(col("knps")),
                kngt=_parse_decimal(col("kngt")),
                knpg=_parse_decimal(col("knpg")),
                hnps=_parse_decimal(col("hnps")),
                hngt=_parse_decimal(col("hngt")),
                hnpg=_parse_decimal(col("hnpg")),
                numb=_parse_decimal(col("numb")),
                doptim=_parse_decimal(col("doptim")),
                lim=lim_v,
                database_version_id=database_version_id,
            )
            pending.append(dp)

        if errors:
            return 0, errors

        try:
            for obj in pending:
                db.session.add(obj)
            db.session.commit()
            return len(pending), []
        except Exception:
            db.session.rollback()
            raise
    finally:
        wb.close()


def import_distribution_parameters_from_upload(file_storage) -> tuple[int, list[str]]:
    raw = file_storage.read()
    if not raw:
        raise ValueError("Пустой файл.")
    return import_distribution_parameters_from_excel(io.BytesIO(raw))
