# -*- coding: utf-8 -*-
"""Экспорт таблицы накопленных денежных доходов населения в Excel."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.common.services.help_services import format_decimal_trim_for_display


def build_accum_monetary_income_excel_stream(context: dict[str, Any]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Накопленные денежные доходы"

    years: list[int] = context.get("display_years") or []
    rd = int(context.get("rounding_digits") or 1)
    yf = context.get("year_features") or {}

    def year_header_cell(y: int) -> str:
        cap = yf.get(y)
        s = str(cap).strip() if cap is not None else ""
        if s:
            return f"{y}\n{s}"
        return str(y)

    ws.cell(row=1, column=1, value="Федеральный округ, млн руб.").font = Font(bold=True)
    for col_idx, year in enumerate(years, start=2):
        cell = ws.cell(row=1, column=col_idx, value=year_header_cell(year))
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row_idx = 2
    for ami_row in context.get("accum_monetary_income_rows") or []:
        label = f"{ami_row.get('abbr', '')} — {ami_row.get('label', '')}"
        ws.cell(row=row_idx, column=1, value=label)
        cells = ami_row.get("cells") or {}
        for col_idx, year in enumerate(years, start=2):
            val = cells.get(year)
            if val is not None:
                disp = format_decimal_trim_for_display(val, digits=rd)
                ws.cell(row=row_idx, column=col_idx, value=disp)
        row_idx += 1

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
