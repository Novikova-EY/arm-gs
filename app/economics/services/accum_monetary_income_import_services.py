# -*- coding: utf-8 -*-
"""Импорт накопленных денежных доходов населения из Excel.

Поддерживаемые форматы:
- экспорт страницы (территория в столбце 1, годы с столбца 2);
- шаблон «Накопленные инвестиции_для загрузки.xlsx» (лист «таблица»):
  аббревиатура ФО в столбце 2, подпись показателя в столбце 3,
  годы с столбца 4; импортируются только строки
  «Накопленные денежные доходы населения».
"""

from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import text

from app.common.services.database_version_services import get_current_version
from app.extensions import db
from config import SCHEMA_REFDATA
from app.economics.services.accum_monetary_income_constants import IMPORT_ROW_LABEL
from app.economics.services.economics_import_versions import (
    all_database_version_ids_for_import,
)
from app.economics.services.accum_monetary_income_logging import (
    log_accum_monetary_income_excel_import,
)
from app.economics.services.accum_monetary_income_services import (
    _get_or_create_fd_row,
    _username,
    federal_districts_for_accum_monetary_income_page,
)
from app.generation.services.machine_services.machine_services import (
    is_same_decimal,
    to_decimal,
)

_UPLOAD_SHEET_NAME_HINTS = ("таблица",)
_UPLOAD_HEADER_HINTS = ("накопленн", "денежн", "доход")


def _normalize_label(text: str | None) -> str:
    return " ".join(str(text or "").split()).strip().lower()


def _parse_year_header(cell_value: Any) -> int | None:
    if cell_value is None:
        return None
    if isinstance(cell_value, (int, float)) and float(cell_value) == int(cell_value):
        y = int(cell_value)
        if 1900 <= y <= 2200:
            return y
    s = str(cell_value).strip()
    m = re.search(r"(19|20)\d{2}", s)
    if m:
        return int(m.group(0))
    try:
        y = int(float(s.replace(",", ".")))
        if 1900 <= y <= 2200:
            return y
    except (ValueError, TypeError):
        return None
    return None


def _parse_cell_value(cell_value: Any) -> Decimal | None:
    if cell_value is None or cell_value == "":
        return None
    if isinstance(cell_value, (int, float)):
        return Decimal(str(cell_value))
    s = str(cell_value).strip().replace("\u00a0", "").replace(" ", "")
    if not s or s in ("—", "-", "–"):
        return None
    s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _build_fd_lookup() -> dict[str, int]:
    lookup: dict[str, int] = {}
    for fd in federal_districts_for_accum_monetary_income_page():
        for candidate in (fd.name_abr, fd.name, fd.name_full):
            if candidate:
                lookup[_normalize_label(candidate)] = fd.id
    return lookup


def _territory_label_tokens(label: str) -> list[str]:
    n = _normalize_label(label)
    tokens = [n]
    for sep in (" — ", " - ", "–", "—"):
        if sep in label:
            left, right = label.split(sep, 1)
            for part in (left, right):
                p = _normalize_label(part)
                if p:
                    tokens.append(p)
    first = n.split()[0] if n else ""
    if first and first not in tokens:
        tokens.append(first)
    return tokens


def _is_fd_marker(label: str, fd_lookup: dict[str, int]) -> bool:
    return any(token in fd_lookup for token in _territory_label_tokens(label))


def _resolve_fd_id(label: str, fd_lookup: dict[str, int]) -> int | None:
    for token in _territory_label_tokens(label):
        fd_id = fd_lookup.get(token)
        if fd_id is not None:
            return fd_id
    return None


def _detect_year_columns(
    sheet: Worksheet, max_scan_rows: int = 8
) -> tuple[int, dict[int, int]]:
    best_row = 0
    best_map: dict[int, int] = {}
    for row_idx in range(1, max_scan_rows + 1):
        year_by_col: dict[int, int] = {}
        for col_idx in range(2, sheet.max_column + 1):
            val = sheet.cell(row=row_idx, column=col_idx).value
            year = _parse_year_header(val)
            if year is not None:
                year_by_col[col_idx] = year
        if len(year_by_col) > len(best_map):
            best_row = row_idx
            best_map = year_by_col
    if len(best_map) >= 2:
        return best_row, best_map
    return 0, {}


def _detect_territory_column(
    sheet: Worksheet,
    header_row: int,
    fd_lookup: dict[str, int],
    *,
    max_scan_rows: int = 20,
) -> int:
    hits_col1 = 0
    hits_col2 = 0
    end_row = min(header_row + max_scan_rows, sheet.max_row)
    for row_idx in range(header_row + 1, end_row + 1):
        for col_idx in (1, 2):
            cell = sheet.cell(row=row_idx, column=col_idx).value
            if cell is None or str(cell).strip() == "":
                continue
            text = str(cell).strip()
            if _is_fd_marker(text, fd_lookup):
                if col_idx == 1:
                    hits_col1 += 1
                else:
                    hits_col2 += 1
    if hits_col2 > hits_col1:
        return 2
    if hits_col1 > 0:
        return 1
    return 2 if hits_col2 > 0 else 1


def _sheet_has_structured_rows(sheet: Worksheet) -> bool:
    for row_idx in range(1, min(sheet.max_row, 200) + 1):
        cell = sheet.cell(row=row_idx, column=3).value
        if cell and IMPORT_ROW_LABEL in _normalize_label(str(cell)):
            return True
    return False


def _sheet_score_for_import(sheet: Worksheet, fd_lookup: dict[str, int]) -> int:
    header_row, year_by_col = _detect_year_columns(sheet)
    if not year_by_col:
        return 0
    territory_col = _detect_territory_column(sheet, header_row, fd_lookup)
    fd_rows = 0
    end_row = min(header_row + 20, sheet.max_row)
    for row_idx in range(header_row + 1, end_row + 1):
        cell = sheet.cell(row=row_idx, column=territory_col).value
        if cell is None or str(cell).strip() == "":
            continue
        if _is_fd_marker(str(cell).strip(), fd_lookup):
            fd_rows += 1
    score = len(year_by_col) + fd_rows * 10
    title = _normalize_label(sheet.title)
    if any(h in title for h in _UPLOAD_SHEET_NAME_HINTS):
        score += 50
    if _sheet_has_structured_rows(sheet):
        score += 100
    return score


def _pick_import_sheet(wb, fd_lookup: dict[str, int]) -> Worksheet:
    best: Worksheet | None = None
    best_score = 0
    for name in wb.sheetnames:
        sheet = wb[name]
        score = _sheet_score_for_import(sheet, fd_lookup)
        if score > best_score:
            best_score = score
            best = sheet
    if best is not None and best_score > 0:
        return best
    if wb.active is not None:
        return wb.active
    raise ValueError("Файл Excel не содержит листов.")


def _parse_flat_sheet(
    sheet: Worksheet,
    fd_lookup: dict[str, int],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    header_row, year_by_col = _detect_year_columns(sheet)
    if not year_by_col:
        raise ValueError("Не найдена строка с годами в шапке таблицы.")

    territory_col = _detect_territory_column(sheet, header_row, fd_lookup)
    first_year_col = min(year_by_col.keys())

    stats: dict[str, Any] = {
        "rows_processed": 0,
        "cells_skipped": 0,
        "territories": 0,
        "warnings": [],
        "sheet_name": sheet.title,
        "header_row": header_row,
        "territory_column": territory_col,
        "year_columns": len(year_by_col),
    }
    parsed_rows: list[dict[str, Any]] = []
    seen_fd: set[int] = set()

    for row_idx in range(header_row + 1, sheet.max_row + 1):
        territory_cell = sheet.cell(row=row_idx, column=territory_col).value
        if territory_cell is None or str(territory_cell).strip() == "":
            continue
        label = str(territory_cell).strip()
        if any(h in _normalize_label(label) for h in _UPLOAD_HEADER_HINTS):
            continue

        fd_id = _resolve_fd_id(label, fd_lookup)
        if fd_id is None:
            stats["warnings"].append(f"Строка {row_idx}: не найден ФО «{label}».")
            continue

        cells: list[tuple[int, Decimal]] = []
        for col_idx, year_n in year_by_col.items():
            if col_idx < first_year_col:
                continue
            val = _parse_cell_value(sheet.cell(row=row_idx, column=col_idx).value)
            if val is not None:
                cells.append((year_n, val))
            else:
                stats["cells_skipped"] += 1

        if not cells:
            continue

        stats["rows_processed"] += 1
        if fd_id not in seen_fd:
            seen_fd.add(fd_id)
            stats["territories"] += 1
        parsed_rows.append({"fd_id": fd_id, "cells": cells, "label": label})

    if not parsed_rows:
        raise ValueError(
            "Не найдено строк с данными по федеральным округам. "
            f"Проверьте лист «{sheet.title}»."
        )

    return parsed_rows, stats


def _parse_structured_upload_sheet(
    sheet: Worksheet,
    fd_lookup: dict[str, int],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    header_row, year_by_col = _detect_year_columns(sheet)
    if not year_by_col:
        raise ValueError("Не найдена строка с годами в шапке таблицы.")

    first_year_col = min(year_by_col.keys())
    stats: dict[str, Any] = {
        "rows_processed": 0,
        "cells_skipped": 0,
        "territories": 0,
        "warnings": [],
        "sheet_name": sheet.title,
        "header_row": header_row,
        "territory_column": 2,
        "year_columns": len(year_by_col),
    }
    parsed_rows: list[dict[str, Any]] = []
    seen_fd: set[int] = set()
    current_fd_id: int | None = None

    for row_idx in range(header_row + 1, sheet.max_row + 1):
        fd_cell = sheet.cell(row=row_idx, column=2).value
        if fd_cell is not None and str(fd_cell).strip():
            fd_id = _resolve_fd_id(str(fd_cell).strip(), fd_lookup)
            if fd_id is not None:
                current_fd_id = fd_id

        metric_cell = sheet.cell(row=row_idx, column=3).value
        if metric_cell is None or str(metric_cell).strip() == "":
            continue
        metric_label = _normalize_label(str(metric_cell))
        if IMPORT_ROW_LABEL not in metric_label:
            continue

        if current_fd_id is None:
            stats["warnings"].append(
                f"Строка {row_idx}: не определён федеральный округ для показателя."
            )
            continue

        cells: list[tuple[int, Decimal]] = []
        for col_idx, year_n in year_by_col.items():
            if col_idx < first_year_col:
                continue
            val = _parse_cell_value(sheet.cell(row=row_idx, column=col_idx).value)
            if val is not None:
                cells.append((year_n, val))
            else:
                stats["cells_skipped"] += 1

        if not cells:
            continue

        stats["rows_processed"] += 1
        if current_fd_id not in seen_fd:
            seen_fd.add(current_fd_id)
            stats["territories"] += 1
        parsed_rows.append(
            {
                "fd_id": current_fd_id,
                "cells": cells,
                "label": str(metric_cell).strip(),
            }
        )

    if not parsed_rows:
        raise ValueError(
            "Не найдено строк «Накопленные денежные доходы населения». "
            f"Проверьте лист «{sheet.title}» (столбец 3, годы с столбца {first_year_col})."
        )

    return parsed_rows, stats


def _parse_accum_monetary_income_sheet(
    sheet: Worksheet,
    fd_lookup: dict[str, int],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if _sheet_has_structured_rows(sheet):
        return _parse_structured_upload_sheet(sheet, fd_lookup)
    return _parse_flat_sheet(sheet, fd_lookup)


def _ensure_years_in_versions(
    version_ids: list[int], year_numbers: set[int]
) -> int:
    if not year_numbers or not version_ids:
        return 0
    ensure_sql = text(
        f"""
        INSERT INTO {SCHEMA_REFDATA}.gs_sys_years (number, database_version_id)
        SELECT :year_number, :version_id
        WHERE NOT EXISTS (
            SELECT 1
            FROM {SCHEMA_REFDATA}.gs_sys_years
            WHERE number = :year_number
              AND database_version_id = :version_id
        )
        """
    )
    created = 0
    for version_id in version_ids:
        for year_n in sorted(year_numbers):
            result = db.session.execute(
                ensure_sql,
                {"year_number": year_n, "version_id": version_id},
            )
            created += result.rowcount or 0
    if created:
        db.session.flush()
    return created


def import_accum_monetary_income_from_xlsx_bytes(raw: bytes) -> dict[str, Any]:
    version_ids = all_database_version_ids_for_import()
    if not version_ids:
        raise ValueError(
            "В системе нет зарегистрированных версий БД — импорт невозможен."
        )

    user = _username()
    fd_lookup = _build_fd_lookup()
    current_version_id = get_current_version()

    wb = load_workbook(filename=BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = _pick_import_sheet(wb, fd_lookup)
        parsed_rows, parse_stats = _parse_accum_monetary_income_sheet(sheet, fd_lookup)
    finally:
        wb.close()

    all_years: set[int] = set()
    for entry in parsed_rows:
        for year_n, _ in entry["cells"]:
            all_years.add(year_n)

    years_created = _ensure_years_in_versions(version_ids, all_years)

    stats = {
        "rows_processed": parse_stats["rows_processed"],
        "cells_written": 0,
        "cells_skipped": parse_stats["cells_skipped"],
        "territories": parse_stats["territories"],
        "database_versions": len(version_ids),
        "warnings": list(parse_stats.get("warnings") or []),
        "sheet_name": parse_stats.get("sheet_name"),
        "years_ensured": years_created,
    }

    for version_id in version_ids:
        for entry in parsed_rows:
            fd_id = entry["fd_id"]
            for year_n, val in entry["cells"]:
                row = _get_or_create_fd_row(version_id, fd_id, year_n, user)
                old_val = row.accum_monetary_income_mln_rub
                if is_same_decimal(to_decimal(old_val), to_decimal(val)):
                    stats["cells_skipped"] += 1
                    continue
                row.accum_monetary_income_mln_rub = val
                row.modified_by = user
                stats["cells_written"] += 1

    db.session.flush()
    if stats["cells_written"]:
        log_accum_monetary_income_excel_import(
            user, stats, database_version_id=current_version_id
        )
    return stats
