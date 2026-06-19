# -*- coding: utf-8 -*-
"""Импорт потребления по ВЭД из Excel."""

from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.common.services.database_version_services import get_current_version
from app.extensions import db
from app.economics.services.economics_import_versions import (
    all_database_version_ids_for_import,
)
from app.economics.services.ved_consumption_constants import (
    FD_TOTAL_ROW_LABEL,
    RUSSIA_TERRITORY_LABELS,
    TOTAL_VED_NAME,
)
from app.economics.services.ved_consumption_logging import (
    log_ved_consumption_excel_import,
)
from app.economics.services.ved_consumption_services import (
    _federal_districts_for_ved_page,
    _get_or_create_fd_row,
    _get_or_create_rf_row,
    _normalize_label,
    _username,
    _ved_types_for_version,
)
from app.generation.services.machine_services.machine_services import (
    is_same_decimal,
    to_decimal,
)
def _parse_year_header(cell_value: Any) -> int | None:
    if cell_value is None:
        return None
    if isinstance(cell_value, (int, float)) and float(cell_value) == int(cell_value):
        y = int(cell_value)
        if 1900 <= y <= 2200:
            return y
    s = str(cell_value).strip()
    m = re.search(r"(19|20)\d{2}", s)
    if m:
        return int(m.group(0))
    try:
        y = int(float(s.replace(",", ".")))
        if 1900 <= y <= 2200:
            return y
    except (ValueError, TypeError):
        return None
    return None


def _parse_cell_value(cell_value: Any) -> Decimal | None:
    if cell_value is None or cell_value == "":
        return None
    if isinstance(cell_value, (int, float)):
        return Decimal(str(cell_value))
    s = str(cell_value).strip().replace("\u00a0", "").replace(" ", "")
    if not s or s in ("—", "-", "–"):
        return None
    s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _build_fd_lookup() -> dict[str, int]:
    lookup: dict[str, int] = {}
    for fd in _federal_districts_for_ved_page():
        for candidate in (fd.name_abr, fd.name, fd.name_full):
            if candidate:
                lookup[_normalize_label(candidate)] = fd.id
    return lookup


def _build_ved_lookup(version_id: int | None) -> dict[str, int]:
    lookup: dict[str, int] = {}
    total_ved_id: int | None = None
    for ved in _ved_types_for_version(version_id):
        if ved.name:
            key = _normalize_label(ved.name)
            lookup[key] = ved.id
            if key == _normalize_label(TOTAL_VED_NAME):
                total_ved_id = ved.id
    if total_ved_id is not None:
        lookup[_normalize_label(FD_TOTAL_ROW_LABEL)] = total_ved_id
    return lookup


def _territory_label_tokens(label: str) -> list[str]:
    """Варианты подписи территории (экспорт «РФ — …», файл «РОССИЯ» в отдельном столбце)."""
    n = _normalize_label(label)
    tokens = [n]
    for sep in (" — ", " - ", "–", "—"):
        if sep in label:
            left, right = label.split(sep, 1)
            for part in (left, right):
                p = _normalize_label(part)
                if p:
                    tokens.append(p)
    first = n.split()[0] if n else ""
    if first and first not in tokens:
        tokens.append(first)
    return tokens


def _ved_label_variants(label: str) -> list[str]:
    n = _normalize_label(label)
    variants = [n]
    stripped = re.sub(r",\s*в том числе:?\s*$", "", n, flags=re.IGNORECASE).strip(" :")
    if stripped and stripped not in variants:
        variants.append(stripped)
    if ";" in n:
        head = n.split(";", 1)[0].strip(" :")
        if head and head not in variants:
            variants.append(head)
    return variants


def _resolve_ved_id(label: str, ved_lookup: dict[str, int]) -> int | None:
    for variant in _ved_label_variants(label):
        ved_id = ved_lookup.get(variant)
        if ved_id is not None:
            return ved_id
    n = _normalize_label(label)
    best_key: str | None = None
    for key in ved_lookup:
        if n.startswith(key) and (best_key is None or len(key) > len(best_key)):
            best_key = key
    if best_key is not None:
        return ved_lookup[best_key]
    return None


def _detect_import_columns(
    sheet,
    header_row: int,
    fd_lookup: dict[str, int],
    *,
    max_scan_rows: int = 40,
) -> tuple[int, int]:
    """(столбец территории, столбец ВЭД): экспорт — (1, 1), шаблон «Структура…» — (2, 3)."""
    territory_hits_col1 = 0
    territory_hits_col2 = 0
    ved_rows_col3 = 0
    end_row = min(header_row + max_scan_rows, sheet.max_row)
    for row_idx in range(header_row + 1, end_row + 1):
        for col_idx, cell in ((1, sheet.cell(row=row_idx, column=1).value), (2, sheet.cell(row=row_idx, column=2).value)):
            if cell is None or str(cell).strip() == "":
                continue
            text = str(cell).strip()
            if _is_russia_marker(text) or _is_fd_marker(text, fd_lookup):
                if col_idx == 1:
                    territory_hits_col1 += 1
                else:
                    territory_hits_col2 += 1
        c3 = sheet.cell(row=row_idx, column=3).value
        if c3 is not None and str(c3).strip():
            ved_rows_col3 += 1
    if territory_hits_col2 >= 1 and ved_rows_col3 >= 3:
        return 2, 3
    return 1, 1


def _detect_year_columns(sheet, max_scan_rows: int = 5) -> tuple[int, dict[int, int]]:
    """Возвращает (номер строки шапки, {col_index: year})."""
    for row_idx in range(1, max_scan_rows + 1):
        year_by_col: dict[int, int] = {}
        for col_idx in range(2, sheet.max_column + 1):
            val = sheet.cell(row=row_idx, column=col_idx).value
            year = _parse_year_header(val)
            if year is not None:
                year_by_col[col_idx] = year
        if len(year_by_col) >= 2:
            return row_idx, year_by_col
    return 0, {}


def _is_russia_marker(label: str) -> bool:
    for token in _territory_label_tokens(label):
        if token in RUSSIA_TERRITORY_LABELS or token.startswith("российская федерация"):
            return True
    return False


def _is_fd_marker(label: str, fd_lookup: dict[str, int]) -> bool:
    return any(token in fd_lookup for token in _territory_label_tokens(label))


def _resolve_fd_id(label: str, fd_lookup: dict[str, int]) -> int | None:
    for token in _territory_label_tokens(label):
        fd_id = fd_lookup.get(token)
        if fd_id is not None:
            return fd_id
    return None


def import_ved_consumption_from_xlsx_bytes(raw: bytes) -> dict[str, Any]:
    version_ids = all_database_version_ids_for_import()
    if not version_ids:
        raise ValueError(
            "В системе нет зарегистрированных версий БД — импорт невозможен."
        )

    user = _username()
    fd_lookup = _build_fd_lookup()
    current_version_id = get_current_version()
    ved_lookup_current = _build_ved_lookup(current_version_id)

    wb = load_workbook(filename=BytesIO(raw), read_only=True, data_only=True)
    sheet = wb.active
    if sheet is None:
        raise ValueError("Файл Excel не содержит листов.")

    header_row, year_by_col = _detect_year_columns(sheet)
    if not year_by_col:
        raise ValueError("Не найдена строка с годами в шапке таблицы.")

    territory_col, ved_col = _detect_import_columns(sheet, header_row, fd_lookup)
    split_territory_ved = territory_col != ved_col

    stats = {
        "rows_processed": 0,
        "cells_written": 0,
        "cells_skipped": 0,
        "territories": 0,
        "database_versions": len(version_ids),
        "warnings": [],
    }

    parsed_rows: list[dict[str, Any]] = []
    current_kind: str | None = None
    current_fd_id: int | None = None

    for row_idx in range(header_row + 1, sheet.max_row + 1):
        territory_cell = sheet.cell(row=row_idx, column=territory_col).value
        territory_label = (
            str(territory_cell).strip()
            if territory_cell is not None and str(territory_cell).strip()
            else None
        )

        if territory_label:
            if _is_russia_marker(territory_label):
                current_kind = "rf"
                current_fd_id = None
                stats["territories"] += 1
                if not split_territory_ved:
                    continue
            elif _is_fd_marker(territory_label, fd_lookup):
                fd_id = _resolve_fd_id(territory_label, fd_lookup)
                if fd_id is None:
                    continue
                current_kind = "fd"
                current_fd_id = fd_id
                stats["territories"] += 1
                if not split_territory_ved:
                    continue

        ved_cell = sheet.cell(row=row_idx, column=ved_col).value
        if ved_cell is None or str(ved_cell).strip() == "":
            continue
        label = str(ved_cell).strip()

        if current_kind is None:
            stats["warnings"].append(f"Строка {row_idx}: нет активной территории для «{label}».")
            continue

        if not split_territory_ved:
            if _is_russia_marker(label):
                current_kind = "rf"
                current_fd_id = None
                stats["territories"] += 1
                continue
            if _is_fd_marker(label, fd_lookup):
                fd_id = _resolve_fd_id(label, fd_lookup)
                if fd_id is None:
                    continue
                current_kind = "fd"
                current_fd_id = fd_id
                stats["territories"] += 1
                continue

        cells: list[tuple[int, Decimal]] = []
        for col_idx, year_n in year_by_col.items():
            val = _parse_cell_value(sheet.cell(row=row_idx, column=col_idx).value)
            if val is not None:
                cells.append((year_n, val))
            else:
                stats["cells_skipped"] += 1

        if not cells:
            continue

        if _resolve_ved_id(label, ved_lookup_current) is None:
            stats["warnings"].append(f"Строка {row_idx}: не найден ВЭД «{label}».")
            stats["cells_skipped"] += len(cells) * len(version_ids)
            continue

        stats["rows_processed"] += 1
        parsed_rows.append(
            {
                "kind": current_kind,
                "fd_id": current_fd_id,
                "ved_label": label,
                "cells": cells,
            }
        )

    for version_id in version_ids:
        ved_lookup = _build_ved_lookup(version_id)
        for entry in parsed_rows:
            ved_id = _resolve_ved_id(entry["ved_label"], ved_lookup)
            if ved_id is None:
                stats["cells_skipped"] += len(entry["cells"])
                continue
            kind = entry["kind"]
            fd_id = entry["fd_id"]
            for year_n, val in entry["cells"]:
                if kind == "rf":
                    row = _get_or_create_rf_row(version_id, ved_id, year_n, user)
                else:
                    assert fd_id is not None
                    row = _get_or_create_fd_row(version_id, fd_id, ved_id, year_n, user)
                old_val = row.energy_consumption_mln_kvt_ch
                if is_same_decimal(to_decimal(old_val), to_decimal(val)):
                    stats["cells_skipped"] += 1
                    continue
                row.energy_consumption_mln_kvt_ch = val
                row.modified_by = user
                stats["cells_written"] += 1

    db.session.flush()
    if stats["cells_written"]:
        log_ved_consumption_excel_import(
            user, stats, database_version_id=current_version_id
        )
        from app.common.services.economics_fd_data_cache import (
            invalidate_economics_fd_data_cache,
        )

        for vid in version_ids:
            invalidate_economics_fd_data_cache(vid, "ved_consumption")
    return stats
