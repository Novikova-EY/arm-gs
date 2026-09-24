from datetime import datetime
from io import BytesIO
import re

from flask import Blueprint, current_app, jsonify, render_template, request, send_file
from flask_login import login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.auth.routes.decorators import roles_required

start_bp = Blueprint('start', __name__)

@start_bp.route('/')
def index():
    return render_template('home.html')


@start_bp.route('/output_documents_transfer')
@login_required
def output_documents_transfer():
    """Страница ссылок JSON-API для выходных документов (доступна всем авторизованным)."""
    from flask import g

    from app.api.services.database_versions_api_services import list_database_versions
    from app.api.services.output_documents_catalog import build_output_document_groups

    versions = list_database_versions().get("rows") or []
    selected_version_id = None
    current = getattr(g, "current_db_version", None)
    version_ids = {int(row["database_version"]) for row in versions}
    try:
        current_id = int(current) if current is not None else None
    except (TypeError, ValueError):
        current_id = None
    if current_id in version_ids:
        selected_version_id = current_id
    elif versions:
        selected_version_id = int(versions[0]["database_version"])

    return render_template(
        "output_documents_transfer.html",
        versions=versions,
        selected_version_id=selected_version_id,
        json_groups=build_output_document_groups(),
    )


@start_bp.route(
    "/output_documents_transfer/export_excel/<dataset>",
    methods=["GET"],
)
@login_required
@roles_required(["admin"])
def export_output_document_excel(dataset: str):
    """Выгрузка одного JSON-набора в Excel (только роль admin)."""
    from app.api.services.database_versions_api_services import (
        UnknownDatabaseVersionError,
    )
    from app.api.services.exchange_api_services import UnknownDatasetError
    from app.api.services.output_documents_excel_export_services import (
        build_dataset_excel_bytes,
        build_export_filename,
        load_output_document_payload_for_excel,
    )

    version_raw = request.args.get("database_version")
    database_version = None
    if version_raw is not None and str(version_raw).strip() != "":
        try:
            database_version = int(version_raw)
        except (TypeError, ValueError):
            return jsonify({"error": "Некорректный database_version."}), 400

    try:
        payload = load_output_document_payload_for_excel(
            dataset,
            database_version=database_version,
        )
    except ValueError:
        return jsonify(
            {"error": "Укажите database_version для выгрузки набора."}
        ), 400
    except UnknownDatabaseVersionError:
        return jsonify(
            {
                "error": "unknown_database_version",
                "database_version": database_version,
            }
        ), 404
    except UnknownDatasetError:
        return jsonify({"error": "unknown_dataset", "dataset": dataset}), 404
    except Exception as exc:
        current_app.logger.error(
            "Ошибка загрузки набора %s для Excel: %s", dataset, exc
        )
        return jsonify({"error": "Не удалось загрузить JSON-набор."}), 500

    try:
        output = build_dataset_excel_bytes(payload)
        filename = build_export_filename(
            str(payload.get("dataset") or dataset),
            database_version=payload.get("database_version", database_version),
        )
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype=(
                "application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet"
            ),
            max_age=0,
        )
    except Exception as exc:
        current_app.logger.error(
            "Ошибка выгрузки Excel для набора %s: %s", dataset, exc
        )
        return jsonify({"error": "Не удалось сформировать Excel-файл."}), 500


def _safe_excel_sheet_name(sheet_name: str) -> str:
    cleaned = re.sub(r'[:\\/?*\[\]]', '_', (sheet_name or '').strip())
    cleaned = cleaned[:31].strip()
    return cleaned or 'Sheet1'


def _safe_download_name(filename: str) -> str:
    cleaned = re.sub(r'[^A-Za-zА-Яа-я0-9._() -]+', '_', (filename or '').strip())
    cleaned = cleaned.strip(' ._')
    return cleaned or f'table_export_{datetime.now().strftime("%Y%m%d_%H%M%S")}'


@start_bp.route('/export_table_excel', methods=['POST'])
@login_required
def export_table_excel():
    """Универсальный экспорт произвольной HTML-таблицы в XLSX."""

    payload = request.get_json(silent=True) or {}
    headers = payload.get('headers') or []
    rows = payload.get('rows') or []
    sheet_name = _safe_excel_sheet_name(payload.get('sheet_name') or 'Данные')
    filename = _safe_download_name(payload.get('filename') or 'table_export')

    if not isinstance(headers, list) or not isinstance(rows, list):
        return jsonify({'error': 'Некорректный формат данных для экспорта.'}), 400

    try:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name

        if headers:
            worksheet.append([str(value) if value is not None else '' for value in headers])
            header_fill = PatternFill(fill_type='solid', fgColor='D1E7DD')
            header_font = Font(bold=True)
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        for row in rows:
            if not isinstance(row, list):
                continue
            worksheet.append([str(value) if value is not None else '' for value in row])

        if headers:
            worksheet.freeze_panes = 'A2'
            worksheet.auto_filter.ref = worksheet.dimensions

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                value = '' if cell.value is None else str(cell.value)
                if len(value) > max_length:
                    max_length = len(value)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 60)

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        if not filename.lower().endswith('.xlsx'):
            filename = f'{filename}.xlsx'

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            max_age=0,
        )
    except Exception as exc:
        current_app.logger.error('Ошибка универсального экспорта таблицы: %s', exc)
        return jsonify({'error': 'Не удалось сформировать Excel-файл.'}), 500
