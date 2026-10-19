# -*- coding: utf-8 -*-
"""Выгрузка JSON-набора выходных документов в Excel (минимальное форматирование)."""
from __future__ import annotations

import json
import re
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.api.services.output_documents_catalog import DATASET_DATABASE_VERSIONS


DATASET_GENERATION_STATIONS = "generation_stations"
GENERATION_STATIONS_CAPACITY_COLUMN = "Установленная мощность, МВт"
GENERATION_STATIONS_CAPACITY_NUMBER_FORMAT = "#,##0.000000"
GENERATION_STATIONS_GENERATION_COLUMN = "Производство электроэнергии в году, МВт*ч"
# В БД выработка хранится числом с плавающей запятой: без формата Excel показывает
# до 15 значащих цифр (16 → "16", 0.316 → "0.316"), без округления значений.
GENERATION_STATIONS_GENERATION_NUMBER_FORMAT = "General"


def _safe_excel_sheet_name(sheet_name: str) -> str:
    cleaned = re.sub(r"[:\\/?*\[\]]", "_", (sheet_name or "").strip())
    cleaned = cleaned[:31].strip()
    return cleaned or "Sheet1"


def _safe_download_name(filename: str) -> str:
    cleaned = re.sub(r"[^A-Za-zА-Яа-я0-9._() -]+", "_", (filename or "").strip())
    cleaned = cleaned.strip(" ._")
    return cleaned or f"dataset_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}"


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _collect_headers(rows: list[Any]) -> list[str]:
    headers: list[str] = []
    seen: set[str] = set()
    for row in rows:
        if not isinstance(row, dict):
            continue
        for key in row.keys():
            name = str(key)
            if name in seen:
                continue
            seen.add(name)
            headers.append(name)
    return headers


def _set_column_number_format(
    worksheet: Any,
    headers: list[str],
    column_name: str,
    number_format: str,
) -> None:
    try:
        column_number = headers.index(column_name) + 1
    except ValueError:
        return
    for row_number in range(2, worksheet.max_row + 1):
        worksheet.cell(row=row_number, column=column_number).number_format = number_format


def build_dataset_excel_bytes(payload: dict[str, Any]) -> BytesIO:
    """Собирает XLSX из оболочки набора (поле ``rows``) с минимальным оформлением."""
    rows = payload.get("rows") if isinstance(payload, dict) else None
    if not isinstance(rows, list):
        rows = []

    headers = _collect_headers(rows)
    sheet_name = _safe_excel_sheet_name(
        str((payload or {}).get("dataset") or "data")
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name

    header_font = Font(bold=True)
    if headers:
        worksheet.append(headers)
        for cell in worksheet[1]:
            cell.font = header_font

    for row in rows:
        if not isinstance(row, dict):
            continue
        worksheet.append([_cell_value(row.get(header)) for header in headers])

    if (payload or {}).get("dataset") == DATASET_GENERATION_STATIONS:
        _set_column_number_format(
            worksheet,
            headers,
            GENERATION_STATIONS_CAPACITY_COLUMN,
            GENERATION_STATIONS_CAPACITY_NUMBER_FORMAT,
        )
        _set_column_number_format(
            worksheet,
            headers,
            GENERATION_STATIONS_GENERATION_COLUMN,
            GENERATION_STATIONS_GENERATION_NUMBER_FORMAT,
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_export_filename(
    dataset: str,
    *,
    database_version: int | None = None,
) -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    parts = [str(dataset or "dataset").strip() or "dataset"]
    if database_version is not None:
        parts.append(f"v{int(database_version)}")
    parts.append(stamp)
    return _safe_download_name("_".join(parts)) + ".xlsx"


def load_output_document_payload_for_excel(
    dataset: str,
    *,
    database_version: int | None = None,
) -> dict[str, Any]:
    """Загружает JSON-набор для выгрузки в Excel.

    Для ``generation_stations`` выработка берётся из БД без округления
    (в JSON-наборе она округляется до 3 знаков — контракт внешнего API не меняем).
    """
    from app.api.services.database_versions_api_services import (
        list_database_versions,
    )
    from app.api.services.exchange_api_services import (
        UnknownDatasetError,
        load_exchange_dataset,
    )

    name = str(dataset or "").strip()
    if not name:
        raise UnknownDatasetError(dataset)

    if name == DATASET_DATABASE_VERSIONS:
        payload = list_database_versions()
        payload = dict(payload)
        payload.setdefault("dataset", DATASET_DATABASE_VERSIONS)
        return payload

    if database_version is None:
        raise ValueError("database_version_required")

    if name == DATASET_GENERATION_STATIONS:
        from app.api.services.generation_objects_exchange_services import (
            load_generation_stations_dataset,
        )

        return load_generation_stations_dataset(
            version_id=int(database_version),
            preserve_generation_precision=True,
        )

    return load_exchange_dataset(name, version_id=int(database_version))
