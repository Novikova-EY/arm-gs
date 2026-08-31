# -*- coding: utf-8 -*-
"""Данные и выгрузка таблицы «Максимумы потребления мощности ОЗП»."""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.common.services.database_version_services import get_current_version
from app.common.services.excel_numeric_cell import (
    excel_number_format,
    excel_numeric_cell_value,
)
from app.common.services.get_services.years.years_get_services import (
    get_filter_end_year,
    get_filter_start_year,
    get_ges_tep_current_price_year_number,
    get_year_list_full,
)
from app.common.services.help_services import (
    format_decimal_for_display,
    parse_decimal_from_display,
)
from app.extensions import db
from app.power_demand.models.ozp_max_power_parameter_model import OzpMaxPowerParameter

GROWTH_ROUNDING_DIGITS = 2
DEFAULT_MW_ROUNDING_DIGITS = -1
PAGE_TITLE = "Максимумы потребления мощности ОЗП"
ROW_MAX_LABEL = "Максимум потребления мощности ОЗП, МВт"
ROW_GROWTH_LABEL = "Прирост к прошлому ОЗП, %"
GROWTH_FORMULA_TEXT = (
    "Прирост к прошлому ОЗП, % = (максимум потребления мощности ОЗП текущего периода / "
    "максимум потребления мощности ОЗП предыдущего периода) × 100 − 100"
)

_HEADER_FILL = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
_HEADER_FONT = Font(bold=True, size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
_CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN = Side(style="thin", color="000000")
_CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BASE_FONT = Font(size=11)


def filter_year_list() -> list[int]:
    nums = sorted({int(y.number) for y in get_year_list_full() if y.number is not None})
    if nums:
        return nums
    return list(range(get_filter_start_year(), get_filter_end_year() + 1))


def period_base_year_n() -> int:
    n = get_ges_tep_current_price_year_number()
    if n is not None:
        return int(n)
    return int(get_filter_end_year())


def parse_rounding_digits(raw: Any) -> int:
    if raw is None or str(raw).strip() == "":
        return DEFAULT_MW_ROUNDING_DIGITS
    try:
        value = int(raw)
    except (TypeError, ValueError):
        return DEFAULT_MW_ROUNDING_DIGITS
    if value == -1:
        return -1
    if value in (0, 1, 2, 3):
        return value
    return DEFAULT_MW_ROUNDING_DIGITS


def parse_year_range(
    start_raw: Any,
    end_raw: Any,
    *,
    years: list[int] | None = None,
) -> tuple[int, int]:
    """Диапазон календарных лет шапки. По умолчанию N−4…N (как на образце 2020–2024)."""
    bounds = years if years is not None else filter_year_list()
    has_start = start_raw not in (None, "")
    has_end = end_raw not in (None, "")
    default_sy = default_ey = None
    if not has_start or not has_end:
        n = period_base_year_n()
        default_sy = n - 4
        default_ey = n
        if bounds:
            lo, hi = bounds[0], bounds[-1]
            default_sy = max(lo, min(default_sy, hi))
            default_ey = max(lo, min(default_ey, hi))
            if default_sy > default_ey:
                default_sy, default_ey = default_ey, default_sy
    try:
        sy = int(start_raw) if has_start else default_sy
    except (TypeError, ValueError):
        sy = default_sy if default_sy is not None else (bounds[0] if bounds else 2020)
    try:
        ey = int(end_raw) if has_end else default_ey
    except (TypeError, ValueError):
        ey = default_ey if default_ey is not None else (bounds[-1] if bounds else sy)
    if sy > ey:
        sy, ey = ey, sy
    if bounds:
        lo, hi = bounds[0], bounds[-1]
        sy = max(lo, min(sy, hi))
        ey = max(lo, min(ey, hi))
    if sy > ey:
        sy, ey = ey, sy
    return sy, ey


def format_ozp_period_label(ozp_start: int, ozp_end: int) -> str:
    return f"{ozp_start}–{ozp_end} гг."


def build_ozp_columns(start_year: int, end_year: int) -> list[dict[str, Any]]:
    """Колонки таблицы: у каждого года один ОЗП, у года конца — два (как на образце)."""
    if start_year > end_year:
        start_year, end_year = end_year, start_year
    columns: list[dict[str, Any]] = []
    for year in range(start_year, end_year):
        ozp_start, ozp_end = year - 1, year
        columns.append(
            {
                "calendar_year": year,
                "ozp_start": ozp_start,
                "ozp_end": ozp_end,
                "label": format_ozp_period_label(ozp_start, ozp_end),
            }
        )
    columns.append(
        {
            "calendar_year": end_year,
            "ozp_start": end_year - 1,
            "ozp_end": end_year,
            "label": format_ozp_period_label(end_year - 1, end_year),
        }
    )
    columns.append(
        {
            "calendar_year": end_year,
            "ozp_start": end_year,
            "ozp_end": end_year + 1,
            "label": format_ozp_period_label(end_year, end_year + 1),
        }
    )
    return columns


def build_year_header_groups(columns: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: list[dict[str, Any]] = []
    for col in columns:
        year = int(col["calendar_year"])
        if groups and groups[-1]["year"] == year:
            groups[-1]["colspan"] += 1
        else:
            groups.append({"year": year, "colspan": 1})
    return groups


def calc_growth_pct(current: Decimal | None, previous: Decimal | None) -> Decimal | None:
    """Текущий / прошлый × 100 − 100; как «Годовой темп прироста, %» на сводке потребления."""
    if current is None or previous is None or previous == 0:
        return None
    return (current / previous) * Decimal(100) - Decimal(100)


def compute_growth_series(
    max_values: list[Decimal | None],
    previous_max: Decimal | None,
) -> list[Decimal | None]:
    prev = previous_max
    out: list[Decimal | None] = []
    for current in max_values:
        out.append(calc_growth_pct(current, prev))
        prev = current
    return out


def format_mw_display(value: Any, rounding_digits: int) -> str:
    if value is None:
        return ""
    return format_decimal_for_display(value, digits=rounding_digits)


def format_growth_display(value: Any) -> str:
    if value is None:
        return "—"
    return format_decimal_for_display(value, digits=GROWTH_ROUNDING_DIGITS)


def _parse_optional_decimal(raw: Any) -> Decimal | None:
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "—", "-", "–"):
        return None
    try:
        return parse_decimal_from_display(raw)
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError("Некорректное число") from None


def load_ozp_rows_by_end_year(
    *, version_id: int | None = None
) -> dict[int, OzpMaxPowerParameter]:
    vid = get_current_version() if version_id is None else version_id
    q = OzpMaxPowerParameter.query
    if vid is None:
        q = q.filter(OzpMaxPowerParameter.database_version_id.is_(None))
    else:
        q = q.filter(OzpMaxPowerParameter.database_version_id == int(vid))
    by_year: dict[int, OzpMaxPowerParameter] = {}
    for row in q.all():
        by_year[int(row.ozp_end_year)] = row
    return by_year


def build_page_context(
    *,
    start_year: int,
    end_year: int,
    rounding_digits: int,
    can_edit: bool,
    filter_years: list[int] | None = None,
) -> dict[str, Any]:
    columns = build_ozp_columns(start_year, end_year)
    year_groups = build_year_header_groups(columns)
    stored = load_ozp_rows_by_end_year()
    max_raw: list[Decimal | None] = []
    max_values: list[str] = []
    max_full: list[str] = []
    for col in columns:
        row = stored.get(int(col["ozp_end"]))
        mw = None if row is None else row.max_power_mw
        max_raw.append(mw)
        max_values.append(format_mw_display(mw, rounding_digits))
        max_full.append(format_mw_display(mw, 0) if mw is not None else "")

    prev_end = int(columns[0]["ozp_end"]) - 1 if columns else None
    prev_row = stored.get(prev_end) if prev_end is not None else None
    prev_max = None if prev_row is None else prev_row.max_power_mw
    growth_raw = compute_growth_series(max_raw, prev_max)
    growth_values = [format_growth_display(v) for v in growth_raw]

    return {
        "page_title": PAGE_TITLE,
        "row_max_label": ROW_MAX_LABEL,
        "row_growth_label": ROW_GROWTH_LABEL,
        "growth_formula_text": GROWTH_FORMULA_TEXT,
        "start_year": start_year,
        "end_year": end_year,
        "rounding_digits": rounding_digits,
        "filter_year_list": filter_years if filter_years is not None else filter_year_list(),
        "columns": columns,
        "year_groups": year_groups,
        "max_values": max_values,
        "growth_values": growth_values,
        "max_full": max_full,
        "prev_max_display": format_mw_display(prev_max, 0) if prev_max is not None else "",
        "can_edit_ozp": can_edit,
        "growth_rounding_digits": GROWTH_ROUNDING_DIGITS,
    }


def upsert_ozp_row(
    ozp_end_year: int,
    *,
    max_power_mw: Decimal | None | object = ...,
    growth_pct: Decimal | None | object = ...,
) -> OzpMaxPowerParameter:
    vid = get_current_version()
    q = OzpMaxPowerParameter.query.filter(
        OzpMaxPowerParameter.ozp_end_year == int(ozp_end_year)
    )
    if vid is None:
        q = q.filter(OzpMaxPowerParameter.database_version_id.is_(None))
    else:
        q = q.filter(OzpMaxPowerParameter.database_version_id == int(vid))
    row = q.order_by(OzpMaxPowerParameter.id.desc()).first()
    if row is None:
        row = OzpMaxPowerParameter(
            ozp_end_year=int(ozp_end_year),
            database_version_id=vid,
        )
        db.session.add(row)
    if max_power_mw is not ...:
        row.max_power_mw = max_power_mw
    if growth_pct is not ...:
        row.growth_pct = growth_pct
    return row


def save_ozp_cells(cells: list[dict[str, Any]]) -> dict[str, Any]:
    """Сохранить ячейки `{ozp_end_year, parameter_key, value}`."""
    grouped: dict[int, dict[str, Any]] = {}
    for item in cells:
        try:
            year = int(item.get("ozp_end_year"))
        except (TypeError, ValueError):
            return {"ok": False, "error": "Неверный период ОЗП"}
        if year < 1900 or year > 2200:
            return {"ok": False, "error": "Неверный период ОЗП"}
        key = str(item.get("parameter_key") or "").strip()
        if key != "max_power":
            return {"ok": False, "error": "Неверный показатель"}
        try:
            parsed = _parse_optional_decimal(item.get("value"))
        except ValueError as exc:
            return {"ok": False, "error": str(exc)}
        grouped.setdefault(year, {})[key] = parsed

    for year, payload in grouped.items():
        upsert_ozp_row(
            year,
            max_power_mw=payload["max_power"] if "max_power" in payload else ...,
        )
    db.session.commit()
    return {"ok": True}


def build_ozp_excel_stream(
    *,
    columns: list[dict[str, Any]],
    year_groups: list[dict[str, Any]],
    stored: dict[int, OzpMaxPowerParameter],
    rounding_digits: int,
) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "ОЗП"

    header1 = [PAGE_TITLE]
    header2 = ["ОЗП"]
    for group in year_groups:
        header1.append(f"{group['year']} г.")
        for _ in range(int(group["colspan"]) - 1):
            header1.append(None)
    for col in columns:
        header2.append(col["label"])

    ws.append(header1)
    ws.append(header2)

    col_idx = 2
    for group in year_groups:
        span = int(group["colspan"])
        if span > 1:
            ws.merge_cells(
                start_row=1,
                start_column=col_idx,
                end_row=1,
                end_column=col_idx + span - 1,
            )
        col_idx += span

    max_row = [ROW_MAX_LABEL]
    growth_row = [ROW_GROWTH_LABEL]
    max_raw: list[Any] = []
    for col in columns:
        row = stored.get(int(col["ozp_end"]))
        mw = None if row is None else row.max_power_mw
        max_raw.append(mw)
        max_row.append(None)
        growth_row.append(None)
    prev_end = int(columns[0]["ozp_end"]) - 1 if columns else None
    prev_row = stored.get(prev_end) if prev_end is not None else None
    prev_max = None if prev_row is None else prev_row.max_power_mw
    growth_raw = compute_growth_series(max_raw, prev_max)
    ws.append(max_row)
    ws.append(growth_row)

    for r in range(1, 5):
        for c in range(1, len(columns) + 2):
            cell = ws.cell(row=r, column=c)
            cell.border = _CELL_BORDER
            cell.font = _HEADER_FONT if r <= 2 else _BASE_FONT
            cell.alignment = _HEADER_ALIGN if r <= 2 or c > 1 else _LEFT_ALIGN
            if r <= 2:
                cell.fill = _HEADER_FILL

    mw_format = excel_number_format(rounding_digits)
    growth_format = excel_number_format(GROWTH_ROUNDING_DIGITS)
    for i, raw in enumerate(max_raw, start=2):
        value, numeric = excel_numeric_cell_value(raw)
        cell = ws.cell(row=3, column=i, value=value)
        cell.alignment = _CENTER_ALIGN
        cell.border = _CELL_BORDER
        cell.font = _BASE_FONT
        if numeric:
            cell.number_format = mw_format
    for i, raw in enumerate(growth_raw, start=2):
        value, numeric = excel_numeric_cell_value(raw)
        cell = ws.cell(row=4, column=i, value=value)
        cell.alignment = _CENTER_ALIGN
        cell.border = _CELL_BORDER
        cell.font = _BASE_FONT
        if numeric:
            cell.number_format = growth_format

    ws.column_dimensions["A"].width = 48
    for i in range(2, len(columns) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    ws.freeze_panes = "B3"

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
