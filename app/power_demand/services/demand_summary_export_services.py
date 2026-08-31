"""Выгрузка сводных таблиц нагрузок в Excel (оформление как на экранной сводке).

Числовые ячейки: полное значение из tooltip + number_format (как на выгрузке «Выработка ЭЭ»).
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.common.services.excel_numeric_cell import (
    excel_number_format,
    excel_numeric_cell_value,
    parse_excel_numeric,
)
from app.power_demand.services.demand_summary_services import (
    CALCULATED_MAX_MW_ROUNDING_DIGITS,
    CALCULATED_MAX_PARAMETER_KEYS,
)

_EXCEL_TEXT_PARAMETER_KEYS = frozenset({"peak_datetime"})

_HEADER_FILL = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
_HEADER_FONT = Font(bold=True, size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

_ENTITY_FILL = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
_ENTITY_FONT = Font(bold=True, size=11)
_ENTITY_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)

_FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
_FILL_STRIPE = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

_THIN = Side(style="thin", color="CCCCCC")
_CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BASE_FONT = Font(size=11)
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _excel_rounding_digits_for_parameter(parameter_key: str, rounding_digits: int) -> int:
    if parameter_key in CALCULATED_MAX_PARAMETER_KEYS:
        return CALCULATED_MAX_MW_ROUNDING_DIGITS
    return rounding_digits


def _excel_raw_year_source(row: dict[str, Any], index: int) -> Any:
    tooltips = row.get("year_numeric_tooltips") or []
    if index < len(tooltips) and tooltips[index]:
        return tooltips[index]
    yvals = row.get("year_values") or []
    if index < len(yvals):
        return yvals[index]
    return None


def _excel_raw_k_source(row: dict[str, Any], index: int) -> Any:
    tooltips = row.get("year_k_full_tooltips") or []
    if index < len(tooltips) and tooltips[index]:
        return tooltips[index]
    ykvals = row.get("year_k_values") or []
    if index < len(ykvals):
        return ykvals[index]
    return None


def _excel_raw_hist_source(row: dict[str, Any]) -> Any:
    tip = row.get("hist_numeric_tooltip")
    if tip:
        return tip
    return row.get("hist_value")


def _excel_cell_value_for_parameter(
    parameter_key: str,
    raw: Any,
    *,
    force_numeric: bool = False,
    verification: bool = False,
) -> tuple[Any, bool]:
    if not force_numeric and parameter_key in _EXCEL_TEXT_PARAMETER_KEYS:
        if raw in (None, ""):
            return "—", False
        return raw, False
    return excel_numeric_cell_value(raw, verification=verification, zero_as_dash=False)


def _excel_hide_plan_year_cell(
    y: int,
    pk: str,
    row: dict[str, Any],
    year_is_plan: dict[int, bool] | None,
    coeff_base_year: int | None,
) -> bool:
    """Пустые «план»-столбцы для всех строк, кроме max_power, РЭС «Совмещенный максимум…» и расчётных строк УЭС."""
    if not year_is_plan or not year_is_plan.get(y):
        return False
    pk = str(pk or "")
    if pk in (
        "cz_total_sum_fo_max_power",
        "cz_total_sum_res_combined_cz",
        "cz_total_imbalance_mw",
    ):
        return False
    if pk == "max_power":
        return False
    if row.get("pd_pd_verify_for_row") or str(pk).startswith("verify_for_"):
        return False
    if pk == "calculated_max_fo_mw":
        return False
    if (
        coeff_base_year is not None
        and (coeff_base_year + 1) <= y <= (coeff_base_year + 6)
        and pk
        in (
            "combined_on_oes",
            "combined_on_ees",
            "combined_on_fo",
            "combined_on_cz",
            "combined_on_ez",
        )
        and row.get("demand_model_name") == "RegionalEnergySystemDemandParameter"
    ):
        return False
    if coeff_base_year is not None:
        rep_med = (coeff_base_year - 9) <= y <= coeff_base_year or (
            (coeff_base_year + 1) <= y <= (coeff_base_year + 6)
        )
        med_only = (coeff_base_year + 1) <= y <= (coeff_base_year + 6)
        dm = row.get("demand_model_name")
        if dm == "UnionEnergySystemDemandParameter":
            if pk in ("calculated_max_power_mw", "calculated_combined_on_ees_mw") and rep_med:
                return False
            if pk == "combined_on_ees" and med_only:
                return False
        if dm == "FederalDistrictDemandParameter":
            if pk in ("calculated_max_power_mw", "calculated_combined_on_cz_mw") and rep_med:
                return False
            if pk == "combined_on_cz" and med_only:
                return False
        if dm == "EnergyZoneDemandParameter":
            if pk in ("calculated_max_power_mw", "calculated_combined_on_ees_mw") and rep_med:
                return False
            if pk == "combined_on_ees" and med_only:
                return False
    return True


def _year_header_text(y: int, year_features: dict[int, Any] | dict[Any, Any] | None) -> str:
    yf = year_features or {}
    cap = yf.get(y)
    s = str(cap).strip() if cap is not None else ""
    if s:
        return f"{y}\n{s}"
    return str(y)


def _coeff_year_segment(y: int, n: int) -> str:
    if (n - 9) <= y <= n:
        return "reporting"
    if (n + 1) <= y <= (n + 6):
        return "medium"
    return "long"


def _coeff_period_groups_for_years(
    years: list[int],
    coeff_base_year: int,
    *,
    include_long: bool,
) -> list[tuple[str, list[int]]]:
    labels = {
        "reporting": "Отчетный период",
        "medium": "Среднесрочный период",
        "long": "Долгосрочный период",
    }
    order = ["reporting", "medium"]
    if include_long:
        order.append("long")
    buckets: dict[str, list[int]] = {k: [] for k in order}
    for y in years:
        seg = _coeff_year_segment(y, coeff_base_year)
        if seg in buckets:
            buckets[seg].append(y)
    return [(labels[k], buckets[k]) for k in order if buckets[k]]


def _style_header_cell(cell) -> None:
    cell.font = _HEADER_FONT
    cell.fill = _HEADER_FILL
    cell.alignment = _HEADER_ALIGN
    cell.border = _CELL_BORDER


def _write_data_cell(
    ws,
    row_idx: int,
    col_idx: int,
    value: Any,
    *,
    is_numeric: bool,
    parameter_key: str,
    rounding_digits: int,
    is_k_column: bool,
    rounding_digits_k: int,
    param_fill: PatternFill,
    verification_row: bool,
) -> None:
    cell = ws.cell(row=row_idx, column=col_idx, value=value)
    cell.fill = param_fill
    cell.border = _CELL_BORDER
    cell.alignment = _CENTER_WRAP

    if verification_row:
        parsed = parse_excel_numeric(value) if is_numeric else None
        if parsed is not None and parsed != 0:
            cell.font = Font(size=11, italic=True, color="FF0000")
        else:
            cell.font = Font(size=11, italic=True)
        if is_numeric:
            rd = rounding_digits_k if is_k_column else _excel_rounding_digits_for_parameter(
                parameter_key, rounding_digits
            )
            cell.number_format = excel_number_format(rd)
        return

    cell.font = _BASE_FONT
    if not is_numeric:
        return
    if parameter_key in _EXCEL_TEXT_PARAMETER_KEYS:
        return

    rd = rounding_digits_k if is_k_column else _excel_rounding_digits_for_parameter(
        parameter_key, rounding_digits
    )
    cell.number_format = excel_number_format(rd)


def build_demand_summary_excel_stream(
    *,
    summary_rows: list[dict[str, Any]],
    years: list[int],
    sheet_title: str,
    year_features: dict[int, Any] | dict[Any, Any] | None = None,
    year_is_plan: dict[int, bool] | None = None,
    export_coeff_k_columns: bool = False,
    coeff_base_year: int | None = None,
    rounding_digits: int = 1,
    rounding_digits_k: int | None = None,
    show_hist_col: bool = True,
    coeff_include_long: bool = False,
) -> BytesIO:
    rd_k = rounding_digits_k if rounding_digits_k is not None else 3
    yf = year_features or {}

    wb = Workbook()
    ws = wb.active
    safe_title = (sheet_title or "Сводка").replace("/", "-").replace("\\", "-")[:31]
    ws.title = safe_title or "Сводка"

    fixed_cols = 2 + (1 if show_hist_col else 0)
    year_col_count = len(years) * (2 if export_coeff_k_columns else 1)
    note_col = fixed_cols + year_col_count + 1

    if export_coeff_k_columns:
        period_groups: list[tuple[str, list[int]]] = []
        if coeff_base_year is not None:
            period_groups = _coeff_period_groups_for_years(
                years, int(coeff_base_year), include_long=coeff_include_long
            )
        header_rows = 2 if period_groups else 1
        data_start_row = header_rows + 1

        if period_groups:
            ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
            ws.cell(row=1, column=1, value="Энергосистема")
            ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
            ws.cell(row=1, column=2, value="Наименование параметров")
            col = fixed_cols + 1
            for label, seg_years in period_groups:
                span = len(seg_years) * 2
                if span > 0:
                    ws.merge_cells(
                        start_row=1, start_column=col, end_row=1, end_column=col + span - 1
                    )
                    ws.cell(row=1, column=col, value=label)
                    col += span
            ws.merge_cells(start_row=1, start_column=note_col, end_row=2, end_column=note_col)
            ws.cell(row=1, column=note_col, value="Примечание")
            col = fixed_cols + 1
            for y in years:
                hdr = _year_header_text(y, yf)
                ws.cell(row=2, column=col, value=f"{hdr}\nk")
                ws.cell(row=2, column=col + 1, value=f"{hdr}\nМВт")
                col += 2
            for r in range(1, 3):
                for c in range(1, note_col + 1):
                    _style_header_cell(ws.cell(row=r, column=c))
            ws.freeze_panes = ws.cell(row=3, column=1).coordinate
        else:
            ws.row_dimensions[1].height = 36
            ws.cell(row=1, column=1, value="Энергосистема")
            ws.cell(row=1, column=2, value="Наименование параметров")
            col = fixed_cols + 1
            for y in years:
                hdr = _year_header_text(y, yf)
                ws.cell(row=1, column=col, value=f"{hdr}\nk")
                ws.cell(row=1, column=col + 1, value=f"{hdr}\nМВт")
                col += 2
            ws.cell(row=1, column=note_col, value="Примечание")
            for c in range(1, note_col + 1):
                _style_header_cell(ws.cell(row=1, column=c))
            ws.freeze_panes = "A2"
    else:
        header_rows = 1
        data_start_row = 2
        ws.row_dimensions[1].height = 36
        ws.cell(row=1, column=1, value="Энергосистема")
        ws.cell(row=1, column=2, value="Наименование параметров")
        col = 3
        if show_hist_col:
            ws.cell(row=1, column=col, value="Исторический собственный максимум")
            col += 1
        for y in years:
            ws.cell(row=1, column=col, value=_year_header_text(y, yf))
            col += 1
        ws.cell(row=1, column=note_col, value="Примечание")
        for c in range(1, note_col + 1):
            _style_header_cell(ws.cell(row=1, column=c))
        ws.freeze_panes = "A2"

    for bi, row in enumerate(summary_rows):
        excel_row = data_start_row + bi
        stripe = bool(bi % 2) and not export_coeff_k_columns
        param_fill = _FILL_STRIPE if stripe else _FILL_WHITE
        verification_row = str(row.get("entity_label") or "").startswith("Проверка ")

        entity_cell = ""
        if row.get("show_entity_cell"):
            depth = int(row.get("entity_depth") or 0)
            label = row.get("entity_label") or ""
            entity_cell = ("  " * depth) + str(label)

        pk = str(row.get("parameter_key") or "")

        c1 = ws.cell(row=excel_row, column=1, value=entity_cell or None)
        c1.font = (
            Font(bold=True, size=11, italic=True) if verification_row else _ENTITY_FONT
        )
        c1.fill = _ENTITY_FILL
        c1.alignment = _ENTITY_ALIGN
        c1.border = _CELL_BORDER

        c2 = ws.cell(row=excel_row, column=2, value=row.get("parameter_label") or "")
        c2.font = Font(size=11, italic=True) if verification_row else _BASE_FONT
        c2.fill = param_fill
        c2.alignment = _LEFT_WRAP
        c2.border = _CELL_BORDER

        col = 3
        if show_hist_col:
            hist_raw = _excel_raw_hist_source(row)
            hist_val, hist_num = _excel_cell_value_for_parameter(
                pk, hist_raw, verification=verification_row
            )
            _write_data_cell(
                ws,
                excel_row,
                col,
                hist_val,
                is_numeric=hist_num,
                parameter_key=pk,
                rounding_digits=rounding_digits,
                is_k_column=False,
                rounding_digits_k=rd_k,
                param_fill=param_fill,
                verification_row=verification_row,
            )
            col += 1

        for i, y in enumerate(years):
            hide_plan = _excel_hide_plan_year_cell(
                y, pk, row, year_is_plan, coeff_base_year
            )
            if export_coeff_k_columns:
                if hide_plan:
                    ws.cell(row=excel_row, column=col, value=None)
                    ws.cell(row=excel_row, column=col + 1, value=None)
                    for cc in (col, col + 1):
                        c = ws.cell(row=excel_row, column=cc)
                        c.fill = param_fill
                        c.border = _CELL_BORDER
                        c.alignment = _CENTER_WRAP
                    col += 2
                    continue
                k_raw = _excel_raw_k_source(row, i)
                k_val, k_num = _excel_cell_value_for_parameter(
                    pk, k_raw, force_numeric=True, verification=verification_row
                )
                _write_data_cell(
                    ws,
                    excel_row,
                    col,
                    k_val,
                    is_numeric=k_num,
                    parameter_key=pk,
                    rounding_digits=rounding_digits,
                    is_k_column=True,
                    rounding_digits_k=rd_k,
                    param_fill=param_fill,
                    verification_row=verification_row,
                )
                mw_raw = _excel_raw_year_source(row, i)
                mw_val, mw_num = _excel_cell_value_for_parameter(
                    pk, mw_raw, verification=verification_row
                )
                _write_data_cell(
                    ws,
                    excel_row,
                    col + 1,
                    mw_val,
                    is_numeric=mw_num,
                    parameter_key=pk,
                    rounding_digits=rounding_digits,
                    is_k_column=False,
                    rounding_digits_k=rd_k,
                    param_fill=param_fill,
                    verification_row=verification_row,
                )
                col += 2
            elif hide_plan:
                c = ws.cell(row=excel_row, column=col, value=None)
                c.fill = param_fill
                c.border = _CELL_BORDER
                c.alignment = _CENTER_WRAP
                col += 1
            else:
                raw = _excel_raw_year_source(row, i)
                cell_val, is_num = _excel_cell_value_for_parameter(
                    pk, raw, verification=verification_row
                )
                _write_data_cell(
                    ws,
                    excel_row,
                    col,
                    cell_val,
                    is_numeric=is_num,
                    parameter_key=pk,
                    rounding_digits=rounding_digits,
                    is_k_column=False,
                    rounding_digits_k=rd_k,
                    param_fill=param_fill,
                    verification_row=verification_row,
                )
                col += 1

        nn = row.get("entity_note_text")
        note_val = "" if nn is None else str(nn)
        c_note = ws.cell(row=excel_row, column=note_col, value=note_val)
        c_note.font = _BASE_FONT
        c_note.fill = param_fill
        c_note.alignment = _LEFT_WRAP
        c_note.border = _CELL_BORDER

    for col_idx in range(1, note_col + 1):
        letter = get_column_letter(col_idx)
        if col_idx == 1:
            ws.column_dimensions[letter].width = 44
        elif col_idx == 2:
            ws.column_dimensions[letter].width = 52
        elif show_hist_col and col_idx == 3:
            ws.column_dimensions[letter].width = 18
        elif col_idx == note_col:
            ws.column_dimensions[letter].width = 28
        else:
            ws.column_dimensions[letter].width = 14 if not export_coeff_k_columns else 12

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
