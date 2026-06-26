# -*- coding: utf-8 -*-
"""Экспорт таблицы электроёмкости в Excel (формат «Таблица 1»)."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.axis import DisplayUnitsLabelList
from openpyxl.chart.marker import Marker
from openpyxl.drawing.colors import ColorChoice
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.electrical_intensity.services.electrical_intensity_constants import (
    REF_ROW_ACCUM_FIXED_CAPITAL,
    REF_ROW_FD_ACCUM_FIXED_CAPITAL,
    ROW_KIND_CALCULATED,
    ROW_KIND_DELTA,
    ROW_KIND_GRAPH_POINT,
    ROW_KIND_INTENSITY,
)
from app.electrical_intensity.services.electrical_intensity_diagnostic_services import (
    compute_intensity_cell_diagnostics,
)

COL_TERRITORY = 1
COL_LABEL = 2
COL_UNIT = 3
YEAR_COL_START = 4

_CHART_STAGING_X_COL = 1
_CHART_STAGING_Y_COL = 2
_CHART_DATA_SHEET_TITLE = "_ei_chart_data"


_FILL_HEADER = PatternFill("solid", fgColor="D1E7DD")
_FILL_BY_CSS: dict[str, PatternFill] = {
    "lt-ei-row-intensity": PatternFill("solid", fgColor="CFE2FF"),
    "lt-ei-row-graph-point": PatternFill("solid", fgColor="F0EAF7"),
    "lt-ei-ref-product-output": PatternFill("solid", fgColor="FCE4E4"),
    "lt-ei-ref-consumption": PatternFill("solid", fgColor="D1E7DD"),
    "lt-ei-ref-accum-capital": PatternFill("solid", fgColor="E2E8F0"),
    "table-warning": PatternFill("solid", fgColor="FFF3CD"),
}
# Стили графиков — как в «02.06.26 Расчет электроемкости. Таблица 1.xlsx»
_CHART_WIDTH = 15
_CHART_HEIGHT = 7.5
_CHART_SCATTER_STYLE = "lineMarker"
_CHART_LEGEND_POSITION = "b"
_CHART_STYLE = 2
_CHART_LINE_WIDTH_EMU = 19050
_CHART_MARKER_LINE_WIDTH_EMU = 9525
_CHART_MARKER_SIZE = 5
_CHART_FACT_COLOR = "4BACC6"
_CHART_CALC_COLOR = "F79646"
_CHART_X_NUMFMT = "#,##0"
_CHART_Y_NUMFMT = "0.0"


def _chart_srgb_color(hex_rgb: str) -> ColorChoice:
    return ColorChoice(srgbClr=hex_rgb.upper())

_PERIOD_RETRO = "Ретроспективный период"
_PERIOD_FACT = "Факт"
_PERIOD_MEDIUM = "Сценарные условия (Базовый вариант)"
_PERIOD_LONG = "Долгосрочный прогноз (Базовый вариант)"


def _export_model_row_label(data_row: dict[str, Any], section: dict[str, Any]) -> str:
    rk = data_row.get("row_kind") or ""
    if rk == ROW_KIND_CALCULATED:
        if section.get("is_population_section"):
            return "Потребление ээ на душу РАСЧЕТНОЕ"
        return "Электроемкость РАСЧЕТНАЯ"
    if rk == ROW_KIND_DELTA:
        return "ДЕЛЬТА"
    return str(data_row.get("row_label") or "")


def _export_model_row_unit(
    data_row: dict[str, Any], section: dict[str, Any], default_unit: str
) -> str:
    rk = data_row.get("row_kind") or ""
    if rk in (ROW_KIND_GRAPH_POINT, ROW_KIND_CALCULATED, ROW_KIND_DELTA):
        return ""
    return str(section.get("unit_label") or default_unit or "")


def _visible_years(context: dict[str, Any]) -> list[int]:
    all_years: list[int] = context.get("years") or context.get("display_years") or []
    visible = context.get("lt_ei_initial_visible_years")
    if visible is None:
        return list(all_years)
    visible_set = set(visible)
    return [y for y in all_years if y in visible_set]


def _last_data_col(years: list[int]) -> int:
    return YEAR_COL_START + len(years) - 1 if years else COL_UNIT


def _fill_for_row_css(row_css: str | None) -> PatternFill | None:
    if not row_css:
        return None
    for token in str(row_css).split():
        fill = _FILL_BY_CSS.get(token)
        if fill is not None:
            return fill
    return None


def _period_segments(
    years: list[int],
    *,
    coeff_base_year: int,
    include_medium: bool,
    include_long: bool,
) -> list[tuple[list[int], str]]:
    if not years:
        return []
    n = int(coeff_base_year)
    retro = [y for y in years if y < n]
    fact = [y for y in years if y == n]
    medium_cap = n + 6
    medium = [y for y in years if n < y <= medium_cap] if include_medium else []
    long = [y for y in years if y > medium_cap] if include_long else []
    if not include_medium:
        long = [y for y in years if y > n]
    segments: list[tuple[list[int], str]] = []
    if retro:
        segments.append((retro, _PERIOD_RETRO))
    if fact:
        segments.append((fact, _PERIOD_FACT))
    if medium:
        segments.append((medium, _PERIOD_MEDIUM))
    if long:
        segments.append((long, _PERIOD_LONG))
    return segments


class _ExcelExportWriter:
    def __init__(
        self,
        ws: Worksheet,
        chart_data_ws: Worksheet,
        *,
        years: list[int],
        year_features: dict[int, str],
        coeff_base_year: int,
        include_medium: bool,
        include_long: bool,
        current_year: int | None,
        page_title: str,
    ) -> None:
        self.ws = ws
        self.chart_data_ws = chart_data_ws
        self.years = years
        self.year_features = year_features
        self.coeff_base_year = coeff_base_year
        self.include_medium = include_medium
        self.include_long = include_long
        self.current_year = current_year
        self.page_title = page_title
        self.last_col = _last_data_col(years)
        self.row_idx = 1
        self._chart_staging_row = 1
        self._territory_abbr_row: int | None = None

    def _apply_fill_row(self, row_idx: int, fill: PatternFill | None) -> None:
        if fill is None:
            return
        for col in range(COL_TERRITORY, self.last_col + 1):
            self.ws.cell(row=row_idx, column=col).fill = fill

    def _merge_banner(self, row_idx: int, *, col_start: int, title: str) -> None:
        if col_start > self.last_col:
            return
        self.ws.merge_cells(
            start_row=row_idx,
            start_column=col_start,
            end_row=row_idx,
            end_column=self.last_col,
        )
        cell = self.ws.cell(row=row_idx, column=col_start, value=title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def write_headers(self) -> None:
        ws = self.ws
        ws.cell(row=1, column=COL_LABEL, value=self.page_title).font = Font(bold=True)
        ws.cell(row=2, column=COL_LABEL, value="Наименование показателя").font = Font(
            bold=True
        )
        ws.cell(row=2, column=COL_UNIT, value="Ед.изм.").font = Font(bold=True)

        year_to_col = {y: YEAR_COL_START + i for i, y in enumerate(self.years)}
        for seg_years, title in _period_segments(
            self.years,
            coeff_base_year=self.coeff_base_year,
            include_medium=self.include_medium,
            include_long=self.include_long,
        ):
            if not seg_years:
                continue
            c0 = year_to_col[seg_years[0]]
            c1 = year_to_col[seg_years[-1]]
            ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c1)
            cell = ws.cell(row=2, column=c0, value=title)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for year in self.years:
            col = year_to_col[year]
            ws.cell(row=3, column=col, value=year).font = Font(bold=True)
            ws.cell(row=3, column=col).alignment = Alignment(
                horizontal="center", vertical="center"
            )

        self._apply_fill_row(2, _FILL_HEADER)
        self._apply_fill_row(3, _FILL_HEADER)
        self.row_idx = 4

    def begin_territory(self, abbr: str) -> None:
        self._territory_abbr_row = self.row_idx
        territory_abbr = "РОССИЯ" if abbr == "РФ" else abbr
        self.ws.cell(row=self.row_idx, column=COL_TERRITORY, value=territory_abbr)
        self._merge_banner(self.row_idx, col_start=COL_LABEL, title=abbr)
        self.row_idx += 1

    def end_territory(self) -> None:
        if self._territory_abbr_row is None:
            return
        end_row = self.row_idx - 1
        if end_row > self._territory_abbr_row:
            self.ws.merge_cells(
                start_row=self._territory_abbr_row,
                start_column=COL_TERRITORY,
                end_row=end_row,
                end_column=COL_TERRITORY,
            )
            cell = self.ws.cell(row=self._territory_abbr_row, column=COL_TERRITORY)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        self._territory_abbr_row = None

    def write_ved_header(self, title: str) -> int:
        chart_anchor = self.row_idx
        self._merge_banner(self.row_idx, col_start=COL_LABEL, title=title)
        self.row_idx += 1
        return chart_anchor

    def write_coef_rows(self, entity: dict[str, Any]) -> None:
        a_val = entity.get("coefficient_a_manual_display", entity.get("coefficient_a_display", ""))
        x_val = entity.get("coefficient_x_display", "")
        for label, value in (("А=", a_val), ("Х=", x_val)):
            self.ws.cell(row=self.row_idx, column=COL_LABEL, value=label)
            if value not in (None, "", "—"):
                self.ws.cell(row=self.row_idx, column=COL_UNIT, value=value)
            self.row_idx += 1

    def _write_year_cells(
        self, row_idx: int, cells_display: dict[Any, Any] | None
    ) -> None:
        cells_display = cells_display or {}
        for i, year in enumerate(self.years):
            disp = cells_display.get(year, "—")
            if disp is None or disp == "—":
                continue
            self.ws.cell(row=row_idx, column=YEAR_COL_START + i, value=disp)

    def write_data_row(
        self,
        *,
        label: str,
        unit: str,
        cells_display: dict[Any, Any] | None,
        row_css: str | None = None,
        ref_row: dict[str, Any] | None = None,
        hide_label: bool = False,
    ) -> int:
        row_idx = self.row_idx
        if not hide_label:
            self.ws.cell(row=row_idx, column=COL_LABEL, value=label)
        unit_text = unit or ""
        if ref_row and ref_row.get("has_coefficient_k"):
            k_disp = ref_row.get("coefficient_k_display", "")
            if k_disp not in (None, "", "—"):
                unit_text = str(k_disp)
        if unit_text:
            self.ws.cell(row=row_idx, column=COL_UNIT, value=unit_text)
        self._write_year_cells(row_idx, cells_display)
        self._apply_fill_row(row_idx, _fill_for_row_css(row_css))
        self.row_idx += 1
        return row_idx

    def _chart_staging_block(
        self, points: list[dict[str, Any]]
    ) -> tuple[int, int, int, int] | None:
        if not points:
            return None
        start = self._chart_staging_row
        row = start
        for pt in points:
            x_raw = pt.get("x")
            y_raw = pt.get("y")
            if x_raw is None or y_raw is None:
                continue
            try:
                x_val = float(x_raw)
                y_val = float(y_raw)
            except (TypeError, ValueError):
                continue
            self.chart_data_ws.cell(row=row, column=_CHART_STAGING_X_COL, value=x_val)
            self.chart_data_ws.cell(row=row, column=_CHART_STAGING_Y_COL, value=y_val)
            row += 1
        end_row = row - 1
        if end_row < start:
            return None
        self._chart_staging_row = end_row + 1
        return _CHART_STAGING_X_COL, _CHART_STAGING_Y_COL, start, end_row

    def attach_scatter_chart(
        self,
        *,
        anchor_row: int,
        chart_payload: dict[str, Any] | None,
        caption: str | None = None,
    ) -> None:
        if not chart_payload:
            return
        fact = chart_payload.get("fact") or []
        calc = chart_payload.get("calc") or []
        if not fact and not calc:
            return

        chart = ScatterChart()
        chart.scatterStyle = _CHART_SCATTER_STYLE
        chart.style = _CHART_STYLE
        chart.roundedCorners = False
        chart.title = (caption or "")[:255]
        chart.x_axis.title = (chart_payload.get("x_axis_label") or "")[:255]
        chart.y_axis.title = (chart_payload.get("y_axis_label") or "")[:255]
        chart.x_axis.axPos = "b"
        chart.y_axis.axPos = "l"
        chart.x_axis.numFmt = _CHART_X_NUMFMT
        chart.x_axis.dispUnits = DisplayUnitsLabelList(builtInUnit="thousands")
        chart.x_axis.majorTickMark = "out"
        chart.x_axis.minorTickMark = "none"
        chart.x_axis.tickLblPos = "nextTo"
        chart.y_axis.numFmt = _CHART_Y_NUMFMT
        chart.y_axis.majorTickMark = "out"
        chart.y_axis.minorTickMark = "none"
        chart.y_axis.tickLblPos = "nextTo"
        chart.width = _CHART_WIDTH
        chart.height = _CHART_HEIGHT
        chart.legend.position = _CHART_LEGEND_POSITION
        chart.visible_cells_only = False

        for points, title, color, markers in (
            (fact, "Фактическая", _CHART_FACT_COLOR, True),
            (calc, "Расчётная", _CHART_CALC_COLOR, False),
        ):
            block = self._chart_staging_block(points)
            if not block:
                continue
            x_col, y_col, lo, hi = block
            if hi < lo:
                continue
            x_ref = Reference(self.chart_data_ws, min_col=x_col, min_row=lo, max_row=hi)
            y_ref = Reference(self.chart_data_ws, min_col=y_col, min_row=lo, max_row=hi)
            ser = Series(y_ref, x_ref, title=title)
            ser.graphicalProperties.line.width = _CHART_LINE_WIDTH_EMU
            if markers:
                ser.marker = Marker("circle", size=_CHART_MARKER_SIZE)
                ser.marker.graphicalProperties.solidFill = _chart_srgb_color(color)
                ser.marker.graphicalProperties.line.width = _CHART_MARKER_LINE_WIDTH_EMU
                ser.marker.graphicalProperties.line.noFill = True
                ser.graphicalProperties.line.noFill = True
            else:
                ser.marker = Marker("none")
                ser.graphicalProperties.line.solidFill = _chart_srgb_color(color)
            chart.series.append(ser)

        if not chart.series:
            return

        anchor_col = self.last_col + 2
        self.ws.add_chart(chart, f"{get_column_letter(anchor_col)}{anchor_row}")

    def write_fd_summary(
        self,
        fd_summary: dict[str, Any],
        *,
        unit_label: str,
        territory_title: str,
    ) -> None:
        summary_chart_row: int | None = None
        for ref_row in fd_summary.get("reference_rows") or []:
            self.write_data_row(
                label=str(ref_row.get("row_label") or ""),
                unit=str(ref_row.get("unit_label") or ""),
                cells_display=ref_row.get("cells_display"),
                row_css=ref_row.get("row_css"),
                ref_row=ref_row,
            )
        if fd_summary.get("has_ei_intensity_block"):
            for data_row in fd_summary.get("rows") or []:
                if data_row.get("row_kind") != ROW_KIND_INTENSITY:
                    continue
                self.write_data_row(
                    label=str(data_row.get("row_label") or ""),
                    unit=unit_label,
                    cells_display=data_row.get("cells_display"),
                    row_css=data_row.get("row_css"),
                )
        scatter = fd_summary.get("scatter_chart")
        if scatter:
            summary_chart_row = self.row_idx
            self.row_idx += 1
            self.attach_scatter_chart(
                anchor_row=summary_chart_row,
                chart_payload=scatter,
                caption=territory_title,
            )

    def _apply_graph_point_diagnostics(
        self,
        row_idx: int,
        *,
        intensity_cells: dict[int, Any],
        investment_cells: dict[int, Any],
    ) -> None:
        notes = compute_intensity_cell_diagnostics(
            display_years=self.years,
            intensity_cells=intensity_cells,
            investment_cells=investment_cells,
            current_year=self.current_year,
        )
        if not notes:
            return
        red_font = Font(color="FF0000")
        for i, year in enumerate(self.years):
            text = notes.get(year)
            if not text:
                continue
            cell = self.ws.cell(row=row_idx, column=YEAR_COL_START + i)
            cell.font = red_font
            try:
                cell.comment = Comment(str(text)[:2000], "Диагностика")
            except (TypeError, ValueError):
                pass

    def write_section(
        self,
        section: dict[str, Any],
        *,
        unit_label: str,
        rf_investment_cells: dict[int, Any],
    ) -> None:
        self.write_ved_header(str(section.get("ved_name") or ""))
        if section.get("has_ei_model_block"):
            self.write_coef_rows(section)
        for ref_row in section.get("reference_rows") or []:
            self.write_data_row(
                label=str(ref_row.get("row_label") or ""),
                unit=str(ref_row.get("unit_label") or ""),
                cells_display=ref_row.get("cells_display"),
                row_css=ref_row.get("row_css"),
            )
        if section.get("has_ei_block"):
            rows = section.get("rows") or []
            for data_row in rows:
                rk = data_row.get("row_kind") or ""
                if rk == ROW_KIND_INTENSITY and not section.get("has_ei_intensity_block"):
                    continue
                if rk != ROW_KIND_INTENSITY and not section.get("has_ei_model_block"):
                    continue
                hide_label = rk == ROW_KIND_GRAPH_POINT
                row_idx = self.write_data_row(
                    label=_export_model_row_label(data_row, section),
                    unit=_export_model_row_unit(data_row, section, unit_label),
                    cells_display=data_row.get("cells_display"),
                    row_css=data_row.get("row_css"),
                    hide_label=hide_label,
                )
                if rk == ROW_KIND_GRAPH_POINT and section.get("has_ei_model_block"):
                    intensity_cells = {
                        r.get("row_kind"): r
                        for r in rows
                        if r.get("row_kind") == ROW_KIND_INTENSITY
                    }
                    int_row = intensity_cells.get(ROW_KIND_INTENSITY) or {}
                    self._apply_graph_point_diagnostics(
                        row_idx,
                        intensity_cells=int_row.get("cells") or {},
                        investment_cells=_investment_cells_from_section(section),
                    )
        scatter = section.get("scatter_chart")
        if scatter:
            chart_row = self.row_idx
            self.row_idx += 1
            self.attach_scatter_chart(
                anchor_row=chart_row,
                chart_payload=scatter,
                caption=str(section.get("ved_name") or ""),
            )

    def write_flat_block_rows(
        self,
        block: dict[str, Any],
        *,
        unit_label: str,
        rf_investment_cells: dict[int, Any],
    ) -> None:
        rows = block.get("rows") or []
        for data_row in rows:
            rk = data_row.get("row_kind") or ""
            hide_label = rk == ROW_KIND_GRAPH_POINT
            label = _export_model_row_label(data_row, block) if rk in (
                ROW_KIND_GRAPH_POINT, ROW_KIND_CALCULATED, ROW_KIND_DELTA
            ) else str(data_row.get("row_label") or "")
            row_idx = self.write_data_row(
                label=label,
                unit=_export_model_row_unit(data_row, block, unit_label),
                cells_display=data_row.get("cells_display"),
                row_css=data_row.get("row_css"),
                hide_label=hide_label,
            )
            if rk == ROW_KIND_GRAPH_POINT:
                intensity_cells = _intensity_cells_from_rows(rows)
                self._apply_graph_point_diagnostics(
                    row_idx,
                    intensity_cells=intensity_cells,
                    investment_cells=rf_investment_cells,
                )

    def apply_layout(self) -> None:
        ws = self.ws
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 52
        ws.column_dimensions["C"].width = 14
        for i in range(len(self.years)):
            col = get_column_letter(YEAR_COL_START + i)
            ws.column_dimensions[col].width = 11


def _intensity_cells_from_rows(rows: list[dict[str, Any]]) -> dict[int, Any]:
    for row in rows:
        if row.get("row_kind") == ROW_KIND_INTENSITY:
            return row.get("cells") or {}
    return {}


def _investment_cells_from_section(section: dict[str, Any]) -> dict[int, Any]:
    for ref_row in section.get("reference_rows") or []:
        if ref_row.get("row_kind") in (
            REF_ROW_ACCUM_FIXED_CAPITAL,
            REF_ROW_FD_ACCUM_FIXED_CAPITAL,
        ):
            return ref_row.get("cells") or {}
    return {}


def build_electrical_intensity_excel_stream(context: dict[str, Any]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "таблица 1"
    chart_data_ws = wb.create_sheet(_CHART_DATA_SHEET_TITLE)
    chart_data_ws.sheet_state = "hidden"

    years = _visible_years(context)
    writer = _ExcelExportWriter(
        ws,
        chart_data_ws,
        years=years,
        year_features=context.get("year_features") or {},
        coeff_base_year=int(context.get("coeff_base_year") or 2025),
        include_medium=bool(context.get("summary_include_medium_years", True)),
        include_long=bool(context.get("summary_include_long_years", True)),
        current_year=context.get("ei_current_year"),
        page_title=str(context.get("page_title") or "ПРИМЕР РАСЧЕТА"),
    )
    writer.write_headers()

    ei_unit = context.get("unit_label") or "кВт.ч./тыс.руб."
    rf_investment = context.get("rf_investment_cells") or {}

    for block in context.get("territory_blocks") or []:
        abbr = str(block.get("abbr") or "")
        writer.begin_territory(abbr)

        fd_summary = block.get("fd_summary") or {}
        if fd_summary:
            territory_title = f"{abbr} — {block.get('label', '')}"
            writer.write_fd_summary(
                fd_summary,
                unit_label=ei_unit,
                territory_title=territory_title,
            )

        if block.get("layout") != "ved_sections" and block.get("has_ei_model_block", True):
            writer.write_coef_rows(block)

        for section in block.get("ved_sections") or []:
            writer.write_section(
                section,
                unit_label=ei_unit,
                rf_investment_cells=rf_investment,
            )

        if block.get("layout") != "ved_sections" and block.get("has_ei_model_block", True):
            writer.write_flat_block_rows(
                block,
                unit_label=ei_unit,
                rf_investment_cells=rf_investment,
            )

        writer.end_territory()

    writer.apply_layout()

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
