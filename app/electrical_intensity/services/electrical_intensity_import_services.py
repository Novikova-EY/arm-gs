# -*- coding: utf-8 -*-
"""Импорт электроёмкости из Excel."""

from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.common.services.database_version_services import get_current_version
from app.extensions import db
from app.electrical_intensity.services.electrical_intensity_constants import (
    EI_MODEL_ROW_KINDS,
    REF_ROW_LABEL_BY_KIND,
    RUSSIA_TERRITORY_LABELS,
    ROW_KIND_INTENSITY,
    ROW_KINDS,
    ROW_LABEL_BY_KIND,
)
from app.electrical_intensity.services.electrical_intensity_services import (
    _ei_row_kinds_for_ved,
    _is_total_consumption_ved,
    _refdata_ved_ids,
    _refdata_ved_types_for_version,
)
from app.refdata.models.economic_activity.economic_activity_type_model import (
    EconomicActivityType,
)
from app.electrical_intensity.services.electrical_intensity_logging import (
    log_electrical_intensity_excel_import,
)
from app.electrical_intensity.services.electrical_intensity_services import (
    _federal_districts_for_page,
    _get_or_create_fd_coef,
    _get_or_create_fd_year_row,
    _get_or_create_rf_coef,
    _get_or_create_rf_year_row,
    _username,
)
from app.economics.services.ved_consumption_services import _normalize_label
from app.generation.services.machine_services.machine_services import (
    is_same_decimal,
    to_decimal,
)


def _parse_year_header(cell_value: Any) -> int | None:
    if cell_value is None:
        return None
    if isinstance(cell_value, (int, float)) and float(cell_value) == int(cell_value):
        y = int(cell_value)
        if 1900 <= y <= 2200:
            return y
    s = str(cell_value).strip()
    m = re.search(r"(19|20)\d{2}", s)
    if m:
        return int(m.group(0))
    try:
        y = int(float(s.replace(",", ".")))
        if 1900 <= y <= 2200:
            return y
    except (ValueError, TypeError):
        return None
    return None


def _parse_cell_value(cell_value: Any) -> Decimal | None:
    if cell_value is None or cell_value == "":
        return None
    if isinstance(cell_value, (int, float)):
        return Decimal(str(cell_value))
    s = str(cell_value).strip().replace("\u00a0", "").replace(" ", "")
    if not s or s in ("—", "-", "–"):
        return None
    s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _build_ved_lookup(version_id: int | None) -> dict[str, int]:
    lookup: dict[str, int] = {}
    for ved in _refdata_ved_types_for_version(version_id):
        for candidate in (ved.name, getattr(ved, "name_2", None)):
            if candidate:
                lookup[_normalize_label(candidate)] = ved.id
    return lookup


def _ref_row_labels_normalized() -> frozenset[str]:
    return frozenset(_normalize_label(v) for v in REF_ROW_LABEL_BY_KIND.values())


def _build_fd_lookup() -> dict[str, int]:
    lookup: dict[str, int] = {}
    for fd in _federal_districts_for_page():
        for candidate in (fd.name_abr, fd.name, fd.name_full):
            if candidate:
                lookup[_normalize_label(candidate)] = fd.id
    return lookup


def _row_kind_from_label(label: str) -> str | None:
    n = _normalize_label(label)
    for kind in ROW_KINDS:
        if _normalize_label(ROW_LABEL_BY_KIND.get(kind, "")) == n:
            return kind
    aliases = {
        "электроемкость": ROW_KIND_INTENSITY,
        "электроемкость расчетная": "calculated",
        "дельта": "delta",
        "δ для электроемкости": "delta",
        "δ": "delta",
    }
    if n.startswith("электроемкость вычисляется"):
        return ROW_KIND_INTENSITY
    return aliases.get(n)


def import_electrical_intensity_from_xlsx_bytes(raw: bytes) -> dict[str, Any]:
    version_id = get_current_version()
    user = _username()
    wb = load_workbook(BytesIO(raw), data_only=True)
    ws = wb.active

    year_cols: dict[int, int] = {}
    header_row = 1
    for col in range(2, (ws.max_column or 1) + 1):
        y = _parse_year_header(ws.cell(row=header_row, column=col).value)
        if y is not None:
            year_cols[col] = y

    if not year_cols:
        raise ValueError("Не найдены столбцы с годами в первой строке.")

    fd_lookup = _build_fd_lookup()
    ved_lookup = _build_ved_lookup(version_id)
    ref_labels = _ref_row_labels_normalized()
    refdata_ids = _refdata_ved_ids(version_id)
    ved_by_id = {v.id: v for v in _refdata_ved_types_for_version(version_id)}
    stats = {
        "territories": 0,
        "rows_processed": 0,
        "cells_written": 0,
        "cells_skipped": 0,
        "warnings": [],
    }
    territories_seen: set[str] = set()

    current_kind: str | None = None
    current_fd_id: int | None = None
    current_ved_id: int | None = None

    for row_idx in range(2, (ws.max_row or 1) + 1):
        label_raw = ws.cell(row=row_idx, column=1).value
        if label_raw is None:
            continue
        label = str(label_raw).strip()
        if not label:
            continue
        label_n = _normalize_label(label)

        if label_n in RUSSIA_TERRITORY_LABELS or label_n.startswith("рф "):
            current_kind = "rf"
            current_fd_id = None
            current_ved_id = None
            territories_seen.add("rf")
            continue

        if label_n in fd_lookup:
            current_kind = "fd"
            current_fd_id = fd_lookup[label_n]
            current_ved_id = None
            territories_seen.add(f"fd:{current_fd_id}")
            continue

        if label_n in ved_lookup:
            current_ved_id = ved_lookup[label_n]
            continue

        if label_n in ref_labels:
            continue

        for sep in (" — ", " - "):
            if sep in label and label_n not in ROW_LABEL_BY_KIND.values():
                parts = [_normalize_label(p) for p in label.split(sep, 1)]
                for p in parts:
                    if p in fd_lookup:
                        current_kind = "fd"
                        current_fd_id = fd_lookup[p]
                        current_ved_id = None
                        territories_seen.add(f"fd:{current_fd_id}")
                        break
                    if p in RUSSIA_TERRITORY_LABELS:
                        current_kind = "rf"
                        current_fd_id = None
                        current_ved_id = None
                        territories_seen.add("rf")
                        break

        skip_ei_for_ved = False
        if current_ved_id is not None:
            ved_obj = ved_by_id.get(current_ved_id)
            if ved_obj is not None and _is_total_consumption_ved(ved_obj):
                skip_ei_for_ved = True

        m_a = re.match(r"^a\s*=\s*(.+)$", label_n, re.IGNORECASE)
        m_x = re.match(r"^x\s*=\s*(.+)$", label_n, re.IGNORECASE)
        if m_a or m_x:
            if current_ved_id is None or current_ved_id not in refdata_ids:
                continue
            if current_kind is None:
                stats["warnings"].append(f"Строка {row_idx}: коэффициент без территории.")
                continue
            val = _parse_cell_value(label.split("=", 1)[-1] if "=" in label else None)
            if val is None and m_a:
                val = _parse_cell_value(ws.cell(row=row_idx, column=2).value)
            if skip_ei_for_ved:
                continue
            if current_kind == "rf":
                coef_row = _get_or_create_rf_coef(version_id, user)
            elif current_ved_id is not None:
                coef_row = _get_or_create_fd_coef(
                    version_id, current_fd_id, current_ved_id, user
                )
            else:
                stats["warnings"].append(
                    f"Строка {row_idx}: коэффициент без ВЭД ({label})."
                )
                continue
            if m_a:
                if not is_same_decimal(to_decimal(coef_row.coefficient_a), to_decimal(val)):
                    coef_row.coefficient_a = val
                    coef_row.modified_by = user
                    stats["cells_written"] += 1
            else:
                if not is_same_decimal(to_decimal(coef_row.coefficient_x), to_decimal(val)):
                    coef_row.coefficient_x = val
                    coef_row.modified_by = user
                    stats["cells_written"] += 1
            continue

        row_kind = _row_kind_from_label(label)
        if row_kind is None:
            continue
        if current_ved_id is not None:
            ved_obj = ved_by_id.get(current_ved_id)
            if ved_obj is not None:
                _, _, allowed_kinds = _ei_row_kinds_for_ved(
                    ved_obj, refdata_ids=refdata_ids
                )
                if row_kind not in allowed_kinds:
                    continue
            elif row_kind in EI_MODEL_ROW_KINDS:
                continue
        if current_kind is None:
            stats["warnings"].append(f"Строка {row_idx}: данные без территории ({label}).")
            continue
        if skip_ei_for_ved:
            continue

        stats["rows_processed"] += 1
        for col, year_n in year_cols.items():
            val = _parse_cell_value(ws.cell(row=row_idx, column=col).value)
            if current_kind == "rf":
                db_row = _get_or_create_rf_year_row(version_id, row_kind, year_n, user)
            elif current_ved_id is not None:
                db_row = _get_or_create_fd_year_row(
                    version_id,
                    current_fd_id,
                    current_ved_id,
                    row_kind,
                    year_n,
                    user,
                )
            else:
                stats["warnings"].append(
                    f"Строка {row_idx}: показатель без ВЭД ({label})."
                )
                break
            if is_same_decimal(to_decimal(db_row.parameter_value), to_decimal(val)):
                stats["cells_skipped"] += 1
                continue
            db_row.parameter_value = val
            db_row.modified_by = user
            stats["cells_written"] += 1

    stats["territories"] = len(territories_seen)
    log_electrical_intensity_excel_import(user, stats, database_version_id=version_id)
    if stats["cells_written"]:
        from app.common.services.economics_fd_data_cache import (
            invalidate_economics_fd_data_cache,
        )

        invalidate_economics_fd_data_cache(
            version_id,
            "ei_year",
            "ei_coef",
            "pop_ei_year",
            "pop_ei_coef",
        )
    return stats
