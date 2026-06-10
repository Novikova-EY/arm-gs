# -*- coding: utf-8 -*-
"""Экспорт таблицы электроёмкости в Excel."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font

from app.energy_consumption.long_term_consumption.services.electrical_intensity_constants import (
    REF_ROW_ACCUM_FIXED_CAPITAL,
    REF_ROW_FD_ACCUM_FIXED_CAPITAL,
    ROW_KIND_INTENSITY,
)
from app.energy_consumption.long_term_consumption.services.electrical_intensity_diagnostic_services import (
    compute_intensity_cell_diagnostics,
)

from app.common.services.help_services import format_decimal_trim_for_display
from app.energy_consumption.long_term_consumption.services.electrical_intensity_constants import (
    EI_GRAPH_POINT_ROUNDING_DIGITS,
    EI_INTENSITY_MAX_FRACTION_DIGITS,
    EI_INTENSITY_VALUE_ROW_KINDS,
    ROW_KIND_GRAPH_POINT,
)


def build_electrical_intensity_excel_stream(context: dict[str, Any]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Электроемкость"

    years: list[int] = context.get("display_years") or []
    rd = int(context.get("rounding_digits") or 1)

    def _cell_display_digits(row_kind: str) -> int:
        if row_kind == ROW_KIND_GRAPH_POINT:
            return EI_GRAPH_POINT_ROUNDING_DIGITS
        if rd == 0 and row_kind in EI_INTENSITY_VALUE_ROW_KINDS:
            return EI_INTENSITY_MAX_FRACTION_DIGITS
        return rd
    yf = context.get("year_features") or {}
    ei_unit = context.get("unit_label") or "кВт.ч./тыс.руб."
    year_col_start = 3
    current_year = context.get("ei_current_year")
    red_font = Font(color="FF0000")

    def _intensity_cells_from_rows(rows: list[dict[str, Any]]) -> dict[int, Any]:
        for row in rows:
            if row.get("row_kind") == ROW_KIND_INTENSITY:
                return row.get("cells") or {}
        return {}

    def _apply_graph_point_diagnostics(
        row_idx: int,
        intensity_cells: dict[int, Any],
        investment_cells: dict[int, Any],
    ) -> None:
        notes = compute_intensity_cell_diagnostics(
            display_years=years,
            intensity_cells=intensity_cells,
            investment_cells=investment_cells,
            current_year=current_year,
        )
        if not notes:
            return
        for col_idx, year in enumerate(years, start=year_col_start):
            text = notes.get(year)
            if not text:
                continue
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = red_font
            cell.comment = Comment(text, "Диагностика")

    def _investment_cells_from_section(section: dict[str, Any]) -> dict[int, Any]:
        for ref_row in section.get("reference_rows") or []:
            if ref_row.get("row_kind") in (
                REF_ROW_ACCUM_FIXED_CAPITAL,
                REF_ROW_FD_ACCUM_FIXED_CAPITAL,
            ):
                return ref_row.get("cells") or {}
        return {}

    def year_header_cell(y: int) -> str:
        cap = yf.get(y)
        s = str(cap).strip() if cap is not None else ""
        if s:
            return f"{y}\n{s}"
        return str(y)

    ws.cell(row=1, column=1, value="Территория / ВЭД / показатель").font = Font(bold=True)
    ws.cell(row=1, column=2, value="ед.изм.").font = Font(bold=True)
    for col_idx, year in enumerate(years, start=year_col_start):
        cell = ws.cell(row=1, column=col_idx, value=year_header_cell(year))
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row_idx = 2
    for block in context.get("territory_blocks") or []:
        title = f"{block.get('abbr', '')} — {block.get('label', '')}"
        ws.cell(row=row_idx, column=1, value=title).font = Font(bold=True)
        row_idx += 1
        fd_summary = block.get("fd_summary") or {}
        for ref_row in fd_summary.get("reference_rows") or []:
            ws.cell(row=row_idx, column=1, value=ref_row.get("row_label") or "")
            ws.cell(row=row_idx, column=2, value=ref_row.get("unit_label") or "")
            cells = ref_row.get("cells") or {}
            for col_idx, year in enumerate(years, start=year_col_start):
                val = cells.get(year)
                if val is not None:
                    disp = format_decimal_trim_for_display(val, digits=rd)
                    ws.cell(row=row_idx, column=col_idx, value=disp)
            row_idx += 1
        if fd_summary.get("has_ei_intensity_block"):
            for data_row in fd_summary.get("rows") or []:
                if data_row.get("row_kind") != ROW_KIND_INTENSITY:
                    continue
                ws.cell(row=row_idx, column=1, value=data_row.get("row_label") or "")
                ws.cell(row=row_idx, column=2, value=ei_unit)
                cells = data_row.get("cells") or {}
                for col_idx, year in enumerate(years, start=year_col_start):
                    val = cells.get(year)
                    if val is not None:
                        disp = format_decimal_trim_for_display(
                            val, digits=_cell_display_digits(ROW_KIND_INTENSITY)
                        )
                        ws.cell(row=row_idx, column=col_idx, value=disp)
                row_idx += 1
        if block.get("layout") != "ved_sections" and block.get("has_ei_model_block", True):
            ws.cell(
                row=row_idx,
                column=1,
                value=(
                    f"A = {block.get('coefficient_a_manual_display', block.get('coefficient_a_display', '—'))}; "
                    f"Арасч. = {block.get('coefficient_a_computed_display', '—')}; "
                    f"X = {block.get('coefficient_x_display', '—')}"
                ),
            )
            ws.cell(row=row_idx, column=2, value="")
            row_idx += 1
        for section in block.get("ved_sections") or []:
            ws.cell(row=row_idx, column=1, value=section.get("ved_name") or "").font = Font(bold=True)
            row_idx += 1
            for ref_row in section.get("reference_rows") or []:
                ws.cell(row=row_idx, column=1, value=ref_row.get("row_label") or "")
                ws.cell(row=row_idx, column=2, value=ref_row.get("unit_label") or "")
                cells = ref_row.get("cells") or {}
                for col_idx, year in enumerate(years, start=year_col_start):
                    val = cells.get(year)
                    if val is not None:
                        disp = format_decimal_trim_for_display(val, digits=rd)
                        ws.cell(row=row_idx, column=col_idx, value=disp)
                row_idx += 1
            if section.get("has_ei_model_block"):
                ws.cell(
                    row=row_idx,
                    column=1,
                    value=(
                        f"A = {section.get('coefficient_a_manual_display', section.get('coefficient_a_display', '—'))}; "
                        f"Арасч. = {section.get('coefficient_a_computed_display', '—')}; "
                        f"X = {section.get('coefficient_x_display', '—')}"
                    ),
                )
                ws.cell(row=row_idx, column=2, value="")
                row_idx += 1
            if section.get("has_ei_block"):
                for data_row in section.get("rows") or []:
                    rk = data_row.get("row_kind") or ""
                    if rk == "intensity" and not section.get("has_ei_intensity_block"):
                        continue
                    if rk != "intensity" and not section.get("has_ei_model_block"):
                        continue
                    ws.cell(row=row_idx, column=1, value=data_row.get("row_label") or "")
                    row_unit = section.get("unit_label") or ei_unit
                    ws.cell(row=row_idx, column=2, value=row_unit)
                    cells = data_row.get("cells") or {}
                    for col_idx, year in enumerate(years, start=year_col_start):
                        val = cells.get(year)
                        if val is not None:
                            disp = format_decimal_trim_for_display(
                                val, digits=_cell_display_digits(rk)
                            )
                            ws.cell(row=row_idx, column=col_idx, value=disp)
                    if rk == ROW_KIND_GRAPH_POINT and section.get("has_ei_model_block"):
                        _apply_graph_point_diagnostics(
                            row_idx,
                            _intensity_cells_from_rows(section.get("rows") or []),
                            _investment_cells_from_section(section),
                        )
                    row_idx += 1
        if block.get("layout") != "ved_sections" and block.get("has_ei_model_block", True):
            for data_row in block.get("rows") or []:
                rk = data_row.get("row_kind") or ""
                ws.cell(row=row_idx, column=1, value=data_row.get("row_label") or "")
                ws.cell(row=row_idx, column=2, value=ei_unit)
                cells = data_row.get("cells") or {}
                for col_idx, year in enumerate(years, start=year_col_start):
                    val = cells.get(year)
                    if val is not None:
                        disp = format_decimal_trim_for_display(
                            val, digits=_cell_display_digits(rk)
                        )
                        ws.cell(row=row_idx, column=col_idx, value=disp)
                if rk == ROW_KIND_GRAPH_POINT:
                    _apply_graph_point_diagnostics(
                        row_idx,
                        _intensity_cells_from_rows(block.get("rows") or []),
                        context.get("rf_investment_cells") or {},
                    )
                row_idx += 1

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
