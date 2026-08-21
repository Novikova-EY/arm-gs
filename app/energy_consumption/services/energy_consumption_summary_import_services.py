# -*- coding: utf-8 -*-
"""Импорт сводных показателей потребления из Excel (листы «млн. кВт.ч» и/или «СиПР»)."""
from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal
from typing import Any, Callable, Optional, Type

from sqlalchemy.exc import IntegrityError
from sqlalchemy import select

from flask import session

from app.common.models.database_version_model import DatabaseVersion
from app.common.perimeter_variant.registry import (
    perimeter_entity_context_for_model,
    perimeter_variant_codes_for_entity,
)
from app.common.services.database_version_filter import (
    filter_by_explicit_db_version,
)
from app.common.services.database_version_services import get_current_version
from app.extensions import db
from app.energy_consumption.services import energy_consumption_parameter_services as ecps
from app.refdata.models.energy_systems.regional_district_regional_energy_system_model import (
    regional_district_regional_energy_system,
)
from app.refdata.models.energy_systems.energy_unit_model import EnergyUnit
from app.refdata.models.energy_systems.regional_energy_system_model import RegionalEnergySystem
from app.refdata.models.energy_systems.energy_system_type_model import EnergySystemType
from app.refdata.models.energy_systems.synchronous_area_model import SynchronousArea
from app.refdata.models.energy_systems.union_energy_system_model import UnionEnergySystem
from app.refdata.models.territories.federal_district_model import FederalDistrict
from app.refdata.models.territories.regional_district_model import RegionalDistrict
from app.refdata.models.years.year_model import Year

from app.energy_consumption.models.energy_systems.ees_russia_energy_consumption_parameter_model import (
    EesRussiaEnergyConsumptionParameter,
)
from app.energy_consumption.models.energy_systems.energy_system_type_energy_consumption_parameter_model import (
    EnergySystemTypeEnergyConsumptionParameter,
)
from app.energy_consumption.models.territories.russia_federation_energy_consumption_parameter_model import (
    RussiaFederationEnergyConsumptionParameter,
)
from app.common.perimeter_variant.constants import (
    CODE_WITH_NT,
    CODE_WITHOUT_NT,
    CODE_WITH_NT_WITH_GAES,
    CODE_WITH_NT_WITHOUT_GAES,
    CODE_WITHOUT_NT_WITH_GAES,
    CODE_WITHOUT_NT_WITHOUT_GAES,
    CODE_WITHOUT_NT_WITH_GAES_WITH_KALININGRAD_ES,
    CODE_WITHOUT_NT_WITH_KALININGRAD_ES,
    CODE_WITHOUT_NT_WITHOUT_KALININGRAD_ES,
    EES_RUSSIA_AGGREGATE_NAME,
    EES_UNIFIED_REF_NAME,
)
from app.common.perimeter_variant.registry import (
    model_supports_perimeter_variant,
    normalize_perimeter_variant_code,
)
from app.energy_consumption.models.energy_systems.energy_unit_energy_consumption_parameter_model import (
    EnergyUnitEnergyConsumptionParameter,
)
from app.energy_consumption.models.energy_systems.regional_energy_system_energy_consumption_parameter_model import (
    RegionalEnergySystemEnergyConsumptionParameter,
)
from app.energy_consumption.models.energy_systems.synchronous_area_energy_consumption_parameter_model import (
    SynchronousAreaEnergyConsumptionParameter,
)
from app.energy_consumption.models.energy_systems.union_energy_system_energy_consumption_parameter_model import (
    UnionEnergySystemEnergyConsumptionParameter,
)
from app.energy_consumption.models.territories.federal_district_energy_consumption_parameter_model import (
    FederalDistrictEnergyConsumptionParameter,
)
from app.energy_consumption.models.territories.regional_district_energy_consumption_parameter_model import (
    RegionalDistrictEnergyConsumptionParameter,
)


SHEET_MLN_KVTCH = "млн. кВт.ч"
SHEET_SIPR = "СиПР"

# Год в шапке: «2025», «2025 г.», «2025г», лат. «2025 g.»; «2025.0» из строкового представления числа Excel.
_YEAR_HEADER_RE = re.compile(
    r"^\s*(\d{4})(?:\s*(?:г|[gG])\.?)?\s*$",
    re.IGNORECASE,
)
_YEAR_HEADER_NUMSTR_RE = re.compile(r"^\s*(\d{4})(?:\.0+)?\s*$")

# Для поиска строки с годами в файле — без совпадения со справочником Years (его может ещё не быть).
_YEAR_HEADER_PLAUSIBLE = range(1990, 2051)


@dataclass(frozen=True)
class _ImportBind:
    demand_model: Type[Any]
    fk_column: str
    parent_id: int
    perimeter_variant_code: str | None = None


_MODEL_FK: dict[Type[Any], str] = {
    UnionEnergySystemEnergyConsumptionParameter: "id_union_energy_system",
    RegionalEnergySystemEnergyConsumptionParameter: "id_regional_energy_system",
    RegionalDistrictEnergyConsumptionParameter: "id_regional_district",
    FederalDistrictEnergyConsumptionParameter: "id_federal_district",
    SynchronousAreaEnergyConsumptionParameter: "id_synchronous_area",
    EnergyUnitEnergyConsumptionParameter: "id_energy_unit",
    EnergySystemTypeEnergyConsumptionParameter: "id_energy_system_type",
}

# Маркер строки параметров без FK родителя (одна строка на год в версии БД).
_FK_AGGREGATE_NO_PARENT = "__aggregate_no_parent__"

_KALININGRAD_ES_LABEL_SUFFIX_WITH = " с ЭС Калининградской области"
_NT_LABEL_INFIX = " без НТ"
_NT_WITH_TERRITORIES_LABEL_INFIX = " с НТ"
_GAES_WITH_CHARGE_LABEL_SUFFIX = " с зарядом ГАЭС"
_GAES_SHORT_LABEL_SUFFIXES = (
    _GAES_WITH_CHARGE_LABEL_SUFFIX,
    " с ГАЭС",
    " без заряда ГАЭС",
)
_VARIANT_COLUMN_HEADER_MARKERS = (
    "perimeter-variants",
    "perimeter variants",
    "вариант периметра",
    "варианты периметра",
)


def _norm_space(s: str) -> str:
    return " ".join(str(s).replace("\xa0", " ").split()).strip()


def _normalize_import_label_typos(label: str) -> str:
    """Исправляет типичные опечатки Excel: латинская «c» перед «НТ» и т.п."""
    s = _norm_space(str(label).replace("\xa0", " "))
    s = re.sub(r"(?i)(?<=\s)c(?=\s*[нn][тt]\b)", "с", s)
    # В сводных Excel встречаются слитные подписи вроде «ЕЭС России сНТ с зарядом ГАЭС».
    # Нормализуем их в канонический вид до разбора суффиксов варианта периметра.
    s = re.sub(r"(?i)(?<=\S)\s*[сc]\s*[нn]\s*[тt]\b", " с НТ", s)
    s = re.sub(r"(?i)(?<=\S)\s*без\s*[нn]\s*[тt]\b", " без НТ", s)
    s = re.sub(r"(?i)(?<=\S)\s*[сc]\s*зарядом\s*гаэс\b", " с зарядом ГАЭС", s)
    s = re.sub(r"(?i)(?<=\S)\s*[сc]\s*гаэс\b", " с ГАЭС", s)
    s = re.sub(r"(?i)(?<=\S)\s*без\s*заряда\s*гаэс\b", " без заряда ГАЭС", s)
    s = re.sub(r"(?i)(?<=\S)\s*без\s*гаэс\b", " без ГАЭС", s)
    s = _norm_space(s)
    return s


def _norm_entity_label(s: str) -> str:
    s_norm = _normalize_import_label_typos(s).casefold().replace("ё", "е")
    s_norm = re.sub(r"\bэнергосистема\b", "эс", s_norm)
    s_norm = re.sub(r"\bобъединенная\s+энергосистема\b", "оэс", s_norm)
    s_norm = re.sub(r"\bобъединённая\s+энергосистема\b", "оэс", s_norm)
    s_norm = re.sub(r"\bфедеральный\s+округ\b", "фо", s_norm)
    s_norm = re.sub(r",?\s*в\s+т\.?\s*ч\.?:?\s*$", "", s_norm)
    s_norm = _norm_space(s_norm.replace(".", " "))
    return s_norm


def _cell_as_text(val: Any) -> str:
    if val is None:
        return ""
    from datetime import date, datetime

    if isinstance(val, datetime):
        return str(val.year)
    if isinstance(val, date):
        return str(val.year)
    return str(val)


def _coerce_cell_to_year_number(val: Any) -> Optional[int]:
    """Распознать календарный год в ячейке по формату (число, «2025 г.», …), без справочника Years."""
    if val is None:
        return None
    if isinstance(val, bool):
        return None
    if isinstance(val, (int, float)):
        try:
            i = int(val)
        except (TypeError, ValueError):
            return None
        return i if i in _YEAR_HEADER_PLAUSIBLE else None
    s = _norm_space(_cell_as_text(val))
    if not s:
        return None
    m = _YEAR_HEADER_RE.match(s) or _YEAR_HEADER_NUMSTR_RE.match(s)
    if not m:
        return None
    y = int(m.group(1))
    return y if y in _YEAR_HEADER_PLAUSIBLE else None


def _parse_numeric_cell(val: Any) -> Optional[Decimal]:
    if val is None:
        return None
    if isinstance(val, bool):
        return None
    if isinstance(val, Decimal):
        return val
    if isinstance(val, (int, float)):
        return Decimal(str(val))
    return ecps.parse_decimal(val)



def _year_numbers_for_version(version_id: int) -> set[int]:
    rows = Year.query.filter(Year.database_version_id == version_id).all()
    return {int(y.number) for y in rows if y.number is not None}


def _database_version_ids_ordered() -> list[int]:
    rows = (
        DatabaseVersion.query.filter(DatabaseVersion.id > 0)
        .order_by(DatabaseVersion.id.asc())
        .all()
    )
    out = [int(r.id) for r in rows]
    if out:
        return out
    cv = get_current_version()
    if cv is not None:
        return [int(cv)]
    return []


def _database_version_ids_for_energy_consumption_import() -> list[int]:
    """Все зарегистрированные версии БД (id > 0) для записи импорта сводки потребления."""
    return _database_version_ids_ordered()


def _find_sheet_title(wb, wanted: str):
    wstrip = wanted.strip()
    for title in wb.sheetnames:
        if title.strip() == wstrip:
            return wb[title]
    low = wstrip.casefold()
    for title in wb.sheetnames:
        if title.strip().casefold() == low:
            return wb[title]
    raise ValueError(f'В книге нет листа «{wanted}».')


def _find_sheet_title_optional(wb, wanted: str):
    """Возвращает лист или None, если в книге нет вкладки с таким именем."""
    wstrip = wanted.strip()
    for title in wb.sheetnames:
        if title.strip() == wstrip:
            return wb[title]
    low = wstrip.casefold()
    for title in wb.sheetnames:
        if title.strip().casefold() == low:
            return wb[title]
    return None


def _sheet_matrix(ws) -> list[tuple[Any, ...]]:
    rows = list(ws.iter_rows(values_only=True))
    return [tuple(r) for r in rows]


def _load_summary_workbook_from_bytes(raw: bytes):
    """Открывает книгу для импорта; ошибки чтения превращает в ValueError с пояснением для пользователя."""
    from io import BytesIO
    from zipfile import BadZipFile

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        return load_workbook(BytesIO(raw), read_only=False, data_only=True)
    except (BadZipFile, InvalidFileException) as e:
        raise ValueError(
            "Не удалось открыть файл как книгу Excel. Нужен формат .xlsx или .xlsm "
            "(не старый .xls). Сохраните документ через «Сохранить как» → «Книга Excel (*.xlsx)»."
        ) from e
    except Exception as e:
        raise ValueError(
            "Не удалось прочитать файл Excel — книга повреждена, неполная или в неподдерживаемом формате. "
            "Откройте её в Excel, сохраните заново как .xlsx и повторите импорт."
        ) from e


def _detect_year_row(matrix: list[tuple[Any, ...]]) -> tuple[int, dict[int, int]]:
    best_row = -1
    best_cols: dict[int, int] = {}
    best_score = 0
    for ri, row in enumerate(matrix):
        cols: dict[int, int] = {}
        seen_years: set[int] = set()
        for ci, val in enumerate(row):
            y = _coerce_cell_to_year_number(val)
            if y is not None:
                if y in seen_years:
                    continue
                seen_years.add(y)
                cols[ci] = y
        if len(cols) > best_score:
            best_score = len(cols)
            best_row = ri
            best_cols = cols
    if best_row < 0 or best_score < 2:
        raise ValueError(
            "Не удалось найти строку с годами в таблице (целые годы или «2016 г.» и т.п.; "
            f"ожидаемый диапазон {_YEAR_HEADER_PLAUSIBLE.start}–{_YEAR_HEADER_PLAUSIBLE.stop - 1}). "
            "При записи учитываются только годы, заведённые в справочнике Years для данной версии БД."
        )
    return best_row, best_cols


def _detect_variant_column(
    matrix: list[tuple[Any, ...]],
    year_row_idx: int,
) -> Optional[int]:
    """Столбец B шаблона: коды вариантов периметра (perimeter-variants)."""
    for ri in range(0, year_row_idx + 1):
        row = matrix[ri] if ri < len(matrix) else ()
        for ci, val in enumerate(row):
            s = _norm_space(_cell_as_text(val)).casefold().replace("ё", "е")
            if any(marker in s for marker in _VARIANT_COLUMN_HEADER_MARKERS):
                return ci
    return None


def _detect_name_column(
    matrix: list[tuple[Any, ...]],
    year_row_idx: int,
    variant_ci: Optional[int] = None,
) -> int:
    needle = "наименование"
    last_hit: Optional[int] = None
    for ri in range(0, year_row_idx + 1):
        row = matrix[ri] if ri < len(matrix) else ()
        for ci, val in enumerate(row):
            if variant_ci is not None and ci <= variant_ci:
                continue
            s = _norm_space(_cell_as_text(val)).casefold()
            if needle in s:
                last_hit = ci
    if last_hit is None:
        raise ValueError('Не найден столбец с заголовком, содержащим «Наименование».')
    return last_hit


def _parse_optional_variant_cell(raw: Any) -> Optional[str]:
    s = _norm_space(_cell_as_text(raw))
    if not s:
        return None
    return normalize_perimeter_variant_code(s, known_only=True)


def _strip_import_variant_suffixes_from_label(label: str) -> str:
    """Базовое имя сущности без суффиксов НТ/ГАЭС/Калининграда из подписи Excel."""
    base, _kal_suffix = _split_kaliningrad_suffix_from_label(
        _normalize_import_label_typos(label)
    )
    changed = True
    while changed:
        changed = False
        for suffix in (
            _GAES_WITH_CHARGE_LABEL_SUFFIX,
            " с ГАЭС",
            " без ГАЭС",
            " без заряда ГАЭС",
            _NT_WITH_TERRITORIES_LABEL_INFIX,
            _NT_LABEL_INFIX,
        ):
            if base.endswith(suffix):
                base = base[: -len(suffix)].strip()
                changed = True
                break
    return _norm_space(base)


def _split_kaliningrad_suffix_from_label(label: str) -> tuple[str, str | None]:
    s = _normalize_import_label_typos(_norm_space(label))
    for suffix in (_KALININGRAD_ES_LABEL_SUFFIX_WITH, " без ЭС Калининградской области"):
        parenthesized = f"({suffix.strip()})"
        if s.endswith(parenthesized):
            base = s[: -len(parenthesized)].strip()
            return (base if base else s), suffix.strip()
        if s.endswith(suffix):
            base = s[: -len(suffix)].strip()
            return (base if base else s), suffix.strip()
    return s, None


def _resolve_row_perimeter_variant_code(
    raw_variant_cell: Any,
    *,
    variant_column_mode: bool,
) -> Optional[str]:
    """Пустая ячейка варианта — без варианта периметра (None), без наследования сверху."""
    if not variant_column_mode:
        return None
    return _parse_optional_variant_cell(raw_variant_cell)


def _entities_with_names(model, version_id: int) -> list[tuple[Any, str]]:
    q = filter_by_explicit_db_version(model.query, model, version_id)
    rows = q.order_by(model.id.asc()).all()
    out: list[tuple[Any, str]] = []
    for r in rows:
        nm = getattr(r, "name", None)
        if nm is None:
            continue
        ns = _norm_space(str(nm))
        if ns:
            out.append((r, ns))
    return out


def _binding_ctx(version_id: int) -> dict[str, list[tuple[Any, str]]]:
    return {
        "ues": _entities_with_names(UnionEnergySystem, version_id),
        "res": _entities_with_names(RegionalEnergySystem, version_id),
        "rd": _entities_with_names(RegionalDistrict, version_id),
        "fd": _entities_with_names(FederalDistrict, version_id),
        "sync_area": _entities_with_names(SynchronousArea, version_id),
        "energy_unit": _entities_with_names(EnergyUnit, version_id),
        "energy_system_type": _entities_with_names(EnergySystemType, version_id),
    }


def _labels_match(excel_label: str, db_name: str) -> bool:
    ex = _norm_entity_label(excel_label)
    db = _norm_entity_label(db_name)
    if not ex or not db:
        return False
    return ex == db


def _pick_match(label: str, pairs: list[tuple[Any, str]]) -> Optional[Any]:
    if not pairs:
        return None
    label_n = _norm_entity_label(label)
    exact = [e for e, n in pairs if _norm_entity_label(n) == label_n]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        exact.sort(key=lambda e: len(getattr(e, "name", "") or ""), reverse=True)
        return exact[0]
    return None


def _pick_exact_match(label: str, pairs: list[tuple[Any, str]]) -> Optional[Any]:
    if not pairs:
        return None
    label_n = _norm_entity_label(label)
    exact = [e for e, n in pairs if _norm_entity_label(n) == label_n]
    if not exact:
        return None
    exact.sort(key=lambda e: len(getattr(e, "name", "") or ""), reverse=True)
    return exact[0]


# Агрегаты верхнего уровня без FK (РФ, ЭЭС России) — вариант из столбца perimeter-variants.
# «ЕЭС России» (EnergySystemType) — отдельно, через справочник типов энергосистемы.
_AGGREGATE_MODEL_BY_BASE_LABEL: dict[str, Type[Any]] = {
    _norm_entity_label(EES_RUSSIA_AGGREGATE_NAME): EesRussiaEnergyConsumptionParameter,
    _norm_entity_label("Россия"): RussiaFederationEnergyConsumptionParameter,
    _norm_entity_label("Россия без НТ"): RussiaFederationEnergyConsumptionParameter,
    _norm_entity_label("Россия с НТ"): RussiaFederationEnergyConsumptionParameter,
}
_EES_UNIFIED_IMPORT_LABEL_CF = _norm_entity_label(EES_UNIFIED_REF_NAME)
_CALCULATED_EES_RUSSIA_BASE_LABEL_CF = _norm_entity_label(EES_RUSSIA_AGGREGATE_NAME)

# Старый шаблон без столбца вариантов: подпись целиком → (модель, код варианта) для агрегатов.
_AGGREGATE_ENERGY_CONSUMPTION_BY_CANONICAL_LABEL: tuple[tuple[str, Type[Any], str | None], ...] = (
    ("Россия", RussiaFederationEnergyConsumptionParameter, CODE_WITHOUT_NT),
    ("Россия без НТ", RussiaFederationEnergyConsumptionParameter, CODE_WITHOUT_NT),
    ("Россия с НТ", RussiaFederationEnergyConsumptionParameter, CODE_WITH_NT),
)

# Старый шаблон: «ЕЭС России» (EnergySystemType) с вариантами ГАЭС в подписи.
_EES_UNIFIED_ENERGY_CONSUMPTION_BY_CANONICAL_LABEL: tuple[tuple[str, str | None], ...] = (
    ("ЕЭС России с НТ с зарядом ГАЭС", CODE_WITH_NT_WITH_GAES),
    ("ЕЭС России без НТ с зарядом ГАЭС", CODE_WITHOUT_NT_WITH_GAES),
)

_IMPORT_SKIPPED_ROW_LABELS = frozenset(
    _norm_entity_label(name)
    for name in (
        "ЭЭС России",
        "ЭЭС России с НТ",
        "ЭЭС России без НТ",
        "Первая синхронная зона",
        "Вторая синхронная зона",
        "Синхронная зона Калининградской области",
    )
)


def _is_ees_unified_import_label(label: str) -> bool:
    """«ЕЭС России» — тип энергосистемы (EnergySystemType), не агрегат ЭЭС России."""
    base = _strip_import_variant_suffixes_from_label(label)
    return _norm_entity_label(base) == _EES_UNIFIED_IMPORT_LABEL_CF


def _is_calculated_ees_russia_import_label(label: str) -> bool:
    """«ЭЭС России» (электроэнергетические системы) — расчёт, не пишется в БД."""
    base = _strip_import_variant_suffixes_from_label(label)
    return _norm_entity_label(base) == _CALCULATED_EES_RUSSIA_BASE_LABEL_CF


def _is_ees_russia_aggregate_import_label(label: str) -> bool:
    """Агрегат «ЭЭС России» с вариантом периметра — импортируется в EesRussia*."""
    base = _strip_import_variant_suffixes_from_label(label)
    return _norm_entity_label(base) == _CALCULATED_EES_RUSSIA_BASE_LABEL_CF


def _resolve_ees_unified_energy_system_type_binding(
    label: str,
    ctx: dict[str, list[tuple[Any, str]]],
    *,
    perimeter_variant_code: str | None = None,
    variant_column_mode: bool = False,
) -> Optional[_ImportBind]:
    """«ЕЭС России» из Excel → ``EnergySystemTypeEnergyConsumptionParameter``."""
    if not _is_ees_unified_import_label(label):
        return None
    est = _pick_exact_match(EES_UNIFIED_REF_NAME, ctx.get("energy_system_type") or [])
    if est is None:
        return None
    pvc = perimeter_variant_code if variant_column_mode else None
    if not variant_column_mode:
        for label_n in _aggregate_import_label_candidates(label):
            for canonical, canonical_pvc in _EES_UNIFIED_ENERGY_CONSUMPTION_BY_CANONICAL_LABEL:
                if label_n == _norm_entity_label(canonical):
                    pvc = canonical_pvc
                    break
    if (
        variant_column_mode
        and pvc is None
        and _aggregate_model_requires_perimeter_variant(EnergySystemTypeEnergyConsumptionParameter)
    ):
        return None
    return _ImportBind(
        EnergySystemTypeEnergyConsumptionParameter,
        _MODEL_FK[EnergySystemTypeEnergyConsumptionParameter],
        int(est.id),
        pvc,
    )


def _label_denotes_ees_unified_gaes_aggregate(label: str) -> bool:
    """Подпись «ЕЭС России … с зарядом ГАЭС» в старом шаблоне без колонки вариантов."""
    label_n = _norm_entity_label(label)
    if "without_gaes" in label_n or "без заряда" in label_n:
        return False
    if not any(
        marker in label_n
        for marker in (
            _norm_entity_label(_GAES_WITH_CHARGE_LABEL_SUFFIX),
            " с зарядом гаэс",
            " с гаэс",
        )
    ):
        return False
    base = _strip_import_variant_suffixes_from_label(label)
    return _is_ees_unified_import_label(base)


def _label_denotes_ees_russia_gaes_aggregate(label: str) -> bool:
    """Подпись «ЭЭС России … с зарядом ГАЭС» в старом шаблоне без колонки вариантов."""
    label_n = _norm_entity_label(label)
    if "without_gaes" in label_n or "без заряда" in label_n:
        return False
    if not any(
        marker in label_n
        for marker in (
            _norm_entity_label(_GAES_WITH_CHARGE_LABEL_SUFFIX),
            " с зарядом гаэс",
            " с гаэс",
        )
    ):
        return False
    base = _strip_import_variant_suffixes_from_label(label)
    return _is_ees_russia_aggregate_import_label(base)


def _is_ees_russia_gaes_aggregate_import_row(
    label: str,
    perimeter_variant_code: str | None,
) -> bool:
    """Строка «ЕЭС России … с зарядом ГАЭС» (EnergySystemType, с/без НТ)."""
    code = str(perimeter_variant_code or "").strip()
    if code:
        if "with_gaes" not in code or "without_gaes" in code:
            return False
        return _is_ees_unified_import_label(label)
    return _label_denotes_ees_unified_gaes_aggregate(label)


def _import_row_year_numeric_pairs(
    row: tuple[Any, ...],
    year_cols: dict[int, int],
    *,
    label: str,
    row_variant: str | None,
    variant_column_mode: bool,
) -> list[tuple[int, Decimal]]:
    """Пары (год, значение) для строки листа.

    Для «ЕЭС России … с зарядом ГАЭС» (``with_nt_with_gaes`` / ``without_nt_with_gaes``)
    числа в шаблоне часто стоят правее, чем у «Россия» / «ЕЭС России» без НТ (пустые
    ячейки в начале блока годов). Разреженные строки привязываем к последним N годам
    сетки; плотный блок со сдвигом — сдвигаем влево на ширину «пустого хвоста».
    """
    del variant_column_mode  # выравнивание одинаково для обоих форматов шаблона
    year_col_indices = sorted(year_cols.keys())
    if not year_col_indices:
        return []

    parsed: list[tuple[int, Decimal]] = []
    for ci in year_col_indices:
        if ci >= len(row):
            continue
        num = _parse_numeric_cell(row[ci])
        if num is None:
            continue
        parsed.append((ci, num))
    if not parsed:
        return []

    first_year_ci = year_col_indices[0]
    first_val_ci = parsed[0][0]
    last_val_ci = parsed[-1][0]
    last_year_ci = year_col_indices[-1]
    values = [val for _, val in parsed]
    if (
        first_val_ci > first_year_ci
        and _is_ees_russia_gaes_aggregate_import_row(label, row_variant)
    ):
        gap = first_val_ci - first_year_ci
        if len(values) == len(year_col_indices):
            shifted: list[tuple[int, Decimal]] = []
            for ci, val in parsed:
                target_ci = ci - gap
                if target_ci in year_cols:
                    shifted.append((year_cols[target_ci], val))
            if len(shifted) == len(values):
                return shifted
        # Разреженные строки у правого края сетки — последние N лет; иначе год берём
        # из фактической колонки (шаблон «Для работы в АРМе»: факт 2024–2025, прогноз пуст).
        if len(values) < len(year_col_indices) and last_val_ci >= last_year_ci:
            target_cols = year_col_indices[-len(values) :]
            return [(year_cols[ci], val) for ci, val in zip(target_cols, values)]

    return [(year_cols[ci], val) for ci, val in parsed]


def _import_row_has_leading_empty_year_cells(
    row: tuple[Any, ...],
    year_cols: dict[int, int],
) -> bool:
    year_col_indices = sorted(year_cols.keys())
    if not year_col_indices:
        return False
    first_year_ci = year_col_indices[0]
    for ci in year_col_indices:
        if ci >= len(row):
            continue
        if _parse_numeric_cell(row[ci]) is not None:
            return ci > first_year_ci
    return False


def _ees_russia_gaes_import_perimeter_variant_from_excel(
    label: str,
    perimeter_variant_code: str | None,
    *,
    row: tuple[Any, ...],
    year_cols: dict[int, int],
) -> str | None:
    """В шаблоне «млн. кВтч» строка «ЕЭС России … с зарядом ГАЭС» иногда помечена ``without_nt`` / ``with_nt``."""
    if not _is_ees_unified_import_label(
        _strip_import_variant_suffixes_from_label(label)
    ):
        return perimeter_variant_code
    if not _import_row_has_leading_empty_year_cells(row, year_cols):
        return perimeter_variant_code
    code = str(perimeter_variant_code or "").strip()
    if code == CODE_WITHOUT_NT:
        return CODE_WITHOUT_NT_WITH_GAES
    if code == CODE_WITH_NT:
        return CODE_WITH_NT_WITH_GAES
    return perimeter_variant_code


def _is_kaliningrad_sync_zone_import_label(label: str) -> bool:
    """«Синхронная зона Калининградской области» — расчётная строка сводки, не из Excel."""
    label_n = _norm_entity_label(label)
    return label_n.startswith("синхронная зона") and "калин" in label_n


def _is_import_skipped_row_label(label: str) -> bool:
    """Строки, исключённые из импорта (устаревший шаблон Excel)."""
    if _is_calculated_ees_russia_import_label(label):
        return True
    if _is_kaliningrad_sync_zone_import_label(label):
        return True
    label_n = _norm_entity_label(label)
    if label_n in _IMPORT_SKIPPED_ROW_LABELS:
        return True
    if "вторая синхронная" in label_n:
        return True
    return False


def _aggregate_import_label_candidates(label: str) -> tuple[str, ...]:
    """Нормализованные подписи для агрегатов: Excel может содержать суффикс «… с зарядом ГАЭС»."""
    label_n = _norm_entity_label(label)
    gaes_tail = _norm_entity_label(_GAES_WITH_CHARGE_LABEL_SUFFIX)
    if label_n.endswith(gaes_tail):
        core = _norm_space(label_n[: -len(gaes_tail)])
        if core and core != label_n:
            return (label_n, core)
    return (label_n,)


def _aggregate_model_requires_perimeter_variant(model_cls: Type[Any]) -> bool:
    ctx = perimeter_entity_context_for_model(model_cls.__name__)
    if ctx is None:
        return False
    return bool(perimeter_variant_codes_for_entity(*ctx))


def _normalize_russia_federation_import_perimeter_variant(
    label: str,
    perimeter_variant_code: str | None,
) -> str | None:
    """Вариант периметра для ``RussiaFederationEnergyConsumptionParameter`` при импорте сводки.

    На листе «млн. кВт.ч» строка «Россия» часто помечена ``with_nt`` в колонке perimeter-variants
    (как в дереве сводки), но в БД базовые значения «млн кВтч» хранятся с ``without_nt``
    (страница «Россия без НТ»). Явная подпись «Россия с НТ» → ``with_nt``.
    """
    label_n = _norm_entity_label(_normalize_import_label_typos(label))
    if label_n == _norm_entity_label("Россия с НТ"):
        return CODE_WITH_NT
    if label_n in (
        _norm_entity_label("Россия"),
        _norm_entity_label("Россия без НТ"),
    ):
        return CODE_WITHOUT_NT
    return perimeter_variant_code


def _resolve_aggregate_energy_consumption_binding(
    label: str,
    *,
    perimeter_variant_code: str | None = None,
    variant_column_mode: bool = False,
) -> Optional[_ImportBind]:
    """Строки-агрегаты в колонке «Наименование» (РФ и ЭЭС), не из справочников."""
    if variant_column_mode:
        base_label = _strip_import_variant_suffixes_from_label(label)
        model_cls = _AGGREGATE_MODEL_BY_BASE_LABEL.get(_norm_entity_label(base_label))
        if model_cls is None:
            model_cls = _AGGREGATE_MODEL_BY_BASE_LABEL.get(_norm_entity_label(label))
        if model_cls is not None:
            if model_cls is RussiaFederationEnergyConsumptionParameter:
                perimeter_variant_code = _normalize_russia_federation_import_perimeter_variant(
                    label, perimeter_variant_code
                )
            if (
                perimeter_variant_code is None
                and _aggregate_model_requires_perimeter_variant(model_cls)
            ):
                # Пустая ячейка perimeter-variants: расчётные агрегаты (ЭЭС России и т.п.)
                # заполняются постобработкой формул после импорта.
                return None
            return _ImportBind(
                model_cls,
                _FK_AGGREGATE_NO_PARENT,
                0,
                perimeter_variant_code,
            )
        return None
    for label_n in _aggregate_import_label_candidates(label):
        for canonical, model_cls, pvc in _AGGREGATE_ENERGY_CONSUMPTION_BY_CANONICAL_LABEL:
            if label_n == _norm_entity_label(canonical):
                return _ImportBind(model_cls, _FK_AGGREGATE_NO_PARENT, 0, pvc)
    return None


def _is_south_ues_import_name(name: str | None) -> bool:
    n = _norm_entity_label(name or "")
    return n == "оэс юга" or (n.startswith("оэс ") and "юга" in n)


def _south_union_energy_system_from_ctx(
    ctx: dict[str, list[tuple[Any, str]]],
) -> tuple[Any, str] | None:
    for ues, name in ctx.get("ues") or []:
        if _is_south_ues_import_name(name):
            return ues, name
    return None


def _south_ues_import_perimeter_variant_from_excel(
    perimeter_variant_code: str | None,
) -> str | None:
    """В шаблоне «млн. кВтч» строка «ОЭС Юга без НТ с зарядом ГАЭС» иногда помечена кодом with_nt_without_gaes."""
    if perimeter_variant_code == CODE_WITH_NT_WITHOUT_GAES:
        return CODE_WITHOUT_NT_WITH_GAES
    return perimeter_variant_code


def _south_ues_gaes_import_label_map(south_name: str) -> dict[str, str]:
    """Нормализованная подпись Excel → ``perimeter_variant_code`` для ОЭС Юга с ГАЭС."""
    base = _norm_space(south_name)
    if not base:
        return {}
    mapping: dict[str, str] = {}
    with_nt_infixes = (_NT_WITH_TERRITORIES_LABEL_INFIX, " c НТ")
    for gaes in _GAES_SHORT_LABEL_SUFFIXES:
        for nt in with_nt_infixes:
            mapping[_norm_entity_label(f"{base}{nt}{gaes}")] = CODE_WITH_NT_WITH_GAES
        mapping[_norm_entity_label(f"{base}{_NT_LABEL_INFIX}{gaes}")] = (
            CODE_WITHOUT_NT_WITH_GAES
        )
    return mapping


def _resolve_south_ues_summary_import_binding(
    label: str,
    ctx: dict[str, list[tuple[Any, str]]],
    *,
    perimeter_variant_code: str | None = None,
    variant_column_mode: bool = False,
) -> Optional[_ImportBind]:
    """Сводная таблица: ОЭС Юга → ``UnionEnergySystemEnergyConsumptionParameter``."""
    pair = _south_union_energy_system_from_ctx(ctx)
    if pair is None:
        return None
    ues, ues_name = pair
    if variant_column_mode:
        match_label = _strip_import_variant_suffixes_from_label(label)
        if _norm_entity_label(match_label) != _norm_entity_label(ues_name):
            return None
        pvc = _south_ues_import_perimeter_variant_from_excel(perimeter_variant_code)
        return _ImportBind(
            UnionEnergySystemEnergyConsumptionParameter,
            _MODEL_FK[UnionEnergySystemEnergyConsumptionParameter],
            int(ues.id),
            pvc,
        )
    label_n = _norm_entity_label(label)
    pvc = _south_ues_gaes_import_label_map(ues_name).get(label_n)
    if pvc is None:
        return None
    return _ImportBind(
        UnionEnergySystemEnergyConsumptionParameter,
        _MODEL_FK[UnionEnergySystemEnergyConsumptionParameter],
        int(ues.id),
        pvc,
    )


def _first_synchronous_area_for_version(
    ctx: dict[str, list[tuple[Any, str]]],
) -> tuple[Any, str] | None:
    pairs = list(ctx.get("sync_area") or [])
    if not pairs:
        return None
    from app.common.services.get_services.energy_systems.synchronous_area_get_services import (
        synchronous_area_display_order_sort_key,
    )

    pairs.sort(key=lambda p: synchronous_area_display_order_sort_key(p[0]))
    return pairs[0]


def _first_sa_import_label_variants(sa_name: str) -> set[str]:
    """Единственная поддерживаемая подпись первой СЗ на листе «млн. кВт.ч» сводной таблицы."""
    base = _norm_space(sa_name)
    if not base:
        return set()
    with_paren = (
        f"{base}{_NT_LABEL_INFIX}{_GAES_WITH_CHARGE_LABEL_SUFFIX} "
        f"({_KALININGRAD_ES_LABEL_SUFFIX_WITH.strip()})"
    )
    without_paren = (
        f"{base}{_NT_LABEL_INFIX}{_GAES_WITH_CHARGE_LABEL_SUFFIX}"
        f"{_KALININGRAD_ES_LABEL_SUFFIX_WITH}"
    )
    return {
        _norm_entity_label(with_paren),
        _norm_entity_label(without_paren),
    }


def _resolve_first_sa_summary_import_binding(
    label: str,
    ctx: dict[str, list[tuple[Any, str]]],
    *,
    perimeter_variant_code: str | None = None,
    variant_column_mode: bool = False,
) -> Optional[_ImportBind]:
    """Сводная таблица: первая СЗ → ``SynchronousAreaEnergyConsumptionParameter``."""
    pair = _first_synchronous_area_for_version(ctx)
    if pair is None:
        return None
    sa, sa_name = pair
    if variant_column_mode:
        match_label = _strip_import_variant_suffixes_from_label(label)
        if _norm_entity_label(match_label) != _norm_entity_label(sa_name):
            return None
        return _ImportBind(
            SynchronousAreaEnergyConsumptionParameter,
            _MODEL_FK[SynchronousAreaEnergyConsumptionParameter],
            int(sa.id),
            perimeter_variant_code,
        )
    label_n = _norm_entity_label(label)
    if label_n not in _first_sa_import_label_variants(sa_name):
        return None
    return _ImportBind(
        SynchronousAreaEnergyConsumptionParameter,
        _MODEL_FK[SynchronousAreaEnergyConsumptionParameter],
        int(sa.id),
        CODE_WITHOUT_NT_WITH_GAES_WITH_KALININGRAD_ES,
    )


def _is_unresolved_first_sa_gaes_kaliningrad_import_label(
    label: str,
    ctx: dict[str, list[tuple[Any, str]]],
) -> bool:
    """Подпись похожа на импорт первой СЗ с ГАЭС, но не сопоставилась — не писать в базовую СЗ."""
    if _resolve_first_sa_summary_import_binding(label, ctx) is not None:
        return False
    label_n = _norm_entity_label(label)
    return "зарядом гаэс" in label_n and "калин" in label_n and "без нт" in label_n


def _resolve_row_binding(
    label: str,
    ctx: dict[str, list[tuple[Any, str]]],
    *,
    perimeter_variant_code: str | None = None,
    variant_column_mode: bool = False,
) -> Optional[_ImportBind]:
    if _is_kaliningrad_sync_zone_import_label(label):
        return None
    if not variant_column_mode and _is_import_skipped_row_label(label):
        return None

    match_label = (
        _strip_import_variant_suffixes_from_label(label)
        if variant_column_mode
        else label
    )

    unified_bind = _resolve_ees_unified_energy_system_type_binding(
        label,
        ctx,
        perimeter_variant_code=perimeter_variant_code,
        variant_column_mode=variant_column_mode,
    )
    if unified_bind is not None:
        return unified_bind

    agg_bind = _resolve_aggregate_energy_consumption_binding(
        label,
        perimeter_variant_code=perimeter_variant_code,
        variant_column_mode=variant_column_mode,
    )
    if agg_bind is not None:
        return agg_bind

    south_ues_bind = _resolve_south_ues_summary_import_binding(
        label,
        ctx,
        perimeter_variant_code=perimeter_variant_code,
        variant_column_mode=variant_column_mode,
    )
    if south_ues_bind is not None:
        return south_ues_bind

    first_sa_bind = _resolve_first_sa_summary_import_binding(
        label,
        ctx,
        perimeter_variant_code=perimeter_variant_code,
        variant_column_mode=variant_column_mode,
    )
    if first_sa_bind is not None:
        return first_sa_bind

    ues = _pick_exact_match(match_label, ctx["ues"])
    if ues is not None:
        return _ImportBind(
            UnionEnergySystemEnergyConsumptionParameter,
            _MODEL_FK[UnionEnergySystemEnergyConsumptionParameter],
            int(ues.id),
            perimeter_variant_code if variant_column_mode else None,
        )
    res = _pick_exact_match(match_label, ctx["res"])
    if res is not None:
        return _ImportBind(
            RegionalEnergySystemEnergyConsumptionParameter,
            _MODEL_FK[RegionalEnergySystemEnergyConsumptionParameter],
            int(res.id),
            perimeter_variant_code if variant_column_mode else None,
        )
    rd = _pick_exact_match(match_label, ctx["rd"])
    if rd is not None:
        return _ImportBind(
            RegionalDistrictEnergyConsumptionParameter,
            _MODEL_FK[RegionalDistrictEnergyConsumptionParameter],
            int(rd.id),
            perimeter_variant_code if variant_column_mode else None,
        )
    fd = _pick_exact_match(match_label, ctx["fd"])
    if fd is not None:
        return _ImportBind(
            FederalDistrictEnergyConsumptionParameter,
            _MODEL_FK[FederalDistrictEnergyConsumptionParameter],
            int(fd.id),
            perimeter_variant_code if variant_column_mode else None,
        )

    sa = _pick_exact_match(match_label, ctx["sync_area"])
    if sa is not None and _is_unresolved_first_sa_gaes_kaliningrad_import_label(label, ctx):
        sa = None
    if sa is not None:
        return _ImportBind(
            SynchronousAreaEnergyConsumptionParameter,
            _MODEL_FK[SynchronousAreaEnergyConsumptionParameter],
            int(sa.id),
            perimeter_variant_code if variant_column_mode else None,
        )
    eu = _pick_exact_match(match_label, ctx["energy_unit"])
    if eu is not None:
        return _ImportBind(
            EnergyUnitEnergyConsumptionParameter,
            _MODEL_FK[EnergyUnitEnergyConsumptionParameter],
            int(eu.id),
            perimeter_variant_code if variant_column_mode else None,
        )

    if variant_column_mode:
        return None

    ues = _pick_match(label, ctx["ues"])
    if ues is not None:
        return _ImportBind(
            UnionEnergySystemEnergyConsumptionParameter,
            _MODEL_FK[UnionEnergySystemEnergyConsumptionParameter],
            int(ues.id),
        )
    res = _pick_match(label, ctx["res"])
    if res is not None:
        return _ImportBind(
            RegionalEnergySystemEnergyConsumptionParameter,
            _MODEL_FK[RegionalEnergySystemEnergyConsumptionParameter],
            int(res.id),
        )
    rd = _pick_match(label, ctx["rd"])
    if rd is not None:
        return _ImportBind(
            RegionalDistrictEnergyConsumptionParameter,
            _MODEL_FK[RegionalDistrictEnergyConsumptionParameter],
            int(rd.id),
        )
    fd = _pick_match(label, ctx["fd"])
    if fd is not None:
        return _ImportBind(
            FederalDistrictEnergyConsumptionParameter,
            _MODEL_FK[FederalDistrictEnergyConsumptionParameter],
            int(fd.id),
        )

    sa = _pick_match(label, ctx["sync_area"])
    if sa is not None and _is_unresolved_first_sa_gaes_kaliningrad_import_label(label, ctx):
        sa = None
    if sa is not None:
        return _ImportBind(
            SynchronousAreaEnergyConsumptionParameter,
            _MODEL_FK[SynchronousAreaEnergyConsumptionParameter],
            int(sa.id),
        )
    eu = _pick_match(label, ctx["energy_unit"])
    if eu is not None:
        return _ImportBind(
            EnergyUnitEnergyConsumptionParameter,
            _MODEL_FK[EnergyUnitEnergyConsumptionParameter],
            int(eu.id),
        )
    return None


# Лист-заглушка «млн. кВт.ч» (несколько строк без сетки годов) — не считаем ошибкой при разборе.
_OPTIONAL_GRID_MAX_ROWS = 6

_SCREEN_EXPORT_ENTITY_HEADER = "энергосистема"
_SCREEN_EXPORT_PARAM_HEADER = "наименование параметров"


def _classify_screen_export_parameter_label(param_label: str) -> str | None:
    """«mln» / «sipr» для строк отчётного листа; формулы и приросты пропускаем."""
    s = _norm_space(param_label).casefold().replace("ё", "е")
    if not s or "прирост" in s or "темп" in s:
        return None
    if "(сипр)" in s and "млн" in s:
        return "sipr"
    if "потребление" in s and "млн" in s and "(сипр)" not in s:
        return "mln"
    return None


def _explicit_perimeter_variant_placeholder_rows(
    acc: defaultdict[tuple[str, int, str | None], Decimal],
) -> set[tuple[str, str]]:
    """Строки Excel с явным кодом варианта, где все годы нулевые (заглушка под другую строку)."""
    by_label_pvc: dict[tuple[str, str], list[Decimal]] = defaultdict(list)
    for (label, _year_n, pvc), total in acc.items():
        if pvc is not None:
            by_label_pvc[(label, pvc)].append(total)
    return {
        key
        for key, totals in by_label_pvc.items()
        if totals and all(val == 0 for val in totals)
    }


def _merge_perimeter_variant_duplicate_year_values(
    acc: defaultdict[tuple[str, int, str | None], Decimal],
) -> defaultdict[tuple[str, int, str | None], Decimal]:
    """Убирает только нулевые заглушки; разные варианты одной сущности сохраняются отдельно.

  Пустая ячейка perimeter-variants → ``None`` (NULL в БД). Для одной сущности могут быть
  параллельные строки без варианта и с ``o1`` (например, «ЭС Камчатского края») — обе
  записываются в свои варианты периметра.

  Строка с явным кодом варианта, нулевая по всем годам, пропускается целиком. Отдельные
  нулевые ячейки внутри не-заглушки также не импортируются.
    """
    placeholder_rows = _explicit_perimeter_variant_placeholder_rows(acc)
    merged: defaultdict[tuple[str, int, str | None], Decimal] = defaultdict(
        lambda: Decimal("0")
    )
    for (label, year_n, pvc), total in acc.items():
        if pvc is not None and (label, pvc) in placeholder_rows:
            continue
        if pvc is not None and total == 0:
            continue
        merged[(label, year_n, pvc)] += total
    return merged


def _aggregate_from_screen_export_matrix(
    matrix: list[tuple[Any, ...]],
) -> tuple[
    defaultdict[tuple[str, int, str | None], Decimal],
    defaultdict[tuple[str, int, str | None], Decimal],
]:
    """Разбор листа «Экспорт в Excel» (колонки Энергосистема / Наименование параметров / годы)."""
    acc_mln: defaultdict[tuple[str, int, str | None], Decimal] = defaultdict(
        lambda: Decimal("0")
    )
    acc_sipr: defaultdict[tuple[str, int, str | None], Decimal] = defaultdict(
        lambda: Decimal("0")
    )
    if not matrix:
        return acc_mln, acc_sipr

    header_row_idx: int | None = None
    entity_ci = 0
    param_ci = 1
    for ri, row in enumerate(matrix[:8]):
        if len(row) < 2:
            continue
        c0 = _norm_space(_cell_as_text(row[0])).casefold().replace("ё", "е")
        c1 = _norm_space(_cell_as_text(row[1])).casefold().replace("ё", "е")
        if _SCREEN_EXPORT_ENTITY_HEADER in c0 and _SCREEN_EXPORT_PARAM_HEADER in c1:
            header_row_idx = ri
            break
    if header_row_idx is None:
        return acc_mln, acc_sipr

    yr_row, year_cols = _detect_year_row(matrix)
    if yr_row != header_row_idx:
        year_cols = {}
        for ci, val in enumerate(matrix[header_row_idx]):
            y = _coerce_cell_to_year_number(val)
            if y is not None:
                year_cols[ci] = y
    if len(year_cols) < 2:
        return acc_mln, acc_sipr

    current_entity = ""
    for ri in range(header_row_idx + 1, len(matrix)):
        row = matrix[ri]
        if entity_ci < len(row):
            ent_raw = _norm_space(_cell_as_text(row[entity_ci]))
            if ent_raw:
                current_entity = _normalize_import_label_typos(ent_raw.strip())
        if not current_entity or param_ci >= len(row):
            continue
        kind = _classify_screen_export_parameter_label(
            _cell_as_text(row[param_ci])
        )
        if kind is None:
            continue
        target = acc_mln if kind == "mln" else acc_sipr
        for year_n, num in _import_row_year_numeric_pairs(
            row,
            year_cols,
            label=current_entity,
            row_variant=None,
            variant_column_mode=False,
        ):
            target[(current_entity, year_n, None)] += num
    return acc_mln, acc_sipr


def _aggregate_sheet_labels(
    matrix: list[tuple[Any, ...]],
    *,
    sheet_title: str = "",
    allow_empty_short_sheet_without_year_grid: bool = False,
) -> tuple[defaultdict[tuple[str, int, str | None], Decimal], bool]:
    """Ключ: (наименование, год, код варианта периметра) → сумма по строкам листа.

    Если allow_empty_short_sheet_without_year_grid=True и в коротком листе нет строки с годами
    (типичный пустой лист-заглушка «млн. кВт.ч»), возвращается пустая сводка без ошибки.
    """
    acc: defaultdict[tuple[str, int, str | None], Decimal] = defaultdict(
        lambda: Decimal("0")
    )
    aggregate_seen_values: dict[tuple[str, int, str | None], Decimal] = {}
    if not matrix:
        return acc, False
    try:
        yr_row, year_cols = _detect_year_row(matrix)
        variant_ci = _detect_variant_column(matrix, yr_row)
        variant_column_mode = variant_ci is not None
        name_ci = _detect_name_column(matrix, yr_row, variant_ci)
    except ValueError:
        if allow_empty_short_sheet_without_year_grid and len(matrix) <= _OPTIONAL_GRID_MAX_ROWS:
            return acc, False
        raise

    for ri in range(yr_row + 1, len(matrix)):
        row = matrix[ri]
        if name_ci >= len(row):
            continue
        raw_name = row[name_ci]
        label = _normalize_import_label_typos(_norm_space(_cell_as_text(raw_name)))
        if not label:
            continue

        raw_variant_cell = (
            row[variant_ci]
            if variant_column_mode and variant_ci is not None and variant_ci < len(row)
            else None
        )
        if variant_column_mode:
            row_variant = _resolve_row_perimeter_variant_code(
                raw_variant_cell,
                variant_column_mode=True,
            )
            row_variant = _ees_russia_gaes_import_perimeter_variant_from_excel(
                label,
                row_variant,
                row=row,
                year_cols=year_cols,
            )
            base_label = _strip_import_variant_suffixes_from_label(label)
            if (
                _AGGREGATE_MODEL_BY_BASE_LABEL.get(_norm_entity_label(base_label))
                is RussiaFederationEnergyConsumptionParameter
            ):
                row_variant = _normalize_russia_federation_import_perimeter_variant(
                    label, row_variant
                )
        else:
            row_variant = None

        if _is_kaliningrad_sync_zone_import_label(label):
            continue
        if _is_calculated_ees_russia_import_label(label):
            if not (
                variant_column_mode
                and row_variant
                and _is_ees_russia_aggregate_import_label(label)
            ):
                continue
        elif not variant_column_mode and _is_import_skipped_row_label(label):
            continue

        aggregate_bind = _resolve_aggregate_energy_consumption_binding(
            label,
            perimeter_variant_code=row_variant,
            variant_column_mode=variant_column_mode,
        )
        for year_n, num in _import_row_year_numeric_pairs(
            row,
            year_cols,
            label=label,
            row_variant=row_variant,
            variant_column_mode=variant_column_mode,
        ):
            if (
                aggregate_bind is not None
                and aggregate_bind.fk_column == _FK_AGGREGATE_NO_PARENT
            ):
                agg_key = (_norm_entity_label(label), int(year_n), row_variant)
                prev = aggregate_seen_values.get(agg_key)
                if prev is None:
                    aggregate_seen_values[agg_key] = num
                    acc[(label, year_n, row_variant)] += num
                    continue
                if prev == num:
                    continue
                sheet_hint = f" на листе «{sheet_title}»" if sheet_title else ""
                variant_hint = f", вариант {row_variant!r}" if row_variant else ""
                raise ValueError(
                    f"Строка «{label}»{variant_hint}{sheet_hint} встречается несколько раз "
                    f"с разными значениями за {year_n} год. Импорт не может однозначно "
                    "выбрать правильное значение."
                )
            acc[(label, year_n, row_variant)] += num
    if variant_column_mode and acc:
        acc = _merge_perimeter_variant_duplicate_year_values(acc)
    return acc, variant_column_mode


def _single_regional_district_id_for_res(
    regional_energy_system_id: int,
    *,
    database_version_id: int,
) -> Optional[int]:
    """
    Если в связях региональная ЭС привязана ровно к одному субъекту РФ и этот субъект
    есть в указанной версии справочника — вернуть его id, иначе None.
    """
    stmt = select(regional_district_regional_energy_system.c.regional_district_id).where(
        regional_district_regional_energy_system.c.regional_energy_system_id
        == regional_energy_system_id
    )
    rows = db.session.execute(stmt).all()
    rd_ids = {int(r[0]) for r in rows if r[0] is not None}
    if len(rd_ids) != 1:
        return None
    rd_id = next(iter(rd_ids))
    rq = RegionalDistrict.query.filter(RegionalDistrict.id == rd_id)
    rq = filter_by_explicit_db_version(rq, RegionalDistrict, database_version_id)
    return rd_id if rq.first() is not None else None


def _import_perimeter_variant_without_entity_bindings(
    model_cls: Type[Any],
    fk_column: str,
    parent_id: int,
    perimeter_variant_code: str | None,
) -> str | None:
    """Для сущностей без привязок в каталоге код из Excel не записывается (NULL в БД)."""
    if perimeter_variant_code is None:
        return None
    if not model_supports_perimeter_variant(model_cls):
        return perimeter_variant_code
    entity_ctx = perimeter_entity_context_for_model(
        model_cls.__name__,
        parent_fk_column=fk_column,
        parent_id=int(parent_id),
    )
    if entity_ctx is None:
        return perimeter_variant_code
    allowed = perimeter_variant_codes_for_entity(*entity_ctx)
    if not allowed:
        return None
    if perimeter_variant_code not in allowed:
        return None
    return perimeter_variant_code


def _regional_district_accepts_import_perimeter_variant(
    rd_id: int,
    perimeter_variant_code: str | None,
) -> bool:
    """Можно ли записать код варианта в параметры субъекта РФ (в т.ч. при зеркале с РЭС).

    Пустой вариант (NULL в БД) допустим всегда: при импорте Excel он не валидируется
  против каталога привязок. Явный код — только если он привязан к этому субъекту.
    """
    if perimeter_variant_code is None:
        return True
    ctx = perimeter_entity_context_for_model(
        RegionalDistrictEnergyConsumptionParameter.__name__,
        parent_fk_column=_MODEL_FK[RegionalDistrictEnergyConsumptionParameter],
        parent_id=int(rd_id),
    )
    if ctx is None:
        return False
    entity_kind, entity_name = ctx
    allowed = perimeter_variant_codes_for_entity(entity_kind, entity_name)
    return perimeter_variant_code in allowed


def _mirror_res_accumulator_rows_to_single_district_rd(
    acc: defaultdict[tuple[Type[Any], str, int, int, str | None], Decimal],
    *,
    database_version_id: int,
    explicit_rd_subject_year_pairs: frozenset[tuple[int, int]],
) -> None:
    """Дублирует суммы импорта по РЭС в параметры единственного субъекта этой РЭС (если он один).

    Если в файле по этому субъекту за тот же год уже есть строка, сопоставленная именно как
    субъект РФ, зеркало не добавляют — без этого значения суммировались бы дважды.
    """
    res_model = RegionalEnergySystemEnergyConsumptionParameter
    fk_res = _MODEL_FK[res_model]
    fk_rd = _MODEL_FK[RegionalDistrictEnergyConsumptionParameter]
    rd_model = RegionalDistrictEnergyConsumptionParameter
    for (model_cls, fk_column, parent_id, year_n, pvc), total in list(acc.items()):
        if model_cls is not res_model or fk_column != fk_res:
            continue
        rd_id = _single_regional_district_id_for_res(
            int(parent_id),
            database_version_id=database_version_id,
        )
        if rd_id is None:
            continue
        if (rd_id, year_n) in explicit_rd_subject_year_pairs:
            continue
        if not _regional_district_accepts_import_perimeter_variant(rd_id, pvc):
            continue
        key_rd = (rd_model, fk_rd, rd_id, year_n, pvc)
        acc[key_rd] += total


def _collapse_labels_to_bind_keys(
    acc_labels: defaultdict[tuple[str, int, str | None], Decimal],
    *,
    ctx: dict[str, list[tuple[Any, str]]],
    years_ok: set[int],
    variant_column_mode: bool,
) -> tuple[
    defaultdict[tuple[Type[Any], str, int, int, str | None], Decimal],
    frozenset[tuple[int, int]],
]:
    """Отбор годов, входящих в справочник Years данной версии; привязка к объектам версии.

    Второй элемент — множество (id_subj_rf, год), для которых в файле уже есть сумма как по
    субъекту РФ (без дубля зеркалом с РЭС).

    Если у сущности нет привязок вариантов периметра, явный код из Excel (например ``o1``)
    сбрасывается в ``NULL``. Параллельная строка без кода и строка с таким кодом не
    суммируются: иначе значения удваиваются (шаблон «Для работы в АРМе» для центральных
    энергорайонов Магадана / Камчатки / Сахалина).
    """
    acc: defaultdict[tuple[Type[Any], str, int, int, str | None], Decimal] = defaultdict(
        lambda: Decimal("0")
    )
    # Ключи, заполненные только после сброса явного PVC → NULL (ещё нет «родной» NULL-строки).
    cleared_pvc_only_keys: set[
        tuple[Type[Any], str, int, int, str | None]
    ] = set()
    explicit_rd_subject_years: set[tuple[int, int]] = set()
    rd_model = RegionalDistrictEnergyConsumptionParameter
    for (label, year_n, pvc), total in acc_labels.items():
        if year_n not in years_ok:
            continue
        bind = _resolve_row_binding(
            label,
            ctx,
            perimeter_variant_code=pvc,
            variant_column_mode=variant_column_mode,
        )
        if bind is None:
            continue
        effective_pvc = bind.perimeter_variant_code
        if (
            effective_pvc is None
            and pvc is None
            and not variant_column_mode
            and model_supports_perimeter_variant(bind.demand_model)
        ):
            entity_ctx = perimeter_entity_context_for_model(
                bind.demand_model.__name__,
                parent_fk_column=bind.fk_column,
                parent_id=int(bind.parent_id),
            )
            if entity_ctx is not None:
                allowed = perimeter_variant_codes_for_entity(*entity_ctx)
                if len(allowed) == 1:
                    effective_pvc = allowed[0]
                    bind = _ImportBind(
                        bind.demand_model,
                        bind.fk_column,
                        bind.parent_id,
                        effective_pvc,
                    )
        pvc_cleared_to_null = False
        if model_supports_perimeter_variant(bind.demand_model):
            cleared_pvc = _import_perimeter_variant_without_entity_bindings(
                bind.demand_model,
                bind.fk_column,
                int(bind.parent_id),
                bind.perimeter_variant_code,
            )
            if cleared_pvc != bind.perimeter_variant_code:
                pvc_cleared_to_null = True
                effective_pvc = cleared_pvc
                bind = _ImportBind(
                    bind.demand_model,
                    bind.fk_column,
                    bind.parent_id,
                    cleared_pvc,
                )
        key = (
            bind.demand_model,
            bind.fk_column,
            bind.parent_id,
            year_n,
            effective_pvc,
        )
        if pvc_cleared_to_null:
            # Явный вариант сброшен в NULL: не суммировать с уже записанной NULL-строкой.
            if key in acc:
                continue
            acc[key] += total
            cleared_pvc_only_keys.add(key)
        elif key in cleared_pvc_only_keys:
            # «Родная» NULL-строка важнее сброшенного o1 — подменяем, не складываем.
            acc[key] = total
            cleared_pvc_only_keys.discard(key)
        else:
            acc[key] += total
        if bind.demand_model is rd_model:
            explicit_rd_subject_years.add((int(bind.parent_id), int(year_n)))
    return acc, frozenset(explicit_rd_subject_years)


def _dedupe_aggregate_energy_consumption_rows_for_year(
    model: Type[Any],
    year_n: int,
    version_id: int,
    *,
    perimeter_variant_code: str | None = None,
) -> None:
    """Standalone-строки дедуплицируются общей логикой, затем выбирается нужный год."""
    ecps.dedupe_duplicate_energy_consumption_rows_without_parent(
        model,
        database_version_id=version_id,
        year_n=year_n,
        perimeter_variant_code=perimeter_variant_code,
    )
    db.session.flush()


def _find_existing_parameter_row(
    model: Type[Any],
    fk_column: str,
    parent_id: int,
    year_n: int,
    version_id: int,
    *,
    perimeter_variant_code: str | None = None,
):
    from app.common.perimeter_variant.registry import filter_query_by_perimeter_variant

    if fk_column == _FK_AGGREGATE_NO_PARENT:
        q = model.query.filter(model.year_number == year_n)
        q = filter_by_explicit_db_version(q, model, version_id)
        q = filter_query_by_perimeter_variant(q, model, perimeter_variant_code)
        return q.first()
    q = model.query.filter(
        getattr(model, fk_column) == parent_id,
        model.year_number == year_n,
    )
    q = filter_by_explicit_db_version(q, model, version_id)
    q = filter_query_by_perimeter_variant(q, model, perimeter_variant_code)
    return q.first()


def _preload_parameter_rows_for_accumulator(
    acc: defaultdict[tuple[Type[Any], str, int, int, str | None], Decimal],
    *,
    version_id: int,
) -> dict[tuple[Type[Any], str, int, int, str | None], Any]:
    """Один/несколько SELECT на модель вместо SELECT на каждую ячейку импорта."""
    from app.common.perimeter_variant.registry import model_supports_perimeter_variant

    by_model_fk: dict[tuple[Type[Any], str], set[tuple[int, int]]] = defaultdict(set)
    for model, fk_column, parent_id, year_n, _pvc in acc.keys():
        by_model_fk[(model, fk_column)].add((int(parent_id), int(year_n)))

    index: dict[tuple[Type[Any], str, int, int, str | None], Any] = {}
    for (model, fk_column), parent_years in by_model_fk.items():
        years = {year_n for _pid, year_n in parent_years}
        supports_pv = model_supports_perimeter_variant(model)
        if fk_column == _FK_AGGREGATE_NO_PARENT:
            q = model.query.filter(model.year_number.in_(years))
            q = filter_by_explicit_db_version(q, model, version_id)
            rows = q.order_by(model.id.asc()).all()
            for row in rows:
                pvc = getattr(row, "perimeter_variant_code", None) if supports_pv else None
                key = (model, fk_column, 0, int(row.year_number), pvc)
                if key not in index:
                    index[key] = row
            continue
        parent_ids = {pid for pid, _year in parent_years}
        q = model.query.filter(
            getattr(model, fk_column).in_(parent_ids),
            model.year_number.in_(years),
        )
        q = filter_by_explicit_db_version(q, model, version_id)
        rows = q.order_by(model.id.asc()).all()
        for row in rows:
            pvc = getattr(row, "perimeter_variant_code", None) if supports_pv else None
            key = (
                model,
                fk_column,
                int(getattr(row, fk_column)),
                int(row.year_number),
                pvc,
            )
            if key not in index:
                index[key] = row
    return index


def _import_username(user: str | None = None) -> str:
    if user:
        return user
    return session.get("username", "Неизвестный пользователь")


def _apply_accumulator_for_version(
    acc: defaultdict[tuple[Type[Any], str, int, int, str | None], Decimal],
    *,
    field_name: str,
    version_id: int,
    preserve_excel_empty_perimeter_variant: bool = False,
    user: str | None = None,
) -> int:
    from app.common.perimeter_variant.registry import model_supports_perimeter_variant

    touched = 0
    import_user = _import_username(user)
    deduped_parents: set[tuple[Any, ...]] = set()
    for (model, fk_column, parent_id, year_n, pvc), total in acc.items():
        # Агрегаты без родителя: один и тот же parent_id (=0) на все годы — включать год в ключ,
        # иначе _dedupe_aggregate_energy_consumption_rows_for_year вызывается только для первого года
        # и лишние строки по другим годам остаются → «удвоение» в суммах/таблицах.
        if fk_column == _FK_AGGREGATE_NO_PARENT:
            dedupe_pk: tuple[Any, ...] = (model, fk_column, parent_id, year_n, pvc)
        elif model_supports_perimeter_variant(model):
            dedupe_pk = (model, fk_column, parent_id, pvc)
        else:
            dedupe_pk = (model, fk_column, parent_id)
        if dedupe_pk not in deduped_parents:
            deduped_parents.add(dedupe_pk)
            if fk_column == _FK_AGGREGATE_NO_PARENT:
                _dedupe_aggregate_energy_consumption_rows_for_year(
                    model, year_n, version_id, perimeter_variant_code=pvc
                )
            else:
                ecps.dedupe_duplicate_energy_consumption_rows_for_parent(
                    model,
                    fk_column,
                    parent_id,
                    database_version_id=version_id,
                    perimeter_variant_code=pvc
                    if model_supports_perimeter_variant(model)
                    else ecps._UNSET,
                )
    db.session.flush()
    existing_by_key = _preload_parameter_rows_for_accumulator(acc, version_id=version_id)

    def _write_bound_row(
        model: Type[Any],
        write_fk: str,
        parent_id: int,
        year_n: int,
        total: Decimal,
        *,
        perimeter_variant_code: str | None = None,
    ) -> None:
        nonlocal touched
        if model_supports_perimeter_variant(model):
            if not (
                preserve_excel_empty_perimeter_variant
                and perimeter_variant_code is None
            ):
                validation_parent_id = (
                    None if write_fk == _FK_AGGREGATE_NO_PARENT else parent_id
                )
                try:
                    perimeter_variant_code = ecps._resolve_perimeter_variant_for_context(
                        perimeter_variant_code,
                        model,
                        None if write_fk == _FK_AGGREGATE_NO_PARENT else write_fk,
                        validation_parent_id,
                    )
                except ValueError:
                    # Расчётная строка сводки для сущности без привязки варианта в каталоге.
                    return
                if perimeter_variant_code is ecps._UNSET:
                    perimeter_variant_code = None
        lookup_parent = 0 if write_fk == _FK_AGGREGATE_NO_PARENT else int(parent_id)
        row_key = (
            model,
            write_fk,
            lookup_parent,
            int(year_n),
            perimeter_variant_code if model_supports_perimeter_variant(model) else None,
        )
        row = existing_by_key.get(row_key)
        if row is None:
            row = model()
            if write_fk != _FK_AGGREGATE_NO_PARENT:
                setattr(row, write_fk, parent_id)
            row.database_version_id = version_id
            row.year_number = year_n
            if model_supports_perimeter_variant(model):
                row.perimeter_variant_code = perimeter_variant_code
            row.created_by = import_user
            db.session.add(row)
            existing_by_key[row_key] = row
        setattr(row, field_name, total)
        row.modified_by = import_user
        touched += 1

    for (model, fk_column, parent_id, year_n, pvc), total in acc.items():
        _write_bound_row(
            model,
            fk_column,
            int(parent_id),
            int(year_n),
            total,
            perimeter_variant_code=pvc if model_supports_perimeter_variant(model) else None,
        )

    db.session.flush()
    return touched


_FORMULA_ROW_PERSIST_PARAMETER_KEYS: tuple[str, ...] = (
    "energy_consumption_mln_kvt_ch",
    "energy_consumption_sipr_mln_kvt_ch",
    "energy_consumption_yoy_growth_pct",
    "energy_consumption_sipr_abs_growth_mln",
    "energy_consumption_sipr_yoy_growth_pct",
)


# Инжект «… без заряда ГАЭС» на /summary/oes|fo|ez — только экранный расчёт.
# У OES-стиля perimeter_variant_code=None совпадает с базовой строкой; запись в БД
# затирала бы ручной ввод (потребление − заряд вместо потребления).
# См. _is_display_only_without_gaes_injected_summary_row в summary_services.


def _accumulator_from_formula_summary_rows(
    summary_rows: list[dict[str, Any]],
    years: list[int],
    *,
    field_name: str,
) -> dict[tuple[Type[Any], str, int, int, str | None], Decimal]:
    from app.energy_consumption.services.energy_consumption_summary_services import (
        _DEMAND_MODEL_CLASS_BY_NAME,
        _is_display_only_without_gaes_injected_summary_row,
        _raw_year_values_from_summary_row,
    )

    acc: dict[tuple[Type[Any], str, int, int, str | None], Decimal] = {}
    for row in summary_rows:
        if not row.get("pd_ec_formula_derived_row"):
            continue
        if _is_display_only_without_gaes_injected_summary_row(row):
            continue
        if row.get("parameter_key") != field_name:
            continue
        demand_model_name = row.get("demand_model_name")
        if not demand_model_name:
            continue
        model_cls = _DEMAND_MODEL_CLASS_BY_NAME.get(str(demand_model_name))
        if model_cls is None:
            continue
        fk_column = row.get("parent_fk_column") or _FK_AGGREGATE_NO_PARENT
        parent_id = int(row.get("parent_id") or 0)
        perimeter_variant_code = row.get("perimeter_variant_code")
        raw_by_year = _raw_year_values_from_summary_row(
            row,
            years,
            ignore_perimeter_variant_year_bounds=True,
        )
        for year in years:
            value = raw_by_year.get(int(year))
            if value is None:
                continue
            acc[(model_cls, fk_column, parent_id, int(year), perimeter_variant_code)] = value
    return acc


def _persist_formula_summary_rows(
    summary_rows: list[dict[str, Any]],
    *,
    database_version_id: int,
    years: list[int],
    skip_accumulator_keys_by_field: dict[str, frozenset[tuple[Any, ...]]] | None = None,
    user: str | None = None,
) -> int:
    if not summary_rows or not years:
        return 0
    skip_by_field = skip_accumulator_keys_by_field or {}
    touched = 0
    for field_name in _FORMULA_ROW_PERSIST_PARAMETER_KEYS:
        acc = _accumulator_from_formula_summary_rows(
            summary_rows,
            years,
            field_name=field_name,
        )
        skip_keys = skip_by_field.get(field_name) or frozenset()
        if skip_keys:
            for skip_key in skip_keys:
                acc.pop(skip_key, None)
        if not acc:
            continue
        touched += _apply_accumulator_for_version(
            defaultdict(lambda: Decimal(0), acc),
            field_name=field_name,
            version_id=database_version_id,
            user=user,
        )
    return touched


def persist_all_energy_consumption_summary_computed_rows(
    *,
    database_version_id: int,
    years: list[int],
    rounding_digits: int = 1,
    skip_accumulator_keys_by_field: dict[str, frozenset[tuple[Any, ...]]] | None = None,
    user: str | None = None,
) -> int:
    """Синхронизация агрегатов ЭЗ/ФО и запись всех расчётных строк сводки в БД."""
    from app.energy_consumption.services.energy_consumption_summary_services import (
        build_all_formula_summary_rows_for_version,
    )

    if not years:
        return 0

    ecps.sync_energy_zone_consumption_aggregates(database_version_id=database_version_id)
    db.session.flush()
    ecps.sync_federal_district_consumption_aggregates(
        database_version_id=database_version_id
    )
    db.session.flush()

    summary_rows = build_all_formula_summary_rows_for_version(
        database_version_id=database_version_id,
        years=years,
        rounding_digits=rounding_digits,
    )
    touched = _persist_formula_summary_rows(
        summary_rows,
        database_version_id=database_version_id,
        years=years,
        skip_accumulator_keys_by_field=skip_accumulator_keys_by_field,
        user=user,
    )
    from app.energy_consumption.services.ec_display_cache import (
        invalidate_energy_consumption_display_caches,
    )

    invalidate_energy_consumption_display_caches(version_id=database_version_id)
    return touched


def persist_all_energy_consumption_summary_computed_rows_for_all_versions(
    *,
    rounding_digits: int = 1,
    years: list[int] | None = None,
    user: str | None = None,
) -> dict[str, Any]:
    """Временно: пересчёт и запись формул сводки во все версии БД (как импорт Excel).

    Позже перейдём на пересчёт только в пределах выбранной версии.
    """
    version_ids = _database_version_ids_for_energy_consumption_import()
    if not years:
        years = None
    total = 0
    processed: list[int] = []
    for vid in version_ids:
        years_ok = (
            sorted({int(y) for y in years})
            if years is not None
            else sorted(_year_numbers_for_version(int(vid)))
        )
        if not years_ok:
            continue
        total += persist_all_energy_consumption_summary_computed_rows(
            database_version_id=int(vid),
            years=years_ok,
            rounding_digits=rounding_digits,
            user=user,
        )
        processed.append(int(vid))
        db.session.flush()
    return {
        "cells_written_formula": total,
        "database_versions_processed": len(processed),
        "database_version_ids": processed,
    }


def persist_summary_table_formula_rows_after_import(
    *,
    database_version_id: int,
    years: list[int],
    rounding_digits: int = 1,
    skip_accumulator_keys_by_field: dict[str, frozenset[tuple[Any, ...]]] | None = None,
) -> int:
    """После импорта Excel: пересчитать формулы сводной таблицы и записать в БД."""
    return persist_all_energy_consumption_summary_computed_rows(
        database_version_id=database_version_id,
        years=years,
        rounding_digits=rounding_digits,
        skip_accumulator_keys_by_field=skip_accumulator_keys_by_field,
    )


def _labels_matched_in_any_version(
    label_keys: set[tuple[str, str | None]],
    version_ids: list[int],
    *,
    variant_column_mode: bool,
) -> set[tuple[str, str | None]]:
    matched: set[tuple[str, str | None]] = set()
    for vid in version_ids:
        ctx = _binding_ctx(vid)
        for lbl, pvc in label_keys:
            if (lbl, pvc) in matched:
                continue
            if _resolve_row_binding(
                lbl,
                ctx,
                perimeter_variant_code=pvc,
                variant_column_mode=variant_column_mode,
            ) is not None:
                matched.add((lbl, pvc))
    return matched


def import_energy_consumption_summary_from_xlsx_bytes(
    raw: bytes,
    *,
    user: str | None = None,
    on_version_progress: Callable[..., None] | None = None,
) -> dict[str, Any]:
    """
    Читает листы «млн. кВт.ч» и/или «СиПР» (какие есть в книге), сопоставляет наименования
    («Россия» / «Россия без НТ» → ``RussiaFederationEnergyConsumptionParameter`` с ``without_nt``,
    даже если в колонке perimeter-variants указан ``with_nt``; «Россия с НТ» → ``with_nt``);
    «ЕЭС России» + варианты с ГАЭС → ``EnergySystemTypeEnergyConsumptionParameter``;
    «ЭЭС России» + варианты → ``EesRussiaEnergyConsumptionParameter``;
    «ОЭС Юга с НТ с ГАЭС» / «… с зарядом ГАЭС» → ``UnionEnergySystemEnergyConsumptionParameter`` с ``with_nt_with_gaes``;
    «ОЭС Юга без НТ с ГАЭС» / «… с зарядом ГАЭС» → ``UnionEnergySystemEnergyConsumptionParameter`` с ``without_nt_with_gaes``;
    «Первая синхронная зона … с зарядом ГАЭС (с ЭС …)» → ``SynchronousAreaEnergyConsumptionParameter``
    с ``without_nt_with_gaes_with_kaliningrad_es``; далее ОЭС → РЭС → субъект РФ → ФО → энергоузел)
    и записывает числа в параметры во все зарегистрированные версии БД.

    Должен присутствовать хотя бы один из двух листов. Отсутствующий лист пропускается.

    Строка, сопоставленная с РЭС, дублируется в параметры субъекта РФ, если в
    gs_sys_regional_district_regional_energy_system у этой РЭС ровно один субъект
    (и он есть в той же версии справочника): лист «млн. кВт.ч» → поле по субъекту,
    лист «СиПР» → поле СиПР по субъекту. Если для того же субъекта и года в файле уже
    есть строка, сопоставленная как субъект РФ (а не как РЭС), копирование с РЭС не
    выполняется — иначе сумма на субъекте удваивается.

    Если лист «млн. кВт.ч» есть, но это короткая заглушка без таблицы, показатели «млн. кВтч»
    из файла не обновляются (как при отсутствии данных на листе).
    """
    version_ids = _database_version_ids_for_energy_consumption_import()
    if not version_ids:
        raise ValueError("Не найдено ни одной версии базы данных для записи импорта.")
    import_user = _import_username(user)
    with ecps.import_actor(import_user):
        return _import_energy_consumption_summary_from_xlsx_bytes_impl(
            raw,
            version_ids=version_ids,
            import_user=import_user,
            on_version_progress=on_version_progress,
        )


def _format_unmatched_label(label: str, pvc: str | None) -> str:
    if pvc:
        return f"{label} [{pvc}]"
    return label


def _emit_import_progress(
    on_version_progress: Callable[..., None] | None,
    done: int,
    total: int,
    detail: str | None = None,
) -> None:
    if on_version_progress is None:
        return
    try:
        on_version_progress(done, total, detail)
    except TypeError:
        # Старые колбэки только (done, total).
        on_version_progress(done, total)


def _import_energy_consumption_summary_from_xlsx_bytes_impl(
    raw: bytes,
    *,
    version_ids: list[int],
    import_user: str,
    on_version_progress: Callable[..., None] | None,
) -> dict[str, Any]:
    versions_total = len(version_ids)
    _emit_import_progress(on_version_progress, 0, versions_total, "чтение Excel…")

    wb = _load_summary_workbook_from_bytes(raw)
    screen_export_mln: defaultdict[tuple[str, int, str | None], Decimal] | None = None
    screen_export_sipr: defaultdict[tuple[str, int, str | None], Decimal] | None = None
    try:
        ws_m = _find_sheet_title_optional(wb, SHEET_MLN_KVTCH)
        ws_s = _find_sheet_title_optional(wb, SHEET_SIPR)
        matrix_m = _sheet_matrix(ws_m) if ws_m is not None else []
        matrix_s = _sheet_matrix(ws_s) if ws_s is not None else []
        if ws_m is None and ws_s is None:
            for title in wb.sheetnames:
                matrix_screen = _sheet_matrix(wb[title])
                mln_part, sipr_part = _aggregate_from_screen_export_matrix(matrix_screen)
                if mln_part or sipr_part:
                    screen_export_mln = mln_part
                    screen_export_sipr = sipr_part
                    break
            if screen_export_mln is None and screen_export_sipr is None:
                raise ValueError(
                    f'В книге нет листов «{SHEET_MLN_KVTCH}» и «{SHEET_SIPR}» для импорта. '
                    "Если файл получен кнопкой «Экспорт в Excel» со старой версии приложения, "
                    "выгрузите его заново (в книге появятся листы для импорта) или используйте "
                    "шаблон СиПР с листами «млн. кВт.ч» и/или «СиПР»."
                )
    finally:
        wb.close()

    if screen_export_mln is not None or screen_export_sipr is not None:
        acc_mln_labels = screen_export_mln or defaultdict(lambda: Decimal("0"))
        acc_sipr_labels = screen_export_sipr or defaultdict(lambda: Decimal("0"))
        variant_column_mode_m = False
        variant_column_mode_s = False
        ws_m = None
        ws_s = None
        matrix_m = []
        matrix_s = []
    else:
        acc_mln_labels, variant_column_mode_m = _aggregate_sheet_labels(
            matrix_m,
            sheet_title=SHEET_MLN_KVTCH,
            allow_empty_short_sheet_without_year_grid=(ws_m is not None),
        )
        acc_sipr_labels, variant_column_mode_s = _aggregate_sheet_labels(
            matrix_s,
            sheet_title=SHEET_SIPR,
        )
    # Важно: режим наличия колонки вариантов периметра может отличаться между листами.
    # Нельзя "склеивать" их через OR, иначе наличие колонки на одном листе заставит другой лист
    # трактоваться как variant-column-mode, что приводит к ошибкам валидации (например, для СЗ).
    variant_column_mode_mln = bool(variant_column_mode_m)
    variant_column_mode_sipr = bool(variant_column_mode_s)

    mln_skipped = (
        ws_m is not None
        and len(matrix_m) > 0
        and len(matrix_m) <= _OPTIONAL_GRID_MAX_ROWS
        and len(acc_mln_labels) == 0
    )
    labels_mln = {(lbl, pvc) for lbl, _, pvc in acc_mln_labels.keys()}
    labels_sipr = {(lbl, pvc) for lbl, _, pvc in acc_sipr_labels.keys()}
    matched_any_mln = _labels_matched_in_any_version(
        labels_mln,
        version_ids,
        variant_column_mode=variant_column_mode_mln,
    )
    matched_any_sipr = _labels_matched_in_any_version(
        labels_sipr,
        version_ids,
        variant_column_mode=variant_column_mode_sipr,
    )
    matched_any = matched_any_mln | matched_any_sipr

    try:
        n_mln = 0
        n_sipr = 0
        n_formula = 0
        for versions_done, vid in enumerate(version_ids, start=1):
            done_before = versions_done - 1
            _emit_import_progress(
                on_version_progress,
                done_before,
                versions_total,
                f"версия {versions_done}/{versions_total}: сопоставление строк",
            )
            ctx = _binding_ctx(vid)
            years_ok = _year_numbers_for_version(vid)
            if not years_ok:
                _emit_import_progress(on_version_progress, versions_done, versions_total)
                continue
            acc_m, explicit_rd_mln = _collapse_labels_to_bind_keys(
                acc_mln_labels,
                ctx=ctx,
                years_ok=years_ok,
                variant_column_mode=variant_column_mode_mln,
            )
            _mirror_res_accumulator_rows_to_single_district_rd(
                acc_m,
                database_version_id=vid,
                explicit_rd_subject_year_pairs=explicit_rd_mln,
            )
            acc_s, explicit_rd_sipr = _collapse_labels_to_bind_keys(
                acc_sipr_labels,
                ctx=ctx,
                years_ok=years_ok,
                variant_column_mode=variant_column_mode_sipr,
            )
            _mirror_res_accumulator_rows_to_single_district_rd(
                acc_s,
                database_version_id=vid,
                explicit_rd_subject_year_pairs=explicit_rd_sipr,
            )
            _emit_import_progress(
                on_version_progress,
                done_before,
                versions_total,
                f"версия {versions_done}/{versions_total}: запись млн. кВтч",
            )
            n_mln += _apply_accumulator_for_version(
                acc_m,
                field_name="energy_consumption_mln_kvt_ch",
                version_id=vid,
                preserve_excel_empty_perimeter_variant=variant_column_mode_mln,
                user=import_user,
            )
            db.session.flush()
            _emit_import_progress(
                on_version_progress,
                done_before,
                versions_total,
                f"версия {versions_done}/{versions_total}: запись СиПР",
            )
            n_sipr += _apply_accumulator_for_version(
                acc_s,
                field_name="energy_consumption_sipr_mln_kvt_ch",
                version_id=vid,
                preserve_excel_empty_perimeter_variant=variant_column_mode_sipr,
                user=import_user,
            )
            db.session.flush()
            _emit_import_progress(
                on_version_progress,
                done_before,
                versions_total,
                f"версия {versions_done}/{versions_total}: пересчёт формул",
            )
            n_formula += persist_all_energy_consumption_summary_computed_rows(
                database_version_id=vid,
                years=sorted(years_ok),
                rounding_digits=1,
                skip_accumulator_keys_by_field={
                    "energy_consumption_mln_kvt_ch": frozenset(acc_m.keys()),
                    "energy_consumption_sipr_mln_kvt_ch": frozenset(acc_s.keys()),
                },
                user=import_user,
            )
            db.session.flush()
            db.session.commit()
            _emit_import_progress(on_version_progress, versions_done, versions_total)
        if db.session.is_active:
            db.session.commit()
    except IntegrityError:
        db.session.rollback()
        raise ValueError("Не удалось сохранить импорт (конфликт целостности данных).") from None

    def _uniq_sorted(xs: set[str], limit: int) -> list[str]:
        out = sorted(xs, key=lambda s: s.casefold())[:limit]
        return out

    stats_out: dict[str, Any] = {
        "cells_written_mln_kvt_ch": n_mln,
        "cells_written_sipr": n_sipr,
        "cells_written_formula": n_formula,
        "database_versions_processed": len(version_ids),
        "unmatched_labels_mln": _uniq_sorted(
            {
                _format_unmatched_label(lbl, pvc)
                for lbl, pvc in labels_mln - matched_any
            },
            80,
        ),
        "unmatched_labels_sipr": _uniq_sorted(
            {
                _format_unmatched_label(lbl, pvc)
                for lbl, pvc in labels_sipr - matched_any
            },
            80,
        ),
        "mln_sheet_skipped_no_data_grid": mln_skipped,
        "mln_sheet_absent": ws_m is None,
        "sipr_sheet_absent": ws_s is None,
        "variant_column_mode_mln": variant_column_mode_mln,
        "variant_column_mode_sipr": variant_column_mode_sipr,
        "variant_column_mode": variant_column_mode_mln or variant_column_mode_sipr,
    }
    try:
        from app.energy_consumption.services.energy_consumption_summary_logging import (
            log_ec_summary_excel_import,
        )

        log_ec_summary_excel_import(
            import_user,
            stats_out,
            database_version_id=version_ids[0] if version_ids else get_current_version(),
        )
    except Exception:
        pass
    return stats_out
