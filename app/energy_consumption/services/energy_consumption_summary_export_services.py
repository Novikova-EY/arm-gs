"""Выгрузка сводных таблиц потребления в Excel (в духе экранной таблицы: шапка, подписи к годам, заливки).

Дополнительно в книгу добавляются листы «млн. кВт.ч» и «СиПР» в формате, совместимом с импортом сводки.
Числовые ячейки основного листа: полное значение + number_format (как на выгрузке «Выработка ЭЭ»).
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.common.services.excel_numeric_cell import (
    excel_number_format,
    excel_numeric_cell_value,
    parse_excel_numeric,
)
from app.energy_consumption.services.energy_consumption_summary_import_services import (
    SHEET_MLN_KVTCH,
    SHEET_SIPR,
)
from app.energy_consumption.services.energy_consumption_summary_services import (
    ENERGY_CONSUMPTION_SIPR_ABS_PARAMETER_KEY,
    ENERGY_CONSUMPTION_SIPR_YOY_PARAMETER_KEY,
    ENERGY_CONSUMPTION_YOY_PARAMETER_KEY,
    GAES_CHARGE_PARAMETER_KEY,
)

_ENERGY_CONSUMPTION_YOY_PARAMETER_KEYS = frozenset(
    {
        ENERGY_CONSUMPTION_YOY_PARAMETER_KEY,
        ENERGY_CONSUMPTION_SIPR_YOY_PARAMETER_KEY,
    }
)
_ENERGY_CONSUMPTION_YOY_DISPLAY_DECIMALS = 2
_EC_SUMMARY_VERIFY_FOR_DISPLAY_DECIMALS = 6
_SIPR_INTEGER_DISPLAY_PARAMETER_KEYS = frozenset(
    {
        "energy_consumption_sipr_mln_kvt_ch",
        ENERGY_CONSUMPTION_SIPR_ABS_PARAMETER_KEY,
        GAES_CHARGE_PARAMETER_KEY,
    }
)

_IMPORT_PARAMETER_KEYS = (
    "energy_consumption_mln_kvt_ch",
    "energy_consumption_sipr_mln_kvt_ch",
)


def _excel_raw_year_source(row: dict[str, Any], index: int) -> Any:
    """Полное значение: tooltip, иначе экранная строка."""
    tooltips = row.get("year_numeric_tooltips") or []
    if index < len(tooltips) and tooltips[index]:
        return tooltips[index]
    yvals = row.get("year_values") or []
    if index < len(yvals):
        return yvals[index]
    return None


def _excel_display_digits_for_row(
    row: dict[str, Any],
    *,
    rounding_digits: int,
    sipr_on: bool,
    summary_table_page: bool,
    verification_row: bool,
) -> int:
    if verification_row:
        return _EC_SUMMARY_VERIFY_FOR_DISPLAY_DECIMALS
    pk = str(row.get("parameter_key") or "")
    if pk in _ENERGY_CONSUMPTION_YOY_PARAMETER_KEYS:
        return _ENERGY_CONSUMPTION_YOY_DISPLAY_DECIMALS
    if sipr_on and (
        row.get("pd_ec_sipr_integer_display_row")
        or (summary_table_page and pk in _SIPR_INTEGER_DISPLAY_PARAMETER_KEYS)
    ):
        return -1
    return rounding_digits


def _import_cell_value(raw: Any) -> Any:
    """Число для листа импорта; пустые и «—» остаются пустыми."""
    return parse_excel_numeric(raw)


def _build_import_format_matrix(
    summary_rows: list[dict[str, Any]],
    years: list[int],
    *,
    parameter_key: str,
) -> list[list[Any]]:
    """Строки листа «млн. кВт.ч» / «СиПР» из тех же summary_rows, что и экранный экспорт."""
    data_rows: list[tuple[str, str | None, list[Any]]] = []
    has_variant = False
    for row in summary_rows:
        if str(row.get("parameter_key") or "") != parameter_key:
            continue
        label = str(row.get("entity_label") or "").strip()
        if not label or label.startswith("Проверка "):
            continue
        pvc = row.get("perimeter_variant_code")
        if pvc not in (None, ""):
            has_variant = True
        raw_vals = [_excel_raw_year_source(row, i) for i in range(len(years))]
        data_rows.append((label, str(pvc) if pvc not in (None, "") else None, raw_vals))

    header: list[Any] = (
        ["perimeter-variants", "Наименование", *years]
        if has_variant
        else ["", "Наименование", *years]
    )
    matrix: list[list[Any]] = [header]
    for label, pvc, raw_vals in data_rows:
        body: list[Any] = [pvc or "", label] if has_variant else ["", label]
        for i, _y in enumerate(years):
            v = raw_vals[i] if i < len(raw_vals) else None
            body.append(_import_cell_value(v))
        matrix.append(body)
    return matrix


def _append_ec_summary_import_sheets(
    wb: Workbook,
    summary_rows: list[dict[str, Any]],
    years: list[int],
) -> None:
    """Листы для повторного импорта (редактирование и «Импорт из Excel»)."""
    mln = _build_import_format_matrix(
        summary_rows, years, parameter_key=_IMPORT_PARAMETER_KEYS[0]
    )
    sipr = _build_import_format_matrix(
        summary_rows, years, parameter_key=_IMPORT_PARAMETER_KEYS[1]
    )
    if len(mln) > 1:
        ws_m = wb.create_sheet(SHEET_MLN_KVTCH)
        for row in mln:
            ws_m.append(row)
    if len(sipr) > 1:
        ws_s = wb.create_sheet(SHEET_SIPR)
        for row in sipr:
            ws_s.append(row)


def _verification_export_year_font(value: Any) -> Font:
    """Курсив; ненулевые значения — красным (как на экране сводки)."""
    parsed = parse_excel_numeric(value)
    if parsed is None or parsed == 0:
        return Font(size=11, italic=True)
    return Font(size=11, italic=True, color="FF0000")


def build_demand_summary_excel_stream(
    *,
    summary_rows: list[dict[str, Any]],
    years: list[int],
    sheet_title: str,
    year_features: dict[int, Any] | dict[Any, Any] | None = None,
    rounding_digits: int = 1,
    sipr_on: bool = False,
    summary_table_page: bool = False,
) -> BytesIO:
    yf = year_features or {}
    wb = Workbook()
    ws = wb.active
    safe_title = (sheet_title or "Сводка").replace("/", "-").replace("\\", "-")[:31]
    ws.title = safe_title or "Сводка"

    header_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
    header_font = Font(bold=True, size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    entity_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    entity_font = Font(bold=True, size=11)
    entity_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    fill_group_param = PatternFill(start_color="EEF5FF", end_color="EEF5FF", fill_type="solid")
    fill_child_param = PatternFill(start_color="FBFCFF", end_color="FBFCFF", fill_type="solid")
    fill_cell_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_cell_stripe = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

    thin = Side(style="thin", color="CCCCCC")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    base_font = Font(size=11)
    center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def year_header_cell(y: int) -> str:
        cap = yf.get(y)
        s = str(cap).strip() if cap is not None else ""
        if s:
            return f"{y}\n{s}"
        return str(y)

    ncols = 2 + len(years)
    ws.append([None] * ncols)
    ws.row_dimensions[1].height = 36

    ws.cell(row=1, column=1, value="Энергосистема")
    ws.cell(row=1, column=2, value="Наименование параметров")
    for j, y in enumerate(years, start=3):
        ws.cell(row=1, column=j, value=year_header_cell(y))

    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = cell_border

    ws.freeze_panes = "A2"

    def body_fills(entity_kind: str, stripe: bool) -> tuple[PatternFill, PatternFill]:
        if entity_kind in ("group", "group-root"):
            return fill_group_param, fill_group_param
        if entity_kind in ("child", "res_subject_sum_check"):
            return fill_child_param, fill_child_param
        param_fill = fill_cell_stripe if stripe else fill_cell_white
        return param_fill, param_fill

    for bi, row in enumerate(summary_rows):
        excel_row = bi + 2
        entity_kind = str(row.get("entity_kind") or "default")
        stripe = bool(bi % 2)
        param_fill, year_fill = body_fills(entity_kind, stripe)
        verification_row = str(row.get("entity_label") or "").startswith("Проверка ")
        digits = _excel_display_digits_for_row(
            row,
            rounding_digits=rounding_digits,
            sipr_on=sipr_on,
            summary_table_page=summary_table_page,
            verification_row=verification_row,
        )

        entity_cell = ""
        if row.get("show_entity_cell"):
            depth = int(row.get("entity_depth") or 0)
            label = row.get("entity_label") or ""
            entity_cell = ("  " * depth) + str(label)

        ws.cell(row=excel_row, column=1, value=entity_cell or None)
        ws.cell(row=excel_row, column=2, value=row.get("parameter_label") or "")

        c1 = ws.cell(row=excel_row, column=1)
        c1.font = Font(bold=True, size=11, italic=True) if verification_row else entity_font
        c1.fill = entity_fill
        c1.alignment = entity_align
        c1.border = cell_border

        c2 = ws.cell(row=excel_row, column=2)
        c2.font = Font(size=11, italic=True) if verification_row else base_font
        c2.fill = param_fill
        c2.alignment = left_wrap
        c2.border = cell_border

        for i, _y in enumerate(years):
            raw = _excel_raw_year_source(row, i)
            cell_value, is_numeric = excel_numeric_cell_value(
                raw, verification=verification_row, zero_as_dash=False
            )
            c = ws.cell(row=excel_row, column=3 + i, value=cell_value)
            c.font = (
                _verification_export_year_font(cell_value)
                if verification_row
                else base_font
            )
            c.fill = year_fill
            c.alignment = center_wrap
            c.border = cell_border
            if is_numeric:
                c.number_format = excel_number_format(digits)

    for col in range(1, ncols + 1):
        letter = get_column_letter(col)
        if col == 1:
            ws.column_dimensions[letter].width = 44
        elif col == 2:
            ws.column_dimensions[letter].width = 52
        else:
            ws.column_dimensions[letter].width = 14

    _append_ec_summary_import_sheets(wb, summary_rows, years)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
