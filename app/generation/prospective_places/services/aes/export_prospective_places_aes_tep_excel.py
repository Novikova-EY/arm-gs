# -*- coding: utf-8 -*-
"""Экспорт одного листа ТЭП АЭС (основные или резервные) в ценах года или исходных."""
from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def write_aes_tep_screen_sheet(
    ws,
    tep_rows: list[dict[str, Any]],
    *,
    sheet_title: str,
    show_price_year_columns: bool = True,
) -> None:
    """Заполняет лист по структуре HTML ``_tep_machines_table.html``."""
    ws.title = sheet_title[:31]
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    header_font = Font(bold=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _hdr(r: int, c: int, val: str) -> None:
        cell = ws.cell(row=r, column=c, value=val)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    row1_titles_single = [
        "№ п/п",
        "Наименование площадки",
        "Станционный номер блока",
        "Тип энергоблока",
        "Мощность энергоблока, МВт",
        "Возможный срок реализации",
        "Срок эксплуатации АЭС, лет",
        "Срок строительства АЭС, лет",
        "Предельное годовое число часов использования мощности энергоблока, час",
    ]
    for col, val in enumerate(row1_titles_single, 1):
        _hdr(1, col, val)
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)

    if show_price_year_columns:
        _hdr(1, 10, "Удельная топливная составляющая эксплуатационных затрат в ценах текущего года")
        ws.merge_cells(start_row=1, start_column=10, end_row=1, end_column=11)
        _hdr(2, 10, "руб./кВт·ч")
        _hdr(2, 11, "год")

        _hdr(1, 12, "Удельные условно постоянные эксплуатационные затраты (без амортизационных отчислений)")
        ws.merge_cells(start_row=1, start_column=12, end_row=1, end_column=13)
        _hdr(2, 12, "тыс. руб./кВт")
        _hdr(2, 13, "год")

        _hdr(1, 14, "Относительная величина расхода электрической энергии на собственные нужды АЭС, %")
        ws.merge_cells(start_row=1, start_column=14, end_row=2, end_column=14)

        _hdr(1, 15, "Удельные капиталовложения в строительство АЭС в ценах текущего года (без НДС)")
        ws.merge_cells(start_row=1, start_column=15, end_row=1, end_column=16)
        _hdr(2, 15, "тыс. руб./кВт")
        _hdr(2, 16, "год")

        _hdr(1, 17, "Удельные затраты на вывод из эксплуатации")
        ws.merge_cells(start_row=1, start_column=17, end_row=1, end_column=18)
        _hdr(2, 17, "тыс. руб./кВт")
        _hdr(2, 18, "год")

        _hdr(1, 19, "Вероятность аварийного состояния")
        ws.merge_cells(start_row=1, start_column=19, end_row=2, end_column=19)

        _hdr(1, 20, "Относительная продолжительность плановых простоев:")
        ws.merge_cells(start_row=1, start_column=20, end_row=1, end_column=21)
        _hdr(2, 20, "ОЗП")
        _hdr(2, 21, "ВЛП")

        total_cols = 21
    else:
        _hdr(
            1,
            10,
            "Удельная топливная составляющая эксплуатационных затрат\nруб./кВт·ч",
        )
        ws.merge_cells(start_row=1, start_column=10, end_row=2, end_column=10)

        _hdr(
            1,
            11,
            "Удельные условно постоянные эксплуатационные затраты (без амортизационных отчислений)\nтыс. руб./кВт",
        )
        ws.merge_cells(start_row=1, start_column=11, end_row=2, end_column=11)

        _hdr(1, 12, "Относительная величина расхода электрической энергии на собственные нужды АЭС, %")
        ws.merge_cells(start_row=1, start_column=12, end_row=2, end_column=12)

        _hdr(1, 13, "Удельные капиталовложения в строительство АЭС (без НДС)\nтыс. руб./кВт")
        ws.merge_cells(start_row=1, start_column=13, end_row=2, end_column=13)

        _hdr(1, 14, "Удельные затраты на вывод из эксплуатации\nтыс. руб./кВт")
        ws.merge_cells(start_row=1, start_column=14, end_row=2, end_column=14)

        _hdr(1, 15, "Вероятность аварийного состояния")
        ws.merge_cells(start_row=1, start_column=15, end_row=2, end_column=15)

        _hdr(1, 16, "Относительная продолжительность плановых простоев:")
        ws.merge_cells(start_row=1, start_column=16, end_row=1, end_column=17)
        _hdr(2, 16, "ОЗП")
        _hdr(2, 17, "ВЛП")

        total_cols = 17

    row_num = 3
    total_mw = 0.0
    for row in tep_rows:
        t = row["tep"]
        m = row["machine"]
        site = row["site_name"]
        if show_price_year_columns:
            vals = [
                row["idx"],
                site,
                m["station_block_number"],
                m["unit_type"],
                m["unit_capacity_mw"],
                m["possible_implementation_period"],
                t["service_life_years"],
                t["construction_period_years"],
                t["max_annual_operating_hours"],
                t["specific_fuel_cost_rub_per_kwh"],
                t["year_specific_fuel_cost"],
                t["specific_fixed_operating_costs_thous_rub_per_kw"],
                t["year_specific_fixed_operating_costs"],
                t["relative_auxiliary_power_consumption_pct"],
                t["specific_capital_investment_thous_rub_per_kw"],
                t["year_specific_capital_investment"],
                t["specific_decommissioning_cost_thous_rub_per_kw"],
                t["year_specific_decommissioning"],
                t["emergency_state_probability"],
                t["ozp"],
                t["vlp"],
            ]
        else:
            vals = [
                row["idx"],
                site,
                m["station_block_number"],
                m["unit_type"],
                m["unit_capacity_mw"],
                m["possible_implementation_period"],
                t["service_life_years"],
                t["construction_period_years"],
                t["max_annual_operating_hours"],
                t["specific_fuel_cost_rub_per_kwh"],
                t["specific_fixed_operating_costs_thous_rub_per_kw"],
                t["relative_auxiliary_power_consumption_pct"],
                t["specific_capital_investment_thous_rub_per_kw"],
                t["specific_decommissioning_cost_thous_rub_per_kw"],
                t["emergency_state_probability"],
                t["ozp"],
                t["vlp"],
            ]
        try:
            u = (m.get("unit_capacity_mw") or "").replace("—", "").strip()
            if u:
                total_mw += float(u.replace(",", "."))
        except (ValueError, TypeError):
            pass
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.border = thin_border
            cell.alignment = center_align
        row_num += 1

    footer = row_num
    if tep_rows:
        ws.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=4)
        c = ws.cell(row=footer, column=1, value="Всего:")
        c.font = header_font
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.border = thin_border
        c.fill = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid")
        ws.cell(row=footer, column=5, value=int(round(total_mw, 0)))
        ws.cell(row=footer, column=5).border = thin_border
        ws.cell(row=footer, column=5).alignment = center_align
        ws.cell(row=footer, column=5).fill = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid")
        ws.merge_cells(start_row=footer, start_column=6, end_row=footer, end_column=total_cols)
        for cidx in range(6, total_cols + 1):
            ws.cell(row=footer, column=cidx).border = thin_border
            ws.cell(row=footer, column=cidx).fill = PatternFill(
                start_color="E2E3E5", end_color="E2E3E5", fill_type="solid"
            )
    else:
        ws.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=total_cols)
        c = ws.cell(row=footer, column=1, value="Нет данных")
        c.alignment = center_align
        c.border = thin_border

    widths_21 = [6, 22, 12, 14, 10, 12, 8, 8, 12, 12, 7, 12, 7, 10, 12, 7, 12, 7, 12, 8, 8]
    widths_17 = [6, 22, 12, 14, 10, 12, 8, 8, 12, 14, 14, 10, 12, 12, 12, 8, 8]
    widths = widths_21 if show_price_year_columns else widths_17
    for i, w in enumerate(widths[:total_cols], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
