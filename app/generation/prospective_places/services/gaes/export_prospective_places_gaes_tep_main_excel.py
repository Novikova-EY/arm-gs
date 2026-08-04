# -*- coding: utf-8 -*-
"""Excel «ТЭП основных площадок ГАЭС»: как на экране — цены текущего года (без столбцов «год» у капзатрат и удельных) или исходные данные с годами."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TEP_MAIN_COLS = 41
TEP_MAIN_SOURCE_COLS = 46


def _thin_border():
    return Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )


def _style_cell(cell, *, font=None, fill=None, align=None, border=None):
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if align is not None:
        cell.alignment = align
    if border is not None:
        cell.border = border


def _apply_border_to_range(ws, min_row, max_row, min_col, max_col, border):
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(row=r, column=c).border = border


def write_tep_main_screen_sheet(
    ws, tep_groups: list[dict], *, source_prices: bool = False
) -> None:
    """ТЭП основных площадок ГАЭС: «текущий год» или «исходные данные» (годы у капзатрат)."""
    if source_prices:
        _write_gaes_tep_main_sheet_source_prices(ws, tep_groups)
    else:
        _write_gaes_tep_main_sheet_current_year_prices(ws, tep_groups)


def _write_gaes_tep_main_sheet_current_year_prices(ws, tep_groups: list[dict]) -> None:
    """Заполняет лист: 3 строки шапки (как в _tep_machines_table_main.html), группы по типу площадки, строки ТЭП."""
    thin = _thin_border()
    header_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    group_fill = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid")
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Шапка: очередность — строка 1; удельные без столбцов «год» (как капзатраты в ценах года); строки 2–3 (rowspan);
    # выработка / остальное как в _tep_machines_table_main.html
    ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=1)
    c = ws.cell(row=1, column=1, value="№ п/п")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 3, 1, 1, thin)

    ws.merge_cells(start_row=1, start_column=2, end_row=3, end_column=2)
    c = ws.cell(row=1, column=2, value="Наименование площадки")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 3, 2, 2, thin)

    ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=23)
    c = ws.cell(row=1, column=3, value="Основные параметры")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 1, 3, 23, thin)

    ws.merge_cells(start_row=1, start_column=24, end_row=1, end_column=35)
    c = ws.cell(
        row=1,
        column=24,
        value="Очередность строительства\n(прирост вводимой мощности), МВт",
    )
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 1, 24, 35, thin)

    ws.merge_cells(start_row=1, start_column=36, end_row=1, end_column=38)
    c = ws.cell(row=1, column=36, value="Капитальные затраты, млн руб.")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 1, 36, 38, thin)

    for y in range(1, 13):
        col = 23 + y
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        c = ws.cell(row=2, column=col, value=f"{y} год")
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 2, 3, col, col, thin)

    for col, text in (
        (36, "Капитальные затраты (без ПИР)"),
        (37, "здания, сооружения,\nоборудование и бассейнами ГАЭС"),
        (38, "СВМ"),
    ):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        c = ws.cell(row=2, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 2, 3, col, col, thin)

    ws.merge_cells(start_row=1, start_column=39, end_row=3, end_column=39)
    c = ws.cell(
        row=1,
        column=39,
        value="Удельные условно-постоянные эксплуатационные затраты (без амортизационных отчислений), млн руб./МВт",
    )
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 3, 39, 39, thin)

    ws.merge_cells(start_row=1, start_column=40, end_row=3, end_column=40)
    c = ws.cell(
        row=1,
        column=40,
        value="Удельные капиталовложения в строительство (с бассейнами), млн руб./МВт",
    )
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 3, 40, 40, thin)

    ws.merge_cells(start_row=1, start_column=41, end_row=3, end_column=41)
    c = ws.cell(row=1, column=41, value="Примечание")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 3, 41, 41, thin)

    ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=6)
    c = ws.cell(row=2, column=3, value="Установленная мощность,\nгенераторный режим, МВт")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 3, 6, thin)

    ws.merge_cells(start_row=2, start_column=7, end_row=2, end_column=10)
    c = ws.cell(row=2, column=7, value="Установленная мощность,\nнасосный режим, МВт")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 7, 10, thin)

    main_detail_labels_row2 = [
        (11, "Количество агрегатов, шт."),
    ]
    for col, text in main_detail_labels_row2:
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        c = ws.cell(row=2, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 2, 3, col, col, thin)

    ws.merge_cells(start_row=2, start_column=12, end_row=2, end_column=13)
    c = ws.cell(row=2, column=12, value="Единичная мощность агрегата,\nМВт")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 12, 13, thin)

    for col, text in (
        (12, "генераторный режим"),
        (13, "насосный режим"),
    ):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    for col, text in ((14, "Тип гидротурбины"), (15, "Срок строительства ГАЭС, лет")):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        c = ws.cell(row=2, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 2, 3, col, col, thin)

    ws.merge_cells(start_row=2, start_column=16, end_row=2, end_column=18)
    c = ws.cell(row=2, column=16, value="Годовая выработка электроэнергии")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 16, 18, thin)

    for col, text in ((16, "всего"), (17, "1 очередь"), (18, "2 очередь")):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    ws.merge_cells(start_row=2, start_column=19, end_row=2, end_column=21)
    c = ws.cell(
        row=2,
        column=19,
        value="Годовое потребление электроэнергии ГАЭС на заряд, млн кВтч",
    )
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 19, 21, thin)

    for col, text in (
        (19, "всего"),
        (20, "1 очередь"),
        (21, "2 очередь"),
    ):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    ws.merge_cells(start_row=2, start_column=22, end_row=2, end_column=23)
    c = ws.cell(row=2, column=22, value="Число часов использования установленной мощности в сутки, час/сутки")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 22, 23, thin)

    for col, text in (
        (22, "генераторный режим"),
        (23, "насосный режим"),
    ):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    for col, text in (
        (3, "всего"),
        (4, "в т.ч. пусковой комплекс"),
        (5, "1 очередь"),
        (6, "2 очередь"),
        (7, "всего"),
        (8, "в т.ч. пусковой комплекс"),
        (9, "1 очередь"),
        (10, "2 очередь"),
    ):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    row_num = 4

    def _write_tf_row(rn: int, tf: dict) -> None:
        vals = [
            tf["installed_capacity_mw_generator_mode"],
            tf["startup_complex_capacity_mw_generator_mode"],
            tf["stage_1_capacity_mw_generator_mode"],
            tf["stage_2_capacity_mw_generator_mode"],
            tf["installed_capacity_mw_pump_mode"],
            tf["startup_complex_capacity_mw_pump_mode"],
            tf["stage_1_capacity_mw_pump_mode"],
            tf["stage_2_capacity_mw_pump_mode"],
            tf["units_count"],
            tf["unit_capacity_mw_generator_mode"],
            tf["unit_capacity_mw_pump_mode"],
            tf["hydro_turbine_type"],
            tf["construction_period_years"],
        ]
        vals.extend(
            [
                tf["generation_average_multiyear_million_kwh"],
                tf["generation_average_multiyear_million_kwh_stage_1"],
                tf["generation_average_multiyear_million_kwh_stage_2"],
                tf["annual_charging_electricity_consumption_million_kwh"],
                tf["annual_charging_electricity_consumption_million_kwh_stage_1"],
                tf["annual_charging_electricity_consumption_million_kwh_stage_2"],
                tf["ccium_turbine_mode"],
                tf["ccium_pump_mode"],
            ]
        )
        for i in range(1, 13):
            vals.append(tf[f"construction_increment_year_{i:02d}_mw"])
        vals.extend(
            [
                tf["capital_cost_wo_pir_total_million_rub"],
                tf["capital_cost_wo_pir_ges_with_reservoir_million_rub"],
                tf["capital_cost_wo_pir_svm_million_rub"],
                tf["specific_semifixed_operating_costs_thous_rub_per_kw"],
                tf["specific_capital_investment_thous_rub_per_kw"],
                tf["note"],
            ]
        )
        note_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        note_col = 3 + len(vals) - 1
        for col, val in enumerate(vals, start=3):
            cell = ws.cell(row=rn, column=col, value=val)
            cell.border = thin
            cell.alignment = note_align if col == note_col else center

    for g in tep_groups:
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=TEP_MAIN_COLS)
        c = ws.cell(row=row_num, column=1, value=g.get("type_label", ""))
        c.font = Font(bold=True, size=14)
        c.alignment = center
        c.fill = group_fill
        c.border = thin
        _apply_border_to_range(ws, row_num, row_num, 1, TEP_MAIN_COLS, thin)
        row_num += 1

        tep_rows = g.get("tep_rows") or []
        if not tep_rows:
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=TEP_MAIN_COLS)
            c = ws.cell(row=row_num, column=1, value="Нет данных")
            c.alignment = center
            c.border = thin
            _apply_border_to_range(ws, row_num, row_num, 1, TEP_MAIN_COLS, thin)
            row_num += 1
            continue

        for tep_row in tep_rows:
            tf = dict(tep_row["tep_full"])
            ty = tep_row.get("tep_full_current_year")
            if ty:
                tf["capital_cost_wo_pir_total_million_rub"] = ty[
                    "capital_cost_wo_pir_total_million_rub"
                ]
                tf["capital_cost_wo_pir_ges_with_reservoir_million_rub"] = ty[
                    "capital_cost_wo_pir_ges_with_reservoir_million_rub"
                ]
                tf["capital_cost_wo_pir_svm_million_rub"] = ty["capital_cost_wo_pir_svm_million_rub"]
                tf["specific_semifixed_operating_costs_thous_rub_per_kw"] = ty[
                    "specific_semifixed_operating_costs_thous_rub_per_kw"
                ]
                tf["specific_capital_investment_thous_rub_per_kw"] = ty[
                    "specific_capital_investment_thous_rub_per_kw"
                ]
            ws.cell(row=row_num, column=1, value=tep_row["idx"])
            ws.cell(row=row_num, column=1).border = thin
            ws.cell(row=row_num, column=1).alignment = center

            if tep_row["show_station_cell"]:
                rs = int(tep_row["station_rowspan"] or 1)
                ws.cell(row=row_num, column=2, value=tep_row["site_name"])
                if rs > 1:
                    ws.merge_cells(
                        start_row=row_num,
                        start_column=2,
                        end_row=row_num + rs - 1,
                        end_column=2,
                    )
                c2 = ws.cell(row=row_num, column=2)
                c2.border = thin
                c2.alignment = left
            _write_tf_row(row_num, tf)
            row_num += 1

    if not tep_groups:
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=TEP_MAIN_COLS)
        c = ws.cell(row=row_num, column=1, value="Нет данных")
        c.alignment = center
        c.border = thin
        _apply_border_to_range(ws, row_num, row_num, 1, TEP_MAIN_COLS, thin)

    widths = [6, 28] + [12] * 15 + [12] * 3 + [12] * 3 + [8] * 12 + [12] * 3 + [12, 12] + [32]
    for i, w in enumerate(widths, start=1):
        if i <= TEP_MAIN_COLS:
            ws.column_dimensions[get_column_letter(i)].width = min(w, 45)


def _write_gaes_tep_main_sheet_source_prices(ws, tep_groups: list[dict]) -> None:
    """Как _tep_machines_table_main_source.html: исходные капзатраты и удельные показатели с годами (46 колонок)."""
    thin = _thin_border()
    header_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    group_fill = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid")
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ncols = TEP_MAIN_SOURCE_COLS

    ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=1)
    c = ws.cell(row=1, column=1, value="№ п/п")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 3, 1, 1, thin)

    ws.merge_cells(start_row=1, start_column=2, end_row=3, end_column=2)
    c = ws.cell(row=1, column=2, value="Наименование площадки")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 3, 2, 2, thin)

    ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=23)
    c = ws.cell(row=1, column=3, value="Основные параметры")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 1, 3, 23, thin)

    ws.merge_cells(start_row=1, start_column=24, end_row=1, end_column=35)
    c = ws.cell(
        row=1,
        column=24,
        value="Очередность строительства\n(прирост вводимой мощности), МВт",
    )
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 1, 24, 35, thin)

    ws.merge_cells(start_row=1, start_column=36, end_row=1, end_column=41)
    c = ws.cell(row=1, column=36, value="Капитальные затраты, млн руб.")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 1, 36, 41, thin)

    for y in range(1, 13):
        col = 23 + y
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        c = ws.cell(row=2, column=col, value=f"{y} год")
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 2, 3, col, col, thin)

    for c0, c1, text in (
        (36, 37, "Капитальные затраты (без ПИР)"),
        (38, 39, "здания, сооружения,\nоборудование и бассейнами ГАЭС"),
        (40, 41, "СВМ"),
    ):
        ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c1)
        c = ws.cell(row=2, column=c0, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 2, 2, c0, c1, thin)

    for col, text in (
        (36, "млн руб."),
        (37, "год"),
        (38, "млн руб."),
        (39, "год"),
        (40, "млн руб."),
        (41, "год"),
    ):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    ws.merge_cells(start_row=1, start_column=42, end_row=2, end_column=43)
    c = ws.cell(
        row=1,
        column=42,
        value="Удельные условно-постоянные эксплуатационные затраты (без амортизационных отчислений),",
    )
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 2, 42, 43, thin)

    ws.merge_cells(start_row=1, start_column=44, end_row=2, end_column=45)
    c = ws.cell(
        row=1,
        column=44,
        value="Удельные капиталовложения в строительство (с бассейнами)",
    )
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 2, 44, 45, thin)

    for col, text in (
        (42, "млн руб./МВт"),
        (43, "год"),
        (44, "млн руб./МВт"),
        (45, "год"),
    ):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    ws.merge_cells(start_row=1, start_column=46, end_row=3, end_column=46)
    c = ws.cell(row=1, column=46, value="Примечание")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 1, 3, 46, 46, thin)

    ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=6)
    c = ws.cell(row=2, column=3, value="Установленная мощность,\nгенераторный режим, МВт")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 3, 6, thin)

    ws.merge_cells(start_row=2, start_column=7, end_row=2, end_column=10)
    c = ws.cell(row=2, column=7, value="Установленная мощность,\nнасосный режим, МВт")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 7, 10, thin)

    main_detail_labels_row2 = [
        (11, "Количество агрегатов, шт."),
    ]
    for col, text in main_detail_labels_row2:
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        c = ws.cell(row=2, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 2, 3, col, col, thin)

    ws.merge_cells(start_row=2, start_column=12, end_row=2, end_column=13)
    c = ws.cell(row=2, column=12, value="Единичная мощность агрегата,\nМВт")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 12, 13, thin)

    for col, text in (
        (12, "генераторный режим"),
        (13, "насосный режим"),
    ):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    for col, text in ((14, "Тип гидротурбины"), (15, "Срок строительства ГАЭС, лет")):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        c = ws.cell(row=2, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 2, 3, col, col, thin)

    ws.merge_cells(start_row=2, start_column=16, end_row=2, end_column=18)
    c = ws.cell(row=2, column=16, value="Годовая выработка электроэнергии")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 16, 18, thin)

    for col, text in ((16, "всего"), (17, "1 очередь"), (18, "2 очередь")):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    ws.merge_cells(start_row=2, start_column=19, end_row=2, end_column=21)
    c = ws.cell(
        row=2,
        column=19,
        value="Годовое потребление электроэнергии ГАЭС на заряд, млн кВтч",
    )
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 19, 21, thin)

    for col, text in (
        (19, "всего"),
        (20, "1 очередь"),
        (21, "2 очередь"),
    ):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    ws.merge_cells(start_row=2, start_column=22, end_row=2, end_column=23)
    c = ws.cell(row=2, column=22, value="Число часов использования установленной мощности в сутки, час/сутки")
    _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
    _apply_border_to_range(ws, 2, 2, 22, 23, thin)

    for col, text in (
        (22, "генераторный режим"),
        (23, "насосный режим"),
    ):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    for col, text in (
        (3, "всего"),
        (4, "в т.ч. пусковой комплекс"),
        (5, "1 очередь"),
        (6, "2 очередь"),
        (7, "всего"),
        (8, "в т.ч. пусковой комплекс"),
        (9, "1 очередь"),
        (10, "2 очередь"),
    ):
        c = ws.cell(row=3, column=col, value=text)
        _style_cell(c, font=header_font, fill=header_fill, align=center, border=thin)
        _apply_border_to_range(ws, 3, 3, col, col, thin)

    row_num = 4

    def _write_tf_row(rn: int, tf: dict) -> None:
        vals = [
            tf["installed_capacity_mw_generator_mode"],
            tf["startup_complex_capacity_mw_generator_mode"],
            tf["stage_1_capacity_mw_generator_mode"],
            tf["stage_2_capacity_mw_generator_mode"],
            tf["installed_capacity_mw_pump_mode"],
            tf["startup_complex_capacity_mw_pump_mode"],
            tf["stage_1_capacity_mw_pump_mode"],
            tf["stage_2_capacity_mw_pump_mode"],
            tf["units_count"],
            tf["unit_capacity_mw_generator_mode"],
            tf["unit_capacity_mw_pump_mode"],
            tf["hydro_turbine_type"],
            tf["construction_period_years"],
        ]
        vals.extend(
            [
                tf["generation_average_multiyear_million_kwh"],
                tf["generation_average_multiyear_million_kwh_stage_1"],
                tf["generation_average_multiyear_million_kwh_stage_2"],
                tf["annual_charging_electricity_consumption_million_kwh"],
                tf["annual_charging_electricity_consumption_million_kwh_stage_1"],
                tf["annual_charging_electricity_consumption_million_kwh_stage_2"],
                tf["ccium_turbine_mode"],
                tf["ccium_pump_mode"],
            ]
        )
        for i in range(1, 13):
            vals.append(tf[f"construction_increment_year_{i:02d}_mw"])
        vals.extend(
            [
                tf["capital_cost_wo_pir_total_million_rub"],
                tf["year_capital_cost_wo_pir_total"],
                tf["capital_cost_wo_pir_ges_with_reservoir_million_rub"],
                tf["year_capital_cost_wo_pir_ges_with_reservoir"],
                tf["capital_cost_wo_pir_svm_million_rub"],
                tf["year_capital_cost_wo_pir_svm"],
                tf["specific_semifixed_operating_costs_thous_rub_per_kw"],
                tf["year_specific_semifixed_operating_costs"],
                tf["specific_capital_investment_thous_rub_per_kw"],
                tf["year_specific_capital_investment"],
                tf["note"],
            ]
        )
        note_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        note_col = 46
        for col, val in enumerate(vals, start=3):
            cell = ws.cell(row=rn, column=col, value=val)
            cell.border = thin
            cell.alignment = note_align if col == note_col else center

    for g in tep_groups:
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=ncols)
        c = ws.cell(row=row_num, column=1, value=g.get("type_label", ""))
        c.font = Font(bold=True, size=14)
        c.alignment = center
        c.fill = group_fill
        c.border = thin
        _apply_border_to_range(ws, row_num, row_num, 1, ncols, thin)
        row_num += 1

        tep_rows = g.get("tep_rows") or []
        if not tep_rows:
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=ncols)
            c = ws.cell(row=row_num, column=1, value="Нет данных")
            c.alignment = center
            c.border = thin
            _apply_border_to_range(ws, row_num, row_num, 1, ncols, thin)
            row_num += 1
            continue

        for tep_row in tep_rows:
            tf = dict(tep_row["tep_full"])
            ws.cell(row=row_num, column=1, value=tep_row["idx"])
            ws.cell(row=row_num, column=1).border = thin
            ws.cell(row=row_num, column=1).alignment = center

            if tep_row["show_station_cell"]:
                rs = int(tep_row["station_rowspan"] or 1)
                ws.cell(row=row_num, column=2, value=tep_row["site_name"])
                if rs > 1:
                    ws.merge_cells(
                        start_row=row_num,
                        start_column=2,
                        end_row=row_num + rs - 1,
                        end_column=2,
                    )
                c2 = ws.cell(row=row_num, column=2)
                c2.border = thin
                c2.alignment = left
            _write_tf_row(row_num, tf)
            row_num += 1

    if not tep_groups:
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=ncols)
        c = ws.cell(row=row_num, column=1, value="Нет данных")
        c.alignment = center
        c.border = thin
        _apply_border_to_range(ws, row_num, row_num, 1, ncols, thin)

    widths = [6, 28] + [12] * 15 + [12] * 3 + [12] * 3 + [8] * 12 + [11, 10, 11, 10, 11, 10] + [12, 10, 12, 10] + [32]
    for i, w in enumerate(widths, start=1):
        if i <= ncols:
            ws.column_dimensions[get_column_letter(i)].width = min(w, 45)
