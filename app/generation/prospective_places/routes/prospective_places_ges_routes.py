# -*- coding: utf-8 -*-
"""Маршруты для перспективных площадок ГЭС (дубликат логики АЭС)."""
from . import prospective_places_bp
from .prospective_places_routes import (
    _annotate_places_with_linked_station_flag,
    _assign_prospective_place_station,
    _get_station_link_context,
    _navigation_redirect_to_station,
    _prospective_place_can_edit,
    no_compress,
)
from flask import jsonify, render_template, send_file, request, redirect, url_for, flash, g
from flask_login import login_required, current_user
from sqlalchemy.orm import joinedload
from io import BytesIO
from datetime import datetime

from app.extensions import db
from app.logs.services.logging_service import log_to_db
from app.common.services.get_services.years.years_get_services import (
    get_ges_tep_current_price_year_number,
)
from app.generation.prospective_places.services.tep_capital_cost_current_year_services import (
    apply_derived_specific_capital_investment_to_ges_tep_source_row,
    fmt_specific_capital_investment_thous_rub_per_kw_derived,
    get_tep_price_coefficient_by_year_map,
    resolve_year_id_for_calendar_year_number,
)
from app.generation.prospective_places.services.ges.machine_prospective_place_ges_sort import (
    sort_machines_by_station_block_number,
)
from app.generation.prospective_places.services.ges.prospective_places_filters_ges_services import (
    build_ges_stations_grouped_by_place_type,
)
from app.generation.prospective_places.forms.decimal_input_display import (
    decimal_for_form_strip_trailing_zeros,
    format_construction_period_years_display,
    normalize_construction_period_years,
)


def _ges_tep_year_select_choices():
    from app.common.services.get_services.years.years_get_services import get_year_list_full

    years = get_year_list_full()
    return [("", "— не указано —")] + [(str(y.id), str(y.number)) for y in years]


def _ges_specific_capital_source_year_display(tep_row) -> str:
    """Календарный год исходных капзатрат «ГЭС (с водохранилищем)» — для подписи у удельных (не год «текущий» из справочника)."""
    if tep_row is None:
        return "—"
    y = getattr(tep_row, "year_capital_cost_wo_pir_ges_with_reservoir", None)
    if y is not None and getattr(y, "number", None) is not None:
        return str(y.number)
    return "—"


def _ges_tep_derive_template_context():
    """Контекст для расчёта удельных капвложений на карточке ГЭС (как в таблице ТЭП)."""
    from app.common.services.get_services.years.years_get_services import get_year_list_full

    ty = get_ges_tep_current_price_year_number()
    coeff = get_tep_price_coefficient_by_year_map()
    coeff_json = {str(k): float(v) for k, v in coeff.items()}
    years_full = get_year_list_full()
    yid_to_num = {
        str(y.id): int(y.number) for y in years_full if y is not None and y.number is not None
    }
    ty_id = resolve_year_id_for_calendar_year_number(ty)
    return {
        "ges_tep_target_price_year_display": str(ty) if ty is not None else "—",
        "ges_tep_price_derive_config": {
            "targetYear": ty,
            "targetYearId": ty_id,
            "coeffByYear": coeff_json,
            "yearIdToNumber": yid_to_num,
        },
    }


def _planned_capacity_mw_int_from_installed_str(value) -> int | None:
    """Целое МВт для поля площадки из строки «Установленная генерирующая мощность, МВт» (ТЭП)."""
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    s = s.replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return int(round(float(s)))
    except (ValueError, TypeError, OverflowError):
        return None


def _set_ges_tep_year_choices(form):
    ch = _ges_tep_year_select_choices()
    form.id_year_capital_cost_wo_pir_total.choices = ch
    form.id_year_capital_cost_wo_pir_ges_with_reservoir.choices = ch
    form.id_year_capital_cost_wo_pir_svm.choices = ch
    form.id_year_specific_semifixed_operating_costs.choices = ch


def _ges_tep_price_coeff_can_edit() -> bool:
    if not current_user.is_authenticated:
        return False
    edit_roles = ("admin", "generation-admin", "generation-editor")
    return any(role in current_user.role_names for role in edit_roles)


def _ges_tep_price_coeff_payload() -> dict:
    from app.generation.prospective_places.services.tep_capital_cost_current_year_services import (
        get_price_coefficients_snapshot,
    )

    data = get_price_coefficients_snapshot()
    data["can_edit"] = _ges_tep_price_coeff_can_edit()
    return data


def _load_prospective_places_ges_stations_and_filters():
    """Станции и контекст фильтров по request.args (список и страницы ТЭП)."""
    from app.generation.prospective_places.models import StationProspectivePlaceGES, ProspectivePlaceGesTepSource
    from app.refdata.models.energy_systems.regional_energy_system_model import (
        RegionalEnergySystem,
    )
    from app.generation.prospective_places.services.ges.prospective_places_filters_ges_services import (
        extract_prospective_places_filters,
        apply_prospective_places_ges_filters,
        get_prospective_places_ges_filter_context,
        apply_prospective_places_ges_station_order,
    )

    filters = extract_prospective_places_filters(request.args)
    query = (
        StationProspectivePlaceGES.query
        .options(
            joinedload(StationProspectivePlaceGES.regional_district),
            joinedload(StationProspectivePlaceGES.prospective_place_type_ges),
            joinedload(StationProspectivePlaceGES.regional_energy_system).joinedload(
                RegionalEnergySystem.union_energy_system
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.prospective_place_type_ges
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.station_prospective_place_ges
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.year_capital_cost_wo_pir_total
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.year_capital_cost_wo_pir_ges_with_reservoir
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.year_capital_cost_wo_pir_svm
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.year_specific_semifixed_operating_costs
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.year_specific_capital_investment
            ),
        )
    )
    query = apply_prospective_places_ges_station_order(query)
    query = apply_prospective_places_ges_filters(query, filters)
    stations = query.all()
    for station in stations:
        sort_machines_by_station_block_number(station)
    filter_context = get_prospective_places_ges_filter_context(filters)
    return stations, filters, filter_context


def _log_prospective_ges_page_view(action_title: str) -> None:
    log_to_db(
        current_user,
        action_title,
        details=f"url={request.full_path}",
        entity_type="prospective_place_ges",
    )


@prospective_places_bp.route("/ges/")
@login_required
def prospective_places_ges():
    """Характеристика актуальных площадок размещения новых ГЭС (список)."""
    stations, filters, filter_context = _load_prospective_places_ges_stations_and_filters()
    _annotate_places_with_linked_station_flag(stations)
    ges_place_type_groups = build_ges_stations_grouped_by_place_type(stations)
    _log_prospective_ges_page_view("Открыта страница: перспективные площадки ГЭС (список)")

    return render_template(
        "generation/prospective_places/ges/prospective_places_ges.html",
        stations=stations,
        ges_place_type_groups=ges_place_type_groups,
        total_count=len(stations),
        has_active_filters=any(
            filters.get(k)
            for k in [
                "energy_system_type_filter",
                "union_energy_system_filter",
                "regional_energy_system_filter",
                "federal_district_filter",
                "regional_district_filter",
                "site_name_filter",
                "prospective_place_type_filter",
            ]
        ),
        **filter_context,
    )


@prospective_places_bp.route("/ges/tep_main/")
@login_required
def prospective_places_ges_tep_main():
    """Перечень исходных ТЭП по всем площадкам ГЭС (как список, полная таблица)."""
    from app.common.services.get_services.years.years_get_services import (
        get_ges_tep_current_price_year_number,
    )
    from app.generation.prospective_places.services.tep_capital_cost_current_year_services import (
        get_tep_main_radio_price_year_number,
        tep_main_rows_have_mixed_source_price_years,
    )
    from app.generation.prospective_places.services.ges.prospective_places_ges_tep_view_services import (
        build_tep_groups_by_station_place_type,
        collect_all_ges_tep_rows,
    )

    stations, filters, filter_context = _load_prospective_places_ges_stations_and_filters()
    ges_place_type_groups = build_ges_stations_grouped_by_place_type(stations)
    all_tep_rows = []
    for g in ges_place_type_groups:
        all_tep_rows.extend(collect_all_ges_tep_rows(g["stations"]))
    current_y = get_ges_tep_current_price_year_number()
    tep_price_label_year = get_tep_main_radio_price_year_number(
        all_tep_rows,
        current_target_year=current_y,
    )
    tep_main_force_source_only = tep_main_rows_have_mixed_source_price_years(all_tep_rows)
    tep_groups, _ = build_tep_groups_by_station_place_type(
        ges_place_type_groups,
        tep_main_current_year_prices=True,
        tep_main_scaling_target_year=tep_price_label_year,
    )
    _log_prospective_ges_page_view("Открыта страница: перспективные площадки ГЭС (перечень ТЭП)")

    return render_template(
        "generation/prospective_places/ges/prospective_places_ges_tep.html",
        **filter_context,
        tep_groups=tep_groups,
        page_title="Перечень исходных ТЭП перспективных площадок ГЭС",
        page_subtitle="",
        toolbar_mode="tep_main",
        has_active_filters=any(
            filters.get(k)
            for k in [
                "energy_system_type_filter",
                "union_energy_system_filter",
                "regional_energy_system_filter",
                "federal_district_filter",
                "regional_district_filter",
                "site_name_filter",
                "prospective_place_type_filter",
            ]
        ),
        tep_price_label_year=tep_price_label_year,
        tep_main_force_source_only=tep_main_force_source_only,
        ges_tep_price_coeff_can_edit=_ges_tep_price_coeff_can_edit(),
    )


@prospective_places_bp.route("/ges/tep_price_conversion_coefficients/", methods=["GET", "POST"])
@login_required
def ges_tep_price_conversion_coefficients_json():
    """JSON: коэффициенты перевода цен и справочник годов; POST — новая запись (редакторы)."""
    from app.generation.prospective_places.services.tep_capital_cost_current_year_services import (
        create_price_conversion_coefficient,
    )

    if request.method == "GET":
        return jsonify(_ges_tep_price_coeff_payload())

    if not _ges_tep_price_coeff_can_edit():
        return jsonify({"error": "Недостаточно прав для изменения."}), 403

    payload = request.get_json(silent=True) or {}
    try:
        id_year = int(payload["id_year"])
    except (KeyError, TypeError, ValueError):
        return jsonify({"error": "Укажите корректный идентификатор года (id_year)."}), 400

    try:
        create_price_conversion_coefficient(id_year, payload.get("coefficient"))
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    out = _ges_tep_price_coeff_payload()
    out["ok"] = True
    return jsonify(out)


@prospective_places_bp.route("/ges/tep_price_conversion_coefficients/<int:row_id>/", methods=["PUT", "DELETE"])
@login_required
def ges_tep_price_conversion_coefficients_item(row_id: int):
    """Изменение или удаление строки коэффициента (редакторы)."""
    from app.generation.prospective_places.services.tep_capital_cost_current_year_services import (
        delete_price_conversion_coefficient,
        update_price_conversion_coefficient_from_request,
    )

    if not _ges_tep_price_coeff_can_edit():
        return jsonify({"error": "Недостаточно прав для изменения."}), 403

    if request.method == "DELETE":
        try:
            delete_price_conversion_coefficient(row_id)
        except LookupError:
            return jsonify({"error": "Запись не найдена."}), 404
        out = _ges_tep_price_coeff_payload()
        out["ok"] = True
        return jsonify(out)

    payload = request.get_json(silent=True) or {}
    try:
        update_price_conversion_coefficient_from_request(row_id, payload)
    except LookupError:
        return jsonify({"error": "Запись не найдена."}), 404
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    out = _ges_tep_price_coeff_payload()
    out["ok"] = True
    return jsonify(out)


@prospective_places_bp.route("/ges/tep_reserve/")
@login_required
def prospective_places_ges_tep_reserve():
    """Старый URL: перечень ТЭП единый для всех площадок ГЭС."""
    dest = url_for("prospective_places_bp.prospective_places_ges_tep_main")
    qs = request.query_string.decode("utf-8")
    if qs:
        dest = f"{dest}?{qs}"
    return redirect(dest)


@prospective_places_bp.route("/ges/<int:id>/", methods=["GET", "POST"])
@login_required
def prospective_place_ges_details(id):
    """Карточка перспективной площадки ГЭС (заполнение и редактирование полей)."""
    from app.generation.prospective_places.models import (
        StationProspectivePlaceGES,
        ProspectivePlaceGesTepSource,
        ProspectivePlaceTypeGES,
    )
    from app.generation.prospective_places.forms import (
        ProspectivePlaceGESEditForm,
        ProspectivePlaceGesTepSourceEditForm,
    )
    from app.refdata.models.energy_systems.regional_energy_system_model import (
        RegionalEnergySystem,
    )
    from app.common.services.get_services.territories.regional_district_get_services import (
        get_regional_district_list_full,
    )
    from app.common.services.get_services.energy_systems.regional_energy_system_get_services import (
        get_regional_energy_system_choices,
        get_res_to_ues_id_map,
    )
    from app.common.services.get_services.energy_systems.union_energy_system_get_services import (
        get_union_energy_systems_map,
    )

    place = (
        StationProspectivePlaceGES.query
        .options(
            joinedload(StationProspectivePlaceGES.regional_district),
            joinedload(StationProspectivePlaceGES.prospective_place_type_ges),
            joinedload(StationProspectivePlaceGES.regional_energy_system).joinedload(
                RegionalEnergySystem.union_energy_system
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.prospective_place_type_ges
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.year_capital_cost_wo_pir_total
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.year_capital_cost_wo_pir_ges_with_reservoir
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.year_capital_cost_wo_pir_svm
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.year_specific_semifixed_operating_costs
            ),
            joinedload(StationProspectivePlaceGES.ges_tep_source_indicators).joinedload(
                ProspectivePlaceGesTepSource.year_specific_capital_investment
            ),
        )
        .get_or_404(id)
    )
    sort_machines_by_station_block_number(place)

    can_edit = _prospective_place_can_edit()

    form = ProspectivePlaceGESEditForm()
    rd_list = get_regional_district_list_full()
    form.id_regional_district.choices = [("", "— Не указано —")] + [
        (str(rd_id), name) for rd_id, name in rd_list
    ]
    res_list = get_regional_energy_system_choices()
    form.id_regional_energy_system.choices = [("", "— Не указано —")] + [
        (str(res_id), name) for res_id, name in res_list
    ]
    ppt_list = ProspectivePlaceTypeGES.query.order_by(ProspectivePlaceTypeGES.name).all()
    form.id_prospective_place_type_ges.choices = [("", "— Не указано —")] + [
        (str(p.id), p.name) for p in ppt_list
    ]

    # Карта для автоподстановки ОЭС по выбранной РЭС (на лету)
    res_to_ues = get_res_to_ues_id_map()
    ues_names = get_union_energy_systems_map()
    res_auto_map = {"": "Нет данных"}
    for res_id, _ in res_list:
        ues_id = res_to_ues.get(res_id)
        ues_name = ues_names.get(ues_id, "Нет данных") if ues_id else "Нет данных"
        res_auto_map[str(res_id)] = ues_name

    if request.method == "GET":
        form.site_name.data = place.site_name
        form.id_regional_district.data = place.id_regional_district
        form.id_regional_energy_system.data = place.id_regional_energy_system
        form.geo_location.data = place.geo_location
        form.project_initiator.data = place.project_initiator
        form.water_body.data = place.water_body
        form.regulation_type.data = place.regulation_type
        form.planned_capacity_mw.data = place.planned_capacity_mw
        form.general_scheme_commissioning_period.data = place.general_scheme_commissioning_period
        form.construction_period_years.data = place.construction_period_years
        form.id_prospective_place_type_ges.data = place.id_prospective_place_type_ges
        form.display_order.data = place.display_order

    tep_form_entries = []
    indicators = list(place.ges_tep_source_indicators or [])
    _tep_idx = 0

    def _ges_tep_form_factory():
        nonlocal _tep_idx
        pfx = f"tep_r{_tep_idx}" if can_edit else None
        _tep_idx += 1
        if pfx:
            return ProspectivePlaceGesTepSourceEditForm(prefix=pfx), pfx
        return ProspectivePlaceGesTepSourceEditForm(), None

    if not indicators:
        if can_edit:
            f, pfx = _ges_tep_form_factory()
            tep_form_entries.append(
                {
                    "form": f,
                    "form_prefix": pfx,
                    "machine_id": 0,
                    "ges_specific_capital_year_display": "—",
                }
            )
    else:
        for tr in indicators:
            f, pfx = _ges_tep_form_factory()
            if request.method == "GET" or not can_edit:
                _populate_ges_tep_form_from_row(f, tr)
            tep_form_entries.append(
                {
                    "form": f,
                    "form_prefix": pfx,
                    "machine_id": tr.id,
                    "ges_specific_capital_year_display": _ges_specific_capital_source_year_display(tr),
                }
            )
        if can_edit and request.args.get("new_tep") == "1":
            f, pfx = _ges_tep_form_factory()
            tep_form_entries.append(
                {
                    "form": f,
                    "form_prefix": pfx,
                    "machine_id": 0,
                    "ges_specific_capital_year_display": "—",
                }
            )

    for entry in tep_form_entries:
        _set_ges_tep_year_choices(entry["form"])

    if request.method == "POST" and can_edit:
        place_ok = form.validate()
        tep_ok = [entry["form"].validate() for entry in tep_form_entries]
        if place_ok and all(tep_ok):
            try:
                place.site_name = (form.site_name.data or "").strip() or None
                place.id_regional_district = form.id_regional_district.data
                place.id_regional_energy_system = form.id_regional_energy_system.data
                place.geo_location = (form.geo_location.data or "").strip() or None
                place.project_initiator = (form.project_initiator.data or "").strip() or None
                place.water_body = (form.water_body.data or "").strip() or None
                place.regulation_type = (form.regulation_type.data or "").strip() or None
                place.planned_capacity_mw = form.planned_capacity_mw.data
                place.general_scheme_commissioning_period = (
                    (form.general_scheme_commissioning_period.data or "").strip() or None
                )
                place.construction_period_years = normalize_construction_period_years(
                    form.construction_period_years.data
                )
                place.id_prospective_place_type_ges = form.id_prospective_place_type_ges.data
                place.display_order = form.display_order.data
                for t in place.ges_tep_source_indicators:
                    t.id_prospective_place_type_ges = place.id_prospective_place_type_ges

                for entry in tep_form_entries:
                    tform = entry["form"]
                    machine_id = entry["machine_id"]
                    if machine_id:
                        tep_row = ProspectivePlaceGesTepSource.query.filter(
                            ProspectivePlaceGesTepSource.id == machine_id,
                            ProspectivePlaceGesTepSource.id_station_prospective_place_ges
                            == place.id,
                        ).first()
                        if not tep_row:
                            raise ValueError("Запись перечня ТЭП не найдена (несогласованные данные).")
                    else:
                        tep_row = None
                    if machine_id and tep_row is not None:
                        _apply_ges_tep_form_to_row(tep_row, tform, place)
                    else:
                        tep_new = ProspectivePlaceGesTepSource(
                            id_station_prospective_place_ges=place.id,
                        )
                        _apply_ges_tep_form_to_row(tep_new, tform, place)
                        db.session.add(tep_new)

                from app.generation.prospective_places.models.prospective_place_hydro_energy_forecast_model import (
                    PLACE_KIND_GES,
                )
                from app.generation.prospective_places.services.hydro_energy_forecast_services import (
                    save_hydro_forecast_from_form,
                )

                hydro_changes = save_hydro_forecast_from_form(
                    current_user,
                    place_kind=PLACE_KIND_GES,
                    place_id=place.id,
                    form_data=request.form,
                    entity_type="prospective_place_ges",
                    commit=False,
                )

                db.session.commit()
                details = f"id={place.id}; site_name={place.site_name!r}"
                if hydro_changes:
                    details = f"{details}; прогноз выработки: {'; '.join(hydro_changes)}"
                log_to_db(
                    current_user,
                    "Изменена перспективная площадка ГЭС (карточка, перечень ТЭП и прогноз выработки)",
                    details=details,
                    entity_type="prospective_place_ges",
                    entity_id=place.id,
                )
                flash(
                    "Перспективная площадка, перечень ТЭП и прогноз выработки успешно сохранены.",
                    "success",
                )
                return redirect(url_for("prospective_places_bp.prospective_place_ges_details", id=place.id))
            except Exception as e:
                db.session.rollback()
                flash(f"Ошибка при сохранении: {str(e)}", "danger")
        else:
            flash("Исправьте ошибки в полях и повторите сохранение.", "danger")

    ctx = _ges_tep_derive_template_context()
    station_link_ctx = _get_station_link_context(place)
    from app.generation.prospective_places.models.prospective_place_hydro_energy_forecast_model import (
        PLACE_KIND_GES,
    )
    from app.generation.prospective_places.services.hydro_energy_forecast_services import (
        build_hydro_forecast_view,
    )

    return render_template(
        "generation/prospective_places/ges/prospective_place_ges_details.html",
        place=place,
        form=form,
        can_edit=can_edit,
        can_edit_hydro_forecast=can_edit,
        res_auto_map=res_auto_map,
        tep_form_entries=tep_form_entries,
        hydro_forecast=build_hydro_forecast_view(PLACE_KIND_GES, place.id),
        **station_link_ctx,
        **ctx,
    )


@prospective_places_bp.route("/ges/<int:id>/link_station/", methods=["POST"])
@login_required
def prospective_place_ges_link_station(id):
    from app.generation.prospective_places.models import StationProspectivePlaceGES

    if not _prospective_place_can_edit():
        flash("Недостаточно прав для изменения связи с электростанцией.", "warning")
        return redirect(url_for("prospective_places_bp.prospective_place_ges_details", id=id))

    place = StationProspectivePlaceGES.query.get_or_404(id)
    station_id_raw = (request.form.get("linked_station_id") or "").strip()

    try:
        station_id = int(station_id_raw) if station_id_raw else None
    except ValueError:
        flash("Некорректно выбрана электростанция.", "danger")
        return redirect(url_for("prospective_places_bp.prospective_place_ges_details", id=id))

    try:
        station = _assign_prospective_place_station(place, station_id)
        db.session.commit()
        if station is None:
            flash("Связь с электростанцией снята.", "success")
        else:
            flash("Связь с электростанцией сохранена.", "success")
        log_to_db(
            current_user,
            "Изменена связь перспективной площадки ГЭС с электростанцией",
            details=(
                f"id_place={place.id}; site_name={place.site_name!r}; "
                f"station_id={(station.id if station else None)}"
            ),
            entity_type="prospective_place_ges",
            entity_id=place.id,
        )
    except (LookupError, ValueError) as e:
        db.session.rollback()
        flash(str(e), "danger")
    except Exception as e:
        db.session.rollback()
        flash(f"Ошибка при сохранении связи: {str(e)}", "danger")

    return redirect(url_for("prospective_places_bp.prospective_place_ges_details", id=id))


@prospective_places_bp.route("/ges/<int:id>/go_station/", methods=["GET"])
@login_required
def prospective_place_ges_go_station(id):
    from app.generation.prospective_places.models import StationProspectivePlaceGES

    place = StationProspectivePlaceGES.query.get_or_404(id)
    return _navigation_redirect_to_station(
        place,
        back_endpoint="prospective_places_bp.prospective_place_ges_details",
        back_id=id,
    )


def _populate_ges_tep_form_from_row(form, tep_row):
    """Заполняет ProspectivePlaceGesTepSourceEditForm из ProspectivePlaceGesTepSource."""
    if not tep_row:
        return
    form.installed_capacity_mw.data = tep_row.installed_capacity_mw
    form.stage_1_capacity_mw.data = tep_row.stage_1_capacity_mw
    form.stage_2_capacity_mw.data = tep_row.stage_2_capacity_mw
    form.startup_complex_capacity_mw.data = tep_row.startup_complex_capacity_mw
    form.units_count.data = tep_row.units_count
    form.unit_capacity_mw.data = tep_row.unit_capacity_mw
    form.hydro_turbine_type.data = tep_row.hydro_turbine_type
    form.construction_period_years.data = tep_row.construction_period_years
    form.construction_increment_year_01_mw.data = tep_row.construction_increment_year_01_mw
    form.construction_increment_year_02_mw.data = tep_row.construction_increment_year_02_mw
    form.construction_increment_year_03_mw.data = tep_row.construction_increment_year_03_mw
    form.construction_increment_year_04_mw.data = tep_row.construction_increment_year_04_mw
    form.construction_increment_year_05_mw.data = tep_row.construction_increment_year_05_mw
    form.construction_increment_year_06_mw.data = tep_row.construction_increment_year_06_mw
    form.construction_increment_year_07_mw.data = tep_row.construction_increment_year_07_mw
    form.construction_increment_year_08_mw.data = tep_row.construction_increment_year_08_mw
    form.construction_increment_year_09_mw.data = tep_row.construction_increment_year_09_mw
    form.construction_increment_year_10_mw.data = tep_row.construction_increment_year_10_mw
    form.construction_increment_year_11_mw.data = tep_row.construction_increment_year_11_mw
    form.construction_increment_year_12_mw.data = tep_row.construction_increment_year_12_mw
    form.specific_semifixed_operating_costs_thous_rub_per_kw.data = (
        tep_row.specific_semifixed_operating_costs_thous_rub_per_kw
    )
    form.id_year_specific_semifixed_operating_costs.data = (
        tep_row.id_year_specific_semifixed_operating_costs
    )
    form.generation_average_multiyear_million_kwh.data = tep_row.generation_average_multiyear_million_kwh
    form.generation_medium_water_50pct_million_kwh.data = tep_row.generation_medium_water_50pct_million_kwh
    form.generation_medium_water_management_year.data = tep_row.generation_medium_water_management_year
    form.generation_low_water_95pct_million_kwh.data = tep_row.generation_low_water_95pct_million_kwh
    form.generation_low_water_management_year.data = tep_row.generation_low_water_management_year
    form.capital_cost_wo_pir_total_million_rub.data = decimal_for_form_strip_trailing_zeros(
        tep_row.capital_cost_wo_pir_total_million_rub
    )
    form.id_year_capital_cost_wo_pir_total.data = tep_row.id_year_capital_cost_wo_pir_total
    form.capital_cost_wo_pir_ges_with_reservoir_million_rub.data = (
        decimal_for_form_strip_trailing_zeros(
            tep_row.capital_cost_wo_pir_ges_with_reservoir_million_rub
        )
    )
    form.id_year_capital_cost_wo_pir_ges_with_reservoir.data = (
        tep_row.id_year_capital_cost_wo_pir_ges_with_reservoir
    )
    form.capital_cost_wo_pir_svm_million_rub.data = decimal_for_form_strip_trailing_zeros(
        tep_row.capital_cost_wo_pir_svm_million_rub
    )
    form.id_year_capital_cost_wo_pir_svm.data = tep_row.id_year_capital_cost_wo_pir_svm
    form.note.data = tep_row.note
    coeff = get_tep_price_coefficient_by_year_map()
    ty = get_ges_tep_current_price_year_number()
    s_derived = fmt_specific_capital_investment_thous_rub_per_kw_derived(
        tep_row,
        installed_capacity_attr="installed_capacity_mw",
        target_year=ty,
        coeff_by_year=coeff,
    )
    form.specific_capital_investment_thous_rub_per_kw.data = None if s_derived == "—" else s_derived
    form.id_year_specific_capital_investment.data = resolve_year_id_for_calendar_year_number(ty)


def _apply_ges_tep_form_to_row(row, form, place):
    """Заполняет ProspectivePlaceGesTepSource из ProspectivePlaceGesTepSourceEditForm.

    Тип площадки задается на карточке площадки; при сохранении ТЭП копируется с площадки.
    """
    def _s(val):
        return (val or "").strip() or None

    row.id_prospective_place_type_ges = place.id_prospective_place_type_ges
    row.installed_capacity_mw = _s(form.installed_capacity_mw.data)
    row.stage_1_capacity_mw = _s(form.stage_1_capacity_mw.data)
    row.stage_2_capacity_mw = _s(form.stage_2_capacity_mw.data)
    row.startup_complex_capacity_mw = _s(form.startup_complex_capacity_mw.data)
    row.units_count = form.units_count.data
    row.unit_capacity_mw = _s(form.unit_capacity_mw.data)
    row.hydro_turbine_type = _s(form.hydro_turbine_type.data)
    row.construction_period_years = normalize_construction_period_years(
        form.construction_period_years.data
    )
    row.construction_increment_year_01_mw = _s(form.construction_increment_year_01_mw.data)
    row.construction_increment_year_02_mw = _s(form.construction_increment_year_02_mw.data)
    row.construction_increment_year_03_mw = _s(form.construction_increment_year_03_mw.data)
    row.construction_increment_year_04_mw = _s(form.construction_increment_year_04_mw.data)
    row.construction_increment_year_05_mw = _s(form.construction_increment_year_05_mw.data)
    row.construction_increment_year_06_mw = _s(form.construction_increment_year_06_mw.data)
    row.construction_increment_year_07_mw = _s(form.construction_increment_year_07_mw.data)
    row.construction_increment_year_08_mw = _s(form.construction_increment_year_08_mw.data)
    row.construction_increment_year_09_mw = _s(form.construction_increment_year_09_mw.data)
    row.construction_increment_year_10_mw = _s(form.construction_increment_year_10_mw.data)
    row.construction_increment_year_11_mw = _s(form.construction_increment_year_11_mw.data)
    row.construction_increment_year_12_mw = _s(form.construction_increment_year_12_mw.data)
    row.specific_semifixed_operating_costs_thous_rub_per_kw = _s(
        form.specific_semifixed_operating_costs_thous_rub_per_kw.data
    )
    row.id_year_specific_semifixed_operating_costs = (
        form.id_year_specific_semifixed_operating_costs.data
    )
    row.generation_average_multiyear_million_kwh = _s(form.generation_average_multiyear_million_kwh.data)
    row.generation_medium_water_50pct_million_kwh = _s(form.generation_medium_water_50pct_million_kwh.data)
    row.generation_medium_water_management_year = _s(form.generation_medium_water_management_year.data)
    row.generation_low_water_95pct_million_kwh = _s(form.generation_low_water_95pct_million_kwh.data)
    row.generation_low_water_management_year = _s(form.generation_low_water_management_year.data)
    row.capital_cost_wo_pir_total_million_rub = form.capital_cost_wo_pir_total_million_rub.data
    row.id_year_capital_cost_wo_pir_total = form.id_year_capital_cost_wo_pir_total.data
    row.capital_cost_wo_pir_ges_with_reservoir_million_rub = (
        form.capital_cost_wo_pir_ges_with_reservoir_million_rub.data
    )
    row.id_year_capital_cost_wo_pir_ges_with_reservoir = (
        form.id_year_capital_cost_wo_pir_ges_with_reservoir.data
    )
    row.capital_cost_wo_pir_svm_million_rub = form.capital_cost_wo_pir_svm_million_rub.data
    row.id_year_capital_cost_wo_pir_svm = form.id_year_capital_cost_wo_pir_svm.data
    apply_derived_specific_capital_investment_to_ges_tep_source_row(row)
    row.note = _s(form.note.data)
    place.planned_capacity_mw = _planned_capacity_mw_int_from_installed_str(row.installed_capacity_mw)


@prospective_places_bp.route("/ges/<int:place_id>/machine/<int:machine_id>/", methods=["GET", "POST"])
@login_required
def machine_prospective_place_ges_details(place_id, machine_id):
    """Добавление/редактирование записи перечня ТЭП (machine_id=0 — добавить)."""
    from app.generation.prospective_places.models import (
        StationProspectivePlaceGES,
        ProspectivePlaceGesTepSource,
    )
    from app.generation.prospective_places.forms import ProspectivePlaceGesTepSourceEditForm

    edit_roles = ["admin", "generation-admin", "generation-editor"]
    can_edit = current_user.is_authenticated and any(
        role in current_user.role_names for role in edit_roles
    )
    if not can_edit:
        flash("Недостаточно прав для редактирования записи перечня ТЭП.", "warning")
        return redirect(url_for("prospective_places_bp.prospective_place_ges_details", id=place_id))

    place = StationProspectivePlaceGES.query.get_or_404(place_id)
    tep_row = ProspectivePlaceGesTepSource.query.filter(
        ProspectivePlaceGesTepSource.id == machine_id,
        ProspectivePlaceGesTepSource.id_station_prospective_place_ges == place_id,
    ).first() if machine_id else None

    if machine_id and not tep_row:
        flash("Запись перечня ТЭП не найдена.", "danger")
        return redirect(url_for("prospective_places_bp.prospective_place_ges_details", id=place_id))

    if request.method == "GET":
        return redirect(url_for("prospective_places_bp.prospective_place_ges_details", id=place_id))

    form = ProspectivePlaceGesTepSourceEditForm()
    _set_ges_tep_year_choices(form)

    if form.validate_on_submit() and can_edit:
        try:
            is_new_tep = tep_row is None
            if tep_row:
                _apply_ges_tep_form_to_row(tep_row, form, place)
            else:
                tep_row = ProspectivePlaceGesTepSource(
                    id_station_prospective_place_ges=place_id,
                )
                _apply_ges_tep_form_to_row(tep_row, form, place)
                db.session.add(tep_row)
            db.session.commit()
            log_to_db(
                current_user,
                "Создана позиция перечня ТЭП (ГЭС)" if is_new_tep else "Изменена позиция перечня ТЭП (ГЭС)",
                details=(
                    f"id_tep={tep_row.id}; id_place={place_id}; site_name={place.site_name!r}"
                ),
                entity_type="ges_tep_source",
                entity_id=tep_row.id,
            )
            flash("Запись перечня ТЭП успешно сохранена.", "success")
            return redirect(url_for("prospective_places_bp.prospective_place_ges_details", id=place_id))
        except Exception as e:
            db.session.rollback()
            flash(f"Ошибка при сохранении: {str(e)}", "danger")

    ctx = _ges_tep_derive_template_context()
    return render_template(
        "generation/prospective_places/ges/machine_prospective_place_ges_details.html",
        place=place,
        tep_row=tep_row,
        form=form,
        ges_specific_capital_year_display=_ges_specific_capital_source_year_display(tep_row),
        **ctx,
    )


@prospective_places_bp.route("/ges/add/", methods=["GET", "POST"])
@login_required
def add_prospective_place_ges():
    """Добавление новой перспективной площадки ГЭС."""
    edit_roles = ["admin", "generation-admin", "generation-editor"]
    can_edit = current_user.is_authenticated and any(
        role in current_user.role_names for role in edit_roles
    )
    if not can_edit:
        flash("Недостаточно прав для добавления перспективной площадки.", "warning")
        return redirect(url_for("prospective_places_bp.prospective_places_ges"))

    from app.generation.prospective_places.forms import ProspectivePlaceGESAddForm
    from app.generation.prospective_places.models import StationProspectivePlaceGES
    from app.common.services.get_services.territories.regional_district_get_services import (
        get_regional_district_list_full,
    )

    form = ProspectivePlaceGESAddForm()
    rd_list = get_regional_district_list_full()
    form.id_regional_district.choices = [("", "— Выберите субъект РФ —")] + [
        (str(rd_id), name) for rd_id, name in rd_list
    ]

    if form.validate_on_submit():
        site_name = (form.site_name.data or "").strip()
        id_regional_district = form.id_regional_district.data

        # Проверка уникальности по (site_name, id_regional_district)
        q = StationProspectivePlaceGES.query
        if site_name:
            q = q.filter(StationProspectivePlaceGES.site_name == site_name)
        else:
            q = q.filter(
                db.or_(
                    StationProspectivePlaceGES.site_name.is_(None),
                    StationProspectivePlaceGES.site_name == "",
                )
            )
        if id_regional_district is not None:
            q = q.filter(
                StationProspectivePlaceGES.id_regional_district == id_regional_district
            )
        else:
            q = q.filter(StationProspectivePlaceGES.id_regional_district.is_(None))

        if q.first():
            flash(
                "Перспективная площадка с таким наименованием и субъектом РФ уже существует.",
                "danger",
            )
            return render_template(
                "generation/prospective_places/ges/prospective_place_ges_add.html",
                form=form,
            )

        try:
            place = StationProspectivePlaceGES(
                site_name=site_name or None,
                id_regional_district=id_regional_district,
            )
            db.session.add(place)
            db.session.commit()
            log_to_db(
                current_user,
                "Создана перспективная площадка ГЭС",
                details=f"id={place.id}; site_name={place.site_name!r}",
                entity_type="prospective_place_ges",
                entity_id=place.id,
            )
            flash("Перспективная площадка успешно создана.", "success")
            return redirect(url_for("prospective_places_bp.prospective_places_ges"))
        except Exception as e:
            db.session.rollback()
            flash(f"Ошибка при создании: {str(e)}", "danger")

    return render_template(
        "generation/prospective_places/ges/prospective_place_ges_add.html",
        form=form,
    )


@prospective_places_bp.route("/ges/export/", methods=["GET"])
@login_required
@no_compress
def export_prospective_places_ges():
    """Экспорт списка ГЭС: те же группы, строки и 13 колонок, что на /ges/ (без дополнительных листов)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    stations, _filters, _fc = _load_prospective_places_ges_stations_and_filters()

    if request.args.get("export_mode") == "tep_main":
        from app.common.services.get_services.years.years_get_services import (
            get_ges_tep_current_price_year_number,
        )
        from app.generation.prospective_places.services.tep_capital_cost_current_year_services import (
            get_tep_main_radio_price_year_number,
        )
        from app.generation.prospective_places.services.ges.export_prospective_places_ges_tep_main_excel import (
            write_tep_main_screen_sheet,
        )
        from app.generation.prospective_places.services.ges.prospective_places_ges_tep_view_services import (
            build_tep_groups_by_station_place_type,
            collect_all_ges_tep_rows,
        )

        source_prices = request.args.get("tep_main_view") == "source"
        ges_place_type_groups = build_ges_stations_grouped_by_place_type(stations)
        all_tep_rows = []
        for g in ges_place_type_groups:
            all_tep_rows.extend(collect_all_ges_tep_rows(g["stations"]))
        current_y = get_ges_tep_current_price_year_number()
        tep_price_label_year = get_tep_main_radio_price_year_number(
            all_tep_rows,
            current_target_year=current_y,
        )
        tep_groups, _ = build_tep_groups_by_station_place_type(
            ges_place_type_groups,
            tep_main_current_year_prices=not source_prices,
            tep_main_scaling_target_year=tep_price_label_year if not source_prices else None,
        )
        wb = Workbook()
        ws = wb.active
        ws.title = "Перечень ТЭП перспективных площадок ГЭС"
        write_tep_main_screen_sheet(ws, tep_groups, source_prices=source_prices)
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        filename = (
            f"ges_tep_main_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"
        )
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename,
        )

    ges_place_type_groups = build_ges_stations_grouped_by_place_type(stations)

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    header_font = Font(bold=True)
    group_font = Font(bold=True, size=12)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    total_cols = 13
    headers = [
        "№ п/п",
        "Наименование площадки",
        "Инициатор проекта",
        "Планируемая установленная генерирующая мощность ГЭС, МВт",
        "Среднемноголетняя выработка электроэнергии, млн кВтч",
        "Вид регулирования",
        "Объединенная энергосистема (ОЭС)",
        "Субъект Российской Федерации",
        "Водный объект",
        "Географическое расположение площадки",
        "Установленная мощность, МВт",
        (
            "Генеральная схема до 2042 года (Период ввода в эксплуатацию).\n"
            "Распоряжение Правительства от 30.12.2024 №4153-р"
        ),
        (
            "Срок строительства по Протоколу Минэнерго от 25.05.2022 №РГ/07-0003пр "
            "(без учета срока выполнения ПИР- 2 года), лет"
        ),
    ]

    def _ues_name(s):
        return (
            s.regional_energy_system.union_energy_system.name
            if s.regional_energy_system and s.regional_energy_system.union_energy_system
            else None
        ) or "—"

    def _rd_name(s):
        return (s.regional_district.name_full if s.regional_district else None) or "—"

    def _fmt_installed(val):
        if val is None or val == "":
            return "—"
        return str(val).replace(".", ",")

    def _fmt_gavg(tr):
        v = getattr(tr, "generation_average_multiyear_million_kwh", None)
        if v is None:
            return "—"
        s = str(v).strip()
        if not s:
            return "—"
        return s.replace(".", ",")

    wb = Workbook()
    ws = wb.active
    ws.title = "Характеристика актуальных площадок размещения новых ГЭС"

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    row_num = 2
    idx = 0

    if not ges_place_type_groups:
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=total_cols)
        c = ws.cell(row=row_num, column=1, value="Нет данных")
        c.alignment = center_align
        c.border = thin_border
        c.font = Font(italic=True, color="808080")
    else:
        for group in ges_place_type_groups:
            ws.merge_cells(
                start_row=row_num, start_column=1, end_row=row_num, end_column=total_cols
            )
            gc = ws.cell(row=row_num, column=1, value=group["type_label"])
            gc.font = group_font
            gc.alignment = center_align
            gc.border = thin_border
            row_num += 1

            for station in group["stations"]:
                idx += 1
                ues = _ues_name(station)
                rd = _rd_name(station)
                site = station.site_name or "—"
                initiator = station.project_initiator or "—"
                water = station.water_body or "—"
                regulation = station.regulation_type or "—"
                geo = station.geo_location or "—"
                planned = station.planned_capacity_mw or "—"
                gen_scheme = station.general_scheme_commissioning_period or "—"
                constr = format_construction_period_years_display(
                    station.construction_period_years
                ) or "—"

                indicators = station.ges_tep_source_indicators or []
                if indicators:
                    first_row = row_num
                    for loop_idx, tr in enumerate(indicators):
                        inst = _fmt_installed(tr.installed_capacity_mw)
                        gavg = _fmt_gavg(tr)
                        if loop_idx == 0:
                            ws.cell(row=row_num, column=1, value=idx).border = thin_border
                            ws.cell(row=row_num, column=1).alignment = center_align
                            ws.cell(row=row_num, column=2, value=site).border = thin_border
                            ws.cell(row=row_num, column=2).alignment = center_align
                            ws.cell(row=row_num, column=3, value=initiator).border = thin_border
                            ws.cell(row=row_num, column=3).alignment = center_align
                            ws.cell(row=row_num, column=4, value=planned).border = thin_border
                            ws.cell(row=row_num, column=4).alignment = center_align
                        ws.cell(row=row_num, column=5, value=gavg).border = thin_border
                        ws.cell(row=row_num, column=5).alignment = center_align
                        if loop_idx == 0:
                            ws.cell(row=row_num, column=6, value=regulation).border = thin_border
                            ws.cell(row=row_num, column=6).alignment = center_align
                            ws.cell(row=row_num, column=7, value=ues).border = thin_border
                            ws.cell(row=row_num, column=7).alignment = center_align
                            ws.cell(row=row_num, column=8, value=rd).border = thin_border
                            ws.cell(row=row_num, column=8).alignment = center_align
                            ws.cell(row=row_num, column=9, value=water).border = thin_border
                            ws.cell(row=row_num, column=9).alignment = center_align
                            ws.cell(row=row_num, column=10, value=geo).border = thin_border
                            ws.cell(row=row_num, column=10).alignment = center_align
                        ws.cell(row=row_num, column=11, value=inst).border = thin_border
                        ws.cell(row=row_num, column=11).alignment = center_align
                        if loop_idx == 0:
                            ws.cell(row=row_num, column=12, value=gen_scheme).border = thin_border
                            ws.cell(row=row_num, column=12).alignment = center_align
                            ws.cell(row=row_num, column=13, value=constr).border = thin_border
                            ws.cell(row=row_num, column=13).alignment = center_align
                        row_num += 1
                    last_row = row_num - 1
                    if last_row > first_row:
                        for c in list(range(1, 5)) + list(range(6, 11)) + [12, 13]:
                            ws.merge_cells(
                                start_row=first_row,
                                start_column=c,
                                end_row=last_row,
                                end_column=c,
                            )
                else:
                    ws.cell(row=row_num, column=1, value=idx).border = thin_border
                    ws.cell(row=row_num, column=1).alignment = center_align
                    ws.cell(row=row_num, column=2, value=site).border = thin_border
                    ws.cell(row=row_num, column=2).alignment = center_align
                    ws.cell(row=row_num, column=3, value=initiator).border = thin_border
                    ws.cell(row=row_num, column=3).alignment = center_align
                    ws.cell(row=row_num, column=4, value=planned).border = thin_border
                    ws.cell(row=row_num, column=4).alignment = center_align
                    em_gavg = ws.cell(row=row_num, column=5, value="—")
                    em_gavg.border = thin_border
                    em_gavg.alignment = center_align
                    em_gavg.font = Font(color="808080")
                    ws.cell(row=row_num, column=6, value=regulation).border = thin_border
                    ws.cell(row=row_num, column=6).alignment = center_align
                    ws.cell(row=row_num, column=7, value=ues).border = thin_border
                    ws.cell(row=row_num, column=7).alignment = center_align
                    ws.cell(row=row_num, column=8, value=rd).border = thin_border
                    ws.cell(row=row_num, column=8).alignment = center_align
                    ws.cell(row=row_num, column=9, value=water).border = thin_border
                    ws.cell(row=row_num, column=9).alignment = center_align
                    ws.cell(row=row_num, column=10, value=geo).border = thin_border
                    ws.cell(row=row_num, column=10).alignment = center_align
                    em_inst = ws.cell(row=row_num, column=11, value="—")
                    em_inst.border = thin_border
                    em_inst.alignment = center_align
                    em_inst.font = Font(color="808080")
                    ws.cell(row=row_num, column=12, value=gen_scheme).border = thin_border
                    ws.cell(row=row_num, column=12).alignment = center_align
                    ws.cell(row=row_num, column=13, value=constr).border = thin_border
                    ws.cell(row=row_num, column=13).alignment = center_align
                    row_num += 1

    widths = [6, 22, 16, 12, 16, 16, 18, 22, 14, 22, 12, 18, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"prospective_places_ges_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )

