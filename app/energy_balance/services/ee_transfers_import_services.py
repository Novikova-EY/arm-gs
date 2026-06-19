# -*- coding: utf-8 -*-
"""Импорт перетоков ЭЭ из сводов ОЭС и сводного файла (лист «Перетоки», xlsx/xls) во все версии БД."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.common.models.database_version_model import DatabaseVersion
from app.common.services.database_version_filter import filter_by_explicit_db_version
from app.common.services.tranzaction_services import _commit_with_retry, quick_fix_seq
from app.extensions import db
from app.energy_balance.models.regional_energy_system_transfer_model import (
    EE_TRANSFER_PERIOD_YEAR,
    RegionalEnergySystemTransfer,
)
from app.energy_balance.services.ee_generation_import_services import (
    _all_database_version_ids,
    _extract_year_number,
    _normalize_name,
    _normalize_text,
    _to_decimal_or_none,
)
from app.generation.services.machine_services.machine_services import is_same_decimal, to_decimal
from app.logs.services.logging_service import log_to_db
from app.energy_balance.services.ee_transfers_resolve_services import (
    TransferImportResolver,
    clear_transfer_import_resolver_cache,
)
from app.refdata.models.energy_systems.union_energy_system_model import UnionEnergySystem
from config import SCHEMA_ENERGY_BALANCE

SHEET_HINT = "переток"
SVOD_FROM_HEADER = "наименование энергосистемы"
SVOD_TO_HEADER_HINTS = ("название объект",)
ANNUAL_COL_INDEX = 2
MONTH_COL_INDEXES: dict[int, int] = {
    4: 1,
    5: 2,
    6: 3,
    8: 4,
    9: 5,
    10: 6,
    12: 7,
    13: 8,
    14: 9,
    16: 10,
    17: 11,
    18: 12,
}
SKIP_TO_NAMES = frozenset({"сальдо", "перетоки"})
SKIP_TO_SUBSTRINGS = ("пригранич",)


@dataclass
class ParsedPeriodValues:
    annual: Decimal | None = None
    months: dict[int, Decimal | None] = field(default_factory=dict)


@dataclass
class ParsedTransferRow:
    from_name: str
    to_name: str
    values: ParsedPeriodValues


@dataclass
class ParsedTransfersWorkbook:
    year_number: int
    source_oes_name: str | None
    transfers: list[ParsedTransferRow] = field(default_factory=list)
    skipped_rows: int = 0
    is_svod_format: bool = False


def _parse_period_values(row: tuple) -> ParsedPeriodValues:
    values = ParsedPeriodValues()
    values.annual = _to_decimal_or_none(row[ANNUAL_COL_INDEX] if len(row) > ANNUAL_COL_INDEX else None)
    for col_idx, month_number in MONTH_COL_INDEXES.items():
        if len(row) <= col_idx:
            continue
        values.months[month_number] = _to_decimal_or_none(row[col_idx])
    return values


def _pick_transfers_sheet(sheetnames: list[str]) -> str:
    for name in sheetnames:
        if SHEET_HINT in _normalize_name(name):
            return name
    raise ValueError("В файле не найден лист «Перетоки».")


def _is_svod_header_row(col_from: str, col_to: str) -> bool:
    if _normalize_name(col_from) != SVOD_FROM_HEADER:
        return False
    to_norm = _normalize_name(col_to)
    return any(hint in to_norm for hint in SVOD_TO_HEADER_HINTS)


def _load_transfers_sheet_rows(file_bytes: bytes, filename: str | None) -> tuple[str, list[tuple]]:
    lower_name = (filename or "").lower()
    if lower_name.endswith(".xls") and not lower_name.endswith(".xlsx"):
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError(
                "Для импорта файлов .xls требуется пакет xlrd. "
                "Установите зависимости: pip install xlrd "
                "или сохраните файл в формате .xlsx."
            ) from exc

        workbook = xlrd.open_workbook(file_contents=file_bytes)
        sheet_name = _pick_transfers_sheet(list(workbook.sheet_names()))
        sheet = workbook.sheet_by_name(sheet_name)
        rows = [tuple(sheet.row_values(row_idx)) for row_idx in range(sheet.nrows)]
        return sheet_name, rows

    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        sheet_name = _pick_transfers_sheet(list(workbook.sheetnames))
        worksheet = workbook[sheet_name]
        rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
        return sheet_name, rows
    finally:
        workbook.close()


def _is_oes_header_name(name: str) -> bool:
    norm = _normalize_name(name)
    return norm.startswith("оэс ") or norm.startswith("титэс ")


def _is_transfer_row(from_name: str, to_name: str) -> bool:
    if not from_name or not to_name:
        return False
    to_norm = _normalize_name(to_name)
    if to_norm in SKIP_TO_NAMES:
        return False
    if any(substr in to_norm for substr in SKIP_TO_SUBSTRINGS):
        return False
    if _is_oes_header_name(from_name) and not to_name:
        return False
    if from_name.strip().startswith("приграничная"):
        return False
    return True


def parse_ee_transfers_workbook(file_bytes: bytes, filename: str | None = None) -> ParsedTransfersWorkbook:
    sheet_name, rows = _load_transfers_sheet_rows(file_bytes, filename)
    year_number = _extract_year_number(filename, sheet_name)

    parsed = ParsedTransfersWorkbook(year_number=year_number, source_oes_name=None)
    header_passed = False

    for row_vals in rows:
        if not any(v is not None and str(v).strip() for v in row_vals[:3]):
            continue

        col_from = _normalize_text(row_vals[0] if len(row_vals) > 0 else None)
        col_to = _normalize_text(row_vals[1] if len(row_vals) > 1 else None)

        if not header_passed:
            if _normalize_name(col_from) == SVOD_FROM_HEADER:
                header_passed = True
                parsed.is_svod_format = _is_svod_header_row(col_from, col_to)
            continue

        if not parsed.is_svod_format and _is_oes_header_name(col_from) and not col_to:
            parsed.source_oes_name = col_from
            continue

        if not _is_transfer_row(col_from, col_to):
            parsed.skipped_rows += 1
            continue

        parsed.transfers.append(
            ParsedTransferRow(
                from_name=col_from,
                to_name=col_to,
                values=_parse_period_values(row_vals),
            )
        )

    if not parsed.transfers:
        raise ValueError("В файле не найдено строк перетоков для импорта.")
    if not parsed.source_oes_name and not parsed.is_svod_format:
        parsed.source_oes_name = sheet_name
    return parsed


def _resolve_ues_id_by_name(oes_name: str | None, version_id: int | None) -> int | None:
    if not oes_name:
        return None
    norm = _normalize_name(oes_name)
    if not norm:
        return None
    ues_list = filter_by_explicit_db_version(
        UnionEnergySystem.query,
        UnionEnergySystem,
        version_id,
    ).all()
    for ues in ues_list:
        for candidate in (ues.name, ues.name_full):
            if candidate and _normalize_name(candidate) == norm:
                return ues.id
    for ues in ues_list:
        if ues.name and norm in _normalize_name(ues.name):
            return ues.id
    return None


def _transfer_record_key(
    lookup: dict[str, Any],
    month_number: int,
) -> tuple[Any, ...]:
    return (
        lookup["id_from_regional_energy_system"],
        lookup["id_from_regional_district"],
        lookup["id_to_regional_energy_system"],
        lookup["id_to_foreign_border_country"],
        lookup["id_to_regional_district"],
        lookup["id_to_energy_unit"],
        lookup["id_union_energy_system"],
        month_number,
    )


def _load_existing_transfers_map(
    *,
    year_number: int,
    version_id: int | None,
) -> dict[tuple[Any, ...], RegionalEnergySystemTransfer]:
    q = RegionalEnergySystemTransfer.query.filter_by(year_number=year_number)
    q = filter_by_explicit_db_version(q, RegionalEnergySystemTransfer, version_id)
    result: dict[tuple[Any, ...], RegionalEnergySystemTransfer] = {}
    for rec in q.all():
        key = (
            rec.id_from_regional_energy_system,
            rec.id_from_regional_district,
            rec.id_to_regional_energy_system,
            rec.id_to_foreign_border_country,
            rec.id_to_regional_district,
            rec.id_to_energy_unit,
            rec.id_union_energy_system,
            rec.month_number,
        )
        result[key] = rec
    return result


def _upsert_transfer_period_values(
    *,
    lookup: dict[str, Any],
    values: ParsedPeriodValues,
    year_number: int,
    version_id: int | None,
    existing_map: dict[tuple[Any, ...], RegionalEnergySystemTransfer],
) -> tuple[int, int, int]:
    inserted = 0
    updated = 0
    unchanged = 0
    periods: list[tuple[int, Decimal | None]] = [(EE_TRANSFER_PERIOD_YEAR, values.annual)]
    periods.extend(sorted(values.months.items()))

    for month_number, new_val in periods:
        if new_val is None:
            continue
        key = _transfer_record_key(lookup, month_number)
        rec = existing_map.get(key)
        old_val = rec.transfer_value if rec else None
        if rec is not None and is_same_decimal(to_decimal(old_val), new_val):
            unchanged += 1
            continue
        if rec is None:
            rec = RegionalEnergySystemTransfer(
                **lookup,
                year_number=year_number,
                month_number=month_number,
            )
            rec.transfer_value = new_val
            rec.database_version_id = version_id
            db.session.add(rec)
            existing_map[key] = rec
            inserted += 1
        else:
            rec.transfer_value = new_val
            db.session.add(rec)
            updated += 1
    return inserted, updated, unchanged


def _import_parsed_transfers_for_version(
    parsed: ParsedTransfersWorkbook,
    version_id: int | None,
) -> tuple[dict[str, int], list[str]]:
    stats: dict[str, int] = defaultdict(int)
    errors: list[str] = []
    seen_errors: set[str] = set()

    def _append_error(message: str) -> None:
        if message in seen_errors or len(errors) >= 30:
            return
        seen_errors.add(message)
        errors.append(message)

    resolver = TransferImportResolver.build(version_id)
    existing_map = _load_existing_transfers_map(
        year_number=parsed.year_number,
        version_id=version_id,
    )

    ues_id = None
    if not parsed.is_svod_format:
        ues_id = _resolve_ues_id_by_name(parsed.source_oes_name, version_id)
        if ues_id is None and parsed.source_oes_name:
            stats["ues_missing"] += 1
            _append_error(
                f"ОЭС «{parsed.source_oes_name}» не найдена "
                f"(версия БД {version_id or 'без версии'})."
            )

    for transfer_row in parsed.transfers:
        from_source = resolver.resolve_from_source(transfer_row.from_name)
        if from_source.res_id is None and from_source.rd_id is None:
            stats["from_source_missing"] += 1
            _append_error(
                f"Источник «{transfer_row.from_name}» не найден "
                f"ни в РЭС, ни в субъектах РФ "
                f"(версия БД {version_id or 'без версии'})."
            )
            continue

        to_target = resolver.resolve_to_target(transfer_row.to_name)
        if (
            to_target.res_id is None
            and to_target.country_id is None
            and to_target.rd_id is None
            and to_target.energy_unit_id is None
        ):
            if version_id is None and (
                resolver.has_country_anchor(transfer_row.to_name)
                or resolver.has_energy_unit_anchor(transfer_row.to_name)
            ):
                stats["legacy_null_version_target_skipped"] += 1
                continue
            stats["to_target_missing"] += 1
            _append_error(
                f"Получатель «{transfer_row.to_name}» не найден "
                f"ни в РЭС, ни в субъектах РФ, ни в зарубежных странах, "
                f"ни в энергорайонах "
                f"(версия БД {version_id or 'без версии'})."
            )
            continue

        lookup = {
            "id_union_energy_system": ues_id,
            "id_from_regional_energy_system": from_source.res_id,
            "id_from_regional_district": from_source.rd_id,
            "id_to_regional_energy_system": to_target.res_id,
            "id_to_foreign_border_country": to_target.country_id,
            "id_to_regional_district": to_target.rd_id,
            "id_to_energy_unit": to_target.energy_unit_id,
            "source_oes_name": parsed.source_oes_name,
        }
        inserted, updated, unchanged = _upsert_transfer_period_values(
            lookup=lookup,
            values=transfer_row.values,
            year_number=parsed.year_number,
            version_id=version_id,
            existing_map=existing_map,
        )
        stats["transfer_rows_inserted"] += inserted
        stats["transfer_rows_updated"] += updated
        stats["transfer_rows_unchanged"] += unchanged

    return dict(stats), errors


def import_ee_transfers_from_excel(file, user: str) -> dict[str, Any]:
    raw = file.read()
    filename = getattr(file, "filename", None)
    parsed = parse_ee_transfers_workbook(raw, filename=filename)

    version_ids = _all_database_version_ids()
    stats: dict[str, int] = defaultdict(int)
    errors: list[str] = []

    try:
        for version_id in version_ids:
            version_stats, version_errors = _import_parsed_transfers_for_version(parsed, version_id)
            for key, value in version_stats.items():
                stats[key] += value
            errors.extend(version_errors)
            if len(errors) > 30:
                errors = errors[:30]
    finally:
        clear_transfer_import_resolver_cache()

    try:
        quick_fix_seq(SCHEMA_ENERGY_BALANCE, "gs_bem_regional_energy_system_transfers", "id")
    except Exception:
        pass

    _commit_with_retry()

    from app.energy_balance.services.energy_balance_cache import clear_energy_balance_cache

    clear_energy_balance_cache()

    source_label = (
        "свод"
        if parsed.is_svod_format
        else (parsed.source_oes_name or "ОЭС")
    )
    inserted = stats.get("transfer_rows_inserted", 0)
    updated = stats.get("transfer_rows_updated", 0)
    unchanged = stats.get("transfer_rows_unchanged", 0)
    changed = inserted + updated
    message = (
        f"Импорт перетоков ЭЭ за {parsed.year_number} г. ({source_label}) завершён. "
        f"Добавлено: {inserted}, обновлено: {updated}, без изменений: {unchanged}. "
        f"Версий БД: {len(version_ids)}."
    )
    if changed == 0 and unchanged > 0:
        message += (
            f" Данные за {parsed.year_number} г. уже есть в базе — "
            "проверьте диапазон лет на странице (должен включать "
            f"{parsed.year_number} г.)."
        )
    log_to_db(
        user,
        "Импорт перетоков ЭЭ из Excel",
        message,
        entity_type="ee_transfers_import",
        entity_id=0,
    )
    return {
        "message": message,
        "stats": dict(stats),
        "errors": errors,
        "errors_count": len(errors),
        "year_number": parsed.year_number,
    }
