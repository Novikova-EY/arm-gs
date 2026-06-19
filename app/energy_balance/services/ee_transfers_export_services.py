# -*- coding: utf-8 -*-
"""Экспорт страницы «Перетоки ЭЭ» в Excel."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.energy_balance.services.ee_transfers_page_services import (
    format_transfer_cell,
    get_ee_transfers_page_data,
)
from app.fuel.services.fuel_exports.hierarchy_excel_layout import (
    FILL_RES_HEADER,
    FILL_UES,
    write_merged_section_row,
)

FILL_TOTAL_RU = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")


def _write_transfer_row(
    ws,
    row_idx: int,
    to_name: str,
    period_columns: list[tuple[int, str]],
    periods: dict,
    rounding_digits: int,
) -> int:
    values = [to_name]
    values.extend(
        format_transfer_cell(periods.get(period_key), rounding_digits)
        for period_key, _label in period_columns
    )
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = Alignment(
            horizontal="left" if col_idx == 1 else "center",
            vertical="center",
            wrap_text=True,
        )
    return row_idx + 1


def _write_total_row(
    ws,
    row_idx: int,
    label: str,
    period_columns: list[tuple[int, str]],
    period_totals: dict,
    rounding_digits: int,
    *,
    fill=None,
) -> int:
    values = [label]
    values.extend(
        format_transfer_cell(period_totals.get(period_key), rounding_digits)
        for period_key, _label in period_columns
    )
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = Alignment(
            horizontal="left" if col_idx == 1 else "center",
            vertical="center",
            wrap_text=True,
        )
        cell.font = Font(bold=True)
        if fill is not None:
            cell.fill = fill
    return row_idx + 1


def export_ee_transfers_to_excel(
    *,
    filters: dict,
    period_mode: str,
    selected_year: int | None,
    start_year: int | None,
    end_year: int | None,
    rounding_digits: int,
    show_totals: bool,
    per_page: int | str = "all",
    page: int = 1,
) -> BytesIO | None:
    """Формирует Excel-файл с тем же содержимым, что отображается на странице."""
    page_data = get_ee_transfers_page_data(
        filters=filters,
        page=page,
        per_page=per_page,
        period_mode=period_mode,
        selected_year=selected_year,
        start_year=start_year,
        end_year=end_year,
        rounding_digits=rounding_digits,
        show_totals=show_totals,
    )

    hierarchy = page_data.get("hierarchy") or []
    period_columns = page_data.get("period_columns") or []
    transfer_aggregates = page_data.get("transfer_aggregates") or {}
    should_show_totals = page_data.get("should_show_totals") or {}

    columns = ["Энергосистема"] + [label for _key, label in period_columns]
    num_cols = len(columns)

    wb = Workbook()
    ws = wb.active
    ws.title = "Перетоки ЭЭ"

    header_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num, title in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_num, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    if not hierarchy:
        ws.cell(row=2, column=1, value="Перетоки не найдены")
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    row_idx = 2
    for ues_block in hierarchy:
        row_idx = write_merged_section_row(
            ws,
            row_idx,
            num_cols,
            ues_block.get("ues_name") or "—",
            FILL_UES,
        )
        for res_block in ues_block.get("res_list") or []:
            row_idx = write_merged_section_row(
                ws,
                row_idx,
                num_cols,
                res_block.get("res_name") or "—",
                FILL_RES_HEADER,
            )
            for transfer_row in res_block.get("transfer_rows") or []:
                row_idx = _write_transfer_row(
                    ws,
                    row_idx,
                    transfer_row.get("to_name") or "—",
                    period_columns,
                    transfer_row.get("periods") or {},
                    rounding_digits,
                )
            group_key = res_block.get("group_key")
            if should_show_totals.get("from_groups", {}).get(group_key):
                res_name = res_block.get("res_name") or "—"
                row_idx = _write_total_row(
                    ws,
                    row_idx,
                    f"{res_name}, всего",
                    period_columns,
                    transfer_aggregates.get("from_groups", {}).get(group_key, {}),
                    rounding_digits,
                    fill=FILL_RES_HEADER,
                )
        ues_id = ues_block.get("ues_id")
        if should_show_totals.get("union_energy_systems", {}).get(ues_id):
            ues_name = ues_block.get("ues_name") or "—"
            row_idx = _write_total_row(
                ws,
                row_idx,
                f"{ues_name}, всего",
                period_columns,
                transfer_aggregates.get("union_energy_systems", {}).get(ues_id, {}),
                rounding_digits,
                fill=FILL_UES,
            )
    if should_show_totals.get("total"):
        row_idx = _write_total_row(
            ws,
            row_idx,
            "Россия, всего",
            period_columns,
            transfer_aggregates.get("total", {}),
            rounding_digits,
            fill=FILL_TOTAL_RU,
        )

    max_lengths = [len(c) for c in columns]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        for i, val in enumerate(row):
            if val is not None and i < len(max_lengths):
                max_lengths[i] = max(max_lengths[i], len(str(val)))
    for col_idx, width in enumerate(max_lengths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(48, width + 2))

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
