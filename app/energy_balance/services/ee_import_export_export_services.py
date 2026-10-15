# -*- coding: utf-8 -*-
"""Выгрузка страницы «Импорт/Экспорт ЭЭ» в Excel."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.energy_balance.services.ee_import_export_page_services import (
    get_ee_import_export_page_data,
)


def export_ee_import_export_to_excel(
    *,
    start_year: int | None,
    end_year: int | None,
    rounding_digits: int,
    direction_filter: list[str] | None = None,
    country_filter: list[int] | None = None,
) -> BytesIO:
    data = get_ee_import_export_page_data(
        start_year=start_year,
        end_year=end_year,
        rounding_digits=rounding_digits,
        direction_filter=direction_filter,
        country_filter=country_filter,
    )
    period_columns = data["period_columns"]
    columns = ["Экспорт/Импорт", "Страна"] + [label for _, label in period_columns]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Импорт-Экспорт ЭЭ"

    header_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    for column_index, title in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=column_index, value=title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if not data["hierarchy"]:
        sheet.cell(row=2, column=1, value="Данные не найдены")
    else:
        row_index = 2
        for oes in data["hierarchy"]:
            sheet.merge_cells(
                start_row=row_index,
                start_column=1,
                end_row=row_index,
                end_column=len(columns),
            )
            title_cell = sheet.cell(row=row_index, column=1, value=oes["ues_name"])
            title_cell.font = Font(bold=True)
            title_cell.fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
            row_index += 1

            for transfer_row in oes["rows"]:
                sheet.cell(
                    row=row_index,
                    column=1,
                    value=transfer_row["direction_label"],
                )
                sheet.cell(row=row_index, column=2, value=transfer_row["country_name"])
                for column_index, (period_key, _label) in enumerate(period_columns, start=3):
                    sheet.cell(
                        row=row_index,
                        column=column_index,
                        value=transfer_row["cell_values"][period_key],
                    )
                row_index += 1


    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal="left" if cell.column <= 2 else "center",
                vertical="center",
                wrap_text=True,
            )
    for column_index in range(1, len(columns) + 1):
        widest = max(
            (len(str(sheet.cell(row=row, column=column_index).value or "")) for row in range(1, sheet.max_row + 1)),
            default=10,
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = max(10, min(48, widest + 2))

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
