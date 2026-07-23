# -*- coding: utf-8 -*-
"""Импорт выработки ЭЭ из сводов ОЭС (xlsx/xls) во все версии БД."""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, Sequence

from openpyxl import load_workbook

from app.common.models.database_version_model import DatabaseVersion
from app.common.services.database_version_filter import filter_by_explicit_db_version
from app.common.services.tranzaction_services import _commit_with_retry, quick_fix_seq
from app.extensions import db
from app.energy_balance.models.espp_energy_generation_model import ESPPEnergyGeneration
from app.energy_balance.models.regional_energy_system_energy_generation_model import (
    RegionalEnergySystemEnergyGeneration,
)
from app.energy_balance.models.station_energy_generation_model import (
    STATION_ENERGY_GENERATION_PERIOD_YEAR,
    StationEnergyGeneration,
)
from app.generation.models.station.station_constants import STATION_SIGN_ESPP
from app.generation.models.station.station_model import Station
from app.generation.services.machine_services.machine_services import is_same_decimal, to_decimal
from app.logs.services.logging_service import log_to_db
from app.refdata.models.energy_systems.regional_energy_system_model import RegionalEnergySystem
from config import SCHEMA_ENERGY_BALANCE

SHEET_HINT = "выработ"
UUID_RE = re.compile(
    r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$",
    re.IGNORECASE,
)
YEAR_RE = re.compile(r"(20\d{2})")

MONTH_HEADER_PATTERNS: dict[int, tuple[str, ...]] = {
    1: ("январ",),
    2: ("феврал",),
    3: ("март", "мар "),
    4: ("апрел", "апр"),
    5: ("май",),
    6: ("июн",),
    7: ("июл",),
    8: ("август", "авг"),
    9: ("сентябр", "сен"),
    10: ("октябр", "окт"),
    11: ("ноябр", "ноя"),
    12: ("декабр", "дек"),
}


@dataclass
class ParsedPeriodValues:
    annual: Decimal | None = None
    months: dict[int, Decimal | None] = field(default_factory=dict)


@dataclass
class ParsedStationRow:
    external_code: str
    values: ParsedPeriodValues


@dataclass
class ParsedMultiYearStationRow:
    external_code: str
    values_by_year: dict[int, Decimal | None]
    kto: str | None = None
    station_sign: str | None = None


@dataclass
class ParsedResRow:
    res_name: str | None
    values: ParsedPeriodValues
    anchor_external_code: str | None = None


@dataclass
class ParsedWorkbook:
    year_number: int
    stations: list[ParsedStationRow] = field(default_factory=list)
    multi_year_stations: list[ParsedMultiYearStationRow] = field(default_factory=list)
    espp_rows: list[ParsedResRow] = field(default_factory=list)
    control_rows: list[ParsedResRow] = field(default_factory=list)
    skipped_rows: int = 0
    is_multi_year: bool = False


@dataclass
class MultiYearColumnLayout:
    header_row_idx: int
    external_code_col: int
    name_col: int | None
    year_cols: dict[int, int]
    kto_col: int | None = None
    prom_sign_col: int | None = None


@dataclass
class GenerationColumnLayout:
    header_row_idx: int
    external_code_cols: list[int]
    annual_col: int | None
    month_cols: dict[int, int]
    res_col: int | None = None
    sign_col: int | None = None
    name_col: int | None = None


def _normalize_text(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").replace("\xa0", " ").strip()
    return " ".join(text.split())


def _normalize_name(value) -> str:
    return _normalize_text(value).casefold()


def _is_uuid(value: str) -> bool:
    if not value:
        return False
    return bool(UUID_RE.match(value.strip()))


def _to_decimal_or_none(value) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    try:
        if isinstance(value, (int, float, Decimal)):
            dec = Decimal(str(value))
        else:
            dec = to_decimal(value)
        if dec is None:
            return None
        return dec
    except (InvalidOperation, ValueError, TypeError):
        return None


def _extract_year_number(filename: str | None, sheet_name: str | None) -> int:
    for source in (filename or "", sheet_name or ""):
        match = YEAR_RE.search(source)
        if match:
            return int(match.group(1))
    raise ValueError("Не удалось определить год из имени файла или листа (ожидается, например, 2024).")


def _pick_generation_sheet(sheetnames: list[str]) -> str:
    for name in sheetnames:
        if SHEET_HINT in _normalize_name(name):
            return name
    return sheetnames[0]


def _is_external_code_header(header_norm: str) -> bool:
    compact = header_norm.replace(" ", "_").replace("-", "_")
    return compact == "external_code"


def _extract_year_from_header(header_norm: str) -> int | None:
    match = YEAR_RE.search(header_norm)
    if not match:
        return None
    return int(match.group(1))


def _is_total_row_name(name: str) -> bool:
    norm = _normalize_name(name)
    return norm.startswith("итого")


def _is_annual_header(header_norm: str) -> bool:
    if "накоплен" in header_norm:
        return True
    if "начала" in header_norm and "год" in header_norm:
        return True
    if "итог" in header_norm and "год" in header_norm:
        return True
    if "годов" in header_norm and "выработ" in header_norm:
        return True
    return header_norm in {"год", "итого за год", "за год"}


def _match_month_number(header_norm: str) -> int | None:
    if not header_norm or "квартал" in header_norm:
        return None
    for month_number, patterns in MONTH_HEADER_PATTERNS.items():
        for pattern in patterns:
            if pattern in header_norm:
                return month_number
    return None


def _is_sign_header(header_norm: str) -> bool:
    return "признак" in header_norm


def _is_kto_header(header_norm: str) -> bool:
    return "код кпо" in header_norm or header_norm == "кпо"


def _is_prom_sign_header(header_norm: str) -> bool:
    return "пром" in header_norm and "предприят" in header_norm


def _normalize_kto(value) -> str | None:
    if value is None:
        return None
    text = _normalize_text(value)
    return text or None


def _parse_prom_sign_value(value) -> str | None:
    norm = _normalize_name(value)
    if norm == "да":
        return "true"
    if norm == "нет":
        return "false"
    return None


def _is_name_header(header_norm: str) -> bool:
    if "признак" in header_norm:
        return False
    # «Наименование Энергосистемы» — колонка РЭС, не название объекта/показателя.
    if "энергосистем" in header_norm:
        return False
    if header_norm == "электростанция":
        return True
    if "название" in header_norm and (
        "объект" in header_norm or "показател" in header_norm
    ):
        return True
    if "наименование" in header_norm:
        return True
    if "выработка" in header_norm and "распредел" in header_norm:
        return True
    return False


def _is_res_header(header_norm: str) -> bool:
    if "рэс" in header_norm:
        return True
    if "энергосистем" in header_norm:
        # «Наименование Энергосистемы», «Региональная энергосистема» и т.п.
        if header_norm.startswith("тип "):
            return False
        return True
    return "региональн" in header_norm and "энергосистем" in header_norm


def _infer_res_col(
    external_code_cols: list[int],
    period_cols: set[int],
) -> int | None:
    if not external_code_cols:
        return 0
    first_code_col = min(external_code_cols)
    meta_cols = [idx for idx in range(first_code_col) if idx not in period_cols]
    return meta_cols[0] if meta_cols else None


def _infer_sign_col(
    external_code_cols: list[int],
    res_col: int | None,
    period_cols: set[int],
) -> int | None:
    if not external_code_cols:
        return None
    first_code_col = min(external_code_cols)
    meta_cols = [
        idx
        for idx in range(first_code_col)
        if idx not in period_cols and idx != res_col
    ]
    return meta_cols[-1] if meta_cols else None


def _infer_name_col(
    external_code_cols: list[int],
    used_cols: set[int],
    header_count: int,
) -> int | None:
    if not external_code_cols:
        return None
    last_code_col = max(external_code_cols)
    for candidate in range(last_code_col + 1, header_count):
        if candidate in used_cols:
            continue
        return candidate
    return None


def _detect_multi_year_column_layout(rows: Sequence[tuple]) -> MultiYearColumnLayout | None:
    """Шаблон с годами в заголовках столбцов (напр. выработка ПАО Сургутнефтегаз)."""
    for row_idx, row_vals in enumerate(rows[:40]):
        headers = [_normalize_name(cell) for cell in row_vals]
        external_code_cols = [idx for idx, header in enumerate(headers) if _is_external_code_header(header)]
        if not external_code_cols:
            continue

        year_cols: dict[int, int] = {}
        for idx, header in enumerate(headers):
            if _is_external_code_header(header):
                continue
            year_number = _extract_year_from_header(header)
            if year_number is not None and _match_month_number(header) is None:
                year_cols[idx] = year_number

        if not year_cols:
            continue

        has_month_cols = any(_match_month_number(header) is not None for header in headers)
        has_annual_col = any(_is_annual_header(header) for header in headers)
        if has_month_cols or has_annual_col:
            continue

        name_col: int | None = None
        kto_col: int | None = None
        prom_sign_col: int | None = None
        for idx, header in enumerate(headers):
            if kto_col is None and _is_kto_header(header):
                kto_col = idx
            if prom_sign_col is None and _is_prom_sign_header(header):
                prom_sign_col = idx
            if name_col is None and _is_name_header(header):
                name_col = idx
        if name_col is None and external_code_cols[0] > 0:
            name_col = external_code_cols[0] - 1

        return MultiYearColumnLayout(
            header_row_idx=row_idx,
            external_code_col=external_code_cols[0],
            name_col=name_col,
            year_cols=year_cols,
            kto_col=kto_col,
            prom_sign_col=prom_sign_col,
        )
    return None


def _detect_column_layout(rows: Sequence[tuple]) -> GenerationColumnLayout:
    for row_idx, row_vals in enumerate(rows[:40]):
        headers = [_normalize_name(cell) for cell in row_vals]
        external_code_cols = [idx for idx, header in enumerate(headers) if _is_external_code_header(header)]
        if not external_code_cols:
            continue

        annual_col: int | None = None
        month_cols: dict[int, int] = {}
        res_col: int | None = None
        sign_col: int | None = None
        name_col: int | None = None

        for idx, header in enumerate(headers):
            if _is_external_code_header(header):
                continue
            if annual_col is None and _is_annual_header(header):
                annual_col = idx
                continue
            month_number = _match_month_number(header)
            if month_number is not None and month_number not in month_cols.values():
                month_cols[idx] = month_number
                continue
            if res_col is None and _is_res_header(header):
                res_col = idx
                continue
            if sign_col is None and _is_sign_header(header):
                sign_col = idx
                continue
            if name_col is None and _is_name_header(header):
                name_col = idx

        period_cols = set(month_cols)
        if annual_col is not None:
            period_cols.add(annual_col)

        if res_col is None:
            res_col = _infer_res_col(external_code_cols, period_cols)
        if sign_col is None:
            sign_col = _infer_sign_col(external_code_cols, res_col, period_cols)
        if name_col is not None and name_col == res_col:
            # Одна и та же колонка не может быть и РЭС, и названием объекта.
            name_col = None
        if name_col is None:
            used_cols = set(external_code_cols) | period_cols | {
                c for c in (res_col, sign_col) if c is not None
            }
            name_col = _infer_name_col(external_code_cols, used_cols, len(headers))

        if annual_col is None and not month_cols:
            raise ValueError(
                "В строке заголовков найден external_code, но не найдены столбцы "
                "годовой или помесячной выработки."
            )

        return GenerationColumnLayout(
            header_row_idx=row_idx,
            external_code_cols=external_code_cols,
            annual_col=annual_col,
            month_cols=month_cols,
            res_col=res_col,
            sign_col=sign_col,
            name_col=name_col,
        )

    raise ValueError(
        "Не найдена строка заголовков с колонкой external_code. "
        "Проверьте, что в файле есть заголовок external_code."
    )


def _cell(row: Sequence, col_idx: int | None):
    if col_idx is None or col_idx < 0 or col_idx >= len(row):
        return None
    return row[col_idx]


def _extract_external_code(row: Sequence, layout: GenerationColumnLayout) -> str | None:
    for col_idx in layout.external_code_cols:
        value = _normalize_text(_cell(row, col_idx))
        if _is_uuid(value):
            return value
    return None


def _parse_period_values(row: Sequence, layout: GenerationColumnLayout) -> ParsedPeriodValues:
    values = ParsedPeriodValues()
    values.annual = _to_decimal_or_none(_cell(row, layout.annual_col))
    for col_idx, month_number in layout.month_cols.items():
        values.months[month_number] = _to_decimal_or_none(_cell(row, col_idx))
    return values


def _is_oes_header_name(name: str) -> bool:
    norm = _normalize_name(name)
    return norm.startswith("оэс ") or norm.startswith("титэс ")


def _looks_like_res_col0(name: str) -> bool:
    norm = _normalize_name(name)
    if not norm:
        return False
    if _is_oes_header_name(norm):
        return False
    if norm.startswith("по "):
        return False
    return True


def _resolve_res_name(col_res: str, current_res_name: str | None) -> str | None:
    if _looks_like_res_col0(col_res):
        return col_res
    if current_res_name and _normalize_name(current_res_name) != _normalize_name(STATION_SIGN_ESPP):
        return current_res_name
    return None


def _is_espp_row(col_sign: str, col_name: str) -> bool:
    if _normalize_name(col_sign) == _normalize_name(STATION_SIGN_ESPP):
        return True
    if not col_sign and _normalize_name(col_name) == _normalize_name(STATION_SIGN_ESPP):
        return True
    return False


def _resolve_espp_res_name(
    col_res: str,
    col_name: str,
    col_sign: str,
    current_res_name: str | None,
) -> str | None:
    """Имя РЭС для строки ЭСПП: отдельная колонка РЭС, дублирование в name или контекст блока."""
    if _normalize_name(col_sign) == _normalize_name(STATION_SIGN_ESPP):
        if (
            _looks_like_res_col0(col_res)
            and _normalize_name(col_res) != _normalize_name(STATION_SIGN_ESPP)
        ):
            return col_res
        if col_name and _normalize_name(col_name) != _normalize_name(STATION_SIGN_ESPP):
            return col_name
        if current_res_name and _normalize_name(current_res_name) != _normalize_name(STATION_SIGN_ESPP):
            return current_res_name
        return None

    if _normalize_name(col_name) == _normalize_name(STATION_SIGN_ESPP):
        if (
            _looks_like_res_col0(col_res)
            and _normalize_name(col_res) != _normalize_name(STATION_SIGN_ESPP)
        ):
            return col_res
        if current_res_name and _normalize_name(current_res_name) != _normalize_name(STATION_SIGN_ESPP):
            return current_res_name
    return None


def _find_next_res_col0(
    rows: list[tuple],
    start_idx: int,
    res_col: int | None,
    limit: int = 12,
) -> str | None:
    for row_vals in rows[start_idx : start_idx + limit]:
        col_res = _normalize_text(_cell(row_vals, res_col))
        if _looks_like_res_col0(col_res):
            return col_res
    return None


def _row_has_data(row_vals: tuple, layout: GenerationColumnLayout) -> bool:
    check_cols = set(layout.external_code_cols)
    if layout.res_col is not None:
        check_cols.add(layout.res_col)
    if layout.sign_col is not None:
        check_cols.add(layout.sign_col)
    if layout.name_col is not None:
        check_cols.add(layout.name_col)
    if layout.annual_col is not None:
        check_cols.add(layout.annual_col)
    check_cols.update(layout.month_cols)
    for col_idx in check_cols:
        value = _cell(row_vals, col_idx)
        if value is not None and str(value).strip():
            return True
    return False


def _load_sheet_rows(file_bytes: bytes, filename: str | None) -> tuple[str, list[tuple]]:
    lower_name = (filename or "").lower()
    if lower_name.endswith(".xls") and not lower_name.endswith(".xlsx"):
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError(
                "Для импорта файлов .xls требуется пакет xlrd. "
                "Установите зависимости: pip install xlrd "
                "или сохраните файл в формате .xlsx."
            ) from exc

        workbook = xlrd.open_workbook(file_contents=file_bytes)
        sheet_name = _pick_generation_sheet(list(workbook.sheet_names()))
        sheet = workbook.sheet_by_name(sheet_name)
        rows = [tuple(sheet.row_values(row_idx)) for row_idx in range(sheet.nrows)]
        return sheet_name, rows

    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        sheet_name = _pick_generation_sheet(list(workbook.sheetnames))
        worksheet = workbook[sheet_name]
        rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
        return sheet_name, rows
    finally:
        workbook.close()


def _parse_multi_year_workbook(rows: Sequence[tuple], layout: MultiYearColumnLayout) -> ParsedWorkbook:
    all_years: set[int] = set()
    parsed = ParsedWorkbook(year_number=0, is_multi_year=True)

    for row_vals in rows[layout.header_row_idx + 1 :]:
        name = _normalize_text(_cell(row_vals, layout.name_col))
        if _is_total_row_name(name):
            continue

        external_code = _normalize_text(_cell(row_vals, layout.external_code_col))
        if not _is_uuid(external_code):
            continue

        values_by_year: dict[int, Decimal | None] = {}
        for col_idx, year_number in layout.year_cols.items():
            value = _to_decimal_or_none(_cell(row_vals, col_idx))
            if value is not None:
                values_by_year[year_number] = value
                all_years.add(year_number)

        if not values_by_year:
            parsed.skipped_rows += 1
            continue

        parsed.multi_year_stations.append(
            ParsedMultiYearStationRow(
                external_code=external_code,
                values_by_year=values_by_year,
                kto=_normalize_kto(_cell(row_vals, layout.kto_col)),
                station_sign=_parse_prom_sign_value(_cell(row_vals, layout.prom_sign_col)),
            )
        )

    if not parsed.multi_year_stations:
        raise ValueError("В файле не найдено строк для импорта выработки ЭЭ.")
    parsed.year_number = min(all_years)
    return parsed


def parse_ee_generation_workbook(file_bytes: bytes, filename: str | None = None) -> ParsedWorkbook:
    sheet_name, rows = _load_sheet_rows(file_bytes, filename)
    multi_year_layout = _detect_multi_year_column_layout(rows)
    if multi_year_layout is not None:
        return _parse_multi_year_workbook(rows, multi_year_layout)

    year_number = _extract_year_number(filename, sheet_name)
    layout = _detect_column_layout(rows)

    parsed = ParsedWorkbook(year_number=year_number)
    current_res_name: str | None = None
    last_station_external_code: str | None = None
    data_rows = [row_vals for row_vals in rows[layout.header_row_idx + 1 :] if _row_has_data(row_vals, layout)]

    for row_idx, row_vals in enumerate(data_rows):
        col_res = _normalize_text(_cell(row_vals, layout.res_col))
        col_sign = _normalize_text(_cell(row_vals, layout.sign_col))
        col_name = _normalize_text(_cell(row_vals, layout.name_col))

        if _looks_like_res_col0(col_res):
            current_res_name = col_res

        external_code = _extract_external_code(row_vals, layout)
        if external_code:
            last_station_external_code = external_code
            parsed.stations.append(
                ParsedStationRow(
                    external_code=external_code,
                    values=_parse_period_values(row_vals, layout),
                )
            )
            continue

        if _is_espp_row(col_sign, col_name):
            res_name = _resolve_espp_res_name(col_res, col_name, col_sign, current_res_name)
            if res_name:
                current_res_name = res_name
            parsed.espp_rows.append(
                ParsedResRow(
                    res_name=res_name,
                    values=_parse_period_values(row_vals, layout),
                    anchor_external_code=last_station_external_code if not res_name else None,
                )
            )
            continue

        if _normalize_name(col_name) == "выработка":
            if _looks_like_res_col0(col_res):
                res_name = col_res
            else:
                res_name = _find_next_res_col0(data_rows, row_idx + 1, layout.res_col)
            if not res_name:
                parsed.skipped_rows += 1
                continue
            current_res_name = res_name
            parsed.control_rows.append(
                ParsedResRow(res_name=res_name, values=_parse_period_values(row_vals, layout))
            )
            continue

        parsed.skipped_rows += 1

    if not parsed.stations and not parsed.espp_rows and not parsed.control_rows:
        raise ValueError("В файле не найдено строк для импорта выработки ЭЭ.")
    return parsed


def _all_database_version_ids() -> list[int | None]:
    version_ids: list[int | None] = [None]
    for dv in DatabaseVersion.query.filter(DatabaseVersion.id.isnot(None)).order_by(DatabaseVersion.id).all():
        if dv.id is not None:
            version_ids.append(dv.id)
    return version_ids


def _build_res_id_by_name_map(version_id: int | None) -> dict[str, int]:
    result: dict[str, int] = {}
    q = filter_by_explicit_db_version(
        RegionalEnergySystem.query,
        RegionalEnergySystem,
        version_id,
    )
    for res in q.all():
        for candidate in (res.name, res.name_full, res.name_rp):
            norm = _normalize_name(candidate)
            if norm and norm not in result:
                result[norm] = res.id
    return result


def _build_stations_by_external_code_map(
    version_id: int | None,
    external_codes: set[str],
) -> dict[str, list[Station]]:
    if not external_codes:
        return {}
    q = Station.query.filter(Station.external_code.in_(external_codes))
    q = filter_by_explicit_db_version(q, Station, version_id)
    result: dict[str, list[Station]] = defaultdict(list)
    for station in q.all():
        if station.external_code:
            result[station.external_code].append(station)
    return result


def _period_items(values: ParsedPeriodValues) -> list[tuple[int, Decimal]]:
    periods: list[tuple[int, Decimal | None]] = [
        (STATION_ENERGY_GENERATION_PERIOD_YEAR, values.annual)
    ]
    periods.extend(sorted(values.months.items()))
    return [(month_number, value) for month_number, value in periods if value is not None]


def _load_generation_map(
    model_cls,
    version_id: int | None,
    year_number: int,
    lookup_field: str,
    lookup_ids: set[int],
) -> dict[tuple[int, int], Any]:
    if not lookup_ids:
        return {}
    q = model_cls.query.filter(
        getattr(model_cls, lookup_field).in_(lookup_ids),
        model_cls.year_number == year_number,
    )
    q = filter_by_explicit_db_version(q, model_cls, version_id)
    return {
        (getattr(rec, lookup_field), rec.month_number): rec
        for rec in q.all()
    }


def _apply_period_values(
    *,
    model_cls,
    lookup_field: str,
    lookup_id: int,
    lookup: dict[str, Any],
    values: ParsedPeriodValues,
    year_number: int,
    version_id: int | None,
    existing_map: dict[tuple[int, int], Any],
    value_attr: str = "electricity_generation",
) -> int:
    updated = 0
    for month_number, new_val in _period_items(values):
        key = (lookup_id, month_number)
        rec = existing_map.get(key)
        old_val = getattr(rec, value_attr, None) if rec else None
        if rec is not None and is_same_decimal(to_decimal(old_val), new_val):
            continue
        if rec is None:
            rec = model_cls(
                **lookup,
                year_number=year_number,
                month_number=month_number,
            )
            setattr(rec, value_attr, new_val)
            rec.database_version_id = version_id
            db.session.add(rec)
            existing_map[key] = rec
        else:
            setattr(rec, value_attr, new_val)
            db.session.add(rec)
        updated += 1
    return updated


def _resolve_espp_res_id(
    espp_row: ParsedResRow,
    res_map: dict[str, int],
    station_map: dict[str, list[Station]],
) -> int | None:
    if espp_row.res_name:
        res_id = res_map.get(_normalize_name(espp_row.res_name))
        if res_id is not None:
            return res_id
    anchor_code = espp_row.anchor_external_code
    if not anchor_code:
        return None
    for station in station_map.get(anchor_code, []):
        if station.id_regional_energy_system is not None:
            return int(station.id_regional_energy_system)
    return None


def _clear_kto_conflicts_in_version(
    kto_value: str,
    version_id: int | None,
    keep_station_id: int,
    kto_registry: dict[str, int],
) -> None:
    """Снимает kto с других станций той же версии (БД + текущий импорт)."""
    prev_station_id = kto_registry.get(kto_value)
    if prev_station_id and prev_station_id != keep_station_id:
        prev_station = db.session.get(Station, prev_station_id)
        if prev_station is not None and prev_station.kto == kto_value:
            prev_station.kto = None
            db.session.add(prev_station)

    conflict_query = Station.query.filter(
        Station.kto == kto_value,
        Station.id != keep_station_id,
    )
    conflict_query = filter_by_explicit_db_version(conflict_query, Station, version_id)
    conflict_query.update({Station.kto: None}, synchronize_session=False)


def _apply_station_kto(
    target: Station,
    kto_value: str,
    version_id: int | None,
    kto_registry: dict[str, int],
) -> bool:
    """Записывает КПО в пределах версии БД, снимая дубликат с других станций."""
    if target.kto == kto_value:
        kto_registry[kto_value] = target.id
        return False

    _clear_kto_conflicts_in_version(
        kto_value,
        version_id,
        target.id,
        kto_registry,
    )
    target.kto = kto_value
    kto_registry[kto_value] = target.id
    db.session.add(target)
    db.session.flush()
    return True


def _import_multi_year_parsed_workbook(
    parsed: ParsedWorkbook,
    version_id: int | None,
) -> tuple[dict[str, int], list[str]]:
    stats: dict[str, int] = defaultdict(int)
    errors: list[str] = []

    external_codes = {row.external_code for row in parsed.multi_year_stations}
    station_map = _build_stations_by_external_code_map(version_id, external_codes)
    station_ids = {station.id for stations in station_map.values() for station in stations}

    all_years = sorted(
        {
            year_number
            for row in parsed.multi_year_stations
            for year_number in row.values_by_year
        }
    )
    station_gen_maps: dict[int, dict[tuple[int, int], Any]] = {}
    for year_number in all_years:
        station_gen_maps[year_number] = _load_generation_map(
            StationEnergyGeneration,
            version_id,
            year_number,
            "id_station",
            station_ids,
        )

    kto_registry: dict[str, int] = {}

    with db.session.no_autoflush:
        for station_row in parsed.multi_year_stations:
            targets = station_map.get(station_row.external_code, [])
            if not targets:
                stats["stations_missing"] += 1
                if len(errors) < 30:
                    errors.append(
                        f"Станция external_code={station_row.external_code} не найдена "
                        f"(версия БД {version_id or 'без версии'})."
                    )
                continue

            for year_number, annual_value in station_row.values_by_year.items():
                if annual_value is None:
                    continue
                values = ParsedPeriodValues(annual=annual_value)
                for target in targets:
                    stats["station_rows_updated"] += _apply_period_values(
                        model_cls=StationEnergyGeneration,
                        lookup_field="id_station",
                        lookup_id=target.id,
                        lookup={"id_station": target.id},
                        values=values,
                        year_number=year_number,
                        version_id=version_id,
                        existing_map=station_gen_maps[year_number],
                    )

            for target in targets:
                station_changed = False
                if station_row.kto is not None:
                    if _apply_station_kto(
                        target,
                        station_row.kto,
                        version_id,
                        kto_registry,
                    ):
                        station_changed = True
                        stats["station_kto_updated"] += 1
                if (
                    station_row.station_sign is not None
                    and target.station_sign != station_row.station_sign
                ):
                    target.station_sign = station_row.station_sign
                    station_changed = True
                    stats["station_sign_updated"] += 1
                if station_changed:
                    db.session.add(target)

    db.session.flush()

    return dict(stats), errors


def _import_parsed_workbook(parsed: ParsedWorkbook, version_id: int | None) -> tuple[dict[str, int], list[str]]:
    if parsed.is_multi_year:
        return _import_multi_year_parsed_workbook(parsed, version_id)

    stats: dict[str, int] = defaultdict(int)
    errors: list[str] = []

    external_codes = {row.external_code for row in parsed.stations}
    external_codes.update(
        row.anchor_external_code
        for row in parsed.espp_rows
        if row.anchor_external_code
    )
    station_map = _build_stations_by_external_code_map(version_id, external_codes)
    res_map = _build_res_id_by_name_map(version_id)

    station_ids = {station.id for stations in station_map.values() for station in stations}
    station_gen_map = _load_generation_map(
        StationEnergyGeneration,
        version_id,
        parsed.year_number,
        "id_station",
        station_ids,
    )

    for station_row in parsed.stations:
        targets = station_map.get(station_row.external_code, [])
        if not targets:
            stats["stations_missing"] += 1
            if len(errors) < 30:
                errors.append(
                    f"Станция external_code={station_row.external_code} не найдена "
                    f"(версия БД {version_id or 'без версии'})."
                )
            continue
        for target in targets:
            with db.session.no_autoflush:
                stats["station_rows_updated"] += _apply_period_values(
                    model_cls=StationEnergyGeneration,
                    lookup_field="id_station",
                    lookup_id=target.id,
                    lookup={"id_station": target.id},
                    values=station_row.values,
                    year_number=parsed.year_number,
                    version_id=version_id,
                    existing_map=station_gen_map,
                )

    espp_res_ids: set[int] = set()
    for espp_row in parsed.espp_rows:
        res_id = _resolve_espp_res_id(espp_row, res_map, station_map)
        if res_id is not None:
            espp_res_ids.add(res_id)
    espp_gen_map = _load_generation_map(
        ESPPEnergyGeneration,
        version_id,
        parsed.year_number,
        "id_regional_energy_system",
        espp_res_ids,
    )

    for espp_row in parsed.espp_rows:
        res_id = _resolve_espp_res_id(espp_row, res_map, station_map)
        if res_id is None:
            stats["espp_res_missing"] += 1
            if len(errors) < 30:
                label = espp_row.res_name or espp_row.anchor_external_code or "?"
                errors.append(
                    f"РЭС для строки ЭСПП «{label}» не найдена "
                    f"(версия БД {version_id or 'без версии'})."
                )
            continue
        stats["espp_rows_updated"] += _apply_period_values(
            model_cls=ESPPEnergyGeneration,
            lookup_field="id_regional_energy_system",
            lookup_id=res_id,
            lookup={"id_regional_energy_system": res_id},
            values=espp_row.values,
            year_number=parsed.year_number,
            version_id=version_id,
            existing_map=espp_gen_map,
        )

    control_res_ids: set[int] = set()
    for control_row in parsed.control_rows:
        res_id = res_map.get(_normalize_name(control_row.res_name))
        if res_id is not None:
            control_res_ids.add(res_id)
    control_gen_map = _load_generation_map(
        RegionalEnergySystemEnergyGeneration,
        version_id,
        parsed.year_number,
        "id_regional_energy_system",
        control_res_ids,
    )

    for control_row in parsed.control_rows:
        res_id = res_map.get(_normalize_name(control_row.res_name))
        if res_id is None:
            stats["control_res_missing"] += 1
            if len(errors) < 30:
                errors.append(
                    f"Контрольная строка: РЭС «{control_row.res_name}» не найдена "
                    f"(версия БД {version_id or 'без версии'})."
                )
            continue
        stats["control_rows_updated"] += _apply_period_values(
            model_cls=RegionalEnergySystemEnergyGeneration,
            lookup_field="id_regional_energy_system",
            lookup_id=res_id,
            lookup={"id_regional_energy_system": res_id},
            values=control_row.values,
            year_number=parsed.year_number,
            version_id=version_id,
            existing_map=control_gen_map,
        )

    return dict(stats), errors


def import_ee_generation_from_excel(file, user: str) -> dict[str, Any]:
    raw = file.read()
    filename = getattr(file, "filename", None)
    parsed = parse_ee_generation_workbook(raw, filename=filename)

    version_ids = _all_database_version_ids()
    stats: dict[str, int] = defaultdict(int)
    errors: list[str] = []

    for version_id in version_ids:
        version_stats, version_errors = _import_parsed_workbook(parsed, version_id)
        for key, value in version_stats.items():
            stats[key] += value
        errors.extend(version_errors)
        if len(errors) > 30:
            errors = errors[:30]

    try:
        quick_fix_seq(SCHEMA_ENERGY_BALANCE, "gs_bem_station_energy_generations", "id")
        quick_fix_seq(SCHEMA_ENERGY_BALANCE, "gs_bem_espp_energy_generations", "id")
        quick_fix_seq(SCHEMA_ENERGY_BALANCE, "gs_bem_regional_energy_system_energy_generations", "id")
    except Exception:
        pass

    _commit_with_retry()

    from app.energy_balance.services.energy_balance_cache import clear_energy_balance_cache

    clear_energy_balance_cache()

    total_updated = (
        stats["station_rows_updated"]
        + stats["espp_rows_updated"]
        + stats["control_rows_updated"]
    )
    if parsed.is_multi_year:
        years = sorted(
            {
                year_number
                for row in parsed.multi_year_stations
                for year_number in row.values_by_year
            }
        )
        if len(years) == 1:
            years_label = f"{years[0]} г."
        else:
            years_label = f"{years[0]}–{years[-1]} гг."
        kto_updated = stats.get("station_kto_updated", 0)
        sign_updated = stats.get("station_sign_updated", 0)
        message = (
            f"Импорт выработки ЭЭ за {years_label} завершён. "
            f"Обновлено записей: {total_updated} "
            f"(станции: {stats['station_rows_updated']}"
            f"{f', КПО: {kto_updated}' if kto_updated else ''}"
            f"{f', признак эл.ст.: {sign_updated}' if sign_updated else ''}). "
            f"Строк в файле: {len(parsed.multi_year_stations)}. "
            f"Версий БД: {len(version_ids)}."
        )
    else:
        message = (
            f"Импорт выработки ЭЭ за {parsed.year_number} г. завершён. "
            f"Обновлено записей: {total_updated} "
            f"(станции: {stats['station_rows_updated']}, "
            f"ЭСПП: {stats['espp_rows_updated']}, "
            f"контроль: {stats['control_rows_updated']}). "
            f"Версий БД: {len(version_ids)}."
        )
    log_to_db(
        user,
        "Импорт выработки ЭЭ из Excel",
        message,
        entity_type="ee_generation_import",
        entity_id=0,
    )
    return {
        "message": message,
        "stats": dict(stats),
        "errors": errors,
        "errors_count": len(errors),
        "year_number": parsed.year_number,
    }
