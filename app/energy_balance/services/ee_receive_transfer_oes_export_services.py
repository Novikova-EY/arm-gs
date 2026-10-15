# -*- coding: utf-8 -*-
"""Экспорт приема/передачи электроэнергии между ОЭС в Excel."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.energy_balance.services.ee_receive_transfer_oes_page_services import (
    get_ee_receive_transfer_oes_page_data,
)


def export_ee_receive_transfer_oes_to_excel(
    *,
    start_year: int,
    end_year: int,
    rounding_digits: int,
    direction_filter: list[str] | None = None,
    oes_filter: list[int] | None = None,
) -> BytesIO:
    data = get_ee_receive_transfer_oes_page_data(
        start_year=start_year,
        end_year=end_year,
        rounding_digits=rounding_digits,
        direction_filter=direction_filter,
        oes_filter=oes_filter,
    )
    period_columns = data["period_columns"]
    columns = ["Прием/Передача", "Наименование ОЭС"] + [label for _, label in period_columns]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Прием-Передача по ОЭС"

    header_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    for column_index, title in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=column_index, value=title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_index = 2
    if not data["hierarchy"]:
        sheet.cell(row=row_index, column=1, value="Данные не найдены")
    else:
        for oes in data["hierarchy"]:
            sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=len(columns))
            title_cell = sheet.cell(row=row_index, column=1, value=oes["ues_name"])
            title_cell.font = Font(bold=True)
            title_cell.fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
            row_index += 1
            for transfer_row in oes["rows"]:
                sheet.cell(row=row_index, column=1, value=transfer_row["direction_label"])
                sheet.cell(row=row_index, column=2, value=transfer_row["country_name"])
                for column_index, (year, _label) in enumerate(period_columns, start=3):
                    sheet.cell(row=row_index, column=column_index, value=transfer_row["cell_values"][year])
                row_index += 1

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal="left" if cell.column <= 2 else "center", vertical="center", wrap_text=True
            )
    for column_index in range(1, len(columns) + 1):
        widest = max(
            (len(str(sheet.cell(row=row, column=column_index).value or "")) for row in range(1, sheet.max_row + 1)),
            default=10,
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = max(10, min(48, widest + 2))

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
