# -*- coding: utf-8 -*-
"""Выгрузка расчета балансов электрической энергии в Excel (все листы макета)."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook

from app.energy_balance.services.ee_balance_page_services import (
    EE_BALANCE_UNIT,
    EUROPEAN_EES_WITHOUT_SOUTH_URAL_NAME,
    EUROPEAN_EES_WITHOUT_SOUTH_URAL_SLUG,
    EUROPEAN_EES_WITHOUT_SOUTH_URAL_SOURCE_SLUGS,
    EUROPEAN_EES_SLUG,
    EES_RUSSIA_WITH_TITES_SLUG,
    SEVERO_ZAPAD_CENTER_SLUG,
    SYNCHRONOUS_ZONES_SLUG,
    URAL_WITHOUT_TYUMEN_SLUG,
    _append_european_ees_table,
    _append_severo_zapad_center_table,
    _append_ural_without_tyumen_table,
    _as_ee_sheet,
    _ees_rossii_with_tites_navigation_sheet,
    _ees_rossii_with_tites_source_slugs,
    _ee_balance_navigation_sheets,
    _european_ees_navigation_sheet,
    _overlay_ural_without_tyumen_manual_generation,
    _severo_zapad_region_sheets,
    _ural_tyumen_energy_system,
    build_ee_balance_tables,
    get_ee_balance_sheets,
    get_ee_balance_year_columns,
    get_ee_balance_year_features,
    group_ee_balance_sheets,
    resolve_ee_balance_rounding_digits,
)
from app.energy_balance.services.power_balance_page_services import (
    tites_east_child_nav_sheets,
)
from app.energy_balance.services.power_balance_excel_services import (
    _safe_sheet_title,
    _write_sheet,
)


def _add_export_sheet(sheets: list[dict[str, Any]], sheet: dict[str, Any] | None) -> None:
    if sheet is None:
        return
    slug = str(sheet.get("slug") or "")
    if not slug or any(str(item.get("slug") or "") == slug for item in sheets):
        return
    sheets.append(_as_ee_sheet(sheet))


def _ee_sheets_in_navigation_order(
    sheets: list[dict[str, Any]],
    tables: dict[str, dict[str, Any]],
    tyumen_sheet: dict[str, Any] | None,
    regional_sheets: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Order worksheets like the page buttons, expanding each button's subtabs inline."""
    sheets_by_slug = {str(sheet.get("slug") or ""): sheet for sheet in sheets}
    ordered: list[dict[str, Any]] = []
    added: set[str] = set()

    def append(
        slug: str,
        display_name: str | None = None,
        *,
        tab_slug: str | None = None,
    ) -> None:
        output_slug = tab_slug or slug
        if not slug or output_slug in added or slug not in tables:
            return
        sheet = dict(tables[slug].get("sheet") or sheets_by_slug.get(slug) or {})
        if not sheet:
            return
        sheet["slug"] = output_slug
        if output_slug != slug:
            sheet["data_slug"] = slug
        if display_name:
            sheet["sheet_name"] = display_name
        ordered.append(sheet)
        added.add(output_slug)

    navigation_base = [
        sheet
        for sheet in sheets
        if str(sheet.get("slug") or "")
        not in {
            EES_RUSSIA_WITH_TITES_SLUG,
            EUROPEAN_EES_SLUG,
            EUROPEAN_EES_WITHOUT_SOUTH_URAL_SLUG,
            SEVERO_ZAPAD_CENTER_SLUG,
            SYNCHRONOUS_ZONES_SLUG,
            URAL_WITHOUT_TYUMEN_SLUG,
        }
    ]
    navigation = _ee_balance_navigation_sheets(navigation_base)
    for group in group_ee_balance_sheets(navigation):
        for nav_sheet in group["sheets"]:
            slug = str(nav_sheet.get("slug") or "")
            if slug == SYNCHRONOUS_ZONES_SLUG:
                continue
            if nav_sheet.get("is_tites_east_hub"):
                for child in tites_east_child_nav_sheets(sheets, slug):
                    append(str(child.get("slug") or ""), child.get("sheet_name"))
                continue
            if slug == "vostok" and slug not in sheets_by_slug:
                append(
                    "2-sz-ees-vostok",
                    nav_sheet.get("sheet_name"),
                    tab_slug=slug,
                )
                continue
            if slug == "ural" and tyumen_sheet is not None:
                append(slug, "ОЭС Урала (с Тюменской ЭС)")
                append(URAL_WITHOUT_TYUMEN_SLUG)
                append(str(tyumen_sheet.get("slug") or ""), "Тюменская ЭС")
                continue
            if slug == "severo-zapad" and regional_sheets:
                append(slug, "ЭС ОЭС Северо-Запада")
                append(SEVERO_ZAPAD_CENTER_SLUG)
                for region in regional_sheets:
                    append(str(region.get("slug") or ""), region.get("sheet_name"))
                continue

            append(slug, nav_sheet.get("sheet_name"))
            if slug == EUROPEAN_EES_SLUG:
                append(EUROPEAN_EES_WITHOUT_SOUTH_URAL_SLUG)

    for sheet in sheets:
        append(str(sheet.get("slug") or ""), sheet.get("sheet_name"))
    return ordered


def _place_export_sheet(
    sheets: list[dict[str, Any]],
    sheet: dict[str, Any],
    after_slugs: tuple[str, ...] = (),
) -> None:
    """Insert a subtab balance right after its host tab, as in the tab buttons."""
    _add_export_sheet(sheets, sheet)
    slug = str(sheet.get("slug") or "")
    current = next(
        (index for index, item in enumerate(sheets) if item.get("slug") == slug),
        None,
    )
    if current is None:
        return
    hosts = [
        index for index, item in enumerate(sheets) if str(item.get("slug")) in after_slugs
    ]
    if not hosts:
        return
    insert_at = max(hosts) + 1
    if insert_at > current:
        insert_at -= 1
    if insert_at != current:
        sheets.insert(insert_at, sheets.pop(current))


def _build_all_ee_balance_export_tables(
    sheets: list[dict[str, Any]],
    years: list[int],
    rounding_digits: int,
) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    """Build regular and separately navigable EE balance tabs for one workbook."""
    export_sheets = [dict(sheet) for sheet in sheets]
    tyumen_sheet = _ural_tyumen_energy_system(export_sheets)
    regional_sheets = _severo_zapad_region_sheets(export_sheets)
    for extra_sheet in [tyumen_sheet, *regional_sheets]:
        _add_export_sheet(export_sheets, extra_sheet)

    ees_with_tites = _ees_rossii_with_tites_navigation_sheet()
    tites_sources = _ees_rossii_with_tites_source_slugs(export_sheets)
    if tites_sources:
        ees_with_tites["source_slugs"] = tites_sources
        _place_export_sheet(export_sheets, ees_with_tites, ("ees-rossii", "ees-tites"))

    european_sheet = _european_ees_navigation_sheet()
    _add_export_sheet(export_sheets, european_sheet)
    european_without_south_ural = dict(european_sheet)
    european_without_south_ural.update(
        slug=EUROPEAN_EES_WITHOUT_SOUTH_URAL_SLUG,
        sheet_name=EUROPEAN_EES_WITHOUT_SOUTH_URAL_NAME,
        title=f"Баланс электрической энергии {EUROPEAN_EES_WITHOUT_SOUTH_URAL_NAME}",
        source_slugs=EUROPEAN_EES_WITHOUT_SOUTH_URAL_SOURCE_SLUGS,
    )
    _place_export_sheet(export_sheets, european_without_south_ural, (EUROPEAN_EES_SLUG,))

    ural_sheet = next(
        (sheet for sheet in export_sheets if sheet.get("slug") == "ural"), None
    )
    if tyumen_sheet is not None and ural_sheet is not None:
        ural_without_tyumen = dict(ural_sheet)
        ural_without_tyumen.update(
            slug=URAL_WITHOUT_TYUMEN_SLUG,
            sheet_name="ОЭС Урала (без Тюменской ЭС)",
            title="Баланс электрической энергии ОЭС Урала (без Тюменской ЭС)",
            parent_slug="ural",
        )
        _place_export_sheet(export_sheets, ural_without_tyumen, ("ural",))

    severo_zapad_sheet = next(
        (sheet for sheet in export_sheets if sheet.get("slug") == "severo-zapad"),
        None,
    )
    if regional_sheets and severo_zapad_sheet is not None:
        center_sheet = dict(severo_zapad_sheet)
        center_sheet.update(
            slug=SEVERO_ZAPAD_CENTER_SLUG,
            sheet_name="ОЭС Северо-Запада (центр)",
            title="Баланс электрической энергии ОЭС Северо-Запада (центр)",
            parent_slug="severo-zapad",
        )
        _place_export_sheet(export_sheets, center_sheet, ("severo-zapad",))

    tables = build_ee_balance_tables(
        years, rounding_digits=rounding_digits, sheets=export_sheets
    )
    if EUROPEAN_EES_SLUG in tables:
        _append_european_ees_table(tables, years, rounding_digits)
        _append_european_ees_table(
            tables,
            years,
            rounding_digits,
            slug=EUROPEAN_EES_WITHOUT_SOUTH_URAL_SLUG,
            sheet_name=EUROPEAN_EES_WITHOUT_SOUTH_URAL_NAME,
            source_slugs=EUROPEAN_EES_WITHOUT_SOUTH_URAL_SOURCE_SLUGS,
        )
    if tyumen_sheet is not None and URAL_WITHOUT_TYUMEN_SLUG in tables:
        _append_ural_without_tyumen_table(
            tables, str(tyumen_sheet["slug"]), years, rounding_digits
        )
        _overlay_ural_without_tyumen_manual_generation(
            tables, years, rounding_digits
        )
    if regional_sheets and SEVERO_ZAPAD_CENTER_SLUG in tables:
        _append_severo_zapad_center_table(
            tables,
            [str(sheet["slug"]) for sheet in regional_sheets],
            years,
            rounding_digits,
        )

    return (
        _ee_sheets_in_navigation_order(
            export_sheets, tables, tyumen_sheet, regional_sheets
        ),
        tables,
    )


def ee_balance_excel_table_title(sheet: dict[str, Any]) -> str:
    """Первая строка листа: «Таблица N - Баланс электрической энергии …, млн.кВт·ч»."""
    number = sheet.get("table_number") or 0
    title = str(
        sheet.get("title") or f"Баланс электрической энергии {sheet.get('sheet_name') or ''}"
    ).strip()
    unit = str(sheet.get("unit") or EE_BALANCE_UNIT).strip() or EE_BALANCE_UNIT
    if title.lower().endswith(f", {unit.lower()}"):
        named = title
    else:
        named = f"{title}, {unit}"
    if number:
        return f"Таблица {int(number)} - {named}"
    return named


def export_ee_balance_to_excel(
    *,
    start_year: Any = None,
    end_year: Any = None,
    rounding_digits: Any = None,
    include_type_breakdown: bool = False,
    show_empty_rows: bool = True,
    show_flow_rows: bool = True,
    tables: dict[str, dict[str, Any]] | None = None,
    sheets: list[dict[str, Any]] | None = None,
    years: list[int] | None = None,
    year_features: dict[int, str] | None = None,
    hydro_year: Any = None,
) -> BytesIO:
    """Файл со вкладкой на каждый лист баланса электрической энергии."""
    del hydro_year
    include_navigation_tabs = sheets is None and tables is None
    digits = resolve_ee_balance_rounding_digits(rounding_digits)
    if years is None:
        years = get_ee_balance_year_columns(start_year, end_year)
    if sheets is None:
        sheets = get_ee_balance_sheets()
    if tables is None:
        if include_navigation_tabs:
            sheets, tables = _build_all_ee_balance_export_tables(
                sheets, years, digits
            )
        else:
            tables = build_ee_balance_tables(
                years, rounding_digits=digits, sheets=sheets
            )
    if year_features is None:
        year_features = get_ee_balance_year_features()

    wb = Workbook()
    used_titles: set[str] = set()
    first = True
    for sheet in sheets:
        if sheet.get("skip_table"):
            continue
        slug = sheet.get("slug")
        payload = (tables or {}).get(sheet.get("data_slug") or slug) or {}
        rows = payload.get("rows") or []
        if first:
            ws = wb.active
            first = False
        else:
            ws = wb.create_sheet()
        ws.title = _safe_sheet_title(str(sheet.get("sheet_name") or slug or "Лист"), used_titles)
        sheet_for_title = dict(payload.get("sheet") or sheet)
        sheet_for_title["title"] = ee_balance_excel_table_title(sheet_for_title).split(" - ", 1)[-1]
        # _write_sheet prepends «Таблица N - » again via power_balance_excel_table_title.
        # Pass a sheet whose title already includes unit; table_number keeps «Таблица N».
        _write_sheet(
            ws,
            sheet=sheet_for_title,
            rows=rows,
            years=years,
            year_features=year_features or {},
            rounding_digits=digits,
            include_type_breakdown=include_type_breakdown,
            show_empty_rows=show_empty_rows,
            show_flow_rows=show_flow_rows,
        )
        # Replace title cell with energy wording (helper writes power-balance title).
        ws.cell(row=1, column=1).value = ee_balance_excel_table_title(sheet_for_title)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
