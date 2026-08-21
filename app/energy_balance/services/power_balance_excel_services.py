# -*- coding: utf-8 -*-
"""Выгрузка расчета балансов мощности в Excel (все листы макета)."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.common.services.excel_numeric_cell import (
    excel_number_format as _excel_number_format,
    excel_numeric_cell_value as _excel_period_cell_value,
)
from app.energy_balance.services.power_balance_page_services import (
    KIND_CHILD,
    KIND_TOTAL,
    KIND_TRANSFER,
    POWER_BALANCE_UNIT,
    build_power_balance_tables,
    get_power_balance_sheets,
    get_power_balance_year_columns,
    get_power_balance_year_features,
    is_power_balance_flow_block_row,
    resolve_power_balance_rounding_digits,
)

HEADER_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="CFF4FC", end_color="CFF4FC", fill_type="solid")
FILL_TRANSFER = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
THIN = Border(
    left=Side(style="thin", color="B7B7B7"),
    right=Side(style="thin", color="B7B7B7"),
    top=Side(style="thin", color="B7B7B7"),
    bottom=Side(style="thin", color="B7B7B7"),
)


def power_balance_excel_table_title(sheet: dict[str, Any]) -> str:
    """Первая строка листа: «Таблица N - Баланс мощности …, МВт»."""
    number = sheet.get("table_number") or 0
    title = str(sheet.get("title") or f"Баланс мощности {sheet.get('sheet_name') or ''}").strip()
    unit = str(sheet.get("unit") or POWER_BALANCE_UNIT).strip() or POWER_BALANCE_UNIT
    if title.lower().endswith(f", {unit.lower()}"):
        named = title
    else:
        named = f"{title}, {unit}"
    if number:
        return f"Таблица {int(number)} - {named}"
    return named


def _safe_sheet_title(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[:\\/?*\[\]]", "_", str(name or "Лист")).strip() or "Лист"
    cleaned = cleaned[:31]
    base = cleaned
    index = 2
    while cleaned in used:
        suffix = f" ({index})"
        cleaned = (base[: 31 - len(suffix)] + suffix)[:31]
        index += 1
    used.add(cleaned)
    return cleaned


def _row_included(
    row: dict[str, Any],
    include_type_breakdown: bool,
    show_empty_rows: bool,
    show_flow_rows: bool = True,
) -> bool:
    if row.get("kind") == KIND_CHILD and not include_type_breakdown:
        return False
    if not show_empty_rows and row.get("hide_when_empty"):
        return False
    if not show_flow_rows and is_power_balance_flow_block_row(row):
        return False
    return True


def _row_fill(row: dict[str, Any]):
    kind = row.get("kind")
    if kind == KIND_TOTAL:
        return FILL_TOTAL
    if kind == KIND_TRANSFER:
        return FILL_TRANSFER
    return None


def _cell_value(row: dict[str, Any], year: int):
    raw_map = row.get("year_values_raw") or {}
    if year in raw_map:
        return raw_map[year]
    display_map = row.get("year_values") or {}
    if year in display_map:
        return display_map[year]
    return None


def _write_sheet(
    ws,
    *,
    sheet: dict[str, Any],
    rows: list[dict[str, Any]],
    years: list[int],
    year_features: dict[int, str],
    rounding_digits: int,
    include_type_breakdown: bool,
    show_empty_rows: bool,
    show_flow_rows: bool = True,
) -> None:
    num_cols = 1 + len(years)
    title = power_balance_excel_table_title(sheet)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(num_cols, 1))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=12)
    title_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22

    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    name_header = ws.cell(row=2, column=1, value="Наименование")
    name_header.font = header_font
    name_header.fill = HEADER_FILL
    name_header.alignment = header_alignment
    name_header.border = THIN
    for col_idx, year in enumerate(years, start=2):
        feature = (year_features or {}).get(year) or ""
        label = f"{year} г."
        if feature:
            label = f"{label}\n{feature}"
        cell = ws.cell(row=2, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = HEADER_FILL
        cell.alignment = header_alignment
        cell.border = THIN
    ws.row_dimensions[2].height = 32 if any((year_features or {}).values()) else 20

    excel_row = 3
    flow_group_child_rows: list[int] = []
    FLOW_GROUP_PARENT_KEY = "flow_total"
    FLOW_GROUP_SKIP_KEYS = {FLOW_GROUP_PARENT_KEY, "surplus_deficit_with_flow"}
    for row in rows:
        if not _row_included(row, include_type_breakdown, show_empty_rows, show_flow_rows):
            continue
        fill = _row_fill(row)
        bold = row.get("kind") == KIND_TOTAL
        italic = bool(row.get("italic"))
        indent = int(row.get("indent") or 0)
        label_cell = ws.cell(row=excel_row, column=1, value=row.get("label") or "")
        label_cell.font = Font(bold=bold, italic=italic, size=11)
        label_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
            indent=min(max(indent, 0), 15),
        )
        label_cell.border = THIN
        if fill is not None:
            label_cell.fill = fill
        for col_idx, year in enumerate(years, start=2):
            value, is_numeric = _excel_period_cell_value(
                _cell_value(row, year),
                zero_as_dash=False,
            )
            if value == "—" and _cell_value(row, year) in (None, ""):
                value = None
                is_numeric = False
            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(bold=bold, italic=italic, size=11)
            cell.border = THIN
            if fill is not None:
                cell.fill = fill
            if is_numeric:
                cell.number_format = _excel_number_format(rounding_digits)
        if (
            show_flow_rows
            and is_power_balance_flow_block_row(row)
            and str(row.get("key") or "") not in FLOW_GROUP_SKIP_KEYS
        ):
            flow_group_child_rows.append(excel_row)
        excel_row += 1

    for row_idx in flow_group_child_rows:
        ws.row_dimensions[row_idx].outline_level = 1
    outline = ws.sheet_properties.outlinePr
    outline.summaryBelow = False

    ws.column_dimensions["A"].width = 72
    for col_idx in range(2, num_cols + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
    ws.freeze_panes = "B3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:2"


def export_power_balance_to_excel(
    *,
    start_year: Any = None,
    end_year: Any = None,
    rounding_digits: Any = None,
    include_type_breakdown: bool = False,
    show_empty_rows: bool = True,
    show_flow_rows: bool = True,
    tables: dict[str, dict[str, Any]] | None = None,
    sheets: list[dict[str, Any]] | None = None,
    years: list[int] | None = None,
    year_features: dict[int, str] | None = None,
) -> BytesIO:
    """Файл со вкладкой на каждый лист баланса мощности."""
    digits = resolve_power_balance_rounding_digits(rounding_digits)
    if years is None:
        years = get_power_balance_year_columns(start_year, end_year)
    if sheets is None:
        sheets = get_power_balance_sheets()
    if tables is None:
        tables = build_power_balance_tables(years, rounding_digits=digits)
    if year_features is None:
        year_features = get_power_balance_year_features()

    wb = Workbook()
    used_titles: set[str] = set()
    first = True
    for sheet in sheets:
        slug = sheet.get("slug")
        payload = (tables or {}).get(slug) or {}
        rows = payload.get("rows") or []
        if first:
            ws = wb.active
            first = False
        else:
            ws = wb.create_sheet()
        ws.title = _safe_sheet_title(str(sheet.get("sheet_name") or slug or "Лист"), used_titles)
        _write_sheet(
            ws,
            sheet=payload.get("sheet") or sheet,
            rows=rows,
            years=years,
            year_features=year_features or {},
            rounding_digits=digits,
            include_type_breakdown=include_type_breakdown,
            show_empty_rows=show_empty_rows,
            show_flow_rows=show_flow_rows,
        )

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
