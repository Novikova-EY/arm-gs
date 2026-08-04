# -*- coding: utf-8 -*-
"""Экспорт страницы «Выработка ЭЭ» в Excel."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.common.services.excel_numeric_cell import (
    excel_number_format as _excel_number_format,
    excel_numeric_cell_value as _excel_period_cell_value,
)
from app.common.services.help_services import format_decimal_for_display
from app.energy_balance.services.station_ee_generation_page_services import (
    get_station_ee_generation_page_data,
    is_verification_nonzero,
)

from app.fuel.services.fuel_exports.hierarchy_excel_layout import (
    FILL_EST,
    FILL_RES_HEADER,
    FILL_STATION_SUMMARY,
    FILL_UES,
    write_merged_section_row,
)

FILL_EU = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
FILL_TOTAL_RU = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
FILL_VERIFY = PatternFill(start_color="E2ECF7", end_color="E2ECF7", fill_type="solid")

STATIC_COLUMNS = [
    "КТО",
    "Электростанция",
    "Тип электростанции",
    "Признак эл.ст.",
    "Тип ТЭС",
    "Тип агрегата ТЭС",
    "Основное топливо",
    "Топливо (по СО ЕЭС)",
]

_PERIOD_COL_START = len(STATIC_COLUMNS) + 1


def _write_period_cell(
    ws,
    row_idx: int,
    col_idx: int,
    value,
    rounding_digits: int,
    *,
    verification: bool = False,
    bold: bool = False,
    fill=None,
    red: bool = False,
) -> None:
    cell_value, is_numeric = _excel_period_cell_value(
        value, verification=verification, zero_as_dash=True
    )
    cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if bold:
        cell.font = Font(bold=True)
    if fill is not None:
        cell.fill = fill
    if red:
        cell.font = Font(bold=bold, color="FF0000")
    if is_numeric:
        cell.number_format = _excel_number_format(rounding_digits)


def _write_data_row(
    ws,
    row_idx: int,
    static_values: list,
    period_values: list | None = None,
    *,
    bold: bool = False,
    fill=None,
    red_period_cols: set[int] | None = None,
    rounding_digits: int = 1,
    verification: bool = False,
) -> None:
    col_idx = 1
    for value in static_values:
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = Alignment(
            horizontal="left" if col_idx == 2 else "center",
            vertical="center",
            wrap_text=True,
        )
        if bold:
            cell.font = Font(bold=True)
        if fill is not None:
            cell.fill = fill
        col_idx += 1

    if period_values is not None:
        for offset, value in enumerate(period_values):
            period_col = _PERIOD_COL_START + offset
            _write_period_cell(
                ws,
                row_idx,
                period_col,
                value,
                rounding_digits,
                verification=verification,
                bold=bold,
                fill=fill,
                red=period_col in (red_period_cols or set()),
            )


def _write_total_row(
    ws,
    row_idx: int,
    label: str,
    period_columns: list[tuple[int, str]],
    period_totals: dict,
    rounding_digits: int,
    *,
    fill=None,
) -> int:
    static_values = ["", label, "", "", "", "", "", ""]
    period_values = [period_totals.get(period_key) for period_key, _label in period_columns]
    _write_data_row(
        ws,
        row_idx,
        static_values,
        period_values,
        bold=True,
        fill=fill,
        rounding_digits=rounding_digits,
    )
    return row_idx + 1


def _write_verification_row(
    ws,
    row_idx: int,
    label: str,
    period_columns: list[tuple[int, str]],
    period_totals: dict,
    rounding_digits: int,
) -> int:
    static_values = ["", label, "", "", "", "", "", ""]
    period_values = []
    red_cols: set[int] = set()
    for idx, (period_key, _label) in enumerate(period_columns, start=_PERIOD_COL_START):
        val = period_totals.get(period_key)
        period_values.append(val)
        if is_verification_nonzero(val):
            red_cols.add(idx)
    _write_data_row(
        ws,
        row_idx,
        static_values,
        period_values,
        bold=True,
        fill=FILL_VERIFY,
        red_period_cols=red_cols,
        rounding_digits=rounding_digits,
        verification=True,
    )
    return row_idx + 1


def _write_sign_group_rows(
    ws,
    row_idx: int,
    sign_group: dict[str, Any],
    period_columns: list[tuple[int, str]],
    rounding_digits: int,
) -> int:
    stations = sign_group.get("stations") or []
    svod_keys = set(sign_group.get("espp_svod_period_keys") or [])
    svod_periods = sign_group.get("periods") or {}
    start_row = row_idx

    for station_idx, station_row in enumerate(stations):
        station_periods = station_row.get("periods") or {}
        static_values = [
            station_row.get("kto_display") or "—",
            station_row.get("station_name") or "—",
            station_row.get("station_type_name") or "—",
            station_row.get("station_sign_display") or "—",
            station_row.get("tes_types") or "—",
            station_row.get("tes_machine_type") or "—",
            station_row.get("primary_fuel") or "—",
            station_row.get("fuel_so") or "—",
        ]
        period_values = []
        for period_key, _label in period_columns:
            if period_key in svod_keys:
                # Значение свода пишем только в первой строке группы; остальные
                # ячейки объединяются ниже.
                period_values.append(
                    svod_periods.get(period_key) if station_idx == 0 else None
                )
            else:
                period_values.append(station_periods.get(period_key))
        _write_data_row(
            ws,
            row_idx,
            static_values,
            period_values,
            rounding_digits=rounding_digits,
        )
        row_idx += 1

    if svod_keys and len(stations) > 1:
        end_row = row_idx - 1
        for offset, (period_key, _label) in enumerate(period_columns):
            if period_key not in svod_keys:
                continue
            col_idx = _PERIOD_COL_START + offset
            ws.merge_cells(
                start_row=start_row,
                start_column=col_idx,
                end_row=end_row,
                end_column=col_idx,
            )
            merged = ws.cell(row=start_row, column=col_idx)
            merged.alignment = Alignment(horizontal="center", vertical="center")

    return row_idx


def export_ee_generation_to_excel(
    *,
    filters: dict,
    period_mode: str,
    selected_year: int | None,
    start_year: int | None,
    end_year: int | None,
    rounding_digits: int,
    show_totals: bool,
    export_verification: bool,
    per_page: int | str = "all",
    page: int = 1,
) -> BytesIO | None:
    """Формирует Excel-файл с тем же содержимым, что отображается на странице."""
    page_data = get_station_ee_generation_page_data(
        filters=filters,
        page=page,
        per_page=per_page,
        period_mode=period_mode,
        selected_year=selected_year,
        start_year=start_year,
        end_year=end_year,
        rounding_digits=rounding_digits,
        show_totals=show_totals,
    )

    hierarchy = page_data.get("hierarchy") or []
    period_columns = page_data.get("period_columns") or []
    generation_aggregates = page_data.get("generation_aggregates") or {}
    should_show_totals = page_data.get("should_show_totals") or {}
    res_show_rd_level_map = page_data.get("res_show_rd_level_map") or {}
    res_verification_by_res = page_data.get("res_verification_by_res") or {}

    columns = STATIC_COLUMNS + [label for _key, label in period_columns]
    num_cols = len(columns)

    wb = Workbook()
    ws = wb.active
    ws.title = "Выработка ЭЭ"

    header_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num, title in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_num, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    if not hierarchy:
        ws.cell(row=2, column=1, value="Электростанции не найдены")
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    row_idx = 2
    for est_block in hierarchy:
        row_idx = write_merged_section_row(
            ws, row_idx, num_cols, est_block.get("est_name") or "—", FILL_EST
        )
        for ues_block in est_block.get("ues_list") or []:
            row_idx = write_merged_section_row(
                ws,
                row_idx,
                num_cols,
                ues_block.get("ues_name") or "—",
                FILL_UES,
            )
            for res_block in ues_block.get("res_list") or []:
                res_id = res_block.get("res_id")
                row_idx = write_merged_section_row(
                    ws,
                    row_idx,
                    num_cols,
                    res_block.get("res_name") or "—",
                    FILL_RES_HEADER,
                )
                show_rd_level = res_show_rd_level_map.get(res_id, False)
                for rd_block in res_block.get("rd_list") or []:
                    rd_id = rd_block.get("rd_id")
                    rd_name = rd_block.get("rd_name") or "—"
                    show_rd_totals = (
                        show_rd_level
                        and should_show_totals.get("regional_districts", {}).get(rd_id)
                    )
                    if show_rd_level and rd_name.lower() != "не указано" and show_rd_totals:
                        row_idx = write_merged_section_row(
                            ws, row_idx, num_cols, rd_name, FILL_STATION_SUMMARY
                        )
                    for eu_block in rd_block.get("eu_list") or []:
                        eu_id = eu_block.get("eu_id")
                        eu_name = eu_block.get("eu_name") or "—"
                        skip_eu = (not eu_id) or eu_name.lower() in ("не указано", "без энергоузла")
                        if not skip_eu and eu_name.lower() != "не указано":
                            row_idx = write_merged_section_row(
                                ws, row_idx, num_cols, eu_name, FILL_EU
                            )
                        for sign_group in eu_block.get("sign_groups") or []:
                            row_idx = _write_sign_group_rows(
                                ws, row_idx, sign_group, period_columns, rounding_digits
                            )
                        if not skip_eu and should_show_totals.get("energy_units", {}).get(eu_id):
                            eu_total_name = eu_name
                            row_idx = _write_total_row(
                                ws,
                                row_idx,
                                f"{eu_total_name}, всего",
                                period_columns,
                                generation_aggregates.get("energy_units", {}).get(eu_id, {}),
                                rounding_digits,
                                fill=FILL_EU,
                            )
                    if show_rd_totals and rd_name.lower() != "не указано":
                        row_idx = _write_total_row(
                            ws,
                            row_idx,
                            f"{rd_name}, всего",
                            period_columns,
                            generation_aggregates.get("regional_districts", {}).get(rd_id, {}),
                            rounding_digits,
                            fill=FILL_STATION_SUMMARY,
                        )
                if should_show_totals.get("regional_energy_systems", {}).get(res_id):
                    res_name = res_block.get("res_name") or "—"
                    row_idx = _write_total_row(
                        ws,
                        row_idx,
                        f"{res_name}, всего",
                        period_columns,
                        generation_aggregates.get("regional_energy_systems", {}).get(res_id, {}),
                        rounding_digits,
                        fill=FILL_RES_HEADER,
                    )
                    if export_verification:
                        row_idx = _write_verification_row(
                            ws,
                            row_idx,
                            f"Проверка для {res_name}",
                            period_columns,
                            res_verification_by_res.get(res_id, {}),
                            rounding_digits,
                        )
            ues_id = ues_block.get("ues_id")
            if should_show_totals.get("union_energy_systems", {}).get(ues_id):
                ues_label = ues_block.get("ues_name") or "—"
                row_idx = _write_total_row(
                    ws,
                    row_idx,
                    f"{ues_label}, всего",
                    period_columns,
                    generation_aggregates.get("union_energy_systems", {}).get(ues_id, {}),
                    rounding_digits,
                    fill=FILL_UES,
                )
    if should_show_totals.get("total"):
        row_idx = _write_total_row(
            ws,
            row_idx,
            "Россия, всего",
            period_columns,
            generation_aggregates.get("total", {}),
            rounding_digits,
            fill=FILL_TOTAL_RU,
        )

    max_lengths = [len(c) for c in columns]
    period_col_offset = len(STATIC_COLUMNS)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i, cell in enumerate(row):
            if i >= len(max_lengths):
                continue
            val = cell.value
            if val is None:
                continue
            if i >= period_col_offset and isinstance(val, (int, float)):
                display = format_decimal_for_display(val, digits=rounding_digits)
            else:
                display = str(val)
            max_lengths[i] = max(max_lengths[i], len(display))
    for col_idx, width in enumerate(max_lengths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(48, width + 2))

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
