# -*- coding: utf-8 -*-
"""Экспорт итоговых сумм выработки ЭЭ в Excel (как на экране)."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.common.services.excel_numeric_cell import (
    excel_number_format,
    excel_numeric_cell_value,
)

_HEADER_FILL = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
_ROW_FILLS = {
    "table-primary": PatternFill(start_color="CFE2FF", end_color="CFE2FF", fill_type="solid"),
    "table-secondary": PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid"),
    "table-info": PatternFill(start_color="CFF4FC", end_color="CFF4FC", fill_type="solid"),
    "table-warning": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
    "table-danger": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
}


def export_ee_generation_totals_summary_to_excel(context: dict[str, Any]) -> BytesIO:
    period_columns = list(context.get("period_columns") or [])
    rows = list(context.get("rows") or [])
    # 0 = «Не округлять» — нельзя подменять через `or 1`
    try:
        rounding_digits = int(context.get("rounding_digits"))
    except (TypeError, ValueError):
        rounding_digits = 1
    if rounding_digits < -1:
        rounding_digits = 1
    year_features = context.get("year_features") or {}

    wb = Workbook()
    ws = wb.active
    ws.title = "Итоговые значения"

    headers = [""] + [str(label) for _key, label in period_columns]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(
            wrap_text=True, horizontal="center", vertical="center"
        )
        if col_idx > 1:
            period_key = period_columns[col_idx - 2][0]
            feature = year_features.get(period_key)
            if feature:
                cell.value = f"{headers[col_idx - 1]}\n{feature}\nмлн кВтч"
            else:
                cell.value = f"{headers[col_idx - 1]}\nмлн кВтч"

    for row in rows:
        if row.get("is_verification") and not context.get("show_verification"):
            continue
        fill = _ROW_FILLS.get(row.get("row_class") or "")
        if row.get("is_verification"):
            fill = PatternFill(start_color="E2ECF7", end_color="E2ECF7", fill_type="solid")
        label = row.get("label") or ""
        if row.get("level") == "res":
            label = f"    {label}"
        excel_row = [label]
        period_values = row.get("period_values") or {}
        for period_key, _label in period_columns:
            excel_row.append(period_values.get(period_key))
        ws.append(excel_row)
        row_idx = ws.max_row
        for col_idx in range(1, len(excel_row) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if fill:
                cell.fill = fill
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = Font(bold=True)
                continue
            value, is_numeric = excel_numeric_cell_value(
                excel_row[col_idx - 1], zero_as_dash=True
            )
            cell.value = value
            if is_numeric:
                cell.number_format = excel_number_format(rounding_digits)

    ws.column_dimensions["A"].width = 42
    for col_idx in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 36

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
