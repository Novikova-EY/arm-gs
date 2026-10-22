# -*- coding: utf-8 -*-
"""Экспорт страницы СЭС/ВЭС в Excel."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.energy_balance.services.ee_balance_ves_ses_page_services import (
    build_ves_ses_page_context,
)

HEADER_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_RUSSIA = PatternFill(start_color="BADBCC", end_color="BADBCC", fill_type="solid")
FILL_UES = PatternFill(start_color="CFE2FF", end_color="CFE2FF", fill_type="solid")
FILL_INCL = PatternFill(start_color="E7F1FF", end_color="E7F1FF", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
FILL_HEADER = PatternFill(start_color="E7F1FF", end_color="E7F1FF", fill_type="solid")
HEADER_FONT = Font(bold=True, size=11)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _year_header(year: int, year_features: dict[int, str]) -> str:
    feature = (year_features or {}).get(int(year)) or ""
    if feature:
        return f"{year}\n{feature}"
    return str(year)


def _cell(ws, row: int, col: int, value: Any, *, fill=None, bold: bool = False, center: bool = True):
    cell = ws.cell(row=row, column=col, value="" if value in (None, "") else value)
    cell.alignment = CENTER if center else LEFT
    if bold:
        cell.font = Font(bold=True)
    if fill is not None:
        cell.fill = fill
    return cell


def _row_fill(kind: str):
    if kind == "russia_total":
        return FILL_RUSSIA
    if kind == "ues_total":
        return FILL_UES
    if kind == "type_total":
        return FILL_INCL
    if kind in {"res_total", "rd_total", "est_total"}:
        return FILL_TOTAL
    return None


def _metric_values(row: dict[str, Any], metric: str, years: list[int]) -> list[Any]:
    display_key = {
        "capacity": "capacity_display",
        "energy": "energy_display",
        "chiim": "chiim_display",
    }[metric]
    values = row.get(display_key) or {}
    return [values.get(year) or "–" for year in years]


OWNER_TITLE = "Выработка электроэнергии электростанцией (от собственника), млн кВтч"


def export_ves_ses_to_excel(**context_kwargs: Any) -> BytesIO:
    context = build_ves_ses_page_context(**context_kwargs)
    years = context.get("years") or []
    owner_years = context.get("owner_years") or []
    year_features = context.get("year_features") or {}
    rows = context.get("rows") or []
    min_label = context.get("min_label") or "Минимальное значение выработки"
    exploitation_date_label = (
        context.get("exploitation_date_label")
        or "Ожидаемая дата ввода в эксплуатацию"
    )
    col_count = 3 + len(years) + 2 + len(owner_years)
    header_rows = 2 if owner_years else 1

    wb = Workbook()
    ws = wb.active
    ws.title = "СЭС и ВЭС"

    def _header_cell(row: int, col: int, value: Any):
        cell = _cell(ws, row, col, value, fill=HEADER_FILL, bold=True)
        cell.font = HEADER_FONT
        return cell

    left_headers = ["Наименование", "Тип станции", "Тип данных"]
    left_headers.extend(_year_header(year, year_features) for year in years)
    left_headers.append(exploitation_date_label)
    left_headers.append(min_label)
    for col, title in enumerate(left_headers, start=1):
        _header_cell(1, col, title)
        if header_rows == 2:
            ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)

    if owner_years:
        owner_start = len(left_headers) + 1
        owner_end = owner_start + len(owner_years) - 1
        for offset, year in enumerate(owner_years):
            col = owner_start + offset
            if offset == 0:
                title_cell = _header_cell(1, col, OWNER_TITLE)
                title_cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            else:
                _header_cell(1, col, None)
            _header_cell(2, col, _year_header(year, year_features))
        if owner_end > owner_start:
            ws.merge_cells(
                start_row=1, start_column=owner_start, end_row=1, end_column=owner_end
            )
        ws.row_dimensions[1].height = 36

    excel_row = header_rows + 1

    def write_metric_row(
        *,
        name: str | None,
        type_name: str | None,
        data_label: str,
        metric: str,
        source: dict[str, Any],
        fill=None,
        bold: bool = False,
        include_min_owner: bool = False,
    ) -> None:
        nonlocal excel_row
        col = 1
        _cell(ws, excel_row, col, name or "", fill=fill, bold=bold, center=False)
        col += 1
        _cell(ws, excel_row, col, type_name or "", fill=fill, bold=bold)
        col += 1
        _cell(ws, excel_row, col, data_label, fill=fill, bold=bold, center=False)
        col += 1
        for value in _metric_values(source, metric, years):
            _cell(ws, excel_row, col, value, fill=fill, bold=bold)
            col += 1
        if include_min_owner:
            _cell(
                ws,
                excel_row,
                col,
                source.get("exploitation_date_expected") or "–",
                fill=fill,
                bold=bold,
            )
            col += 1
            _cell(ws, excel_row, col, source.get("min_display") or "–", fill=fill, bold=bold)
            col += 1
            owner = source.get("owner_display") or {}
            for year in owner_years:
                _cell(ws, excel_row, col, owner.get(year) or "–", fill=fill, bold=bold)
                col += 1
        else:
            _cell(ws, excel_row, col, "–", fill=fill, bold=bold)
            col += 1
            _cell(ws, excel_row, col, "–", fill=fill, bold=bold)
            col += 1
            for _year in owner_years:
                _cell(ws, excel_row, col, "–", fill=fill, bold=bold)
                col += 1
        excel_row += 1

    for row in rows:
        kind = row.get("kind")
        if kind in {"est_header", "ues_header", "res_header", "rd_header"}:
            ws.merge_cells(
                start_row=excel_row, start_column=1, end_row=excel_row, end_column=col_count
            )
            cell = _cell(ws, excel_row, 1, row.get("label") or "", fill=FILL_HEADER, bold=True, center=False)
            cell.font = Font(bold=True)
            excel_row += 1
            continue
        fill = _row_fill(kind or "")
        bold = kind != "station"
        name = row.get("name") or row.get("label") or ""
        type_name = row.get("type_name") or ""
        start = excel_row
        write_metric_row(
            name=name,
            type_name=type_name,
            data_label=row.get("capacity_label") or "",
            metric="capacity",
            source=row,
            fill=fill,
            bold=bold,
        )
        write_metric_row(
            name=None,
            type_name=None,
            data_label=row.get("energy_label") or "",
            metric="energy",
            source=row,
            fill=fill,
            bold=bold,
            include_min_owner=True,
        )
        write_metric_row(
            name=None,
            type_name=None,
            data_label=row.get("chiim_label") or "",
            metric="chiim",
            source=row,
            fill=fill,
            bold=bold,
        )
        if excel_row - start == 3:
            ws.merge_cells(start_row=start, start_column=1, end_row=start + 2, end_column=1)
            ws.merge_cells(start_row=start, start_column=2, end_row=start + 2, end_column=2)

    ws.freeze_panes = f"D{header_rows + 1}"
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 28
    for col in range(4, col_count + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
