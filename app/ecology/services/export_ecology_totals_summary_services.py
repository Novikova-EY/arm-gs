# -*- coding: utf-8 -*-
"""Экспорт итоговых сумм модуля «Экология» в Excel (как на экране)."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
_ROW_FILLS = {
    "table-primary": PatternFill(start_color="CFE2FF", end_color="CFE2FF", fill_type="solid"),
    "table-success": PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid"),
    "table-secondary": PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid"),
    "table-info": PatternFill(start_color="CFF4FC", end_color="CFF4FC", fill_type="solid"),
    "table-warning": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
    "table-danger": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
}
_STATION_KEYS = frozenset({"NAME", "Stancia"})


def export_ecology_totals_summary_to_excel(context: dict[str, Any]) -> BytesIO:
    """Выгрузка таблицы итоговых сумм экологии: шапка и блоки как на экране."""
    columns = list(context.get("columns") or [])
    header_rows = list(context.get("header_rows") or [])
    blocks = list(context.get("blocks") or [])
    rowspan_keys = set(context.get("rowspan_column_keys") or [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Итоговые суммы"

    col_count = len(columns) or 1
    # Упрощённая шапка: последняя неключевая строка с подписями столбцов,
    # либо по ключам/лейблам visible columns.
    label_by_key = {c.get("key"): (c.get("label") or c.get("key") or "") for c in columns}
    header_labels: list[str] = []
    for column in columns:
        key = column.get("key")
        if key in _STATION_KEYS:
            header_labels.append("")
            continue
        header_labels.append(str(label_by_key.get(key) or key or ""))

    # Если есть header_rows – берём текст из последней строки с column_key
    if header_rows:
        by_key_text: dict[str, str] = {}
        for row in header_rows:
            for cell in row or []:
                key = cell.get("column_key")
                if not key:
                    continue
                if key in _STATION_KEYS:
                    by_key_text[key] = ""
                else:
                    by_key_text[key] = str(cell.get("text") or "")
        header_labels = [
            by_key_text.get(column.get("key"), header_labels[idx] if idx < len(header_labels) else "")
            for idx, column in enumerate(columns)
        ]

    ws.append(header_labels)
    for col_idx in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

    for block in blocks:
        fill = _ROW_FILLS.get(block.get("row_class") or "")
        for row in block.get("rows") or []:
            cells = list(row.get("cells") or [])
            # Выравниваем длину под columns
            if len(cells) < col_count:
                cells.extend([""] * (col_count - len(cells)))
            elif len(cells) > col_count:
                cells = cells[:col_count]
            ws.append(cells)
            excel_row = ws.max_row
            for col_idx, column in enumerate(columns, start=1):
                cell = ws.cell(row=excel_row, column=col_idx)
                key = column.get("key")
                is_station = key in _STATION_KEYS
                cell.alignment = Alignment(
                    horizontal="left" if is_station else "center",
                    vertical="center",
                    wrap_text=False,
                )
                if fill:
                    cell.fill = fill
                if is_station or key in rowspan_keys:
                    cell.font = Font(bold=True)

    for col_idx, column in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        key = column.get("key")
        if key in _STATION_KEYS:
            ws.column_dimensions[letter].width = 28
        elif key == "Year":
            ws.column_dimensions[letter].width = 10
        else:
            ws.column_dimensions[letter].width = 14

    freeze_col = 1
    for idx, column in enumerate(columns, start=1):
        if column.get("key") == "Year":
            freeze_col = idx + 1
            break
        if column.get("sticky"):
            freeze_col = idx + 1
    ws.freeze_panes = f"{get_column_letter(freeze_col)}2"

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
