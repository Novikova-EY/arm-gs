# -*- coding: utf-8 -*-
"""Экспорт отчетных данных ГИС ТЭК в Excel."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from werkzeug.datastructures import MultiDict
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.ecology.services.gis_tek_constants import PAGE_TITLES, SHEET_NAMES, ZV_NORM_ATTR
from app.ecology.services.gis_tek_page_services import (
    build_gis_tek_page_context,
    norm_zv_caption,
)


def _style_header_cell(cell, fill, font, alignment):
    cell.fill = fill
    cell.font = font
    cell.alignment = alignment


def build_gis_tek_excel_stream(kind: str, args, user=None) -> BytesIO:
    export_args = MultiDict(args)
    export_args["per_page"] = "all"
    context = build_gis_tek_page_context(kind, export_args, user=user)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAMES[kind][:31]

    header_fill = PatternFill("solid", fgColor="C6EFCE")
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.cell(1, 2, PAGE_TITLES[kind])

    identity_columns = list(context.get("identity_columns") or [])
    # В выгрузке к каждой станции добавляем её ОЭС и РЭС.
    insert_at = next(
        (index + 1 for index, column in enumerate(identity_columns) if column[0] == "station_numb"),
        0,
    )
    export_columns = list(identity_columns)
    export_columns[insert_at:insert_at] = [
        ("ues_name", "ОЭС", False),
        ("res_name", "РЭС", False),
    ]
    identity_headers = [label for _key, label, _num in export_columns]
    identity_index_by_key = {
        key: index for index, (key, _label, _num) in enumerate(export_columns, start=1)
    }
    groups = list(context.get("value_column_groups") or [])
    value_columns = list(context.get("value_columns") or [])

    norm_col: int | None = None
    if groups:
        banner = str(context.get("value_column_banner") or "").strip()
        header_top = 3
        group_row = header_top + 1 if banner else header_top
        header_bottom = group_row + 1
        for col, title in enumerate(identity_headers, start=1):
            cell = ws.cell(header_top, col, title)
            _style_header_cell(cell, header_fill, header_font, center)
            ws.merge_cells(
                start_row=header_top,
                start_column=col,
                end_row=header_bottom,
                end_column=col,
            )
        value_start = len(identity_headers) + 1
        if banner:
            cell = ws.cell(header_top, value_start, banner)
            _style_header_cell(cell, header_fill, header_font, center)
            ws.merge_cells(
                start_row=header_top,
                start_column=value_start,
                end_row=header_top,
                end_column=value_start + len(value_columns) - 1,
            )
        col = value_start
        for group_label, group_cols in groups:
            n = len(group_cols)
            if n == 1 and group_cols[0][0] == ZV_NORM_ATTR:
                # Столбец «Нормируемые ЗВ» подписываем перечнем выбранных слагаемых.
                norm_col = col
                group_label = norm_zv_caption(context.get("norm_zv_selected"))
            cell = ws.cell(group_row, col, group_label)
            _style_header_cell(cell, header_fill, header_font, center)
            if n == 1 and not group_cols[0][1]:
                ws.merge_cells(
                    start_row=group_row,
                    start_column=col,
                    end_row=header_bottom,
                    end_column=col,
                )
                col += 1
                continue
            if n > 1:
                ws.merge_cells(
                    start_row=group_row,
                    start_column=col,
                    end_row=group_row,
                    end_column=col + n - 1,
                )
            for offset, (_attr, label) in enumerate(group_cols):
                sub = ws.cell(header_bottom, col + offset, label)
                _style_header_cell(sub, header_fill, header_font, center)
            col += n
        if norm_col is not None:
            # Подпись слагаемых в шапке занимает несколько строк.
            ws.row_dimensions[group_row].height = 30
            ws.row_dimensions[header_bottom].height = 30
        data_row = header_bottom + 1
        total_cols = col - 1
    else:
        headers = identity_headers + [label for _attr, label in value_columns]
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(3, col, title)
            _style_header_cell(cell, header_fill, header_font, center)
        data_row = 4
        total_cols = len(headers)

    row_idx = data_row
    show_empty_rows = bool(context.get("show_empty_rows"))
    identity_count = len(export_columns)

    def _write_summary(
        summary: dict | None,
        ues_name: str = "",
        res_name: str = "",
    ) -> None:
        nonlocal row_idx
        if not summary:
            return
        if not show_empty_rows and summary.get("is_empty"):
            return
        ws.cell(row_idx, 1, summary.get("label"))
        if "ues_name" in identity_index_by_key:
            ws.cell(row_idx, identity_index_by_key["ues_name"], ues_name or None)
        if "res_name" in identity_index_by_key:
            ws.cell(row_idx, identity_index_by_key["res_name"], res_name or None)
        metrics = summary.get("metrics") or {}
        for offset, (attr, _label) in enumerate(value_columns):
            ws.cell(row_idx, identity_count + 1 + offset, metrics.get(attr))
        row_idx += 1

    for ues in context["hierarchy"]:
        for res in ues["res_list"]:
            for row in res["rows"]:
                if not show_empty_rows and row.get("is_empty"):
                    continue
                values = []
                for key, _label, _num in export_columns:
                    val = row.get(key)
                    if key == "station_affiliation" and val == "—":
                        val = None
                    values.append(val)
                values += [row["metrics"].get(attr) for attr, _label in value_columns]
                for col, value in enumerate(values, start=1):
                    ws.cell(row_idx, col, value)
                row_idx += 1
            _write_summary(res.get("summary"), ues.get("ues_name"), res.get("res_name"))
        _write_summary(ues.get("summary"), ues.get("ues_name"))

    for col in range(1, total_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 60 if col == norm_col else 18

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
