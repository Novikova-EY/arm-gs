# -*- coding: utf-8 -*-
"""Выгрузка страниц анализа ПГ/ЗВ (/ecology/analysis/...) в Excel.

Помимо столбцов страницы в файл добавляются ОЭС и РЭС каждой станции.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.ecology.services.gis_tes_fuel_comparison_constants import KEY_LABELS

_STATION_KEYS = frozenset({"NAME", "Stancia"})
_SPAN_KEYS = frozenset({"NAME", "Stancia", "numb1120"})
_KEYS_MARKER = "#keys"
_HEADER_FILL = PatternFill("solid", fgColor="D1E7DD")
_SUMMARY_FILL = PatternFill("solid", fgColor="FFF3CD")
_HEADER_FONT = Font(bold=True)


def _header_labels(context: dict[str, Any]) -> dict[str, str]:
    """Тексты шапки по ключу столбца (последнее непустое значение)."""
    labels: dict[str, str] = {}
    for row in context.get("header_rows") or []:
        for cell in row or []:
            key = str(cell.get("column_key") or "")
            text = str(cell.get("text") or "").strip()
            if key and text:
                labels[key] = text
    return labels


def _column_labels(
    columns: list[dict[str, Any]],
    header_labels: dict[str, str],
) -> list[str]:
    labels: list[str] = []
    for column in columns:
        key = str(column.get("key") or "")
        label = (
            header_labels.get(key)
            or str(column.get("label") or "").strip()
            or KEY_LABELS.get(key, key)
        )
        labels.append(label)
    return labels


def _excel_value(raw: Any) -> Any:
    if raw is None:
        return None
    if isinstance(raw, Decimal):
        return float(raw)
    if isinstance(raw, (int, float, str, bool)):
        return raw
    return str(raw)


def _row_values(
    row: dict[str, Any],
    group: dict[str, Any],
    columns: list[dict[str, Any]],
    display_cells: list[Any],
) -> list[Any]:
    """Значения строки: числа – сырыми, остальные – как на экране."""
    filled = (group.get("filled_by_year") or {}).get(row.get("year")) or []
    values: list[Any] = []
    for index, column in enumerate(columns):
        cell_text = display_cells[index] if index < len(display_cells) else ""
        key = str(column.get("key") or "")
        text = str(cell_text).strip()
        if key in {"Year", "numb1120"} and text.isdigit():
            values.append(int(text))
            continue
        source_index = column.get("source_index")
        if (
            column.get("numeric")
            and isinstance(source_index, int)
            and source_index < len(filled)
            and filled[source_index] is not None
        ):
            values.append(_excel_value(filled[source_index]))
            continue
        values.append(cell_text)
    return values


def _write_summary(
    ws,
    summary: dict[str, Any] | None,
    row_idx: int,
    columns: list[dict[str, Any]],
    ues_name: str,
    res_name: str,
    *,
    show_empty_rows: bool,
) -> int:
    if not summary:
        return row_idx
    for row in summary.get("rows") or []:
        if not show_empty_rows and row.get("is_empty"):
            continue
        cells = list(row.get("cells") or [])
        values = [ues_name, res_name] + [
            cells[index] if index < len(cells) else "" for index in range(len(columns))
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col, value)
            cell.fill = _SUMMARY_FILL
        row_idx += 1
    return row_idx


def export_comparison_table_to_excel(context: dict[str, Any]) -> BytesIO:
    """Собрать XLSX по контексту страницы анализа (иерархия + ОЭС/РЭС)."""
    columns = list(context.get("columns") or [])
    sheet = context.get("sheet") or {}
    page_title = str(context.get("page_title") or sheet.get("page_title") or "Данные")
    hierarchy = list(context.get("hierarchy") or [])
    labels = _column_labels(columns, _header_labels(context))
    if "norm_zv_selected" in context:
        selected = context["norm_zv_selected"]
        parts = [label for attr, label in context.get("norm_zv_options", []) if attr in selected]
        for index, column in enumerate(columns):
            if column.get("key") == "ZV_normir":
                labels[index] += " (" + (" + ".join(parts) or "слагаемые не выбраны") + ")"

    wb = Workbook()
    ws = wb.active
    ws.title = (str(sheet.get("nav_label") or page_title) or "Данные")[:31]
    ws.cell(1, 2, page_title).font = Font(bold=True, size=12)

    header_row = 3
    # Служебная скрытая строка с кодами столбцов: по ней импорт находит столбцы,
    # даже если подписи в шапке повторяются (например, «Расход топлива, т у.т.»).
    ws.cell(2, 1, _KEYS_MARKER)
    for index, column in enumerate(columns, start=3):
        ws.cell(2, index, str(column.get("key") or ""))
    ws.row_dimensions[2].hidden = True

    for col, title in enumerate(["ОЭС", "РЭС", *labels], start=1):
        cell = ws.cell(header_row, col, title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_idx = header_row + 1
    show_empty_rows = bool(context.get("show_empty_rows"))
    for est_block in hierarchy:
        for ues_block in est_block.get("ues_list") or []:
            ues_name = str(ues_block.get("ues_name") or "")
            for res_block in ues_block.get("res_list") or []:
                res_name = str(res_block.get("res_name") or "")
                for group in res_block.get("groups") or []:
                    group_rows = list(group.get("rows") or [])
                    span_values: dict[str, Any] = {}
                    for row in group_rows:
                        if not row.get("is_group_first"):
                            continue
                        cells = list(row.get("cells") or [])
                        for index, column in enumerate(columns):
                            key = column.get("key")
                            if key in _SPAN_KEYS and index < len(cells):
                                span_values[key] = cells[index]
                        break
                    for row in group_rows:
                        if not show_empty_rows and row.get("is_empty"):
                            continue
                        cells = list(row.get("cells") or [])
                        for index, column in enumerate(columns):
                            key = column.get("key")
                            if (
                                key in _SPAN_KEYS
                                and index < len(cells)
                                and not str(cells[index] or "").strip()
                            ):
                                cells[index] = span_values.get(key, "")
                        values = [ues_name, res_name] + _row_values(
                            row, group, columns, cells
                        )
                        for col, value in enumerate(values, start=1):
                            ws.cell(row_idx, col, value)
                        row_idx += 1
                row_idx = _write_summary(
                    ws,
                    res_block.get("summary"),
                    row_idx,
                    columns,
                    ues_name,
                    res_name,
                    show_empty_rows=show_empty_rows,
                )
            row_idx = _write_summary(
                ws,
                ues_block.get("summary"),
                row_idx,
                columns,
                ues_name,
                "",
                show_empty_rows=show_empty_rows,
            )
        row_idx = _write_summary(
            ws,
            est_block.get("summary"),
            row_idx,
            columns,
            "",
            "",
            show_empty_rows=show_empty_rows,
        )

    ws.freeze_panes = f"C{header_row + 1}"
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
    )
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 32
    for index, column in enumerate(columns, start=3):
        key = str(column.get("key") or "")
        if key in _STATION_KEYS:
            width = 30
        elif key == "Year":
            width = 10
        else:
            width = 14
        ws.column_dimensions[get_column_letter(index)].width = width

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
