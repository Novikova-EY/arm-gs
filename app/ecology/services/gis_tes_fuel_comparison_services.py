# -*- coding: utf-8 -*-
"""Чтение макета сопоставления ГИС ТЭС / Топливо и сборка страницы."""
from __future__ import annotations

import functools
import logging
import math
from collections import defaultdict
from copy import deepcopy
from decimal import ROUND_HALF_UP, Decimal, localcontext
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.common.services.get_services.years.years_get_services import get_year_feature_dict
from app.ecology.services.ecology_decimal_helpers import (
    SIGNIFICANT_DIGITS,
    decimal_or_none as _decimal_or_none,
    drop_float_noise,
)
from app.ecology.services.ecology_request_helpers import (
    arg_values as _arg_values,
    parse_truthy_arg,
)
from app.ecology.services.ecology_text_helpers import text_contains
from app.ecology.services.gis_tes_fuel_comparison_constants import (
    COLUMN_FORMULA_KEYS,
    COMPARISON_SHEETS,
    DEFAULT_PER_PAGE,
    DEFAULT_SHEET_SLUG,
    DIFF_HIGHLIGHT_COLUMN_KEYS,
    EMISSION_COEFFICIENT_COLUMN_KEYS,
    FUEL_OVERLAY_BY_KEY,
    FUEL_TCE_HEADER_META,
    FUEL_TCE_OVERLAY_BY_KEY,
    FUEL_TEP_HEADER_KEYS,
    HIDDEN_TABLE_COLUMN_KEYS,
    REPORTED_TABLE_COLUMN_KEYS,
    KCAL_PER_KWH,
    KEY_LABELS,
    NETWORK_XLSX_PATH,
    ANALYSIS_SHEET_SLUGS,
    BEZ_KOT_SHEET_SLUGS,
    PG_ANALYSIS_SHEET_SLUGS,
    PG_DERIVED_COLUMN_KEYS,
    PG_DIFF_HEADER_KEYS,
    PG_ECOLOGY_EDIT_RELATED_KEYS,
    PG_ECOLOGY_EDITABLE_KEYS,
    PG_ECOLOGY_OVERLAY_BY_KEY,
    PG_ECOLOGY_VALUE_FIELDS,
    ZV_ECOLOGY_EDIT_RELATED_KEYS,
    ZV_ECOLOGY_EDITABLE_KEYS,
    ZV_ECOLOGY_OVERLAY_BY_KEY,
    ZV_ECOLOGY_VALUE_FIELDS,
    PG_FUEL_HEADER_GROUP_KEYS,
    PG_FUEL_HEADER_GROUPS,
    SKIP_SHEET_NAME,
    GENERATION_STATION_UNKNOWN_NAME,
    STATION_AFFILIATION_AFTER_KEY,
    STATION_AFFILIATION_COLUMN_KEY,
    STATION_AFFILIATION_LABEL,
    STATION_AFFILIATION_SHEET_SLUGS,
    STATION_AFFILIATION_UNKNOWN,
    STATION_GENERATION_COLUMNS,
    STATION_GENERATION_COLUMN_KEYS,
    STATION_GENERATION_COLUMN_KEY,
    STATION_GENERATION_COLUMN_LABELS,
    STATION_GENERATION_SHEET_SLUGS,
    STATION_RUST_COLUMN_KEY,
    ZV_ANALYSIS_SHEET_SLUGS,
    ZV_COLUMN_FORMULA_KEYS,
    ZV_DISPLAY_COLUMN_SWAPS,
    ZV_EE_SPECIFIC_AFTER_KEY,
    ZV_EE_SPECIFIC_KEY,
    ZV_ENERGY_SPECIFIC_BANNER_TITLE,
    ZV_ENERGY_SPECIFIC_LEAF_LABELS,
    ZV_GASEOUS_PART_KEYS,
    ZV_GASEOUS_TOTAL_KEY,
    ZV_GAS_FUEL_KEYS,
    ZV_MAZUT_FUEL_KEYS,
    ZV_NORM_DEFAULT_PARTS,
    ZV_SOLID_FUEL_KEYS,
    ZV_SOLID_KEY,
    ZV_SPECIFIC_ROUNDING_COLUMN_KEYS,
    ZV_TOTAL_KEY,
    STICKY_COL_COUNT,
    STICKY_COLUMN_KEYS,
    TERRITORY_FILTER_KEYS,
    UNSPECIFIED_LEVEL_NAME,
    VED_FILTER_OPTIONS,
    VED_FILTER_TO_VALUES,
    YEAR_COL_INDEX,
    bundled_xlsx_path,
    get_comparison_sheet,
    list_comparison_sheets_for_group,
    pg_ecology_header_labels,
    zv_emissions_banner_keys,
    zv_energy_specific_banner_keys,
    zv_gaseous_header_group,
    zv_leaf_header_groups,
)
from app.ecology.services.gis_tek_page_services import (
    norm_zv_options,
    parse_norm_zv_parts,
)
from app.ecology.services.operating_tes_page_services import (
    format_cell,
    resolve_rounding_digits,
)

_logger = logging.getLogger(__name__)


def _year_features() -> dict[Any, Any]:
    try:
        return get_year_feature_dict() or {}
    except Exception:
        _logger.exception("ecology comparison: year features unavailable")
        return {}


def resolve_source_path(explicit: str | Path | None = None) -> Path:
    if explicit:
        path = Path(explicit)
        if path.is_file():
            return path
        raise FileNotFoundError(f"Файл сопоставления не найден: {path}")
    bundled = bundled_xlsx_path()
    if bundled.is_file():
        return bundled
    if NETWORK_XLSX_PATH.is_file():
        return NETWORK_XLSX_PATH
    raise FileNotFoundError(
        "Не найден файл «Данные по ЗВ и ПГ 2017-2023 гг.xlsx». "
        f"Ожидается {bundled} или {NETWORK_XLSX_PATH}."
    )


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _as_year(value: Any) -> int | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return None
    try:
        year = int(value)
    except (TypeError, ValueError):
        return None
    if 1990 <= year <= 2100:
        return year
    return None


def _normalize_row(row: Any, width: int | None = None) -> list[Any]:
    values = list(row) if row is not None else []
    if width is None:
        return values
    if len(values) < width:
        values.extend([None] * (width - len(values)))
    return values[:width]


def _used_width(rows: list[list[Any]]) -> int:
    last = 0
    for row in rows:
        for index, value in enumerate(row, start=1):
            if not _is_blank(value):
                last = max(last, index)
    return last


def _looks_like_keys(row: list[Any]) -> bool:
    first = str(row[0] if row else "").strip()
    return first in {"Fed_okr", "BLST"}


def _row_year(row: list[Any], columns: list[dict[str, Any]] | None = None) -> int | None:
    """Год строки: по ключу Year, иначе по фиксированному индексу Excel-макета."""
    if columns:
        year_idx = _column_index(columns, {"Year"})
        if year_idx is not None and year_idx < len(row):
            year = _as_year(row[year_idx])
            if year is not None:
                return year
    if len(row) <= YEAR_COL_INDEX:
        return None
    return _as_year(row[YEAR_COL_INDEX])


def _cell_text(value: Any) -> str:
    if _is_blank(value):
        return ""
    return str(value).strip()


def build_header_cells(grid: list[list[Any]]) -> list[list[dict[str, Any]]]:
    """Собрать ячейки thead с rowspan/colspan по пустым клеткам макета Excel."""
    if not grid:
        return []
    ncols = max(len(row) for row in grid)
    table = [_normalize_row(row, ncols) for row in grid]
    nrows = len(table)
    claimed = [[False] * ncols for _ in range(nrows)]
    result: list[list[dict[str, Any]]] = [[] for _ in range(nrows)]
    for row_i in range(nrows):
        for col_i in range(ncols):
            if claimed[row_i][col_i]:
                continue
            text = _cell_text(table[row_i][col_i])
            if not text:
                claimed[row_i][col_i] = True
                continue
            colspan = 1
            while (
                col_i + colspan < ncols
                and _is_blank(table[row_i][col_i + colspan])
                and not claimed[row_i][col_i + colspan]
            ):
                colspan += 1
            rowspan = 1
            while row_i + rowspan < nrows:
                below = table[row_i + rowspan][col_i : col_i + colspan]
                if any(not _is_blank(item) for item in below):
                    break
                if any(claimed[row_i + rowspan][col_i + offset] for offset in range(colspan)):
                    break
                rowspan += 1
            for delta_r in range(rowspan):
                for delta_c in range(colspan):
                    claimed[row_i + delta_r][col_i + delta_c] = True
            result[row_i].append(
                {
                    "text": text,
                    "rowspan": rowspan,
                    "colspan": colspan,
                    "sticky": col_i < STICKY_COL_COUNT,
                    "col_index": col_i,
                }
            )
    return [row for row in result if row]


def _ensure_leaf_labels(grid: list[list[Any]], keys: list[Any]) -> list[list[Any]]:
    ncols = max((len(row) for row in grid), default=0)
    ncols = max(ncols, len(keys))
    if ncols <= 0:
        return grid
    padded = [_normalize_row(row, ncols) for row in grid]
    if not padded:
        padded = [[None] * ncols]
    for col_i in range(ncols):
        if any(not _is_blank(row[col_i]) for row in padded):
            continue
        key = keys[col_i] if col_i < len(keys) else None
        label = KEY_LABELS.get(str(key), str(key) if key not in (None, "") else "")
        if label:
            padded[0][col_i] = label
    return padded


def _column_is_numeric(rows: list[list[Any]], col_i: int, key: str) -> bool:
    if col_i == YEAR_COL_INDEX or key in {"Year", "numb1120", "obor", "OBOR"}:
        return False
    seen_number = False
    for row in rows:
        if col_i >= len(row):
            continue
        value = row[col_i]
        if _is_blank(value) or isinstance(value, bool):
            continue
        if isinstance(value, str):
            return False
        if isinstance(value, (int, float, Decimal)):
            seen_number = True
            continue
        return False
    return seen_number


def _column_key_at(keys: list[Any], index: int) -> str:
    if index < len(keys) and keys[index] not in (None, ""):
        return str(keys[index]).strip()
    return f"c{index}"


def _build_key_header_row(keys: list[Any], width: int) -> list[dict[str, Any]]:
    """Последняя строка шапки: технические имена из 1-й строки Excel."""
    row: list[dict[str, Any]] = []
    for index in range(width):
        key = _column_key_at(keys, index)
        row.append(
            {
                "text": key,
                "rowspan": 1,
                "colspan": 1,
                "sticky": index < STICKY_COL_COUNT,
                "col_index": index,
                "is_key": True,
                "column_key": key,
            }
        )
    return row


def parse_sheet_rows(raw_rows: list[list[Any]], sheet_meta: dict[str, Any]) -> dict[str, Any]:
    width = _used_width(raw_rows)
    rows = [_normalize_row(row, width) for row in raw_rows]
    data_idx = next((index for index, row in enumerate(rows) if _row_year(row) is not None), None)
    if data_idx is None:
        return {
            "slug": sheet_meta["slug"],
            "sheet_name": sheet_meta["sheet_name"],
            "page_title": sheet_meta["page_title"],
            "header_rows": [],
            "columns": [],
            "rows": [],
            "years": [],
        }
    header_rows = rows[:data_idx]
    data_rows = [row for row in rows[data_idx:] if any(not _is_blank(item) for item in row)]
    keys: list[Any] = []
    display_headers = header_rows
    if header_rows and _looks_like_keys(header_rows[0]):
        keys = list(header_rows[0])
        display_headers = header_rows[1:]
    if not keys:
        keys = [f"c{index}" for index in range(width)]
    display_headers = _ensure_leaf_labels(display_headers, keys)
    header_cells = build_header_cells(display_headers)
    numeric_flags = [
        _column_is_numeric(data_rows, index, _column_key_at(keys, index))
        for index in range(width)
    ]
    years = sorted(
        {
            year
            for row in data_rows
            if (year := _row_year(row)) is not None
        }
    )
    columns = [
            {
                "key": _column_key_at(keys, index),
                "numeric": numeric_flags[index],
                "sticky": index < STICKY_COL_COUNT,
            }
            for index in range(width)
        ]
    _relabel_headers(header_cells, columns)
    if keys:
        header_cells.append(_build_key_header_row(keys, width))
    return {
        "slug": sheet_meta["slug"],
        "sheet_name": sheet_meta["sheet_name"],
        "page_title": sheet_meta["page_title"],
        "header_rows": header_cells,
        "columns": columns,
        "rows": data_rows,
        "years": years,
    }


def _relabel_headers(header_rows: list[list[dict[str, Any]]], columns: list[dict[str, Any]]) -> None:
    wanted = {
        index: KEY_LABELS[col["key"]]
        for index, col in enumerate(columns)
        if col["key"] == "numb1120" and col["key"] in KEY_LABELS
    }
    if not wanted:
        return
    for row in header_rows:
        for cell in row:
            if cell.get("is_key"):
                continue
            col_i = cell.get("col_index")
            if col_i in wanted and cell.get("colspan", 1) == 1:
                cell["text"] = wanted[col_i]


def _fuel_tep_header_meta() -> dict[str, dict[str, str]]:
    """Подписи Nust/E/Q как на stations_equipment_group_fuel_params."""
    from app.fuel.models.fue_equipment_group_fuel_param_model import EquipmentGroupFuelParam
    from app.fuel.services.fuel_header_access_codes import fuel_header_code

    labels = {
        "nust": EquipmentGroupFuelParam.NUST_COLUMN_LABEL,
        "e": EquipmentGroupFuelParam.E_COLUMN_LABEL,
        "q": EquipmentGroupFuelParam.Q_COLUMN_LABEL,
    }
    out: dict[str, dict[str, str]] = {}
    for excel_key, attr in FUEL_OVERLAY_BY_KEY.items():
        if excel_key not in FUEL_TEP_HEADER_KEYS:
            continue
        out[excel_key] = {
            "label": labels[attr],
            "access_code": fuel_header_code(attr),
        }
    return out


def apply_fuel_tep_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
) -> None:
    """Серый фон и подписи ТЭП (nust/E/Q) в шапке видимой таблицы."""
    meta_by_key = _fuel_tep_header_meta()
    for column in columns:
        key = column.get("key")
        meta = meta_by_key.get(key) if isinstance(key, str) else None
        column["fuel_tep_highlight"] = meta is not None
    for row in header_rows:
        for cell in row:
            key = str(cell.get("column_key") or "")
            meta = meta_by_key.get(key)
            if not meta:
                cell["fuel_tep_highlight"] = False
                continue
            cell["fuel_tep_highlight"] = True
            if cell.get("is_key"):
                cell["text"] = meta["access_code"]
                cell.pop("access_code", None)
            else:
                cell["text"] = meta["label"]
                cell["access_code"] = meta["access_code"]


def apply_fuel_tce_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
) -> None:
    """Подписи и серый фон столбцов расхода топлива из ТЭП (GAZ/MAZUT/UGOL/PROCH)."""
    highlight_keys = frozenset(FUEL_TCE_OVERLAY_BY_KEY) | frozenset(FUEL_TCE_HEADER_META)
    if not highlight_keys:
        return
    for column in columns:
        key = column.get("key")
        meta = FUEL_TCE_HEADER_META.get(key) if isinstance(key, str) else None
        if meta:
            column["label"] = meta["label"]
            column["access_code"] = meta["access_code"]
        if isinstance(key, str) and key in highlight_keys:
            column["fuel_tep_highlight"] = True
    for row in header_rows:
        for cell in row:
            key = str(cell.get("column_key") or "")
            if key not in highlight_keys:
                continue
            if int(cell.get("colspan") or 1) != 1:
                continue
            cell["fuel_tep_highlight"] = True
            meta = FUEL_TCE_HEADER_META.get(key)
            if not meta:
                continue
            if cell.get("is_key"):
                cell["text"] = meta["access_code"]
                cell.pop("access_code", None)
            else:
                cell["text"] = meta["label"]
                cell["access_code"] = meta["access_code"]


def apply_pg_ecology_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
) -> None:
    """Подписи 2-й строки шапки ПГ без привязки к виду топлива."""
    labels = pg_ecology_header_labels()
    if not labels:
        return
    for column in columns:
        key = column.get("key")
        if isinstance(key, str) and key in labels:
            column["label"] = labels[key]
    for row in header_rows:
        for cell in row:
            if cell.get("is_key"):
                continue
            if cell.get("colspan", 1) != 1:
                continue
            key = str(cell.get("column_key") or "")
            label = labels.get(key)
            if label:
                cell["text"] = label


def _header_caption_rows(
    header_rows: list[list[dict[str, Any]]],
) -> tuple[list[list[dict[str, Any]]], list[dict[str, Any]] | None]:
    if header_rows and header_rows[-1] and header_rows[-1][0].get("is_key"):
        return header_rows[:-1], header_rows[-1]
    return header_rows, None


def drop_key_header_row(header_rows: list[list[dict[str, Any]]]) -> None:
    """Убирает техническую строку ключей Excel (GAZ, d_gaz, …) из шапки страницы."""
    caption_rows, key_row = _header_caption_rows(header_rows)
    if key_row is None:
        return
    caption_by_key: dict[str, dict[str, Any]] = {}
    for row in caption_rows:
        for cell in row:
            key = str(cell.get("column_key") or "")
            if key:
                caption_by_key.setdefault(key, cell)
    for cell in key_row:
        key = str(cell.get("column_key") or "")
        target = caption_by_key.get(key)
        if not target:
            continue
        tooltip = cell.get("formula_tooltip")
        if tooltip and not target.get("formula_tooltip"):
            target["formula_tooltip"] = tooltip
    header_rows[:] = caption_rows


def _top_cell_covers_keys(
    cell: dict[str, Any],
    columns: list[dict[str, Any]],
    index_by_key: dict[str, int],
) -> set[str]:
    key = str(cell.get("column_key") or "")
    if key:
        return {key} if key in index_by_key else set()
    col_i = cell.get("col_index")
    span = int(cell.get("colspan") or 1)
    if col_i is None:
        return set()
    covered: set[str] = set()
    # col_index после prepare_visible: sticky – sticky_index, иначе исходный Excel-индекс.
    for column in columns:
        source = column.get("source_index")
        sticky = column.get("sticky_index")
        pos = sticky if column.get("sticky") and sticky is not None else source
        if pos is None:
            continue
        if int(col_i) <= int(pos) < int(col_i) + span:
            covered.add(str(column["key"]))
    return covered


def _pop_header_cell(row: list[dict[str, Any]] | None, key: str) -> dict[str, Any] | None:
    if not row:
        return None
    for index, cell in enumerate(row):
        if str(cell.get("column_key") or "") == key:
            return row.pop(index)
    return None


def apply_pg_fuel_group_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
) -> None:
    """Заменяет строку «Отчет»/«Макет Экология» на группы Газ / Мазут / Твердое топливо."""
    caption_rows, _key_row = _header_caption_rows(header_rows)
    if not caption_rows:
        return
    top = caption_rows[0]
    mid = caption_rows[1] if len(caption_rows) > 1 else None
    caption_span = len(caption_rows)
    index_by_key = {
        str(col["key"]): index
        for index, col in enumerate(columns)
        if col.get("key") not in (None, "")
    }
    group_start = {
        keys[0]: (title, keys)
        for title, keys in PG_FUEL_HEADER_GROUPS
        if keys and keys[0] in index_by_key
    }

    # Подписи d_* из верхней строки переносим на среднюю (под новую группу).
    if mid is not None:
        mid_keys = {
            str(cell.get("column_key") or "")
            for cell in mid
            if cell.get("column_key")
        }
        for cell in list(top):
            key = str(cell.get("column_key") or "")
            if key not in PG_DIFF_HEADER_KEYS or key in mid_keys:
                continue
            moved = dict(cell)
            moved["rowspan"] = 1
            moved["colspan"] = 1
            insert_at = 0
            target_index = index_by_key.get(key, 0)
            for pos, mid_cell in enumerate(mid):
                mid_key = str(mid_cell.get("column_key") or "")
                mid_index = index_by_key.get(mid_key, -1)
                if mid_index < target_index:
                    insert_at = pos + 1
            mid.insert(insert_at, moved)
            mid_keys.add(key)

    kept: list[dict[str, Any]] = []
    for cell in top:
        covered = _top_cell_covers_keys(cell, columns, index_by_key)
        if covered & (PG_FUEL_HEADER_GROUP_KEYS | PG_DIFF_HEADER_KEYS):
            continue
        text = str(cell.get("text") or "").strip().casefold()
        if text in {"отчет", "макет экология"}:
            continue
        kept.append(cell)

    rebuilt: list[dict[str, Any]] = []
    emitted_groups: set[str] = set()
    kept_by_first_key: dict[str, dict[str, Any]] = {}
    for cell in kept:
        covered = _top_cell_covers_keys(cell, columns, index_by_key)
        if len(covered) == 1:
            kept_by_first_key[next(iter(covered))] = cell

    for column in columns:
        key = str(column.get("key") or "")
        if key in FUEL_TEP_HEADER_KEYS:
            if key in emitted_groups:
                continue
            cell = _pop_header_cell(mid, key) or kept_by_first_key.get(key)
            if cell is None:
                continue
            promoted = dict(cell)
            promoted["rowspan"] = caption_span
            promoted["colspan"] = 1
            rebuilt.append(promoted)
            emitted_groups.add(key)
            continue
        if key in PG_FUEL_HEADER_GROUP_KEYS:
            spec = group_start.get(key)
            if not spec or key in emitted_groups:
                continue
            title, keys = spec
            present = [item for item in keys if item in index_by_key]
            if not present:
                continue
            rebuilt.append(
                {
                    "text": title,
                    "rowspan": 1,
                    "colspan": len(present),
                    "sticky": False,
                    "col_index": column.get("source_index"),
                    "column_key": present[0],
                    "diff_highlight": False,
                    "fuel_tep_highlight": False,
                }
            )
            emitted_groups.update(present)
            continue
        cell = kept_by_first_key.get(key)
        if cell is not None:
            rebuilt.append(cell)

    caption_rows[0][:] = rebuilt
    _remove_pg_diff_subrow(header_rows)


def _cell_sort_index(
    cell: dict[str, Any],
    columns: list[dict[str, Any]],
    index_by_key: dict[str, int],
) -> int:
    covered = _top_cell_covers_keys(cell, columns, index_by_key)
    if covered:
        return min(index_by_key[key] for key in covered if key in index_by_key)
    col_i = cell.get("col_index")
    return int(col_i) if col_i is not None else 10**9


def _insert_index_for_column(
    row: list[dict[str, Any]],
    target_index: int,
    columns: list[dict[str, Any]],
    index_by_key: dict[str, int],
) -> int:
    insert_at = 0
    for pos, cell in enumerate(row):
        if _cell_sort_index(cell, columns, index_by_key) < target_index:
            insert_at = pos + 1
        else:
            break
    return insert_at


def _cells_without_keys(
    cell: dict[str, Any],
    columns: list[dict[str, Any]],
    index_by_key: dict[str, int],
    drop_keys: set[str],
) -> list[dict[str, Any]]:
    covered = _top_cell_covers_keys(cell, columns, index_by_key)
    if not covered:
        return [cell]
    if not (covered & drop_keys):
        return [cell]
    leftover = [
        str(col["key"])
        for col in columns
        if str(col.get("key") or "") in (covered - drop_keys)
    ]
    if not leftover:
        return []
    rewritten = dict(cell)
    rewritten["colspan"] = len(leftover)
    first = leftover[0]
    rewritten["column_key"] = first if len(leftover) == 1 else None
    column = columns[index_by_key[first]]
    rewritten["col_index"] = (
        column.get("sticky_index")
        if column.get("sticky") and column.get("sticky_index") is not None
        else column.get("source_index")
    )
    return [rewritten]


def apply_zv_column_group_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
    title: str,
    group_keys: tuple[str, ...],
    leaf_labels: dict[str, str],
) -> None:
    """Двухрядная шапка: объединённый заголовок группы и короткие подписи столбцов."""
    caption_rows, _key_row = _header_caption_rows(header_rows)
    if not caption_rows:
        return
    index_by_key = {
        str(col["key"]): index
        for index, col in enumerate(columns)
        if col.get("key") not in (None, "")
    }
    present = [key for key in group_keys if key in index_by_key]
    if not present:
        return
    group_set = set(present)
    added_leaf_row = False
    if len(caption_rows) == 1:
        header_rows.insert(1, [])
        added_leaf_row = True
        caption_rows, _key_row = _header_caption_rows(header_rows)
    caption_span = len(caption_rows)
    for row in caption_rows:
        rewritten: list[dict[str, Any]] = []
        for cell in row:
            rewritten.extend(_cells_without_keys(cell, columns, index_by_key, group_set))
        row[:] = rewritten
    top = caption_rows[0]
    leaf = caption_rows[-1]
    if added_leaf_row:
        for cell in top:
            cell["rowspan"] = int(cell.get("rowspan") or 1) + 1
    first_index = min(index_by_key[key] for key in present)
    first_col = columns[index_by_key[present[0]]]
    is_fuel_tce_group = all(
        key in FUEL_TCE_OVERLAY_BY_KEY or key in FUEL_TCE_HEADER_META
        for key in present
    )
    group_cell = {
        "text": title,
        "rowspan": max(caption_span - 1, 1),
        "colspan": len(present),
        "sticky": False,
        "col_index": first_col.get("source_index"),
        "column_key": None,
        "diff_highlight": False,
        "fuel_tep_highlight": is_fuel_tce_group,
    }
    top.insert(_insert_index_for_column(top, first_index, columns, index_by_key), group_cell)
    leaf_cells = []
    for key in present:
        column = columns[index_by_key[key]]
        label = leaf_labels.get(key) or KEY_LABELS.get(key) or key
        column["label"] = label
        meta = FUEL_TCE_HEADER_META.get(key)
        is_fuel_tce = key in FUEL_TCE_OVERLAY_BY_KEY or key in FUEL_TCE_HEADER_META
        if is_fuel_tce:
            column["fuel_tep_highlight"] = True
            if meta:
                column["access_code"] = meta["access_code"]
        leaf_cell: dict[str, Any] = {
            "text": label,
            "rowspan": 1,
            "colspan": 1,
            "sticky": False,
            "col_index": column.get("source_index"),
            "column_key": key,
            "diff_highlight": False,
            "fuel_tep_highlight": is_fuel_tce,
        }
        if meta:
            leaf_cell["access_code"] = meta["access_code"]
        leaf_cells.append(leaf_cell)
    insert_at = _insert_index_for_column(leaf, first_index, columns, index_by_key)
    for offset, cell in enumerate(leaf_cells):
        leaf.insert(insert_at + offset, cell)


def apply_zv_gaseous_group_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
) -> None:
    """Двухрядная шапка «Газообразные и жидкие вещества» как на gis_tek/pollutants."""
    title, group_keys, leaf_labels = zv_gaseous_header_group()
    apply_zv_column_group_headers(columns, header_rows, title, group_keys, leaf_labels)


def apply_zv_leaf_group_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
) -> None:
    """Группы шапки ЗВ: газообразные, удельные и нормируемые по ингредиентам."""
    for title, group_keys, leaf_labels in zv_leaf_header_groups():
        apply_zv_column_group_headers(columns, header_rows, title, group_keys, leaf_labels)


def apply_zv_top_banner_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
    banners: list[tuple[str, tuple[str, ...]]],
) -> None:
    """Одна верхняя строка шапки с одной или несколькими объединёнными группами."""
    caption_rows, _key_row = _header_caption_rows(header_rows)
    if not caption_rows:
        return
    index_by_key = {
        str(col["key"]): index
        for index, col in enumerate(columns)
        if col.get("key") not in (None, "")
    }
    resolved: list[tuple[str, list[str]]] = []
    banner_set: set[str] = set()
    for title, keys in banners:
        present = [key for key in keys if key in index_by_key]
        if not present:
            continue
        resolved.append((title, present))
        banner_set.update(present)
    if not resolved:
        return
    header_rows.insert(0, [])
    caption_rows, _key_row = _header_caption_rows(header_rows)
    banner_row = caption_rows[0]
    group_row = caption_rows[1]
    moved: list[dict[str, Any]] = []
    kept: list[dict[str, Any]] = []
    for cell in group_row:
        covered = _top_cell_covers_keys(cell, columns, index_by_key)
        if covered and not (covered & banner_set):
            cell["rowspan"] = int(cell.get("rowspan") or 1) + 1
            moved.append(cell)
            continue
        if not covered:
            cell["rowspan"] = int(cell.get("rowspan") or 1) + 1
            moved.append(cell)
            continue
        kept.append(cell)
    group_row[:] = kept
    banner_row[:] = moved
    for title, present in resolved:
        first_index = min(index_by_key[key] for key in present)
        first_col = columns[index_by_key[present[0]]]
        banner_cell = {
            "text": title,
            "rowspan": 1,
            "colspan": len(present),
            "sticky": False,
            "col_index": first_col.get("source_index"),
            "column_key": None,
            "diff_highlight": False,
            "fuel_tep_highlight": False,
        }
        banner_row.insert(
            _insert_index_for_column(banner_row, first_index, columns, index_by_key),
            banner_cell,
        )


def _relabel_zv_header_leaves(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
    labels: dict[str, str],
) -> None:
    for column in columns:
        key = str(column.get("key") or "")
        if key in labels:
            column["label"] = labels[key]
    caption_rows, _key_row = _header_caption_rows(header_rows)
    for row in caption_rows:
        for cell in row:
            key = str(cell.get("column_key") or "")
            if key in labels:
                cell["text"] = labels[key]


def _zv_total_leaf_label() -> str:
    from app.ecology.services.gis_tek_constants import ZV_VALUE_COLUMNS

    return next(label for attr, label in ZV_VALUE_COLUMNS if attr == "total")


def apply_zv_emissions_banner_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
) -> None:
    """Верхняя строка «Всего выброшено в атмосферу ЗВ за год, т» над всего / твёрдыми / газообразными."""
    from app.ecology.services.gis_tek_constants import ZV_EMISSIONS_BANNER_TITLE

    apply_zv_top_banner_headers(
        columns,
        header_rows,
        [(ZV_EMISSIONS_BANNER_TITLE, zv_emissions_banner_keys())],
    )
    _relabel_zv_header_leaves(columns, header_rows, {ZV_TOTAL_KEY: _zv_total_leaf_label()})


def apply_zv_analysis_banner_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
) -> None:
    """Верхняя шапка ЗВ: выбросы за год и удельные по электрической энергии."""
    from app.ecology.services.gis_tek_constants import ZV_EMISSIONS_BANNER_TITLE

    apply_zv_top_banner_headers(
        columns,
        header_rows,
        [
            (ZV_EMISSIONS_BANNER_TITLE, zv_emissions_banner_keys()),
            (ZV_ENERGY_SPECIFIC_BANNER_TITLE, zv_energy_specific_banner_keys()),
        ],
    )
    _relabel_zv_header_leaves(
        columns,
        header_rows,
        {ZV_TOTAL_KEY: _zv_total_leaf_label(), **ZV_ENERGY_SPECIFIC_LEAF_LABELS},
    )


def _remove_pg_diff_subrow(header_rows: list[list[dict[str, Any]]]) -> None:
    """Убирает подстроку шапки «6тп-Экол» и уменьшает rowspan выше."""
    caption_rows, _key_row = _header_caption_rows(header_rows)
    if not caption_rows:
        return
    to_remove: list[list[dict[str, Any]]] = []
    for row in caption_rows:
        texts = {
            str(cell.get("text") or "").strip().casefold().replace(" ", "")
            for cell in row
        }
        texts.discard("")
        if texts and texts <= {"6тп-экол"}:
            to_remove.append(row)
    if not to_remove:
        return
    for row in to_remove:
        header_rows.remove(row)
    removed = len(to_remove)
    for row in header_rows:
        if row and row[0].get("is_key"):
            continue
        for cell in row:
            rowspan = int(cell.get("rowspan") or 1)
            if rowspan > removed:
                cell["rowspan"] = rowspan - removed


def _column_formula_texts() -> dict[str, str]:
    from app.ecology.services.formula_text.operating_tes_formula_text_services import (
        ot_formula_text,
    )

    out: dict[str, str] = {}
    formula_keys = {**COLUMN_FORMULA_KEYS, **ZV_COLUMN_FORMULA_KEYS}
    for column_key, formula_key in formula_keys.items():
        text = str(ot_formula_text(formula_key) or "").strip()
        if text:
            out[column_key] = text
    return out


def _attach_column_formula_tooltips(
    header_rows: list[list[dict[str, Any]]],
    formula_by_column_key: dict[str, str],
) -> None:
    """Вешает текст формулы на первую подпись столбца; иначе – на строку ключей."""
    if not formula_by_column_key:
        return
    attached: set[str] = set()
    for row in header_rows:
        for cell in row:
            if cell.get("is_key"):
                continue
            key = str(cell.get("column_key") or "")
            text = formula_by_column_key.get(key)
            if not text or key in attached:
                continue
            cell["formula_tooltip"] = text
            attached.add(key)
    for row in header_rows:
        for cell in row:
            if not cell.get("is_key"):
                continue
            key = str(cell.get("column_key") or "")
            if key in attached:
                continue
            text = formula_by_column_key.get(key)
            if not text:
                continue
            cell["formula_tooltip"] = text
            attached.add(key)


def _column_index(columns: list[dict[str, Any]], keys: set[str]) -> int | None:
    for index, column in enumerate(columns):
        if column["key"] in keys:
            return index
    return None


def _row_numb(row: list[Any], columns: list[dict[str, Any]]) -> int | None:
    index = _column_index(columns, {"numb1120"})
    if index is None or index >= len(row):
        return None
    value = row[index]
    if _is_blank(value):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def apply_fuel_indicators(
    raw_row: list[Any],
    columns: list[dict[str, Any]],
    overlay: dict[tuple[int, int], dict[str, Any]],
) -> list[Any]:
    """Подставляет ФО/ОЭС/субъект/станцию/Nust/E/Q и расходы топлива (т у.т.) из ТЭП."""
    numb = _row_numb(raw_row, columns)
    year = _row_year(raw_row, columns)
    if numb is None or year is None:
        return raw_row
    fuel = overlay.get((numb, year))
    if not fuel:
        return raw_row
    out = list(raw_row)
    if len(out) < len(columns):
        out.extend([None] * (len(columns) - len(out)))
    for index, column in enumerate(columns):
        key = column["key"]
        field = FUEL_OVERLAY_BY_KEY.get(key) or FUEL_TCE_OVERLAY_BY_KEY.get(key)
        if not field:
            continue
        value = fuel.get(field)
        if value is not None:
            out[index] = value
    return out


def apply_pg_ecology_overlay(
    raw_row: list[Any],
    columns: list[dict[str, Any]],
    overlay: dict[tuple[int, int], dict[str, Any]],
) -> list[Any]:
    """Подставляет GAZ_ek/CO2_*/MAZUT_ek/TvTopl_ek из параметров или отчетных ПГ."""
    numb = _row_numb(raw_row, columns)
    year = _row_year(raw_row, columns)
    if numb is None or year is None:
        return raw_row
    values = overlay.get((numb, year))
    if not values:
        return raw_row
    out = list(raw_row)
    if len(out) < len(columns):
        out.extend([None] * (len(columns) - len(out)))
    for index, column in enumerate(columns):
        field = PG_ECOLOGY_OVERLAY_BY_KEY.get(column["key"])
        if not field or field not in values:
            continue
        out[index] = values[field]
    return out


def apply_zv_ecology_overlay(
    raw_row: list[Any],
    columns: list[dict[str, Any]],
    overlay: dict[tuple[int, int], dict[str, Any]],
) -> list[Any]:
    """Подставляет выбросы ЗВ из отчетных данных ГИС ТЭК."""
    numb = _row_numb(raw_row, columns)
    year = _row_year(raw_row, columns)
    if numb is None or year is None:
        return raw_row
    values = overlay.get((numb, year))
    if not values:
        return raw_row
    out = list(raw_row)
    if len(out) < len(columns):
        out.extend([None] * (len(columns) - len(out)))
    for index, column in enumerate(columns):
        field = ZV_ECOLOGY_OVERLAY_BY_KEY.get(column["key"])
        if not field or field not in values:
            continue
        out[index] = values[field]
    return out


def _pg_values_from_row(row: Any, *, include_none: bool = False) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for field in PG_ECOLOGY_VALUE_FIELDS:
        value = getattr(row, field, None)
        if include_none or value is not None:
            result[field] = value
    return result


def load_pg_ecology_overlay_by_numb_year(
    sheet_slug: str,
    numbs: set[int],
    years: list[int],
) -> dict[tuple[int, int], dict[str, Any]]:
    """Параметры анализа ПГ (если есть) иначе отчетные ПГ: ключ (numb1120, год)."""
    slug = str(sheet_slug or "").strip()
    if slug not in PG_ANALYSIS_SHEET_SLUGS or not numbs or not years:
        return {}
    try:
        from flask import has_app_context

        from app.common.services.database_version_filter import get_current_db_version_id
        from app.ecology.models.pg_param_calc_model import PGParamCalc
        from app.ecology.models.pg_report_gis_tek_model import PGReportGisTek

        if not has_app_context():
            return {}
        version_id = get_current_db_version_id()
        numb_list = list(numbs)
        year_list = list(years)

        report_q = PGReportGisTek.query.filter(
            PGReportGisTek.station_numb.in_(numb_list),
            PGReportGisTek.year_number.in_(year_list),
        )
        param_q = PGParamCalc.query.filter(
            PGParamCalc.sheet_slug == slug,
            PGParamCalc.station_numb.in_(numb_list),
            PGParamCalc.year_number.in_(year_list),
        )
        if version_id is not None:
            report_q = report_q.filter(
                PGReportGisTek.database_version_id == version_id
            )
            param_q = param_q.filter(PGParamCalc.database_version_id == version_id)

        result: dict[tuple[int, int], dict[str, Any]] = {}
        for report in report_q.all():
            # Отчет: только заполненные поля (пустые не затирают Excel).
            result[(int(report.station_numb), int(report.year_number))] = _pg_values_from_row(
                report,
                include_none=False,
            )
        for param in param_q.all():
            # Параметры страницы перекрывают отчет/Excel целиком (в т.ч. None).
            result[(int(param.station_numb), int(param.year_number))] = _pg_values_from_row(
                param,
                include_none=True,
            )
        return result
    except Exception:
        _logger.exception(
            "ecology comparison: PG overlay load failed (sheet=%s)", sheet_slug
        )
        return {}


def _zv_values_from_row(row: Any, *, include_none: bool = False) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for field in ZV_ECOLOGY_VALUE_FIELDS:
        value = getattr(row, field, None)
        if include_none or value is not None:
            result[field] = value
    return result


def load_zv_ecology_overlay_by_numb_year(
    sheet_slug: str,
    numbs: set[int],
    years: list[int],
) -> dict[tuple[int, int], dict[str, Any]]:
    """Параметры анализа ЗВ (если есть) иначе отчетные ЗВ: ключ (numb1120, год)."""
    slug = str(sheet_slug or "").strip()
    if slug not in ZV_ANALYSIS_SHEET_SLUGS or not numbs or not years:
        return {}
    try:
        from flask import has_app_context

        from app.common.services.database_version_filter import get_current_db_version_id
        from app.ecology.models.zv_param_calc_model import ZVParamCalc
        from app.ecology.models.zv_report_gis_tek_model import ZVReportGisTek

        if not has_app_context():
            return {}
        version_id = get_current_db_version_id()
        numb_list = list(numbs)
        year_list = list(years)

        report_q = ZVReportGisTek.query.filter(
            ZVReportGisTek.station_numb.in_(numb_list),
            ZVReportGisTek.year_number.in_(year_list),
        )
        param_q = ZVParamCalc.query.filter(
            ZVParamCalc.sheet_slug == slug,
            ZVParamCalc.station_numb.in_(numb_list),
            ZVParamCalc.year_number.in_(year_list),
        )
        if version_id is not None:
            report_q = report_q.filter(
                ZVReportGisTek.database_version_id == version_id
            )
            param_q = param_q.filter(ZVParamCalc.database_version_id == version_id)

        result: dict[tuple[int, int], dict[str, Any]] = {}
        for report in report_q.all():
            # Отчет: только заполненные поля (пустые не затирают Excel).
            result[(int(report.station_numb), int(report.year_number))] = _zv_values_from_row(
                report,
                include_none=False,
            )
        for param in param_q.all():
            # Параметры страницы перекрывают отчет/Excel целиком (в т.ч. None).
            result[(int(param.station_numb), int(param.year_number))] = _zv_values_from_row(
                param,
                include_none=True,
            )
        return result
    except Exception:
        _logger.exception(
            "ecology comparison: ZV overlay load failed (sheet=%s)", sheet_slug
        )
        return {}


# Строки для пересчёта ЗВ после правки: значения ЗВ, топливо/ТЭП и расчётные итоги.
ZV_RELATED_CALC_KEYS: tuple[str, ...] = (
    "TVERD_V",
    "SO2_gaz",
    "CO_gaz",
    "NOx_gaz",
    "UGLEVOD",
    "LOS",
    "PROCH_ekol",
    "GAZ_6tp",
    "MAZUT_6tp",
    "TVERD_TOPL_6tp",
    "W_en",
    "Teplo",
) + tuple(ZV_COLUMN_FORMULA_KEYS)


def _zv_ecology_values_by_key(
    zv_values: dict[str, Any] | None,
    fuel: dict[str, Any] | None,
    norm_zv_parts: list[str] | tuple[str, ...] = ZV_NORM_DEFAULT_PARTS,
) -> dict[str, Any]:
    """Свод значений по ключам столбцов для пересчёта итогов и удельных после правки."""
    by_key: dict[str, Any] = {
        excel_key: (zv_values or {}).get(field)
        for excel_key, field in ZV_ECOLOGY_OVERLAY_BY_KEY.items()
    }
    by_key["GAZ_6tp"] = (fuel or {}).get("gaz_all_tce")
    by_key["MAZUT_6tp"] = (fuel or {}).get("mazut_tce")
    by_key["TVERD_TOPL_6tp"] = (fuel or {}).get("solid_tce")
    by_key["W_en"] = (fuel or {}).get("e")
    by_key["Teplo"] = (fuel or {}).get("q")
    columns = [
        {"key": key, "source_index": index}
        for index, key in enumerate(ZV_RELATED_CALC_KEYS)
    ]
    filled = _compute_derived_columns(
        [by_key.get(key) for key in ZV_RELATED_CALC_KEYS],
        columns,
        norm_zv_parts,
    )
    return dict(zip(ZV_RELATED_CALC_KEYS, filled))


def _load_hidden_boiler_numbs(sheet_slug: str) -> set[int]:
    """
    NUMB1120 котельных, которые на этой вкладке не показываем
    (см. hidden_boiler_numbs_for_sheet).

    Без контекста приложения (юнит-тесты) отбор пропускается.
    """
    from flask import has_app_context

    from app.common.services.database_version_filter import get_current_db_version_id
    from app.ecology.services.gis_tes_pg_from_fuel_services import (
        hidden_boiler_numbs_for_sheet,
    )

    if not has_app_context():
        return set()
    return hidden_boiler_numbs_for_sheet(sheet_slug, get_current_db_version_id())


def _drop_boiler_groups(
    groups: list[dict[str, Any]],
    sheet_slug: str,
) -> list[dict[str, Any]]:
    """Убирает строки котельных, которых на этой вкладке быть не должно."""
    hidden_boiler_numbs = _load_hidden_boiler_numbs(sheet_slug)
    if not hidden_boiler_numbs:
        return groups
    return [
        group for group in groups if group.get("numb") not in hidden_boiler_numbs
    ]


def _pg_ecology_values_by_key(
    pg_values: dict[str, Any] | None,
    fuel: dict[str, Any] | None,
) -> dict[str, Any]:
    """Свод значений по ключам столбцов для пересчёта d_* / K_* после правки."""
    by_key: dict[str, Any] = {}
    for excel_key, field in PG_ECOLOGY_OVERLAY_BY_KEY.items():
        by_key[excel_key] = (pg_values or {}).get(field)
    for excel_key, field in FUEL_TCE_OVERLAY_BY_KEY.items():
        by_key[excel_key] = (fuel or {}).get(field)
    by_key["W_en"] = (fuel or {}).get("e")
    by_key["Teplo"] = (fuel or {}).get("q")
    by_key["d_mazut"] = _safe_subtract(by_key.get("MAZUT_6tp"), by_key.get("MAZUT_ek"))
    by_key["d_gaz"] = _safe_subtract(by_key.get("GAZ_ESTEST"), by_key.get("GAZ_ek"))
    by_key["d_tv"] = _safe_subtract(by_key.get("TvTopl_6tp"), by_key.get("TvTopl_ek"))
    by_key["K_gaz_ek"] = _safe_divide(by_key.get("CO2_gaz"), by_key.get("GAZ_ek"))
    by_key["K_mazut_ek"] = _safe_divide(by_key.get("CO2_mazut"), by_key.get("MAZUT_ek"))
    by_key["K_ugol_ek"] = _safe_divide(by_key.get("CO2_ugol"), by_key.get("TvTopl_ek"))
    by_key["CO2"] = _safe_sum(
        by_key.get("CO2_gaz"),
        by_key.get("CO2_mazut"),
        by_key.get("CO2_ugol"),
    )
    fuel_ek_total = _safe_sum(
        by_key.get("GAZ_ek"),
        by_key.get("MAZUT_ek"),
        by_key.get("TvTopl_ek"),
    )
    co2_kg = None if _is_blank(by_key.get("CO2")) else float(by_key["CO2"]) * 1000.0
    by_key["UD_CO2_ek"] = _safe_divide(co2_kg, fuel_ek_total)
    by_key["UD_CO2_ener"] = _safe_divide(
        by_key.get("CO2"),
        _brought_energy_mln_kwh(by_key.get("W_en"), by_key.get("Teplo")),
    )
    return by_key


def format_related_pg_ecology_cells(
    sheet_slug: str,
    station_numb: int,
    year_number: int,
    column_key: str,
    rounding_digits: int | str | None = None,
    emission_rounding_digits: int | str | None = None,
) -> dict[str, dict[str, Any]]:
    """Форматирует отредактированную ячейку и все расчётные d_* / K_* / CO2 / UD_*."""
    key = str(column_key or "").strip()
    related = PG_ECOLOGY_EDIT_RELATED_KEYS.get(key)
    if not related:
        related = (key,) + tuple(PG_DERIVED_COLUMN_KEYS) if key else tuple(PG_DERIVED_COLUMN_KEYS)
    # Уникальный порядок: правка, затем расчётные.
    ordered: list[str] = []
    for item in related:
        if item and item not in ordered:
            ordered.append(item)
    digits = resolve_rounding_digits(rounding_digits)
    emission_digits = resolve_optional_rounding_digits(emission_rounding_digits)
    numb = int(station_numb)
    year = int(year_number)
    pg = load_pg_ecology_overlay_by_numb_year(sheet_slug, {numb}, [year]).get(
        (numb, year)
    )
    fuel = load_fuel_indicators_by_numb_year({numb}, [year]).get((numb, year))
    by_key = _pg_ecology_values_by_key(pg, fuel)
    result: dict[str, dict[str, Any]] = {}
    for rel_key in ordered:
        text = _format_value(
            by_key.get(rel_key),
            True,
            column_rounding_digits({"key": rel_key}, digits, emission_digits),
        )
        result[rel_key] = {
            "text": text,
            "title": _format_full_value(by_key.get(rel_key)),
            "nonzero": rel_key in DIFF_HIGHLIGHT_COLUMN_KEYS and _display_is_nonzero(text),
        }
    return result


def save_analysis_pg_ecology_cell(
    sheet_slug: str,
    station_numb: int,
    year_number: int,
    column_key: str,
    value: Any,
    rounding_digits: int | str | None = None,
    emission_rounding_digits: int | str | None = None,
) -> dict[str, dict[str, Any]]:
    """Сохраняет параметр ПГ и возвращает обновлённые отображаемые ячейки."""
    key = str(column_key or "").strip()
    field = PG_ECOLOGY_OVERLAY_BY_KEY.get(key)
    if not field:
        raise ValueError(f"Столбец недоступен для редактирования: {column_key!r}")
    from app.ecology.services.analysis_pg_param_services import (
        upsert_analysis_pg_param_values,
    )

    upsert_analysis_pg_param_values(
        sheet_slug,
        int(station_numb),
        int(year_number),
        {field: value},
    )
    return format_related_pg_ecology_cells(
        sheet_slug,
        int(station_numb),
        int(year_number),
        key,
        rounding_digits,
        emission_rounding_digits,
    )


def format_related_zv_ecology_cells(
    sheet_slug: str,
    station_numb: int,
    year_number: int,
    column_key: str,
    rounding_digits: int | str | None = None,
    specific_rounding_digits: int | str | None = None,
    norm_zv_parts: list[str] | tuple[str, ...] = ZV_NORM_DEFAULT_PARTS,
) -> dict[str, dict[str, Any]]:
    """Форматирует отредактированную ячейку ЗВ и все расчётные итоги/удельные."""
    key = str(column_key or "").strip()
    related = ZV_ECOLOGY_EDIT_RELATED_KEYS.get(key)
    if not related:
        related = (key,) + tuple(ZV_COLUMN_FORMULA_KEYS) if key else tuple(ZV_COLUMN_FORMULA_KEYS)
    # Уникальный порядок: правка, затем расчётные.
    ordered: list[str] = []
    for item in related:
        if item and item not in ordered:
            ordered.append(item)
    digits = resolve_rounding_digits(rounding_digits)
    specific_digits = resolve_optional_rounding_digits(specific_rounding_digits)
    numb = int(station_numb)
    year = int(year_number)
    zv = load_zv_ecology_overlay_by_numb_year(sheet_slug, {numb}, [year]).get(
        (numb, year)
    )
    fuel = load_fuel_indicators_by_numb_year({numb}, [year]).get((numb, year))
    by_key = _zv_ecology_values_by_key(zv, fuel, norm_zv_parts)
    result: dict[str, dict[str, Any]] = {}
    for rel_key in ordered:
        text = _format_value(
            by_key.get(rel_key),
            True,
            column_rounding_digits({"key": rel_key}, digits, specific_digits=specific_digits),
        )
        result[rel_key] = {
            "text": text,
            "title": _format_full_value(by_key.get(rel_key)),
            "nonzero": rel_key in DIFF_HIGHLIGHT_COLUMN_KEYS and _display_is_nonzero(text),
        }
    return result


def save_analysis_zv_ecology_cell(
    sheet_slug: str,
    station_numb: int,
    year_number: int,
    column_key: str,
    value: Any,
    rounding_digits: int | str | None = None,
    specific_rounding_digits: int | str | None = None,
    norm_zv_parts: list[str] | tuple[str, ...] = ZV_NORM_DEFAULT_PARTS,
) -> dict[str, dict[str, Any]]:
    """Сохраняет параметр ЗВ и возвращает обновлённые отображаемые ячейки."""
    key = str(column_key or "").strip()
    field = ZV_ECOLOGY_OVERLAY_BY_KEY.get(key)
    if not field:
        raise ValueError(f"Столбец недоступен для редактирования: {column_key!r}")
    from app.ecology.services.analysis_zv_param_services import (
        upsert_analysis_zv_param_values,
    )

    upsert_analysis_zv_param_values(
        sheet_slug,
        int(station_numb),
        int(year_number),
        {field: value},
    )
    return format_related_zv_ecology_cells(
        sheet_slug,
        int(station_numb),
        int(year_number),
        key,
        rounding_digits,
        specific_rounding_digits,
        norm_zv_parts,
    )


def save_analysis_ecology_cell(
    sheet_slug: str,
    station_numb: int,
    year_number: int,
    column_key: str,
    value: Any,
    rounding_digits: int | str | None = None,
    emission_rounding_digits: int | str | None = None,
    specific_rounding_digits: int | str | None = None,
    norm_zv_parts: list[str] | tuple[str, ...] = ZV_NORM_DEFAULT_PARTS,
) -> dict[str, dict[str, Any]]:
    """Сохраняет правку страницы анализа ЗВ или ПГ."""
    slug = str(sheet_slug or "").strip()
    if slug in PG_ANALYSIS_SHEET_SLUGS:
        return save_analysis_pg_ecology_cell(
            slug,
            station_numb,
            year_number,
            column_key,
            value,
            rounding_digits,
            emission_rounding_digits,
        )
    if slug in ZV_ANALYSIS_SHEET_SLUGS:
        return save_analysis_zv_ecology_cell(
            slug,
            station_numb,
            year_number,
            column_key,
            value,
            rounding_digits,
            specific_rounding_digits,
            norm_zv_parts,
        )
    raise ValueError(f"Лист недоступен для редактирования: {sheet_slug!r}")


def _sum_optional(*values: Any) -> Decimal | None:
    total: Decimal | None = None
    for value in values:
        part = _decimal_or_none(value)
        if part is None:
            continue
        total = part if total is None else total + part
    return total


def _thousand_tut_to_tce(value: Any) -> Decimal | None:
    """тыс. т у.т. (Топливо) → т у.т. (страницы анализа)."""
    amount = _decimal_or_none(value)
    if amount is None:
        return None
    return amount * Decimal("1000")


def _fuel_indicator_payload(param: Any, group: Any) -> dict[str, Any]:
    district = getattr(group, "regional_district", None)
    federal = getattr(district, "federal_district", None) if district is not None else None
    sync_area = getattr(district, "synchronous_area", None) if district is not None else None
    res = getattr(group, "regional_energy_system", None)
    ues = getattr(res, "union_energy_system", None) if res is not None else None
    est = getattr(ues, "energy_system_type", None) if ues is not None else None
    gk = getattr(group, "gen_company_external_mapping", None)
    station_name = (getattr(group, "name", None) or getattr(param, "name", None) or "").strip()
    ugol = getattr(param, "ugol", None)
    torf = getattr(param, "torf", None)
    mazut = getattr(param, "mazut", None)
    gaz = getattr(param, "gaz", None)
    isk_gaz = getattr(param, "isk_gaz", None)
    proch = getattr(param, "proch", None)
    return {
        "ved": getattr(param, "ved", None),
        "federal_district": (getattr(federal, "name", None) or "").strip() or None,
        "ues": (getattr(ues, "name", None) or "").strip() or None,
        "res_name": (getattr(res, "name", None) or "").strip() or None,
        "est_name": (getattr(est, "name", None) or "").strip() or None,
        "regional_district": (getattr(district, "name", None) or "").strip() or None,
        "synchronous_area": (getattr(sync_area, "name", None) or "").strip() or None,
        "station_name": station_name or None,
        "gen_company_name": (getattr(gk, "external_name", None) or "").strip() or None,
        "energy_system_type_id": getattr(ues, "id_energy_system_type", None),
        "ues_id": getattr(ues, "id", None),
        "res_id": getattr(res, "id", None),
        "federal_district_id": getattr(federal, "id", None) or getattr(district, "id_federal_district", None),
        "regional_district_id": getattr(district, "id", None),
        "synchronous_area_id": getattr(sync_area, "id", None)
        or getattr(district, "id_synchronous_area", None),
        "nust": getattr(param, "nust", None),
        "e": getattr(param, "e", None),
        "q": getattr(param, "q", None),
        "ugol_tce": _thousand_tut_to_tce(ugol),
        "solid_tce": _thousand_tut_to_tce(_sum_optional(ugol, torf)),
        "mazut_tce": _thousand_tut_to_tce(mazut),
        "gaz_estest_tce": _thousand_tut_to_tce(gaz),
        "gaz_all_tce": _thousand_tut_to_tce(_sum_optional(gaz, isk_gaz)),
        "proch_tce": _thousand_tut_to_tce(proch),
    }


def _prefer_fuel_payload(current: dict[str, Any] | None, incoming: dict[str, Any]) -> dict[str, Any]:
    if current is None:
        return incoming
    from app.ecology.services.operating_tes_constants import VED_PARENT_SHELL

    current_ved = current.get("ved")
    incoming_ved = incoming.get("ved")
    if current_ved == VED_PARENT_SHELL and incoming_ved != VED_PARENT_SHELL:
        return incoming
    return current


def load_fuel_indicators_by_numb_year(
    numbs: set[int],
    years: list[int],
) -> dict[tuple[int, int], dict[str, Any]]:
    """ТЭП групп оборудования текущей версии БД: ключ (numb1120, год)."""
    if not numbs or not years:
        return {}
    try:
        from flask import has_app_context
        from sqlalchemy import or_
        from sqlalchemy.orm import selectinload

        from app.common.services.database_version_filter import get_current_db_version_id
        from app.extensions import db
        from app.fuel.models.fue_equipment_group_fuel_param_model import EquipmentGroupFuelParam
        from app.fuel.models.fue_equipment_group_model import EquipmentGroup
        from app.refdata.models.energy_systems.regional_energy_system_model import (
            RegionalEnergySystem,
        )
        from app.refdata.models.energy_systems.union_energy_system_model import (
            UnionEnergySystem,
        )
        from app.refdata.models.territories.regional_district_model import RegionalDistrict

        if not has_app_context():
            return {}
        version_id = get_current_db_version_id()
        numb_list = list(numbs)
        query = (
            db.session.query(EquipmentGroupFuelParam, EquipmentGroup)
            .join(EquipmentGroup, EquipmentGroup.id == EquipmentGroupFuelParam.equipment_group_id)
            .options(
                selectinload(EquipmentGroup.regional_energy_system)
                .selectinload(RegionalEnergySystem.union_energy_system)
                .selectinload(UnionEnergySystem.energy_system_type),
                selectinload(EquipmentGroup.regional_district).selectinload(
                    RegionalDistrict.federal_district
                ),
                selectinload(EquipmentGroup.regional_district).selectinload(
                    RegionalDistrict.synchronous_area
                ),
                selectinload(EquipmentGroup.gen_company_external_mapping),
            )
            .filter(EquipmentGroupFuelParam.year_number.in_(years))
            .filter(
                or_(
                    EquipmentGroupFuelParam.numb1120.in_(numb_list),
                    EquipmentGroup.numb.in_(numb_list),
                )
            )
        )
        if version_id is not None:
            query = query.filter(EquipmentGroup.database_version_id == version_id)
            query = query.filter(
                or_(
                    EquipmentGroupFuelParam.database_version_id == version_id,
                    EquipmentGroupFuelParam.database_version_id.is_(None),
                )
            )
        out: dict[tuple[int, int], dict[str, Any]] = {}
        for param, group in query.all():
            year = getattr(param, "year_number", None)
            if year is None:
                continue
            year_int = int(year)
            candidates: list[int] = []
            for raw in (getattr(param, "numb1120", None), getattr(group, "numb", None)):
                try:
                    numb = int(raw)
                except (TypeError, ValueError):
                    continue
                if numb in numbs:
                    candidates.append(numb)
            if not candidates:
                continue
            payload = _fuel_indicator_payload(param, group)
            for numb in candidates:
                key = (numb, year_int)
                out[key] = _prefer_fuel_payload(out.get(key), payload)
        return out
    except Exception:
        _logger.exception("ecology comparison: fuel indicators load failed")
        return {}


_SHEET_LAYOUT_PREVIEW_ROWS = 40


def _read_workbook_rows(path: Path) -> dict[str, list[list[Any]]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        by_name: dict[str, list[list[Any]]] = {}
        for sheet_name in workbook.sheetnames:
            if sheet_name == SKIP_SHEET_NAME:
                continue
            worksheet = workbook[sheet_name]
            by_name[sheet_name] = [list(row) for row in worksheet.iter_rows(values_only=True)]
        return by_name
    finally:
        workbook.close()


def _read_sheet_preview(path: Path, sheet_name: str, max_rows: int) -> list[list[Any]] | None:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return None
        worksheet = workbook[sheet_name]
        rows: list[list[Any]] = []
        for index, row in enumerate(worksheet.iter_rows(values_only=True)):
            rows.append(list(row))
            if index + 1 >= max_rows:
                break
        return rows
    finally:
        workbook.close()


@functools.lru_cache(maxsize=4)
def _parse_workbook_cached(path_str: str, mtime: float) -> dict[str, dict[str, Any]]:
    del mtime
    by_name = _read_workbook_rows(Path(path_str))
    parsed: dict[str, dict[str, Any]] = {}
    for meta in COMPARISON_SHEETS:
        raw = by_name.get(meta["excel_name"])
        if raw is None:
            continue
        parsed[meta["slug"]] = parse_sheet_rows(raw, meta)
    return parsed


@functools.lru_cache(maxsize=16)
def _parse_sheet_layout_cached(
    path_str: str,
    mtime: float,
    excel_name: str,
) -> dict[str, Any] | None:
    del mtime
    meta = next((item for item in COMPARISON_SHEETS if item["excel_name"] == excel_name), None)
    if meta is None:
        return None
    raw = _read_sheet_preview(Path(path_str), excel_name, _SHEET_LAYOUT_PREVIEW_ROWS)
    if raw is None:
        return None
    parsed = parse_sheet_rows(raw, meta)
    parsed["rows"] = []
    return parsed


def load_comparison_sheets(path: str | Path | None = None) -> dict[str, dict[str, Any]]:
    resolved = resolve_source_path(path)
    return _parse_workbook_cached(str(resolved), resolved.stat().st_mtime)


def load_comparison_sheet_layout(
    sheet_slug: str,
    path: str | Path | None = None,
) -> dict[str, Any] | None:
    """Заголовки и столбцы листа без чтения всех строк макета."""
    sheet = get_comparison_sheet(sheet_slug)
    if sheet is None:
        return None
    resolved = resolve_source_path(path)
    return _parse_sheet_layout_cached(
        str(resolved),
        resolved.stat().st_mtime,
        sheet["excel_name"],
    )


def _is_empty_numeric_row(values: list[Any], columns: list[dict[str, Any]]) -> bool:
    has_numeric = False
    for index, column in enumerate(columns):
        if not column["numeric"]:
            continue
        if str(column.get("key") or "") in STATION_GENERATION_COLUMN_KEYS:
            # Столбцы сверки с «Генерацией» не влияют на скрытие пустых строк.
            continue
        has_numeric = True
        value = values[index] if index < len(values) else None
        if _is_blank(value):
            continue
        try:
            if Decimal(str(value)) != 0:
                return False
        except Exception:
            return False
    return has_numeric


def column_rounding_digits(
    column: dict[str, Any],
    digits: int,
    emission_digits: int | None = None,
    specific_digits: int | None = None,
) -> int:
    """Знаки после запятой для столбца: у «Коэффициента эмиссии» и «Удельных» свои."""
    key = str(column.get("key") or "")
    if emission_digits is not None and key in EMISSION_COEFFICIENT_COLUMN_KEYS:
        return emission_digits
    if specific_digits is not None and key in ZV_SPECIFIC_ROUNDING_COLUMN_KEYS:
        return specific_digits
    return digits


def resolve_optional_rounding_digits(value: Any) -> int | None:
    """Знаки отдельного селектора; пусто – столбец округляется как основной «Округл»."""
    if value in (None, ""):
        return None
    return resolve_rounding_digits(value)


def _format_value(value: Any, numeric: bool, digits: int) -> str:
    if _is_blank(value):
        return ""
    if not numeric:
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    if digits == 0:
        # «Не округлять» показывает значение целиком, поэтому оставляем только
        # значащие цифры: хвост вроде 802393,6000000001 – шум исходного double,
        # а не данные (как на страницах ГИС ТЭК).
        cleaned = drop_float_noise(value)
        if cleaned is not None:
            value = cleaned
    return format_cell(value, digits=digits)


def _format_full_value(value: Any) -> str:
    """Значение целиком для подсказки при наведении: как в БД/расчёте, без округления."""
    if _is_blank(value):
        return ""
    return format_cell(value, digits=0)


# Столбцы-идентификаторы строки: подсказка со значением целиком им не нужна.
FULL_VALUE_TOOLTIP_SKIP_KEYS = frozenset({"numb1120", "Year", "NAME", "Stancia"})


def cell_titles(filled: list[Any], columns: list[dict[str, Any]]) -> list[str]:
    """Подсказки к ячейкам строки: значение ячейки без выбранного округления."""
    titles: list[str] = []
    for column in columns:
        key = str(column.get("key") or "")
        index = column.get("source_index")
        if not column.get("numeric") or key in FULL_VALUE_TOOLTIP_SKIP_KEYS or index is None:
            titles.append("")
            continue
        value = filled[index] if index < len(filled) else None
        titles.append(_format_full_value(value))
    return titles


def _display_is_nonzero(display: Any) -> bool:
    text = str(display or "").strip().replace(" ", "").replace("\u00a0", "").replace(",", ".")
    if not text:
        return False
    try:
        return float(text) != 0.0
    except (TypeError, ValueError):
        return False


_TERRITORIAL_LIST_KEYS = (
    "energy_system_type_list",
    "union_energy_system_list",
    "regional_energy_system_list",
    "federal_district_list",
    "regional_district_list",
)
_TERRITORIAL_MAP_KEYS = (
    "regional_energy_system_mapping",
    "regional_district_mapping",
    "est_to_ues_mapping",
    "est_to_res_mapping",
    "est_to_rd_mapping",
    "est_to_fd_mapping",
    "ues_to_est_mapping",
    "ues_to_res_mapping",
    "ues_to_rd_mapping",
    "ues_to_fd_mapping",
    "res_to_est_mapping",
    "res_to_ues_mapping_one",
    "res_to_rd_mapping",
    "res_to_fd_mapping",
    "rd_to_fd_mapping_one",
    "rd_to_res_mapping",
    "rd_to_ues_mapping",
    "rd_to_est_mapping",
    "fd_to_rd_mapping",
    "fd_to_res_mapping",
    "fd_to_ues_mapping",
    "fd_to_est_mapping",
)
_ID_FILTER_TO_FIELD = {
    "energy_system_type_filter": "energy_system_type_id",
    "union_energy_system_filter": "ues_id",
    "regional_energy_system_filter": "res_id",
    "federal_district_filter": "federal_district_id",
    "regional_district_filter": "regional_district_id",
}


def _clean_number(value: Any) -> Decimal | None:
    """Число без хвоста float-шума источника; None – пусто или не число.

    Расчётные столбцы страницы (d_*, K_*, CO2, UD_*, ЗВ) считаем от очищенных
    значений: иначе разность 3121 и 3120,2000000000003 даёт 0,7999999999997272.
    """
    if _is_blank(value) or isinstance(value, bool):
        return None
    return drop_float_noise(value)


def _significant_float(value: Decimal) -> float:
    """Расчётный столбец страницы – float, как и значения макета, но без шума."""
    return float(drop_float_noise(value))


def _safe_subtract(a: Any, b: Any) -> Any:
    """Безопасное вычитание: возвращает разность или None, если любое значение пусто."""
    left = _clean_number(a)
    right = _clean_number(b)
    if left is None or right is None:
        return None
    return _significant_float(left - right)


def _safe_divide(numerator: Any, denominator: Any) -> Any:
    """Безопасное деление: None при пустых значениях или нулевом знаменателе."""
    left = _clean_number(numerator)
    right = _clean_number(denominator)
    if left is None or right is None or right == 0:
        return None
    with localcontext() as ctx:
        ctx.prec = SIGNIFICANT_DIGITS
        ctx.rounding = ROUND_HALF_UP
        return float(+(left / right))


def _safe_sum(*values: Any) -> Any:
    """Сумма числовых слагаемых; None, если все пустые или есть нечисло."""
    total: Decimal | None = None
    for value in values:
        if _is_blank(value) or isinstance(value, bool):
            continue
        part = _clean_number(value)
        if part is None:
            return None
        total = part if total is None else total + part
    if total is None:
        return None
    return _significant_float(total)


def _brought_energy_mln_kwh(w_en: Any, teplo: Any) -> Any:
    """Приведённая энергия, млн кВт·ч: W_en + Teplo / 859,8 · 1000."""
    energy = _clean_number(w_en)
    heat = _clean_number(teplo)
    if energy is None and heat is None:
        return None
    with localcontext() as ctx:
        ctx.prec = SIGNIFICANT_DIGITS
        ctx.rounding = ROUND_HALF_UP
        total = (energy or Decimal(0)) + (heat or Decimal(0)) / Decimal(str(KCAL_PER_KWH)) * 1000
        return float(+total)


def _compute_differences(
    filled: list[Any],
    source_columns: list[dict[str, Any]],
) -> list[Any]:
    """Пересчитывает разностные столбцы: d_mazut, d_gaz, d_tv."""
    out = list(filled)
    for col_key, parent_key, child_key in [
        ("d_mazut", "MAZUT_6tp", "MAZUT_ek"),
        ("d_gaz", "GAZ_ESTEST", "GAZ_ek"),
        ("d_tv", "TvTopl_6tp", "TvTopl_ek"),
    ]:
        idx = _column_index(source_columns, {col_key})
        if idx is None:
            continue
        p_idx = _column_index(source_columns, {parent_key})
        c_idx = _column_index(source_columns, {child_key})
        if p_idx is None or c_idx is None:
            continue
        diff = _safe_subtract(filled[p_idx], filled[c_idx])
        if len(out) <= idx:
            out.extend([None] * (idx - len(out) + 1))
        out[idx] = diff
    return out


def _compute_emission_coefficients(
    filled: list[Any],
    source_columns: list[dict[str, Any]],
) -> list[Any]:
    """Коэффициенты эмиссии: выбросы СО2 / расход топлива."""
    out = list(filled)
    for col_key, co2_key, fuel_key in [
        ("K_gaz_ek", "CO2_gaz", "GAZ_ek"),
        ("K_mazut_ek", "CO2_mazut", "MAZUT_ek"),
        ("K_ugol_ek", "CO2_ugol", "TvTopl_ek"),
    ]:
        idx = _column_index(source_columns, {col_key})
        if idx is None:
            continue
        co2_idx = _column_index(source_columns, {co2_key})
        fuel_idx = _column_index(source_columns, {fuel_key})
        if co2_idx is None or fuel_idx is None:
            continue
        value = _safe_divide(filled[co2_idx], filled[fuel_idx])
        if len(out) <= idx:
            out.extend([None] * (idx - len(out) + 1))
        out[idx] = value
    return out


def _compute_total_co2(
    filled: list[Any],
    source_columns: list[dict[str, Any]],
) -> list[Any]:
    """Выброшено ПГ всего: CO2 = CO2_gaz + CO2_mazut + CO2_ugol."""
    out = list(filled)
    idx = _column_index(source_columns, {"CO2"})
    if idx is None:
        return out
    parts = []
    for key in ("CO2_gaz", "CO2_mazut", "CO2_ugol"):
        part_idx = _column_index(source_columns, {key})
        if part_idx is None:
            return out
        parts.append(filled[part_idx] if part_idx < len(filled) else None)
    total = _safe_sum(*parts)
    if len(out) <= idx:
        out.extend([None] * (idx - len(out) + 1))
    out[idx] = total
    return out


def _compute_specific_co2(
    filled: list[Any],
    source_columns: list[dict[str, Any]],
) -> list[Any]:
    """Удельные выбросы ПГ: UD_CO2_ek и UD_CO2_ener."""
    out = list(filled)
    co2_idx = _column_index(source_columns, {"CO2"})
    co2 = filled[co2_idx] if co2_idx is not None and co2_idx < len(filled) else None

    ud_ek_idx = _column_index(source_columns, {"UD_CO2_ek"})
    if ud_ek_idx is not None:
        fuel_parts = []
        missing = False
        for key in ("GAZ_ek", "MAZUT_ek", "TvTopl_ek"):
            part_idx = _column_index(source_columns, {key})
            if part_idx is None:
                missing = True
                break
            fuel_parts.append(filled[part_idx] if part_idx < len(filled) else None)
        value = None
        if not missing:
            fuel_total = _safe_sum(*fuel_parts)
            co2_value = _clean_number(co2)
            co2_kg = None if co2_value is None else co2_value * 1000
            value = _safe_divide(co2_kg, fuel_total)
        if len(out) <= ud_ek_idx:
            out.extend([None] * (ud_ek_idx - len(out) + 1))
        out[ud_ek_idx] = value

    ud_en_idx = _column_index(source_columns, {"UD_CO2_ener"})
    if ud_en_idx is not None:
        w_idx = _column_index(source_columns, {"W_en"})
        q_idx = _column_index(source_columns, {"Teplo"})
        value = None
        if w_idx is not None and q_idx is not None:
            w_en = filled[w_idx] if w_idx < len(filled) else None
            teplo = filled[q_idx] if q_idx < len(filled) else None
            value = _safe_divide(co2, _brought_energy_mln_kwh(w_en, teplo))
        if len(out) <= ud_en_idx:
            out.extend([None] * (ud_en_idx - len(out) + 1))
        out[ud_en_idx] = value
    return out


def _compute_gaseous_and_liquid(
    filled: list[Any],
    source_columns: list[dict[str, Any]],
) -> list[Any]:
    """Газообразные и жидкие, всего: сумма SO2 + CO + NOx + углеводороды + ЛОС + прочие."""
    out = list(filled)
    idx = _column_index(source_columns, {ZV_GASEOUS_TOTAL_KEY})
    if idx is None:
        return out
    parts: list[Any] = []
    found_part = False
    for key in ZV_GASEOUS_PART_KEYS:
        part_idx = _column_index(source_columns, {key})
        if part_idx is None:
            continue
        found_part = True
        parts.append(filled[part_idx] if part_idx < len(filled) else None)
    if not found_part:
        return out
    total = _safe_sum(*parts)
    if len(out) <= idx:
        out.extend([None] * (idx - len(out) + 1))
    out[idx] = total
    return out


def _compute_zv_total(
    filled: list[Any],
    source_columns: list[dict[str, Any]],
) -> list[Any]:
    """Всего ЗВ: твёрдые вещества + газообразные и жидкие (всего)."""
    out = list(filled)
    idx = _column_index(source_columns, {ZV_TOTAL_KEY})
    if idx is None:
        return out
    parts: list[Any] = []
    found_part = False
    for key in (ZV_SOLID_KEY, ZV_GASEOUS_TOTAL_KEY):
        part_idx = _column_index(source_columns, {key})
        if part_idx is None:
            continue
        found_part = True
        parts.append(filled[part_idx] if part_idx < len(filled) else None)
    if not found_part:
        return out
    total = _safe_sum(*parts)
    if len(out) <= idx:
        out.extend([None] * (idx - len(out) + 1))
    out[idx] = total
    return out


def _filled_at(filled: list[Any], columns: list[dict[str, Any]], keys: set[str]) -> Any:
    idx = _column_index(columns, keys)
    if idx is None or idx >= len(filled):
        return None
    return filled[idx]


def _assign_filled(
    out: list[Any],
    columns: list[dict[str, Any]],
    key: str,
    value: Any,
) -> None:
    idx = _column_index(columns, {key})
    if idx is None:
        return
    if len(out) <= idx:
        out.extend([None] * (idx - len(out) + 1))
    out[idx] = value


def _kg_per_tce(mass_t: Any, fuel_tce: Any) -> Any:
    """т → кг/т у.т.: масса, т · 1000 / расход топлива, т у.т."""
    mass = _clean_number(mass_t)
    if mass is None:
        return None
    return _safe_divide(mass * 1000, fuel_tce)


def _compute_zv_specific_metrics(
    filled: list[Any],
    source_columns: list[dict[str, Any]],
    norm_zv_parts: list[str] | tuple[str, ...] = ZV_NORM_DEFAULT_PARTS,
) -> list[Any]:
    """Приведённая энергия и удельные ЗВ: на энергию, на выработку и на топливо."""
    out = list(filled)
    solid = _filled_at(out, source_columns, {ZV_SOLID_KEY})
    so2 = _filled_at(out, source_columns, {"SO2_gaz"})
    nox = _filled_at(out, source_columns, {"NOx_gaz"})
    zv_normir = _safe_sum(*(
        _filled_at(out, source_columns, {key})
        for key, attr in ZV_ECOLOGY_OVERLAY_BY_KEY.items()
        if attr in norm_zv_parts
    ))
    _assign_filled(out, source_columns, "ZV_normir", zv_normir)

    w_en = _filled_at(out, source_columns, {"W_en"})
    teplo = _filled_at(out, source_columns, {"Teplo"})
    energy = _brought_energy_mln_kwh(w_en, teplo)
    _assign_filled(out, source_columns, "energy", energy)

    total_zv = _filled_at(out, source_columns, {ZV_TOTAL_KEY})
    _assign_filled(out, source_columns, "udel_na_energ", _safe_divide(total_zv, energy))
    _assign_filled(out, source_columns, ZV_EE_SPECIFIC_KEY, _safe_divide(zv_normir, w_en))
    _assign_filled(out, source_columns, "normir_ZV_na_energ", _safe_divide(zv_normir, energy))
    _assign_filled(
        out, source_columns, "normir_TVERD_na_energ", _safe_divide(solid, energy)
    )
    _assign_filled(out, source_columns, "normir_SO2_na_energ", _safe_divide(so2, energy))
    _assign_filled(out, source_columns, "normir_NOx_na_energ", _safe_divide(nox, energy))

    mazut = _filled_at(out, source_columns, set(ZV_MAZUT_FUEL_KEYS))
    coal = _filled_at(out, source_columns, set(ZV_SOLID_FUEL_KEYS))
    gas = _filled_at(out, source_columns, set(ZV_GAS_FUEL_KEYS))
    solid_and_oil = _safe_sum(mazut, coal)
    all_fuel = _safe_sum(gas, mazut, coal)
    _assign_filled(out, source_columns, "udel_tv", _kg_per_tce(solid, solid_and_oil))
    _assign_filled(out, source_columns, "udel_SO2", _kg_per_tce(so2, solid_and_oil))
    _assign_filled(out, source_columns, "udel_NOx", _kg_per_tce(nox, all_fuel))
    return out


def insert_column_after(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
    after_key: str,
    new_column: dict[str, Any],
    caption_text: str,
    data_rows: list[list[Any]] | None = None,
) -> None:
    """Вставляет столбец после after_key в список колонок, шапку и строки данных."""
    new_key = str(new_column.get("key") or "")
    if not new_key or any(str(col.get("key") or "") == new_key for col in columns):
        return
    after = next((i for i, col in enumerate(columns) if col.get("key") == after_key), None)
    if after is None:
        return
    columns.insert(after + 1, dict(new_column))
    for row in data_rows or []:
        while len(row) <= after:
            row.append(None)
        row.insert(after + 1, None)
    for row in header_rows:
        insert_at = None
        is_key_row = bool(row and row[0].get("is_key"))
        for pos, cell in enumerate(row):
            col_i = cell.get("col_index")
            if col_i is not None and int(col_i) > after:
                cell["col_index"] = int(col_i) + 1
            if str(cell.get("column_key") or "") == after_key or col_i == after:
                insert_at = pos + 1
        if insert_at is None:
            continue
        row.insert(
            insert_at,
            {
                "text": new_key if is_key_row else caption_text,
                "rowspan": 1,
                "colspan": 1,
                "sticky": False,
                "col_index": after + 1,
                "column_key": new_key,
                "is_key": is_key_row,
            },
        )


def ensure_zv_electricity_specific_column(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
    data_rows: list[list[Any]] | None = None,
) -> None:
    """Столбец «на выработку электроэнергии» рядом с удельными на приведённую."""
    insert_column_after(
        columns,
        header_rows,
        ZV_EE_SPECIFIC_AFTER_KEY,
        {"key": ZV_EE_SPECIFIC_KEY, "numeric": True, "sticky": False},
        ZV_ENERGY_SPECIFIC_LEAF_LABELS[ZV_EE_SPECIFIC_KEY],
        data_rows=data_rows,
    )


def _first_column_key(
    columns: list[dict[str, Any]],
    keys: tuple[str, ...],
) -> str | None:
    """Первый из ключей, который есть среди столбцов."""
    present = {str(column.get("key") or "") for column in columns}
    return next((key for key in keys if key in present), None)


def station_generation_check_pairs(
    columns: list[dict[str, Any]],
) -> dict[str, str]:
    """Столбец «Генерации» → столбец макета, с которым сверяется ячейка страницы."""
    pairs: dict[str, str] = {}
    for key, _header, macro_keys in STATION_GENERATION_COLUMNS:
        counterpart = _first_column_key(columns, macro_keys)
        if counterpart is not None and _column_index(columns, {key}) is not None:
            pairs[key] = counterpart
    return pairs


def ensure_station_generation_columns(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
    data_rows: list[list[Any]] | None = None,
) -> None:
    """Столбцы «Руст, МВт» и «Выработка ЭЭ, млн кВтч (Генерация)» сразу после столбцов ТЭП."""
    for key, header, macro_keys in STATION_GENERATION_COLUMNS:
        after_key = _first_column_key(columns, macro_keys)
        if after_key is None:
            continue
        insert_column_after(
            columns,
            header_rows,
            after_key,
            {"key": key, "numeric": True, "sticky": False},
            header,
            data_rows=data_rows,
        )


def ensure_station_affiliation_column(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
    data_rows: list[list[Any]] | None = None,
) -> None:
    """Столбец «Принадлежность» сразу после «Отпуск ТЭ, тыс.Гкал» (Q)."""
    insert_column_after(
        columns,
        header_rows,
        STATION_AFFILIATION_AFTER_KEY,
        {"key": STATION_AFFILIATION_COLUMN_KEY, "numeric": False, "sticky": False},
        STATION_AFFILIATION_LABEL,
        data_rows=data_rows,
    )


def apply_full_height_headers(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
    labels: dict[str, str],
    *,
    highlight_flag: str | None = None,
) -> None:
    """Подписи вставленных столбцов: одна ячейка на всю высоту шапки, как у столбцов ТЭП."""
    if not labels:
        return
    index_by_key = {
        str(column["key"]): index
        for index, column in enumerate(columns)
        if column.get("key")
    }
    for column in columns:
        key = str(column.get("key") or "")
        if key not in labels:
            continue
        column["label"] = labels[key]
        if highlight_flag:
            column[highlight_flag] = True
    if not header_rows:
        return
    for row in header_rows:
        row[:] = [
            cell for cell in row if str(cell.get("column_key") or "") not in labels
        ]
    top = header_rows[0]
    for key, header in labels.items():
        index = index_by_key.get(key)
        if index is None:
            continue
        cell = {
            "text": header,
            "rowspan": len(header_rows),
            "colspan": 1,
            "sticky": False,
            "col_index": columns[index].get("source_index"),
            "column_key": key,
            "diff_highlight": False,
            "fuel_tep_highlight": False,
        }
        if highlight_flag:
            cell[highlight_flag] = True
        top.insert(_insert_index_for_column(top, index, columns, index_by_key), cell)


def station_generation_mismatch(
    cells: list[Any],
    columns: list[dict[str, Any]],
    check_pairs: dict[str, str],
) -> list[bool]:
    """Ячейки новых столбцов, где станционное значение не совпало со столбцом макета."""
    index_by_key = {
        str(column.get("key") or ""): index for index, column in enumerate(columns)
    }
    flags = [False] * len(cells)
    for key, counterpart in check_pairs.items():
        index = index_by_key.get(key)
        target = index_by_key.get(counterpart)
        if index is None or target is None or index >= len(cells) or target >= len(cells):
            continue
        flags[index] = bool(cells[index]) and cells[index] != cells[target]
    return flags


def apply_station_generation_overlay(
    filled: list[Any],
    columns: list[dict[str, Any]],
    overlay: dict[str, Any],
) -> list[Any]:
    """Название станции «Генерации» и её «Руст»/выработка для строки года."""
    names = overlay.get("names_by_numb") or {}
    values = overlay.get("values_by_numb_year") or {}
    if not names and not values:
        return filled
    out = list(filled)
    numb = _row_numb(out, columns)
    name = names.get(numb) if numb is not None else None
    if name:
        for name_key in ("NAME", "Stancia"):
            index = _column_index(columns, {name_key})
            if index is not None and index < len(out):
                out[index] = name
    year = _row_year(out, columns)
    row_values = values.get((numb, year)) if numb is not None and year is not None else None
    for key, value in (row_values or {}).items():
        index = _column_index(columns, {key})
        if index is not None and index < len(out):
            out[index] = value
    return out


def _generation_stations_by_numb(
    numbs: list[int],
    version_id: int | None,
) -> dict[int, dict[int, str]]:
    """NUMB группы оборудования → {ID станции «Генерации»: её название}."""
    from app.ecology.services.gis_tek_page_services import (
        _equipment_group_query,
        linked_generation_stations,
    )
    from app.fuel.models.fue_equipment_group_model import EquipmentGroup

    groups = (
        _equipment_group_query(version_id)
        .filter(EquipmentGroup.numb.in_(numbs))
        .all()
    )
    stations_by_numb: dict[int, dict[int, str]] = {}
    for group in groups:
        linked = stations_by_numb.setdefault(int(group.numb), {})
        for station in linked_generation_stations(group, version_id):
            linked[int(station["id"])] = str(station["name"])
    return stations_by_numb


def _station_links_by_numb(
    stations_by_numb: dict[int, dict[int, str]],
) -> dict[int, list[dict[str, Any]]]:
    """Ссылки на карточку станции: порядок имён как в столбце «Электростанция»."""
    return {
        numb: [
            {"id": station_id, "name": name}
            for station_id, name in sorted(
                stations.items(), key=lambda item: (item[1].casefold(), item[0])
            )
        ]
        for numb, stations in stations_by_numb.items()
    }


def load_generation_station_links_by_numb(
    numbs: set[int],
) -> dict[int, list[dict[str, Any]]]:
    """Станции «Генерации», привязанные к группам оборудования (по их NUMB).

    Пустой список у NUMB – группа не привязана к станции: на страницах анализа
    «С котельными» такие группы не показываем.
    """
    numb_list = sorted({int(numb) for numb in numbs or [] if numb is not None})
    if not numb_list:
        return {}
    from flask import has_app_context

    if not has_app_context():
        return {}

    from app.common.services.database_version_filter import get_current_db_version_id

    return _station_links_by_numb(
        _generation_stations_by_numb(numb_list, get_current_db_version_id())
    )


def load_station_generation_overlay(
    numbs: set[int],
    years: list[int],
) -> dict[str, Any]:
    """Станции «Генерации» по группам оборудования текущей версии БД.

    Названия те же, что в столбце «Электростанция (Генерация)» на /ecology/gis_tek/,
    значения – как в карточке станции: «Руст, МВт» из строки «Всего по электростанции»
    и годовая выработка из таблицы «Выработка электроэнергии электростанцией, млн кВтч».
    Несколько станций у одной группы оборудования суммируются.
    """
    numb_list = sorted({int(numb) for numb in numbs or [] if numb is not None})
    year_list = sorted({int(year) for year in years or [] if year is not None})
    if not numb_list or not year_list:
        return {}
    from flask import has_app_context

    if not has_app_context():
        return {}

    from app.common.services.database_version_filter import get_current_db_version_id
    from app.generation.services.station_services.station_power_aggregation import (
        load_stations_powers_by_year,
    )
    from app.generation.services.station_services.station_services import (
        stations_annual_energy_by_year,
    )

    version_id = get_current_db_version_id()
    stations_by_numb = _generation_stations_by_numb(numb_list, version_id)
    names_by_numb = {
        numb: ", ".join(sorted(stations.values(), key=str.casefold))
        or GENERATION_STATION_UNKNOWN_NAME
        for numb, stations in stations_by_numb.items()
    }
    # Те же станции для гиперссылок на карточку: порядок как в names_by_numb.
    station_links_by_numb = _station_links_by_numb(stations_by_numb)
    station_ids = {
        station_id for stations in stations_by_numb.values() for station_id in stations
    }
    powers = load_stations_powers_by_year(station_ids, year_list, version_id)
    energy = stations_annual_energy_by_year(
        station_ids, year_list[0], year_list[-1], version_id
    )
    values: dict[tuple[int, int], dict[str, Any]] = {}
    for numb, stations in stations_by_numb.items():
        for year in year_list:
            rust = _safe_sum(
                *(
                    powers.get(station_id, {}).get(year, {}).get("p_ust")
                    for station_id in stations
                )
            )
            generation = _safe_sum(
                *(energy.get(station_id, {}).get(year) for station_id in stations)
            )
            if rust is None and generation is None:
                continue
            values[(numb, year)] = {
                STATION_RUST_COLUMN_KEY: rust,
                STATION_GENERATION_COLUMN_KEY: generation,
            }
    return {
        "names_by_numb": names_by_numb,
        "station_links_by_numb": station_links_by_numb,
        "values_by_numb_year": values,
    }


def load_station_affiliations_by_numb_year(
    numbs: set[int],
    years: list[int],
) -> dict[tuple[int, int], str]:
    """«Принадлежность» групп оборудования по годам – как на страницах ГИС ТЭК.

    Тип станции берётся из кода VED строки «Топливо» за год; у оболочки составной
    станции своего ved нет, поэтому тип берётся у её дочерних групп, если он у них
    один (та же логика, что на /ecology/gis_tek/pg/ и /ecology/gis_tek/pollutants/).
    """
    numb_list = sorted({int(numb) for numb in numbs or [] if numb is not None})
    year_list = sorted({int(year) for year in years or [] if year is not None})
    if not numb_list or not year_list:
        return {}
    from flask import has_app_context

    if not has_app_context():
        return {}

    from sqlalchemy import or_

    from app.common.services.database_version_filter import get_current_db_version_id
    from app.ecology.services.gis_tek_page_services import (
        _equipment_group_query,
        station_affiliation,
        station_types_by_equipment_group,
    )
    from app.fuel.models.fue_equipment_group_model import EquipmentGroup

    version_id = get_current_db_version_id()
    wanted = set(numb_list)
    groups = (
        _equipment_group_query(version_id)
        .filter(or_(EquipmentGroup.numb.in_(numb_list), EquipmentGroup.main.in_(numb_list)))
        .all()
    )
    out: dict[tuple[int, int], str] = {}
    for year in year_list:
        types_by_group = station_types_by_equipment_group(groups, year, version_id)
        for group in groups:
            try:
                numb = int(group.numb)
            except (TypeError, ValueError):
                continue
            if numb not in wanted:
                continue
            value = station_affiliation(types_by_group.get(group.id))
            current = out.get((numb, year))
            if current is None or (
                current == STATION_AFFILIATION_UNKNOWN
                and value != STATION_AFFILIATION_UNKNOWN
            ):
                out[(numb, year)] = value
    return out


def apply_station_affiliation_cells(
    filled: list[Any],
    columns: list[dict[str, Any]],
    affiliations: dict[tuple[int, int], str],
) -> list[Any]:
    """Подставляет «Принадлежность» строки года; без данных ГИС ТЭК – «–»."""
    index = _column_index(columns, {STATION_AFFILIATION_COLUMN_KEY})
    if index is None or index >= len(filled):
        return filled
    numb = _row_numb(filled, columns)
    year = _row_year(filled, columns)
    value = affiliations.get((numb, year)) if numb is not None and year is not None else None
    out = list(filled)
    out[index] = value or STATION_AFFILIATION_UNKNOWN
    return out


def _compute_derived_columns(
    filled: list[Any],
    source_columns: list[dict[str, Any]],
    norm_zv_parts: list[str] | tuple[str, ...] = ZV_NORM_DEFAULT_PARTS,
) -> list[Any]:
    """Разности 6тп−экология, коэффициенты эмиссии, итоги ПГ/ЗВ и удельные."""
    out = _compute_differences(filled, source_columns)
    out = _compute_emission_coefficients(out, source_columns)
    out = _compute_total_co2(out, source_columns)
    out = _compute_specific_co2(out, source_columns)
    out = _compute_gaseous_and_liquid(out, source_columns)
    out = _compute_zv_total(out, source_columns)
    return _compute_zv_specific_metrics(out, source_columns, norm_zv_parts)


def _empty_territorial_filter_payload() -> dict[str, Any]:
    data: dict[str, Any] = {key: [] for key in _TERRITORIAL_LIST_KEYS}
    data.update({key: {} for key in _TERRITORIAL_MAP_KEYS})
    for key in TERRITORY_FILTER_KEYS[:5]:
        data[key] = []
    data["station_name_filter"] = ""
    data["gen_company_filter"] = ""
    data["station_header_filter"] = ""
    data["affiliation_filter"] = ""
    return data


def _as_int_ids(values: list[Any]) -> list[int]:
    out: list[int] = []
    seen: set[int] = set()
    for value in values:
        try:
            number = int(value)
        except (TypeError, ValueError):
            continue
        if number in seen:
            continue
        seen.add(number)
        out.append(number)
    return out


def parse_territory_filters(args: Any = None) -> dict[str, Any]:
    filters: dict[str, Any] = {
        key: _as_int_ids(_arg_values(args, key)) for key in TERRITORY_FILTER_KEYS[:5]
    }
    station_values = _arg_values(args, "station_name_filter")
    gen_values = _arg_values(args, "gen_company_filter")
    filters["station_name_filter"] = str(station_values[0]).strip() if station_values else ""
    filters["gen_company_filter"] = str(gen_values[0]).strip() if gen_values else ""
    station_header_values = _arg_values(args, "station_header_filter")
    filters["station_header_filter"] = (
        str(station_header_values[0]).strip() if station_header_values else ""
    )
    affiliation_values = _arg_values(args, "affiliation_filter")
    filters["affiliation_filter"] = (
        str(affiliation_values[0]).strip() if affiliation_values else ""
    )
    allowed_ved = {key for key, _label in VED_FILTER_OPTIONS}
    ved_raw = [str(value).strip() for value in _arg_values(args, "ved_filter")]
    ved_value = next((value for value in ved_raw if value in allowed_ved), "")
    filters["ved_filter"] = ved_value
    return filters


def load_territorial_filter_context(filters: dict[str, Any] | None = None) -> dict[str, Any]:
    selected = filters or parse_territory_filters(None)
    data = _empty_territorial_filter_payload()
    try:
        from flask import has_app_context

        if has_app_context():
            from app.generation.services.station_services.filters_services import (
                get_territorial_filter_reference_data,
            )

            data.update(get_territorial_filter_reference_data())
    except Exception:
        _logger.exception("ecology comparison: territorial filter context unavailable")
    for key in TERRITORY_FILTER_KEYS[:5]:
        data[key] = list(selected.get(key) or [])
    data["station_name_filter"] = selected.get("station_name_filter") or ""
    data["gen_company_filter"] = selected.get("gen_company_filter") or ""
    data["station_header_filter"] = selected.get("station_header_filter") or ""
    data["affiliation_filter"] = selected.get("affiliation_filter") or ""
    data["ved_filter"] = selected.get("ved_filter") or ""
    data["ved_filter_options"] = [
        {"value": value, "label": label} for value, label in VED_FILTER_OPTIONS
    ]
    return data


def _keep_query_args(
    start: int | None,
    end: int | None,
    digits: int | None,
    view_year: int | None,
    filters: dict[str, Any],
    *,
    include_years: bool = True,
    include_view: bool = True,
    per_page: Any = None,
    emission_digits: int | None = None,
    specific_digits: int | None = None,
) -> dict[str, Any]:
    query: dict[str, Any] = {}
    if include_years:
        if start is not None:
            query["start_year"] = start
        if end is not None:
            query["end_year"] = end
        if include_view and view_year is not None:
            query["view_year"] = view_year
    if digits is not None:
        query["rounding_digits"] = digits
    if emission_digits is not None:
        query["emission_rounding_digits"] = emission_digits
    if specific_digits is not None:
        query["specific_rounding_digits"] = specific_digits
    for key in TERRITORY_FILTER_KEYS[:5]:
        values = filters.get(key) or []
        if values:
            query[key] = values
    station = (filters.get("station_name_filter") or "").strip()
    if station:
        query["station_name_filter"] = station
    station_header = str(filters.get("station_header_filter") or "").strip()
    if station_header:
        query["station_header_filter"] = station_header
    affiliation = str(filters.get("affiliation_filter") or "").strip()
    if affiliation:
        query["affiliation_filter"] = affiliation
    gen_company = (filters.get("gen_company_filter") or "").strip()
    if gen_company:
        query["gen_company_filter"] = gen_company
    ved_filter = str(filters.get("ved_filter") or "").strip()
    if ved_filter:
        query["ved_filter"] = ved_filter
    if per_page not in (None, ""):
        query["per_page"] = per_page
    numb_filter = filters.get("numb1120_filter") or []
    if numb_filter:
        query["numb1120_filter"] = numb_filter
    if "norm_zv" in filters:
        query["norm_zv"] = filters["norm_zv"] or [""]
    query["show_totals"] = "1" if filters.get("show_totals") else "0"
    query["show_empty_rows"] = "1" if filters.get("show_empty_rows") else "0"
    query["show_reported"] = "1" if filters.get("show_reported") else "0"
    return query


def prepare_visible_table(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
    extra_hidden_keys: frozenset[str] | set[str] | None = None,
) -> tuple[list[dict[str, Any]], list[list[dict[str, Any]]], int]:
    """Скрыть ФО/ОЭС/субъект и пересчитать липкие колонки без изменения кэша парсера."""
    hidden_keys = HIDDEN_TABLE_COLUMN_KEYS
    if extra_hidden_keys:
        hidden_keys = hidden_keys | frozenset(extra_hidden_keys)
    copied: list[dict[str, Any]] = []
    hidden_idx: set[int] = set()
    for index, column in enumerate(columns):
        item = dict(column)
        item["hidden"] = item["key"] in hidden_keys
        item["sticky"] = (not item["hidden"]) and item["key"] in STICKY_COLUMN_KEYS
        item["diff_highlight"] = item["key"] in DIFF_HIGHLIGHT_COLUMN_KEYS
        if item["hidden"]:
            hidden_idx.add(index)
        copied.append(item)

    sticky_n = 0
    for item in copied:
        if item["hidden"] or not item["sticky"]:
            item["sticky_index"] = None
            continue
        item["sticky_index"] = sticky_n
        sticky_n += 1

    visible: list[dict[str, Any]] = []
    for index, item in enumerate(copied):
        if item["hidden"]:
            continue
        shown = dict(item)
        shown["source_index"] = index
        visible.append(shown)

    headers: list[list[dict[str, Any]]] = []
    for row in header_rows:
        new_row: list[dict[str, Any]] = []
        for cell in row:
            col_i = int(cell.get("col_index") or 0)
            span = int(cell.get("colspan") or 1)
            covered = range(col_i, col_i + span)
            if all(idx in hidden_idx for idx in covered):
                continue
            new_cell = dict(cell)
            if span == 1 and col_i < len(copied) and not copied[col_i]["hidden"]:
                new_cell["column_key"] = copied[col_i]["key"]
                new_cell["sticky"] = copied[col_i]["sticky"]
                new_cell["diff_highlight"] = copied[col_i]["diff_highlight"]
                if copied[col_i]["sticky"]:
                    new_cell["col_index"] = copied[col_i]["sticky_index"]
            else:
                new_cell["sticky"] = False
                new_cell["diff_highlight"] = False
            new_row.append(new_cell)
        if new_row:
            headers.append(new_row)
    return visible, headers, sticky_n


def swap_visible_columns(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
    key_a: str,
    key_b: str,
) -> None:
    """Меняет местами два видимых столбца и их ячейки шапки."""
    idx_a = next((i for i, col in enumerate(columns) if col.get("key") == key_a), None)
    idx_b = next((i for i, col in enumerate(columns) if col.get("key") == key_b), None)
    if idx_a is None or idx_b is None or idx_a == idx_b:
        return
    columns[idx_a], columns[idx_b] = columns[idx_b], columns[idx_a]
    for row in header_rows:
        pos_a = next((i for i, cell in enumerate(row) if cell.get("column_key") == key_a), None)
        pos_b = next((i for i, cell in enumerate(row) if cell.get("column_key") == key_b), None)
        if pos_a is None or pos_b is None:
            continue
        row[pos_a], row[pos_b] = row[pos_b], row[pos_a]


def apply_zv_display_column_order(
    columns: list[dict[str, Any]],
    header_rows: list[list[dict[str, Any]]],
) -> None:
    """Упорядочивает столбцы ЗВ, оставляя всего ЗВ на приведённую энергию справа."""
    for key_a, key_b in ZV_DISPLAY_COLUMN_SWAPS:
        idx_a = next((i for i, col in enumerate(columns) if col.get("key") == key_a), None)
        idx_b = next((i for i, col in enumerate(columns) if col.get("key") == key_b), None)
        if idx_a is None or idx_b is None or idx_a >= idx_b:
            continue
        swap_visible_columns(columns, header_rows, key_a, key_b)
    index = next((i for i, col in enumerate(columns) if col.get("key") == "udel_na_energ"), None)
    if index is not None:
        columns.append(columns.pop(index))
        for row in header_rows:
            cell = _pop_header_cell(row, "udel_na_energ")
            if cell is not None:
                row.append(cell)


def _text_contains(needle: str, haystack: Any) -> bool:
    return text_contains(needle, haystack, blank_as_empty=True)


def row_matches_territory_filters(
    filled: list[Any],
    columns: list[dict[str, Any]],
    fuel: dict[str, Any] | None,
    filters: dict[str, Any],
) -> bool:
    for key, field in _ID_FILTER_TO_FIELD.items():
        selected = filters.get(key) or []
        if not selected:
            continue
        value = (fuel or {}).get(field)
        try:
            value_id = int(value)
        except (TypeError, ValueError):
            return False
        if value_id not in {int(item) for item in selected}:
            return False

    station_q = (filters.get("station_name_filter") or "").strip()
    if station_q:
        names: list[Any] = []
        if fuel:
            names.append(fuel.get("station_name"))
        station_idx = _column_index(columns, {"NAME", "Stancia"})
        if station_idx is not None and station_idx < len(filled):
            names.append(filled[station_idx])
        if not any(_text_contains(station_q, name) for name in names):
            return False

    gen_q = (filters.get("gen_company_filter") or "").strip()
    if gen_q:
        names = []
        if fuel:
            names.append(fuel.get("gen_company_name"))
        gen_idx = _column_index(columns, {"Gen_komp"})
        if gen_idx is not None and gen_idx < len(filled):
            names.append(filled[gen_idx])
        if not any(_text_contains(gen_q, name) for name in names):
            return False

    ved_group = str(filters.get("ved_filter") or "").strip()
    if ved_group:
        allowed = set(VED_FILTER_TO_VALUES.get(ved_group, ()))
        if allowed:
            raw_ved = (fuel or {}).get("ved")
            try:
                ved_value = int(raw_ved)
            except (TypeError, ValueError):
                return False
            if ved_value not in allowed:
                return False
    return True


ROWSPAN_COLUMN_KEYS = frozenset({"NAME", "Stancia", "numb1120"})
# При суммировании по регионам эти столбцы не складываем – пересчитываем формулами.
SUMMARY_RECOMPUTE_KEYS = frozenset(PG_DERIVED_COLUMN_KEYS) | frozenset(ZV_COLUMN_FORMULA_KEYS)
SUMMARY_IDENTITY_KEYS = frozenset(
    {
        "NAME",
        "Stancia",
        "numb1120",
        "Year",
        "Fed_okr",
        "nameoes",
        "oblname",
        "oblast",
        "Gen_komp",
        "SF6",
        "BLST",
        "c32",
        "OBOR",
        "obor",
    }
)


def parse_show_totals_arg(args: Any = None) -> bool:
    return parse_truthy_arg(args, "show_totals", use_getlist=True)


def parse_show_empty_rows_arg(args: Any = None) -> bool:
    """Кнопка «Пустые строки»: по умолчанию не нажата (скрывать)."""
    return parse_truthy_arg(args, "show_empty_rows", use_getlist=True)


def parse_show_reported_arg(args: Any = None) -> bool:
    """Кнопка «Отчитавшиеся»: по умолчанию не нажата (столбцы скрыты)."""
    return parse_truthy_arg(args, "show_reported", use_getlist=True)


def _region_keys(meta: dict[str, Any]) -> dict[str, tuple]:
    return {
        "res": (meta.get("est_id"), meta.get("ues_id"), meta.get("res_id")),
        "ues": (meta.get("est_id"), meta.get("ues_id")),
        "est": (meta.get("est_id"),),
    }


def _sum_filled_rows(
    filled_list: list[list[Any]],
    source_columns: list[dict[str, Any]],
    norm_zv_parts: list[str] | tuple[str, ...] = ZV_NORM_DEFAULT_PARTS,
) -> list[Any] | None:
    if not filled_list:
        return None
    width = max(len(source_columns), max((len(row) for row in filled_list), default=0))
    out: list[Any] = [None] * width
    for index, column in enumerate(source_columns):
        key = str(column.get("key") or "")
        if key in SUMMARY_IDENTITY_KEYS or key in SUMMARY_RECOMPUTE_KEYS:
            continue
        if not column.get("numeric"):
            continue
        parts = [
            row[index] if index < len(row) else None
            for row in filled_list
        ]
        out[index] = _safe_sum(*parts)
    # Пересчёт d_*/K_*/CO2/UD_* поверх сумм слагаемых (если столбцы есть).
    return _compute_derived_columns(out, source_columns, norm_zv_parts)


def _format_summary_row(
    *,
    label: str,
    year: int,
    filled: list[Any],
    source_columns: list[dict[str, Any]],
    visible_columns: list[dict[str, Any]],
    digits: int,
    key: str,
    is_first: bool,
    emission_digits: int | None = None,
    specific_digits: int | None = None,
) -> dict[str, Any]:
    year_idx = _column_index(source_columns, {"Year"})
    if year_idx is not None:
        if len(filled) <= year_idx:
            filled.extend([None] * (year_idx - len(filled) + 1))
        filled[year_idx] = year
    for name_key in ("NAME", "Stancia"):
        name_idx = _column_index(source_columns, {name_key})
        if name_idx is None:
            continue
        if len(filled) <= name_idx:
            filled.extend([None] * (name_idx - len(filled) + 1))
        filled[name_idx] = label
    numb_idx = _column_index(source_columns, {"numb1120"})
    if numb_idx is not None:
        if len(filled) <= numb_idx:
            filled.extend([None] * (numb_idx - len(filled) + 1))
        filled[numb_idx] = None
    cells = [
        _format_value(
            filled[column["source_index"]] if column["source_index"] < len(filled) else None,
            column["numeric"],
            column_rounding_digits(column, digits, emission_digits, specific_digits),
        )
        for column in visible_columns
    ]
    titles = cell_titles(filled, visible_columns)
    # Подпись итога всегда в столбце станции (даже когда в filled уже есть).
    for index, column in enumerate(visible_columns):
        if column.get("key") in ("NAME", "Stancia"):
            cells[index] = label
        elif column.get("key") == "numb1120":
            cells[index] = ""
        elif column.get("key") == "Year":
            cells[index] = str(year)
    cell_nonzero = [
        bool(column.get("diff_highlight")) and _display_is_nonzero(text)
        for column, text in zip(visible_columns, cells)
    ]
    cell_mismatch = station_generation_mismatch(
        cells,
        visible_columns,
        station_generation_check_pairs(visible_columns),
    )
    visible_filled = [
        filled[column["source_index"]] if column["source_index"] < len(filled) else None
        for column in visible_columns
    ]
    return {
        "key": key,
        "year": year,
        "cells": cells,
        "titles": titles,
        "cell_nonzero": cell_nonzero,
        "cell_mismatch": cell_mismatch,
        "is_empty": _is_empty_numeric_row(visible_filled, visible_columns),
        "is_group_first": is_first,
        "is_summary": True,
        "summary_label": label,
    }


def build_region_summary_block(
    groups: list[dict[str, Any]],
    label: str,
    *,
    source_columns: list[dict[str, Any]],
    visible_columns: list[dict[str, Any]],
    digits: int,
    key_prefix: str,
    emission_digits: int | None = None,
    specific_digits: int | None = None,
    norm_zv_parts: list[str] | tuple[str, ...] = ZV_NORM_DEFAULT_PARTS,
) -> dict[str, Any] | None:
    """Итоговые строки уровня (РЭС/ОЭС/ЕЭС) по годам."""
    by_year: dict[int, list[list[Any]]] = defaultdict(list)
    for group in groups:
        for year, filled in (group.get("filled_by_year") or {}).items():
            if year is None or not filled:
                continue
            by_year[int(year)].append(list(filled))
    if not by_year:
        return None
    summary_label = f"{label}, всего"
    rows: list[dict[str, Any]] = []
    for loop_i, year in enumerate(sorted(by_year)):
        summed = _sum_filled_rows(by_year[year], source_columns, norm_zv_parts)
        if summed is None:
            continue
        rows.append(
            _format_summary_row(
                label=summary_label,
                year=year,
                filled=summed,
                source_columns=source_columns,
                visible_columns=visible_columns,
                digits=digits,
                key=f"{key_prefix}-{year}",
                is_first=loop_i == 0,
                emission_digits=emission_digits,
                specific_digits=specific_digits,
            )
        )
    if not rows:
        return None
    return {
        "label": summary_label,
        "rowspan": len(rows),
        "rows": rows,
    }


def _index_groups_by_region(
    groups: list[dict[str, Any]],
) -> tuple[dict[tuple, list[dict[str, Any]]], dict[tuple, list[dict[str, Any]]], dict[tuple, list[dict[str, Any]]]]:
    by_res: dict[tuple, list[dict[str, Any]]] = defaultdict(list)
    by_ues: dict[tuple, list[dict[str, Any]]] = defaultdict(list)
    by_est: dict[tuple, list[dict[str, Any]]] = defaultdict(list)
    for group in groups:
        keys = _region_keys(group.get("meta") or {})
        by_res[keys["res"]].append(group)
        by_ues[keys["ues"]].append(group)
        by_est[keys["est"]].append(group)
    return by_res, by_ues, by_est


def attach_region_summaries(
    hierarchy: list[dict[str, Any]],
    all_groups: list[dict[str, Any]],
    *,
    source_columns: list[dict[str, Any]],
    visible_columns: list[dict[str, Any]],
    digits: int,
    emission_digits: int | None = None,
    specific_digits: int | None = None,
    norm_zv_parts: list[str] | tuple[str, ...] = ZV_NORM_DEFAULT_PARTS,
) -> list[dict[str, Any]]:
    """Добавляет итоги РЭС/ОЭС/ЕЭС, когда на странице последний блок региона."""
    by_res, by_ues, by_est = _index_groups_by_region(all_groups)
    for est_block in hierarchy:
        for ues_block in est_block.get("ues_list") or []:
            for res_block in ues_block.get("res_list") or []:
                page_groups = res_block.get("groups") or []
                if not page_groups:
                    continue
                res_key = _region_keys(page_groups[-1].get("meta") or {})["res"]
                full_res = by_res.get(res_key) or []
                if full_res and page_groups[-1].get("id") == full_res[-1].get("id"):
                    res_block["summary"] = build_region_summary_block(
                        full_res,
                        res_block.get("res_name") or UNSPECIFIED_LEVEL_NAME,
                        source_columns=source_columns,
                        visible_columns=visible_columns,
                        digits=digits,
                        key_prefix=f"sum-res-{res_key}",
                        emission_digits=emission_digits,
                        specific_digits=specific_digits,
                        norm_zv_parts=norm_zv_parts,
                    )
            ues_res_list = ues_block.get("res_list") or []
            if not ues_res_list:
                continue
            last_res_groups = ues_res_list[-1].get("groups") or []
            if not last_res_groups:
                continue
            ues_key = _region_keys(last_res_groups[-1].get("meta") or {})["ues"]
            full_ues = by_ues.get(ues_key) or []
            if full_ues and last_res_groups[-1].get("id") == full_ues[-1].get("id"):
                ues_block["summary"] = build_region_summary_block(
                    full_ues,
                    ues_block.get("ues_name") or UNSPECIFIED_LEVEL_NAME,
                    source_columns=source_columns,
                    visible_columns=visible_columns,
                    digits=digits,
                    key_prefix=f"sum-ues-{ues_key}",
                    emission_digits=emission_digits,
                    specific_digits=specific_digits,
                    norm_zv_parts=norm_zv_parts,
                )
        est_ues_list = est_block.get("ues_list") or []
        if not est_ues_list:
            continue
        last_ues_res = (est_ues_list[-1].get("res_list") or [])
        if not last_ues_res:
            continue
        last_est_groups = last_ues_res[-1].get("groups") or []
        if not last_est_groups:
            continue
        est_key = _region_keys(last_est_groups[-1].get("meta") or {})["est"]
        full_est = by_est.get(est_key) or []
        if full_est and last_est_groups[-1].get("id") == full_est[-1].get("id"):
            est_block["summary"] = build_region_summary_block(
                full_est,
                est_block.get("est_name") or UNSPECIFIED_LEVEL_NAME,
                source_columns=source_columns,
                visible_columns=visible_columns,
                digits=digits,
                key_prefix=f"sum-est-{est_key}",
                emission_digits=emission_digits,
                specific_digits=specific_digits,
                norm_zv_parts=norm_zv_parts,
            )
    return hierarchy


def _cell_by_key(filled: list[Any], columns: list[dict[str, Any]], keys: set[str]) -> str:
    index = _column_index(columns, keys)
    if index is None or index >= len(filled):
        return ""
    return _cell_text(filled[index])


def parse_page_arg(args: Any = None) -> int:
    values = _arg_values(args, "page")
    if not values:
        return 1
    try:
        page = int(values[0])
    except (TypeError, ValueError):
        return 1
    return page if page > 0 else 1


def parse_per_page_arg(args: Any = None, default: int = DEFAULT_PER_PAGE) -> int | str:
    values = _arg_values(args, "per_page")
    if not values:
        return default
    raw = str(values[0]).strip().lower()
    if raw == "all":
        return "all"
    try:
        number = int(raw)
    except (TypeError, ValueError):
        return default
    if number <= 0:
        return default
    return number


def _unspecified(name: Any) -> str:
    text = _cell_text(name)
    return text or UNSPECIFIED_LEVEL_NAME


def _level_id(value_id: Any, name: str) -> Any:
    if value_id is not None:
        return value_id
    if name and name != UNSPECIFIED_LEVEL_NAME:
        return name
    return None


def _group_meta(fuel: dict[str, Any] | None, filled: list[Any], columns: list[dict[str, Any]]) -> dict[str, Any]:
    excel_ues = _cell_by_key(filled, columns, {"nameoes"})
    excel_station = _cell_by_key(filled, columns, {"NAME", "Stancia"})
    excel_fo = _cell_by_key(filled, columns, {"Fed_okr"})
    fuel = fuel or {}
    est_name = _unspecified(fuel.get("est_name"))
    ues_name = _unspecified(fuel.get("ues") or excel_ues)
    res_name = _unspecified(fuel.get("res_name"))
    fo_name = _unspecified(fuel.get("federal_district") or excel_fo)
    sa_name = _unspecified(fuel.get("synchronous_area"))
    return {
        "est_id": _level_id(fuel.get("energy_system_type_id"), est_name),
        "est_name": est_name,
        "ues_id": _level_id(fuel.get("ues_id"), ues_name),
        "ues_name": ues_name,
        "res_id": _level_id(fuel.get("res_id"), res_name),
        "res_name": res_name,
        "fo_id": _level_id(fuel.get("federal_district_id"), fo_name),
        "fo_name": fo_name,
        "sa_id": _level_id(fuel.get("synchronous_area_id"), sa_name),
        "sa_name": sa_name,
        "station_name": _unspecified(fuel.get("station_name") or excel_station),
    }


def _group_sort_key(group: dict[str, Any]) -> tuple:
    meta = group["meta"]
    numb = group.get("numb")
    numb_key = numb if isinstance(numb, int) else 10**9
    return (
        (meta.get("est_name") or "").casefold(),
        (meta.get("ues_name") or "").casefold(),
        (meta.get("res_name") or "").casefold(),
        (meta.get("station_name") or "").casefold(),
        numb_key,
    )


def _group_station_name(group: dict[str, Any]) -> str:
    """Название станции группы – как в столбце «Электростанция»."""
    return _cell_text((group.get("meta") or {}).get("station_name"))


def _group_affiliation_sets(
    groups: list[dict[str, Any]],
    years: list[int],
) -> list[set[str]]:
    """«Принадлежность» каждой группы (по годам) – как в столбце таблицы."""
    numbs: set[int] = set()
    for group in groups:
        for numb in group.get("member_numbs") or [group.get("numb")]:
            if isinstance(numb, int):
                numbs.add(numb)
    mapping = load_station_affiliations_by_numb_year(numbs, years) if years else {}
    result: list[set[str]] = []
    for group in groups:
        values: set[str] = set()
        for numb in group.get("member_numbs") or [group.get("numb")]:
            if not isinstance(numb, int):
                continue
            for year in years:
                value = mapping.get((numb, year))
                if value:
                    values.add(value)
        result.append(values or {STATION_AFFILIATION_UNKNOWN})
    return result


def apply_header_filters(
    groups: list[dict[str, Any]],
    visible_columns: list[dict[str, Any]],
    filters: dict[str, Any],
    years: list[int],
) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    """Отбирает группы по фильтрам в шапке и возвращает варианты для списков."""
    station_header_filter = str(filters.get("station_header_filter") or "").strip()
    affiliation_filter = str(filters.get("affiliation_filter") or "").strip()
    station_choices = sorted(
        {name for group in groups if (name := _group_station_name(group))},
        key=str.casefold,
    )
    has_affiliation_column = (
        _column_index(visible_columns, {STATION_AFFILIATION_COLUMN_KEY}) is not None
    )
    affiliation_sets = (
        _group_affiliation_sets(groups, years) if has_affiliation_column else []
    )
    # «–» – не принадлежность, а отсутствие данных ГИС ТЭК, поэтому в списке фильтра
    # его нет (как на /ecology/gis_tek/pg/ и /ecology/gis_tek/pollutants/).
    affiliation_choices = sorted(
        {
            value
            for values in affiliation_sets
            for value in values
            if value != STATION_AFFILIATION_UNKNOWN
        },
        key=str.casefold,
    )
    if station_header_filter:
        groups = [
            group
            for group in groups
            if _group_station_name(group) == station_header_filter
        ]
    if affiliation_filter and affiliation_sets:
        groups = [
            group
            for group, values in zip(groups, affiliation_sets)
            if affiliation_filter in values
        ]
    return groups, station_choices, affiliation_choices


def _paginate_groups(
    groups: list[dict[str, Any]],
    page: int,
    per_page: int | str,
) -> tuple[list[dict[str, Any]], int, int, int, int | str]:
    total = len(groups)
    if per_page == "all":
        return groups, 1, 1, total, "all"
    size = int(per_page) if int(per_page) > 0 else DEFAULT_PER_PAGE
    total_pages = max(1, math.ceil(total / size) if total else 1)
    current = min(max(1, page), total_pages)
    start = (current - 1) * size
    return groups[start : start + size], current, total_pages, total, size


def _build_hierarchy(groups: list[dict[str, Any]]) -> list[dict[str, Any]]:
    est_map: dict[tuple[Any, str], dict[str, Any]] = {}
    est_order: list[tuple[Any, str]] = []
    for group in groups:
        meta = group["meta"]
        est_key = (meta.get("est_id"), meta.get("est_name") or UNSPECIFIED_LEVEL_NAME)
        ues_key = (meta.get("ues_id"), meta.get("ues_name") or UNSPECIFIED_LEVEL_NAME)
        res_key = (meta.get("res_id"), meta.get("res_name") or UNSPECIFIED_LEVEL_NAME)
        if est_key not in est_map:
            est_map[est_key] = {
                "est_id": meta.get("est_id"),
                "est_name": est_key[1],
                "ues_map": {},
                "ues_order": [],
            }
            est_order.append(est_key)
        est_block = est_map[est_key]
        if ues_key not in est_block["ues_map"]:
            est_block["ues_map"][ues_key] = {
                "ues_id": meta.get("ues_id"),
                "ues_name": ues_key[1],
                "res_map": {},
                "res_order": [],
            }
            est_block["ues_order"].append(ues_key)
        ues_block = est_block["ues_map"][ues_key]
        if res_key not in ues_block["res_map"]:
            ues_block["res_map"][res_key] = {
                "res_id": meta.get("res_id"),
                "res_name": res_key[1],
                "groups": [],
            }
            ues_block["res_order"].append(res_key)
        ues_block["res_map"][res_key]["groups"].append(group)

    hierarchy: list[dict[str, Any]] = []
    for est_key in est_order:
        est_block = est_map[est_key]
        ues_list = []
        for ues_key in est_block["ues_order"]:
            ues_block = est_block["ues_map"][ues_key]
            res_list = [ues_block["res_map"][res_key] for res_key in ues_block["res_order"]]
            ues_list.append(
                {
                    "ues_id": ues_block["ues_id"],
                    "ues_name": ues_block["ues_name"],
                    "res_list": res_list,
                }
            )
        hierarchy.append(
            {
                "est_id": est_block["est_id"],
                "est_name": est_block["est_name"],
                "ues_list": ues_list,
            }
        )
    return hierarchy


def build_comparison_table_context(
    sheet_slug: str,
    start_year: Any = None,
    end_year: Any = None,
    rounding_digits: Any = None,
    emission_rounding_digits: Any = None,
    specific_rounding_digits: Any = None,
    view_year: Any = None,
    source_path: str | Path | None = None,
    request_args: Any = None,
) -> dict[str, Any] | None:
    sheet = get_comparison_sheet(sheet_slug or DEFAULT_SHEET_SLUG)
    if sheet is None:
        return None
    if sheet["slug"] in PG_ANALYSIS_SHEET_SLUGS:
        parsed = load_comparison_sheet_layout(sheet["slug"], source_path)
    else:
        parsed_all = load_comparison_sheets(source_path)
        parsed = parsed_all.get(sheet["slug"])
    if parsed is None:
        return None
    if sheet["slug"] in BEZ_KOT_SHEET_SLUGS:
        # Обе вкладки используют шапку «С котельными», но свой состав групп и параметры.
        layout = load_comparison_sheet_layout(f"{sheet['group']}-vse", source_path)
        if layout is None:
            return None
        parsed = {
            **parsed,
            "columns": deepcopy(layout["columns"]),
            "header_rows": deepcopy(layout["header_rows"]),
            "rows": [],
        }
    years_all = list(parsed["years"])
    planning_start = None
    planning_end = None
    if sheet["slug"] in ANALYSIS_SHEET_SLUGS:
        try:
            from app.common.services.get_services.years.years_get_services import (
                get_filter_end_year,
                get_filter_start_year,
                get_year_numbers_sorted_for_current_db_version,
            )

            catalog = get_year_numbers_sorted_for_current_db_version()
            if catalog:
                years_all = list(catalog)
            planning_start = get_filter_start_year()
            planning_end = get_filter_end_year()
        except Exception:
            feature_years = sorted(
                int(year)
                for year in (_year_features() or {})
                if str(year).isdigit() or isinstance(year, int)
            )
            if feature_years:
                years_all = feature_years
    if not years_all:
        years_all = []
    default_start = (
        planning_start if planning_start is not None else (years_all[0] if years_all else None)
    )
    default_end = (
        planning_end if planning_end is not None else (years_all[-1] if years_all else None)
    )
    if years_all and default_start is not None:
        default_start = min(max(int(default_start), years_all[0]), years_all[-1])
    if years_all and default_end is not None:
        default_end = min(max(int(default_end), years_all[0]), years_all[-1])

    def _clip_year(raw: Any, fallback: int | None) -> int | None:
        if fallback is None:
            return None
        try:
            year = int(raw)
        except (TypeError, ValueError):
            return fallback
        return min(max(year, years_all[0]), years_all[-1]) if years_all else fallback

    start = _clip_year(start_year, default_start) if start_year not in (None, "") else default_start
    end = _clip_year(end_year, default_end) if end_year not in (None, "") else default_end
    if start is not None and end is not None and start > end:
        start, end = end, start
    range_years = [year for year in years_all if start is not None and end is not None and start <= year <= end]
    digits = resolve_rounding_digits(rounding_digits)
    emission_digits = (
        resolve_rounding_digits(emission_rounding_digits)
        if emission_rounding_digits not in (None, "")
        else digits
    )
    specific_digits = (
        resolve_rounding_digits(specific_rounding_digits)
        if specific_rounding_digits not in (None, "")
        else digits
    )
    source_columns = parsed["columns"]
    if sheet["slug"] in ZV_ANALYSIS_SHEET_SLUGS:
        ensure_zv_electricity_specific_column(
            source_columns,
            parsed["header_rows"],
            parsed.get("rows"),
        )
    if sheet["slug"] in STATION_GENERATION_SHEET_SLUGS:
        ensure_station_generation_columns(
            source_columns,
            parsed["header_rows"],
            parsed.get("rows"),
        )
    if sheet["slug"] in STATION_AFFILIATION_SHEET_SLUGS:
        ensure_station_affiliation_column(
            source_columns,
            parsed["header_rows"],
            parsed.get("rows"),
        )
    show_reported = parse_show_reported_arg(request_args)
    has_reported_columns = any(
        str(column.get("key") or "") in REPORTED_TABLE_COLUMN_KEYS
        for column in source_columns
    )
    extra_hidden = frozenset() if show_reported else REPORTED_TABLE_COLUMN_KEYS
    visible_columns, header_rows, sticky_col_count = prepare_visible_table(
        source_columns,
        parsed["header_rows"],
        extra_hidden_keys=extra_hidden,
    )
    apply_fuel_tep_headers(visible_columns, header_rows)
    apply_fuel_tce_headers(visible_columns, header_rows)
    if sheet["slug"] in PG_ANALYSIS_SHEET_SLUGS:
        apply_pg_fuel_group_headers(visible_columns, header_rows)
        apply_pg_ecology_headers(visible_columns, header_rows)
        for column in visible_columns:
            if column.get("key") in PG_ECOLOGY_EDITABLE_KEYS:
                column["editable"] = True
    if sheet["slug"] in ZV_ANALYSIS_SHEET_SLUGS:
        apply_zv_display_column_order(visible_columns, header_rows)
        apply_zv_leaf_group_headers(visible_columns, header_rows)
        apply_zv_analysis_banner_headers(visible_columns, header_rows)
        for column in visible_columns:
            if column.get("key") in ZV_ECOLOGY_EDITABLE_KEYS:
                column["editable"] = True
    norm_zv_selected = parse_norm_zv_parts(request_args, ZV_NORM_DEFAULT_PARTS)
    _attach_column_formula_tooltips(header_rows, _column_formula_texts())
    if sheet["slug"] in ZV_ANALYSIS_SHEET_SLUGS:
        formula = " + ".join(label for attr, label in norm_zv_options() if attr in norm_zv_selected)
        for row in header_rows:
            for cell in row:
                if cell.get("column_key") == "ZV_normir":
                    cell["formula_tooltip"] = formula or "Слагаемые не выбраны"
    if sheet["slug"] in PG_ANALYSIS_SHEET_SLUGS or sheet["slug"] in ZV_ANALYSIS_SHEET_SLUGS:
        drop_key_header_row(header_rows)
    apply_full_height_headers(
        visible_columns,
        header_rows,
        STATION_GENERATION_COLUMN_LABELS,
        highlight_flag="station_generation_highlight",
    )
    apply_full_height_headers(
        visible_columns,
        header_rows,
        {STATION_AFFILIATION_COLUMN_KEY: STATION_AFFILIATION_LABEL},
    )
    station_check_pairs = station_generation_check_pairs(visible_columns)
    territory_filters = parse_territory_filters(request_args)
    if sheet["slug"] in ZV_ANALYSIS_SHEET_SLUGS:
        territory_filters["norm_zv"] = norm_zv_selected
    numb1120_filter = _as_int_ids(_arg_values(request_args, "numb1120_filter"))
    territory_filters["numb1120_filter"] = numb1120_filter
    page = parse_page_arg(request_args)
    per_page = parse_per_page_arg(request_args)
    show_totals = parse_show_totals_arg(request_args)
    show_empty_rows = parse_show_empty_rows_arg(request_args)
    territory_filters["show_totals"] = show_totals
    territory_filters["show_empty_rows"] = show_empty_rows
    territory_filters["show_reported"] = show_reported

    if sheet["slug"] in ANALYSIS_SHEET_SLUGS and start is not None and end is not None:
        from app.ecology.services.gis_tes_pg_from_fuel_services import (
            collect_pg_groups_from_fuel,
            fill_pg_group_year_rows,
        )

        groups, numb1120_filter_choices = collect_pg_groups_from_fuel(
            sheet["slug"],
            start_year=start,
            end_year=end,
            territory_filters=territory_filters,
            source_columns=source_columns,
            visible_columns=visible_columns,
            digits=digits,
            hydrate=False,
            emission_digits=emission_digits,
            specific_digits=specific_digits,
        )
        if numb1120_filter:
            wanted_numbs = set(numb1120_filter)
            # У сводной строки составной станции в numb – код головы, а фильтр может
            # содержать код любой её подгруппы.
            groups = [
                group
                for group in groups
                if group.get("numb") in wanted_numbs
                or wanted_numbs.intersection(group.get("member_numbs") or ())
            ]
            numb1120_filter_choices = sorted(
                {n for n in numb1120_filter_choices if n in wanted_numbs}
                or {
                    group["numb"]
                    for group in groups
                    if isinstance(group.get("numb"), int)
                }
            )
        elif not numb1120_filter_choices:
            numb1120_filter_choices = sorted(
                {group["numb"] for group in groups if isinstance(group.get("numb"), int)}
            )
    else:
        visible_raw: list[tuple[int, list[Any]]] = []
        numbs: set[int] = set()
        for index, raw in enumerate(parsed["rows"]):
            year = _row_year(raw)
            if year is not None and range_years and year not in range_years:
                continue
            if year is None and range_years:
                continue
            visible_raw.append((index, raw))
            numb = _row_numb(raw, source_columns)
            if numb is not None:
                numbs.add(numb)
        overlay = load_fuel_indicators_by_numb_year(numbs, range_years)
        pg_overlay = (
            load_pg_ecology_overlay_by_numb_year(sheet["slug"], numbs, range_years)
            if sheet["slug"] in PG_ANALYSIS_SHEET_SLUGS
            else {}
        )
        station_generation_overlay = (
            load_station_generation_overlay(numbs, range_years)
            if sheet["slug"] in STATION_GENERATION_SHEET_SLUGS
            else {}
        )
        station_links_by_numb = (
            station_generation_overlay.get("station_links_by_numb") or {}
        )
        station_affiliations = (
            load_station_affiliations_by_numb_year(
                numbs, sorted({year for _numb, year in overlay})
            )
            if sheet["slug"] in STATION_AFFILIATION_SHEET_SLUGS
            else {}
        )
        grouped: dict[Any, list[dict[str, Any]]] = defaultdict(list)
        for index, raw in visible_raw:
            year = _row_year(raw)
            filled = apply_fuel_indicators(raw, source_columns, overlay)
            if pg_overlay:
                filled = apply_pg_ecology_overlay(filled, source_columns, pg_overlay)
            if station_generation_overlay:
                filled = apply_station_generation_overlay(
                    filled, source_columns, station_generation_overlay
                )
            if station_affiliations:
                filled = apply_station_affiliation_cells(
                    filled, source_columns, station_affiliations
                )
            numb = _row_numb(raw, source_columns)
            fuel = overlay.get((numb, year)) if numb is not None and year is not None else None
            if not row_matches_territory_filters(filled, source_columns, fuel, territory_filters):
                continue
            filled = _compute_derived_columns(filled, source_columns, norm_zv_selected)
            cells = [
                _format_value(
                    filled[column["source_index"]] if column["source_index"] < len(filled) else None,
                    column["numeric"],
                    column_rounding_digits(
                        column, digits, emission_digits, specific_digits
                    ),
                )
                for column in visible_columns
            ]
            titles = cell_titles(filled, visible_columns)
            cell_nonzero = [
                bool(column.get("diff_highlight")) and _display_is_nonzero(text)
                for column, text in zip(visible_columns, cells)
            ]
            cell_mismatch = station_generation_mismatch(
                cells, visible_columns, station_check_pairs
            )
            visible_filled = [
                filled[column["source_index"]] if column["source_index"] < len(filled) else None
                for column in visible_columns
            ]
            group_id = numb if numb is not None else f"row-{index}"
            grouped[group_id].append(
                {
                    "key": f"{sheet['slug']}-{index}-{year or ''}",
                    "year": year,
                    "cells": cells,
                    "titles": titles,
                    "cell_nonzero": cell_nonzero,
                    "cell_mismatch": cell_mismatch,
                    "is_empty": _is_empty_numeric_row(visible_filled, visible_columns),
                    "fuel": fuel,
                    "filled": filled,
                }
            )
        groups = []
        for group_id, items in grouped.items():
            items.sort(key=lambda item: (item["year"] is None, item["year"] or 0))
            meta_src = next(
                (item for item in items if (item.get("fuel") or {}).get("ues_id") is not None),
                items[0],
            )
            meta = _group_meta(meta_src.get("fuel"), meta_src.get("filled") or [], source_columns)
            groups.append(
                {
                    "id": group_id,
                    "numb": group_id if isinstance(group_id, int) else None,
                    "meta": meta,
                    "station_links": station_links_by_numb.get(group_id) or [],
                    "rowspan": len(items),
                    "filled_by_year": {
                        item["year"]: item["filled"]
                        for item in items
                        if item.get("year") is not None and item.get("filled") is not None
                    },
                    "rows": [
                        {
                            "key": item["key"],
                            "year": item["year"],
                            "cells": item["cells"],
                            "titles": item["titles"],
                            "cell_nonzero": item["cell_nonzero"],
                            "cell_mismatch": item["cell_mismatch"],
                            "is_empty": item["is_empty"],
                            "is_group_first": loop_i == 0,
                        }
                        for loop_i, item in enumerate(items)
                    ],
                }
            )
        groups.sort(key=_group_sort_key)
        if sheet["slug"] in ZV_ANALYSIS_SHEET_SLUGS:
            groups = _drop_boiler_groups(groups, sheet["slug"])
        numb1120_filter_choices = sorted(
            {group["numb"] for group in groups if isinstance(group.get("numb"), int)}
        )
        if numb1120_filter:
            wanted_numbs = set(numb1120_filter)
            groups = [group for group in groups if group.get("numb") in wanted_numbs]

    # Фильтры в шапке: «Электростанция» и «Принадлежность».
    station_header_filter = str(
        territory_filters.get("station_header_filter") or ""
    ).strip()
    affiliation_filter = str(territory_filters.get("affiliation_filter") or "").strip()
    groups, station_filter_choices, affiliation_filter_choices = apply_header_filters(
        groups, visible_columns, territory_filters, list(range_years)
    )

    page_groups, current_page, total_pages, total_count, per_page_value = _paginate_groups(
        groups, page, per_page
    )
    if sheet["slug"] in ANALYSIS_SHEET_SLUGS and start is not None and end is not None:
        from app.ecology.services.gis_tes_pg_from_fuel_services import (
            fill_pg_group_year_rows,
        )

        hydrate_groups = groups if (show_totals or per_page_value == "all") else page_groups
        fill_pg_group_year_rows(
            hydrate_groups,
            sheet_slug=sheet["slug"],
            start_year=start,
            end_year=end,
            territory_filters=territory_filters,
            source_columns=source_columns,
            visible_columns=visible_columns,
            digits=digits,
            emission_digits=emission_digits,
            specific_digits=specific_digits,
        )
    hierarchy = _build_hierarchy(page_groups)
    if show_totals:
        attach_region_summaries(
            hierarchy,
            groups,
            source_columns=source_columns,
            visible_columns=visible_columns,
            digits=digits,
            emission_digits=emission_digits,
            specific_digits=specific_digits,
            norm_zv_parts=norm_zv_selected,
        )
    keep_query = _keep_query_args(
        start,
        end,
        digits,
        None,
        territory_filters,
        include_view=False,
        per_page=per_page_value,
        emission_digits=emission_digits,
        specific_digits=specific_digits,
    )
    context = {
        "page_title": sheet["page_title"],
        "sheet": sheet,
        "comparison_sheets": list_comparison_sheets_for_group(sheet["group"]),
        "header_rows": header_rows,
        "columns": visible_columns,
        "hierarchy": hierarchy,
        "years": years_all,
        "range_years": range_years,
        "filter_year_list": years_all,
        "start_year": start,
        "end_year": end,
        "rounding_digits": digits,
        "emission_rounding_digits": emission_digits,
        "specific_rounding_digits": specific_digits,
        "year_features": _year_features(),
        "sticky_col_count": sticky_col_count,
        "keep_query": keep_query,
        "keep_query_with_view": keep_query,
        "year_reset_query": _keep_query_args(
            start,
            end,
            digits,
            None,
            territory_filters,
            include_years=False,
            per_page=per_page_value,
            emission_digits=emission_digits,
            specific_digits=specific_digits,
        ),
        "current_page": current_page,
        "total_pages": total_pages,
        "total_count": total_count,
        "per_page": per_page_value,
        "show_all": per_page_value == "all",
        "default_per_page": str(DEFAULT_PER_PAGE),
        "rowspan_column_keys": ROWSPAN_COLUMN_KEYS,
        "numb1120_filter": numb1120_filter,
        "numb1120_filter_choices": numb1120_filter_choices,
        "station_header_filter": station_header_filter,
        "station_filter_choices": station_filter_choices,
        "affiliation_filter": affiliation_filter,
        "affiliation_filter_choices": affiliation_filter_choices,
        "show_totals": show_totals,
        "show_empty_rows": show_empty_rows,
        "show_reported": show_reported,
        "has_reported_columns": has_reported_columns,
        "station_check_pairs": station_check_pairs,
        "norm_zv_options": norm_zv_options() if sheet["slug"] in ZV_ANALYSIS_SHEET_SLUGS else [],
        "norm_zv_selected": norm_zv_selected,
        "norm_zv_default": list(ZV_NORM_DEFAULT_PARTS),
        "norm_zv_label": KEY_LABELS["ZV_normir"],
        "has_active_filters": bool(
            numb1120_filter
            or station_header_filter
            or affiliation_filter
            or (territory_filters.get("station_name_filter") or "").strip()
            or (territory_filters.get("gen_company_filter") or "").strip()
            or str(territory_filters.get("ved_filter") or "").strip()
            or any(territory_filters.get(key) for key in TERRITORY_FILTER_KEYS[:5])
        ),
    }
    context.update(load_territorial_filter_context(territory_filters))
    return context


def clear_comparison_cache() -> None:
    _parse_workbook_cached.cache_clear()
    _parse_sheet_layout_cached.cache_clear()
